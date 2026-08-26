"""Read-only adapter for the effective Zhaoshang result workbook format."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import AbstractSet, Any

from openpyxl import load_workbook

from app.services.saas_audit_core import normalize_cell_str, parse_decimal

# Single source of truth for the Chinese classification literals. Both result
# readers must emit identical semantics; the identity reader once carried
# double-encoded mojibake copies of these strings (P1-1).
PRICE_SOURCE_ZHAOSHANG = "招商局"
PRICE_SOURCE_PRICE_NOT_VISIBLE = "招商局-价格不可见"
RESULT_STATUS_SUCCESS = "成功"
RESULT_STATUS_UNKNOWN = "未知"


def normalize_result_url(value: Any) -> str:
    """The legacy result key is the exact URL after surrounding whitespace only."""
    return normalize_cell_str(value)


def read_template_code_url_map(template_path: Path | str | None) -> dict[str, str]:
    if not template_path:
        return {}
    path = Path(template_path)
    if not path.exists():
        return {}
    workbook = load_workbook(filename=str(path), data_only=True, read_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            return {}
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        except StopIteration:
            return {}

        def find(*names: str) -> int:
            return next((headers.index(name) for name in names if name in headers), -1)

        code_index = find("商品编码", "skuId", "sku_id", "code")
        url_index = find("url", "URL", "链接", "商品URL", "商品链接")
        if code_index < 0 or url_index < 0:
            return {}
        mapping: dict[str, str] = {}
        for row in iterator:
            if row is None or code_index >= len(row) or url_index >= len(row):
                continue
            code = normalize_cell_str(row[code_index])
            url = normalize_result_url(row[url_index])
            if code and url and code not in mapping:
                mapping[code] = url
        return mapping
    finally:
        workbook.close()


def _template_code_urls(template_path: Path | str | None) -> dict[str, set[str]]:
    """Return every template URL for a SKU; ambiguity is preserved, not hidden."""
    if not template_path:
        return {}
    path = Path(template_path)
    if not path.exists():
        return {}
    workbook = load_workbook(filename=str(path), data_only=True, read_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            return {}
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        except StopIteration:
            return {}

        def find(*names: str) -> int:
            return next((headers.index(name) for name in names if name in headers), -1)

        code_index = find("\u5546\u54c1\u7f16\u7801", "skuId", "sku_id", "code")
        url_index = find("url", "URL", "\u94fe\u63a5", "\u5546\u54c1URL", "\u5546\u54c1\u94fe\u63a5")
        if code_index < 0 or url_index < 0:
            return {}
        result: dict[str, set[str]] = {}
        for row in iterator:
            if row is None or code_index >= len(row) or url_index >= len(row):
                continue
            code = normalize_cell_str(row[code_index])
            url = normalize_result_url(row[url_index])
            if code and url:
                result.setdefault(code, set()).add(url)
        return result
    finally:
        workbook.close()


def read_zhaoshang_result_xlsx_by_identity(
    path: Path | str,
    template_path: Path | str | None = None,
    *,
    code_index: int | None = None,
    include_excel_rows: AbstractSet[int] | None = None,
) -> dict[tuple[str, str], dict[str, Any]]:
    """Read prices keyed by the verified ``(skuId, exact URL)`` crawl identity.

    URL-only exports are accepted only when the template proves the URL maps
    to one SKU. SKU-only exports are accepted only when that SKU has one URL.
    This avoids silently merging distinct products sharing a product page URL.
    """
    result_path = Path(path)
    if not result_path.exists():
        return {}
    code_to_urls = _template_code_urls(template_path)
    url_to_codes: dict[str, set[str]] = {}
    for code, urls in code_to_urls.items():
        for url in urls:
            url_to_codes.setdefault(url, set()).add(code)
    workbook = load_workbook(filename=str(result_path), data_only=True, read_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            return {}
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        except StopIteration:
            return {}

        def find(*names: str) -> int:
            return next((headers.index(name) for name in names if name in headers), -1)

        url_index = find("\u539f\u59cb\u94fe\u63a5", "\u5546\u54c1\u8be6\u60c5\u5730\u5740", "url", "URL", "\u5546\u54c1\u94fe\u63a5", "\u94fe\u63a5", "\u5546\u54c1URL")
        selected_code_index = (
            code_index
            if code_index is not None
            else find("\u5546\u54c1\u7f16\u7801", "\u5546\u54c1\u6700\u5c0f\u54c1\u7c7b\u5355\u5143id", "skuId", "sku_id", "code")
        )
        price_index = find("\u5546\u54c1\u4ef7\u683c", "\u4ef7\u683c", "\u91c7\u96c6\u4ef7\u683c", "page_price")
        status_index = find("status", "crawl_status")
        reason_index = find("reason", "failure_reason")
        if url_index < 0 and (selected_code_index < 0 or not code_to_urls):
            return {}

        results: dict[tuple[str, str], dict[str, Any]] = {}
        for excel_row, row in enumerate(iterator, start=2):
            if include_excel_rows is not None and excel_row not in include_excel_rows:
                continue
            if row is None:
                continue
            url = normalize_result_url(row[url_index]) if 0 <= url_index < len(row) else ""
            code = (
                normalize_cell_str(row[selected_code_index])
                if 0 <= selected_code_index < len(row)
                else ""
            )
            if not url:
                urls = code_to_urls.get(code, set())
                if len(urls) != 1:
                    raise ValueError("result workbook SKU fallback is ambiguous")
                url = next(iter(urls))
            if not code:
                codes = url_to_codes.get(url, set())
                if len(codes) != 1:
                    raise ValueError("result workbook URL fallback is ambiguous")
                code = next(iter(codes))
            identity = (code, url)
            price: Decimal | None = parse_decimal(row[price_index]) if 0 <= price_index < len(row) else None
            status = str(row[status_index]).strip() if 0 <= status_index < len(row) and row[status_index] is not None else ""
            reason = str(row[reason_index]).strip() if 0 <= reason_index < len(row) and row[reason_index] is not None else ""
            if price is not None and price > 0:
                page_status, source = "OK", PRICE_SOURCE_ZHAOSHANG
            elif price is not None and price == 0:
                page_status, source = "PRICE_IS_ZERO", PRICE_SOURCE_ZHAOSHANG
            elif RESULT_STATUS_SUCCESS in status:
                page_status, source = "PRICE_MISSING", PRICE_SOURCE_PRICE_NOT_VISIBLE
            else:
                page_status, source = "FAILED", PRICE_SOURCE_ZHAOSHANG + "-" + (reason or status or RESULT_STATUS_UNKNOWN)
            if identity in results:
                raise ValueError("result workbook contains duplicate crawl identity")
            results[identity] = {
                "product_code": code,
                "page_price_value": price,
                "page_price_status": page_status,
                "page_price_source": source,
            }
        return results
    finally:
        workbook.close()


def read_zhaoshang_result_xlsx(
    path: Path | str,
    template_path: Path | str | None = None,
    *,
    code_index: int | None = None,
    include_excel_rows: AbstractSet[int] | None = None,
) -> dict[str, dict[str, Any]]:
    """Read the currently effective result mapping, including template code fallback."""
    result_path = Path(path)
    if not result_path.exists():
        return {}
    code_to_url = read_template_code_url_map(template_path)
    workbook = load_workbook(filename=str(result_path), data_only=True, read_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            return {}
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        except StopIteration:
            return {}

        def find(*names: str) -> int:
            return next((headers.index(name) for name in names if name in headers), -1)

        url_index = find("原始链接", "商品详情地址", "url", "URL", "商品链接", "链接", "商品URL")
        selected_code_index = (
            code_index
            if code_index is not None
            else find("商品编码", "商品最小品类单元id", "skuId", "sku_id", "code")
        )
        price_index = find("商品价格", "价格", "采集价格", "page_price")
        status_index = find("status", "crawl_status")
        reason_index = find("reason", "failure_reason")
        if url_index < 0 and (selected_code_index < 0 or not code_to_url):
            return {}

        results: dict[str, dict[str, Any]] = {}
        for excel_row, row in enumerate(iterator, start=2):
            if include_excel_rows is not None and excel_row not in include_excel_rows:
                continue
            if row is None:
                continue
            url = normalize_result_url(row[url_index]) if 0 <= url_index < len(row) else ""
            code = (
                normalize_cell_str(row[selected_code_index])
                if 0 <= selected_code_index < len(row)
                else ""
            )
            if not url:
                url = code_to_url.get(code, "")
            if not url:
                continue
            price: Decimal | None = parse_decimal(row[price_index]) if 0 <= price_index < len(row) else None
            status = str(row[status_index]).strip() if 0 <= status_index < len(row) and row[status_index] is not None else ""
            reason = str(row[reason_index]).strip() if 0 <= reason_index < len(row) and row[reason_index] is not None else ""
            if price is not None and price > 0:
                page_status, source = "OK", PRICE_SOURCE_ZHAOSHANG
            elif price is not None and price == 0:
                page_status, source = "PRICE_IS_ZERO", PRICE_SOURCE_ZHAOSHANG
            elif RESULT_STATUS_SUCCESS in status:
                page_status, source = "PRICE_MISSING", PRICE_SOURCE_PRICE_NOT_VISIBLE
            else:
                page_status, source = "FAILED", PRICE_SOURCE_ZHAOSHANG + "-" + (reason or status or RESULT_STATUS_UNKNOWN)
            results[url] = {"product_code": code, "page_price_value": price, "page_price_status": page_status, "page_price_source": source}
        return results
    finally:
        workbook.close()
