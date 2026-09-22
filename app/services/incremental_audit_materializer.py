"""Offline, atomic workbook publication for a frozen incremental audit snapshot."""

from __future__ import annotations

import dataclasses
import hashlib
import json
import os
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from enum import Enum
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

from openpyxl import load_workbook


TENANT_ID = "188605"
COMMODITY_POOL_ID = "222"
MANIFEST_SCHEMA_VERSION = 1
EMBEDDED_MANIFEST_PREFIX = "incremental_audit_input_manifest_sha256="
WRITER_CONTRACT_VERSION = "legacy-workbook-v1"


class MaterializationValidationError(ValueError):
    pass


class ArtifactConflictError(RuntimeError):
    pass


@dataclass(frozen=True)
class FrozenAuditArtifact:
    business_date: str
    tenant_id: str
    commodity_pool_id: str
    output_identity: str
    frozen_snapshot_id: str
    frozen_snapshot_hash: str
    frozen_snapshot_evidence: Mapping[str, Any]
    headers: Sequence[str]
    frozen_rows: Sequence[Mapping[str, Any]]
    enriched_rows: Sequence[Mapping[str, Any]]
    audit_response: Any


@dataclass(frozen=True)
class MaterializationResult:
    output_path: Path
    manifest_path: Path
    input_manifest_sha256: str
    manifest_sha256: str
    workbook_sha256: str
    row_count: int
    reused: bool
    recovered_sidecar: bool = False


@dataclass(frozen=True)
class WorkbookProjection:
    keys: tuple[tuple[str, str, str], ...]
    record_results_sha256: str
    summary_business_sha256: str


def _text(value: object) -> str:
    return str(value or "").strip()


def _jsonable(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Enum):
        return _jsonable(value.value)
    if isinstance(value, Mapping):
        return {str(key): _jsonable(item) for key, item in sorted(value.items(), key=lambda item: str(item[0]))}
    if isinstance(value, (list, tuple)):
        return [_jsonable(item) for item in value]
    if dataclasses.is_dataclass(value):
        return _jsonable(dataclasses.asdict(value))
    if hasattr(value, "model_dump"):
        return _jsonable(value.model_dump(mode="json"))
    if hasattr(value, "dict"):
        return _jsonable(value.dict())
    raise MaterializationValidationError(f"cannot serialize artifact value of type {type(value).__name__}")


def _canonical_hash(value: Any) -> str:
    encoded = json.dumps(_jsonable(value), ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


class LegacyWorkbookMaterializer:
    def __init__(
        self,
        *,
        writer: Callable[[Path, list[str], list[dict[str, Any]], list[dict[str, Any]], Any], None] | None = None,
        after_workbook_replace: Callable[[], None] | None = None,
    ) -> None:
        self._writer = writer
        self._after_workbook_replace = after_workbook_replace

    @staticmethod
    def _audit_records(audit_response: Any) -> list[Any]:
        if isinstance(audit_response, Mapping):
            records = audit_response.get("record_results")
        else:
            records = getattr(audit_response, "record_results", None)
        if records is None:
            raise MaterializationValidationError("audit response must contain record_results")
        if not isinstance(records, (list, tuple)):
            raise MaterializationValidationError("audit response record_results must be a sequence")
        return list(records)

    @staticmethod
    def _row_key(row: Mapping[str, Any]) -> tuple[str, str, str]:
        return (_text(row.get("vCommodityId")), _text(row.get("skuId")), _text(row.get("extLink")))

    def _validate_artifact(self, artifact: FrozenAuditArtifact, output_path: Path) -> None:
        if _text(artifact.tenant_id) != TENANT_ID or _text(artifact.commodity_pool_id) != COMMODITY_POOL_ID:
            raise MaterializationValidationError("artifact scope must remain tenant 188605 / pool 222")
        try:
            date.fromisoformat(_text(artifact.business_date))
        except ValueError as exc:
            raise MaterializationValidationError("business_date must be ISO YYYY-MM-DD") from exc
        if not _text(artifact.output_identity):
            raise MaterializationValidationError("output_identity must be non-empty and stable")
        if not _text(artifact.frozen_snapshot_id) or not _text(artifact.frozen_snapshot_hash):
            raise MaterializationValidationError("frozen snapshot identity and hash are required")
        if not output_path.name or output_path.suffix.lower() != ".xlsx":
            raise MaterializationValidationError("output path must be a non-empty .xlsx path")
        records = self._audit_records(artifact.audit_response)
        length = len(artifact.frozen_rows)
        if not length == len(artifact.enriched_rows) == len(records):
            raise MaterializationValidationError("frozen rows, enriched rows, and audit records must have equal length")
        if any(not isinstance(row, Mapping) for row in artifact.frozen_rows):
            raise MaterializationValidationError("frozen rows must be dictionaries")
        if not artifact.headers:
            raise MaterializationValidationError("legacy headers are required")

    def _input_manifest_hash(self, artifact: FrozenAuditArtifact) -> str:
        ordered_rows = []
        for row in artifact.frozen_rows:
            ordered_rows.append(
                {
                    "key": self._row_key(row),
                    "fingerprint": _text(row.get("source_fingerprint") or row.get("__source_fingerprint"))
                    or _canonical_hash(row),
                }
            )
        return _canonical_hash(
            {
                "schema_version": MANIFEST_SCHEMA_VERSION,
                "writer_contract_version": WRITER_CONTRACT_VERSION,
                "business_date": _text(artifact.business_date),
                "scope": {"tenant_id": TENANT_ID, "commodity_pool_id": COMMODITY_POOL_ID},
                "output_identity": _text(artifact.output_identity),
                "frozen_snapshot": {
                    "id": _text(artifact.frozen_snapshot_id),
                    "hash": _text(artifact.frozen_snapshot_hash),
                    "evidence": artifact.frozen_snapshot_evidence,
                },
                "ordered_rows": ordered_rows,
                "headers": list(artifact.headers),
                "enriched_rows": artifact.enriched_rows,
                "audit_response": artifact.audit_response,
            }
        )

    @staticmethod
    def _manifest_path(output_path: Path) -> Path:
        return output_path.with_name(output_path.name + ".manifest.json")

    @staticmethod
    def _embedded_manifest_hash(path: Path) -> str:
        workbook = load_workbook(path, read_only=False, data_only=True)
        try:
            description = _text(workbook.properties.description)
            if description.startswith(EMBEDDED_MANIFEST_PREFIX):
                return description[len(EMBEDDED_MANIFEST_PREFIX):]
            return ""
        finally:
            workbook.close()

    @staticmethod
    def _embed_manifest_hash(path: Path, input_manifest_sha256: str) -> None:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            workbook.properties.description = EMBEDDED_MANIFEST_PREFIX + input_manifest_sha256
            workbook.save(path)
        finally:
            workbook.close()

    @staticmethod
    def _header_index(headers: list[str], accepted: set[str]) -> int:
        return next((index for index, header in enumerate(headers) if header in accepted), -1)

    @staticmethod
    def _cell_value(value: Any) -> Any:
        """Keep a JSON-stable representation of the values users can observe in Excel."""
        if value is None or (isinstance(value, str) and value == ""):
            return ""
        return _jsonable(value)

    @staticmethod
    def _is_volatile_summary_header(value: object) -> bool:
        text = _text(value).lower()
        return any(token in text for token in ("time", "timestamp", "时间", "日期", "生成"))

    def _summary_business_projection(self, sheet: Any) -> list[list[Any]]:
        rows = [[self._cell_value(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        if not rows:
            return []
        volatile_columns = {
            index for index, header in enumerate(rows[0]) if self._is_volatile_summary_header(header)
        }
        return [
            [value for index, value in enumerate(row) if index not in volatile_columns]
            for row in rows
        ]

    @staticmethod
    def _legacy_display_header(header: str) -> str:
        # This is a runtime-only import.  It keeps the materializer tied to the
        # legacy workbook contract without importing its CLI during module load.
        from app.services.audit_workbook import _display_header

        return _display_header(header)

    @staticmethod
    def _legacy_cell_value(value: Any) -> Any:
        from app.services.audit_workbook import _format_cell_value

        return _format_cell_value(value)

    def _validate_frozen_source_matrix(
        self,
        headers: list[str],
        record_rows: list[tuple[Any, ...]],
        artifact: FrozenAuditArtifact,
    ) -> None:
        expected_headers = [self._legacy_display_header(str(header)) for header in artifact.headers]
        raw_headers = [str(header) for header in artifact.headers]
        if headers[:len(expected_headers)] not in (expected_headers, raw_headers):
            raise MaterializationValidationError("record_results does not preserve the legacy source header matrix")
        if len(record_rows) != len(artifact.frozen_rows):
            raise MaterializationValidationError("record_results row count does not equal frozen row count")
        for row_index, source_row in enumerate(artifact.frozen_rows):
            actual = record_rows[row_index]
            expected = [self._cell_value(self._legacy_cell_value(source_row.get(header, ""))) for header in artifact.headers]
            received = [self._cell_value(actual[index] if index < len(actual) else None) for index in range(len(expected))]
            if received != expected:
                raise MaterializationValidationError(
                    f"record_results source cell matrix differs from frozen row at index {row_index}"
                )

    def _validate_workbook(self, path: Path, artifact: FrozenAuditArtifact) -> WorkbookProjection:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            required_sheets = {"summary", "record_results"}
            if not required_sheets.issubset(set(workbook.sheetnames)):
                raise MaterializationValidationError("legacy workbook required sheets are missing")
            sheet = workbook["record_results"]
            iterator = sheet.iter_rows(values_only=True)
            try:
                headers = [_text(value) for value in next(iterator)]
            except StopIteration as exc:
                raise MaterializationValidationError("record_results sheet is empty") from exc
            v_index = self._header_index(headers, {"vCommodityId", "SaaS商品ID", "SaaS 商品ID"})
            sku_index = self._header_index(headers, {"skuId", "SKU ID", "SKU编号", "SKU编码"})
            url_index = self._header_index(headers, {"extLink", "商品链接", "互联网链接", "URL", "url", "链接"})
            if min(v_index, sku_index, url_index) < 0:
                raise MaterializationValidationError("record_results does not expose required business key columns")
            record_rows = list(iterator)
            actual_keys = []
            for row in record_rows:
                actual_keys.append(
                    (
                        _text(row[v_index] if v_index < len(row) else ""),
                        _text(row[sku_index] if sku_index < len(row) else ""),
                        _text(row[url_index] if url_index < len(row) else ""),
                    )
                )
            self._validate_frozen_source_matrix(headers, record_rows, artifact)
            record_matrix = [[self._cell_value(value) for value in row] for row in record_rows]
            summary_projection = self._summary_business_projection(workbook["summary"])
        finally:
            workbook.close()
        expected_keys = [self._row_key(row) for row in artifact.frozen_rows]
        if len(actual_keys) != len(expected_keys):
            raise MaterializationValidationError("record_results row count does not equal frozen row count")
        if actual_keys != expected_keys:
            raise MaterializationValidationError("record_results ordered business keys differ from frozen rows")
        return WorkbookProjection(
            tuple(actual_keys),
            _canonical_hash(record_matrix),
            _canonical_hash(summary_projection),
        )

    @staticmethod
    def _ordered_key_hash(keys: Sequence[tuple[str, str, str]]) -> str:
        return _canonical_hash(list(keys))

    def _manifest_payload(
        self,
        artifact: FrozenAuditArtifact,
        input_manifest_sha256: str,
        workbook_sha256: str,
        projection: WorkbookProjection,
    ) -> dict[str, Any]:
        return {
            "schema_version": MANIFEST_SCHEMA_VERSION,
            "input_manifest_sha256": input_manifest_sha256,
            "workbook_sha256": workbook_sha256,
            "row_count": len(projection.keys),
            "ordered_key_sha256": self._ordered_key_hash(projection.keys),
            "record_results_sha256": projection.record_results_sha256,
            "summary_business_sha256": projection.summary_business_sha256,
            "business_date": _text(artifact.business_date),
            "writer_contract_version": WRITER_CONTRACT_VERSION,
            "scope": {"tenant_id": TENANT_ID, "commodity_pool_id": COMMODITY_POOL_ID},
            "output_identity": _text(artifact.output_identity),
            "frozen_snapshot_id": _text(artifact.frozen_snapshot_id),
            "frozen_snapshot_hash": _text(artifact.frozen_snapshot_hash),
        }

    @staticmethod
    def _write_manifest_atomically(path: Path, payload: Mapping[str, Any]) -> str:
        path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
        try:
            encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
            with temp_path.open("xb") as stream:
                stream.write(encoded)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temp_path, path)
            return hashlib.sha256(encoded).hexdigest()
        finally:
            temp_path.unlink(missing_ok=True)

    def _result_from_existing(
        self,
        output_path: Path,
        manifest_path: Path,
        artifact: FrozenAuditArtifact,
        input_manifest_sha256: str,
    ) -> MaterializationResult:
        try:
            projection = self._validate_workbook(output_path, artifact)
        except MaterializationValidationError as exc:
            raise ArtifactConflictError("published workbook no longer matches its frozen artifact") from exc
        workbook_sha256 = _file_sha256(output_path)
        if manifest_path.is_file():
            try:
                manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError) as exc:
                raise ArtifactConflictError("published artifact has an unreadable manifest") from exc
            if not isinstance(manifest, dict) or manifest.get("input_manifest_sha256") != input_manifest_sha256:
                raise ArtifactConflictError("published artifact input manifest conflicts with requested frozen input")
            if manifest.get("workbook_sha256") != workbook_sha256:
                raise ArtifactConflictError("published artifact workbook hash differs from its manifest")
            if manifest.get("ordered_key_sha256") != self._ordered_key_hash(projection.keys):
                raise ArtifactConflictError("published artifact ordered keys differ from its manifest")
            if manifest.get("writer_contract_version") != WRITER_CONTRACT_VERSION:
                raise ArtifactConflictError("published artifact writer contract differs from this materializer")
            if manifest.get("record_results_sha256") != projection.record_results_sha256:
                raise ArtifactConflictError("published artifact record_results matrix differs from its manifest")
            if manifest.get("summary_business_sha256") != projection.summary_business_sha256:
                raise ArtifactConflictError("published artifact summary statistics differ from its manifest")
            return MaterializationResult(
                output_path, manifest_path, input_manifest_sha256, _file_sha256(manifest_path), workbook_sha256, len(projection.keys), True
            )

        if self._embedded_manifest_hash(output_path) != input_manifest_sha256:
            raise ArtifactConflictError("published workbook lacks embedded evidence for this frozen input")
        payload = self._manifest_payload(artifact, input_manifest_sha256, workbook_sha256, projection)
        manifest_sha256 = self._write_manifest_atomically(manifest_path, payload)
        return MaterializationResult(
            output_path, manifest_path, input_manifest_sha256, manifest_sha256, workbook_sha256, len(projection.keys), True, True
        )

    @staticmethod
    def _default_writer(path: Path, headers: list[str], rows: list[dict[str, Any]], enriched: list[dict[str, Any]], audit_response: Any) -> None:
        from app.services.audit_workbook import write_audit_workbook

        write_audit_workbook(path, headers, rows, enriched, audit_response)

    def materialize(self, artifact: FrozenAuditArtifact, output_path: Path | str) -> MaterializationResult:
        output_path = Path(output_path)
        self._validate_artifact(artifact, output_path)
        input_manifest_sha256 = self._input_manifest_hash(artifact)
        manifest_path = self._manifest_path(output_path)
        if output_path.exists():
            return self._result_from_existing(output_path, manifest_path, artifact, input_manifest_sha256)
        if manifest_path.exists():
            raise ArtifactConflictError("manifest exists without its published workbook")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = output_path.with_name(f".{output_path.stem}.{uuid.uuid4().hex}.tmp{output_path.suffix}")
        writer = self._writer or self._default_writer
        try:
            writer(
                temp_path,
                list(artifact.headers),
                [dict(row) for row in artifact.frozen_rows],
                [dict(row) for row in artifact.enriched_rows],
                artifact.audit_response,
            )
            self._embed_manifest_hash(temp_path, input_manifest_sha256)
            projection = self._validate_workbook(temp_path, artifact)
            os.replace(temp_path, output_path)
            if self._after_workbook_replace is not None:
                self._after_workbook_replace()
            workbook_sha256 = _file_sha256(output_path)
            payload = self._manifest_payload(artifact, input_manifest_sha256, workbook_sha256, projection)
            manifest_sha256 = self._write_manifest_atomically(manifest_path, payload)
            return MaterializationResult(
                output_path, manifest_path, input_manifest_sha256, manifest_sha256, workbook_sha256, len(projection.keys), False
            )
        finally:
            temp_path.unlink(missing_ok=True)
