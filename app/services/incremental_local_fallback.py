"""Durable timeout policy and orchestration for incremental local fallback."""
from __future__ import annotations

from collections import Counter
from datetime import UTC, datetime, time as datetime_time
from decimal import Decimal, InvalidOperation
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
from typing import Any, Callable, Mapping, Sequence
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook

from app.services.daily_audit_store import DailyAuditStore, LeaseToken
from app.services.page_price_semantics import valid_page_price_result
from scripts.filesystem_safety import is_reparse_point

DAYTIME_60M = "DAYTIME_60M"
PEAK_3M = "PEAK_3M"
AFTER_1800 = "AFTER_1800"
SHANGHAI = ZoneInfo("Asia/Shanghai")
SUPPORTED_MALLS = frozenset({
    "nbdeli", "guangbo", "ofs", "lxwlmall", "qstbg", "ycw", "comix", "tima",
})
MAX_ATTEMPTS = 2
MIN_EXECUTION_WINDOW_SECONDS = 60
MAX_CHILD_TIMEOUT_SECONDS = 600
EXPECTED_CRAWLER_REVISION = "9e4c1c1b182adcf3776a8ae5fcbbd6ba59a5911b"
DEFAULT_CRAWLER_ROOT = Path(r"E:\.A中核运维\cnnc-mall-crawler-releases\multisite-prod-20260924")
DEFAULT_CRAWLER_PYTHON = Path(r"E:\PythonEnvs\cnnc-mall-crawler-prod-9e4c1c\Scripts\python.exe")
SITE_ALIASES = {"ofs": "officemate"}
SITE_HOSTS = {
    "nbdeli": {"b2b.nbdeli.com"},
    "guangbo": {"btob.guangbo.net"},
    "officemate": {"officemate.cn", "www.officemate.cn"},
    "lxwlmall": {"www.lxwlmall.com"},
    "qstbg": {"qstbg.com", "www.qstbg.com"},
    "ycw": {"ycw.cn", "www.ycw.cn"},
    "comix": {"qx.com", "www.qx.com", "comix.com.cn", "www.comix.com.cn"},
    "tima": {"tima365.com", "www.tima365.com"},
}


class CrawlerOutcomes(dict):
    """Carry private artifact ownership without adding paths to business results."""

    artifact_dir: Path | None = None


def _plain_artifact_folder(folder: Path) -> bool:
    absolute = folder.absolute()
    return (
        absolute.parent.name == "local_fallback"
        and (os.name != "nt" or absolute.drive.upper() == "E:")
        and not any(is_reparse_point(path) for path in (absolute, *absolute.parents))
    )


def _cleanup_settled_artifacts(outcomes: Mapping, states: Sequence[str]) -> None:
    folder = getattr(outcomes, "artifact_dir", None)
    if folder is None or not states or any(state not in {"SUCCESS", "TERMINAL_FAILED"} for state in states):
        return
    folder = Path(folder)
    if not _plain_artifact_folder(folder):
        return
    if any(state != "SUCCESS" for state in states):
        # Durable errors and original workbooks remain available for diagnosis.
        # This marker is written only AFTER every claim has committed terminally.
        marker = folder / ".durably-settled.json"
        if is_reparse_point(marker):
            return
        marker.write_text(json.dumps({"version": 1, "states": list(states)}), encoding="utf-8")
        reviewed = []
        for old in folder.parent.glob("*/.durably-settled.json"):
            if is_reparse_point(old) or not _plain_artifact_folder(old.parent):
                continue
            try:
                proof = json.loads(old.read_text(encoding="utf-8"))
                # A terminal crawler error is NOT a completed human review.
                if proof.get("diagnostics_reviewed") is not True or proof.get("version") != 1 or not proof.get("states") or any(state not in {"SUCCESS", "TERMINAL_FAILED"} for state in proof["states"]):
                    continue
                reviewed.append(old)
            except (OSError, ValueError):
                continue
        for old in sorted(reviewed, key=lambda path: path.stat().st_mtime_ns)[:-5]:
            if _remove_known_artifacts(old.parent):
                old.unlink()
                if not any(old.parent.iterdir()):
                    old.parent.rmdir()
        return
    _remove_known_artifacts(folder)
    if not any(folder.iterdir()):
        folder.rmdir()


def _remove_known_artifacts(folder: Path) -> bool:
    if not _plain_artifact_folder(folder):
        return False
    names = ["manifest.xlsx"]
    for site in SITE_HOSTS:
        names.extend((f"{site}_input.xlsx", f"{site}_output.xlsx"))
    targets = [folder / name for name in (*names, "_tmp", "_pycache")]
    for child in targets:
        if is_reparse_point(child):
            return False
        if child.is_dir() and any(is_reparse_point(path) for path in child.rglob("*")):
            return False
    for name in names:
        (folder / name).unlink(missing_ok=True)
    for name in ("_tmp", "_pycache"):
        child = folder / name
        if child.is_dir() and not is_reparse_point(child):
            shutil.rmtree(child)
    return True


def _moment(value: datetime | str | None) -> datetime | None:
    if value is None:
        return None
    parsed = value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise ValueError("fallback timestamps must include an explicit offset")
    return parsed.astimezone(UTC)


def eligibility_reason(
    now: datetime,
    remote_started_at: datetime | str,
    last_remote_check_at: datetime | str | None,
) -> str | None:
    """Return the audited reason for the current window, or ``None``."""
    current = _moment(now)
    started = _moment(remote_started_at)
    previous = _moment(last_remote_check_at)
    assert current is not None and started is not None
    local = current.astimezone(SHANGHAI)
    local_time = local.timetz().replace(tzinfo=None)
    if local_time >= datetime_time(18, 0):
        return AFTER_1800
    if local_time >= datetime_time(17, 0):
        peak_start = local.replace(hour=17, minute=0, second=0, microsecond=0).astimezone(UTC)
        if previous is not None and previous >= peak_start and (current - previous).total_seconds() >= 180:
            return PEAK_3M
        return None
    if (current - started).total_seconds() >= 3600:
        return DAYTIME_60M
    return None


def _metrics(store: DailyAuditStore, token: LeaseToken) -> dict[str, Any]:
    rows = store.list_local_fallback_work(token)
    unique: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows:
        key = (str(row["mall_key"]), str(row["exact_url"]))
        item = unique.setdefault(key, dict(row))
        reason = item.get("eligible_reason") or row.get("eligible_reason")
        if str(row.get("crawl_state")) == "SUCCESS":
            item.update(row)
        if reason:
            item["eligible_reason"] = reason
    states: Counter[str] = Counter()
    reasons: Counter[str] = Counter()
    sites: dict[str, dict[str, int]] = {}
    links: list[dict[str, Any]] = []
    for row in unique.values():
        source = ""
        if str(row.get("crawl_state")) == "SUCCESS":
            try:
                result = json.loads(str(row.get("crawl_result_json") or "{}"))
            except (TypeError, ValueError):
                result = {}
            if isinstance(result, Mapping):
                source = str(result.get("page_price_source") or "")
        if str(row.get("crawl_state")) == "SUCCESS":
            state = "SUCCESS" if source.startswith("本地兜底:") else "REMOTE_SUCCESS"
        else:
            state = str(row["state"])
        states[state] += 1
        if row.get("eligible_reason"):
            reasons[str(row["eligible_reason"])] += 1
        site = sites.setdefault(
            str(row["mall_key"]),
            {"pending": 0, "running": 0, "success": 0, "remote_success": 0, "failed": 0},
        )
        effective_state = state
        if effective_state == "CLAIMED":
            site["running"] += 1
        elif effective_state == "SUCCESS":
            site["success"] += 1
        elif effective_state == "REMOTE_SUCCESS":
            site["remote_success"] += 1
        elif effective_state in {"RETRYABLE_FAILED", "TERMINAL_FAILED"}:
            site["failed"] += 1
        else:
            site["pending"] += 1
        links.append({
            "mall_key": str(row["mall_key"]),
            "exact_url": str(row["exact_url"]),
            "state": effective_state,
            "source": source or "待定",
            "remote_job_id": row.get("remote_job_id"),
        })
    return {
        "total": len(unique),
        "pending": states["PENDING"],
        "running": states["CLAIMED"],
        "success": states["SUCCESS"],
        "remote_success": states["REMOTE_SUCCESS"],
        "failed": states["RETRYABLE_FAILED"] + states["TERMINAL_FAILED"],
        "reasons": dict(reasons),
        "sites": sites,
        "links": links,
    }


def _normalized_site(mall_key: str) -> str:
    return SITE_ALIASES.get(mall_key, mall_key)


def _url_matches_site(url: str, site: str) -> bool:
    return (urlparse(url).hostname or "").lower() in SITE_HOSTS.get(site, set())


def read_crawler_output(
    path: Path,
    site: str,
    expected_urls: Sequence[str],
) -> dict[str, dict[str, Any]]:
    """Strictly reconcile one child workbook with its locked exact URLs."""
    expected = list(expected_urls)
    if len(expected) != len(set(expected)):
        raise ValueError("锁定输入 URL 重复")
    if not path.is_file():
        raise ValueError(f"本地 crawler 输出不存在: {path}")
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"本地 crawler 输出不是有效 XLSX: {path}") from exc
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in (next(rows, ()) or ())]
        required = {"商品价格", "商品链接", "采集状态", "失败原因"}
        if not required.issubset(headers):
            raise ValueError("本地 crawler 输出表头不完整")
        index = {name: headers.index(name) for name in required}
        results: dict[str, dict[str, Any]] = {}
        for values in rows:
            url = str(values[index["商品链接"]] or "").strip()
            if not url:
                continue
            if url in results:
                raise ValueError(f"本地 crawler 输出 URL 重复: {url}")
            if not _url_matches_site(url, site):
                raise ValueError(f"本地 crawler 输出 URL 与站点不匹配: {site} {url}")
            status = str(values[index["采集状态"]] or "").strip().upper()
            reason = str(values[index["失败原因"]] or "").strip()
            if status == "SUCCESS":
                try:
                    price = Decimal(str(values[index["商品价格"]]))
                except (InvalidOperation, TypeError, ValueError) as exc:
                    raise ValueError(f"本地 crawler SUCCESS 价格非法: {url}") from exc
                if not price.is_finite() or price < 0:
                    raise ValueError(f"本地 crawler SUCCESS 价格非法: {url}")
                results[url] = {
                    "result": {
                        "page_price_value": price,
                        "page_price_status": "PRICE_IS_ZERO" if price == 0 else "OK",
                        "page_price_source": f"本地兜底:{site}",
                    }
                }
            else:
                risk = status == "SKIPPED_RISK" or any(
                    marker in reason.upper() for marker in ("RISK", "CAPTCHA", "LOGIN", "验证码", "登录")
                )
                retryable = status in {"FAILED", "UNKNOWN"} and any(
                    marker in reason.upper() for marker in ("NETWORK", "TIMEOUT", "TEMPORARY", "网络", "超时")
                )
                results[url] = {
                    "error": ("RISK_STOP" if risk else status or "INVALID_RESPONSE") + (f": {reason}" if reason else ""),
                    "retryable": retryable and not risk,
                }
        if set(results) != set(expected):
            raise ValueError(
                "本地 crawler 输出未与锁定 URL 集合逐项对账: "
                f"missing={len(set(expected) - set(results))} foreign={len(set(results) - set(expected))}"
            )
        return results
    finally:
        workbook.close()


def _child_env(run_dir: Path) -> dict[str, str]:
    if os.name == "nt" and run_dir.resolve().drive.upper() != "E:":
        raise ValueError("local fallback temporary paths must be on E drive")
    allowed = {
        "SYSTEMROOT", "WINDIR", "PATH", "PATHEXT", "COMSPEC", "USERPROFILE",
        "LOCALAPPDATA", "APPDATA", "PROGRAMFILES", "PROGRAMFILES(X86)",
        "PROGRAMDATA", "HOMEDRIVE", "HOMEPATH",
    }
    env = {key: value for key, value in os.environ.items() if key.upper() in allowed}
    temp = run_dir / "_tmp"
    pycache = Path(os.environ.get("CNNC_PYTHON_CACHE_ROOT", r"E:\CodexData\cache\python-pyc\cnnc")) / f"{sys.version_info.major}.{sys.version_info.minor}"
    if os.name == "nt" and (not pycache.is_absolute() or pycache.resolve().drive.upper() != "E:"):
        raise ValueError("local fallback bytecode cache must be an absolute E-drive path")
    temp.mkdir(parents=True, exist_ok=True)
    pycache.mkdir(parents=True, exist_ok=True)
    env.update({
        "TEMP": str(temp.resolve()), "TMP": str(temp.resolve()),
        "PYTHONPYCACHEPREFIX": str(pycache.resolve()), "PYTHONIOENCODING": "utf-8",
    })
    return env


def verify_crawler_revision(
    crawler_root: Path,
    *,
    expected_revision: str = EXPECTED_CRAWLER_REVISION,
    runner: Callable[..., subprocess.CompletedProcess] = subprocess.run,
) -> str:
    root = Path(crawler_root).resolve()
    if not (root / "mall_crawler" / "cli.py").is_file():
        raise ValueError(f"固定 crawler 入口不存在: {root}")
    env = _child_env(root / "_run" / "revision-check")
    revision = runner(
        ["git", "rev-parse", "HEAD"], cwd=str(root), env=env, check=False,
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    )
    actual = str(revision.stdout or "").strip().lower()
    status = runner(
        ["git", "status", "--porcelain"], cwd=str(root), env=env, check=False,
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    )
    if revision.returncode != 0 or not re.fullmatch(r"[0-9a-f]{40}", actual):
        raise ValueError("固定 crawler Git revision 无法验证")
    if status.returncode != 0 or str(status.stdout or "").strip():
        raise ValueError("固定 crawler 工作区不干净")
    if actual != expected_revision.lower():
        raise ValueError(f"固定 crawler revision 不匹配: expected={expected_revision} actual={actual}")
    return actual


def execute_pinned_crawler(
    claims: Sequence[Mapping[str, Any]],
    *,
    run_dir: Path,
    tick_id: str,
    crawler_root: Path = DEFAULT_CRAWLER_ROOT,
    timeout_seconds: float | Callable[[], float] = MAX_CHILD_TIMEOUT_SECONDS,
) -> Mapping[int, Mapping[str, Any]]:
    """Run one locked multi-site child with an allowlisted environment."""
    claim_list = [dict(claim) for claim in claims]
    if not claim_list:
        return {}
    grouped: dict[str, list[dict[str, Any]]] = {}
    for claim in claim_list:
        site = _normalized_site(str(claim["mall_key"]))
        grouped.setdefault(site, []).append(claim)
    outcomes: CrawlerOutcomes = CrawlerOutcomes()
    by_site: dict[str, list[dict[str, Any]]] = {}
    for site, site_claims in grouped.items():
        urls = [str(claim["exact_url"]) for claim in site_claims]
        if site not in SITE_HOSTS or any(not _url_matches_site(url, site) for url in urls):
            error = "INVALID_SCOPE"
        elif len(urls) != len(set(urls)):
            error = "MANIFEST_MISMATCH: duplicate exact URL"
        else:
            by_site[site] = site_claims
            continue
        for claim in site_claims:
            outcomes[int(claim["id"])] = {"error": error, "retryable": False}
    if not by_site:
        return outcomes
    live_timeout = timeout_seconds() if callable(timeout_seconds) else timeout_seconds
    if live_timeout < MIN_EXECUTION_WINDOW_SECONDS:
        outcomes.update({int(claim["id"]): {"deferred": True} for claims_for_site in by_site.values() for claim in claims_for_site})
        return outcomes
    revision = verify_crawler_revision(crawler_root)
    live_timeout = timeout_seconds() if callable(timeout_seconds) else timeout_seconds
    if live_timeout < MIN_EXECUTION_WINDOW_SECONDS:
        outcomes.update({int(claim["id"]): {"deferred": True} for claims_for_site in by_site.values() for claim in claims_for_site})
        return outcomes
    # The caller-selected runtime may be an intentional junction. Capture its
    # physical E-drive boundary once, but never follow links inside owned work.
    folder = Path(run_dir).resolve() / "local_fallback" / re.sub(r"[^A-Za-z0-9_.-]+", "_", tick_id)
    if os.name == "nt" and folder.resolve().drive.upper() != "E:":
        raise ValueError("local fallback artifacts must be on E drive")
    if is_reparse_point(folder.parent) or is_reparse_point(folder):
        raise ValueError("local fallback owned artifacts must not be reparse points")
    folder.mkdir(parents=True, exist_ok=True)
    outcomes.artifact_dir = folder
    manifest = Workbook()
    sheet = manifest.active
    sheet.title = "jobs"
    sheet.append(["site", "input", "output"])
    outputs: dict[str, Path] = {}
    for site, site_claims in sorted(by_site.items()):
        input_path = (folder / f"{site}_input.xlsx").resolve()
        output_path = (folder / f"{site}_output.xlsx").resolve()
        book = Workbook()
        input_sheet = book.active
        input_sheet.append(["URL"])
        for claim in site_claims:
            input_sheet.append([str(claim["exact_url"])])
        book.save(input_path)
        sheet.append([site, str(input_path), str(output_path)])
        outputs[site] = output_path
    manifest_path = (folder / "manifest.xlsx").resolve()
    manifest.save(manifest_path)
    env = _child_env(folder)
    live_timeout = timeout_seconds() if callable(timeout_seconds) else timeout_seconds
    if live_timeout < MIN_EXECUTION_WINDOW_SECONDS:
        outcomes.update({int(claim["id"]): {"deferred": True} for claims_for_site in by_site.values() for claim in claims_for_site})
        return outcomes
    try:
        completed = subprocess.run(
            [str(DEFAULT_CRAWLER_PYTHON), "-m", "mall_crawler.cli", "--manifest", str(manifest_path)],
            cwd=str(Path(crawler_root).resolve()), env=env, check=False,
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            timeout=min(MAX_CHILD_TIMEOUT_SECONDS, live_timeout),
        )
    except subprocess.TimeoutExpired:
        outcomes.update({int(claim["id"]): {"error": "TEMPORARY_TIMEOUT", "retryable": True} for claims_for_site in by_site.values() for claim in claims_for_site})
        return outcomes
    for site, site_claims in by_site.items():
        expected = [str(claim["exact_url"]) for claim in site_claims]
        try:
            parsed = read_crawler_output(outputs[site], site, expected)
        except Exception as exc:
            if not outputs[site].is_file() and completed.returncode != 0:
                risk = completed.returncode == 3
                error = "RISK_STOP" if risk else f"PROCESS_FAILED: exit={completed.returncode}"
                retryable = not risk
            else:
                error = f"MANIFEST_MISMATCH: {exc}"
                retryable = False
            for claim in site_claims:
                outcomes[int(claim["id"])] = {"error": error, "retryable": retryable}
            continue
        for claim in site_claims:
            outcome = dict(parsed[str(claim["exact_url"])])
            if "result" in outcome:
                outcome["crawler_revision"] = revision
            outcomes[int(claim["id"])] = outcome
    return outcomes


def run_incremental_local_fallback(
    store: DailyAuditStore,
    token: LeaseToken,
    *,
    now: datetime,
    tick_id: str,
    final_remote_check: Callable[[Sequence[Mapping[str, Any]]], set[int]],
    execute: Callable[[Sequence[Mapping[str, Any]]], Mapping[int, Mapping[str, Any]]],
    checked_remote_job_ids: set[int] | frozenset[int] = frozenset(),
    execution_budget_seconds: Callable[[], float] | None = None,
) -> dict[str, Any]:
    """Observe, final-check, claim, execute and settle exact unresolved URLs."""
    current = _moment(now)
    assert current is not None
    store.recover_interrupted_local_fallback_claims(
        token, tick_id, current.isoformat(timespec="seconds"), max_attempts=MAX_ATTEMPTS
    )
    opening_checked = {int(value) for value in checked_remote_job_ids}
    observed = store.observe_local_fallback_remote_checks(
        token,
        current.isoformat(timespec="seconds"),
        SUPPORTED_MALLS,
        checked_remote_job_ids=opening_checked,
    )
    eligible: list[dict[str, Any]] = []
    peak_initial_checks: list[dict[str, Any]] = []
    seen_urls: set[str] = set()
    local_time = current.astimezone(SHANGHAI)
    peak_start = local_time.replace(hour=17, minute=0, second=0, microsecond=0).astimezone(UTC)
    in_peak = datetime_time(17, 0) <= local_time.timetz().replace(tzinfo=None) < datetime_time(18, 0)
    for row in observed:
        reason = eligibility_reason(current, row["remote_started_at"], row.get("last_remote_check_at"))
        exact_url = str(row["exact_url"])
        if exact_url in seen_urls:
            continue
        previous = _moment(row.get("last_remote_check_at"))
        remote_job_id = int(row["remote_job_id"])
        if (
            reason is None
            and in_peak
            and (previous is None or previous < peak_start)
            and remote_job_id not in opening_checked
        ):
            seen_urls.add(exact_url)
            peak_initial_checks.append(dict(row))
            continue
        if reason is None:
            continue
        seen_urls.add(exact_url)
        eligible.append({**row, "eligible_reason": reason})
    if peak_initial_checks:
        try:
            checked = {int(value) for value in final_remote_check(peak_initial_checks)}
        except Exception:
            return _metrics(store, token)
        if checked:
            store.observe_local_fallback_remote_checks(
                token,
                current.isoformat(timespec="seconds"),
                SUPPORTED_MALLS,
                checked_remote_job_ids=checked,
            )
    if not eligible:
        return _metrics(store, token)

    # The final check is separate from the opening sweep. A local claim
    # requires proof that this exact remote job completed that final query.
    final_candidates = eligible
    final_checked: set[int] = set()
    if final_candidates:
        try:
            final_checked = {int(value) for value in final_remote_check(final_candidates)}
        except Exception:
            return _metrics(store, token)
        if final_checked:
            store.observe_local_fallback_remote_checks(
                token,
                current.isoformat(timespec="seconds"),
                SUPPORTED_MALLS,
                checked_remote_job_ids=final_checked,
            )
    if execution_budget_seconds is not None and execution_budget_seconds() < MIN_EXECUTION_WINDOW_SECONDS:
        return _metrics(store, token)
    claims: list[dict[str, Any]] = []
    for row in eligible:
        if execution_budget_seconds is not None and execution_budget_seconds() < MIN_EXECUTION_WINDOW_SECONDS:
            for claim in claims:
                store.defer_unstarted_local_fallback(token, int(claim["id"]), tick_id)
            return _metrics(store, token)
        if int(row["remote_job_id"]) not in final_checked:
            continue
        current_work = store.get_crawl_work(int(row["id"]))
        if current_work is None or current_work["state"] == "SUCCESS":
            continue
        claim = store.claim_local_fallback(
            token,
            int(row["id"]),
            str(row["eligible_reason"]),
            tick_id,
            current.isoformat(timespec="seconds"),
            max_attempts=MAX_ATTEMPTS,
        )
        if claim is not None:
            claims.append(claim)
    if not claims:
        return _metrics(store, token)
    try:
        outcomes = execute(tuple(claims))
    except Exception as exc:
        outcomes = {
            int(claim["id"]): {"error": f"PROCESS_FAILED: {type(exc).__name__}: {exc}", "retryable": True}
            for claim in claims
        }
    settled_states: list[str] = []
    for claim in claims:
        fallback_id = int(claim["id"])
        outcome = outcomes.get(fallback_id) or outcomes.get(str(fallback_id))  # type: ignore[arg-type]
        if isinstance(outcome, Mapping) and outcome.get("deferred") is True:
            store.defer_unstarted_local_fallback(token, fallback_id, tick_id)
            settled_states.append("PENDING")
            continue
        if not isinstance(outcome, Mapping):
            settled = store.fail_local_fallback(
                token, fallback_id, "PROCESS_FAILED: missing child outcome", retryable=True,
                max_attempts=MAX_ATTEMPTS,
            )
            settled_states.append(str(settled["state"]))
            continue
        result = outcome.get("result")
        if isinstance(result, Mapping) and valid_page_price_result(result):
            settled = store.complete_local_fallback(
                token,
                fallback_id,
                result,
                crawler_revision=str(outcome.get("crawler_revision") or "injected-test-executor"),
                completed_at=current.isoformat(timespec="seconds"),
            )
            settled_states.append(str(settled["state"]))
            continue
        settled = store.fail_local_fallback(
            token,
            fallback_id,
            str(outcome.get("error") or "INVALID_RESPONSE: no valid exact price"),
            retryable=bool(outcome.get("retryable", False)),
            max_attempts=MAX_ATTEMPTS,
        )
        settled_states.append(str(settled["state"]))
    try:
        _cleanup_settled_artifacts(outcomes, settled_states)
    except OSError:
        # Cleanup must not turn committed business success into a retry.
        pass
    return _metrics(store, token)
