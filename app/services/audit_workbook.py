# Canonical reusable capability; original script retains compatibility imports.
from __future__ import annotations
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from app.schemas.workflow import WorkflowAction
from app.services import saas_audit_core as _audit_core
from app.integrations.zhaoshang_mall_mapping import is_zhaoshang_domain

normalize_cell_str = _audit_core.normalize_cell_str

parse_decimal = _audit_core.parse_decimal

compute_over_30_check = _audit_core.compute_over_30_check

compute_saas_price_vs_customer = _audit_core.compute_saas_price_vs_customer

compute_price_changed = _audit_core.compute_price_changed

compute_final_decision = _audit_core.compute_final_decision

RESULT_APPENDED_COLUMNS = [
    "fourCategoryId",
    "fourCategoryName",
    "customer_lookup_status",
    "list_reference_price",
    "customerPrice",
    "customerPriceStatus",
    "customerPriceSource",
    "customerPriceWarning",
    "customerPriceEvidence",
    "json_price",
    "saas_price_vs_customer",
    "title_match_result",
    "title_match_confidence",
    "title_match_reason",
    "sku_link_check",
    "price_check",
    "price_vs_customer_check",
    "primary_action",
    "suggested_actions",
    "reason",
    "evidence",
    "needs_human_approval",
    "execution_allowed",
    "executed",
    "page_price_value",
    "page_price_status",
    "page_price_source",
    "page_vs_customer_check",
    "over_30pct",
    "price_changed",
    "final_decision",
]

HIGHLIGHTED_RESULT_COLUMNS = {
    "customerPrice",
    "customerPriceStatus",
    "primary_action",
    "price_vs_customer_check",
    "saas_price_vs_customer",
    "page_price_value",
    "over_30pct",
    "price_changed",
    "final_decision",
}

COLUMN_DISPLAY_NAMES = {
    # 原始 SaaS 列
    "vCommodityId": "SaaS商品ID",
    "documentId": "文档ID",
    "queryStr": "客户标题",
    "sassVCommodityName": "SaaS商品名称",
    "skuId": "SKU编号",
    "skuName": "互联网标题",
    "vendorName": "供应商",
    "shopName": "店铺名称",
    "price": "SaaS价格",
    "pageActualPrice": "页面实际价格",
    "extLink": "互联网链接",
    "commentNum": "评论数",
    "createTime": "创建时间",
    "updateTime": "更新时间",
    "modelAuditResult": "模型审核结果",
    "isMatch": "是否匹配",
    "customerPrice": "客户价",
    # 客户列
    "fourCategoryId": "四级类目ID",
    "fourCategoryName": "四级类目名称",
    "customer_lookup_status": "客户数据匹配状态",
    "list_reference_price": "客户列表参考价",
    # 客户价详情
    "customerPriceStatus": "客户价状态",
    "customerPriceSource": "客户价来源",
    "customerPriceWarning": "客户价警告",
    "customerPriceEvidence": "客户价证据",
    # 审查结果
    "json_price": "JSON价格",
    "title_match_result": "标题匹配结果",
    "title_match_confidence": "标题匹配置信度",
    "title_match_reason": "标题匹配原因",
    "sku_link_check": "SKU链接检查",
    "price_check": "价格检查",
    "price_vs_customer_check": "价格对比客户价",
    "primary_action": "建议操作",
    "suggested_actions": "建议操作列表",
    "reason": "原因",
    "evidence": "证据",
    "needs_human_approval": "需人工审批",
    "execution_allowed": "允许执行",
    "executed": "已执行",
    # 爬虫
    "page_price_value": "爬取页面价格",
    "page_price_status": "页面价格状态",
    "page_price_source": "页面价格来源",
    # 30% 高价
    "over_30pct": "超过客户价30%",
    "saas_price_vs_customer": "SaaS价对比客户价",
    "price_changed": "价格变动",
    "final_decision": "最终决策",
    # SAAS_LIST 特有
    "commodityName": "商品名称",
    "referenceLink": "参考链接",
    "categoryId": "类目ID",
    "categoryName": "类目名称",
    "materialCode": "物料编码",
    "materialDesc": "物料描述",
    "quasiStandardCommodityCode": "准标准商品编码",
}

def _display_header(name: str) -> str:
    """英文字段名 → 中文显示名；未映射的保持原文。"""
    return COLUMN_DISPLAY_NAMES.get(name, name)

def _display_value(value: Any) -> Any:
    """枚举值 → 中文显示。"""
    if isinstance(value, str) and value in VALUE_DISPLAY_NAMES:
        return VALUE_DISPLAY_NAMES[value]
    return value

HEADER_FONT = Font(bold=True)

HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

HIGHLIGHT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

OVER_30_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

VALUE_DISPLAY_NAMES = {
    # WorkflowAction
    "APPROVE": "建议通过",
    "REJECT": "建议驳回",
    "MANUAL_CHECK": "人工核查",
    "WAIT_CRAWLER": "等待爬取",
    "SUPPLEMENT_SAME_PRODUCT": "补充同款",
    "DO_NOT_PUSH": "不推送",
    "FEEDBACK_CUSTOMER": "反馈客户",
    "READY_TO_PUSH": "准备推送",
    "IMPORT_CUSTOMER_LINKS": "导入客户链接",
    "UNBIND_WRONG_RELATIONS": "解绑错误关系",
    "RECALCULATE_INDEX": "重新计算指数",
    "PUSH_INDEX": "推送指数",
    "ASK_CUSTOMER_RECHECK": "请客户复核",
    "MANUAL_VERIFY": "人工核验",
    "FORWARD_TO_WORK01": "转交Work01",
    "NO_ACTION": "无操作",
    "WAIT_COLLECTION": "等待采集",
    "RETRY_PUSH": "重试推送",
    "CONFIRM": "确认",
    "ALERT": "告警",
    "EXEMPTION": "豁免",
    "QUERY": "查询",
    "EXPORT": "导出",
    # title_match_result
    "SAME": "同款",
    "NOT_SAME": "非同款",
    "MANUAL_REVIEW": "人工复核",
    "MODEL_UNAVAILABLE": "模型不可用",
    # customerPriceStatus
    "FOUND": "已找到",
    "FOUND_SINGLE_REFERENCE": "仅找到单个参考价",
    "FOUND_CACHED": "已找到(缓存)",
    "EXISTS": "已存在",
    "NOT_PROVIDED": "未提供",
    "NOT_FOUND": "未找到",
    "CUSTOMER_NOT_FOUND": "客户侧查无此单据",
    "LOOKUP_FAILED": "匹配失败",
    "DETAIL_REQUEST_FAILED": "详情请求失败",
    "DETAIL_REQUEST_SKIPPED": "详情请求跳过",
    "DETAIL_PARSE_FAILED": "详情解析失败",
    "DETAIL_SINGLE_REFERENCE": "详情单参考价",
    "DETAIL_REFERENCE_INFOS": "详情多参考价",
    "EXCEL_INPUT": "Excel输入",
    # customer_lookup_status
    "MATCHED": "已匹配",
    "MATCHED_NO_CATEGORY": "已匹配(无类目)",
    # sku_link_check
    "MATCH": "一致",
    "MISMATCH": "不一致",
    "UNKNOWN": "未知",
    # price_check
    "JSON_PRICE_MISSING": "JSON价格为空",
    "WAIT_PAGE_PRICE_CHECK": "待页面价格核验",
    # price_vs_customer_check
    "WITHIN_THRESHOLD": "阈值内",
    "HIGHER_THAN_CUSTOMER_PRICE_30": "超过客户价30%",
    "CUSTOMER_PRICE_MISSING": "客户价缺失",
    "CUSTOMER_PRICE_INVALID": "客户价无效",
    "CUSTOMER_ABSENT_SKIPPED": "客户侧无此单据，本轮无需处理",
    "CANDIDATE_PRICE_MISSING": "候选价格缺失",
    "CANDIDATE_PRICE_INVALID": "候选价格无效",
    "PAGE_PRICE_MISSING": "页面价格缺失",
    # page_price_status
    "PENDING_CRAWL": "待爬取",
    "SKIPPED_JD": "跳过(京东)",
    "SKIPPED_ZHAOSHANG": "招商局处理",
    "SKIPPED_NO_LINK": "跳过(无链接)",
    "SKIPPED_MODEL_NOT_SAME": "跳过(模型判非同款)",
    "SKIPPED": "跳过",
    "OK": "成功",
    "FAILED": "失败",
    "CRAWLER_PROCESS_FAILED": "爬虫进程失败",
    "PRICE_MISSING": "价格不可见",
    "NO_RESULT": "无采集结果",
    "PRICE_IS_ZERO": "价格为0",
    # over_30pct
    "TRUE": "是",
    "FALSE": "否",
}

def _format_cell_value(value: Any) -> Any:
    if value is None:
        return ""
    value = _display_value(value)
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, WorkflowAction):
        return _display_value(value.value)
    if isinstance(value, list):
        rendered: list[str] = []
        for item in value:
            if isinstance(item, WorkflowAction):
                rendered.append(_display_value(item.value))
            else:
                rendered.append(str(_display_value(item)))
        return "；".join(rendered)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return value

def _build_summary_counts(audit_response: Any, enriched: list[dict[str, Any]]) -> dict[str, Any]:
    summary = audit_response.summary or {}
    record_results = audit_response.record_results or []
    customer_price_missing = sum(
        1 for info in enriched
        if info.get("customerPriceStatus") in {"NOT_PROVIDED", "DETAIL_REQUEST_SKIPPED", "NOT_FOUND", "DETAIL_PARSE_FAILED"}
    )
    customer_price_failed = sum(
        1 for info in enriched if info.get("customerPriceStatus") == "DETAIL_REQUEST_FAILED"
    )
    customer_price_found = sum(
        1 for info in enriched if info.get("customerPriceStatus") == "FOUND"
    )
    customer_price_single = sum(
        1 for info in enriched if info.get("customerPriceStatus") == "FOUND_SINGLE_REFERENCE"
    )
    return {
        "总记录数": int(summary.get("total_records", len(record_results))),
        "建议通过数量": int(summary.get("approve_suggestion_count", 0)),
        "人工核查数量": int(summary.get("manual_check_count", 0)),
        "建议驳回数量": int(summary.get("reject_suggestion_count", 0)),
        "客户价高于30%数量": int(summary.get("higher_than_customer_price_30_count", 0)),
        "客户价FOUND数量": customer_price_found,
        "客户价SINGLE数量": customer_price_single,
        "客户价缺失数量": customer_price_missing,
        "客户价请求失败数量": customer_price_failed,
        "价格不一致数量": int(summary.get("price_mismatch_count", 0)),
        "输出时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

def _set_column_widths(sheet: Any, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        display = _display_header(header)
        width = min(max(len(display) * 2 + 4, 12), 60)
        if header in {"reason", "evidence", "customerPriceEvidence", "extLink", "queryStr", "skuName", "customerPriceWarning", "title_match_reason"}:
            width = 50
        sheet.column_dimensions[letter].width = width

def _write_summary_sheet(workbook: Workbook, summary_counts: dict[str, Any]) -> None:
    sheet = workbook.active
    sheet.title = "summary"
    headers = list(summary_counts.keys())
    sheet.append(headers)
    sheet.append([_format_cell_value(summary_counts[key]) for key in headers])
    for column_index, _ in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    _set_column_widths(sheet, headers)

def _collect_title_reason(audit_result):
    if audit_result is None:
        return ""
    chunks = []
    for item in audit_result.evidence or []:
        text = str(item)
        if text.startswith("title_model.reason="):
            chunks.append(text[len("title_model.reason="):])
        elif text.startswith("title_model.risk="):
            chunks.append("risk=" + text[len("title_model.risk="):])
    return ";".join(chunks)

def _is_zhaoshang_mall(url: str) -> bool:
    """检查 URL 是否属于招商局任务系统商城（不走 multi_mall_crawler）。"""
    return is_zhaoshang_domain(url)

def _compute_page_price_status(row, audit_result):
    # 如果已被 apply_crawl_results_to_rows 写入最终结果，直接用
    if "__page_price_status" in row:
        return {
            "page_price_value": row.get("__page_price_value"),
            "page_price_status": row.get("__page_price_status", ""),
            "page_price_source": row.get("__page_price_source", ""),
        }
    ext_link = normalize_cell_str(row.get("extLink"))
    if not ext_link:
        return {
            "page_price_value": None,
            "page_price_status": "SKIPPED_NO_LINK",
            "page_price_source": "",
        }
    if _is_zhaoshang_mall(ext_link):
        return {
            "page_price_value": None,
            "page_price_status": "SKIPPED_ZHAOSHANG",
            "page_price_source": "招商局",
        }
    # 注意：模型 NOT_SAME 也要爬取价格（业务要求）
    # 不再跳过爬虫，让 over_30pct + saas_price_vs_customer + price_changed 一起兜底判断
    return {
        "page_price_value": None,
        "page_price_status": "PENDING_CRAWL",
        "page_price_source": "",
    }

def _write_record_results_sheet(
    workbook: Workbook,
    headers: list[str],
    rows: list[dict[str, Any]],
    enriched: list[dict[str, Any]],
    audit_response: Any,
) -> None:
    sheet = workbook.create_sheet("record_results")
    base_headers = [header for header in headers if header not in RESULT_APPENDED_COLUMNS]
    final_headers = base_headers + RESULT_APPENDED_COLUMNS
    display_headers = [_display_header(h) for h in final_headers]
    sheet.append(display_headers)
    for column_index, header in enumerate(display_headers, start=1):
        cell = sheet.cell(row=1, column=column_index)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN

    record_results = list(audit_response.record_results or [])
    if len(record_results) != len(rows):
        while len(record_results) < len(rows):
            record_results.append(None)

    for row_index, (row, info, audit_result) in enumerate(zip(rows, enriched, record_results), start=2):
        merged: dict[str, Any] = {}
        for header in base_headers:
            merged[header] = row.get(header, "")

        merged["fourCategoryId"] = row.get("fourCategoryId", "")
        merged["fourCategoryName"] = row.get("fourCategoryName", "")
        merged["customer_lookup_status"] = row.get("customer_lookup_status", "")
        merged["list_reference_price"] = row.get("list_reference_price")

        merged["customerPrice"] = info.get("customerPrice")
        merged["customerPriceStatus"] = info.get("customerPriceStatus", "")
        merged["customerPriceSource"] = info.get("customerPriceSource", "")
        merged["customerPriceWarning"] = info.get("customerPriceWarning", "")
        merged["customerPriceEvidence"] = info.get("customerPriceEvidence", "")

        if audit_result is not None:
            merged["json_price"] = audit_result.json_price
            merged["title_match_result"] = audit_result.title_match_result
            merged["title_match_confidence"] = audit_result.title_match_confidence
            merged["title_match_reason"] = _collect_title_reason(audit_result)
            merged["sku_link_check"] = audit_result.sku_link_check
            merged["price_check"] = audit_result.price_check
            merged["price_vs_customer_check"] = audit_result.price_vs_customer_check
            merged["primary_action"] = audit_result.primary_action
            merged["suggested_actions"] = audit_result.suggested_actions
            merged["reason"] = audit_result.reason
            merged["evidence"] = audit_result.evidence
            merged["needs_human_approval"] = audit_result.needs_human_approval
            merged["execution_allowed"] = audit_result.execution_allowed
            merged["executed"] = audit_result.executed
        else:
            for column in (
                "json_price",
                "title_match_result",
                "title_match_confidence",
                "title_match_reason",
                "sku_link_check",
                "price_check",
                "price_vs_customer_check",
                "primary_action",
                "suggested_actions",
                "reason",
                "evidence",
                "needs_human_approval",
                "execution_allowed",
                "executed",
            ):
                merged[column] = ""

        page_status_info = _compute_page_price_status(row, audit_result)
        merged["page_price_value"] = page_status_info["page_price_value"]
        merged["page_price_status"] = page_status_info["page_price_status"]
        merged["page_price_source"] = page_status_info["page_price_source"]

        page_vs_check, over_30 = compute_over_30_check(
            customer_price=info.get("customerPrice"),
            page_price=page_status_info["page_price_value"],
            page_price_status=page_status_info["page_price_status"],
        )
        merged["page_vs_customer_check"] = page_vs_check
        merged["over_30pct"] = over_30

        # SaaS 价 vs 客户价预检（爬虫前就能算）
        saas_price_for_check = row.get("price")
        if not isinstance(saas_price_for_check, Decimal):
            saas_price_for_check = parse_decimal(saas_price_for_check)
        saas_vs_cust_status, _saas_vs_cust_evidence = compute_saas_price_vs_customer(
            customer_price=info.get("customerPrice"),
            saas_price=saas_price_for_check,
        )
        merged["saas_price_vs_customer"] = saas_vs_cust_status

        # 价格变动（SaaS json 里的价 vs 爬虫拿到的页面价）
        price_changed_status, _pc_evidence = compute_price_changed(
            saas_price=saas_price_for_check,
            page_price=page_status_info["page_price_value"],
            page_price_status=page_status_info["page_price_status"],
        )
        merged["price_changed"] = price_changed_status

        # 综合最终决策
        merged["final_decision"] = compute_final_decision(
            saas_vs_customer_status=saas_vs_cust_status,
            price_changed_status=price_changed_status,
            over_30pct_value=over_30,
        )

        row_is_over = (
            over_30 == "TRUE"
            or saas_vs_cust_status == "SAAS_HIGHER_THAN_CUSTOMER_30"
            or price_changed_status == "CHANGED"
            or page_status_info["page_price_status"] == "PRICE_IS_ZERO"
        )
        for column_index, header in enumerate(final_headers, start=1):
            cell = sheet.cell(row=row_index, column=column_index)
            value = merged.get(header, "")
            # Prices are numeric workbook evidence, not display text.  Keep
            # Decimal precision for all calculations above, then hand Excel a
            # number at the final serialization boundary.
            if (
                header == "customerPrice"
                and isinstance(value, Decimal)
                and info.get("__numeric_customer_price") is True
            ):
                value = float(value)
            cell.value = _format_cell_value(value)
            if row_is_over:
                cell.fill = OVER_30_FILL
            elif header in HIGHLIGHTED_RESULT_COLUMNS:
                cell.fill = HIGHLIGHT_FILL

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(final_headers))}{max(len(rows) + 1, 1)}"
    _set_column_widths(sheet, final_headers)

def write_audit_workbook(
    output_path: Path,
    headers: list[str],
    rows: list[dict[str, Any]],
    enriched: list[dict[str, Any]],
    audit_response: Any,
) -> None:
    if output_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"输出文件必须是 xlsx：{output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary_counts = _build_summary_counts(audit_response, enriched)
    _write_summary_sheet(workbook, summary_counts)
    _write_record_results_sheet(workbook, headers, rows, enriched, audit_response)
    workbook.save(str(output_path))
