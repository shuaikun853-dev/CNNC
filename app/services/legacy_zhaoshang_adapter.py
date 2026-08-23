"""Narrow create-only / wait-export-only seam for the legacy Zhaoshang tool."""

from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import re
import subprocess
import sys
import time
import uuid
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass, replace
from datetime import datetime, timedelta, timezone
from enum import Enum
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping
from urllib.parse import parse_qsl, urlparse

from openpyxl import Workbook, load_workbook

try:
    from pymysql.constants import CR as PyMySQLClientError
    from pymysql.constants import ER as PyMySQLServerError
    from pymysql.err import InterfaceError as PyMySQLInterfaceError
    from pymysql.err import OperationalError as PyMySQLOperationalError
except ImportError:  # The legacy priority backend is optional outside production.
    _PYMYSQL_INTERFACE_PRIORITY_ERRORS: tuple[type[BaseException], ...] = ()
    _PYMYSQL_OPERATIONAL_PRIORITY_ERRORS: tuple[type[BaseException], ...] = ()
    _RETRYABLE_PYMYSQL_OPERATIONAL_CODES: frozenset[int] = frozenset()
else:
    _PYMYSQL_INTERFACE_PRIORITY_ERRORS = (PyMySQLInterfaceError,)
    _PYMYSQL_OPERATIONAL_PRIORITY_ERRORS = (PyMySQLOperationalError,)
    _RETRYABLE_PYMYSQL_OPERATIONAL_CODES = frozenset(
        {
            PyMySQLServerError.CON_COUNT_ERROR,
            PyMySQLServerError.SERVER_SHUTDOWN,
            PyMySQLServerError.TOO_MANY_USER_CONNECTIONS,
            PyMySQLServerError.LOCK_WAIT_TIMEOUT,
            PyMySQLServerError.LOCK_TABLE_FULL,
            PyMySQLServerError.LOCK_DEADLOCK,
            PyMySQLClientError.CR_CONNECTION_ERROR,
            PyMySQLClientError.CR_CONN_HOST_ERROR,
            PyMySQLClientError.CR_IPSOCK_ERROR,
            PyMySQLClientError.CR_SERVER_GONE_ERROR,
            PyMySQLClientError.CR_SERVER_LOST,
            PyMySQLClientError.CR_SERVER_LOST_EXTENDED,
        }
    )

from app.services.zhaoshang_result_adapter import (
    read_zhaoshang_result_xlsx,
    read_zhaoshang_result_xlsx_by_identity,
)
from app.services.saas_audit_core import parse_decimal
from scripts.auth_manager import AuthManager
from scripts.curl_utils import internal_subprocess_env
from scripts.monitor_stb_fallback import load_zhaoshang_db_env
from scripts.zhaoshang_mall_mapping import (
    ZhaoshangMallMapping,
    load_zhaoshang_mall_mappings,
)


EXTERNAL_TOOL_PATH = Path(
    r"E:\.A2026\招商局自动建任务\external_tools\crawl_task_tool\auto_crawl_task.py"
)
DEFAULT_DB_ENV_PATH = Path(__file__).resolve().parents[2] / "data" / "ZHAOSHANG_DB.env"
TEMPLATE_HEADERS = ("商品编码", "url")
RESULT_URL_HEADERS = frozenset({"原始链接", "商品详情地址", "url", "URL", "商品链接", "链接", "商品URL"})
RESULT_CODE_HEADERS = frozenset({
    "商品编码",
    "商品最小品类单元Id",
    "商品最小品类单元id",
    "skuId",
    "sku_id",
    "code",
})
RESULT_EXECUTION_HEADERS = frozenset({"执行编号", "executionNumber", "execution_number"})
RESULT_BATCH_HEADERS = frozenset({"批次号", "batchNumber", "batch_number", "BATCH_NUMBER"})
RESULT_MALL_HEADERS = frozenset({"商城标识", "mall", "mallLogo", "mall_key", "MALL_LOGO"})
RETRYABLE_PRIORITY_ERROR_DIAGNOSTICS = frozenset(
    {"TimeoutError", "ConnectionError", "InterfaceError"}
    | {
        f"OperationalError[{code}]"
        for code in _RETRYABLE_PYMYSQL_OPERATIONAL_CODES
    }
)


def retryable_priority_error_diagnostic(exc: BaseException) -> str | None:
    """Classify only explicit transient types used by the priority backend."""
    if isinstance(exc, TimeoutError):
        return "TimeoutError"
    if isinstance(exc, ConnectionError):
        return "ConnectionError"
    if _PYMYSQL_INTERFACE_PRIORITY_ERRORS and isinstance(
        exc, _PYMYSQL_INTERFACE_PRIORITY_ERRORS
    ):
        return "InterfaceError"
    if _PYMYSQL_OPERATIONAL_PRIORITY_ERRORS and isinstance(
        exc, _PYMYSQL_OPERATIONAL_PRIORITY_ERRORS
    ):
        code = exc.args[0] if exc.args else None
        if type(code) is int and code in _RETRYABLE_PYMYSQL_OPERATIONAL_CODES:
            return f"OperationalError[{code}]"
    return None


def is_retryable_priority_error_diagnostic(value: object) -> bool:
    """Validate a previously persisted normalized priority diagnostic."""
    return str(value or "").strip() in RETRYABLE_PRIORITY_ERROR_DIAGNOSTICS


# The legacy priority core raises PriorityUpdateZeroRowsError from both
# lookup-exhaustion sites ("rows not materialized yet") and several fail-closed
# verification sites (lock-set drift, UPDATE rowcount, post-commit readback).
# Only the two anchored lookup messages below prove REMOTE_NOT_MATERIALIZED;
# every other raise keeps the fail-closed manual path.
_REMOTE_NOT_MATERIALIZED_TYPE_NAMES = frozenset(
    {
        "PriorityUpdateZeroRowsError",
        "PriorityRowsNotMaterializedError",
        "PriorityRowsIncompleteError",
    }
)
REMOTE_NOT_MATERIALIZED_DIAGNOSTICS = frozenset(
    {"PriorityRowsNotMaterialized", "PriorityRowsIncomplete"}
)
# The pre-classification production build persisted only the bare exception
# type name for every PriorityUpdateZeroRowsError site, so that legacy string
# cannot by itself distinguish zero-rows from verification failures. It is
# admitted only by the zero-rows recovery, which re-derives the classification
# from a live probe before any state transition.
REMOTE_NOT_MATERIALIZED_LEGACY_DIAGNOSTIC = "PriorityUpdateZeroRowsError"


def remote_not_materialized_diagnostic(exc: BaseException) -> str | None:
    """Classify only bounded-lookup exhaustion for the exact task identity.

    Zero rows after the bounded retry loop and an incomplete-but-valid row set
    both mean the remote scheduler has not finished materializing this exact
    taskNumber/mall; identity or UPDATE verification anomalies are intentionally
    excluded and stay fail-closed manual.
    """
    if type(exc).__name__ not in _REMOTE_NOT_MATERIALIZED_TYPE_NAMES:
        return None
    message = str(exc).strip()
    if message.startswith("no matching rows for mall/task_number/create_time window"):
        return "PriorityRowsNotMaterialized"
    if message.startswith("incomplete subtask rows: matched="):
        return "PriorityRowsIncomplete"
    return None


def _deterministic_wait_export_startup_failure(exc: BaseException) -> bool:
    """Deterministically missing local prerequisites stay manual.

    A missing external tool or a missing required file cannot heal on a later
    wait-export pass.  Any other startup failure (transient auth refresh,
    pipe/spawn hiccup, exhausted command deadline) only proves this one
    execution failed, so the exact remote identity keeps its bounded retry.
    """
    if isinstance(exc, FileNotFoundError):
        return True
    message = str(exc)
    if message == "legacy Zhaoshang external tool is unavailable":
        return True
    return message.endswith("FileNotFoundError")


class SubmitStatus(str, Enum):
    SUBMITTED = "SUBMITTED"
    SUBMIT_UNKNOWN = "SUBMIT_UNKNOWN"
    NOT_ATTEMPTED = "NOT_ATTEMPTED"


class HarvestStatus(str, Enum):
    EXPORTED = "EXPORTED"
    RUNNING = "RUNNING"
    MANUAL = "MANUAL"
    FAILED = "FAILED"


class ReconcileStatus(str, Enum):
    FOUND = "FOUND"
    NOT_FOUND = "NOT_FOUND"
    MANUAL = "MANUAL"


class PriorityStatus(str, Enum):
    UPDATED = "PRIORITY_UPDATED"
    ALREADY_HIGHEST = "PRIORITY_ALREADY_HIGHEST"
    UNCONFIRMED = "PRIORITY_UNCONFIRMED"


class LegacyZhaoshangAdapterError(RuntimeError):
    pass


class TemplateConflictError(LegacyZhaoshangAdapterError):
    pass


class PriorityEvidenceConflictError(LegacyZhaoshangAdapterError):
    pass


class SubmissionEvidenceConflictError(LegacyZhaoshangAdapterError):
    pass


class LegacyPreSendFailure(LegacyZhaoshangAdapterError):
    """The legacy subprocess provably never started; no remote POST occurred."""


@dataclass(frozen=True)
class CrawlRequest:
    mall_key: str
    mall_source_id: str
    mall_name: str
    sku_id: str
    url: str
    backend_version: str

    @property
    def identity(self) -> tuple[str, str, str, str, str]:
        return (self.mall_key, self.mall_source_id, self.sku_id, self.url.strip(), self.backend_version)


@dataclass(frozen=True)
class PreparedLegacyJob:
    business_date: str
    mall_key: str
    mall_source_id: str
    mall_name: str
    task_name: str
    template_path: Path
    result_path: Path
    batch_state_path: Path
    requests: tuple[CrawlRequest, ...]
    logical_job_key: str
    payload_sha256: str


@dataclass(frozen=True)
class LegacySubmitOutcome:
    status: SubmitStatus
    task_id: str | None = None
    task_number: str | None = None
    priority_evidence: str | None = None
    diagnostic: str = ""


@dataclass(frozen=True)
class LegacyHarvestOutcome:
    status: HarvestStatus
    result_map: Mapping[str, Mapping[str, Any]] | None = None
    # URL-only maps remain for legacy callers. Incremental consumers must use
    # this source identity to avoid merging two SKUs with one exact URL.
    result_by_identity: Mapping[tuple[str, str], Mapping[str, Any]] | None = None
    reused_existing: bool = False
    diagnostic: str = ""
    validation_summary: Mapping[str, int] | None = None
    no_result_identities: frozenset[tuple[str, str]] = frozenset()
    manual_identities: Mapping[tuple[str, str], str] | None = None


@dataclass(frozen=True)
class PartialExactHarvestOutcome:
    exact_matches: Mapping[tuple[str, str], Mapping[str, Any]]
    exact_absent: frozenset[tuple[str, str]]
    ambiguous: Mapping[tuple[str, str], str]
    validation_summary: Mapping[str, int]


@dataclass(frozen=True)
class LegacyReconcileOutcome:
    status: ReconcileStatus
    task_id: str | None = None
    task_number: str | None = None
    diagnostic: str = ""


@dataclass(frozen=True)
class LegacyPriorityOutcome:
    status: PriorityStatus
    weight_type: int | None = None
    diagnostic: str = ""
    retryable: bool = False


@dataclass(frozen=True)
class ExportNoPriceProbe:
    """Read-only live evidence that one finished task exports no price rows."""

    task_status_rsp_desc: str
    progress: Mapping[str, int]
    export_rsp_desc: str
    export_count: int
    remote_calls: int


class _BoundedReadOnlySession:
    """Clamp legacy helper HTTP calls to the tick's remaining reconcile budget."""

    def __init__(
        self,
        session: Any,
        timeout_seconds: float,
        *,
        monotonic: Callable[[], float] = time.monotonic,
        deadline: float | None = None,
    ) -> None:
        self._session = session
        self._monotonic = monotonic
        self._deadline = (
            float(deadline)
            if deadline is not None
            else self._monotonic() + float(timeout_seconds)
        )

    def _bounded_timeout(self, requested: Any) -> float:
        remaining = self._deadline - self._monotonic()
        if remaining <= 0:
            raise TimeoutError("legacy read-only reconciliation deadline is exhausted")
        try:
            requested_timeout = float(requested)
        except (TypeError, ValueError):
            requested_timeout = remaining
        return min(remaining, max(requested_timeout, 0.001))

    def post(self, *args: Any, **kwargs: Any) -> Any:
        kwargs["timeout"] = self._bounded_timeout(kwargs.get("timeout"))
        return self._session.post(*args, **kwargs)

    def get(self, *args: Any, **kwargs: Any) -> Any:
        kwargs["timeout"] = self._bounded_timeout(kwargs.get("timeout"))
        return self._session.get(*args, **kwargs)

    def __getattr__(self, name: str) -> Any:
        return getattr(self._session, name)


def _text(value: object) -> str:
    return str(value or "").strip()


def _canonical_hash(value: object) -> str:
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


@contextmanager
def _temporary_environment(values: Mapping[str, str]):
    """Expose state-backed auth only while constructing the legacy session."""
    previous = {key: os.environ.get(key) for key in values}
    try:
        os.environ.update({key: str(value) for key, value in values.items()})
        yield
    finally:
        for key, value in previous.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value


def _host_matches(host: str, domains: Iterable[str]) -> bool:
    host = _text(host).lower().lstrip(".")
    return bool(host) and any(
        host == str(domain).lower().lstrip(".")
        or host.endswith("." + str(domain).lower().lstrip("."))
        for domain in domains
        if str(domain).strip()
    )


class CnncLegacyZhaoshangAdapter:
    def __init__(
        self,
        *,
        mappings: Iterable[ZhaoshangMallMapping] | None = None,
        runner: Callable[[list[str], float, Mapping[str, str]], Any] | None = None,
        reconciler: Callable[[str, str, float], Mapping[str, Any] | None] | None = None,
        priority_runner: Callable[[PreparedLegacyJob, str, str, float], Any] | None = None,
        export_probe: Callable[[PreparedLegacyJob, str, str, float], Any] | None = None,
        event_recorder: Callable[..., Any] | None = None,
        auth_manager: AuthManager | None = None,
        external_tool_path: Path = EXTERNAL_TOOL_PATH,
        db_env_path: Path | None = DEFAULT_DB_ENV_PATH,
        monotonic: Callable[[], float] = time.monotonic,
    ) -> None:
        self._mappings = tuple(mappings) if mappings is not None else tuple(load_zhaoshang_mall_mappings())
        self._runner = runner
        self._reconciler = reconciler
        self._priority_runner = priority_runner
        self._export_probe = export_probe
        self._event_recorder = event_recorder
        self._auth_manager = auth_manager
        self._external_tool_path = Path(external_tool_path)
        self._db_env_path = Path(db_env_path) if db_env_path is not None else None
        self._monotonic = monotonic
        self._host_mapping_cache: dict[str, ZhaoshangMallMapping | None] = {}

    def _mapping_for_host(self, host: str) -> ZhaoshangMallMapping | None:
        normalized = _text(host).lower().lstrip(".")
        if not normalized:
            return None
        if normalized not in self._host_mapping_cache:
            self._host_mapping_cache[normalized] = next(
                (item for item in self._mappings if _host_matches(normalized, item.domains)),
                None,
            )
        return self._host_mapping_cache[normalized]

    def request_from_saas_row(self, row: Mapping[str, Any], *, backend_version: str) -> CrawlRequest | None:
        url = _text(row.get("extLink"))
        vendor = _text(row.get("vendorName")).lower()
        mapping: ZhaoshangMallMapping | None = None
        host = (urlparse(url).hostname or "").lower().lstrip(".")
        if host:
            mapping = self._mapping_for_host(host)
        elif vendor:
            mapping = next(
                (
                    item
                    for item in self._mappings
                    if any(keyword.lower() in vendor for keyword in item.vendor_keywords)
                ),
                None,
            )
        if mapping is None:
            return None
        return CrawlRequest(
            mall_key=mapping.mall_key,
            mall_source_id=mapping.source_id,
            mall_name=mapping.mall_name,
            sku_id=_text(row.get("skuId")),
            url=url,
            backend_version=_text(backend_version),
        )

    @staticmethod
    def _canonical_requests(requests: Iterable[CrawlRequest]) -> tuple[CrawlRequest, ...]:
        unique: dict[tuple[str, str, str, str, str], CrawlRequest] = {}
        for crawl_request in requests:
            normalized = CrawlRequest(
                mall_key=_text(crawl_request.mall_key),
                mall_source_id=_text(crawl_request.mall_source_id),
                mall_name=_text(crawl_request.mall_name),
                sku_id=_text(crawl_request.sku_id),
                url=_text(crawl_request.url),
                backend_version=_text(crawl_request.backend_version),
            )
            if not normalized.sku_id:
                raise ValueError("Zhaoshang crawl request skuId must not be blank")
            if not all((normalized.mall_key, normalized.mall_source_id, normalized.mall_name, normalized.backend_version)):
                raise ValueError("Zhaoshang crawl request mall identity and backend version are required")
            unique.setdefault(normalized.identity, normalized)
        if not unique:
            raise ValueError("Zhaoshang job must contain at least one crawl request")
        return tuple(unique[key] for key in sorted(unique))

    @staticmethod
    def _expected_template_rows(requests: Iterable[CrawlRequest]) -> list[tuple[str, str]]:
        return [(request.sku_id, request.url) for request in requests]

    @staticmethod
    def _read_template_rows(path: Path) -> tuple[tuple[str, ...], list[tuple[str, str]], str]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet is None:
                return (), [], ""
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = tuple(_text(value) for value in next(rows))
            except StopIteration:
                return (), [], ""
            values = [(_text(row[0]) if len(row) > 0 else "", _text(row[1]) if len(row) > 1 else "") for row in rows]
            return headers, values, _text(workbook.properties.description)
        finally:
            workbook.close()

    @staticmethod
    def _template_manifest(payload_sha256: str) -> str:
        return f"legacy_zhaoshang_payload_sha256={payload_sha256}"

    @classmethod
    def _verify_existing_template(cls, path: Path, requests: tuple[CrawlRequest, ...], payload_sha256: str) -> None:
        try:
            headers, rows, description = cls._read_template_rows(path)
        except Exception as exc:
            # A zero-byte file may be a stale placeholder from the previous
            # implementation, but it carries no proof that it is ours.  Never
            # overwrite it blindly.
            raise TemplateConflictError("existing Zhaoshang template is unreadable and cannot be safely replaced") from exc
        if (
            headers != TEMPLATE_HEADERS
            or rows != cls._expected_template_rows(requests)
            or description != cls._template_manifest(payload_sha256)
        ):
            raise TemplateConflictError("existing Zhaoshang template differs from durable payload evidence")

    def _write_or_verify_template(self, path: Path, requests: tuple[CrawlRequest, ...], payload_sha256: str) -> None:
        expected_rows = self._expected_template_rows(requests)
        if path.exists():
            self._verify_existing_template(path, requests, payload_sha256)
            return

        path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
        try:
            with temp_path.open("xb"):
                pass
            workbook = Workbook()
            try:
                sheet = workbook.active
                sheet.title = "商品编码导入模板"
                sheet.append(list(TEMPLATE_HEADERS))
                for row in expected_rows:
                    sheet.append(list(row))
                workbook.properties.description = self._template_manifest(payload_sha256)
                workbook.save(temp_path)
            finally:
                workbook.close()
            # Workbook.save owns its file handle.  Flush the completed temp
            # file before the atomic publication step.
            with temp_path.open("r+b") as stream:
                os.fsync(stream.fileno())
            if path.exists():
                self._verify_existing_template(path, requests, payload_sha256)
                return
            for attempt in range(5):
                try:
                    os.replace(temp_path, path)
                    break
                except PermissionError:
                    if attempt == 4:
                        raise
                    time.sleep(0.1 * (attempt + 1))
            self._verify_existing_template(path, requests, payload_sha256)
        finally:
            try:
                temp_path.unlink(missing_ok=True)
            except PermissionError:
                # Preserve the original publication error if a Windows scanner
                # still holds the temporary workbook during cleanup.
                pass

    def prepare_job(
        self,
        business_date: str,
        requests: Iterable[CrawlRequest],
        template_path: Path | str,
        result_path: Path | str,
        task_name: str,
        batch_state_path: Path | str,
    ) -> PreparedLegacyJob:
        canonical_requests = self._canonical_requests(requests)
        mall_identities = {(item.mall_key, item.mall_source_id, item.mall_name) for item in canonical_requests}
        if len(mall_identities) != 1:
            raise ValueError("each legacy Zhaoshang job must contain exactly one mall")
        mall_key, mall_source_id, mall_name = next(iter(mall_identities))
        business_date = _text(business_date)
        task_name = _text(task_name)
        if not business_date or not task_name:
            raise ValueError("business_date and task_name are required")
        payload = {
            "business_date": business_date,
            "mall": {"mall_key": mall_key, "mall_source_id": mall_source_id, "mall_name": mall_name},
            "requests": [
                {
                    "mall_key": item.mall_key,
                    "sku_id": item.sku_id,
                    "url": item.url,
                    "backend_version": item.backend_version,
                }
                for item in canonical_requests
            ],
        }
        logical_job_key = _canonical_hash(payload)
        payload_sha256 = _canonical_hash({"headers": TEMPLATE_HEADERS, "payload": payload})
        prepared = PreparedLegacyJob(
            business_date=business_date,
            mall_key=mall_key,
            mall_source_id=mall_source_id,
            mall_name=mall_name,
            task_name=task_name,
            template_path=Path(template_path),
            result_path=Path(result_path),
            batch_state_path=Path(batch_state_path),
            requests=canonical_requests,
            logical_job_key=logical_job_key,
            payload_sha256=payload_sha256,
        )
        if not prepared.batch_state_path.is_file():
            raise ValueError("batch state file must exist before preparing a legacy job")
        self._write_or_verify_template(prepared.template_path, canonical_requests, payload_sha256)
        return prepared

    @staticmethod
    def _read_batch_events(job: PreparedLegacyJob) -> list[dict[str, Any]]:
        try:
            lines = job.batch_state_path.read_text(encoding="utf-8").splitlines()
        except OSError:
            return []
        events: list[dict[str, Any]] = []
        for line in lines:
            stripped = line.strip()
            if not (stripped.startswith("- `") and stripped.endswith("`")):
                continue
            try:
                event = json.loads(stripped[3:-1])
            except (TypeError, ValueError):
                continue
            if isinstance(event, dict) and event.get("task_name") == job.task_name and event.get("mall") == job.mall_key:
                events.append(event)
        return events
    @staticmethod
    def _priority_evidence(
        events: Iterable[Mapping[str, Any]],
    ) -> tuple[str | None, str | None, str | None, bool]:
        selected: tuple[str, str, str, bool] | None = None
        incomplete_evidence: str | None = None
        incomplete_retryable = False
        for event in events:
            action = _text(event.get("action"))
            evidence: str | None = None
            retryable = False
            if action == "PRIORITY_UPDATED" and int(event.get("affected_rows") or 0) > 0:
                evidence = "updated"
            elif action == "PRIORITY_ALREADY_HIGHEST" and int(event.get("already_highest_count") or 0) > 0:
                evidence = "already_highest_priority"
            elif action == "PRIORITY_UNCONFIRMED":
                evidence = "completed_unconfirmed"
                retryable = event.get("retryable") is True
            if evidence is None:
                continue
            task_id = _text(event.get("task_id"))
            task_number = _text(event.get("task_number"))
            if not task_id or not task_number:
                incomplete_evidence = evidence
                incomplete_retryable = retryable
                continue
            candidate = (task_id, task_number, evidence, retryable)
            if selected is not None and selected[:2] != candidate[:2]:
                raise PriorityEvidenceConflictError(
                    "conflicting accepted priority evidence identities"
                )
            selected = candidate
        if selected is not None:
            return selected
        return None, None, incomplete_evidence, incomplete_retryable

    @staticmethod
    def _submission_evidence(
        job: PreparedLegacyJob,
        events: Iterable[Mapping[str, Any]],
    ) -> tuple[str | None, str | None]:
        created: set[tuple[str, str]] = set()
        enabled: set[tuple[str, str]] = set()
        for event in events:
            action = _text(event.get("action"))
            if action not in {"CREATED", "ENABLED"}:
                continue
            task_id = _text(event.get("task_id"))
            task_number = _text(event.get("task_number"))
            if not task_id or not task_number:
                raise SubmissionEvidenceConflictError(f"{action} evidence has incomplete remote identity")
            if action == "CREATED":
                expected_count = event.get("expected_subtask_count")
                if expected_count not in (None, "") and int(expected_count) != len(job.requests):
                    raise SubmissionEvidenceConflictError("CREATED evidence does not match immutable request count")
                created.add((task_id, task_number))
            else:
                enabled.add((task_id, task_number))
        if not created or not enabled:
            return None, None
        if len(created) != 1 or created != enabled:
            raise SubmissionEvidenceConflictError("CREATED and ENABLED evidence has conflicting remote identity")
        return next(iter(created))

    @staticmethod
    def _submission_marker_time(
        events: Iterable[Mapping[str, Any]],
        task_id: str,
        task_number: str,
    ) -> datetime:
        markers: set[datetime] = set()
        takeover_markers: set[datetime] = set()
        for event in events:
            if _text(event.get("action")) != "CREATED":
                continue
            if (_text(event.get("task_id")), _text(event.get("task_number"))) != (task_id, task_number):
                continue
            takeover_timestamp = _text(event.get("submission_marker_time"))
            timestamp = takeover_timestamp or _text(event.get("timestamp"))
            if not timestamp:
                raise SubmissionEvidenceConflictError("CREATED evidence has no durable timestamp")
            try:
                marker = datetime.fromisoformat(timestamp)
            except ValueError as exc:
                raise SubmissionEvidenceConflictError("CREATED evidence timestamp is invalid") from exc
            if marker.tzinfo is not None and marker.utcoffset() is not None:
                marker = marker.astimezone(timezone(timedelta(hours=8))).replace(tzinfo=None)
            (takeover_markers if takeover_timestamp else markers).add(marker)
        selected = takeover_markers or markers
        if len(selected) != 1:
            raise SubmissionEvidenceConflictError("CREATED evidence has conflicting durable timestamps")
        return next(iter(selected))

    @staticmethod
    def _submission_intent_marker(
        job: PreparedLegacyJob,
        events: Iterable[Mapping[str, Any]],
    ) -> str:
        markers: set[str] = set()
        for event in events:
            if _text(event.get("action")) != "ADD_INTENT":
                continue
            expected_count = event.get("expected_subtask_count")
            if expected_count not in (None, "") and int(expected_count) != len(job.requests):
                raise SubmissionEvidenceConflictError(
                    "ADD_INTENT evidence does not match immutable request count"
                )
            timestamp = _text(event.get("timestamp"))
            if not timestamp:
                raise SubmissionEvidenceConflictError(
                    "ADD_INTENT evidence has no durable timestamp"
                )
            try:
                datetime.fromisoformat(timestamp)
            except ValueError as exc:
                raise SubmissionEvidenceConflictError(
                    "ADD_INTENT evidence timestamp is invalid"
                ) from exc
            markers.add(timestamp)
        if len(markers) != 1:
            raise SubmissionEvidenceConflictError(
                "exact takeover requires one durable ADD_INTENT marker"
            )
        return next(iter(markers))

    def _record_batch_event(self, path: Path, **event: Any) -> None:
        if self._event_recorder is not None:
            self._event_recorder(path, **event)
            return
        if not self._external_tool_path.is_file():
            raise FileNotFoundError("legacy Zhaoshang event recorder is unavailable")
        module = self._load_module(
            self._external_tool_path,
            "legacy_zhaoshang_reconcile_events",
        )
        module.record_batch_state_event(path, **event)

    def _validate_submit_inputs(self, job: PreparedLegacyJob, timeout_seconds: float) -> str | None:
        if timeout_seconds <= 0:
            return "submit timeout must be positive"
        if not job.batch_state_path.is_file():
            return "batch state file is missing"
        if not job.template_path.is_file():
            return "template is missing"
        try:
            headers, rows, description = self._read_template_rows(job.template_path)
        except Exception:
            return "template cannot be validated"
        if (
            headers != TEMPLATE_HEADERS
            or rows != self._expected_template_rows(job.requests)
            or description != self._template_manifest(job.payload_sha256)
        ):
            return "template no longer matches durable payload"
        return None

    @staticmethod
    def _create_args(job: PreparedLegacyJob) -> list[str]:
        return [
            sys.executable,
            str(EXTERNAL_TOOL_PATH),
            "--task-name", job.task_name,
            "--excel", str(job.template_path),
            "--mall", job.mall_key,
            "--output", str(job.result_path),
            "--task-weight", "0",
            "--create-only",
            "--batch-state-file", str(job.batch_state_path),
        ]

    @staticmethod
    def _harvest_args(job: PreparedLegacyJob, timeout_seconds: float) -> list[str]:
        return [
            sys.executable,
            str(EXTERNAL_TOOL_PATH),
            "--task-name", job.task_name,
            "--excel", str(job.template_path),
            "--mall", job.mall_key,
            "--output", str(job.result_path),
            "--task-weight", "0",
            "--wait-timeout", str(max(1, int(timeout_seconds))),
            "--wait-export-only",
            "--batch-state-file", str(job.batch_state_path),
        ]

    def _production_runner(self, args: list[str], timeout_seconds: float) -> subprocess.CompletedProcess[str]:
        # Pre-send phase: everything before Popen() is provably local, so any
        # failure here means no child process exists and the remote /add
        # cannot have happened.
        deadline = self._monotonic() + timeout_seconds
        try:
            if not self._external_tool_path.is_file():
                raise LegacyPreSendFailure("legacy Zhaoshang external tool is unavailable")
            auth_manager = self._auth_manager or AuthManager()
            environment = internal_subprocess_env()
            environment.update(auth_manager.export_environment("zhaoshang", refresh=True))
            if self._db_env_path is not None:
                environment.update(load_zhaoshang_db_env(self._db_env_path))
            for key in tuple(environment):
                if key.startswith("TITLE_MATCH_MODEL_"):
                    environment.pop(key, None)
            remaining = deadline - self._monotonic()
            if remaining <= 0:
                raise TimeoutError("legacy Zhaoshang auth refresh exhausted the command deadline")
            args = [str(self._external_tool_path) if value == str(EXTERNAL_TOOL_PATH) else value for value in args]
            try:
                process = subprocess.Popen(
                    args,
                    cwd=str(self._external_tool_path.parent),
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                    env=environment,
                    shell=False,
                )
            except OSError as exc:
                # Popen only raises here when the child could not be created.
                raise LegacyPreSendFailure(f"legacy Zhaoshang subprocess could not be started: {exc}") from exc
        except LegacyPreSendFailure:
            raise
        except Exception as exc:
            raise LegacyPreSendFailure(f"legacy Zhaoshang pre-send setup failed: {type(exc).__name__}") from exc
        # Post-send boundary: the child exists, so any failure below may
        # already have dispatched the remote create and must stay SUBMIT_UNKNOWN.
        try:
            stdout, stderr = process.communicate(timeout=remaining)
        except subprocess.TimeoutExpired as exc:
            process.kill()
            exc.stdout, exc.stderr = process.communicate()
            raise
        except BaseException:
            process.kill()
            process.wait()
            raise
        return subprocess.CompletedProcess(args, process.poll(), stdout, stderr)

    def _run(self, args: list[str], timeout_seconds: float) -> Any:
        if self._runner is not None:
            return self._runner(args, timeout_seconds, {})
        return self._production_runner(args, timeout_seconds)

    @staticmethod
    def _load_module(path: Path, name: str) -> Any:
        spec = importlib.util.spec_from_file_location(name, path)
        if spec is None or spec.loader is None:
            raise RuntimeError(f"cannot load legacy Zhaoshang module: {path.name}")
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        try:
            spec.loader.exec_module(module)
        except Exception:
            sys.modules.pop(name, None)
            raise
        return module

    def _production_priority_runner(
        self,
        job: PreparedLegacyJob,
        task_id: str,
        task_number: str,
        timeout_seconds: float,
    ) -> Mapping[str, Any]:
        if timeout_seconds <= 0:
            raise TimeoutError("legacy priority deadline is exhausted")
        if not self._external_tool_path.is_file():
            raise FileNotFoundError("legacy Zhaoshang external tool is unavailable")
        if self._db_env_path is None:
            raise RuntimeError("legacy Zhaoshang priority database environment is unavailable")
        auto_module = self._load_module(self._external_tool_path, "legacy_zhaoshang_priority_events")
        try:
            priority_module = self._load_module(
                self._external_tool_path.with_name("zhaoshang_priority_boost.py"),
                "legacy_zhaoshang_priority_core",
            )
            db_environment = load_zhaoshang_db_env(self._db_env_path)
            with _temporary_environment(db_environment):
                marker_time = self._submission_marker_time(
                    self._read_batch_events(job),
                    task_id,
                    task_number,
                )
                # A re-check must find the same already-created task, not guess
                # by widening the identity: task_num and mall stay exact and
                # window_start keeps guarding against earlier tasks. Only the
                # upper bound stretches to now so a subtask materialized after
                # marker+15min is still recoverable on a later probe.
                window_end = max(
                    marker_time + timedelta(minutes=15),
                    datetime.now() + timedelta(seconds=30),
                )
                target = priority_module.PriorityTarget(
                    task_name=job.task_name,
                    task_number=task_number,
                    mall_code=auto_module.priority_mall_code_for_key(job.mall_key),
                    marker_time=marker_time,
                    window_start=marker_time - timedelta(minutes=2),
                    window_end=window_end,
                )
                connection = priority_module.connect_from_env()
                read_timeout_seconds = max(30, min(300, int(timeout_seconds)))
                if not hasattr(connection, "_read_timeout"):
                    raise RuntimeError("legacy priority connection has no bounded read timeout")
                connection._read_timeout = read_timeout_seconds
                try:
                    result = priority_module.boost_zhaoshang_task_priority(
                        connection,
                        target,
                        expected_subtask_count=len(job.requests),
                        max_select_retries=max(1, min(60, int(timeout_seconds / 2) or 1)),
                        retry_sleep_seconds=min(2.0, max(0.0, timeout_seconds / 60)),
                    )
                finally:
                    if hasattr(connection, "close"):
                        connection.close()
        except Exception as exc:
            retryable_diagnostic = retryable_priority_error_diagnostic(exc)
            stall_diagnostic = remote_not_materialized_diagnostic(exc)
            try:
                auto_module.record_batch_state_event(
                    job.batch_state_path,
                    mall_key=job.mall_key,
                    task_name=job.task_name,
                    action="PRIORITY_UNCONFIRMED",
                    task_id=task_id,
                    task_number=task_number,
                    reason=retryable_diagnostic or stall_diagnostic or type(exc).__name__,
                    retryable=retryable_diagnostic is not None or stall_diagnostic is not None,
                    remote_stall=stall_diagnostic,
                )
            except Exception:
                pass
            raise

        priority_state = _text(result.get("priority_state"))
        details: dict[str, Any] = {}
        if priority_state == "updated" and int(result.get("affected_rows") or 0) > 0:
            action = "PRIORITY_UPDATED"
            details["affected_rows"] = int(result["affected_rows"])
        elif priority_state == "already_highest_priority" and int(result.get("already_highest_count") or 0) > 0:
            action = "PRIORITY_ALREADY_HIGHEST"
            details["already_highest_count"] = int(result["already_highest_count"])
        else:
            raise RuntimeError("legacy priority core returned no accepted verification evidence")
        auto_module.record_batch_state_event(
            job.batch_state_path,
            mall_key=job.mall_key,
            task_name=job.task_name,
            action=action,
            task_id=task_id,
            task_number=task_number,
            **details,
        )
        return result

    def submit(self, job: PreparedLegacyJob, *, timeout_seconds: float) -> LegacySubmitOutcome:
        invalid_reason = self._validate_submit_inputs(job, timeout_seconds)
        if invalid_reason:
            return LegacySubmitOutcome(SubmitStatus.NOT_ATTEMPTED, diagnostic=invalid_reason)
        process = None
        run_diagnostic = ""
        pre_send_diagnostic = ""
        try:
            process = self._run(self._create_args(job), timeout_seconds)
        except LegacyPreSendFailure as exc:
            pre_send_diagnostic = str(exc)
        except (subprocess.TimeoutExpired, TimeoutError, OSError):
            run_diagnostic = "legacy create-only completion is unknown"
        events = self._read_batch_events(job)
        try:
            task_id, task_number = self._submission_evidence(job, events)
        except SubmissionEvidenceConflictError as exc:
            return LegacySubmitOutcome(SubmitStatus.SUBMIT_UNKNOWN, diagnostic=str(exc))
        if task_id and task_number:
            priority_evidence = None
            diagnostic = run_diagnostic
            try:
                (
                    priority_task_id,
                    priority_task_number,
                    priority_evidence,
                    _priority_retryable,
                ) = self._priority_evidence(events)
            except PriorityEvidenceConflictError as exc:
                diagnostic = str(exc)
            if priority_evidence is not None and (priority_task_id, priority_task_number) != (task_id, task_number):
                priority_evidence = None
                diagnostic = "priority evidence conflicts with confirmed remote identity"
            return LegacySubmitOutcome(
                SubmitStatus.SUBMITTED,
                task_id,
                task_number,
                priority_evidence,
                diagnostic,
            )
        if pre_send_diagnostic:
            # Durable CREATED+ENABLED evidence already won above. With no such
            # evidence and a provably un-started subprocess, no POST began.
            return LegacySubmitOutcome(SubmitStatus.NOT_ATTEMPTED, diagnostic=pre_send_diagnostic)
        if run_diagnostic:
            return LegacySubmitOutcome(SubmitStatus.SUBMIT_UNKNOWN, diagnostic=run_diagnostic)
        if int(getattr(process, "returncode", -1)) != 0:
            return LegacySubmitOutcome(SubmitStatus.SUBMIT_UNKNOWN, diagnostic="legacy create-only returned a nonzero exit")
        return LegacySubmitOutcome(SubmitStatus.SUBMIT_UNKNOWN, diagnostic="legacy create-only lacks CREATED + ENABLED evidence")

    @staticmethod
    def _identity_from_remote(remote: Mapping[str, Any]) -> tuple[str | None, str | None]:
        task_id = _text(remote.get("taskId") or remote.get("id") or remote.get("task_id")) or None
        task_number = _text(remote.get("taskNumber") or remote.get("taskNo") or remote.get("task_number")) or None
        return task_id, task_number

    @staticmethod
    def _one_reconciled_remote(remote: Any) -> Mapping[str, Any] | None:
        if remote is None:
            return None
        if isinstance(remote, Mapping):
            return remote
        if isinstance(remote, (list, tuple)):
            if not remote:
                return None
            if len(remote) != 1 or not isinstance(remote[0], Mapping):
                raise SubmissionEvidenceConflictError(
                    "exact task reconciliation returned multiple or malformed matches"
                )
            return remote[0]
        raise SubmissionEvidenceConflictError(
            "exact task reconciliation returned an invalid result shape"
        )

    @staticmethod
    def _remote_takeover_conflict(
        job: PreparedLegacyJob,
        remote: Mapping[str, Any],
    ) -> str:
        remote_task_name = _text(
            remote.get("taskName") or remote.get("task_name")
        )
        if remote_task_name and remote_task_name != job.task_name:
            return "exact task reconciliation returned a different task name"
        remote_mall = _text(
            remote.get("mallLogo")
            or remote.get("identification")
            or remote.get("sourceIdentification")
            or remote.get("mall_key")
        )
        if remote_mall and remote_mall != job.mall_key:
            return "exact task reconciliation returned a different mall"
        for field in (
            "expectedSubtaskCount",
            "expected_subtask_count",
            "subTaskCount",
            "subtaskCount",
            "taskCount",
        ):
            value = remote.get(field)
            if value in (None, ""):
                continue
            try:
                count = int(value)
            except (TypeError, ValueError):
                return "exact task reconciliation returned an invalid subtask count"
            if count != len(job.requests):
                return "exact task reconciliation returned a different subtask count"
            break
        return ""

    @staticmethod
    def _remote_is_enabled(remote: Mapping[str, Any]) -> bool:
        enabled = remote.get("enabled")
        if isinstance(enabled, bool):
            return enabled
        state = _text(
            remote.get("state")
            or remote.get("taskState")
            or remote.get("task_state")
            or remote.get("enableState")
        ).lower()
        return state in {"1", "enabled", "enable", "running", "active"}

    def _default_reconciler(self, task_name: str, mall_key: str, timeout_seconds: float) -> Mapping[str, Any] | None:
        if timeout_seconds <= 0:
            raise ValueError("reconciliation timeout must be positive")
        deadline = self._monotonic() + timeout_seconds
        if not self._external_tool_path.is_file():
            raise FileNotFoundError("legacy Zhaoshang external tool is unavailable")
        spec = importlib.util.spec_from_file_location("legacy_zhaoshang_read_only", self._external_tool_path)
        if spec is None or spec.loader is None:
            raise RuntimeError("cannot load legacy Zhaoshang reconciliation helpers")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        auth_manager = self._auth_manager or AuthManager()
        auth_environment = auth_manager.export_environment("zhaoshang", refresh=True)
        remaining = deadline - self._monotonic()
        if remaining <= 0:
            raise TimeoutError("legacy reconciliation auth refresh exhausted the deadline")
        with _temporary_environment(auth_environment):
            session = _BoundedReadOnlySession(
                module.build_session(),
                remaining,
                monotonic=self._monotonic,
                deadline=deadline,
            )
            mall_code = module.priority_mall_code_for_key(mall_key)
            try:
                return module.check_existing_task(session, task_name, mall_code)
            except RuntimeError as exc:
                raise SubmissionEvidenceConflictError(str(exc)) from exc

    def reconcile_unknown(self, job: PreparedLegacyJob, *, timeout_seconds: float) -> LegacyReconcileOutcome:
        if timeout_seconds <= 0:
            return LegacyReconcileOutcome(ReconcileStatus.MANUAL, diagnostic="reconciliation timeout must be positive")
        reconciler = self._reconciler or self._default_reconciler
        try:
            remote = self._one_reconciled_remote(
                reconciler(job.task_name, job.mall_key, timeout_seconds)
            )
        except SubmissionEvidenceConflictError as exc:
            return LegacyReconcileOutcome(ReconcileStatus.MANUAL, diagnostic=str(exc))
        except (FileNotFoundError, ImportError, AttributeError):
            return LegacyReconcileOutcome(ReconcileStatus.MANUAL, diagnostic="exact task reconciliation tooling is unavailable")
        except Exception:
            # A read timeout, authentication refresh race, or other transient
            # failure proves neither absence nor existence. Keep the durable
            # submit unknown until the outer harvest attempts/elapsed cap.
            return LegacyReconcileOutcome(ReconcileStatus.NOT_FOUND, diagnostic="exact task reconciliation is not yet conclusive")
        if remote is None:
            return LegacyReconcileOutcome(ReconcileStatus.NOT_FOUND)
        conflict = self._remote_takeover_conflict(job, remote)
        if conflict:
            return LegacyReconcileOutcome(ReconcileStatus.MANUAL, diagnostic=conflict)
        task_id, task_number = self._identity_from_remote(remote)
        if not task_id or not task_number:
            return LegacyReconcileOutcome(ReconcileStatus.MANUAL, diagnostic="exact task reconciliation returned incomplete identity")
        try:
            events = self._read_batch_events(job)
            event_task_id, event_task_number = self._submission_evidence(job, events)
        except SubmissionEvidenceConflictError as exc:
            return LegacyReconcileOutcome(ReconcileStatus.MANUAL, diagnostic=str(exc))
        if event_task_id and event_task_number:
            if (event_task_id, event_task_number) != (task_id, task_number):
                return LegacyReconcileOutcome(ReconcileStatus.MANUAL, diagnostic="submission evidence conflicts with exact remote identity")
            return LegacyReconcileOutcome(ReconcileStatus.FOUND, task_id, task_number)

        partial_identities = {
            (_text(event.get("task_id")), _text(event.get("task_number")))
            for event in events
            if _text(event.get("action")) in {"CREATED", "ENABLED"}
        }
        partial_identities.discard(("", ""))
        if partial_identities and partial_identities != {(task_id, task_number)}:
            return LegacyReconcileOutcome(ReconcileStatus.MANUAL, diagnostic="submission evidence conflicts with exact remote identity")
        if not self._remote_is_enabled(remote):
            return LegacyReconcileOutcome(
                ReconcileStatus.MANUAL,
                diagnostic="exact remote task is not durably confirmed enabled",
            )
        try:
            marker_time = self._submission_intent_marker(job, events)
            common = {
                "mall_key": job.mall_key,
                "task_name": job.task_name,
                "task_id": task_id,
                "task_number": task_number,
                "submission_marker_time": marker_time,
                "recovered_from": "SUBMIT_UNKNOWN_EXACT_LOOKUP",
            }
            self._record_batch_event(
                job.batch_state_path,
                action="CREATED",
                expected_subtask_count=len(job.requests),
                **common,
            )
            self._record_batch_event(
                job.batch_state_path,
                action="ENABLED",
                **common,
            )
            adopted_task_id, adopted_task_number = self._submission_evidence(
                job,
                self._read_batch_events(job),
            )
        except SubmissionEvidenceConflictError as exc:
            return LegacyReconcileOutcome(ReconcileStatus.MANUAL, diagnostic=str(exc))
        except Exception:
            return LegacyReconcileOutcome(
                ReconcileStatus.NOT_FOUND,
                diagnostic="exact remote task was found but durable takeover evidence could not be written",
            )
        if (adopted_task_id, adopted_task_number) != (task_id, task_number):
            return LegacyReconcileOutcome(
                ReconcileStatus.NOT_FOUND,
                diagnostic="exact remote task was found but durable takeover evidence is incomplete",
            )
        return LegacyReconcileOutcome(ReconcileStatus.FOUND, task_id, task_number)

    def prioritize(
        self,
        job: PreparedLegacyJob,
        *,
        task_id: str,
        task_number: str,
        timeout_seconds: float,
        retry_unconfirmed: bool = False,
    ) -> LegacyPriorityOutcome:
        if timeout_seconds <= 0:
            return LegacyPriorityOutcome(PriorityStatus.UNCONFIRMED, diagnostic="priority timeout must be positive")
        events = self._read_batch_events(job)
        try:
            created_task_id, created_task_number = self._submission_evidence(job, events)
        except SubmissionEvidenceConflictError as exc:
            return LegacyPriorityOutcome(PriorityStatus.UNCONFIRMED, diagnostic=str(exc))
        if (created_task_id, created_task_number) != (_text(task_id), _text(task_number)):
            return LegacyPriorityOutcome(PriorityStatus.UNCONFIRMED, diagnostic="priority identity does not match confirmed submission")

        def existing_outcome() -> LegacyPriorityOutcome | None:
            try:
                (
                    event_task_id,
                    event_task_number,
                    evidence,
                    evidence_retryable,
                ) = self._priority_evidence(self._read_batch_events(job))
            except PriorityEvidenceConflictError as exc:
                return LegacyPriorityOutcome(PriorityStatus.UNCONFIRMED, diagnostic=str(exc))
            if evidence is None:
                return None
            if (event_task_id, event_task_number) != (created_task_id, created_task_number):
                return LegacyPriorityOutcome(PriorityStatus.UNCONFIRMED, diagnostic="priority evidence conflicts with confirmed submission")
            if evidence == "updated":
                return LegacyPriorityOutcome(PriorityStatus.UPDATED, weight_type=-1)
            if evidence == "already_highest_priority":
                return LegacyPriorityOutcome(PriorityStatus.ALREADY_HIGHEST, weight_type=-1)
            if retry_unconfirmed or evidence_retryable:
                return None
            return LegacyPriorityOutcome(PriorityStatus.UNCONFIRMED, diagnostic="highest priority was not confirmed")

        existing = existing_outcome()
        if existing is not None:
            return existing
        priority_runner = self._priority_runner or self._production_priority_runner
        try:
            priority_runner(job, _text(task_id), _text(task_number), timeout_seconds)
        except Exception as exc:
            retryable_diagnostic = retryable_priority_error_diagnostic(exc)
            stall_diagnostic = remote_not_materialized_diagnostic(exc)
            return LegacyPriorityOutcome(
                PriorityStatus.UNCONFIRMED,
                diagnostic=retryable_diagnostic or stall_diagnostic or type(exc).__name__,
                retryable=retryable_diagnostic is not None or stall_diagnostic is not None,
            )
        return existing_outcome() or LegacyPriorityOutcome(
            PriorityStatus.UNCONFIRMED,
            diagnostic="priority runner returned without durable evidence",
        )

    @staticmethod
    def _result_column_index(headers: list[str], allowed: frozenset[str]) -> int:
        return next((index for index, value in enumerate(headers) if value in allowed), -1)

    @staticmethod
    def _result_column_indices(headers: list[str], allowed: frozenset[str]) -> list[int]:
        return [index for index, value in enumerate(headers) if value in allowed]

    @staticmethod
    def _matches_legacy_export_url(durable_url: str, exported_url: str) -> bool:
        """Match exact URLs plus strict, allowlisted legacy export rewrites."""
        if durable_url == exported_url:
            return True
        if not durable_url or not exported_url:
            return False
        durable = urlparse(durable_url)
        exported = urlparse(exported_url)
        durable_query = parse_qsl(durable.query, keep_blank_values=True)
        exported_query = parse_qsl(exported.query, keep_blank_values=True)
        if (
            durable.scheme == exported.scheme == "https"
            and durable.netloc == exported.netloc == "www.qstbg.com"
            and durable.path == "/gd/detail.do"
            and exported.path == "/gd/initParameter.do"
            and not durable.fragment
            and exported.fragment == "title_detail"
            and len({key for key, _value in durable_query}) == len(durable_query)
            and len({key for key, _value in exported_query}) == len(exported_query)
        ):
            durable_params = dict(durable_query)
            exported_params = dict(exported_query)
            goods_id = durable_params.get("goods_id", "")
            if (
                goods_id
                and set(durable_params).issubset({"goods_id", "formid"})
                and durable_params.get("formid", "searchform") == "searchform"
                and exported_params == {"goods_id": goods_id, "formid": "form1"}
            ):
                return True
        jd_hosts = {"i-item.jd.com", "item.jd.com"}
        durable_jd_product = re.fullmatch(r"/([0-9]+)\.html", durable.path)
        exported_jd_product = re.fullmatch(r"/([0-9]+)\.html", exported.path)
        if (
            durable.scheme == exported.scheme == "https"
            and {durable.netloc, exported.netloc} == jd_hosts
            and not durable.query
            and not exported.query
            and not durable.fragment
            and not exported.fragment
            and durable_jd_product is not None
            and exported_jd_product is not None
            and durable_jd_product.group(1) == exported_jd_product.group(1)
        ):
            return True
        if not durable.path.endswith("/detail/inner"):
            return False
        canonical_path = durable.path[: -len("/inner")]
        return exported == durable._replace(path=canonical_path)

    def exact_export_remote_identity(self, job: PreparedLegacyJob) -> str | None:
        """Return the exact exported remote identity, or ``None`` on any conflict."""
        events = self._read_batch_events(job)
        try:
            task_id, task_number = self._submission_evidence(job, events)
        except SubmissionEvidenceConflictError:
            return None
        if not task_id or not task_number:
            return None
        expected_path = os.path.normcase(str(job.result_path.resolve()))
        for event in reversed(events):
            if _text(event.get("action")) != "EXPORTED":
                continue
            if (_text(event.get("task_id")), _text(event.get("task_number"))) != (
                task_id,
                task_number,
            ):
                continue
            raw_path = _text(event.get("output_path"))
            if raw_path and os.path.normcase(str(Path(raw_path).resolve())) == expected_path:
                return f"{job.mall_key}|{task_id}|{task_number}"
        return None

    def exact_priority_remote_identity(self, job: PreparedLegacyJob) -> str | None:
        """Return an exact identity only for accepted durable priority evidence."""
        events = self._read_batch_events(job)
        try:
            submitted = self._submission_evidence(job, events)
            task_id, task_number, evidence, _retryable = self._priority_evidence(events)
        except (PriorityEvidenceConflictError, SubmissionEvidenceConflictError):
            return None
        if (
            not task_id
            or not task_number
            or submitted != (task_id, task_number)
            or evidence not in {"updated", "already_highest_priority"}
        ):
            return None
        return f"{job.mall_key}|{task_id}|{task_number}"

    def probe_export_no_price(
        self,
        job: PreparedLegacyJob,
        task_id: str,
        task_number: str,
        timeout_seconds: float,
    ) -> ExportNoPriceProbe:
        """Re-verify live terminal no-price evidence for one exact existing task.

        Read-only against the crawl platform: the exact existing task is looked
        up by durable name, its identity must match, and the status / progress /
        export endpoints are queried directly.  The export response body is
        inspected but never written anywhere.
        """
        if self._export_probe is not None:
            return self._export_probe(job, task_id, task_number, timeout_seconds)
        return self._production_export_no_price_probe(job, task_id, task_number, timeout_seconds)

    def _production_export_no_price_probe(
        self,
        job: PreparedLegacyJob,
        task_id: str,
        task_number: str,
        timeout_seconds: float,
    ) -> ExportNoPriceProbe:
        if timeout_seconds <= 0:
            raise TimeoutError("export no-price probe timeout must be positive")
        if not self._external_tool_path.is_file():
            raise FileNotFoundError("legacy Zhaoshang external tool is unavailable")
        deadline = self._monotonic() + float(timeout_seconds)
        tool = self._load_module(self._external_tool_path, "legacy_zhaoshang_export_no_price_probe")
        auth_manager = self._auth_manager or AuthManager()
        environment = internal_subprocess_env()
        environment.update(auth_manager.export_environment("zhaoshang", refresh=True))
        remaining = deadline - self._monotonic()
        if remaining <= 0:
            raise TimeoutError("legacy Zhaoshang auth refresh exhausted the probe deadline")
        # The legacy session constructor reads exactly these two auth variables.
        session_environment = {
            key: str(environment[key])
            for key in ("BLADE_AUTH", "COOKIE")
            if environment.get(key)
        }
        if len(session_environment) != 2:
            raise LegacyZhaoshangAdapterError("legacy Zhaoshang session auth is unavailable")
        with _temporary_environment(session_environment):
            session = tool.build_session()
            mall_code = tool.priority_mall_code_for_key(job.mall_key)
            resolved_task_name, _batch_number = tool.build_browser_task_identifiers(
                job.task_name,
                job.mall_key,
            )
            task = tool.check_existing_task(session, resolved_task_name, mall_code)
            remote_calls = 1
            if task is None:
                raise LegacyZhaoshangAdapterError(
                    "export no-price probe found no exact existing task"
                )
            live_task_id, live_task_number, mall_logo = tool.get_task_identity(task, mall_code)
            if str(live_task_id) != str(task_id) or str(live_task_number) != str(task_number):
                raise LegacyZhaoshangAdapterError(
                    "live task identity does not match the durable remote identity"
                )
            remaining = deadline - self._monotonic()
            if remaining <= 0:
                raise TimeoutError("export no-price probe deadline exhausted before status query")
            status = tool.query_task_status(session, str(live_task_id))
            remote_calls += 1
            remaining = deadline - self._monotonic()
            if remaining <= 0:
                raise TimeoutError("export no-price probe deadline exhausted before progress query")
            # The progress endpoint keys on the task record's batchNumber, not
            # the RW- taskNumber; mirror run_auto_task's identifier choice.
            batch_number = _text(task.get("batchNumber")) or resolved_task_name
            progress_json = tool.query_task_progress(session, batch_number)
            remote_calls += 1
            remaining = deadline - self._monotonic()
            if remaining <= 0:
                raise TimeoutError("export no-price probe deadline exhausted before export query")
            response = session.post(
                tool.EXPORT_TASK_URL,
                json={
                    "taskNumber": str(live_task_number),
                    "taskId": str(live_task_id),
                    "mallLogo": mall_logo,
                },
                timeout=min(60.0, remaining),
            )
            response.raise_for_status()
            remote_calls += 1
        content_type = response.headers.get("Content-Type", "") or ""
        content_disposition = response.headers.get("Content-Disposition", "") or ""
        if (
            "spreadsheetml" in content_type
            or "application/vnd.ms-excel" in content_type
            or "attachment" in content_disposition
            or response.content[:2] == b"PK"
        ):
            raise LegacyZhaoshangAdapterError(
                "export endpoint returned a result workbook; terminal no-price resolution does not apply"
            )
        try:
            data = response.json()
        except ValueError as exc:
            raise LegacyZhaoshangAdapterError(
                "export endpoint returned neither a workbook nor JSON"
            ) from exc
        if not isinstance(data, dict):
            raise LegacyZhaoshangAdapterError("export endpoint returned an unexpected payload")
        embedded = data.get("data")
        if isinstance(embedded, dict) and any(
            _text(embedded.get(key)) for key in ("url", "fileUrl", "downloadUrl", "exportUrl")
        ):
            raise LegacyZhaoshangAdapterError(
                "export endpoint returned a result download; terminal no-price resolution does not apply"
            )
        if isinstance(embedded, str) and any(token in embedded for token in (".xlsx", "download", "export")):
            raise LegacyZhaoshangAdapterError(
                "export endpoint returned a result download; terminal no-price resolution does not apply"
            )
        raw_count = data.get("count")
        if raw_count is None:
            # An exact no-price answer may omit the count field; only an
            # empty data payload keeps that shape admissible.
            if embedded not in (None, {}, []):
                raise LegacyZhaoshangAdapterError(
                    "export endpoint returned rows without a count"
                )
            export_count = 0
        else:
            try:
                export_count = int(raw_count)
            except (TypeError, ValueError) as exc:
                raise LegacyZhaoshangAdapterError(
                    "export endpoint returned an unparseable count"
                ) from exc
        return ExportNoPriceProbe(
            task_status_rsp_desc=_text(status.get("rspDesc")),
            progress=dict(tool.parse_task_progress_desc(_text(progress_json.get("rspDesc")))),
            export_rsp_desc=_text(data.get("rspDesc")),
            export_count=export_count,
            remote_calls=remote_calls,
        )

    @classmethod
    def _validate_result_workbook(cls, job: PreparedLegacyJob) -> tuple[
        Mapping[str, Mapping[str, Any]],
        Mapping[tuple[str, str], Mapping[str, Any]],
        Mapping[str, int],
    ]:
        if not job.result_path.is_file():
            raise ValueError("result workbook does not exist")
        workbook = load_workbook(job.result_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet is None:
                raise ValueError("result workbook has no active sheet")
            iterator = sheet.iter_rows(values_only=True)
            try:
                headers = [_text(value) for value in next(iterator)]
            except StopIteration as exc:
                raise ValueError("result workbook is empty") from exc
            url_index = cls._result_column_index(headers, RESULT_URL_HEADERS)
            code_indices = cls._result_column_indices(headers, RESULT_CODE_HEADERS)
            if url_index < 0 and not code_indices:
                raise ValueError("result workbook header is not recognized")
            workbook_rows = [
                (excel_row, row)
                for excel_row, row in enumerate(iterator, start=2)
                if any(_text(value) for value in row)
            ]
        finally:
            workbook.close()
        execution_index = cls._result_column_index(headers, RESULT_EXECUTION_HEADERS)
        batch_index = cls._result_column_index(headers, RESULT_BATCH_HEADERS)
        mall_index = cls._result_column_index(headers, RESULT_MALL_HEADERS)
        provenance_present = min(execution_index, batch_index, mall_index) >= 0
        if not provenance_present:
            raise ValueError("result workbook provenance columns are required")
        foreign_filtered = 0
        current_rows: list[tuple[int, tuple[Any, ...]]] = []
        current_executions: set[str] = set()
        for excel_row, row in workbook_rows:
            execution_number = _text(
                row[execution_index] if execution_index < len(row) else ""
            )
            if not execution_number:
                raise ValueError("result workbook row has no execution provenance")
            batch_number = _text(row[batch_index] if batch_index < len(row) else "")
            if not batch_number:
                raise ValueError("result workbook row has no batch provenance")
            if batch_number != job.task_name:
                foreign_filtered += 1
                continue
            mall_key = _text(row[mall_index] if mall_index < len(row) else "")
            if mall_key != job.mall_key:
                raise ValueError("result workbook current batch mall does not match durable job")
            if execution_number in current_executions:
                raise ValueError("result workbook current batch execution is duplicated")
            current_executions.add(execution_number)
            current_rows.append((excel_row, row))
        workbook_rows = current_rows
        included_excel_rows = {excel_row for excel_row, _row in workbook_rows}
        result_rows = [row for _excel_row, row in workbook_rows]
        if len(result_rows) != len(job.requests):
            raise ValueError("result workbook row count does not match request count")

        expected = Counter((request.sku_id, request.url.strip()) for request in job.requests)
        by_sku: dict[str, set[str]] = {}
        by_url: dict[str, set[str]] = {}
        for sku_id, exact_url in expected:
            by_sku.setdefault(sku_id, set()).add(exact_url)
            by_url.setdefault(exact_url, set()).add(sku_id)

        def durable_identity(sku_id: str, exact_url: str) -> tuple[str, str]:
            identity = (sku_id, exact_url)
            if identity in expected:
                return identity
            conflicting_skus = {
                durable_sku
                for durable_sku, durable_urls in by_sku.items()
                for durable_url in durable_urls
                if durable_sku != sku_id
                and cls._matches_legacy_export_url(durable_url, exact_url)
            }
            if conflicting_skus:
                raise ValueError(
                    "result workbook skuId conflicts with an exact durable URL identity"
                )
            urls = by_sku.get(sku_id, set())
            if sku_id and urls == {""}:
                return (sku_id, "")
            if sku_id and len(urls) == 1:
                durable_url = next(iter(urls))
                if cls._matches_legacy_export_url(durable_url, exact_url):
                    return (sku_id, durable_url)
            raise ValueError(
                "result workbook contains an unexpected skuId / exact URL identity"
            )

        def identities_for(code_index: int) -> list[tuple[str, str]]:
            observed: list[tuple[str, str]] = []
            for row in result_rows:
                sku_id = _text(row[code_index] if 0 <= code_index < len(row) else "")
                exact_url = _text(row[url_index] if 0 <= url_index < len(row) else "")
                if sku_id and exact_url:
                    # Canonicalized URLs remain safe only when the SKU is
                    # unique and the URL does not identify another request.
                    identity = durable_identity(sku_id, exact_url)
                elif sku_id:
                    urls = by_sku.get(sku_id, set())
                    if len(urls) != 1:
                        raise ValueError("result workbook skuId fallback is ambiguous")
                    identity = (sku_id, next(iter(urls)))
                elif exact_url:
                    skus = by_url.get(exact_url, set())
                    if len(skus) != 1:
                        raise ValueError("result workbook exact URL fallback is ambiguous")
                    identity = (next(iter(skus)), exact_url)
                else:
                    raise ValueError("result workbook row has no parseable crawl identity")
                observed.append(identity)
            observed_counter = Counter(observed)
            if observed_counter != expected:
                raise ValueError(
                    "result workbook crawl identities do not match durable request identities"
                )
            return observed

        matching_candidates: list[
            tuple[int, list[tuple[str, str]], tuple[str, ...]]
        ] = []
        for candidate_index in code_indices or [-1]:
            try:
                matching_candidates.append(
                    (
                        candidate_index,
                        identities_for(candidate_index),
                        tuple(
                            _text(
                                row[candidate_index]
                                if 0 <= candidate_index < len(row)
                                else ""
                            )
                            for row in result_rows
                        ),
                    )
                )
            except ValueError:
                continue
        if not matching_candidates:
            raise ValueError("no SKU candidate column matches durable request identities")
        equivalence_classes: dict[
            tuple[str, ...], list[tuple[int, list[tuple[str, str]]]]
        ] = {}
        for candidate_index, identities, value_vector in matching_candidates:
            equivalence_classes.setdefault(value_vector, []).append(
                (candidate_index, identities)
            )
        if len(equivalence_classes) != 1:
            raise ValueError("multiple SKU candidate columns match durable request identities")
        code_index, observed = next(iter(equivalence_classes.values()))[0]
        observed_counter = Counter(observed)

        observed_result_by_identity = read_zhaoshang_result_xlsx_by_identity(
            job.result_path,
            template_path=job.template_path,
            code_index=code_index,
            include_excel_rows=included_excel_rows,
        )
        result_by_identity: dict[tuple[str, str], Mapping[str, Any]] = {}
        for observed_identity, value in observed_result_by_identity.items():
            sku_id, exact_url = observed_identity
            mapped_identity = durable_identity(sku_id, exact_url)
            if mapped_identity in result_by_identity:
                raise ValueError("result workbook result identity is ambiguous")
            result_by_identity[mapped_identity] = value
        if Counter(result_by_identity.keys()) != observed_counter:
            raise ValueError("result workbook has no complete parseable crawl identity mapping")
        result_map = read_zhaoshang_result_xlsx(
            job.result_path,
            template_path=job.template_path,
            code_index=code_index,
            include_excel_rows=included_excel_rows,
        )
        for (_sku_id, durable_url), value in result_by_identity.items():
            if durable_url:
                result_map.setdefault(durable_url, value)
        expected_urls = {url for _sku_id, url in expected if url}
        if expected_urls and (
            not result_map or not expected_urls.issubset(result_map)
        ):
            raise ValueError("result workbook has no complete parseable URL result mapping")
        return result_map, result_by_identity, {
            "expected": len(job.requests),
            "matched": len(result_by_identity),
            "missing": 0,
            "duplicate": 0,
            "unexpected": 0,
            "foreign_filtered": foreign_filtered,
        }

    def harvest_existing(self, job: PreparedLegacyJob) -> LegacyHarvestOutcome:
        """Validate one existing export without making a remote call."""
        if not job.result_path.is_file():
            return LegacyHarvestOutcome(
                HarvestStatus.FAILED,
                diagnostic="existing result workbook is unavailable",
            )
        if self.exact_export_remote_identity(job) is None:
            try:
                result_map, result_by_identity, validation_summary = (
                    self._validate_result_workbook(job)
                )
            except Exception:
                return LegacyHarvestOutcome(
                    HarvestStatus.FAILED,
                    diagnostic="existing result workbook is invalid",
                )
            return LegacyHarvestOutcome(
                HarvestStatus.EXPORTED,
                result_map=result_map,
                result_by_identity=result_by_identity,
                reused_existing=True,
                validation_summary=validation_summary,
            )
        try:
            classification = self.classify_existing_export_partial(job)
        except Exception:
            return LegacyHarvestOutcome(
                HarvestStatus.FAILED,
                diagnostic="existing result workbook is invalid",
            )
        result_by_identity = dict(classification.exact_matches)
        result_map: dict[str, Mapping[str, Any]] = {}
        for (_sku_id, durable_url), result in result_by_identity.items():
            exported_url = _text(result.get("exported_url"))
            if exported_url:
                result_map.setdefault(exported_url, result)
            if durable_url:
                result_map.setdefault(durable_url, result)
        return LegacyHarvestOutcome(
            HarvestStatus.EXPORTED,
            result_map=result_map,
            result_by_identity=result_by_identity,
            reused_existing=True,
            validation_summary=classification.validation_summary,
            no_result_identities=classification.exact_absent,
            manual_identities=classification.ambiguous,
        )

    def validate_existing_export_exact_requests(
        self,
        job: PreparedLegacyJob,
        requests: Sequence[CrawlRequest],
    ) -> Mapping[tuple[str, str], Mapping[str, Any]]:
        """Validate that an existing export contains exactly the supplied requests."""
        selected = tuple(requests)
        if not selected:
            raise ValueError("exact export validation requires at least one request")
        full = Counter(request.identity for request in job.requests)
        subset = Counter(request.identity for request in selected)
        if subset - full:
            raise ValueError("exact export validation requests are outside the durable job")
        scoped_job = replace(job, requests=selected)
        _result_map, result_by_identity, _validation_summary = self._validate_result_workbook(scoped_job)
        return result_by_identity

    @classmethod
    def _classify_existing_export_partial_legacy(
        cls,
        job: PreparedLegacyJob,
    ) -> PartialExactHarvestOutcome:
        """Classify a provenance-valid partial export without weakening identity rules."""
        if not job.result_path.is_file():
            raise ValueError("result workbook does not exist")
        workbook = load_workbook(job.result_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet is None:
                raise ValueError("result workbook has no active sheet")
            iterator = sheet.iter_rows(values_only=True)
            try:
                headers = [_text(value) for value in next(iterator)]
            except StopIteration as exc:
                raise ValueError("result workbook is empty") from exc
            url_index = cls._result_column_index(headers, RESULT_URL_HEADERS)
            code_indices = cls._result_column_indices(headers, RESULT_CODE_HEADERS)
            if url_index < 0 and not code_indices:
                raise ValueError("result workbook header is not recognized")
            workbook_rows = [
                (excel_row, row)
                for excel_row, row in enumerate(iterator, start=2)
                if any(_text(value) for value in row)
            ]
        finally:
            workbook.close()

        execution_index = cls._result_column_index(headers, RESULT_EXECUTION_HEADERS)
        batch_index = cls._result_column_index(headers, RESULT_BATCH_HEADERS)
        mall_index = cls._result_column_index(headers, RESULT_MALL_HEADERS)
        if min(execution_index, batch_index, mall_index) < 0:
            raise ValueError("result workbook provenance columns are required")
        foreign_filtered = 0
        current_rows: list[tuple[int, tuple[Any, ...]]] = []
        current_executions: set[str] = set()
        for excel_row, row in workbook_rows:
            execution_number = _text(
                row[execution_index] if execution_index < len(row) else ""
            )
            batch_number = _text(row[batch_index] if batch_index < len(row) else "")
            if not execution_number:
                raise ValueError("result workbook row has no execution provenance")
            if not batch_number:
                raise ValueError("result workbook row has no batch provenance")
            if batch_number != job.task_name:
                foreign_filtered += 1
                continue
            mall_key = _text(row[mall_index] if mall_index < len(row) else "")
            if mall_key != job.mall_key:
                raise ValueError("result workbook current batch mall does not match durable job")
            if execution_number in current_executions:
                raise ValueError("result workbook current batch execution is duplicated")
            current_executions.add(execution_number)
            current_rows.append((excel_row, row))

        expected_counter = Counter(
            (request.sku_id, request.url.strip()) for request in job.requests
        )
        if not expected_counter or any(count != 1 for count in expected_counter.values()):
            raise ValueError("partial export requires unique durable request identities")
        expected = set(expected_counter)
        by_sku: dict[str, set[str]] = {}
        by_url: dict[str, set[str]] = {}
        for sku_id, exact_url in expected:
            by_sku.setdefault(sku_id, set()).add(exact_url)
            by_url.setdefault(exact_url, set()).add(sku_id)

        def durable_identity(sku_id: str, exported_url: str) -> tuple[str, str]:
            identity = (sku_id, exported_url)
            if identity in expected:
                return identity
            conflicting_skus = {
                durable_sku
                for durable_sku, durable_urls in by_sku.items()
                for durable_url in durable_urls
                if durable_sku != sku_id
                and cls._matches_legacy_export_url(durable_url, exported_url)
            }
            if conflicting_skus:
                raise ValueError(
                    "result workbook skuId conflicts with an exact durable URL identity"
                )
            urls = by_sku.get(sku_id, set())
            if sku_id and urls == {""}:
                return (sku_id, "")
            if sku_id and len(urls) == 1:
                durable_url = next(iter(urls))
                if cls._matches_legacy_export_url(durable_url, exported_url):
                    return (sku_id, durable_url)
            raise ValueError(
                "result workbook contains an unexpected skuId / exact URL identity"
            )

        def identities_for(code_index: int) -> tuple[tuple[str, str], ...]:
            observed: list[tuple[str, str]] = []
            for _excel_row, row in current_rows:
                sku_id = _text(row[code_index] if 0 <= code_index < len(row) else "")
                exported_url = _text(row[url_index] if 0 <= url_index < len(row) else "")
                if sku_id and exported_url:
                    identity = durable_identity(sku_id, exported_url)
                elif sku_id:
                    urls = by_sku.get(sku_id, set())
                    if len(urls) != 1:
                        raise ValueError("result workbook skuId fallback is ambiguous")
                    identity = (sku_id, next(iter(urls)))
                elif exported_url:
                    skus = by_url.get(exported_url, set())
                    if len(skus) != 1:
                        raise ValueError("result workbook exact URL fallback is ambiguous")
                    identity = (next(iter(skus)), exported_url)
                else:
                    raise ValueError("result workbook row has no parseable crawl identity")
                observed.append(identity)
            return tuple(observed)

        equivalence_classes: dict[tuple[tuple[str, str], ...], int] = {}
        for candidate_index in code_indices or [-1]:
            try:
                vector = identities_for(candidate_index)
            except ValueError:
                continue
            equivalence_classes.setdefault(vector, candidate_index)
        if not equivalence_classes:
            raise ValueError("no SKU candidate column maps the partial durable identities")

        primary_vector, primary_code_index = next(iter(equivalence_classes.items()))
        primary_counts = Counter(primary_vector)
        alternative_identities = {
            identity
            for vector in list(equivalence_classes)[1:]
            for identity in vector
        }
        price_index = cls._result_column_index(
            headers, frozenset({"商品价格", "价格", "采集价格", "page_price"})
        )
        status_index = cls._result_column_index(headers, frozenset({"status", "crawl_status"}))
        reason_index = cls._result_column_index(headers, frozenset({"reason", "failure_reason"}))
        exact_matches: dict[tuple[str, str], Mapping[str, Any]] = {}
        ambiguous: dict[tuple[str, str], str] = {}
        for row_index, identity in enumerate(primary_vector):
            if primary_counts[identity] != 1:
                ambiguous[identity] = "conflicting exact export rows"
                continue
            row = current_rows[row_index][1]
            price = parse_decimal(row[price_index]) if 0 <= price_index < len(row) else None
            status = _text(row[status_index] if 0 <= status_index < len(row) else "")
            reason = _text(row[reason_index] if 0 <= reason_index < len(row) else "")
            if price is not None and price > 0:
                page_status, source = "OK", "招商局"
            elif price is not None and price == 0:
                page_status, source = "PRICE_IS_ZERO", "招商局"
            elif "成功" in status:
                page_status, source = "PRICE_MISSING", "招商局-价格不可见"
            else:
                page_status, source = "FAILED", "招商局-" + (reason or status or "未知")
            exact_matches[identity] = {
                "product_code": _text(
                    row[primary_code_index]
                    if 0 <= primary_code_index < len(row)
                    else identity[0]
                ),
                "page_price_value": price,
                "page_price_status": page_status,
                "page_price_source": source,
            }

        exact_absent: set[tuple[str, str]] = set()
        for identity in expected - set(exact_matches) - set(ambiguous):
            if identity in alternative_identities:
                ambiguous[identity] = "SKU candidate equivalence classes disagree"
            else:
                exact_absent.add(identity)
        return PartialExactHarvestOutcome(
            exact_matches=exact_matches,
            exact_absent=frozenset(exact_absent),
            ambiguous=ambiguous,
            validation_summary={
                "expected": len(expected),
                "matched": len(exact_matches),
                "missing": len(exact_absent),
                "ambiguous": len(ambiguous),
                "candidate_classes": len(equivalence_classes),
                "foreign_filtered": foreign_filtered,
            },
        )

    def classify_existing_export_partial(
        self,
        job: PreparedLegacyJob,
    ) -> PartialExactHarvestOutcome:
        """Map task-owned result payloads onto immutable remote request groups."""
        if not job.result_path.is_file():
            raise ValueError("result workbook does not exist")
        events = self._read_batch_events(job)
        try:
            _task_id, task_number = self._submission_evidence(job, events)
        except SubmissionEvidenceConflictError as exc:
            raise ValueError(str(exc)) from exc
        if not task_number:
            raise ValueError("result workbook requires exact remote task provenance")

        workbook = load_workbook(job.result_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet is None:
                raise ValueError("result workbook has no active sheet")
            iterator = sheet.iter_rows(values_only=True)
            try:
                headers = [_text(value) for value in next(iterator)]
            except StopIteration as exc:
                raise ValueError("result workbook is empty") from exc
            workbook_rows = [
                (excel_row, row)
                for excel_row, row in enumerate(iterator, start=2)
                if any(_text(value) for value in row)
            ]
        finally:
            workbook.close()

        url_index = self._result_column_index(headers, RESULT_URL_HEADERS)
        code_indices = self._result_column_indices(headers, RESULT_CODE_HEADERS)
        execution_index = self._result_column_index(headers, RESULT_EXECUTION_HEADERS)
        batch_index = self._result_column_index(headers, RESULT_BATCH_HEADERS)
        mall_index = self._result_column_index(headers, RESULT_MALL_HEADERS)
        price_index = self._result_column_index(
            headers,
            frozenset({"\u5546\u54c1\u4ef7\u683c", "\u4ef7\u683c", "\u91c7\u96c6\u4ef7\u683c", "page_price"}),
        )
        status_index = self._result_column_index(
            headers, frozenset({"status", "crawl_status"})
        )
        reason_index = self._result_column_index(
            headers, frozenset({"reason", "failure_reason"})
        )
        if min(execution_index, batch_index, mall_index) < 0:
            raise ValueError("result workbook provenance columns are required")
        if url_index < 0 and not code_indices:
            raise ValueError("result workbook header is not recognized")

        foreign_filtered = 0
        current_rows: list[tuple[int, tuple[Any, ...], str]] = []
        current_executions: dict[str, Any] = {}
        for excel_row, row in workbook_rows:
            execution_number = _text(
                row[execution_index] if execution_index < len(row) else ""
            )
            batch_number = _text(row[batch_index] if batch_index < len(row) else "")
            if not execution_number:
                raise ValueError("result workbook row has no execution provenance")
            if not batch_number:
                raise ValueError("result workbook row has no batch provenance")
            if batch_number != job.task_name:
                foreign_filtered += 1
                continue
            if not execution_number.startswith(f"{task_number}-"):
                raise ValueError(
                    "result workbook execution does not belong to the exact remote task"
                )
            mall_key = _text(row[mall_index] if mall_index < len(row) else "")
            if mall_key != job.mall_key:
                raise ValueError("result workbook current batch mall does not match durable job")
            price = parse_decimal(row[price_index]) if 0 <= price_index < len(row) else None
            if execution_number in current_executions:
                if current_executions[execution_number] != price:
                    raise ValueError("result workbook execution has conflicting prices")
                raise ValueError("result workbook current batch execution is duplicated")
            current_executions[execution_number] = price
            current_rows.append((excel_row, row, execution_number))

        expected_counter = Counter(
            (request.sku_id, request.url.strip()) for request in job.requests
        )
        if not expected_counter or any(count != 1 for count in expected_counter.values()):
            raise ValueError("partial export requires unique durable request identities")
        expected = set(expected_counter)
        primary_code_index = code_indices[0] if code_indices else -1
        exact_matches: dict[tuple[str, str], Mapping[str, Any]] = {}
        ambiguous: dict[tuple[str, str], str] = {}
        unassigned_rows = 0

        for excel_row, row, execution_number in current_rows:
            exported_codes = {
                _text(row[index] if index < len(row) else "")
                for index in code_indices
                if _text(row[index] if index < len(row) else "")
            }
            exported_url = _text(row[url_index] if 0 <= url_index < len(row) else "")
            code_matches = {
                identity for identity in expected if identity[0] in exported_codes
            }
            url_matches = {
                identity
                for identity in expected
                if self._matches_legacy_export_url(identity[1], exported_url)
            }
            if code_matches:
                if url_matches and code_matches.isdisjoint(url_matches):
                    for identity in code_matches | url_matches:
                        ambiguous[identity] = (
                            "exported business fields map to conflicting durable request groups"
                        )
                    continue
                targets = code_matches
            elif len(url_matches) == 1:
                targets = url_matches
            elif len(url_matches) > 1:
                for identity in url_matches:
                    ambiguous[identity] = "exported URL maps to multiple durable request groups"
                continue
            elif len(expected) == 1:
                targets = expected
            else:
                unassigned_rows += 1
                continue

            price = parse_decimal(row[price_index]) if 0 <= price_index < len(row) else None
            status = _text(row[status_index] if 0 <= status_index < len(row) else "")
            reason = _text(row[reason_index] if 0 <= reason_index < len(row) else "")
            if price is not None and price > 0:
                page_status, source = "OK", "\u62db\u5546\u5c40"
            elif price is not None and price == 0:
                page_status, source = "PRICE_IS_ZERO", "\u62db\u5546\u5c40"
            elif "\u6210\u529f" in status:
                page_status, source = (
                    "PRICE_MISSING",
                    "\u62db\u5546\u5c40-\u4ef7\u683c\u4e0d\u53ef\u89c1",
                )
            else:
                page_status = "FAILED"
                source = "\u62db\u5546\u5c40-" + (reason or status or "\u672a\u77e5")
            result = {
                "product_code": _text(
                    row[primary_code_index]
                    if 0 <= primary_code_index < len(row)
                    else ""
                ),
                "exported_sku_candidates": sorted(exported_codes),
                "exported_url": exported_url,
                "page_price_value": price,
                "page_price_status": page_status,
                "page_price_source": source,
                "remote_execution": execution_number,
                "remote_batch": job.task_name,
                "remote_mall": job.mall_key,
                "export_excel_row": excel_row,
            }
            for identity in targets:
                if identity in ambiguous:
                    continue
                if identity in exact_matches:
                    exact_matches.pop(identity, None)
                    ambiguous[identity] = (
                        "multiple remote executions compete for one durable request"
                    )
                    continue
                exact_matches[identity] = result

        unresolved = expected - set(exact_matches) - set(ambiguous)
        if unassigned_rows:
            if not unresolved:
                raise ValueError(
                    "provenance-valid export has rows outside durable request groups"
                )
            for identity in unresolved:
                ambiguous[identity] = (
                    "provenance-valid result row lacks durable request-group mapping"
                )
        exact_absent = expected - set(exact_matches) - set(ambiguous)
        return PartialExactHarvestOutcome(
            exact_matches=exact_matches,
            exact_absent=frozenset(exact_absent),
            ambiguous=ambiguous,
            validation_summary={
                "expected": len(expected),
                "matched": len(exact_matches),
                "missing": len(exact_absent),
                "ambiguous": len(ambiguous),
                "candidate_classes": len(code_indices),
                "foreign_filtered": foreign_filtered,
            },
        )

    def harvest(self, job: PreparedLegacyJob, *, timeout_seconds: float) -> LegacyHarvestOutcome:
        if timeout_seconds <= 0:
            # Budget exhaustion is transient: the same exact remote identity
            # stays harvestable and the next sweep retries the bounded wait.
            return LegacyHarvestOutcome(
                HarvestStatus.RUNNING,
                diagnostic="legacy wait-export-only had no budget this pass",
            )
        if job.result_path.exists():
            return self.harvest_existing(job)

        def validated_export_after_runner_failure() -> LegacyHarvestOutcome | None:
            if not job.result_path.is_file():
                return None
            outcome = self.harvest_existing(job)
            return outcome if outcome.status is HarvestStatus.EXPORTED else None

        try:
            process = self._run(self._harvest_args(job, timeout_seconds), timeout_seconds)
        except (subprocess.TimeoutExpired, TimeoutError):
            validated = validated_export_after_runner_failure()
            if validated is not None:
                return validated
            return LegacyHarvestOutcome(HarvestStatus.RUNNING, diagnostic="legacy wait-export-only is still running")
        except (OSError, LegacyPreSendFailure) as exc:
            if _deterministic_wait_export_startup_failure(exc):
                return LegacyHarvestOutcome(HarvestStatus.MANUAL, diagnostic="legacy wait-export-only could not be started")
            # One failed execution against an already-confirmed remote task is
            # transient; the retry only ever re-addresses the same identity.
            return LegacyHarvestOutcome(
                HarvestStatus.RUNNING,
                diagnostic="legacy wait-export-only could not be started; retry scheduled on the same remote identity",
            )
        if int(getattr(process, "returncode", -1)) != 0:
            validated = validated_export_after_runner_failure()
            if validated is not None:
                return validated
            if job.result_path.is_file():
                # The child failed AND left a workbook that does not validate:
                # deterministic evidence, a rerun cannot fix it.
                return LegacyHarvestOutcome(HarvestStatus.MANUAL, diagnostic="legacy wait-export-only did not produce an export")
            return LegacyHarvestOutcome(
                HarvestStatus.RUNNING,
                diagnostic="legacy wait-export-only failed transiently; retry scheduled on the same remote identity",
            )
        if not job.result_path.exists():
            return LegacyHarvestOutcome(HarvestStatus.RUNNING, diagnostic="legacy export is not available yet")
        return self.harvest_existing(job)
