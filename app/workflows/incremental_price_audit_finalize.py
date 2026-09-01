"""Build and atomically publish the formal workbook for one frozen daily run."""

from __future__ import annotations

import copy
import hashlib
import json
import os
import uuid
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from app.services.daily_audit_store import (
    DailyAuditStore,
    LeaseToken,
    LeaseUnavailable,
    ProcessLock,
)
from app.services.incremental_audit_materializer import (
    FrozenAuditArtifact,
    LegacyWorkbookMaterializer,
    MaterializationResult,
)
from app.services.saas_audit_core import build_records, normalize_cell_str, parse_decimal
from app.services.title_match_model_adapter import TitleMatchModelResult
from app.workflows.work_01_saas_approval_json_audit import audit_saas_approval_records


@dataclass(frozen=True)
class FinalizeReadiness:
    frozen_rows: int = 0
    detail_pending: int = 0
    detail_terminal: int = 0
    model_pending: int = 0
    model_terminal: int = 0
    crawl_pending: int = 0
    crawl_terminal: int = 0
    crawl_unplanned: int = 0
    crawl_invalid_price: int = 0
    crawl_no_result: int = 0
    remote_pending: int = 0
    remote_manual: int = 0
    remote_invalid: int = 0
    # Additive force-finalize decomposition of the pending counts above.  The
    # coarse counters keep their exact ordinary-finalize meaning; these buckets
    # split the same rows/jobs by whether the work is an already-submitted
    # Zhaoshang long tail or something force-finalize must still refuse.
    crawl_awaiting_remote: int = 0
    crawl_submit_unknown: int = 0
    crawl_pending_waiting_submit: int = 0
    crawl_pending_awaiting_priority: int = 0
    crawl_pending_blocked: int = 0
    remote_awaiting_remote: int = 0
    remote_intent: int = 0
    remote_submit_unknown: int = 0
    remote_priority_pending: int = 0

    @property
    def ready(self) -> bool:
        return not any(
            (
                self.detail_pending,
                self.model_pending,
                self.model_terminal,
                self.crawl_pending,
                self.crawl_terminal,
                self.crawl_unplanned,
                self.crawl_invalid_price,
                self.remote_pending,
                self.remote_manual,
                self.remote_invalid,
            )
        )


@dataclass(frozen=True)
class FrozenArtifactBuild:
    readiness: FinalizeReadiness
    artifact: FrozenAuditArtifact | None


@dataclass(frozen=True)
class FinalizePublication:
    readiness: FinalizeReadiness
    materialization: MaterializationResult | None


@dataclass(frozen=True)
class ReviewPublication:
    output_path: Path
    row_count: int
    manual_count: int
    sha256: str


class FinalizeBlocked(RuntimeError):
    def __init__(self, message: str, readiness: FinalizeReadiness | None = None) -> None:
        self.readiness = readiness
        super().__init__(message)


class FinalizeAlreadyRunning(RuntimeError):
    pass


def is_recoverable_finalize_wait(readiness: FinalizeReadiness | None) -> bool:
    """True when only pending (not manual/terminal) work delays the final."""
    if readiness is None:
        return False
    has_pending = any((
        readiness.detail_pending,
        readiness.model_pending,
        readiness.crawl_pending,
        readiness.remote_pending,
    ))
    has_actionable_blocker = any((
        readiness.model_terminal,
        readiness.crawl_terminal,
        readiness.crawl_unplanned,
        readiness.crawl_invalid_price,
        readiness.remote_manual,
        readiness.remote_invalid,
    ))
    return has_pending and not has_actionable_blocker


# Crawl rows in any of these states keep the ordinary coarse crawl_pending
# count; force-finalize further splits them into tolerable vs intolerable.
CRAWL_PENDING_STATES = frozenset({
    "PENDING",
    "PRIORITY_PENDING",
    "SUBMITTED",
    "RUNNING",
    "EXPORTED",
    "RETRYABLE_FAILED",
    "SUBMIT_UNKNOWN",
})
# Remote jobs already submitted to Zhaoshang and still awaiting execution or
# harvest.  These are the only blockers force-finalize may give up waiting on.
INFLIGHT_REMOTE_STATES = frozenset({"SUBMITTED", "RUNNING", "EXPORTED"})
UNSUBMITTED_REMOTE_STATES = frozenset({"SUBMIT_INTENT", "SUBMIT_UNKNOWN", "PRIORITY_PENDING"})

FORCE_FINALIZE_REASON = "operator accepted limited remote crawl incompleteness"
FORCE_FINALIZE_INCOMPLETE_TASKS_KEY = "force_finalize_incomplete_tasks"

# Readiness buckets that force-finalize must never tolerate.  Detail/Model
# evidence, unplanned or invalid crawl work, unknown submission outcomes,
# manual/terminal states, and not-yet-submitted intent all stay blocking.
NON_TOLERABLE_BLOCKERS = (
    "detail_pending",
    "detail_terminal",
    "model_pending",
    "model_terminal",
    "crawl_terminal",
    "crawl_unplanned",
    "crawl_invalid_price",
    "crawl_submit_unknown",
    "crawl_pending_waiting_submit",
    "crawl_pending_awaiting_priority",
    "crawl_pending_blocked",
    "remote_invalid",
    "remote_manual",
    "remote_intent",
    "remote_submit_unknown",
    "remote_priority_pending",
)


def _crawl_awaiting_remote(
    crawl: Mapping[str, Any],
    remote_jobs: Mapping[int, Mapping[str, Any]],
) -> bool:
    """True when this crawl row is bound to an in-flight Zhaoshang task."""
    state = normalize_cell_str(crawl.get("state"))
    if state in INFLIGHT_REMOTE_STATES:
        return True
    if state in {"PENDING", "PRIORITY_PENDING", "RETRYABLE_FAILED"}:
        job = remote_jobs.get(int(crawl.get("remote_job_id") or 0))
        return bool(job is not None and normalize_cell_str(job.get("state")) in INFLIGHT_REMOTE_STATES)
    return False


def _decode_json(value: object, expected: type, field: str) -> Any:
    try:
        decoded = json.loads(str(value or ""))
    except (TypeError, ValueError) as exc:
        raise FinalizeBlocked(f"frozen {field} is not valid JSON") from exc
    if not isinstance(decoded, expected):
        raise FinalizeBlocked(f"frozen {field} has the wrong JSON shape")
    return decoded


def _stage_maps(store: DailyAuditStore, token: LeaseToken) -> tuple[dict[int, Mapping[str, Any]], dict[int, Mapping[str, Any]]]:
    details = {int(row["audit_item_id"]): row for row in store.list_stage_work(token, stage="DETAIL")}
    models = {int(row["audit_item_id"]): row for row in store.list_stage_work(token, stage="MODEL")}
    return details, models


def _model_key(row: Mapping[str, Any]) -> tuple[str, str, str, str]:
    return tuple(
        normalize_cell_str(row.get(name))
        for name in ("queryStr", "skuName", "modelAuditResult", "isMatch")
    )  # type: ignore[return-value]


def _detail_failure(error_text: object) -> dict[str, Any]:
    return {
        "customerPrice": None,
        "customerPriceStatus": "DETAIL_REQUEST_FAILED",
        "customerPriceSource": "",
        "customerPriceWarning": f"增量详情阶段失败：{normalize_cell_str(error_text) or 'ExternalError'}",
        "customerPriceEvidence": "",
    }


def _crawl_key(row: Mapping[str, Any]) -> tuple[str, str, str]:
    return (
        normalize_cell_str(row.get("vCommodityId") or row.get("documentId")),
        normalize_cell_str(row.get("skuId")),
        normalize_cell_str(row.get("extLink")),
    )


def _valid_crawl_price(result: Mapping[str, Any]) -> bool:
    status = normalize_cell_str(result.get("page_price_status"))
    try:
        value = Decimal(str(result.get("page_price_value")))
    except (InvalidOperation, TypeError, ValueError):
        return False
    return (status == "OK" and value > 0) or (status == "PRICE_IS_ZERO" and value == 0)


EXPORT_NO_PRICE_TERMINAL_EVIDENCE_TYPE = "EXPORT_NO_PRICE_TERMINAL"
EXPORT_NO_PRICE_RSP_DESC = "未检索到商品价格信息"
RECOVERY_EXHAUSTED_EVIDENCE_TYPE = "REMOTE_STALL_RECOVERY_EXHAUSTED"
RECOVERY_EXHAUSTED_REMOTE_STATE = "RECOVERY_EXHAUSTED"


def _valid_no_result(result: Mapping[str, Any]) -> bool:
    evidence_type = normalize_cell_str(result.get("evidence_type"))
    if evidence_type == EXPORT_NO_PRICE_TERMINAL_EVIDENCE_TYPE:
        try:
            export_count = int(result.get("export_count"))
        except (TypeError, ValueError):
            return False
        export_valid = (
            normalize_cell_str(result.get("export_rspDesc")) == EXPORT_NO_PRICE_RSP_DESC
            and export_count == 0
        )
    elif evidence_type == RECOVERY_EXHAUSTED_EVIDENCE_TYPE:
        try:
            export_count = int(result.get("export_count"))
        except (TypeError, ValueError):
            return False
        export_valid = (
            normalize_cell_str(result.get("remote_state")) == RECOVERY_EXHAUSTED_REMOTE_STATE
            and int(result.get("recovery_attempt") or 0) == 1
            and bool(result.get("disable_confirmed"))
            and bool(result.get("enable_confirmed"))
            and bool(result.get("post_recovery_harvest_checked"))
            and export_count == 0
        )
    else:
        export_valid = evidence_type == "EXPORTED_NO_MATCHING_ROW"
    return (
        normalize_cell_str(result.get("page_price_status")) == "NO_RESULT"
        and result.get("page_price_value") is None
        and export_valid
        and bool(normalize_cell_str(result.get("no_result_reason")))
    )


def _valid_recovery_exhausted_evidence(remote_job: Mapping[str, Any]) -> bool:
    """Accept the local give-up terminal only with its complete proof."""
    if not bool(int(remote_job.get("post_attempted") or 0)):
        return False
    if not normalize_cell_str(remote_job.get("remote_identity")):
        return False
    if int(remote_job.get("recovery_attempts") or 0) != 1:
        return False
    try:
        evidence = json.loads(str(remote_job.get("recovery_evidence_json") or ""))
    except (TypeError, ValueError, json.JSONDecodeError):
        return False
    if not isinstance(evidence, dict):
        return False
    restart = evidence.get("restart")
    if not isinstance(restart, dict) or normalize_cell_str(restart.get("phase")) != "ENABLE_CONFIRMED":
        return False
    disable_evidence = restart.get("disable_evidence")
    enable_evidence = restart.get("enable_evidence")
    if not (
        isinstance(disable_evidence, dict)
        and disable_evidence.get("confirmed_disabled") is True
        and isinstance(enable_evidence, dict)
        and enable_evidence.get("confirmed_enabled") is True
    ):
        return False
    try:
        datetime.fromisoformat(str(restart.get("recovery_at") or ""))
    except (TypeError, ValueError):
        return False
    return True


def _conflicting_export_price_reason(
    remote_job: Mapping[str, Any] | None,
    crawl: Mapping[str, Any] | None,
) -> str:
    if remote_job is None or crawl is None:
        return ""
    result_path = Path(normalize_cell_str(remote_job.get("result_path")))
    if not result_path.is_file():
        return ""
    workbook = load_workbook(result_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            return ""
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [normalize_cell_str(value) for value in next(iterator)]
        except StopIteration:
            return ""
        code_headers = {
            "商品编码",
            "商品最小品类单元Id",
            "商品最小品类单元id",
            "skuId",
            "sku_id",
            "code",
        }
        url_headers = {"原始链接", "商品详情地址", "url", "URL", "商品链接", "链接", "商品URL"}
        price_headers = {"商品价格", "价格", "采集价格", "page_price"}
        code_indices = [index for index, header in enumerate(headers) if header in code_headers]
        url_index = next((index for index, header in enumerate(headers) if header in url_headers), -1)
        price_index = next((index for index, header in enumerate(headers) if header in price_headers), -1)
        if not code_indices or url_index < 0 or price_index < 0:
            return ""
        sku_id = normalize_cell_str(crawl.get("sku_id"))
        exact_url = normalize_cell_str(crawl.get("exact_url"))
        prices_by_identity: dict[tuple[str, str], list[Decimal]] = {}
        for row in iterator:
            row_url = normalize_cell_str(row[url_index] if url_index < len(row) else "")
            price = parse_decimal(row[price_index] if price_index < len(row) else None)
            if not row_url or price is None:
                continue
            for index in code_indices:
                row_code = normalize_cell_str(row[index] if index < len(row) else "")
                if not row_code:
                    continue
                identity_prices = prices_by_identity.setdefault((row_code, row_url), [])
                if price not in identity_prices:
                    identity_prices.append(price)
    finally:
        workbook.close()
    prices = prices_by_identity.get((sku_id, exact_url), [])
    if len(prices) <= 1:
        prices = next(
            (values for values in prices_by_identity.values() if len(values) > 1),
            [],
        )
    if len(prices) <= 1:
        return ""
    rendered = " / ".join(format(price, ".2f") for price in prices)
    return f"同一 exact identity 存在冲突价格：{rendered}"


def build_frozen_audit_artifact(
    store: DailyAuditStore,
    token: LeaseToken,
    *,
    output_identity: str,
    review_mode: bool = False,
    tolerate_pending: bool = False,
    snapshot_id: int | None = None,
) -> FrozenArtifactBuild:
    run = store.get_run(token.business_date)
    if run is None or int(run["id"]) != token.run_id or (
        snapshot_id is None and run["frozen_snapshot_id"] is None
    ):
        raise FinalizeBlocked("daily run has no frozen authoritative snapshot")
    resolved_snapshot_id = (
        int(run["frozen_snapshot_id"]) if snapshot_id is None else int(snapshot_id)
    )
    snapshot = store.get_snapshot(resolved_snapshot_id)
    if snapshot is None:
        if snapshot_id is None:
            raise FinalizeBlocked("frozen snapshot metadata is missing")
        raise FinalizeBlocked("candidate snapshot metadata is missing")
    if snapshot_id is not None and int(snapshot["run_id"]) != int(run["id"]):
        raise FinalizeBlocked("candidate snapshot does not belong to this run")
    headers = _decode_json(snapshot.get("headers_json"), list, "headers")
    evidence = _decode_json(snapshot.get("evidence_json"), dict, "evidence")
    if not headers:
        raise FinalizeBlocked("frozen snapshot did not persist legacy headers")

    items = store.ordered_snapshot_items(int(snapshot["id"]))
    details, models = _stage_maps(store, token)
    crawls_by_key: dict[tuple[str, str, str], list[Mapping[str, Any]]] = {}
    for crawl in store.list_crawl_work(token):
        key = (
            normalize_cell_str(crawl.get("source_id")),
            normalize_cell_str(crawl.get("sku_id")),
            normalize_cell_str(crawl.get("exact_url")),
        )
        crawls_by_key.setdefault(key, []).append(crawl)
    remote_jobs = {int(job["id"]): job for job in store.list_remote_jobs(token)}
    replacements_by_old = {
        int(replacement["old_remote_job_id"]): replacement
        for replacement in store.list_remote_replacements(token.run_id)
    }

    counts = {
        "frozen_rows": len(items),
        "detail_pending": 0,
        "detail_terminal": 0,
        "model_pending": 0,
        "model_terminal": 0,
        "crawl_pending": 0,
        "crawl_terminal": 0,
        "crawl_unplanned": 0,
        "crawl_invalid_price": 0,
        "crawl_no_result": 0,
        "remote_pending": 0,
        "remote_manual": 0,
        "remote_invalid": 0,
        "crawl_awaiting_remote": 0,
        "crawl_submit_unknown": 0,
        "crawl_pending_waiting_submit": 0,
        "crawl_pending_awaiting_priority": 0,
        "crawl_pending_blocked": 0,
        "remote_awaiting_remote": 0,
        "remote_intent": 0,
        "remote_submit_unknown": 0,
        "remote_priority_pending": 0,
    }
    relevant_remote_job_ids: set[int] = set()
    rows: list[dict[str, Any]] = []
    enriched: list[dict[str, Any]] = []
    model_results: dict[tuple[str, str, str, str], TitleMatchModelResult] = {}

    for item in items:
        item_id = int(item["id"])
        row = dict(item["row"])
        row["__source_fingerprint"] = item["source_fingerprint"]
        detail = details.get(item_id)
        if detail is None or detail["state"] in {"PENDING", "CLAIMED", "RETRYABLE_FAILED"}:
            counts["detail_pending"] += 1
            info = {}
        elif detail["state"] == "TERMINAL_FAILED":
            counts["detail_terminal"] += 1
            info = _detail_failure(detail.get("error_text"))
        elif detail["state"] == "SUCCESS":
            info = _decode_json(detail.get("result_json"), dict, "detail result")
            if "customerPrice" in info:
                info = {**info, "customerPrice": parse_decimal(info.get("customerPrice"))}
        else:
            counts["detail_pending"] += 1
            info = {}
        enriched.append(info)

        model = models.get(item_id)
        if model is None or model["state"] in {"PENDING", "CLAIMED", "RETRYABLE_FAILED"}:
            counts["model_pending"] += 1
        elif model["state"] == "TERMINAL_FAILED":
            counts["model_terminal"] += 1
        elif model["state"] == "SUCCESS":
            parsed = TitleMatchModelResult.model_validate(
                _decode_json(model.get("result_json"), dict, "model result")
            )
            key = _model_key(row)
            previous = model_results.get(key)
            if previous is not None and previous != parsed:
                raise FinalizeBlocked("one model fingerprint produced conflicting persisted results")
            model_results[key] = parsed
        else:
            counts["model_pending"] += 1

        crawl_candidates = crawls_by_key.get(_crawl_key(row), [])
        successful = [crawl for crawl in crawl_candidates if crawl["state"] == "SUCCESS"]
        no_results = [crawl for crawl in crawl_candidates if crawl["state"] == "NO_RESULT"]
        required = bool(row.get("__crawl_required", bool(normalize_cell_str(row.get("extLink")))))
        if required:
            relevant_remote_job_ids.update(
                int(crawl["remote_job_id"])
                for crawl in crawl_candidates
                if crawl.get("remote_job_id") is not None
            )
        review_selected: Mapping[str, Any] | None = None
        review_status = "NOT_REQUIRED"
        review_price: Decimal | None = None
        review_reason = ""
        if successful:
            selected = max(successful, key=lambda value: int(value["id"]))
            review_selected = selected
            crawl_result = _decode_json(selected.get("result_json"), dict, "crawl result")
            if _valid_crawl_price(crawl_result):
                review_price = parse_decimal(crawl_result.get("page_price_value"))
                row["__page_price_value"] = review_price
                row["__page_price_status"] = crawl_result.get("page_price_status")
                row["__page_price_source"] = crawl_result.get("page_price_source")
                review_status = "SUCCESS"
            elif required:
                counts["crawl_invalid_price"] += 1
                review_status = "MANUAL_REVIEW"
                review_reason = "crawl result has no valid exact price"
        elif no_results:
            selected = max(no_results, key=lambda value: int(value["id"]))
            review_selected = selected
            crawl_result = _decode_json(selected.get("result_json"), dict, "crawl no-result evidence")
            if _valid_no_result(crawl_result):
                row["__page_price_value"] = None
                row["__page_price_status"] = "NO_RESULT"
                row["__page_price_source"] = crawl_result.get("page_price_source") or "招商局（无采集结果）"
                counts["crawl_no_result"] += 1
                review_status = "NO_RESULT"
            elif required:
                counts["crawl_invalid_price"] += 1
                review_status = "MANUAL_REVIEW"
                review_reason = "crawl NO_RESULT evidence is invalid"
        elif required:
            review_selected = max(crawl_candidates, key=lambda value: int(value["id"])) if crawl_candidates else None
            review_status = "MANUAL_REVIEW"
            if not crawl_candidates:
                counts["crawl_unplanned"] += 1
                review_reason = "required crawl identity has no durable crawl work"
            elif any(crawl["state"] in CRAWL_PENDING_STATES for crawl in crawl_candidates):
                counts["crawl_pending"] += 1
                pending_candidates = [
                    crawl for crawl in crawl_candidates if crawl["state"] in CRAWL_PENDING_STATES
                ]
                if any(crawl["state"] == "SUBMIT_UNKNOWN" for crawl in pending_candidates):
                    counts["crawl_submit_unknown"] += 1
                elif all(_crawl_awaiting_remote(crawl, remote_jobs) for crawl in pending_candidates):
                    counts["crawl_awaiting_remote"] += 1
                elif all(crawl.get("remote_job_id") is None for crawl in pending_candidates):
                    counts["crawl_pending_waiting_submit"] += 1
                elif all(
                    normalize_cell_str(
                        (remote_jobs.get(int(crawl.get("remote_job_id") or 0)) or {}).get("state")
                    )
                    in UNSUBMITTED_REMOTE_STATES
                    for crawl in pending_candidates
                ):
                    counts["crawl_pending_awaiting_priority"] += 1
                else:
                    counts["crawl_pending_blocked"] += 1
                review_reason = normalize_cell_str(review_selected.get("error_text")) if review_selected else ""
            else:
                counts["crawl_terminal"] += 1
                review_reason = normalize_cell_str(review_selected.get("error_text")) if review_selected else ""
            if review_selected is not None:
                remote_job_id = int(review_selected.get("remote_job_id") or 0)
                remote_job = remote_jobs.get(remote_job_id)
                conflict_reason = _conflicting_export_price_reason(
                    remote_job,
                    review_selected,
                )
                if conflict_reason:
                    review_reason = conflict_reason
            if not review_reason and review_selected is not None:
                remote_job_id = int(review_selected.get("remote_job_id") or 0)
                remote_job = remote_jobs.get(remote_job_id)
                review_reason = normalize_cell_str(
                    remote_job.get("error_text") if remote_job else ""
                )
            if not review_reason and review_selected is not None:
                review_reason = f"crawl state={review_selected['state']}"
        if review_mode:
            if review_status == "MANUAL_REVIEW":
                row["__page_price_value"] = None
                row["__page_price_status"] = "MANUAL_REVIEW"
                row["__page_price_source"] = ""
            row["__review_result_status"] = review_status
            row["__review_price_result"] = review_price
            row["__review_exception_reason"] = review_reason
            row["__review_remote_job_id"] = (
                int(review_selected.get("remote_job_id") or 0)
                if review_selected is not None and review_selected.get("remote_job_id")
                else None
            )
            row["__review_crawl_work_id"] = (
                int(review_selected["id"]) if review_selected is not None else None
            )
            row["__review_needs_human"] = review_status == "MANUAL_REVIEW"
        rows.append(row)

    pending_remote_states = {
        "SUBMIT_INTENT",
        "SUBMIT_UNKNOWN",
        "PRIORITY_PENDING",
        "SUBMITTED",
        "RUNNING",
        "EXPORTED",
    }
    for job_id in relevant_remote_job_ids:
        remote_job = remote_jobs.get(job_id)
        state = normalize_cell_str(remote_job.get("state") if remote_job else "")
        if state == "COMPLETED":
            if (
                remote_job is not None
                and bool(int(remote_job.get("post_attempted") or 0))
                and bool(normalize_cell_str(remote_job.get("remote_identity")))
            ):
                continue
            counts["remote_invalid"] += 1
            continue
        if state == "RECOVERY_EXHAUSTED":
            # A local give-up is acceptable only with its complete durable
            # recovery proof; anything less stays an invalid-state blocker.
            if remote_job is not None and _valid_recovery_exhausted_evidence(remote_job):
                continue
            counts["remote_invalid"] += 1
            continue
        if state == "SUPERSEDED":
            # A manual rebuild durably retired this task; its replacement job
            # carries the remaining work, so the old task neither blocks nor
            # counts as pending/manual. Without a recorded replacement whose
            # new job still exists, SUPERSEDED is not trusted evidence.
            replacement = replacements_by_old.get(job_id)
            if replacement is not None and int(replacement["new_remote_job_id"]) in remote_jobs:
                continue
            counts["remote_invalid"] += 1
            continue
        if state in pending_remote_states:
            counts["remote_pending"] += 1
            if state == "SUBMIT_INTENT":
                counts["remote_intent"] += 1
            elif state == "SUBMIT_UNKNOWN":
                counts["remote_submit_unknown"] += 1
            elif state == "PRIORITY_PENDING":
                counts["remote_priority_pending"] += 1
            else:
                counts["remote_awaiting_remote"] += 1
        elif state == "MANUAL_REQUIRED":
            counts["remote_manual"] += 1
        else:
            counts["remote_invalid"] += 1

    readiness = FinalizeReadiness(**counts)
    if not readiness.ready and not review_mode and not tolerate_pending:
        return FrozenArtifactBuild(readiness, None)
    if review_mode and (
        readiness.detail_pending
        or readiness.detail_terminal
        or readiness.model_pending
        or readiness.model_terminal
    ):
        raise FinalizeBlocked(
            "review export requires completed Detail and Model evidence",
            readiness,
        )

    records = build_records(rows, enriched)
    audit_response = audit_saas_approval_records(
        records,
        use_model=True,
        use_link_audit=True,
        use_crawl_result=False,
        precomputed_title_model_results=model_results,
    )
    artifact = FrozenAuditArtifact(
        business_date=token.business_date,
        tenant_id=str(run["tenant_id"]),
        commodity_pool_id=str(run["pool_id"]),
        output_identity=output_identity,
        frozen_snapshot_id=str(snapshot["id"]),
        frozen_snapshot_hash=str(snapshot["ordered_hash"]),
        frozen_snapshot_evidence={
            "observed_at": snapshot["observed_at"],
            "row_count": snapshot["row_count"],
            "source": evidence,
        },
        headers=headers,
        frozen_rows=rows,
        enriched_rows=enriched,
        audit_response=audit_response,
    )
    return FrozenArtifactBuild(readiness, artifact)


REVIEW_COLUMNS = (
    ("结果状态", "__review_result_status"),
    ("价格结果", "__review_price_result"),
    ("异常原因", "__review_exception_reason"),
    ("remote_job_id", "__review_remote_job_id"),
    ("crawl_work_id", "__review_crawl_work_id"),
    ("是否需要人工确认", "__review_needs_human"),
)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_review_workbook(path: Path, artifact: FrozenAuditArtifact) -> None:
    from scripts.audit_saas_excel_file import write_audit_workbook

    write_audit_workbook(
        path,
        list(artifact.headers),
        [dict(row) for row in artifact.frozen_rows],
        [dict(row) for row in artifact.enriched_rows],
        artifact.audit_response,
    )
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        sheet = workbook["record_results"]
        first_review_column = sheet.max_column + 1
        source_header = sheet.cell(row=1, column=sheet.max_column)
        for offset, (display_name, _source_name) in enumerate(REVIEW_COLUMNS):
            cell = sheet.cell(row=1, column=first_review_column + offset, value=display_name)
            cell.font = copy.copy(source_header.font)
            cell.fill = copy.copy(source_header.fill)
            cell.border = copy.copy(source_header.border)
            cell.alignment = copy.copy(source_header.alignment)
        manual_fill = PatternFill("solid", fgColor="FFF2CC")
        for row_index, source_row in enumerate(artifact.frozen_rows, start=2):
            for offset, (_display_name, source_name) in enumerate(REVIEW_COLUMNS):
                value = source_row.get(source_name)
                if source_name == "__review_needs_human":
                    value = "是" if value else "否"
                elif source_name == "__review_price_result" and value is not None:
                    value = float(value)
                cell = sheet.cell(
                    row=row_index,
                    column=first_review_column + offset,
                    value=value,
                )
                if source_name == "__review_price_result":
                    cell.number_format = "0.00"
                if source_row.get("__review_result_status") == "MANUAL_REVIEW":
                    cell.fill = manual_fill
        widths = (18, 14, 58, 16, 16, 18)
        for offset, width in enumerate(widths):
            sheet.column_dimensions[
                sheet.cell(row=1, column=first_review_column + offset).column_letter
            ].width = width
        sheet.auto_filter.ref = (
            f"A1:{sheet.cell(row=1, column=sheet.max_column).column_letter}"
            f"{max(sheet.max_row, 1)}"
        )
        summary = workbook["summary"]
        status_counts: dict[str, int] = {}
        for row in artifact.frozen_rows:
            status = normalize_cell_str(row.get("__review_result_status"))
            status_counts[status] = status_counts.get(status, 0) + 1
        summary.append(["复核表 frozen snapshot", artifact.frozen_snapshot_id])
        for status in ("SUCCESS", "NO_RESULT", "NOT_REQUIRED", "MANUAL_REVIEW"):
            summary.append([f"复核结果 {status}", status_counts.get(status, 0)])
        summary.column_dimensions["A"].width = 34
        summary.column_dimensions["B"].width = max(
            float(summary.column_dimensions["B"].width or 0),
            18,
        )
        workbook.properties.description = (
            "incremental_audit_review_snapshot_id="
            f"{artifact.frozen_snapshot_id};hash={artifact.frozen_snapshot_hash}"
        )
        workbook.save(path)
    finally:
        workbook.close()


def _validate_review_workbook(
    path: Path,
    artifact: FrozenAuditArtifact,
) -> tuple[int, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "record_results" not in workbook.sheetnames:
            raise FinalizeBlocked("review workbook record_results sheet is missing")
        rows = list(workbook["record_results"].iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise FinalizeBlocked("review workbook record_results sheet is empty")
    headers = [normalize_cell_str(value) for value in rows[0]]
    review_headers = [name for name, _source in REVIEW_COLUMNS]
    if headers[-len(review_headers):] != review_headers:
        raise FinalizeBlocked("review workbook audit columns are missing or reordered")
    data_rows = rows[1:]
    if len(data_rows) != len(artifact.frozen_rows):
        raise FinalizeBlocked("review workbook row count differs from frozen snapshot")
    status_index = headers.index("结果状态")
    price_index = headers.index("价格结果")
    manual_index = headers.index("是否需要人工确认")
    manual_count = 0
    for row in data_rows:
        status = normalize_cell_str(row[status_index])
        price = row[price_index]
        manual = normalize_cell_str(row[manual_index])
        if status == "MANUAL_REVIEW":
            manual_count += 1
            if price is not None:
                raise FinalizeBlocked("manual review row must not contain a selected price")
            if manual != "是":
                raise FinalizeBlocked("manual review row lacks its human-confirmation flag")
        if status == "NO_RESULT" and price is not None:
            raise FinalizeBlocked("NO_RESULT review row must keep price blank")
    return len(data_rows), manual_count


def materialize_frozen_review(
    store: DailyAuditStore,
    process_lock: ProcessLock,
    business_date: str,
    output_path: Path | str,
) -> ReviewPublication:
    """Publish a complete review workbook without touching final durable fields."""
    resolved_output = Path(output_path).resolve()
    if "final" in resolved_output.stem.lower() or resolved_output.suffix.lower() != ".xlsx":
        raise FinalizeBlocked("review output must be a non-final .xlsx path")
    if resolved_output.exists():
        raise FinalizeBlocked("review output already exists; refusing to overwrite it")
    try:
        process_lock.acquire()
    except LeaseUnavailable as exc:
        raise FinalizeAlreadyRunning("another tick/finalize owns the process lock") from exc
    token: LeaseToken | None = None
    temp_path: Path | None = None
    try:
        run = store.get_run(business_date)
        if run is None or run["frozen_snapshot_id"] is None:
            raise FinalizeBlocked("daily run has no frozen authoritative snapshot")
        token = store.acquire_lease(
            business_date,
            f"review-{os.getpid()}",
            ttl_seconds=600,
            process_lock=process_lock,
        )
        built = build_frozen_audit_artifact(
            store,
            token,
            output_identity=f"daily-review-{business_date.replace('-', '')}",
            review_mode=True,
        )
        if built.artifact is None:
            raise FinalizeBlocked("frozen review artifact could not be built", built.readiness)
        resolved_output.parent.mkdir(parents=True, exist_ok=True)
        temp_path = resolved_output.with_name(
            f".{resolved_output.stem}.{uuid.uuid4().hex}.tmp.xlsx"
        )
        _write_review_workbook(temp_path, built.artifact)
        row_count, manual_count = _validate_review_workbook(temp_path, built.artifact)
        os.replace(temp_path, resolved_output)
        return ReviewPublication(
            output_path=resolved_output,
            row_count=row_count,
            manual_count=manual_count,
            sha256=_file_sha256(resolved_output),
        )
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)
        if token is not None:
            store.release_lease(token)
        process_lock.release()


def materialize_frozen_daily_run(
    store: DailyAuditStore,
    process_lock: ProcessLock,
    business_date: str,
    output_path: Path | str,
    *,
    materializer: LegacyWorkbookMaterializer | None = None,
) -> FinalizePublication:
    try:
        process_lock.acquire()
    except LeaseUnavailable as exc:
        raise FinalizeAlreadyRunning("another tick/finalize owns the process lock") from exc
    token: LeaseToken | None = None
    try:
        run = store.get_run(business_date)
        if run is None or run["frozen_snapshot_id"] is None:
            raise FinalizeBlocked("daily run has no frozen authoritative snapshot")
        token = store.acquire_lease(
            business_date,
            f"finalize-{os.getpid()}",
            ttl_seconds=600,
            allow_unexpired_takeover=True,
            process_lock=process_lock,
        )
        built = build_frozen_audit_artifact(
            store,
            token,
            output_identity=f"daily-final-{business_date.replace('-', '')}",
        )
        if not built.readiness.ready or built.artifact is None:
            raise FinalizeBlocked("frozen snapshot still has mandatory pending or invalid work", built.readiness)
        publisher = materializer or LegacyWorkbookMaterializer()
        resolved_output = Path(output_path).resolve()
        store.prepare_finalization(token, str(resolved_output))
        result = publisher.materialize(built.artifact, resolved_output)
        store.mark_finalized(token, str(resolved_output), result.manifest_sha256)
        store.update_daily_metrics(token, {"finalized_rows": result.row_count})
        return FinalizePublication(built.readiness, result)
    finally:
        if token is not None:
            store.release_lease(token)
        process_lock.release()


class ForceFinalizeRejected(RuntimeError):
    """A guarded refusal to force-finalize; never retries or mutates state."""

    def __init__(
        self,
        error_type: str,
        message: str,
        readiness: FinalizeReadiness | None = None,
    ) -> None:
        self.error_type = error_type
        self.readiness = readiness
        super().__init__(f"{error_type}: {message}")


@dataclass(frozen=True)
class ForceFinalizePublication:
    readiness: FinalizeReadiness
    materialization: MaterializationResult | None
    forced_incomplete_count: int
    remote_incomplete_job_count: int
    blocker_count: int
    incomplete_tasks: tuple[Mapping[str, Any], ...]
    already_finalized: bool
    frozen_snapshot_id: int | None
    output_path: Path | None
    forced: bool


def _run_metrics(run: Mapping[str, Any]) -> dict[str, Any]:
    try:
        decoded = json.loads(str(run.get("metrics_json") or "{}"))
    except (TypeError, ValueError):
        return {}
    return decoded if isinstance(decoded, dict) else {}


def _parse_remote_identity(identity: object) -> tuple[str | None, str | None]:
    parts = [normalize_cell_str(part) for part in str(identity or "").split("|")]
    if len(parts) == 3 and all(parts):
        return parts[1], parts[2]
    return None, None


def _collect_incomplete_tasks(store: DailyAuditStore, token: LeaseToken) -> tuple[dict[str, Any], ...]:
    """Describe the in-flight Zhaoshang tasks force-finalize gives up waiting on.

    Only durable evidence is recorded; nothing here mutates remote or crawl
    state, cancels remote work, or fabricates results.
    """
    inflight_jobs = store.list_remote_jobs(token, states=set(INFLIGHT_REMOTE_STATES))
    if not inflight_jobs:
        return ()
    unfinished_crawls: dict[int, list[Mapping[str, Any]]] = {}
    for crawl in store.list_crawl_work(token):
        if crawl.get("remote_job_id") is None:
            continue
        if normalize_cell_str(crawl.get("state")) in {"SUCCESS", "NO_RESULT", "TERMINAL_FAILED"}:
            continue
        unfinished_crawls.setdefault(int(crawl["remote_job_id"]), []).append(crawl)
    tasks: list[dict[str, Any]] = []
    for job in inflight_jobs:
        items = unfinished_crawls.get(int(job["id"]), [])
        if not items:
            continue
        try:
            payload = json.loads(str(job.get("payload_json") or ""))
        except (TypeError, ValueError):
            payload = None
        task_id, task_number = _parse_remote_identity(job.get("remote_identity"))
        tasks.append(
            {
                "remote_job_id": int(job["id"]),
                "logical_key": job.get("logical_key"),
                "task_name": job.get("task_name"),
                "mall_code": job.get("mall_key"),
                "task_id": task_id,
                "task_number": task_number,
                "remote_identity": job.get("remote_identity"),
                "remote_state": job.get("state"),
                "error_text": job.get("error_text"),
                "expected_subtask_count": len(payload) if isinstance(payload, list) else None,
                "crawl_item_count": len(items),
                "crawl_items": [
                    {
                        "crawl_work_id": int(crawl["id"]),
                        "source_id": crawl.get("source_id"),
                        "sku_id": crawl.get("sku_id"),
                        "exact_url": crawl.get("exact_url"),
                        "mall_code": crawl.get("mall_key"),
                        "crawl_state": crawl.get("state"),
                    }
                    for crawl in items
                ],
            }
        )
    return tuple(tasks)


def _already_finalized_publication(run: Mapping[str, Any]) -> ForceFinalizePublication:
    metrics = _run_metrics(run)
    return ForceFinalizePublication(
        readiness=FinalizeReadiness(),
        materialization=None,
        forced_incomplete_count=int(metrics.get("forced_incomplete_item_count", metrics.get("forced_incomplete_count", 0)) or 0),
        remote_incomplete_job_count=int(metrics.get("remote_incomplete_job_count", 0) or 0),
        blocker_count=int(metrics.get("blocker_count", 0) or 0),
        incomplete_tasks=(),
        already_finalized=True,
        frozen_snapshot_id=int(run["frozen_snapshot_id"]) if run.get("frozen_snapshot_id") is not None else None,
        output_path=Path(str(run["final_output_path"])) if run.get("final_output_path") else None,
        forced=bool(metrics.get("force_finalize", False)),
    )


def materialize_forced_frozen_daily_run(
    store: DailyAuditStore,
    process_lock: ProcessLock,
    business_date: str,
    output_path: Path | str,
    *,
    max_incomplete: int = 10,
    materializer: LegacyWorkbookMaterializer | None = None,
) -> ForceFinalizePublication:
    """Publish the formal final workbook while tolerating bounded remote long tail.

    Preconditions: the day must already be frozen, no production writer may be
    active, and every remaining blocker must be an already-submitted Zhaoshang
    crawl/remote item.  ``max_incomplete`` bounds the number of business rows
    without a crawl result; in-flight remote jobs are counted separately.
    Rows without crawl results keep the existing writer "no result" semantics;
    nothing is fabricated, and the durable remote state survives untouched as
    history.
    """
    if max_incomplete < 0:
        raise ValueError("max_incomplete must be zero or positive")
    resolved_output = Path(output_path).resolve()
    run = store.get_run(business_date)
    if run is None or run.get("frozen_snapshot_id") is None:
        raise ForceFinalizeRejected(
            "FORCE_FINALIZE_NOT_FROZEN",
            "daily run has no frozen authoritative snapshot; freeze first",
        )
    if run["state"] == "FINALIZED":
        return _already_finalized_publication(run)
    try:
        process_lock.acquire()
    except LeaseUnavailable as exc:
        raise ForceFinalizeRejected(
            "FORCE_FINALIZE_PRODUCTION_WRITER_ACTIVE",
            "another tick/closeout/finalize owns the production writer lock",
        ) from exc
    token: LeaseToken | None = None
    try:
        run = store.get_run(business_date)
        if run is None or run.get("frozen_snapshot_id") is None:
            raise ForceFinalizeRejected(
                "FORCE_FINALIZE_NOT_FROZEN",
                "daily run has no frozen authoritative snapshot; freeze first",
            )
        if run["state"] == "FINALIZED":
            return _already_finalized_publication(run)
        token = store.acquire_lease(
            business_date,
            f"force-finalize-{os.getpid()}",
            ttl_seconds=600,
            allow_unexpired_takeover=True,
            process_lock=process_lock,
        )
        built = build_frozen_audit_artifact(
            store,
            token,
            output_identity=f"daily-final-{business_date.replace('-', '')}",
            tolerate_pending=True,
        )
        readiness = built.readiness
        if built.artifact is None:
            raise ForceFinalizeRejected(
                "FORCE_FINALIZE_BLOCKED",
                "frozen artifact could not be built",
                readiness,
            )
        blockers = {
            name: getattr(readiness, name)
            for name in NON_TOLERABLE_BLOCKERS
            if getattr(readiness, name)
        }
        if blockers:
            rendered = ", ".join(f"{name}={value}" for name, value in sorted(blockers.items()))
            raise ForceFinalizeRejected(
                "FORCE_FINALIZE_BLOCKED",
                f"non-tolerable blockers remain: {rendered}",
                readiness,
            )
        # Business threshold: at most N data rows may lack a crawl result.
        # In-flight remote jobs are recorded separately and never consume the
        # item budget.
        item_count = readiness.crawl_awaiting_remote
        remote_job_count = readiness.remote_awaiting_remote
        blocker_count = item_count + remote_job_count
        if item_count > max_incomplete:
            raise ForceFinalizeRejected(
                "FORCE_FINALIZE_INCOMPLETE_LIMIT_EXCEEDED",
                f"tolerable incomplete items {item_count} (plus {remote_job_count} in-flight remote jobs)"
                f" exceed max_incomplete {max_incomplete}",
                readiness,
            )
        incomplete_tasks = _collect_incomplete_tasks(store, token)
        publisher = materializer or LegacyWorkbookMaterializer()
        store.prepare_finalization(token, str(resolved_output))
        result = publisher.materialize(built.artifact, resolved_output)
        store.mark_finalized(token, str(resolved_output), result.manifest_sha256)
        store.update_daily_metrics(
            token,
            {
                "finalized_rows": result.row_count,
                "forced": True,
                "forced_incomplete_count": item_count,
                "forced_incomplete_item_count": item_count,
                "remote_incomplete_job_count": remote_job_count,
                "blocker_count": blocker_count,
                "force_finalize": True,
                "force_finalize_at": datetime.now(UTC).isoformat(timespec="seconds"),
                "force_finalize_max_incomplete": max_incomplete,
                "force_finalize_actual_incomplete": item_count,
                "force_finalize_reason": FORCE_FINALIZE_REASON,
                "frozen_snapshot_id": int(run["frozen_snapshot_id"]),
                "frozen_row_count": readiness.frozen_rows,
                "blocker_summary_at_force_finalize": asdict(readiness),
                FORCE_FINALIZE_INCOMPLETE_TASKS_KEY: [dict(task) for task in incomplete_tasks],
            },
        )
        return ForceFinalizePublication(
            readiness=readiness,
            materialization=result,
            forced_incomplete_count=item_count,
            remote_incomplete_job_count=remote_job_count,
            blocker_count=blocker_count,
            incomplete_tasks=incomplete_tasks,
            already_finalized=False,
            frozen_snapshot_id=int(run["frozen_snapshot_id"]),
            output_path=result.output_path,
            forced=True,
        )
    finally:
        if token is not None:
            store.release_lease(token)
        process_lock.release()
