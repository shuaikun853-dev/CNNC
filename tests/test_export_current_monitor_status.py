import importlib.util
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPT_DIR))
SPEC = importlib.util.spec_from_file_location(
    "export_current_monitor_status",
    SCRIPT_DIR / "export_current_monitor_status.py",
)
assert SPEC and SPEC.loader
monitor_export = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(monitor_export)


@pytest.mark.parametrize("value", [1, 0.01, "12.50", " 3 "])
def test_valid_new_price_accepts_positive_numbers(value):
    assert monitor_export._is_valid_new_price(value) is True


@pytest.mark.parametrize(
    "value",
    [
        None,
        "",
        " ",
        "null",
        "None",
        "abc",
        "NaN",
        "Infinity",
        0,
        "0",
        -1,
        "-0.01",
        True,
    ],
)
def test_valid_new_price_rejects_missing_non_numeric_and_non_positive_values(
    value,
):
    assert monitor_export._is_valid_new_price(value) is False


@pytest.mark.parametrize(
    ("row", "pushable", "category"),
    [
        (
            {
                "isSync": 1,
                "isSyncStandard": 1,
                "syncPriceStatus": 1,
                "newPrice": "25.80",
            },
            "是",
            "可推送",
        ),
        (
            {
                "isSync": 1,
                "isSyncStandard": 1,
                "syncPriceStatus": "",
                "newPrice": 8,
            },
            "是",
            "可推送",
        ),
        (
            {
                "isSync": 1,
                "isSyncStandard": 1,
                "syncPriceStatus": 2,
                "newPrice": 8,
            },
            "是",
            "可推送",
        ),
        (
            {
                "isSync": 1,
                "isSyncStandard": 1,
                "syncPriceStatus": 5,
                "newPrice": 8,
            },
            "否",
            "价格检测校验失败",
        ),
        (
            {
                "isSync": 1,
                "isSyncStandard": 1,
                "syncPriceStatus": 1,
                "newPrice": "",
            },
            "否",
            "推送价格为空或无效",
        ),
        (
            {
                "isSync": 0,
                "isSyncStandard": 1,
                "syncPriceStatus": 1,
                "newPrice": 8,
            },
            "否",
            "采集未成功",
        ),
        (
            {
                "isSync": 1,
                "isSyncStandard": 2,
                "syncPriceStatus": 1,
                "newPrice": 8,
            },
            "否",
            "同款匹配异常/未成功",
        ),
    ],
)
def test_status_requires_successful_states_and_valid_push_price(
    row,
    pushable,
    category,
):
    result = monitor_export._status(row)
    assert result[3] == pushable
    assert result[4] == category


def test_build_monitor_values_exposes_new_price_and_push_decision():
    values = monitor_export._build_monitor_values(
        {
            "isSync": 1,
            "isSyncStandard": 1,
            "syncPriceStatus": 1,
            "newPrice": "19.90",
        }
    )

    assert values[3] == "19.90"
    assert values[4] == "是"
    assert values[5] == "可推送"
    assert values[7] == "已匹配"


def test_write_workbook_reconciles_new_price_classification(tmp_path):
    records = [
        [
            "A",
            *[""] * 12,
            "采集成功",
            "同款匹配成功",
            "价格计算成功",
            10,
            "是",
            "可推送",
            "",
            "已匹配",
        ],
        [
            "B",
            *[""] * 12,
            "采集成功",
            "同款匹配成功",
            "无价格",
            "",
            "否",
            "推送价格为空或无效",
            "推送价格为空或无效",
            "已匹配",
        ],
        [
            "C",
            *[""] * 12,
            "",
            "",
            "",
            "",
            "否",
            "SaaS监控缺失",
            "中核存在但SaaS监控未返回",
            "缺失",
        ],
    ]
    output = tmp_path / "monitor.xlsx"

    monitor_export._write_workbook(records, output)

    workbook = load_workbook(output, read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "全量监控结果",
        "可推送数据",
        "需处理数据",
        "SaaS监控缺失",
        "状态统计",
    ]
    headers = [
        cell.value
        for cell in workbook["全量监控结果"][1]
    ]
    assert headers[16] == "newPrice"
    assert workbook["可推送数据"].max_row - 1 == 1
    assert workbook["需处理数据"].max_row - 1 == 1
    assert workbook["SaaS监控缺失"].max_row - 1 == 1
