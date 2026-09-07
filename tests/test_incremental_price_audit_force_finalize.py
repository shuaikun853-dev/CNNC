"""RED-first tests for the guarded force-finalize command.

Force-finalize publishes the formal daily workbook from the frozen snapshot
while tolerating a bounded number of already-submitted Zhaoshang remote
long-tail items.  It must never fabricate crawl evidence, must reject every
non-long-tail blocker, and must stay idempotent and immutable after publish.
"""

from __future__ import annotations

import json
import os
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.daily_audit_store import DailyAuditStore, ProcessLock
from app.services.incremental_audit_runtime import (
    run_force_finalize_command,
    run_runtime_command,
)
from app.services.incremental_audit_status import collect_runtime_status
from app.services.incremental_audit_materializer import LegacyWorkbookMaterializer
from app.workflows.incremental_price_audit_finalize import (
    ForceFinalizeRejected,
    build_frozen_audit_artifact,
    materialize_forced_frozen_daily_run,
    materialize_frozen_daily_run,
)
from scripts.audit_saas_excel_file import RESULT_APPENDED_COLUMNS

pytestmark = pytest.mark.critical_safety


DATE = "2026-08-24"
NO_RESULT_REASON = "legacy export missing safe skuId / exact URL result"


def _row(v="v-1", sku="sku-1", url="https://shop.example/item", **extra):
    return {
        "vCommodityId": v,
        "documentId": v,
        "skuId": sku,
        "extLink": url,
        "queryStr": "客户标题",
        "sassVCommodityName": "客户标题",
        "skuName": "互联网标题",
        "vendorName": "商城",
        "shopName": "店铺",
        "price": "10",
        "fourCategoryId": "cat-1",
        "fourCategoryName": "分类",
        "list_reference_price": "10",
        "customer_lookup_status": "MATCHED",
        "__crawl_required": bool(url),
        **extra,
    }


def _seed_force(tmp_path: Path, *, rows=None, freeze=True):
    """Frozen run with completed Detail/Model and crawl work left PENDING."""
    rows = list(rows or [_row(), _row("v-2", "sku-2", "https://shop.example/other")])
    runtime_root = tmp_path / "runtime"
    run_dir = runtime_root / DATE
    run_dir.mkdir(parents=True, exist_ok=True)
    store = DailyAuditStore(run_dir / "state.sqlite3")
    store.ensure_run(DATE, cutoff_at=f"{DATE}T10:00:00+00:00")
    token = store.acquire_lease(DATE, "seed", 60)
    # Mirror the real source export: audit-appended columns are never part of
    # the frozen header matrix the final writer preserves positionally.
    source_headers = [
        name for name in rows[0]
        if not name.startswith("__") and name not in RESULT_APPENDED_COLUMNS
    ]
    snapshot = store.ingest_snapshot(
        token,
        rows,
        f"{DATE}T10:01:00+00:00",
        headers=source_headers,
        evidence={"observed_count": len(rows)},
    )["snapshot"]
    if freeze:
        store.freeze_snapshot(token, snapshot["id"])
    for item in {item["id"]: item for item in store.ordered_snapshot_items(snapshot["id"])}.values():
        detail = store.ensure_stage_work(token, item["id"], "DETAIL", f"detail-{item['id']}")
        store.claim_stage_work(token, detail["id"], max_attempts=3, max_elapsed_seconds=300)
        store.complete_stage_work(token, detail["id"], {
            "customerPrice": "10",
            "customerPriceStatus": "FOUND",
            "customerPriceSource": "DETAIL_REFERENCE_INFOS",
            "customerPriceWarning": "",
            "customerPriceEvidence": "customer=10",
        })
        model = store.ensure_stage_work(token, item["id"], "MODEL", f"model-{item['id']}")
        store.claim_stage_work(token, model["id"], max_attempts=1, max_elapsed_seconds=300)
        store.complete_stage_work(token, model["id"], {
            "match_result": "SAME", "confidence": 0.95, "model_match": "exact", "reason": "一致",
        })
    crawls = []
    for row in rows:
        if row.get("__crawl_required"):
            crawls.append(store.ensure_crawl_work(
                token, "mall", row["skuId"], row["extLink"], "legacy-v1",
                f"crawl-fp-{row['skuId']}",
                source_id=row["vCommodityId"],
            ))
    store.release_lease(token)
    return store, snapshot, crawls


def _link_inflight_job(store, tmp_path, crawls, *, index=0, remote_state="SUBMITTED", crawl_state="SUBMITTED"):
    """Bind crawls to a durable remote job already submitted to Zhaoshang."""
    token = store.acquire_lease(DATE, "link", 60)
    job = store.prepare_remote_intent(
        token,
        f"mall:hash-{index}",
        f"daily-{DATE.replace('-', '')}-mall-{index}",
        f"payload-{index}",
        mall_key="mall",
        result_path=str(tmp_path / f"result-{index}.xlsx"),
        payload_json=[{"sku_id": crawl["sku_id"], "url": crawl["exact_url"]} for crawl in crawls],
    )
    store.claim_remote_submission(token, int(job["id"]))
    store.reconcile_remote_job(
        token,
        int(job["id"]),
        remote_identity=f"mall|task-id-{index}|task-number-{index}",
        state=remote_state,
    )
    for crawl in crawls:
        store.link_crawl_work_to_remote_job(token, int(crawl["id"]), int(job["id"]))
    if crawl_state is not None:
        store.advance_crawl_work_for_job(token, int(job["id"]), crawl_state)
    store.release_lease(token)
    return store.get_remote_job(int(job["id"]))


def _complete_crawl(store, crawl, value="9.5"):
    token = store.acquire_lease(DATE, "complete", 60)
    store.complete_crawl_work(token, crawl["id"], {
        "page_price_value": value, "page_price_status": "OK", "page_price_source": "招商局",
    })
    store.release_lease(token)


def _complete_no_result_crawl(store, tmp_path, crawl, *, index=0):
    token = store.acquire_lease(DATE, "no-result", 60)
    job = store.prepare_remote_intent(
        token,
        f"mall:no-result-{index}",
        f"daily-{DATE.replace('-', '')}-no-result-{index}",
        f"payload-no-result-{index}",
        result_path=str(tmp_path / f"no-result-{index}.xlsx"),
    )
    store.claim_remote_submission(token, int(job["id"]))
    store.reconcile_remote_job(
        token,
        int(job["id"]),
        remote_identity=f"mall|task-id-nr-{index}|task-number-nr-{index}",
        state="COMPLETED",
        result={"https://other/item": {"price": "12.30"}},
    )
    store.link_crawl_work_to_remote_job(token, int(crawl["id"]), int(job["id"]))
    store.fail_crawl_work(token, int(crawl["id"]), NO_RESULT_REASON, retryable=False)
    store.complete_crawl_work_without_result(token, int(crawl["id"]), reason=NO_RESULT_REASON)
    store.release_lease(token)


def _metrics(store) -> dict:
    return json.loads(store.get_run(DATE)["metrics_json"] or "{}")


# Case 1: normal force finalize over a bounded in-flight remote long tail.
def test_force_finalize_publishes_formal_workbook_with_force_audit(tmp_path):
    store, _snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)
    output = tmp_path / "formal.xlsx"

    publication = materialize_forced_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=2,
    )

    assert publication.already_finalized is False
    # Business threshold counts data rows without a crawl result; the one
    # in-flight remote job is recorded separately and never eats the budget.
    assert publication.forced_incomplete_count == 2
    assert publication.remote_incomplete_job_count == 1
    assert publication.blocker_count == 3
    assert publication.materialization is not None
    assert publication.materialization.output_path == output.resolve()
    assert output.is_file()
    assert output.with_name(output.name + ".manifest.json").is_file()
    run = store.get_run(DATE)
    assert run["state"] == "FINALIZED"
    assert run["final_output_path"] == str(output.resolve())
    assert run["final_manifest_hash"]
    metrics = _metrics(store)
    assert metrics["force_finalize"] is True
    assert metrics["forced"] is True
    assert metrics["forced_incomplete_count"] == 2
    assert metrics["forced_incomplete_item_count"] == 2
    assert metrics["remote_incomplete_job_count"] == 1
    assert metrics["blocker_count"] == 3
    assert metrics["force_finalize_max_incomplete"] == 2
    assert metrics["force_finalize_actual_incomplete"] == 2
    assert metrics["force_finalize_reason"] == "operator accepted limited remote crawl incompleteness"
    assert metrics["force_finalize_at"]
    assert metrics["frozen_snapshot_id"] == run["frozen_snapshot_id"]
    assert metrics["frozen_row_count"] == 2
    assert metrics["blocker_summary_at_force_finalize"]["crawl_awaiting_remote"] == 2
    assert metrics["blocker_summary_at_force_finalize"]["remote_awaiting_remote"] == 1
    tasks = metrics["force_finalize_incomplete_tasks"]
    assert len(tasks) == 1
    task = tasks[0]
    assert task["remote_state"] == "SUBMITTED"
    assert task["task_name"]
    assert task["mall_code"] == "mall"
    assert task["remote_identity"]
    assert task["expected_subtask_count"] == 2
    assert task["crawl_item_count"] == 2
    assert {item["crawl_state"] for item in task["crawl_items"]} == {"SUBMITTED"}


# Business example: 10 crawl rows + 1 remote job is allowed at max_incomplete=10.
def test_force_finalize_item_budget_is_not_consumed_by_remote_jobs(tmp_path):
    rows = [_row(f"v-{index}", f"sku-{index}", f"https://shop.example/item-{index}") for index in range(1, 11)]
    store, _snapshot, crawls = _seed_force(tmp_path, rows=rows)
    _link_inflight_job(store, tmp_path, crawls)
    output = tmp_path / "formal.xlsx"

    publication = materialize_forced_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
    )

    assert publication.forced_incomplete_count == 10
    assert publication.remote_incomplete_job_count == 1
    assert publication.blocker_count == 11  # 11 blockers, still allowed
    assert store.get_run(DATE)["state"] == "FINALIZED"
    assert output.is_file()


# Case 2: tolerable incomplete items above the threshold must be refused.
def test_force_finalize_rejects_when_incomplete_exceeds_max(tmp_path):
    rows = [_row(f"v-{index}", f"sku-{index}", f"https://shop.example/item-{index}") for index in range(1, 12)]
    store, _snapshot, crawls = _seed_force(tmp_path, rows=rows)
    _link_inflight_job(store, tmp_path, crawls)
    output = tmp_path / "formal.xlsx"

    with pytest.raises(ForceFinalizeRejected) as excinfo:
        materialize_forced_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
        )

    assert excinfo.value.error_type == "FORCE_FINALIZE_INCOMPLETE_LIMIT_EXCEEDED"
    assert excinfo.value.readiness is not None
    assert excinfo.value.readiness.crawl_awaiting_remote == 11
    assert excinfo.value.readiness.remote_awaiting_remote == 1
    assert not output.exists()
    assert store.get_run(DATE)["state"] != "FINALIZED"


# Case 3: a single model_terminal blocker must never be tolerated.
def test_force_finalize_rejects_model_terminal(tmp_path):
    store, snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)
    token = store.acquire_lease(DATE, "break-model", 60)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    model = next(
        row for row in store.list_stage_work(token, stage="MODEL")
        if row["audit_item_id"] == item["id"]
    )
    store.ensure_stage_work(token, item["id"], "MODEL", model["fingerprint"] + "-changed")
    assert store.claim_stage_work(token, model["id"], max_attempts=1, max_elapsed_seconds=1) is not None
    store.fail_stage_work(
        token, model["id"], elapsed_seconds=2, error="ModelTerminalError",
        retryable=False, max_attempts=1, max_elapsed_seconds=1,
    )
    store.release_lease(token)
    output = tmp_path / "formal.xlsx"

    with pytest.raises(ForceFinalizeRejected) as excinfo:
        materialize_forced_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
        )

    assert excinfo.value.error_type == "FORCE_FINALIZE_BLOCKED"
    assert excinfo.value.readiness.model_terminal == 1
    assert not output.exists()


# Case 4: remote_unknown (SUBMIT_UNKNOWN) must be refused.
def test_force_finalize_rejects_remote_unknown(tmp_path):
    store, _snapshot, crawls = _seed_force(tmp_path)
    job = _link_inflight_job(store, tmp_path, crawls)
    token = store.acquire_lease(DATE, "unknown", 60)
    store.mark_submit_unknown(token, int(job["id"]), "submit dispatch interrupted")
    store.release_lease(token)
    output = tmp_path / "formal.xlsx"

    with pytest.raises(ForceFinalizeRejected) as excinfo:
        materialize_forced_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
        )

    assert excinfo.value.error_type == "FORCE_FINALIZE_BLOCKED"
    assert excinfo.value.readiness.remote_submit_unknown == 1
    assert not output.exists()


# Case 5: remote MANUAL_REQUIRED must be refused.
def test_force_finalize_rejects_remote_manual(tmp_path):
    store, _snapshot, crawls = _seed_force(tmp_path)
    job = _link_inflight_job(store, tmp_path, crawls)
    token = store.acquire_lease(DATE, "manual", 60)
    store.mark_remote_manual(token, int(job["id"]), "remote task requires manual intervention")
    store.release_lease(token)
    output = tmp_path / "formal.xlsx"

    with pytest.raises(ForceFinalizeRejected) as excinfo:
        materialize_forced_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
        )

    assert excinfo.value.error_type == "FORCE_FINALIZE_BLOCKED"
    assert excinfo.value.readiness.remote_manual == 1
    assert not output.exists()


# Case 6: crawl-level SUBMIT_UNKNOWN must be refused.
def test_force_finalize_rejects_crawl_submit_unknown(tmp_path):
    store, _snapshot, crawls = _seed_force(tmp_path)
    job = _link_inflight_job(store, tmp_path, crawls, crawl_state=None)
    token = store.acquire_lease(DATE, "crawl-unknown", 60)
    store.advance_crawl_work_for_job(token, int(job["id"]), "SUBMIT_UNKNOWN")
    store.release_lease(token)
    output = tmp_path / "formal.xlsx"

    with pytest.raises(ForceFinalizeRejected) as excinfo:
        materialize_forced_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
        )

    assert excinfo.value.error_type == "FORCE_FINALIZE_BLOCKED"
    assert excinfo.value.readiness.crawl_submit_unknown == 2
    assert not output.exists()


# Case 7: no frozen snapshot yet must be refused outright.
def test_force_finalize_requires_frozen_snapshot(tmp_path):
    store, _snapshot, _crawls = _seed_force(tmp_path, freeze=False)
    output = tmp_path / "formal.xlsx"

    with pytest.raises(ForceFinalizeRejected) as excinfo:
        materialize_forced_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
        )

    assert excinfo.value.error_type == "FORCE_FINALIZE_NOT_FROZEN"
    assert not output.exists()


# Case 8: an active production writer must be refused.
def test_force_finalize_rejects_while_production_writer_active(tmp_path):
    store, _snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)
    output = tmp_path / "formal.xlsx"
    busy = ProcessLock(tmp_path / "finalize.lock")
    busy.acquire(owner={"command": "closeout"})

    with pytest.raises(ForceFinalizeRejected) as excinfo:
        materialize_forced_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
        )

    assert excinfo.value.error_type == "FORCE_FINALIZE_PRODUCTION_WRITER_ACTIVE"
    assert not output.exists()
    assert store.get_run(DATE)["state"] != "FINALIZED"
    busy.release()


# Case 9: rerunning the identical command must be idempotent.
def test_force_finalize_is_idempotent_after_success(tmp_path):
    store, _snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)
    output = tmp_path / "formal.xlsx"
    first = materialize_forced_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
    )
    published = output.read_bytes()
    metrics_before = store.get_run(DATE)["metrics_json"]

    second = materialize_forced_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
    )

    assert second.already_finalized is True
    assert second.forced is True
    assert second.forced_incomplete_count == 2
    assert second.remote_incomplete_job_count == 1
    assert second.blocker_count == 3
    assert second.frozen_snapshot_id == first.frozen_snapshot_id
    assert second.output_path == output.resolve()
    assert output.read_bytes() == published
    assert store.get_run(DATE)["metrics_json"] == metrics_before


# Case 10: every frozen row stays in the workbook; unfinished rows keep the
# existing "no crawl result" semantics instead of fabricated prices.
def test_force_finalize_workbook_covers_whole_frozen_snapshot(tmp_path):
    rows = [
        _row("v-1", "sku-1", "https://shop.example/one"),
        _row("v-2", "sku-2", "https://shop.example/two"),
        _row("v-3", "sku-3", "https://shop.example/three"),
        _row("v-4", "sku-4", "https://shop.example/four"),
    ]
    store, _snapshot, crawls = _seed_force(tmp_path, rows=rows)
    _complete_crawl(store, crawls[0], "9.5")
    _complete_no_result_crawl(store, tmp_path, crawls[1], index=1)
    _link_inflight_job(store, tmp_path, crawls[2:], index=2)
    output = tmp_path / "formal.xlsx"

    publication = materialize_forced_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
    )

    assert publication.forced_incomplete_count == 2  # the 2 unfinished rows
    assert publication.remote_incomplete_job_count == 1
    assert publication.blocker_count == 3
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        sheet = workbook["record_results"]
        matrix = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = [str(value) for value in matrix[0]]
    data_rows = matrix[1:]
    assert len(data_rows) == 4  # the workbook still covers the entire frozen snapshot
    v_index = next(index for index, header in enumerate(headers) if header in {"vCommodityId", "SaaS商品ID"})
    assert [row[v_index] for row in data_rows] == ["v-1", "v-2", "v-3", "v-4"]
    price_index = headers.index("爬取页面价格")
    status_index = headers.index("页面价格状态")
    assert str(data_rows[0][price_index]).strip() == "9.5"
    assert data_rows[0][status_index] == "成功"
    assert not data_rows[1][price_index]  # NO_RESULT keeps its blank price
    assert data_rows[1][status_index] == "无采集结果"
    for pending_row in data_rows[2:]:
        assert not pending_row[price_index]  # no fabricated price
        assert pending_row[status_index] in {"待爬取", "招商局处理"}  # existing no-result wording


# Section 9: a late harvest after force finalize must not rewrite the final.
def test_late_remote_completion_never_rewrites_published_final(tmp_path):
    store, _snapshot, crawls = _seed_force(tmp_path)
    job = _link_inflight_job(store, tmp_path, crawls)
    output = tmp_path / "formal.xlsx"
    materialize_forced_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
    )
    published = output.read_bytes()

    # The Zhaoshang remote task finishes much later and a harvest is replayed.
    token = store.acquire_lease(DATE, "late-harvest", 60)
    store.reconcile_remote_job(
        token, int(job["id"]),
        remote_identity=f"mall|task-id-0|task-number-0",
        state="COMPLETED",
        result={"https://shop.example/item": {"price": "1.10"}},
    )
    for crawl in crawls:
        store.complete_crawl_work(token, crawl["id"], {
            "page_price_value": "1.10", "page_price_status": "OK", "page_price_source": "招商局",
        })
    store.release_lease(token)

    replayed = materialize_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), DATE, output,
    )
    assert replayed.materialization is not None and replayed.materialization.reused is True
    assert output.read_bytes() == published
    rerun = materialize_forced_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
    )
    assert rerun.already_finalized is True
    assert output.read_bytes() == published


# Ordinary finalize keeps refusing a day that only force-finalize may publish.
def test_ordinary_finalize_still_blocks_the_same_snapshot(tmp_path):
    store, _snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)
    token = store.acquire_lease(DATE, "readiness", 60)
    built = build_frozen_audit_artifact(store, token, output_identity=f"daily-final-{DATE.replace('-', '')}")
    store.release_lease(token)

    assert built.readiness.ready is False  # unchanged ordinary finalize semantics
    assert built.readiness.crawl_pending == 2
    assert built.readiness.remote_pending == 1
    assert built.artifact is None


# ---- Customer-price gates (plan §6, matrix D force half) -------------------


def _rework_detail(store, item_id, *, result=None, error=None):
    """Replace one DETAIL row's durable evidence with a new outcome."""
    token = store.acquire_lease(DATE, "rework-detail", 60)
    detail = next(
        row for row in store.list_stage_work(token, stage="DETAIL")
        if row["audit_item_id"] == item_id
    )
    detail = store.ensure_stage_work(
        token, item_id, "DETAIL", str(detail["fingerprint"]) + "-rework"
    )
    claimed = store.claim_stage_work(
        token, detail["id"], max_attempts=1, max_elapsed_seconds=300
    )
    assert claimed is not None
    if error is not None:
        store.fail_stage_work(
            token, detail["id"], elapsed_seconds=1, error=error, retryable=True,
            max_attempts=1, max_elapsed_seconds=1,
        )
    else:
        store.complete_stage_work(token, detail["id"], result)
    store.release_lease(token)


def _proven_absent_result(status):
    return {
        "customerPrice": None,
        "customerPriceStatus": status,
        "customerPriceSource": "",
        "customerPriceWarning": "authoritative customer export has no such document",
        "customerPriceEvidence": "",
    }


@pytest.mark.parametrize(
    ("rework", "expected_blocker", "expected_count"),
    [
        (
            {"error": "customerPriceStatus=NOT_FOUND；详情接口未返回有效参考价"},
            "customer_price_missing",
            1,
        ),
        (
            {"error": "TimeoutError"},
            "customer_price_failed",
            1,
        ),
        (
            {"result": {**_proven_absent_result("DETAIL_REQUEST_FAILED"), "customerPriceWarning": "rspCode=500"}},
            "customer_price_failed",
            1,
        ),
        (
            {"result": {
                "customerPrice": "0",
                "customerPriceStatus": "FOUND",
                "customerPriceSource": "DETAIL_REFERENCE_INFOS",
                "customerPriceWarning": "",
                "customerPriceEvidence": "参考价1=0",
            }},
            "customer_price_invalid",
            1,
        ),
    ],
)
def test_force_finalize_never_tolerates_unresolved_customer_price(
    tmp_path, rework, expected_blocker, expected_count
):
    store, snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    _rework_detail(store, item["id"], **rework)
    output = tmp_path / "formal.xlsx"

    with pytest.raises(ForceFinalizeRejected) as excinfo:
        materialize_forced_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
        )

    assert excinfo.value.error_type == "FORCE_FINALIZE_BLOCKED"
    assert getattr(excinfo.value.readiness, expected_blocker) == expected_count
    assert not output.exists()
    assert store.get_run(DATE)["state"] != "FINALIZED"


@pytest.mark.parametrize("status", ["LOOKUP_FAILED", "DETAIL_REQUEST_SKIPPED", "NOT_PROVIDED"])
def test_force_finalize_never_tolerates_proven_absent_customer_price(tmp_path, status):
    """Round-2 fail-closed ruling (flipped from round-1 tolerance): a required
    row without a valid customer price blocks force-finalize too, even when
    the absence is deterministic and evidence-complete."""
    store, snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)
    for item in store.ordered_snapshot_items(snapshot["id"]):
        _rework_detail(store, item["id"], result=_proven_absent_result(status))
    output = tmp_path / "formal.xlsx"

    with pytest.raises(ForceFinalizeRejected) as excinfo:
        materialize_forced_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), DATE, output, max_incomplete=10,
        )

    assert excinfo.value.error_type == "FORCE_FINALIZE_BLOCKED"
    assert excinfo.value.readiness.customer_price_absent_proven == 2
    assert not output.exists()
    assert store.get_run(DATE)["state"] != "FINALIZED"


# CLI wiring.
def test_cli_force_finalize_outputs_json_and_exit_zero(tmp_path, capsys):
    from scripts.incremental_price_audit import main

    store, _snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)

    code = main([
        "force-finalize",
        "--business-date", DATE,
        "--runtime-root", str(tmp_path / "runtime"),
        "--max-incomplete", "10",
    ])

    assert code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["business_date"] == DATE
    assert payload["outcome"] == "FORCE_FINALIZED"
    assert payload["frozen_snapshot_id"] == store.get_run(DATE)["frozen_snapshot_id"]
    assert payload["frozen_row_count"] == 2
    assert payload["forced_incomplete_count"] == 2
    assert payload["remote_incomplete_job_count"] == 1
    assert payload["blocker_count"] == 3
    assert payload["max_incomplete"] == 10
    assert payload["manifest_recorded"] is True
    assert payload["finalized"] is True
    assert Path(payload["output_path"]).is_file()


def test_cli_force_finalize_failure_reports_error_type_and_exit_two(tmp_path, capsys):
    from scripts.incremental_price_audit import main

    _seed_force(tmp_path, freeze=False)

    code = main([
        "force-finalize",
        "--business-date", DATE,
        "--runtime-root", str(tmp_path / "runtime"),
        "--max-incomplete", "10",
    ])

    assert code == 2
    payload = json.loads(capsys.readouterr().out)
    assert payload["error_type"] == "FORCE_FINALIZE_NOT_FROZEN"


# Production-writer mutual exclusion, direction 1: while force-finalize is
# inside its critical section, a real closeout must be refused entry.
def test_closeout_cannot_enter_while_force_finalize_holds_production_lock(tmp_path, monkeypatch):
    store, _snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)
    runtime_root = tmp_path / "runtime"
    lock_path = runtime_root / DATE / "incremental_price_audit.lock"
    original_materialize = LegacyWorkbookMaterializer.materialize
    probe: dict = {}

    def probing_materialize(self, artifact, output_path):
        # Runs inside force-finalize's critical section: readiness is already
        # validated, the final writer is about to publish.  A closeout arriving
        # now must hit the production writer lock and be refused entry.
        result = run_runtime_command("closeout", business_date=DATE, runtime_root=runtime_root)
        probe["outcome"] = result.tick.outcome
        probe["lock_conflict"] = result.lock_conflict
        return original_materialize(self, artifact, output_path)

    monkeypatch.setattr(LegacyWorkbookMaterializer, "materialize", probing_materialize)
    publication = materialize_forced_frozen_daily_run(
        store, ProcessLock(lock_path), DATE, runtime_root / DATE / "formal.xlsx", max_incomplete=10,
    )

    assert publication.forced is True
    assert probe["outcome"] == "ALREADY_RUNNING"
    assert probe["lock_conflict"] is not None
    assert probe["lock_conflict"]["requested"] == "closeout"
    assert store.get_run(DATE)["state"] == "FINALIZED"


# Production-writer mutual exclusion, direction 2: while a closeout owns the
# same per-day production writer lock, force-finalize must refuse to run.
def test_force_finalize_refused_while_closeout_holds_production_lock(tmp_path):
    store, _snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)
    runtime_root = tmp_path / "runtime"
    closeout_lock = ProcessLock(runtime_root / DATE / "incremental_price_audit.lock")
    closeout_lock.acquire(owner={
        "business_date": DATE,
        "pid": os.getpid(),
        "command": "closeout",
        "started_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "invocation": "scheduler",
    })

    result = run_force_finalize_command(
        business_date=DATE, runtime_root=runtime_root, max_incomplete=10,
    )

    assert result["outcome"] == "FORCE_FINALIZE_PRODUCTION_WRITER_ACTIVE"
    assert result["error_type"] == "FORCE_FINALIZE_PRODUCTION_WRITER_ACTIVE"
    assert store.get_run(DATE)["state"] != "FINALIZED"
    assert not (runtime_root / DATE / "saas_audit_final_20260824.xlsx").exists()
    closeout_lock.release()


# Status semantics after a forced finalize: durable crawl/remote work keeps its
# real in-flight states, but status reports the day as forced-finalized and
# never prompts rerunning finalize or delivering the day again.
def test_status_after_forced_finalize_reports_done_without_rerun_prompts(tmp_path):
    store, _snapshot, crawls = _seed_force(tmp_path)
    _link_inflight_job(store, tmp_path, crawls)
    runtime_root = tmp_path / "runtime"
    result = run_force_finalize_command(
        business_date=DATE, runtime_root=runtime_root, max_incomplete=10,
    )
    assert result["outcome"] == "FORCE_FINALIZED"

    status = collect_runtime_status(business_date=DATE, runtime_root=runtime_root)

    assert status["state"] == "FINALIZED"
    assert status["final"]["finalized"] is True
    assert status["final"]["forced"] is True
    assert status["final"]["manifest_recorded"] is True
    assert status["final"]["output_exists"] is True
    assert status["final"]["output_path"].endswith("saas_audit_final_20260824.xlsx")
    # Durable history keeps its real in-flight states...
    assert status["crawl"]["running"] == 2
    assert status["remote"]["awaiting_harvest"] == 1
    # ...while next_actions no longer suggest any rerun/delivery work.
    assert status["next_actions"] == ["none; the frozen daily workbook is finalized"]
