from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.services.daily_audit_store import DailyAuditStore, FrozenSnapshotConflict, ProcessLock
from app.services.incremental_audit_materializer import LegacyWorkbookMaterializer
from app.workflows.incremental_price_audit_finalize import (
    FinalizeBlocked,
    build_frozen_audit_artifact,
    materialize_frozen_review,
    materialize_frozen_daily_run,
)


def _row(v="v-1", sku="sku-1", url="https://shop.example/item", **extra):
    return {
        "vCommodityId": v,
        "documentId": v,
        "skuId": sku,
        "extLink": url,
        "queryStr": "客户标题",
        "sassVCommodityName": "客户标题",
        "skuName": "互联网标题",
        "vendorName": "商城",
        "shopName": "店铺",
        "price": "10",
        "fourCategoryId": "cat-1",
        "fourCategoryName": "分类",
        "list_reference_price": "10",
        "customer_lookup_status": "MATCHED",
        "__crawl_required": bool(url),
        **extra,
    }


def _seed(
    tmp_path: Path,
    rows=None,
    crawl_result=None,
    *,
    crawl_no_result=False,
    crawl_export_no_price=False,
    crawl_pending=False,
):
    rows = list(rows or [_row(), _row()])
    store = DailyAuditStore(tmp_path / "state.sqlite")
    store.ensure_run("2026-08-11", cutoff_at="2026-08-11T10:00:00+00:00")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    snapshot = store.ingest_snapshot(
        token,
        rows,
        "2026-08-11T10:01:00+00:00",
        headers=[name for name in rows[0] if not name.startswith("__")],
        evidence={"observed_count": len(rows)},
    )["snapshot"]
    store.freeze_snapshot(token, snapshot["id"])
    items = store.ordered_snapshot_items(snapshot["id"])
    for item in {item["id"]: item for item in items}.values():
        detail = store.ensure_stage_work(token, item["id"], "DETAIL", f"detail-{item['id']}")
        store.claim_stage_work(token, detail["id"], max_attempts=3, max_elapsed_seconds=300)
        store.complete_stage_work(token, detail["id"], {
            "customerPrice": "10",
            "customerPriceStatus": "FOUND",
            "customerPriceSource": "DETAIL_REFERENCE_INFOS",
            "customerPriceWarning": "",
            "customerPriceEvidence": "customer=10",
        })
        model = store.ensure_stage_work(token, item["id"], "MODEL", f"model-{item['id']}")
        store.claim_stage_work(token, model["id"], max_attempts=1, max_elapsed_seconds=300)
        store.complete_stage_work(token, model["id"], {
            "match_result": "SAME", "confidence": 0.95, "model_match": "exact", "reason": "一致",
        })
    if rows[0].get("__crawl_required"):
        crawl = store.ensure_crawl_work(
            token, "mall", rows[0]["skuId"], rows[0]["extLink"], "legacy-v1", "crawl-fp",
            source_id=rows[0]["vCommodityId"],
        )
        if crawl_no_result:
            job = store.prepare_remote_intent(
                token,
                "mall:hash",
                "daily-20260811-mall",
                "payload-a",
                result_path=str(tmp_path / "mall-result.xlsx"),
            )
            store.claim_remote_submission(token, int(job["id"]))
            store.reconcile_remote_job(
                token,
                int(job["id"]),
                remote_identity="mall|task-id|task-number",
                state="COMPLETED",
                result={"https://other/item": {"price": "12.30"}},
            )
            store.link_crawl_work_to_remote_job(token, int(crawl["id"]), int(job["id"]))
            error = "legacy export missing safe skuId / exact URL result"
            store.fail_crawl_work(token, int(crawl["id"]), error, retryable=False)
            store.complete_crawl_work_without_result(token, int(crawl["id"]), reason=error)
        elif crawl_export_no_price:
            from app.services.incremental_audit_recovery import (
                EXPORT_NO_PRICE_CRAWL_REASON,
                WAIT_EXPORT_NO_EXPORT_DIAGNOSTIC,
            )

            job = store.prepare_remote_intent(
                token,
                "mall:hash",
                "daily-20260811-mall",
                "payload-a",
                result_path=str(tmp_path / "mall-result.xlsx"),
            )
            store.claim_remote_submission(token, int(job["id"]))
            store.reconcile_remote_job(
                token,
                int(job["id"]),
                remote_identity="mall|task-id|task-number",
                state="SUBMITTED",
            )
            store.link_crawl_work_to_remote_job(token, int(crawl["id"]), int(job["id"]))
            store.mark_remote_manual(token, int(job["id"]), WAIT_EXPORT_NO_EXPORT_DIAGNOSTIC)
            store.complete_crawl_work_export_no_price(
                token,
                int(crawl["id"]),
                remote_job_id=int(job["id"]),
                remote_identity="mall|task-id|task-number",
                reason=EXPORT_NO_PRICE_CRAWL_REASON,
                export_rsp_desc="未检索到商品价格信息",
                export_count=0,
            )
            store.complete_manual_export_no_price_job(
                token,
                int(job["id"]),
                remote_identity="mall|task-id|task-number",
                result={
                    "evidence_type": "EXPORT_NO_PRICE_TERMINAL",
                    "export_rspDesc": "未检索到商品价格信息",
                    "export_count": 0,
                },
            )
        elif not crawl_pending:
            store.complete_crawl_work(token, crawl["id"], crawl_result or {
                "page_price_value": "9.5", "page_price_status": "OK", "page_price_source": "招商局",
            })
    store.release_lease(token)
    return store, snapshot


def _writer(path, headers, rows, enriched, audit_response):
    from scripts.audit_saas_excel_file import _format_cell_value

    workbook = Workbook()
    summary = workbook.active
    summary.title = "summary"
    summary.append(["total"])
    summary.append([len(rows)])
    sheet = workbook.create_sheet("record_results")
    sheet.append([*headers, "客户价", "最终决策"])
    for row, info, result in zip(rows, enriched, audit_response.record_results, strict=True):
        sheet.append([
            *(_format_cell_value(row.get(header, "")) for header in headers),
            str(info.get("customerPrice") or ""), result.primary_action.value,
        ])
    workbook.save(path)
    workbook.close()


def test_build_artifact_replays_duplicate_frozen_order_and_precomputed_results(tmp_path):
    store, snapshot = _seed(tmp_path)
    token = store.acquire_lease("2026-08-11", "build", 60)

    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    assert built.readiness.ready is True
    assert [row["vCommodityId"] for row in built.artifact.frozen_rows] == ["v-1", "v-1"]
    assert len(built.artifact.audit_response.record_results) == 2
    assert all(row["__page_price_status"] == "OK" for row in built.artifact.frozen_rows)
    assert built.artifact.frozen_snapshot_id == str(snapshot["id"])
    store.release_lease(token)


def test_review_export_keeps_unresolved_rows_without_finalizing_durable_run(tmp_path):
    store, _snapshot = _seed(tmp_path, crawl_pending=True)
    before = store.get_run("2026-08-11")
    output = tmp_path / "saas_audit_review_20260811.xlsx"

    publication = materialize_frozen_review(
        store,
        ProcessLock(tmp_path / "review.lock"),
        "2026-08-11",
        output,
    )

    after = store.get_run("2026-08-11")
    assert publication.row_count == 2
    assert publication.manual_count == 2
    assert output.is_file()
    assert not output.with_name(output.name + ".manifest.json").exists()
    assert after["state"] == before["state"]
    assert after["frozen_snapshot_id"] == before["frozen_snapshot_id"]
    assert after["final_output_path"] is None
    assert after["final_manifest_hash"] is None

    workbook = load_workbook(output, read_only=False, data_only=True)
    try:
        sheet = workbook["record_results"]
        rows = list(sheet.iter_rows(values_only=True))
        assert workbook["summary"].column_dimensions["A"].width >= 30
        assert workbook["summary"].column_dimensions["B"].width >= 18
    finally:
        workbook.close()
    headers = list(rows[0])
    status_index = headers.index("结果状态")
    price_index = headers.index("价格结果")
    manual_index = headers.index("是否需要人工确认")
    assert [row[status_index] for row in rows[1:]] == ["MANUAL_REVIEW"] * 2
    assert [row[price_index] for row in rows[1:]] == [None, None]
    assert [row[manual_index] for row in rows[1:]] == ["是", "是"]


def test_review_export_requires_completed_detail_evidence(tmp_path):
    store, snapshot = _seed(tmp_path)
    token = store.acquire_lease("2026-08-11", "terminal-detail", 60)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    detail = next(
        row
        for row in store.list_stage_work(token, stage="DETAIL")
        if row["audit_item_id"] == item["id"]
    )
    store.ensure_stage_work(token, item["id"], "DETAIL", detail["fingerprint"] + "-changed")
    assert store.claim_stage_work(
        token, detail["id"], max_attempts=1, max_elapsed_seconds=1
    ) is not None
    store.fail_stage_work(
        token,
        detail["id"],
        elapsed_seconds=2,
        error="TimeoutError",
        retryable=True,
        max_attempts=1,
        max_elapsed_seconds=1,
    )
    store.release_lease(token)

    with pytest.raises(FinalizeBlocked, match="completed Detail and Model evidence"):
        materialize_frozen_review(
            store,
            ProcessLock(tmp_path / "review.lock"),
            "2026-08-11",
            tmp_path / "saas_audit_review_20260811.xlsx",
        )


@pytest.mark.parametrize(
    ("crawl_result", "crawl_no_result", "expected_status", "expected_price"),
    [
        (
            {
                "page_price_value": "0",
                "page_price_status": "PRICE_IS_ZERO",
                "page_price_source": "招商局",
            },
            False,
            "SUCCESS",
            0,
        ),
        (None, True, "NO_RESULT", None),
    ],
)
def test_review_export_distinguishes_real_zero_from_no_result(
    tmp_path,
    crawl_result,
    crawl_no_result,
    expected_status,
    expected_price,
):
    store, _snapshot = _seed(
        tmp_path,
        crawl_result=crawl_result,
        crawl_no_result=crawl_no_result,
    )
    output = tmp_path / "saas_audit_review_20260811.xlsx"

    materialize_frozen_review(
        store,
        ProcessLock(tmp_path / "review.lock"),
        "2026-08-11",
        output,
    )

    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        rows = list(workbook["record_results"].iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = list(rows[0])
    status_index = headers.index("结果状态")
    price_index = headers.index("价格结果")
    assert [row[status_index] for row in rows[1:]] == [expected_status] * 2
    assert [row[price_index] for row in rows[1:]] == [expected_price] * 2


def test_review_export_reports_conflicting_prices_without_selecting_one(tmp_path):
    store, _snapshot = _seed(
        tmp_path,
        rows=[
            _row("v-1", "sku-1", "https://shop.example/item"),
            _row("v-2", "sku-2", "https://shop.example/other"),
        ],
        crawl_pending=True,
    )
    result_path = tmp_path / "conflicting-result.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品编码", "商品最小品类单元id", "商品详情地址", "商品价格"])
    sheet.append(["sku-1", "sku-1", "https://shop.example/item", 1139])
    sheet.append(["reference-row", "sku-1", "https://shop.example/item", 692.22])
    sheet.append(["sku-2", "sku-2", "https://shop.example/other", 20])
    workbook.save(result_path)
    workbook.close()
    token = store.acquire_lease("2026-08-11", "link-manual", 60)
    first_crawl = store.list_crawl_work(token)[0]
    second_crawl = store.ensure_crawl_work(
        token,
        "mall",
        "sku-2",
        "https://shop.example/other",
        "legacy-v1",
        "crawl-fp-2",
        source_id="v-2",
    )
    job = store.prepare_remote_intent(
        token,
        "mall:manual",
        "task-manual",
        "payload-manual",
        result_path=str(result_path),
    )
    store.claim_remote_submission(token, int(job["id"]))
    store.reconcile_remote_job(
        token,
        int(job["id"]),
        remote_identity="mall|task-id|task-number",
        state="SUBMITTED",
    )
    store.link_crawl_work_to_remote_job(token, int(first_crawl["id"]), int(job["id"]))
    store.link_crawl_work_to_remote_job(token, int(second_crawl["id"]), int(job["id"]))
    store.mark_remote_manual(token, int(job["id"]), "existing result workbook is invalid")
    store.release_lease(token)
    output = tmp_path / "saas_audit_review_20260811.xlsx"

    materialize_frozen_review(
        store,
        ProcessLock(tmp_path / "review.lock"),
        "2026-08-11",
        output,
    )

    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        rows = list(workbook["record_results"].iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = list(rows[0])
    reason_index = headers.index("异常原因")
    price_index = headers.index("价格结果")
    assert [row[reason_index] for row in rows[1:]] == [
        "同一 exact identity 存在冲突价格：1139.00 / 692.22"
    ] * 2
    assert [row[price_index] for row in rows[1:]] == [None, None]


def test_model_pending_and_unplanned_required_crawl_block_official_output(tmp_path):
    store, snapshot = _seed(tmp_path, rows=[_row("v-1", "s-1"), _row("v-2", "s-2", "https://other/item")])
    token = store.acquire_lease("2026-08-11", "break", 60)
    second = store.ordered_snapshot_items(snapshot["id"])[1]
    model = next(row for row in store.list_stage_work(token, stage="MODEL") if row["audit_item_id"] == second["id"])
    # A changed fingerprint resets the second model to durable PENDING.
    store.ensure_stage_work(token, second["id"], "MODEL", model["fingerprint"] + "-changed")

    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    assert built.readiness.ready is False
    assert built.readiness.model_pending == 1
    assert built.readiness.crawl_unplanned == 1
    store.release_lease(token)


def test_detail_terminal_blocks_ordinary_finalize_and_counts_failed_price(tmp_path):
    """Flipped with the customer-price gates: terminal DETAIL is no longer
    materializable; the row's price evidence classifies as customer_price_failed."""
    store, snapshot = _seed(tmp_path, crawl_result={
        "page_price_value": None, "page_price_status": "PRICE_MISSING", "page_price_source": "招商局-价格不可见",
    })
    token = store.acquire_lease("2026-08-11", "terminal", 60)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    detail = next(row for row in store.list_stage_work(token, stage="DETAIL") if row["audit_item_id"] == item["id"])
    store.ensure_stage_work(token, item["id"], "DETAIL", detail["fingerprint"] + "-changed")
    claimed = store.claim_stage_work(token, detail["id"], max_attempts=1, max_elapsed_seconds=1)
    assert claimed is not None
    store.fail_stage_work(
        token, detail["id"], elapsed_seconds=2, error="TimeoutError", retryable=True,
        max_attempts=1, max_elapsed_seconds=1,
    )
    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    assert built.readiness.detail_terminal >= 1
    assert built.readiness.customer_price_failed >= 1
    assert built.readiness.crawl_invalid_price >= 1
    assert built.readiness.ready is False
    store.release_lease(token)


def test_detail_terminal_keeps_reference_price_blank_but_blocks_final(tmp_path):
    """Flipped with the customer-price gates: the terminal DETAIL row keeps its
    blank reference price, but the day no longer finalizes without it."""
    store, snapshot = _seed(tmp_path, rows=[_row(url="")])
    token = store.acquire_lease("2026-08-11", "terminal-detail", 60)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    detail = next(
        row for row in store.list_stage_work(token, stage="DETAIL")
        if row["audit_item_id"] == item["id"]
    )
    detail = store.ensure_stage_work(
        token, item["id"], "DETAIL", detail["fingerprint"] + "-changed"
    )
    claimed = store.claim_stage_work(
        token, detail["id"], max_attempts=1, max_elapsed_seconds=1
    )
    assert claimed is not None
    store.fail_stage_work(
        token,
        detail["id"],
        elapsed_seconds=2,
        error="TimeoutError",
        retryable=True,
        max_attempts=1,
        max_elapsed_seconds=1,
    )

    built = build_frozen_audit_artifact(
        store, token, output_identity="daily-final-20260811"
    )

    assert built.readiness.detail_terminal == 1
    assert built.readiness.customer_price_failed == 1
    assert built.readiness.ready is False
    assert built.artifact is None
    store.release_lease(token)


def test_explicit_no_result_is_materializable_without_faking_zero_price(tmp_path):
    store, _snapshot = _seed(tmp_path, crawl_no_result=True)
    token = store.acquire_lease("2026-08-11", "no-result", 60)

    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    assert built.readiness.ready is True
    assert built.readiness.crawl_no_result == 2
    assert built.readiness.crawl_terminal == 0
    assert built.artifact is not None
    assert all(row["__page_price_value"] is None for row in built.artifact.frozen_rows)
    assert all(row["__page_price_status"] == "NO_RESULT" for row in built.artifact.frozen_rows)
    assert all(row["__page_price_source"] == "招商局（无采集结果）" for row in built.artifact.frozen_rows)
    store.release_lease(token)


def test_export_no_price_terminal_no_result_is_materializable(tmp_path):
    store, _snapshot = _seed(tmp_path, crawl_export_no_price=True)
    token = store.acquire_lease("2026-08-11", "export-no-price", 60)

    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    assert built.readiness.ready is True
    assert built.readiness.crawl_no_result == 2
    assert built.readiness.crawl_invalid_price == 0
    assert built.readiness.remote_manual == 0
    assert built.artifact is not None
    assert all(row["__page_price_value"] is None for row in built.artifact.frozen_rows)
    assert all(row["__page_price_status"] == "NO_RESULT" for row in built.artifact.frozen_rows)
    assert all(row["__page_price_source"] == "招商局（无采集结果）" for row in built.artifact.frozen_rows)
    store.release_lease(token)


def test_valid_no_result_accepts_both_terminal_evidence_types():
    from app.workflows.incremental_price_audit_finalize import _valid_no_result

    exported = {
        "page_price_status": "NO_RESULT",
        "page_price_value": None,
        "evidence_type": "EXPORTED_NO_MATCHING_ROW",
        "no_result_reason": "legacy export missing safe skuId / exact URL result",
    }
    assert _valid_no_result(exported) is True

    export_no_price = {
        "page_price_status": "NO_RESULT",
        "page_price_value": None,
        "evidence_type": "EXPORT_NO_PRICE_TERMINAL",
        "no_result_reason": "remote task finished and export endpoint reported no price rows",
        "export_rspDesc": "未检索到商品价格信息",
        "export_count": 0,
    }
    assert _valid_no_result(export_no_price) is True

    assert _valid_no_result({**export_no_price, "export_count": 3}) is False
    assert _valid_no_result({**export_no_price, "export_rspDesc": "导出失败"}) is False
    assert _valid_no_result({**export_no_price, "export_count": None}) is False
    assert _valid_no_result({**export_no_price, "page_price_value": "0"}) is False
    assert _valid_no_result({**export_no_price, "no_result_reason": ""}) is False


def test_workbook_display_name_for_no_result_is_explicit():
    from scripts.audit_saas_excel_file import _format_cell_value

    assert _format_cell_value("NO_RESULT") == "无采集结果"


def test_materialize_marks_run_finalized_and_same_finalize_reuses(tmp_path):
    store, _snapshot = _seed(tmp_path, rows=[_row(url="")])
    output = tmp_path / "formal.xlsx"
    materializer = LegacyWorkbookMaterializer(writer=_writer)

    first = materialize_frozen_daily_run(
        store,
        ProcessLock(tmp_path / "finalize.lock"),
        "2026-08-11",
        output,
        materializer=materializer,
    )
    second = materialize_frozen_daily_run(
        store,
        ProcessLock(tmp_path / "finalize.lock"),
        "2026-08-11",
        output,
        materializer=materializer,
    )

    assert first.materialization is not None and first.materialization.reused is False
    assert second.materialization is not None and second.materialization.reused is True
    run = store.get_run("2026-08-11")
    assert run["state"] == "FINALIZED"
    assert run["final_output_path"] == str(output.resolve())
    assert json.loads(run["metrics_json"])["finalized_rows"] == 1


def test_finalized_run_rejects_different_output_before_materializer_writes(tmp_path):
    store, _snapshot = _seed(tmp_path, rows=[_row(url="")])
    first_output = tmp_path / "formal.xlsx"
    conflicting_output = tmp_path / "other.xlsx"
    materializer = LegacyWorkbookMaterializer(writer=_writer)

    materialize_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), "2026-08-11", first_output,
        materializer=materializer,
    )

    with pytest.raises(FrozenSnapshotConflict):
        materialize_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), "2026-08-11", conflicting_output,
            materializer=materializer,
        )

    assert not conflicting_output.exists()
    assert not conflicting_output.with_name(conflicting_output.name + ".manifest.json").exists()


def test_crash_after_publish_before_store_mark_recovers_idempotently(tmp_path, monkeypatch):
    store, _snapshot = _seed(tmp_path, rows=[_row(url="")])
    output = tmp_path / "formal.xlsx"
    materializer = LegacyWorkbookMaterializer(writer=_writer)
    original = store.mark_finalized
    calls = 0

    def crash_once(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise SystemExit("crash before DB mark")
        return original(*args, **kwargs)

    monkeypatch.setattr(store, "mark_finalized", crash_once)
    with pytest.raises(SystemExit):
        materialize_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), "2026-08-11", output,
            materializer=materializer,
        )
    published = output.read_bytes()
    interrupted = store.get_run("2026-08-11")
    assert interrupted["state"] == "FINALIZING"
    assert interrupted["final_output_path"] == str(output.resolve())

    recovered = materialize_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), "2026-08-11", output,
        materializer=materializer,
    )

    assert recovered.materialization is not None and recovered.materialization.reused is True
    assert output.read_bytes() == published
    assert store.get_run("2026-08-11")["state"] == "FINALIZED"


def test_blocked_readiness_never_writes_formal_workbook(tmp_path):
    store, snapshot = _seed(tmp_path, rows=[_row(url="")])
    token = store.acquire_lease("2026-08-11", "reset", 60)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    model = store.list_stage_work(token, stage="MODEL")[0]
    store.ensure_stage_work(token, item["id"], "MODEL", model["fingerprint"] + "-new")
    store.release_lease(token)
    output = tmp_path / "formal.xlsx"

    with pytest.raises(FinalizeBlocked):
        materialize_frozen_daily_run(
            store, ProcessLock(tmp_path / "finalize.lock"), "2026-08-11", output,
            materializer=LegacyWorkbookMaterializer(writer=_writer),
        )

    assert not output.exists()


# ---- Customer-price gates (plan §6, matrix D) ------------------------------


def _rework_detail(store, token, item_id, *, result=None, error=None):
    """Give the item's DETAIL row fresh evidence (business gate test helper)."""
    detail = next(
        row for row in store.list_stage_work(token, stage="DETAIL")
        if row["audit_item_id"] == item_id
    )
    detail = store.ensure_stage_work(
        token, item_id, "DETAIL", detail["fingerprint"] + "-rework"
    )
    claimed = store.claim_stage_work(
        token, detail["id"], max_attempts=1, max_elapsed_seconds=300
    )
    assert claimed is not None
    if error is not None:
        store.fail_stage_work(
            token, detail["id"], elapsed_seconds=1, error=error, retryable=True,
            max_attempts=1, max_elapsed_seconds=1,
        )
    else:
        store.complete_stage_work(token, detail["id"], result)


NOT_FOUND_ERROR = "customerPriceStatus=NOT_FOUND；详情接口未返回有效参考价"


def test_customer_price_pending_blocks_finalize(tmp_path):
    store, snapshot = _seed(tmp_path, rows=[_row(url="")])
    token = store.acquire_lease("2026-08-11", "pending", 60)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    detail = next(row for row in store.list_stage_work(token, stage="DETAIL") if row["audit_item_id"] == item["id"])
    # A changed fingerprint resets the row to durable PENDING.
    store.ensure_stage_work(token, item["id"], "DETAIL", detail["fingerprint"] + "-pending")

    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    assert built.readiness.customer_price_pending == 1
    assert built.readiness.detail_pending == 1
    assert built.readiness.ready is False
    store.release_lease(token)


def test_customer_price_missing_after_not_found_exhaustion_blocks_finalize(tmp_path):
    store, snapshot = _seed(tmp_path, rows=[_row(url="")])
    token = store.acquire_lease("2026-08-11", "missing", 60)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    _rework_detail(store, token, item["id"], error=NOT_FOUND_ERROR)

    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    assert built.readiness.customer_price_missing == 1
    assert built.readiness.detail_terminal == 1
    assert built.readiness.ready is False
    assert built.artifact is None
    store.release_lease(token)


def test_legacy_success_without_price_blocks_finalize(tmp_path):
    """Old invalid-success rows (SUCCESS + status failed + price None) block."""
    store, snapshot = _seed(tmp_path, rows=[_row(url="")])
    token = store.acquire_lease("2026-08-11", "legacy", 60)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    _rework_detail(store, token, item["id"], result={
        "customerPrice": None,
        "customerPriceStatus": "DETAIL_REQUEST_FAILED",
        "customerPriceSource": "",
        "customerPriceWarning": "详情接口业务失败：rspCode=500",
        "customerPriceEvidence": "",
    })

    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    assert built.readiness.customer_price_failed == 1
    assert built.readiness.ready is False
    assert built.artifact is None
    store.release_lease(token)


def test_zero_price_success_counts_invalid_and_blocks_finalize(tmp_path):
    store, snapshot = _seed(tmp_path, rows=[_row(url="")])
    token = store.acquire_lease("2026-08-11", "invalid", 60)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    _rework_detail(store, token, item["id"], result={
        "customerPrice": "0",
        "customerPriceStatus": "FOUND",
        "customerPriceSource": "DETAIL_REFERENCE_INFOS",
        "customerPriceWarning": "",
        "customerPriceEvidence": "参考价1=0",
    })

    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    assert built.readiness.customer_price_invalid == 1
    assert built.readiness.ready is False
    store.release_lease(token)


def test_valid_customer_price_keeps_finalize_ready(tmp_path):
    store, _snapshot = _seed(tmp_path, rows=[_row(url="")])
    token = store.acquire_lease("2026-08-11", "valid", 60)

    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    assert built.readiness.ready is True
    assert built.readiness.customer_price_missing == 0
    assert built.readiness.customer_price_failed == 0
    assert built.readiness.customer_price_invalid == 0
    assert built.readiness.customer_price_absent_proven == 0
    assert built.artifact is not None
    assert built.artifact.enriched_rows[0]["customerPrice"] is not None
    store.release_lease(token)


@pytest.mark.parametrize("status", ["LOOKUP_FAILED", "DETAIL_REQUEST_SKIPPED", "NOT_PROVIDED"])
def test_proven_absent_customer_price_blocks_finalize(tmp_path, status):
    """Round-2 fail-closed ruling: a required row with no valid customer price
    blocks formal finalization even when the absence is deterministic and
    evidence-complete.  (Flipped from the round-1 'does not block' contract.)"""
    store, snapshot = _seed(tmp_path, rows=[_row(url="")])
    token = store.acquire_lease("2026-08-11", "absent", 60)
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    _rework_detail(store, token, item["id"], result={
        "customerPrice": None,
        "customerPriceStatus": status,
        "customerPriceSource": "",
        "customerPriceWarning": "authoritative customer export has no such document",
        "customerPriceEvidence": "",
    })

    built = build_frozen_audit_artifact(store, token, output_identity="daily-final-20260811")

    # The enriched evidence is preserved for review, but the formal final is blocked.
    assert built.readiness.ready is False
    assert built.readiness.customer_price_absent_proven == 1
    assert built.artifact is None
    store.release_lease(token)


def test_detail_failure_keeps_decoded_status_from_error_prefix():
    from app.workflows.incremental_price_audit_finalize import _detail_failure

    decoded = _detail_failure(NOT_FOUND_ERROR)
    assert decoded["customerPriceStatus"] == "NOT_FOUND"
    assert decoded["customerPrice"] is None
    assert "详情接口未返回有效参考价" in decoded["customerPriceWarning"]

    fallback = _detail_failure("TimeoutError")
    assert fallback["customerPriceStatus"] == "DETAIL_REQUEST_FAILED"
    assert fallback["customerPriceWarning"].startswith("增量详情阶段失败：TimeoutError")


# ---- is_recoverable_finalize_wait (round 2, matrix E) ----------------------


def test_recoverable_finalize_wait_is_true_for_pure_pending_only():
    from app.workflows.incremental_price_audit_finalize import FinalizeReadiness, is_recoverable_finalize_wait

    assert is_recoverable_finalize_wait(None) is False
    assert is_recoverable_finalize_wait(FinalizeReadiness()) is False
    assert is_recoverable_finalize_wait(FinalizeReadiness(detail_pending=1)) is True
    assert is_recoverable_finalize_wait(FinalizeReadiness(crawl_pending=2, remote_pending=1)) is True


@pytest.mark.parametrize(
    "blocker",
    [
        {"detail_terminal": 1},
        {"customer_price_missing": 1},
        {"customer_price_failed": 1},
        {"customer_price_invalid": 1},
        {"customer_price_absent_proven": 1},
        {"model_terminal": 1},
        {"crawl_terminal": 1},
        {"remote_manual": 1},
        {"remote_invalid": 1},
        {"crawl_unplanned": 1},
        {"crawl_invalid_price": 1},
    ],
)
def test_actionable_blockers_are_never_a_recoverable_wait(blocker):
    """Any actionable blocker (including every customer_price_* bucket and
    detail_terminal) must stop the closeout from idling as if it were only a
    pending wait."""
    from app.workflows.incremental_price_audit_finalize import FinalizeReadiness, is_recoverable_finalize_wait

    readiness = FinalizeReadiness(detail_pending=1, **blocker)
    assert is_recoverable_finalize_wait(readiness) is False


def test_absent_proven_alone_is_an_actionable_blocker_not_a_wait():
    """Even without pending work, absent_proven must read as not-recoverable:
    the closeout must not idle-wait on a gate only a human can resolve."""
    from app.workflows.incremental_price_audit_finalize import FinalizeReadiness, is_recoverable_finalize_wait

    assert is_recoverable_finalize_wait(FinalizeReadiness(customer_price_absent_proven=1)) is False
