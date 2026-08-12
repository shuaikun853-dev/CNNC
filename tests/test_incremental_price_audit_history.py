"""Read-only regression evidence from the fixed July-August 2026 snapshots."""
from __future__ import annotations

from collections import Counter, defaultdict
import hashlib
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.daily_audit_store import DailyAuditStore, ProcessLock
from app.services.incremental_audit_materializer import LegacyWorkbookMaterializer
from app.services.saas_audit_core import (
    build_customer_index,
    customer_fingerprint,
    detail_fingerprint,
    join_saas_with_customer,
    model_fingerprint,
    normalize_cell_str,
    read_saas_pending_json,
    source_fingerprint,
)
from app.services import saas_audit_core as audit_core
from app.workflows.incremental_price_audit import (
    CrawlMember,
    CrawlPlan,
    DiscoveredState,
    TickConfig,
    TickDependencies,
    run_incremental_tick,
)
from app.workflows.incremental_price_audit_finalize import (
    build_frozen_audit_artifact,
    materialize_frozen_daily_run,
)
import scripts.audit_saas_excel_file as legacy_facade


RUNTIME_ROOT = Path("runtime")

# This manifest is deliberately fixed; later runtime artifacts must not change
# the historical evidence used to justify the identity/fingerprint design.
HISTORY_SNAPSHOTS = (
    "price_audit_20260713_162725", "price_audit_20260714_103026",
    "price_audit_20260714_104455", "price_audit_20260714_104718",
    "price_audit_20260714_113217", "price_audit_20260714_113444",
    "price_audit_20260714_144431", "price_audit_20260714_145856",
    "price_audit_20260714_173317", "price_audit_20260715_135858",
    "price_audit_20260715_170347", "price_audit_20260715_194302",
    "price_audit_20260716_162220", "price_audit_20260716_180728",
    "price_audit_20260717_183234", "price_audit_20260717_183356_514814",
    "price_audit_20260721_222632", "price_audit_20260723_141533",
    "price_audit_20260723_173442", "price_audit_20260724_195418",
    "price_audit_20260724_203830", "price_audit_20260727_182731",
    "price_audit_20260728_175338", "price_audit_20260729_181034",
    "price_audit_20260730_175552", "price_audit_20260731_231019",
    "price_audit_20260805_181050", "price_audit_20260806_202120",
    "price_audit_20260808_185049", "price_audit_20260810_175003",
    "price_audit_20260811_180427",
)

FINAL_EQUIVALENCE = (
    ("price_audit_20260714_145856", "saas_current.json", "saas_audit_final_20260714_145856.xlsx"),
    ("price_audit_20260714_173317", "saas_current.json", "saas_audit_final_20260714_173317.xlsx"),
    ("price_audit_20260715_135858", "saas_current.json", "saas_audit_final_20260715_135858.xlsx"),
    ("price_audit_20260715_170347", "saas_json.json", "saas_audit_final_20260715_170347.xlsx"),
    ("price_audit_20260715_194302", "saas_json.json", "saas_audit_final_20260715_194302.xlsx"),
    ("price_audit_20260716_162220", "saas_json.json", "saas_audit_final_20260716_162220.xlsx"),
    ("price_audit_20260716_180728", "saas_json.json", "saas_audit_final_20260716_180728.xlsx"),
    ("price_audit_20260721_222632", "saas_json.json", "saas_audit_final_20260721_222632.xlsx"),
    ("price_audit_20260723_141533", "saas_json.json", "saas_audit_final_20260723_141533.xlsx"),
    ("price_audit_20260723_173442", "saas_json.json", "saas_audit_final_20260723_173442.xlsx"),
    ("price_audit_20260724_195418", "saas_json.json", "saas_audit_final_20260724_195418.xlsx"),
    ("price_audit_20260727_182731", "saas_json.json", "saas_audit_final_20260727_182731.xlsx"),
    ("price_audit_20260731_231019", "saas_json.json", "saas_audit_final_20260731_231019.xlsx"),
    ("price_audit_20260805_181050", "saas_json.json", "saas_audit_final_20260805_181050.xlsx"),
    ("price_audit_20260806_202120", "saas_json.json", "saas_audit_final_20260806_202120.xlsx"),
    ("price_audit_20260810_175003", "saas_json.json", "saas_audit_final_20260810_175003.xlsx"),
    ("price_audit_20260811_180427", "saas_json.json", "saas_audit_final_20260811_180427.xlsx"),
)


def _require(paths: list[Path]) -> None:
    missing = [str(path) for path in paths if not path.exists()]
    if not paths or missing:
        pytest.skip("fixed historical fixture is unavailable")


def _snapshot_path(batch: str) -> Path:
    directory = RUNTIME_ROOT / batch
    for name in ("saas_json.json", "saas_current.json"):
        path = directory / name
        if path.exists():
            return path
    return directory / "saas_json.json"


def _customer_path(batch: str) -> Path | None:
    directory = RUNTIME_ROOT / batch
    for name in ("customer_json.json", "customer_current.json"):
        path = directory / name
        if path.exists():
            return path
    return None


def _key(row: dict[str, object]) -> tuple[str, str, str]:
    return (
        normalize_cell_str(row.get("vCommodityId")),
        normalize_cell_str(row.get("skuId")),
        normalize_cell_str(row.get("extLink")),
    )


def _excel_keys(path: Path) -> list[tuple[str, str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["record_results"]
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator)]
        columns = [headers.index(legacy_facade._display_header(name)) for name in ("vCommodityId", "skuId", "extLink")]
        return [tuple(normalize_cell_str(row[index]) for index in columns) for row in iterator]
    finally:
        workbook.close()


def test_fixed_history_recomputes_identity_evidence():
    paths = [_snapshot_path(batch) for batch in HISTORY_SNAPSHOTS]
    _require(paths)
    snapshots = [read_saas_pending_json(path)[1] for path in paths]
    rows = [row for snapshot in snapshots for row in snapshot]
    nonempty_urls = {row["extLink"] for row in rows if row["extLink"]}
    duplicate_groups = sum(
        1
        for snapshot in snapshots
        for count in Counter((row["vCommodityId"], row["skuId"]) for row in snapshot).values()
        if count > 1
    )
    created_before_scheduler_day = sum(
        1
        for batch, snapshot in zip(HISTORY_SNAPSHOTS, snapshots, strict=True)
        for row in snapshot
        if normalize_cell_str(row["createTime"]).replace("-", "")[:8] < batch.removeprefix("price_audit_")[:8]
    )
    adjacent_changed_source_sku_sets = 0
    for index, (before, after) in enumerate(zip(snapshots, snapshots[1:], strict=False)):
        before_day = HISTORY_SNAPSHOTS[index].removeprefix("price_audit_")[:8]
        after_day = HISTORY_SNAPSHOTS[index + 1].removeprefix("price_audit_")[:8]
        if before_day != after_day:
            continue
        before_sets: dict[str, set[str]] = defaultdict(set)
        after_sets: dict[str, set[str]] = defaultdict(set)
        for row in before:
            before_sets[row["vCommodityId"]].add(row["skuId"])
        for row in after:
            after_sets[row["vCommodityId"]].add(row["skuId"])
        adjacent_changed_source_sku_sets += sum(
            before_sets[source_id] != after_sets[source_id]
            for source_id in before_sets.keys() & after_sets.keys()
        )
    skus_by_source: dict[str, set[str]] = defaultdict(set)
    sources_by_sku: dict[str, set[str]] = defaultdict(set)
    sources_by_url: dict[str, set[str]] = defaultdict(set)
    skus_by_url: dict[str, set[str]] = defaultdict(set)
    urls_by_relationship: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in rows:
        source_id = row["vCommodityId"]
        sku_id = row["skuId"]
        exact_url = row["extLink"]
        skus_by_source[source_id].add(sku_id)
        sources_by_sku[sku_id].add(source_id)
        if exact_url:
            sources_by_url[exact_url].add(source_id)
            skus_by_url[exact_url].add(sku_id)
            urls_by_relationship[(source_id, sku_id)].add(exact_url)
    assert len(rows) == 10_014
    assert sum(not row["vCommodityId"] for row in rows) == 0
    assert sum(not row["skuId"] for row in rows) == 0
    assert sum(not row["extLink"] for row in rows) == 8
    assert len({row["vCommodityId"] for row in rows}) == 1_218
    assert len({row["skuId"] for row in rows}) == 7_045
    assert len(nonempty_urls) == 7_025
    assert len({(row["vCommodityId"], row["skuId"]) for row in rows}) == 7_400
    assert duplicate_groups == 47
    assert created_before_scheduler_day == 2_342
    assert adjacent_changed_source_sku_sets == 50
    assert sum(len(values) > 1 for values in skus_by_source.values()) == 696
    assert sum(len(values) > 1 for values in sources_by_sku.values()) == 307
    assert sum(len(values) > 1 for values in sources_by_url.values()) == 304
    assert sum(len(values) > 1 for values in skus_by_url.values()) == 14
    assert sum(len(values) > 1 for values in urls_by_relationship.values()) == 0


def test_fixed_history_recomputes_customer_document_relationship_evidence():
    paired = [
        (batch, _snapshot_path(batch), customer_path)
        for batch in HISTORY_SNAPSHOTS
        if (customer_path := _customer_path(batch)) is not None
    ]
    _require([path for _batch, saas_path, customer_path in paired for path in (saas_path, customer_path)])

    customer_rows_total = 0
    customer_document_ids: set[str] = set()
    customer_document_missing = 0
    paired_saas_total = 0
    matched = 0
    reference_id_matches_sku = 0
    reference_link_matches_url = 0
    for _batch, saas_path, customer_path in paired:
        raw_customer_rows = audit_core._collect_rows_from_segments(
            audit_core._load_json_file_segments(customer_path)
        )
        customer_rows_total += len(raw_customer_rows)
        customer_by_document: dict[str, dict[str, object]] = {}
        for row in raw_customer_rows:
            document_id = normalize_cell_str(row.get("documentId"))
            customer_document_missing += int(not document_id)
            if document_id:
                customer_document_ids.add(document_id)
                customer_by_document[document_id] = row

        saas_rows = read_saas_pending_json(saas_path)[1]
        paired_saas_total += len(saas_rows)
        for row in saas_rows:
            customer = customer_by_document.get(row["vCommodityId"])
            if customer is None:
                continue
            matched += 1
            reference_id_matches_sku += int(
                normalize_cell_str(customer.get("referenceId")) == row["skuId"]
            )
            reference_link_matches_url += int(
                normalize_cell_str(customer.get("referenceLink")) == row["extLink"]
            )

    all_saas_rows = [
        row
        for batch in HISTORY_SNAPSHOTS
        for row in read_saas_pending_json(_snapshot_path(batch))[1]
    ]
    assert len(paired) == 23
    assert customer_rows_total == 5_837
    assert customer_document_missing == 0
    assert len(customer_document_ids) == 2_516
    assert paired_saas_total == 7_792
    assert matched == 6_624
    assert paired_saas_total - matched == 1_168
    assert reference_id_matches_sku == 0
    assert reference_link_matches_url == 83
    assert sum(not row["documentId"] for row in all_saas_rows) == 10_014


def test_fixed_history_proves_cutoff_is_post_18_complete_snapshot_not_create_time_filter():
    before_path = RUNTIME_ROOT / "price_audit_20260810_175003" / "superseded_332" / "saas_json.json"
    after_path = RUNTIME_ROOT / "price_audit_20260810_175003" / "saas_json.json"
    july_path = RUNTIME_ROOT / "price_audit_20260724_195418" / "saas_json.json"
    _require([before_path, after_path, july_path])

    before = read_saas_pending_json(before_path)[1]
    after = read_saas_pending_json(after_path)[1]
    existing = Counter(_key(row) for row in before)
    delta: list[dict[str, object]] = []
    for row in after:
        key = _key(row)
        if existing[key]:
            existing[key] -= 1
        else:
            delta.append(row)
    july = read_saas_pending_json(july_path)[1]

    assert (len(before), len(after), len(delta)) == (332, 401, 69)
    assert sum(normalize_cell_str(row.get("createTime"))[11:19] >= "18:00:00" for row in delta) == 3
    assert len(july) == 84
    assert sum(normalize_cell_str(row.get("createTime"))[11:19] >= "18:00:00" for row in july) == 41


def test_fixed_legacy_finals_equal_frozen_ordered_snapshot_members(tmp_path):
    fixture_paths = [RUNTIME_ROOT / batch / name for batch, _, name in FINAL_EQUIVALENCE]
    source_paths = [RUNTIME_ROOT / batch / name for batch, name, _ in FINAL_EQUIVALENCE]
    _require(fixture_paths + source_paths)
    total_rows = 0
    for index, (batch, source_name, final_name) in enumerate(FINAL_EQUIVALENCE, start=1):
        source_path = RUNTIME_ROOT / batch / source_name
        _, rows = read_saas_pending_json(source_path)
        business_date = batch.removeprefix("price_audit_")[:8]
        business_date = f"{business_date[:4]}-{business_date[4:6]}-{business_date[6:]}"
        store = DailyAuditStore(tmp_path / f"history_{index}.sqlite")
        store.ensure_run(business_date)
        token = store.acquire_lease(business_date, owner=f"history-{index}", ttl_seconds=60)
        ingested = store.ingest_snapshot(
            token,
            rows,
            source_fingerprints=[source_fingerprint(row) for row in rows],
            customer_fingerprints=["history-no-customer"] * len(rows),
        )
        snapshot_id = int(ingested["snapshot"]["id"])
        store.freeze_snapshot(token, snapshot_id)
        frozen_keys = [_key(item["row"]) for item in store.ordered_snapshot_items(snapshot_id)]
        legacy_keys = _excel_keys(RUNTIME_ROOT / batch / final_name)
        assert frozen_keys == legacy_keys
        assert Counter(frozen_keys) == Counter(legacy_keys)
        total_rows += len(frozen_keys)
        store.release_lease(token)
    assert len(FINAL_EQUIVALENCE) == 17
    assert total_rows == 5_718


def _sheet_values(path: Path, name: str) -> list[tuple[object, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [tuple(row) for row in workbook[name].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _summary_business_projection(path: Path) -> list[list[object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return LegacyWorkbookMaterializer()._summary_business_projection(workbook["summary"])
    finally:
        workbook.close()


def test_real_history_multitick_incremental_materializes_like_legacy_writer(tmp_path):
    """Compare the shared business artifact through both workbook writers."""
    batch = "price_audit_20260811_180427"
    saas_path = RUNTIME_ROOT / batch / "saas_json.json"
    customer_path = RUNTIME_ROOT / batch / "customer_json.json"
    _require([saas_path, customer_path])
    headers, saas_rows = read_saas_pending_json(saas_path)
    customer = build_customer_index(customer_path)
    # Keep the regression quick while forcing the real cross-source identity
    # edges into the sample rather than relying on an unrepresentative prefix.
    skus_by_url: dict[str, set[str]] = defaultdict(set)
    sources_by_sku_url: dict[tuple[str, str], set[str]] = defaultdict(set)
    duplicate_relationships = Counter((row["vCommodityId"], row["skuId"]) for row in saas_rows)
    for row in saas_rows:
        if row["extLink"]:
            skus_by_url[row["extLink"]].add(row["skuId"])
            sources_by_sku_url[(row["skuId"], row["extLink"])].add(row["vCommodityId"])
    edge_indices = {
        index
        for index, row in enumerate(saas_rows)
        if duplicate_relationships[(row["vCommodityId"], row["skuId"])] > 1
        or len(skus_by_url[row["extLink"]]) > 1
        or len(sources_by_sku_url[(row["skuId"], row["extLink"])]) > 1
    }
    selected_indices = sorted(list(edge_indices) + [
        index for index in range(len(saas_rows)) if index not in edge_indices
    ][: 80 - len(edge_indices)])
    assert len(selected_indices) == 80
    assert edge_indices <= set(selected_indices)
    saas_rows = [saas_rows[index] for index in selected_indices]
    rows = join_saas_with_customer(saas_rows, customer)
    source_fps = [source_fingerprint(row) for row in saas_rows]
    customer_fps = [customer_fingerprint(customer.get(row["vCommodityId"], {})) for row in saas_rows]
    prefixes = (rows[:20], rows[:45], rows[:65], rows)
    prefix_source_fps = (source_fps[:20], source_fps[:45], source_fps[:65], source_fps)
    prefix_customer_fps = (customer_fps[:20], customer_fps[:45], customer_fps[:65], customer_fps)
    observed = iter((
        datetime(2026, 8, 11, 8, 0, tzinfo=UTC),
        datetime(2026, 8, 11, 11, 0, tzinfo=UTC),
        datetime(2026, 8, 11, 14, 0, tzinfo=UTC),
        datetime(2026, 8, 11, 18, 1, tzinfo=UTC),
    ))
    snapshots = iter(zip(prefixes, prefix_source_fps, prefix_customer_fps, strict=True))
    detail_calls: list[str] = []
    model_requests: list[list[str]] = []
    posted_identities: list[tuple[str, str, str, str, str]] = []

    def discover(_timeout: float) -> DiscoveredState:
        snapshot_rows, snapshot_sources, snapshot_customers = next(snapshots)
        return DiscoveredState(
            observed_at=next(observed), headers=headers, rows=snapshot_rows,
            source_fingerprints=snapshot_sources, customer_fingerprints=snapshot_customers,
            evidence={"fixture": batch},
        )

    def detail(row, _timeout):
        detail_calls.append(str(row["vCommodityId"]))
        return {
            "customerPrice": "10", "customerPriceStatus": "FIXTURE",
            "customerPriceSource": "history-fixture", "customerPriceWarning": "", "customerPriceEvidence": "fixture",
        }

    def model_batch(batch_rows, _timeout):
        model_requests.append([str(row["skuId"]) for row in batch_rows])
        return {
            index: {"match_result": "SAME", "confidence": 0.9, "model_match": "exact", "spec_match": True}
            for index, _row_value in enumerate(batch_rows)
        }

    def plan(items, _timeout):
        members = [
            CrawlMember("history-mall", str(item["row"]["vCommodityId"]), str(item["row"]["skuId"]), str(item["row"]["extLink"]), "history-v1", item["row"])
            for item in items if normalize_cell_str(item["row"].get("extLink"))
        ]
        return [CrawlPlan("history-mall", "proposal", "", "", members=tuple(members))] if members else []

    def finalize(plan, members, _timeout):
        identities = [(member.mall_key, member.source_id, member.sku_id, member.exact_url.strip(), member.backend_version) for member in members]
        digest = hashlib.sha256(repr(identities).encode("utf-8")).hexdigest()
        return CrawlPlan(
            plan.mall_key, f"history:{digest}", f"history-{digest[:12]}", digest,
            payload_json={"members": identities}, members=tuple(members),
        )

    def submit(job, crawls, _timeout):
        posted_identities.extend(
            (str(crawl["mall_key"]), str(crawl["source_id"]), str(crawl["sku_id"]), str(crawl["exact_url"]).strip(), str(crawl["backend_version"]))
            for crawl in crawls
        )
        return {"remote_identity": f"history-{job['id']}", "state": "SUBMITTED"}

    def harvest(jobs, _timeout):
        return {
            int(job["id"]): {
                "state": "COMPLETED", "remote_identity": str(job["remote_identity"]),
                "crawl_results": {
                    int(crawl["id"]): {"page_price_value": Decimal("9.5"), "page_price_status": "OK", "page_price_source": "history-fixture"}
                    for crawl in job["crawl_work"]
                },
            }
            for job in jobs
        }

    store = DailyAuditStore(tmp_path / "history-state.sqlite")
    lock = ProcessLock(tmp_path / "history.lock")
    config = TickConfig(
        business_date="2026-08-11", runtime_budget_seconds=300, heavy_work_reserve_seconds=1,
        lease_ttl_seconds=60, cutoff_at=datetime(2026, 8, 11, 18, 0, tzinfo=UTC), model_batch_size=10,
    )
    deps = TickDependencies(
        discover=discover, harvest=harvest, detail=detail, model_batch=model_batch,
        plan_crawl=plan, finalize_crawl_plan=finalize, submit_crawl=submit,
        wall_clock=lambda: datetime(2026, 8, 11, 18, 1, tzinfo=UTC),
    )
    for _ in range(3):
        run_incremental_tick(store, lock, config, deps)
    frozen_tick = run_incremental_tick(store, lock, config, deps, finalize=True)
    assert frozen_tick.frozen_snapshot_id is not None
    # The final cutoff delta is harvested by the next finite scheduler tick.
    completed_tick = run_incremental_tick(store, lock, config, deps, finalize=True)
    assert completed_tick.outcome == "COMPLETED"

    token = store.acquire_lease("2026-08-11", "assert-history", 60)
    frozen = store.get_snapshot(int(frozen_tick.frozen_snapshot_id))
    assert frozen is not None
    frozen_items = store.ordered_snapshot_items(int(frozen["id"]))
    expected_detail_fingerprints = {
        detail_fingerprint(item["row"], config.business_date, config.detail_schema_version)
        for item in frozen_items
    }
    expected_model_fingerprints = {
        model_fingerprint(item["row"], config.model_name, config.model_prompt_version, config.model_schema_version)
        for item in frozen_items
    }
    expected_crawl_identities = {
        ("history-mall", str(item["row"]["vCommodityId"]), str(item["row"]["skuId"]), normalize_cell_str(item["row"].get("extLink")), "history-v1")
        for item in frozen_items if normalize_cell_str(item["row"].get("extLink"))
    }
    artifact = build_frozen_audit_artifact(store, token, output_identity="history-equivalence").artifact
    store.release_lease(token)
    assert artifact is not None
    assert len(detail_calls) == len(expected_detail_fingerprints)
    assert sum(map(len, model_requests)) == len(expected_model_fingerprints)
    # New prefixes are submitted on separate finite ticks, but no title input
    # is sent twice; batching still avoids the legacy one-POST-per-row shape.
    assert len(model_requests) < len(expected_model_fingerprints)
    # Durable crawl requests retain source_id, while one backend result may
    # safely satisfy several v rows only on an exact (SKU, URL) match.
    assert {
        (mall, sku, url, backend)
        for mall, _source, sku, url, backend in posted_identities
    } == {
        (mall, sku, url, backend)
        for mall, _source, sku, url, backend in expected_crawl_identities
    }
    assert set(posted_identities) <= expected_crawl_identities
    assert len(posted_identities) == len({
        (mall, sku, url, backend)
        for mall, _source, sku, url, backend in expected_crawl_identities
    })
    assert len(posted_identities) == len(expected_crawl_identities) - 1

    actual = tmp_path / "incremental.xlsx"
    publication = materialize_frozen_daily_run(store, lock, "2026-08-11", actual)
    assert publication.materialization is not None
    direct = tmp_path / "legacy-writer.xlsx"
    legacy_facade.write_audit_workbook(
        direct, list(artifact.headers), [dict(row) for row in artifact.frozen_rows],
        [dict(row) for row in artifact.enriched_rows], artifact.audit_response,
    )
    assert _sheet_values(actual, "record_results") == _sheet_values(direct, "record_results")
    assert _summary_business_projection(actual) == _summary_business_projection(direct)
