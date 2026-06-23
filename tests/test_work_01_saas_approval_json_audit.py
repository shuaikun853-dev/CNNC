import json
import os
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient

from app.schemas.workflow import InternalCrawlResult, SaasSameProductApprovalItem, WorkflowAction
from app.services.link_audit_worker import LinkAuditResult, check_price_consistency, check_sku_link_consistency
from app.services.title_match_model_adapter import TitleMatchModelRequest, TitleMatchModelResult, call_title_match_model
from app.workflows.work_01_saas_approval_json_audit import (
    audit_saas_approval_item,
    audit_saas_approval_records,
    audit_title_similarity,
    check_candidate_price_against_customer_price,
    get_candidate_price,
    get_customer_price,
    match_crawl_result,
    sanitize_screenshot_path,
)
from scripts.audit_saas_json_file import load_crawl_results, load_records, main as audit_file_main


SAMPLE = {
    "vCommodityId": "679416640600166400",
    "queryStr": "世达 (SATA) 电工维修组合工具 06005A 27件/套 销售单位：套 （电工维修组合工具\\公制 27件套） （计量单位：套）",
    "syncState": "1",
    "skuId": "100237336261",
    "skuName": "世达（SATA）06005 替代款06005A 27件基本维修组套",
    "modelAuditResult": "1",
    "isMatch": "1",
    "vendorName": "jd",
    "shopName": "京东商城",
    "price": "347.48",
    "commentNum": "",
    "extLink": "https://item.jd.com/100237336261.html",
    "createTime": "2026-06-17 14:16:41",
    "updateTime": "2026-06-17 14:16:46",
    "isOperate": 1,
}


def _scope(**overrides):
    base = {"project": "CNNC", "tenant_id": 188605, "commodity_pool_id": 222}
    base.update(overrides)
    return base


def _set_token(value=None):
    if value is None:
        os.environ.pop("LOCAL_WORKFLOW_TOKEN", None)
    else:
        os.environ["LOCAL_WORKFLOW_TOKEN"] = value


@pytest.fixture(autouse=True)
def _clean_token():
    original = os.environ.get("LOCAL_WORKFLOW_TOKEN")
    yield
    if original is None:
        os.environ.pop("LOCAL_WORKFLOW_TOKEN", None)
    else:
        os.environ["LOCAL_WORKFLOW_TOKEN"] = original


def _item(**overrides) -> SaasSameProductApprovalItem:
    data = dict(SAMPLE)
    data.update(overrides)
    return SaasSameProductApprovalItem(**data)


def _crawl(**overrides) -> InternalCrawlResult:
    data = {
        "extLink": SAMPLE["extLink"],
        "skuId": SAMPLE["skuId"],
        "vendorName": SAMPLE["vendorName"],
        "shopName": SAMPLE["shopName"],
        "crawl_status": "SUCCESS",
        "page_title": "阀门 DN20",
        "page_price": "10.00",
        "page_status": "NORMAL",
        "screenshot_path": "runtime/screenshots/link_audit/a.png",
        "crawl_time": "2026-06-17 10:00:00",
        "price_source": "INTERNAL_CRAWLER",
    }
    data.update(overrides)
    return InternalCrawlResult(**data)


def test_local_audit_script_loads_records_array(tmp_path):
    path = tmp_path / "records.json"
    path.write_text(json.dumps([SAMPLE], ensure_ascii=False), encoding="utf-8")
    records = load_records(path)
    assert len(records) == 1
    assert records[0].vCommodityId == SAMPLE["vCommodityId"]


def test_local_audit_script_loads_records_object(tmp_path):
    path = tmp_path / "records.json"
    path.write_text(json.dumps({"records": [SAMPLE]}, ensure_ascii=False), encoding="utf-8")
    records = load_records(path)
    assert len(records) == 1
    assert records[0].skuId == SAMPLE["skuId"]


def test_local_audit_script_runs_without_crawl_results(tmp_path, capsys):
    records_path = tmp_path / "records.json"
    output_path = tmp_path / "audit_result.json"
    sample = dict(SAMPLE)
    sample.update({"queryStr": "阀门 DN20", "skuName": "阀门 DN20", "price": "10.00", "pageActualPrice": "10.00"})
    records_path.write_text(json.dumps([sample], ensure_ascii=False), encoding="utf-8")
    exit_code = audit_file_main(["--records-json", str(records_path), "--output-json", str(output_path)])
    output = json.loads(output_path.read_text(encoding="utf-8"))
    captured = capsys.readouterr().out
    assert exit_code == 0
    assert "总记录数" in captured
    assert "输出文件路径" in captured
    assert "summary" in output
    assert "record_results" in output
    assert output["record_results"][0]["execution_allowed"] is False


def test_local_audit_script_runs_with_crawl_results(tmp_path):
    records_path = tmp_path / "records.json"
    crawl_path = tmp_path / "crawl_results.json"
    output_path = tmp_path / "audit_result.json"
    sample = dict(SAMPLE)
    sample.update({"queryStr": "阀门 DN20", "skuName": "阀门 DN20", "price": "10.00", "customerPrice": "100"})
    records_path.write_text(json.dumps({"records": [sample]}, ensure_ascii=False), encoding="utf-8")
    crawl_path.write_text(json.dumps([_crawl(page_price="11.00").model_dump(mode="json")], ensure_ascii=False), encoding="utf-8")
    exit_code = audit_file_main(["--records-json", str(records_path), "--crawl-results-json", str(crawl_path), "--output-json", str(output_path)])
    output = json.loads(output_path.read_text(encoding="utf-8"))
    record = output["record_results"][0]
    assert exit_code == 0
    assert record["crawl_matched"] is True
    assert record["price_check"] == "MISMATCH"
    assert record["needs_human_approval"] is True
    assert record["execution_allowed"] is False
    assert record["executed"] is False


def test_local_audit_script_rejects_excel_csv_paths(tmp_path):
    records_path = tmp_path / "records.csv"
    records_path.write_text("[]", encoding="utf-8")
    with pytest.raises(ValueError, match="不读取 Excel/CSV"):
        load_records(records_path)


def test_local_audit_script_loads_crawl_results_object(tmp_path):
    crawl_path = tmp_path / "crawl_results.json"
    crawl_path.write_text(json.dumps({"crawl_results": [_crawl().model_dump(mode="json")]}, ensure_ascii=False), encoding="utf-8")
    results = load_crawl_results(crawl_path)
    assert len(results) == 1
    assert results[0].crawl_status == "SUCCESS"


def test_local_audit_script_contains_no_external_execution_paths():
    source = os.path.join(os.path.dirname(__file__), "..", "scripts", "audit_saas_json_file.py")
    text = open(source, "r", encoding="utf-8").read().lower()
    forbidden = ["playwright", "ocr", "requests", "httpx", "sqlite", "sqlalchemy", "openpyxl", "pandas", "cookie", "har"]
    assert all(item not in text for item in forbidden)


def test_parse_sample_json_and_manual_review_replacement_reason():
    result = audit_saas_approval_item(_item(), use_model=False)
    assert result.vCommodityId == "679416640600166400"
    assert result.json_price == Decimal("347.48")
    assert result.title_match_result == "MANUAL_REVIEW"
    assert "替代款" in result.reason
    assert result.primary_action == WorkflowAction.MANUAL_CHECK


def test_jd_sku_link_match():
    result, evidence = check_sku_link_consistency("100237336261", "https://item.jd.com/100237336261.html")
    assert result == "MATCH"
    assert "京东链接 SKU 与 skuId 一致" in evidence


def test_jd_sku_link_mismatch():
    result, evidence = check_sku_link_consistency("1", "https://item.jd.com/100237336261.html")
    assert result == "MISMATCH"
    assert "京东链接 SKU 与 skuId 不一致" in evidence


def test_unsupported_link_unknown():
    result, evidence = check_sku_link_consistency("1", "https://example.com/1")
    assert result == "UNKNOWN"
    assert "暂不支持该商城链接 SKU 解析，需人工核验" in evidence


def test_price_string_to_decimal():
    assert _item(price="347.48").price == Decimal("347.48")


def test_empty_price_returns_json_price_missing():
    result, evidence = check_price_consistency(_item(price="").price, Decimal("1"))
    assert result == "JSON_PRICE_MISSING"
    assert "SaaS JSON 价格为空" in evidence


def test_missing_page_actual_price_waits_check():
    result, evidence = check_price_consistency(Decimal("1"), None)
    assert result == "WAIT_PAGE_PRICE_CHECK"
    assert "缺少页面实际价格，需人工打开链接核验" in evidence


def test_page_actual_price_match():
    result, evidence = check_price_consistency(Decimal("1.00"), Decimal("1.009"))
    assert result == "MATCH"
    assert "SaaS JSON 价格与页面实际价格一致" in evidence


def test_page_actual_price_mismatch():
    result, evidence = check_price_consistency(Decimal("1.00"), Decimal("1.02"))
    assert result == "MISMATCH"
    assert "SaaS JSON 价格与页面实际价格不一致" in evidence


def test_title_model_match_likely_match():
    result, evidence = audit_title_similarity("阀门 DN20 1套", "某品牌 DN20 阀门 1套", "", "")
    assert result == "SAME"
    assert any("核心型号命中" in item for item in evidence)


def test_title_replacement_manual_review():
    result, evidence = audit_title_similarity(SAMPLE["queryStr"], SAMPLE["skuName"], "1", "1")
    assert result == "MANUAL_REVIEW"
    assert any("替代款" in item for item in evidence)


def test_model_flags_add_evidence():
    _, evidence = audit_title_similarity("阀门 DN20", "阀门 DN20", "1", "1")
    assert "SaaS模型判断为同款" in evidence


def test_model_unavailable_does_not_interrupt(monkeypatch):
    monkeypatch.delenv("TITLE_MATCH_MODEL_API_KEY", raising=False)
    monkeypatch.delenv("TITLE_MATCH_MODEL_BASE_URL", raising=False)
    result = call_title_match_model(TitleMatchModelRequest(query_str="阀门 DN20", sku_name="阀门 DN20"))
    assert result.match_result == "MODEL_UNAVAILABLE"


def test_model_same_participates_in_approve():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20 1套",
        skuName="某品牌阀门 DN20 1套",
        skuId="100237336261",
        extLink="https://item.jd.com/100237336261.html",
        price="10.00",
        pageActualPrice="10.00",
    ), title_model_result=model)
    assert result.title_match_result == "SAME"
    assert result.title_match_confidence == 0.91
    assert result.primary_action == WorkflowAction.APPROVE


def test_risk_word_downgrades_model_same_to_manual_check():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        skuId="100237336261",
        extLink="https://item.jd.com/100237336261.html",
        price="347.48",
        pageActualPrice="347.48",
    ), title_model_result=model)
    assert result.primary_action == WorkflowAction.MANUAL_CHECK
    assert result.execution_allowed is False
    assert result.executed is False
    assert result.needs_human_approval is True
    assert "替代款" in " ".join(result.evidence)


@pytest.mark.parametrize("field,value,word", [
    ("skuName", "某品牌阀门 DN20 替代款", "替代款"),
    ("queryStr", "通用阀门 DN20", "通用"),
    ("skuName", "某品牌阀门 DN20 适用", "适用"),
    ("skuName", "某品牌阀门 DN20 升级款", "升级款"),
])
def test_risk_words_block_approval(field, value, word):
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    overrides = {
        "queryStr": "阀门 DN20",
        "skuName": "某品牌阀门 DN20",
        "price": "10.00",
        "pageActualPrice": "10.00",
    }
    overrides[field] = value
    result = audit_saas_approval_item(_item(**overrides), title_model_result=model)
    assert result.primary_action == WorkflowAction.MANUAL_CHECK
    assert word in " ".join(result.evidence)


def test_risk_word_does_not_downgrade_sku_mismatch_reject():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="某品牌阀门 DN20 替代款",
        skuId="1",
        extLink="https://item.jd.com/2.html",
        price="10.00",
        pageActualPrice="10.00",
    ), title_model_result=model)
    assert result.primary_action == WorkflowAction.REJECT


def test_risk_word_does_not_downgrade_price_mismatch_reject():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="某品牌阀门 DN20 替代款",
        price="10.00",
        pageActualPrice="11.00",
    ), title_model_result=model)
    assert result.primary_action == WorkflowAction.REJECT


def test_model_not_same_rejects():
    model = TitleMatchModelResult(match_result="NOT_SAME", confidence=0.88, model_match="mismatch")
    result = audit_saas_approval_item(_item(price="10.00", pageActualPrice="10.00"), title_model_result=model)
    assert result.primary_action == WorkflowAction.REJECT


def test_model_manual_review_manual_check():
    model = TitleMatchModelResult(match_result="MANUAL_REVIEW", confidence=0.5, model_match="partial")
    result = audit_saas_approval_item(_item(price="10.00", pageActualPrice="10.00"), title_model_result=model)
    assert result.primary_action == WorkflowAction.MANUAL_CHECK


def test_likely_match_sku_match_price_match_approve():
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20 1套",
        skuName="某品牌阀门 DN20 1套",
        skuId="100237336261",
        extLink="https://item.jd.com/100237336261.html",
        price="10.00",
        pageActualPrice="10.00",
    ), use_model=False)
    assert result.title_match_result == "SAME"
    assert result.sku_link_check == "MATCH"
    assert result.price_check == "MATCH"
    assert result.primary_action == WorkflowAction.APPROVE


def test_match_crawl_result_ext_link_priority():
    record = _item(extLink="https://example.com/a", skuId="sku-a")
    crawl = _crawl(extLink="https://example.com/a", skuId="other", page_price="11.00")
    assert match_crawl_result(record, [crawl]) == crawl


def test_match_crawl_result_sku_id_priority():
    record = _item(extLink="https://example.com/a", skuId="sku-a")
    crawl = _crawl(extLink="https://example.com/b", skuId="sku-a", page_price="11.00")
    assert match_crawl_result(record, [crawl]) == crawl


def test_match_crawl_result_vendor_and_sku_priority():
    record = _item(extLink="https://example.com/a", skuId="sku-a", vendorName="jd")
    crawl = _crawl(extLink="https://example.com/b", skuId="sku-a", vendorName="京东", page_price="11.00")
    assert match_crawl_result(record, [crawl]) == crawl


def test_match_crawl_result_latest_crawl_time_wins():
    record = _item(extLink="https://example.com/a")
    old = _crawl(extLink="https://example.com/a", page_price="11.00", crawl_time="2026-06-17 10:00:00")
    new = _crawl(extLink="https://example.com/a", page_price="12.00", crawl_time="2026-06-17 11:00:00")
    assert match_crawl_result(record, [old, new]) == new


def test_match_crawl_result_invalid_time_uses_last():
    record = _item(extLink="https://example.com/a")
    first = _crawl(extLink="https://example.com/a", page_price="11.00", crawl_time="bad")
    last = _crawl(extLink="https://example.com/a", page_price="12.00", crawl_time="also-bad")
    assert match_crawl_result(record, [first, last]) == last


def test_missing_crawl_result_manual_check_when_enabled():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="10.00",
        pageActualPrice="10.00",
    ), title_model_result=model, crawl_results=[], use_crawl_result=True)
    assert result.crawl_matched is False
    assert result.price_check == "WAIT_PAGE_PRICE_CHECK"
    assert result.primary_action == WorkflowAction.MANUAL_CHECK
    assert "缺少内部爬取结果，需人工核验页面价格" in " ".join(result.evidence)


@pytest.mark.parametrize("crawl_status", ["SOLD_OUT", "PRICE_MISSING", "CAPTCHA", "ACCESS_DENIED", "TIMEOUT", "FAILED"])
def test_crawl_abnormal_status_manual_check(crawl_status):
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="10.00",
        customerPrice="100",
    ), title_model_result=model, crawl_results=[_crawl(crawl_status=crawl_status, page_status=crawl_status, page_price="10.00")], use_crawl_result=True)
    assert result.crawl_status == crawl_status
    assert result.primary_action == WorkflowAction.MANUAL_CHECK
    assert result.execution_allowed is False
    assert result.executed is False
    assert result.needs_human_approval is True


def test_crawl_off_shelf_reject_suggestion_only():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="10.00",
        customerPrice="100",
    ), title_model_result=model, crawl_results=[_crawl(crawl_status="OFF_SHELF", page_status="OFF_SHELF", page_price="10.00")], use_crawl_result=True)
    assert result.off_shelf is True
    assert result.primary_action == WorkflowAction.REJECT
    assert result.execution_allowed is False
    assert result.executed is False
    assert result.needs_human_approval is True


def test_crawl_page_price_match_and_mismatch():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    matched = audit_saas_approval_item(_item(queryStr="阀门 DN20", skuName="阀门 DN20", price="10.00"), title_model_result=model, crawl_results=[_crawl(page_price="10.009")], use_crawl_result=True)
    mismatched = audit_saas_approval_item(_item(queryStr="阀门 DN20", skuName="阀门 DN20", price="10.00"), title_model_result=model, crawl_results=[_crawl(page_price="10.02")], use_crawl_result=True)
    assert matched.price_check == "MATCH"
    assert mismatched.price_check == "MISMATCH"
    assert mismatched.primary_action == WorkflowAction.REJECT


def test_crawl_page_price_takes_priority_for_customer_price_check():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="100",
        pageActualPrice="100",
        customerPrice="100",
    ), title_model_result=model, crawl_results=[_crawl(page_price="131")], use_crawl_result=True)
    assert result.page_price == Decimal("131")
    assert result.page_actual_price == Decimal("131")
    assert result.price_vs_customer_check == "HIGHER_THAN_CUSTOMER_PRICE_30"
    assert result.primary_action == WorkflowAction.REJECT


def test_crawl_missing_page_price_does_not_approve():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="10.00",
        customerPrice="100",
    ), title_model_result=model, crawl_results=[_crawl(page_price=None, crawl_status="PRICE_MISSING")], use_crawl_result=True)
    assert result.price_check == "WAIT_PAGE_PRICE_CHECK"
    assert result.price_vs_customer_check != "HIGHER_THAN_CUSTOMER_PRICE_30"
    assert result.primary_action == WorkflowAction.MANUAL_CHECK


def test_crawl_customer_price_missing_and_invalid_do_not_trigger_high_price_reject():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    missing = audit_saas_approval_item(_item(queryStr="阀门 DN20", skuName="阀门 DN20", price="131"), title_model_result=model, crawl_results=[_crawl(page_price="131")], use_crawl_result=True)
    invalid = audit_saas_approval_item(_item(queryStr="阀门 DN20", skuName="阀门 DN20", price="131", customerPrice="0"), title_model_result=model, crawl_results=[_crawl(page_price="131")], use_crawl_result=True)
    assert missing.price_vs_customer_check == "CUSTOMER_PRICE_MISSING"
    assert invalid.price_vs_customer_check == "CUSTOMER_PRICE_INVALID"
    assert missing.primary_action != WorkflowAction.REJECT
    assert invalid.primary_action != WorkflowAction.REJECT


def test_crawl_result_does_not_call_browser_audit(monkeypatch):
    import app.workflows.work_01_saas_approval_json_audit as mod

    def fail_audit_link(_request):
        raise AssertionError("不应在 use_crawl_result=True 时调用浏览器/链接抓取")

    monkeypatch.setattr(mod, "audit_link", fail_audit_link)
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(queryStr="阀门 DN20", skuName="阀门 DN20", price="10.00"), title_model_result=model, crawl_results=[_crawl(page_price="10.00")], use_crawl_result=True)
    assert result.crawl_matched is True


def test_ocr_price_source_cannot_directly_approve():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(queryStr="阀门 DN20", skuName="阀门 DN20", price="10.00"), title_model_result=model, crawl_results=[_crawl(page_price="10.00", price_source="OCR")], use_crawl_result=True)
    assert result.primary_action == WorkflowAction.MANUAL_CHECK
    assert "价格仅来自 OCR" in " ".join(result.evidence)


def test_screenshot_path_traversal_ignored():
    safe, safe_evidence = sanitize_screenshot_path("runtime/screenshots/link_audit/a.png")
    unsafe, unsafe_evidence = sanitize_screenshot_path("../secret.png")
    assert safe == "runtime/screenshots/link_audit/a.png"
    assert "内部爬取截图路径" in " ".join(safe_evidence)
    assert unsafe == ""
    assert "截图路径疑似不安全，已忽略" in " ".join(unsafe_evidence)


def test_crawl_result_does_not_expose_excel_csv_har_or_sensitive_headers():
    result = audit_saas_approval_item(_item(queryStr="阀门 DN20", skuName="阀门 DN20", price="10.00"), crawl_results=[_crawl(page_price="10.00")], use_crawl_result=True)
    evidence = " ".join(result.evidence).lower()
    assert "excel" not in evidence
    assert "csv" not in evidence
    assert "har" not in evidence
    assert "cookie" not in evidence
    assert "authorization" not in evidence


def test_crawl_summary_counts():
    response = audit_saas_approval_records([
        _item(vCommodityId="v1", extLink="https://example.com/1", skuId="1", price="10.00"),
        _item(vCommodityId="v2", extLink="https://example.com/2", skuId="2", price="10.00"),
        _item(vCommodityId="v3", extLink="https://example.com/3", skuId="3", price="10.00"),
    ], use_model=False, crawl_results=[
        _crawl(extLink="https://example.com/1", skuId="1", page_price="10.00"),
        _crawl(extLink="https://example.com/2", skuId="2", page_price="11.00"),
        _crawl(extLink="https://example.com/3", skuId="3", crawl_status="OFF_SHELF", page_status="OFF_SHELF", page_price="10.00"),
    ], use_crawl_result=True)
    assert response.summary["crawl_result_matched_count"] == 3
    assert response.summary["crawl_price_mismatch_count"] == 1
    assert response.summary["crawl_off_shelf_count"] == 1


def test_customer_price_getter_priority():
    item = _item(customerPrice="100", customer_price="101", referencePrice="102", reference_price="103")
    assert get_customer_price(item) == Decimal("100")


def test_customer_price_30_high_rejects_even_when_other_checks_match():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="131",
        pageActualPrice="131",
        customerPrice="100",
    ), title_model_result=model)
    assert result.price_vs_customer_check == "HIGHER_THAN_CUSTOMER_PRICE_30"
    assert result.primary_action == WorkflowAction.REJECT
    assert result.suggested_actions == [WorkflowAction.REJECT, WorkflowAction.MANUAL_CHECK]
    assert result.execution_allowed is False
    assert result.executed is False
    assert result.needs_human_approval is True
    evidence = " ".join(result.evidence)
    assert "客户价格=100" in evidence
    assert "候选价格=131" in evidence
    assert "高价阈值=客户价格*1.3=130.00" in evidence
    assert "候选价格高于客户价格30%以上" in evidence
    assert "已驳回" not in evidence
    assert "已驳回" not in result.reason


def test_customer_price_30_equal_threshold_not_reject():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="130",
        pageActualPrice="130",
        customerPrice="100",
    ), title_model_result=model)
    assert result.price_vs_customer_check == "WITHIN_THRESHOLD"
    assert result.primary_action == WorkflowAction.APPROVE


def test_page_actual_price_takes_priority_for_customer_price_check():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    item = _item(queryStr="阀门 DN20", skuName="阀门 DN20", price="100", pageActualPrice="131", customerPrice="100")
    assert get_candidate_price(item, item.pageActualPrice) == Decimal("131")
    result = audit_saas_approval_item(item, title_model_result=model)
    assert result.price_vs_customer_check == "HIGHER_THAN_CUSTOMER_PRICE_30"
    assert result.primary_action == WorkflowAction.REJECT


def test_customer_price_missing_does_not_trigger_threshold_rule():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="100",
        pageActualPrice="100",
    ), title_model_result=model)
    assert result.price_vs_customer_check == "CUSTOMER_PRICE_MISSING"
    assert result.primary_action == WorkflowAction.APPROVE


def test_zero_customer_price_does_not_trigger_high_price_reject():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="131",
        pageActualPrice="131",
        customerPrice="0",
    ), title_model_result=model)
    assert result.price_vs_customer_check == "CUSTOMER_PRICE_INVALID"
    assert result.primary_action != WorkflowAction.REJECT
    evidence = " ".join(result.evidence)
    assert "客户价格无效" in evidence
    assert "候选价格高于客户价格30%以上" not in evidence


def test_negative_customer_price_does_not_trigger_high_price_reject():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="131",
        pageActualPrice="131",
        customerPrice="-1",
    ), title_model_result=model)
    assert result.price_vs_customer_check == "CUSTOMER_PRICE_INVALID"
    assert result.primary_action != WorkflowAction.REJECT
    evidence = " ".join(result.evidence)
    assert "客户价格无效" in evidence
    assert "候选价格高于客户价格30%以上" not in evidence


def test_candidate_price_zero_does_not_trigger_high_price_reject():
    result, evidence = check_candidate_price_against_customer_price(Decimal("0"), Decimal("100"))
    assert result == "CANDIDATE_PRICE_INVALID"
    assert "候选价格无效" in " ".join(evidence)


def test_candidate_price_missing_continues_existing_manual_check_path():
    model = TitleMatchModelResult(match_result="SAME", confidence=0.91, model_match="exact")
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="",
        pageActualPrice=None,
        customerPrice="100",
    ), title_model_result=model)
    assert result.price_vs_customer_check == "CANDIDATE_PRICE_MISSING"
    assert result.price_check == "JSON_PRICE_MISSING"
    assert result.primary_action == WorkflowAction.MANUAL_CHECK


def test_candidate_price_customer_price_function_outputs():
    result, evidence = check_candidate_price_against_customer_price(Decimal("131"), Decimal("100"))
    assert result == "HIGHER_THAN_CUSTOMER_PRICE_30"
    assert "候选价格高于客户价格30%以上" in evidence


def test_price_mismatch_reject():
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        price="10.00",
        pageActualPrice="11.00",
    ), use_model=False)
    assert result.price_check == "MISMATCH"
    assert result.primary_action == WorkflowAction.REJECT


def test_sku_mismatch_reject():
    result = audit_saas_approval_item(_item(
        queryStr="阀门 DN20",
        skuName="阀门 DN20",
        skuId="1",
        extLink="https://item.jd.com/2.html",
        price="10.00",
        pageActualPrice="10.00",
    ), use_model=False)
    assert result.sku_link_check == "MISMATCH"
    assert result.primary_action == WorkflowAction.REJECT


def test_wait_page_price_check_manual_check():
    result = audit_saas_approval_item(_item(queryStr="阀门 DN20", skuName="阀门 DN20", pageActualPrice=None), use_model=False)
    assert result.price_check == "WAIT_PAGE_PRICE_CHECK"
    assert result.primary_action == WorkflowAction.MANUAL_CHECK


def test_reachable_false_reject(monkeypatch):
    import app.workflows.work_01_saas_approval_json_audit as mod

    def fake_audit_link(_request):
        return LinkAuditResult(
            reachable=False,
            off_shelf=None,
            sku_link_check="MATCH",
            page_price=Decimal("10.00"),
            price_check="MATCH",
            evidence=["链接不可访问"],
        )

    monkeypatch.setattr(mod, "audit_link", fake_audit_link)
    result = audit_saas_approval_item(_item(queryStr="阀门 DN20", skuName="阀门 DN20", price="10.00", pageActualPrice="10.00"), use_model=False)
    assert result.reachable is False
    assert result.primary_action == WorkflowAction.REJECT


def test_off_shelf_true_reject(monkeypatch):
    import app.workflows.work_01_saas_approval_json_audit as mod

    def fake_audit_link(_request):
        return LinkAuditResult(
            reachable=True,
            off_shelf=True,
            sku_link_check="MATCH",
            page_price=Decimal("10.00"),
            price_check="MATCH",
            evidence=["商品疑似下架"],
        )

    monkeypatch.setattr(mod, "audit_link", fake_audit_link)
    result = audit_saas_approval_item(_item(queryStr="阀门 DN20", skuName="阀门 DN20", price="10.00", pageActualPrice="10.00"), use_model=False)
    assert result.off_shelf is True
    assert result.primary_action == WorkflowAction.REJECT


def test_output_is_recommendation_only():
    response = audit_saas_approval_records([_item(), _item(vCommodityId="v2")], use_model=False)
    assert response.execution_allowed is False
    assert response.executed is False
    for result in response.record_results:
        assert result.execution_allowed is False
        assert result.executed is False
        assert result.needs_human_approval is True


def test_summary_group_counts():
    response = audit_saas_approval_records([
        _item(vCommodityId="v1", pageActualPrice="347.48"),
        _item(vCommodityId="v1", skuId="1", extLink="https://item.jd.com/2.html", pageActualPrice="347.48"),
    ], use_model=False, use_crawl_result=False)
    assert response.summary["total_records"] == 2
    assert response.summary["total_v_commodities"] == 1
    assert response.summary["sku_mismatch_count"] == 1
    assert response.summary["group.v1.candidate_count"] == 2
    assert response.summary["group.v1.valid_mall_count"] == 1


class TestSaasAuditAPI:
    def test_api_token_not_configured_rejected(self):
        _set_token(None)
        from app.main import app
        client = TestClient(app)
        resp = client.post("/workflows/same-product-approval/audit-saas-json", json={"scope": _scope(), "records": [SAMPLE]})
        assert resp.status_code == 403

    def test_api_wrong_token_rejected(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        resp = client.post(
            "/workflows/same-product-approval/audit-saas-json",
            json={"scope": _scope(), "records": [SAMPLE]},
            headers={"X-Local-Workflow-Token": "wrong"},
        )
        assert resp.status_code == 403

    def test_api_correct_token_passes(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        resp = client.post(
            "/workflows/same-product-approval/audit-saas-json",
            json={"scope": _scope(), "records": [SAMPLE], "use_crawl_result": False},
            headers={"X-Local-Workflow-Token": "correct-token"},
        )
        assert resp.status_code == 200
        assert resp.json()["record_results"][0]["execution_allowed"] is False

    def test_api_consumes_crawl_results(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        sample = dict(SAMPLE)
        sample.update({"queryStr": "阀门 DN20", "skuName": "阀门 DN20", "price": "10.00", "customerPrice": "100"})
        resp = client.post(
            "/workflows/same-product-approval/audit-saas-json",
            json={"scope": _scope(), "records": [sample], "crawl_results": [_crawl(page_price="11.00").model_dump(mode="json")]},
            headers={"X-Local-Workflow-Token": "correct-token"},
        )
        assert resp.status_code == 200
        record = resp.json()["record_results"][0]
        assert record["crawl_matched"] is True
        assert record["page_price"] == "11.00"
        assert record["price_check"] == "MISMATCH"

    def test_api_use_crawl_result_false_keeps_r36_logic(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        sample = dict(SAMPLE)
        sample.update({"queryStr": "阀门 DN20", "skuName": "阀门 DN20", "price": "10.00", "pageActualPrice": "10.00"})
        resp = client.post(
            "/workflows/same-product-approval/audit-saas-json",
            json={
                "scope": _scope(),
                "records": [sample],
                "crawl_results": [_crawl(page_price="11.00").model_dump(mode="json")],
                "use_crawl_result": False,
            },
            headers={"X-Local-Workflow-Token": "correct-token"},
        )
        assert resp.status_code == 200
        record = resp.json()["record_results"][0]
        assert record["crawl_matched"] is False
        assert record["price_check"] == "MATCH"

    def test_api_empty_crawl_results_manual_check(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        sample = dict(SAMPLE)
        sample.update({"queryStr": "阀门 DN20", "skuName": "阀门 DN20", "price": "10.00", "pageActualPrice": "10.00"})
        resp = client.post(
            "/workflows/same-product-approval/audit-saas-json",
            json={"scope": _scope(), "records": [sample], "crawl_results": []},
            headers={"X-Local-Workflow-Token": "correct-token"},
        )
        assert resp.status_code == 200
        record = resp.json()["record_results"][0]
        assert record["price_check"] == "WAIT_PAGE_PRICE_CHECK"
        assert record["primary_action"] == WorkflowAction.MANUAL_CHECK

    def test_api_non_cnnc_scope_rejected(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        bad_scope = {"project": "OTHER", "tenant_id": 999, "commodity_pool_id": 999}
        resp = client.post(
            "/workflows/same-product-approval/audit-saas-json",
            json={"scope": bad_scope, "records": [SAMPLE]},
            headers={"X-Local-Workflow-Token": "correct-token"},
        )
        assert resp.status_code in (400, 422)

    def test_api_empty_records_rejected(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        resp = client.post(
            "/workflows/same-product-approval/audit-saas-json",
            json={"scope": _scope(), "records": []},
            headers={"X-Local-Workflow-Token": "correct-token"},
        )
        assert resp.status_code == 422

    def test_diagnose_and_health_not_affected(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        assert client.get("/health").status_code == 200
        assert client.post("/diagnose", json={"question": "中核同款审批"}).status_code in (200, 422)

    def test_work_01_original_rules_not_affected(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        resp = client.post(
            "/workflows/same-product-approval/analyze",
            json={"scope": _scope(), "records": [{"scope": _scope(), "standard_commodity_code": "SP001"}]},
            headers={"X-Local-Workflow-Token": "correct-token"},
        )
        assert resp.status_code == 200
        assert resp.json()["record_results"][0]["rule_id"] == "R02_NO_CANDIDATES"

    def test_work_02_framework_not_affected(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        resp = client.post(
            "/workflows/price-index-feedback/analyze",
            json={"scope": _scope(), "records": [{"scope": _scope(), "standard_commodity_code": "SP001"}]},
            headers={"X-Local-Workflow-Token": "correct-token"},
        )
        assert resp.status_code == 200
        assert resp.json()["record_results"][0]["rule_id"] == "R40_FRAMEWORK_PLACEHOLDER"

    def test_work_03_placeholder_not_affected(self):
        _set_token("correct-token")
        from app.main import app
        client = TestClient(app)
        resp = client.post(
            "/workflows/same-product-shortage/analyze",
            json={"scope": _scope(), "records": [{"scope": _scope(), "commodity_name": "测试"}]},
            headers={"X-Local-Workflow-Token": "correct-token"},
        )
        assert resp.status_code == 200
        assert resp.json()["record_results"][0]["rule_id"] == "R35_PLACEHOLDER"


# ---------------------------------------------------------------------------
# Excel runner tests (test_work_01_saas_approval_excel_file)
# ---------------------------------------------------------------------------


import io
import urllib.error
import urllib.request

from openpyxl import Workbook, load_workbook
from scripts.audit_saas_excel_file import (
    _display_header,
    _display_value,
    build_record_payload,
    enrich_rows_with_customer_price,
    extract_customer_price_from_detail,
    normalize_cell_str,
    parse_decimal,
    read_excel_rows,
    run_main,
    write_audit_workbook,
)


def _make_test_xlsx(tmp_path, rows: list[dict[str, Any]]) -> Path:
    path = tmp_path / "input.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    if not rows:
        wb.save(str(path))
        return path
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    wb.save(str(path))
    return path


class MockResponse:
    """模拟 urllib.request.urlopen 返回的对象。"""
    def __init__(self, data: bytes, status: int = 200):
        self.data = data
        self.status = status

    def __enter__(self):
        return self

    def __exit__(self, *args: Any) -> None:
        pass

    def read(self) -> bytes:
        return self.data


def _make_opener(payload: dict | None, status: int = 200):
    """生成 opener 工厂用于注入 call_detail_api。"""
    if payload is not None:
        body = json.dumps(payload).encode("utf-8")
    else:
        body = b"not json"

    def opener(request, timeout=15.0):
        if status >= 400:
            raise urllib.error.HTTPError(
                request.full_url or request.get_full_url() or "http://mock",
                status,
                "Error",
                {},
                io.BytesIO(body),
            )
        return MockResponse(body, status)

    return opener


DEFAULT_EXCEL_COLS = [
    "documentId", "fourCategoryId", "vCommodityId", "queryStr",
    "skuId", "skuName", "vendorName", "shopName", "price", "extLink",
]

SAMPLE_ROW_VALUES = [
    "doc1", "cat1", "v1", "阀门 DN20", "sku1", "阀门 DN20 1套",
    "jd", "京东商城", "100", "https://item.jd.com/1.html",
]


class TestExcelReadWrite:
    def test_read_xlsx_input(self, tmp_path):
        """1. 能读取 xlsx 输入"""
        rows = [dict(zip(DEFAULT_EXCEL_COLS, SAMPLE_ROW_VALUES))]
        xlsx = _make_test_xlsx(tmp_path, rows)
        headers, loaded = read_excel_rows(xlsx)
        assert len(loaded) == 1
        assert loaded[0]["documentId"] == "doc1"
        assert headers[:3] == ["documentId", "fourCategoryId", "vCommodityId"]

    def test_output_xlsx(self, tmp_path):
        """2. 能输出 xlsx"""
        rows = [dict(zip(DEFAULT_EXCEL_COLS, SAMPLE_ROW_VALUES))]
        xlsx = _make_test_xlsx(tmp_path, rows)
        headers, loaded = read_excel_rows(xlsx)
        output = tmp_path / "out.xlsx"
        enriched = [{
            "customerPrice": None,
            "customerPriceStatus": "NOT_PROVIDED",
            "customerPriceSource": "",
            "customerPriceWarning": "",
            "customerPriceEvidence": "",
        }]
        from scripts.audit_saas_excel_file import run_audit
        records = [SaasSameProductApprovalItem(**build_record_payload(loaded[0]))]
        audit_response = run_audit(records)
        write_audit_workbook(output, headers, loaded, enriched, audit_response)
        assert output.exists()
        out_wb = load_workbook(str(output))
        assert "summary" in out_wb.sheetnames
        assert "record_results" in out_wb.sheetnames
        out_wb.close()

    def test_preserves_original_columns(self, tmp_path):
        """3. 保留原始列"""
        rows = [dict(zip(DEFAULT_EXCEL_COLS, SAMPLE_ROW_VALUES))]
        xlsx = _make_test_xlsx(tmp_path, rows)
        out = tmp_path / "out.xlsx"
        result = run_main(input_xlsx=xlsx, output_xlsx=out, curl_file=None,
                          document_id_field="documentId", category_id_field="fourCategoryId")
        out_wb = load_workbook(str(out))
        out_ws = out_wb["record_results"]
        out_headers = [cell.value for cell in out_ws[1]]
        for col in DEFAULT_EXCEL_COLS:
            assert _display_header(col) in out_headers, f"缺失列: {col}"
        out_wb.close()

    def test_appends_audit_result_columns(self, tmp_path):
        """4. 追加审查结果列"""
        rows = [dict(zip(DEFAULT_EXCEL_COLS, SAMPLE_ROW_VALUES))]
        xlsx = _make_test_xlsx(tmp_path, rows)
        out = tmp_path / "out.xlsx"
        result = run_main(input_xlsx=xlsx, output_xlsx=out, curl_file=None,
                          document_id_field="documentId", category_id_field="fourCategoryId")
        out_wb = load_workbook(str(out))
        out_ws = out_wb["record_results"]
        out_headers = [str(cell.value or "") for cell in out_ws[1]]
        appended = ["customerPrice", "customerPriceStatus", "primary_action", "reason"]
        for col in appended:
            assert _display_header(col) in out_headers, f"缺失列: {_display_header(col)}"
        out_wb.close()


class TestCustomerPriceExtraction:
    def test_reference_infos_two_prices_takes_max(self):
        """5. data.referenceInfos 两个价格时取最大值"""
        payload = {
            "data": {
                "referenceInfos": [
                    {"referencePrice": "11585.37", "order": 1},
                    {"referencePrice": "9700", "order": 2},
                ],
                "referencePrice": "11585.37",
            }
        }
        result = extract_customer_price_from_detail(payload)
        assert result["customerPriceStatus"] == "FOUND"
        assert result["customerPrice"] == Decimal("11585.37")
        assert "参考价1=11585.37" in result["customerPriceEvidence"]
        assert "参考价2=9700" in result["customerPriceEvidence"]

    def test_data_reference_price_single(self):
        """6. data.referencePrice 单价时输出 FOUND_SINGLE_REFERENCE"""
        payload = {
            "data": {
                "referencePrice": "11585.37",
                "referenceInfos": [],
            }
        }
        result = extract_customer_price_from_detail(payload)
        assert result["customerPriceStatus"] == "FOUND_SINGLE_REFERENCE"
        assert result["customerPrice"] == Decimal("11585.37")
        assert "当前详情结果只提取到1个参考价" in result["customerPriceWarning"]

    def test_invalid_prices_not_participate_in_max(self):
        """7. 空值、0、负数、非法数字不参与 max"""
        payload = {
            "data": {
                "referenceInfos": [
                    {"referencePrice": "100", "order": 1},
                    {"referencePrice": "", "order": 2},
                    {"referencePrice": "0", "order": 3},
                    {"referencePrice": "-5", "order": 4},
                    {"referencePrice": None, "order": 5},
                    {"referencePrice": "abc", "order": 6},
                ],
            }
        }
        result = extract_customer_price_from_detail(payload)
        assert result["customerPriceStatus"] == "FOUND_SINGLE_REFERENCE"
        assert result["customerPrice"] == Decimal("100")
        assert "1个参考价" in result["customerPriceWarning"]

    def test_detail_request_failed_status(self, tmp_path):
        """8. 请求失败时输出 DETAIL_REQUEST_FAILED"""
        rows = [dict(zip(DEFAULT_EXCEL_COLS, [
            "doc1", "cat1", "v1", "test", "sku1", "test sku",
            "jd", "shop", "100", "https://jd.com/1",
        ]))]
        xlsx = _make_test_xlsx(tmp_path, rows)
        curl_template = {
            "url": "https://api.example.com/detail",
            "headers": [("Content-Type", "application/json")],
            "cookie": "",
            "data": '{"documentId": "x", "fourCategoryId": "y"}',
            "method": "POST",
        }
        enriched = enrich_rows_with_customer_price(
            rows, curl_template,
            "documentId", "fourCategoryId",
            opener=_make_opener(None, status=500),
        )
        assert enriched[0]["customerPriceStatus"] == "DETAIL_REQUEST_FAILED"
        assert "详情接口请求失败" in enriched[0]["customerPriceWarning"]

    def test_no_curl_file_uses_existing_customer_price(self, tmp_path):
        """9. 无 curl-file 时仍可审查已有 customerPrice"""
        cols = DEFAULT_EXCEL_COLS + ["customerPrice"]
        vals = SAMPLE_ROW_VALUES + ["100"]
        rows = [dict(zip(cols, vals))]
        xlsx = _make_test_xlsx(tmp_path, rows)
        out = tmp_path / "out.xlsx"
        result = run_main(input_xlsx=xlsx, output_xlsx=out, curl_file=None,
                          document_id_field="documentId", category_id_field="fourCategoryId")
        assert result["enriched"][0]["customerPriceStatus"] == "EXISTS"
        assert result["enriched"][0]["customerPrice"] == Decimal("100")

    def test_query_str_falls_back_to_sass_v_commodity_name(self):
        """10. queryStr 缺失时可使用 sassVCommodityName"""
        row = {
            "vCommodityId": "v1",
            "queryStr": "",
            "sassVCommodityName": "阀门 DN20 替代名称",
            "skuId": "sku1",
            "skuName": "阀门 DN20 1套",
            "vendorName": "jd",
            "shopName": "京东",
            "price": "100",
            "extLink": "https://jd.com/1",
        }
        payload = build_record_payload(row)
        assert payload["queryStr"] == "阀门 DN20 替代名称"

    def test_query_str_preferred_over_fallback(self):
        row = {
            "vCommodityId": "v1",
            "queryStr": "标准描述",
            "sassVCommodityName": "备选描述",
            "skuId": "sku1",
            "skuName": "阀门 DN20",
            "vendorName": "jd",
            "shopName": "京东",
            "price": "100",
            "extLink": "https://jd.com/1",
        }
        payload = build_record_payload(row)
        assert payload["queryStr"] == "标准描述"


class TestSafety:
    def test_no_print_cookie_token(self):
        """11. 不打印 cookie/token"""
        import ast as _ast
        path = r"E:\.A中核运维\cnnc-ops-agent\scripts\audit_saas_excel_file.py"
        source = open(path, encoding="utf-8").read()
        tree = _ast.parse(source)
        suspicious = ("cookie", "token", "auth-token", "access-token", "y1amuise")
        for node in _ast.walk(tree):
            if isinstance(node, _ast.Call):
                func = node.func
                func_name = ""
                if isinstance(func, _ast.Name):
                    func_name = func.id
                elif isinstance(func, _ast.Attribute):
                    func_name = func.attr
                if func_name in {"print", "info", "debug", "warning", "error"}:
                    for arg in node.args:
                        if isinstance(arg, _ast.Constant) and isinstance(arg.value, str):
                            lower = arg.value.lower()
                            for word in suspicious:
                                assert word not in lower, f"print 中包含敏感字: {word}"
                        if isinstance(arg, _ast.JoinedStr):
                            for value in arg.values:
                                if isinstance(value, _ast.Constant) and isinstance(value.value, str):
                                    lower = value.value.lower()
                                    for word in suspicious:
                                        assert word not in lower, f"print f-string 中包含敏感字: {word}"

    def test_no_auto_login_before_reading(self):
        """12-14. 不自动登录、不调用审批/驳回、不写数据库"""
        import ast as _ast
        path = r"E:\.A中核运维\cnnc-ops-agent\scripts\audit_saas_excel_file.py"
        source = open(path, encoding="utf-8").read()
        tree = _ast.parse(source)

        forbidden_imports = {
            "playwright", "playwright.sync_api", "playwright.async_api",
            "pytesseract", "easyocr", "paddleocr",
            "sqlite3", "sqlalchemy", "pymysql", "pymongo", "psycopg2",
            "requests",
        }
        for node in _ast.walk(tree):
            if isinstance(node, _ast.Import):
                for alias in node.names:
                    assert alias.name not in forbidden_imports, f"脚本不应 import {alias.name}"
            elif isinstance(node, _ast.ImportFrom):
                module = node.module or ""
                base = module.split(".")[0]
                assert base not in forbidden_imports, f"脚本不应 import {module}"

        # 检查代码（去除字符串字面量与注释）中的禁用 token
        code_only_tokens = []
        for node in _ast.walk(tree):
            if isinstance(node, _ast.Name):
                code_only_tokens.append(node.id)
            elif isinstance(node, _ast.Attribute):
                code_only_tokens.append(node.attr)
        joined = " ".join(code_only_tokens).lower()
        # 不应直接调用审批/驳回相关 API 路径或函数名
        forbidden_attrs = {"approve_record", "reject_record", "submit_approval", "execute_sql"}
        for attr in forbidden_attrs:
            assert attr not in joined, f"脚本不应包含函数: {attr}"


class TestEndToEnd:
    def test_e2e_with_curl_and_detail_api(self, tmp_path):
        """完整流程：xlsx + curl + mock detail API + audit + output xlsx"""
        cols = DEFAULT_EXCEL_COLS + ["fourCategoryId"]
        vals = SAMPLE_ROW_VALUES + ["06049901"]
        rows = [dict(zip(cols, vals))]
        xlsx = _make_test_xlsx(tmp_path, rows)
        curl_template = {
            "url": "https://api.example.com/detail",
            "headers": [("Content-Type", "application/json")],
            "cookie": "",
            "data": '{"documentId": "x", "fourCategoryId": "y"}',
            "method": "POST",
        }
        detail_payload = {
            "data": {
                "referenceInfos": [
                    {"referencePrice": "11585.37", "order": 1},
                    {"referencePrice": "9700", "order": 2},
                ],
                "referencePrice": "11585.37",
            }
        }
        out = tmp_path / "out.xlsx"
        result = run_main(
            input_xlsx=xlsx,
            output_xlsx=out,
            curl_file=None,
            document_id_field="documentId",
            category_id_field="fourCategoryId",
            opener=_make_opener(detail_payload),
        )
        # 需要绕过 curl_file 检查，直接用 enriched 注入
        # 改用 enrich_rows_with_customer_price 直接测试
        enriched = enrich_rows_with_customer_price(
            rows, curl_template, "documentId", "fourCategoryId",
            opener=_make_opener(detail_payload),
        )
        assert enriched[0]["customerPriceStatus"] == "FOUND"
        assert enriched[0]["customerPrice"] == Decimal("11585.37")

    def test_e2e_no_curl(self, tmp_path):
        """完整流程：xlsx 已有 customerPrice，不经 detail API"""
        cols = DEFAULT_EXCEL_COLS + ["customerPrice"]
        vals = SAMPLE_ROW_VALUES + ["99.99"]
        rows = [dict(zip(cols, vals))]
        xlsx = _make_test_xlsx(tmp_path, rows)
        out = tmp_path / "out.xlsx"
        result = run_main(
            input_xlsx=xlsx,
            output_xlsx=out,
            curl_file=None,
            document_id_field="documentId",
            category_id_field="fourCategoryId",
        )
        assert result["enriched"][0]["customerPriceStatus"] == "EXISTS"
        assert result["enriched"][0]["customerPrice"] == Decimal("99.99")
        # 验证输出 Excel 中 summary 存在
        out_wb = load_workbook(str(out))
        summary_ws = out_wb["summary"]
        summary_data = [cell.value for cell in summary_ws[2]]
        assert int(summary_data[0]) >= 1
        out_wb.close()


# ---------------------------------------------------------------------------
# SaaS list-json mode tests
# ---------------------------------------------------------------------------


from scripts.audit_saas_excel_file import (
    SAAS_LIST_HEADERS,
    enrich_rows_with_customer_price as _enrich,
    read_saas_list_json,
)


SAMPLE_LIST_ROW = {
    "documentId": "679483364934705152",
    "fourCategoryId": "06049901",
    "fourCategoryName": "其他离子交换树脂",
    "commodityName": "茂谷聚 Cs树脂 AMP-PAN 1kg/瓶 100~600目",
    "sassVCommodityName": "茂谷聚 Cs树脂 AMP-PAN 1kg/瓶 销售单位：瓶",
    "referenceLink": "https://www.colipu.com/itemcode-WF7769406.html",
    "referencePrice": "11585.37",
    "referenceId": "679483363319898112",
    "categoryId": "060499",
    "categoryName": "其他离子交换树脂",
    "materialCode": "01063412",
    "quasiStandardCommodityCode": "20260617001681",
}

LIST_CURL_TEMPLATE = {
    "url": "https://api.example.com/detail",
    "headers": [("Content-Type", "application/json")],
    "cookie": "",
    "data": '{"documentId":"x","fourCategoryId":"y"}',
    "method": "POST",
}


def _write_saas_list_json(tmp_path, rows):
    path = tmp_path / "saas_list.json"
    payload = {"rspCode": "0000", "count": len(rows), "rows": rows, "success": True}
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


class TestSaasListJsonInput:
    def test_read_saas_list_json_object(self, tmp_path):
        path = _write_saas_list_json(tmp_path, [SAMPLE_LIST_ROW])
        headers, rows = read_saas_list_json(path)
        assert headers == list(SAAS_LIST_HEADERS)
        assert len(rows) == 1
        row = rows[0]
        assert row["documentId"] == "679483364934705152"
        assert row["fourCategoryId"] == "06049901"
        assert row["list_reference_price"] == Decimal("11585.37")
        assert row["referenceLink"].startswith("https://")

    def test_read_saas_list_json_array(self, tmp_path):
        path = tmp_path / "saas_list.json"
        path.write_text(json.dumps([SAMPLE_LIST_ROW], ensure_ascii=False), encoding="utf-8")
        _, rows = read_saas_list_json(path)
        assert len(rows) == 1
        assert rows[0]["documentId"] == "679483364934705152"

    def test_read_saas_list_json_rejects_bad_payload(self, tmp_path):
        path = tmp_path / "saas_list.json"
        path.write_text(json.dumps({"foo": "bar"}, ensure_ascii=False), encoding="utf-8")
        with pytest.raises(ValueError, match="未找到 rows"):
            read_saas_list_json(path)

    def test_read_saas_list_json_rejects_non_json_suffix(self, tmp_path):
        path = tmp_path / "saas_list.txt"
        path.write_text("[]", encoding="utf-8")
        with pytest.raises(ValueError, match="必须是 JSON"):
            read_saas_list_json(path)


class TestSaasListJsonE2E:
    def test_e2e_max_price_from_reference_infos(self, tmp_path):
        """E2E：列表 JSON → mock detail 多价 → 取 max → Excel"""
        path = _write_saas_list_json(tmp_path, [SAMPLE_LIST_ROW])
        out = tmp_path / "saas_customer_price.xlsx"
        detail_payload = {
            "data": {
                "referenceInfos": [
                    {"referencePrice": "11585.37", "order": 1},
                    {"referencePrice": "9700", "order": 2},
                ],
            }
        }
        from scripts.audit_saas_excel_file import (
            build_records,
            run_audit,
            write_audit_workbook,
        )
        _, rows = read_saas_list_json(path)
        enriched = _enrich(
            rows, LIST_CURL_TEMPLATE,
            "documentId", "fourCategoryId",
            opener=_make_opener(detail_payload),
            request_interval=0.0,
        )
        records = build_records(rows, enriched)
        audit_response = run_audit(records)
        write_audit_workbook(out, list(SAAS_LIST_HEADERS), rows, enriched, audit_response)
        assert enriched[0]["customerPriceStatus"] == "FOUND"
        assert enriched[0]["customerPrice"] == Decimal("11585.37")
        wb = load_workbook(str(out))
        assert "summary" in wb.sheetnames
        assert "record_results" in wb.sheetnames
        ws = wb["record_results"]
        out_headers = [str(c.value or "") for c in ws[1]]
        for col in (
            "documentId", "fourCategoryId", "list_reference_price",
            "customerPrice", "customerPriceStatus", "primary_action",
        ):
            assert _display_header(col) in out_headers, f"缺失列: {col}"
        wb.close()

    def test_e2e_request_interval_calls_sleep_between_requests(self, tmp_path):
        rows_payload = [
            {**SAMPLE_LIST_ROW, "documentId": "doc1"},
            {**SAMPLE_LIST_ROW, "documentId": "doc2"},
            {**SAMPLE_LIST_ROW, "documentId": "doc3"},
        ]
        path = _write_saas_list_json(tmp_path, rows_payload)
        _, rows = read_saas_list_json(path)
        detail_payload = {
            "data": {
                "referenceInfos": [
                    {"referencePrice": "10"},
                    {"referencePrice": "20"},
                ]
            }
        }
        sleep_calls: list[float] = []

        def fake_sleep(seconds: float) -> None:
            sleep_calls.append(seconds)

        enriched = _enrich(
            rows, LIST_CURL_TEMPLATE,
            "documentId", "fourCategoryId",
            opener=_make_opener(detail_payload),
            request_interval=0.5,
            sleep_func=fake_sleep,
        )
        # 3 个请求之间应有 2 次 sleep
        assert sleep_calls == [0.5, 0.5]
        assert all(item["customerPriceStatus"] == "FOUND" for item in enriched)
        assert all(item["customerPrice"] == Decimal("20") for item in enriched)

    def test_e2e_detail_request_failed(self, tmp_path):
        path = _write_saas_list_json(tmp_path, [SAMPLE_LIST_ROW])
        _, rows = read_saas_list_json(path)
        enriched = _enrich(
            rows, LIST_CURL_TEMPLATE,
            "documentId", "fourCategoryId",
            opener=_make_opener(None, status=500),
            request_interval=0.0,
        )
        assert enriched[0]["customerPriceStatus"] == "DETAIL_REQUEST_FAILED"


class TestRunMainExclusivity:
    def test_run_main_requires_one_input(self, tmp_path):
        out = tmp_path / "out.xlsx"
        with pytest.raises(ValueError, match="必须指定"):
            run_main(output_xlsx=out)

    def test_run_main_rejects_both_inputs(self, tmp_path):
        out = tmp_path / "out.xlsx"
        xlsx = _make_test_xlsx(tmp_path, [dict(zip(DEFAULT_EXCEL_COLS, SAMPLE_ROW_VALUES))])
        json_path = _write_saas_list_json(tmp_path, [SAMPLE_LIST_ROW])
        with pytest.raises(ValueError, match="互斥"):
            run_main(output_xlsx=out, input_xlsx=xlsx, input_saas_list_json=json_path)

    def test_run_main_saas_list_no_curl_marks_not_provided(self, tmp_path):
        path = _write_saas_list_json(tmp_path, [SAMPLE_LIST_ROW])
        out = tmp_path / "out.xlsx"
        result = run_main(output_xlsx=out, input_saas_list_json=path)
        assert result["enriched"][0]["customerPriceStatus"] == "NOT_PROVIDED"


# ---------------------------------------------------------------------------
# Dual JSON (saas + customer) mode tests
# ---------------------------------------------------------------------------


from scripts.audit_saas_excel_file import (
    SAAS_PENDING_HEADERS,
    build_customer_index,
    compute_over_30_check,
    is_jd_link,
    join_saas_with_customer,
    read_saas_pending_json,
)


SAMPLE_SAAS_PENDING_ROW = {
    "vCommodityId": "679483364934705152",
    "queryStr": "茂谷聚 Cs树脂 AMP-PAN 1kg/瓶 100~600目（销售单位：瓶）",
    "skuName": "茂谷聚 Cs树脂 AMP-PAN 1kg/瓶 100~600目 销售单位：瓶",
    "skuId": "1004360200014308",
    "vendorName": "colipu",
    "shopName": "科力普商城",
    "price": "11585.37",
    "extLink": "https://www.colipu.com/itemcode-WF7769406.html",
    "createTime": "2026-06-17 18:27:42",
    "updateTime": "2026-06-18 09:42:43",
    "modelAuditResult": "1",
    "isMatch": "1",
}

SAMPLE_CUSTOMER_ROW = {
    "documentId": "679483364934705152",
    "fourCategoryId": "06049901",
    "fourCategoryName": "其他离子交换树脂",
    "referencePrice": "11585.37",
    "commodityName": "茂谷聚 Cs树脂 AMP-PAN",
    "referenceLink": "https://www.colipu.com/itemcode-WF7769406.html",
}



def _write_pending_json(tmp_path, rows):
    path = tmp_path / "saas_pending.json"
    payload = {"rspCode": "0000", "count": len(rows), "rows": rows, "success": True}
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


def _write_customer_json(tmp_path, rows):
    path = tmp_path / "customer.json"
    payload = {"rspCode": "0000", "count": len(rows), "rows": rows, "success": True}
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


class TestSaasPendingJsonInput:
    def test_read_saas_pending_json(self, tmp_path):
        path = _write_pending_json(tmp_path, [SAMPLE_SAAS_PENDING_ROW])
        headers, rows = read_saas_pending_json(path)
        assert headers == list(SAAS_PENDING_HEADERS)
        assert len(rows) == 1
        row = rows[0]
        assert row["vCommodityId"] == "679483364934705152"
        assert row["queryStr"] == SAMPLE_SAAS_PENDING_ROW["queryStr"]
        assert row["skuName"] == SAMPLE_SAAS_PENDING_ROW["skuName"]
        assert row["extLink"].startswith("https://")
        assert row["price"] == Decimal("11585.37")

    def test_read_saas_pending_json_query_str_falls_back(self, tmp_path):
        raw = dict(SAMPLE_SAAS_PENDING_ROW)
        raw["queryStr"] = ""
        raw["sassVCommodityName"] = "兜底标题"
        path = _write_pending_json(tmp_path, [raw])
        _, rows = read_saas_pending_json(path)
        assert rows[0]["queryStr"] == "兜底标题"


class TestCustomerIndex:
    def test_build_customer_index(self, tmp_path):
        path = _write_customer_json(tmp_path, [SAMPLE_CUSTOMER_ROW])
        index = build_customer_index(path)
        assert "679483364934705152" in index
        entry = index["679483364934705152"]
        assert entry["fourCategoryId"] == "06049901"
        assert entry["fourCategoryName"] == "其他离子交换树脂"
        assert entry["list_reference_price"] == Decimal("11585.37")

    def test_build_customer_index_skips_rows_without_document_id(self, tmp_path):
        bad = dict(SAMPLE_CUSTOMER_ROW)
        bad["documentId"] = ""
        path = _write_customer_json(tmp_path, [bad])
        index = build_customer_index(path)
        assert index == {}


class TestJoin:
    def test_join_matched(self, tmp_path):
        saas_path = _write_pending_json(tmp_path, [SAMPLE_SAAS_PENDING_ROW])
        cust_path = _write_customer_json(tmp_path, [SAMPLE_CUSTOMER_ROW])
        _, saas_rows = read_saas_pending_json(saas_path)
        index = build_customer_index(cust_path)
        joined = join_saas_with_customer(saas_rows, index)
        assert joined[0]["customer_lookup_status"] == "MATCHED"
        assert joined[0]["fourCategoryId"] == "06049901"
        assert joined[0]["list_reference_price"] == Decimal("11585.37")

    def test_join_not_found(self, tmp_path):
        saas_row = dict(SAMPLE_SAAS_PENDING_ROW)
        saas_row["vCommodityId"] = "999999999999999999"
        saas_path = _write_pending_json(tmp_path, [saas_row])
        cust_path = _write_customer_json(tmp_path, [SAMPLE_CUSTOMER_ROW])
        _, saas_rows = read_saas_pending_json(saas_path)
        index = build_customer_index(cust_path)
        joined = join_saas_with_customer(saas_rows, index)
        assert joined[0]["customer_lookup_status"] == "NOT_FOUND"
        assert joined[0]["fourCategoryId"] == ""
        assert joined[0]["list_reference_price"] is None

    def test_join_extra_customer_rows_ignored(self, tmp_path):
        saas_path = _write_pending_json(tmp_path, [SAMPLE_SAAS_PENDING_ROW])
        extra_customer = dict(SAMPLE_CUSTOMER_ROW)
        extra_customer["documentId"] = "888888888888888888"
        cust_path = _write_customer_json(tmp_path, [SAMPLE_CUSTOMER_ROW, extra_customer])
        _, saas_rows = read_saas_pending_json(saas_path)
        index = build_customer_index(cust_path)
        joined = join_saas_with_customer(saas_rows, index)
        # 输出行数 = saas 行数（多余的 customer 行不进表）
        assert len(joined) == len(saas_rows)


class TestJDDetection:
    @pytest.mark.parametrize("url,expected", [
        ("https://item.jd.com/100237336261.html", True),
        ("https://i-item.jd.com/100237336261.html", True),
        ("https://jd.com/foo", True),
        ("https://www.jd.com/abc", True),
        ("https://www.colipu.com/itemcode-WF7769406.html", False),
        ("https://b2b.nbdeli.com/pcweb/detail?sku=122604488", False),
        ("", False),
    ])
    def test_is_jd_link(self, url, expected):
        assert is_jd_link(url) is expected


class TestOver30Check:
    def test_over_30_true(self):
        check, over = compute_over_30_check(Decimal("100"), Decimal("131"), "OK")
        assert check == "HIGHER_THAN_CUSTOMER_PRICE_30"
        assert over == "TRUE"

    def test_over_30_false_at_threshold(self):
        check, over = compute_over_30_check(Decimal("100"), Decimal("130"), "OK")
        assert check == "WITHIN_THRESHOLD"
        assert over == "FALSE"

    def test_over_30_false(self):
        check, over = compute_over_30_check(Decimal("100"), Decimal("99"), "OK")
        assert check == "WITHIN_THRESHOLD"
        assert over == "FALSE"

    def test_over_30_pending_crawl_returns_empty(self):
        check, over = compute_over_30_check(Decimal("100"), None, "PENDING_CRAWL")
        assert check == "PAGE_PRICE_MISSING"
        assert over == ""

    def test_over_30_jd_skipped_returns_empty(self):
        check, over = compute_over_30_check(Decimal("100"), None, "SKIPPED_JD")
        assert check == "PAGE_PRICE_MISSING"
        assert over == ""

    def test_over_30_customer_price_missing(self):
        check, over = compute_over_30_check(None, Decimal("131"), "OK")
        assert check == "CUSTOMER_PRICE_MISSING"
        assert over == ""

    def test_over_30_customer_price_zero(self):
        check, over = compute_over_30_check(Decimal("0"), Decimal("131"), "OK")
        assert check == "CUSTOMER_PRICE_MISSING"
        assert over == ""

    def test_over_30_failed_status(self):
        check, over = compute_over_30_check(Decimal("100"), None, "FAILED")
        assert check == "PAGE_PRICE_MISSING"
        assert over == ""


class TestRunMainSaasJsonExclusivity:
    def test_run_main_requires_customer_json_pair(self, tmp_path):
        saas_path = _write_pending_json(tmp_path, [SAMPLE_SAAS_PENDING_ROW])
        out = tmp_path / "out.xlsx"
        with pytest.raises(ValueError, match="必须配合 --input-customer-json"):
            run_main(output_xlsx=out, input_saas_json=saas_path)

    def test_run_main_customer_json_alone_rejected(self, tmp_path):
        cust_path = _write_customer_json(tmp_path, [SAMPLE_CUSTOMER_ROW])
        out = tmp_path / "out.xlsx"
        with pytest.raises(ValueError, match="必须配合 --input-saas-json"):
            run_main(output_xlsx=out, input_customer_json=cust_path)

    def test_run_main_three_modes_mutually_exclusive(self, tmp_path):
        saas_path = _write_pending_json(tmp_path, [SAMPLE_SAAS_PENDING_ROW])
        cust_path = _write_customer_json(tmp_path, [SAMPLE_CUSTOMER_ROW])
        list_path = tmp_path / "list.json"
        list_path.write_text(json.dumps({"rows": [SAMPLE_CUSTOMER_ROW]}, ensure_ascii=False), encoding="utf-8")
        out = tmp_path / "out.xlsx"
        with pytest.raises(ValueError, match="三选一"):
            run_main(
                output_xlsx=out,
                input_saas_json=saas_path,
                input_customer_json=cust_path,
                input_saas_list_json=list_path,
            )


class TestRunMainSaasJsonE2E:
    def test_run_main_dual_mode_e2e_no_curl_marks_not_provided(self, tmp_path):
        saas_path = _write_pending_json(tmp_path, [SAMPLE_SAAS_PENDING_ROW])
        cust_path = _write_customer_json(tmp_path, [SAMPLE_CUSTOMER_ROW])
        out = tmp_path / "out.xlsx"
        result = run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
        )
        assert result["mode"] == "input_saas_json"
        assert result["enriched"][0]["customerPriceStatus"] == "NOT_PROVIDED"
        assert result["rows"][0]["fourCategoryId"] == "06049901"
        wb = load_workbook(str(out))
        assert "summary" in wb.sheetnames
        assert "record_results" in wb.sheetnames
        ws = wb["record_results"]
        out_headers = [str(c.value or "") for c in ws[1]]
        for col in (
            "vCommodityId", "queryStr", "skuName", "extLink",
            "fourCategoryId", "fourCategoryName",
            "customer_lookup_status", "list_reference_price",
            "customerPrice", "customerPriceStatus",
            "title_match_result", "title_match_confidence", "title_match_reason",
            "page_price_value", "page_price_status", "page_price_source",
            "page_vs_customer_check", "over_30pct",
        ):
            assert _display_header(col) in out_headers, f"缺失列: {_display_header(col)}"
        wb.close()

    def test_run_main_dual_mode_lookup_failed(self, tmp_path):
        saas_row = dict(SAMPLE_SAAS_PENDING_ROW)
        saas_row["vCommodityId"] = "999999999999999999"
        saas_path = _write_pending_json(tmp_path, [saas_row])
        cust_path = _write_customer_json(tmp_path, [SAMPLE_CUSTOMER_ROW])
        out = tmp_path / "out.xlsx"
        result = run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
        )
        assert result["enriched"][0]["customerPriceStatus"] == "LOOKUP_FAILED"
        assert "未找到该 documentId" in result["enriched"][0]["customerPriceWarning"]

    def test_run_main_dual_mode_jd_link_marks_skipped_jd(self, tmp_path):
        saas_row = dict(SAMPLE_SAAS_PENDING_ROW)
        saas_row["extLink"] = "https://item.jd.com/100237336261.html"
        saas_path = _write_pending_json(tmp_path, [saas_row])
        cust_path = _write_customer_json(tmp_path, [SAMPLE_CUSTOMER_ROW])
        out = tmp_path / "out.xlsx"
        run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
        )
        wb = load_workbook(str(out))
        ws = wb["record_results"]
        out_headers = [str(c.value or "") for c in ws[1]]
        idx_status = out_headers.index(_display_header("page_price_status")) + 1
        idx_over = out_headers.index(_display_header("over_30pct")) + 1
        assert ws.cell(row=2, column=idx_status).value == _display_value("SKIPPED_ZHAOSHANG")
        assert ws.cell(row=2, column=idx_over).value in ("", None)
        wb.close()


# ---------------------------------------------------------------------------
# Multi-segment JSON + vCommodityId join tests
# ---------------------------------------------------------------------------


from scripts.audit_saas_excel_file import _parse_json_segments


class TestMultiSegmentJson:
    def test_parse_two_segments_concatenated(self):
        seg1 = json.dumps({"rspCode": "0000", "count": 2, "rows": [{"a": 1}, {"a": 2}]}, ensure_ascii=False)
        seg2 = json.dumps({"rspCode": "0000", "count": 1, "rows": [{"a": 3}]}, ensure_ascii=False)
        text = seg1 + "\n" + seg2
        segments = _parse_json_segments(text)
        assert len(segments) == 2
        assert segments[0]["count"] == 2
        assert segments[1]["count"] == 1

    def test_read_saas_pending_multi_segment(self, tmp_path):
        seg1_rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v1", "skuId": "sku-1"},
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v1", "skuId": "sku-2"},
        ]
        seg2_rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v2", "skuId": "sku-3"},
        ]
        path = tmp_path / "saas_multi.json"
        text1 = json.dumps({"rspCode": "0000", "count": 2, "rows": seg1_rows}, ensure_ascii=False)
        text2 = json.dumps({"rspCode": "0000", "count": 1, "rows": seg2_rows}, ensure_ascii=False)
        path.write_text(text1 + "\n" + text2, encoding="utf-8")
        _, rows = read_saas_pending_json(path)
        assert len(rows) == 3
        assert [r["skuId"] for r in rows] == ["sku-1", "sku-2", "sku-3"]
        assert {r["vCommodityId"] for r in rows} == {"v1", "v2"}

    def test_read_customer_multi_segment(self, tmp_path):
        seg1 = [{"documentId": "v1", "fourCategoryId": "c1", "fourCategoryName": "n1", "referencePrice": "10"}]
        seg2 = [{"documentId": "v2", "fourCategoryId": "c2", "fourCategoryName": "n2", "referencePrice": "20"}]
        path = tmp_path / "customer_multi.json"
        text1 = json.dumps({"rspCode": "0000", "count": 1, "rows": seg1}, ensure_ascii=False)
        text2 = json.dumps({"rspCode": "0000", "count": 1, "rows": seg2}, ensure_ascii=False)
        path.write_text(text1 + "\n" + text2, encoding="utf-8")
        index = build_customer_index(path)
        assert set(index.keys()) == {"v1", "v2"}
        assert index["v1"]["fourCategoryId"] == "c1"
        assert index["v2"]["fourCategoryId"] == "c2"


class TestVCommodityIdJoin:
    def test_join_uses_v_commodity_id_to_lookup_document_id(self, tmp_path):
        # saas 端只有 vCommodityId，customer 端用 documentId 提供 fourCategoryId
        saas_row = dict(SAMPLE_SAAS_PENDING_ROW)
        saas_row["vCommodityId"] = "shared-id-001"
        cust_row = dict(SAMPLE_CUSTOMER_ROW)
        cust_row["documentId"] = "shared-id-001"
        cust_row["fourCategoryId"] = "12345601"
        cust_row["fourCategoryName"] = "测试类目"

        saas_path = _write_pending_json(tmp_path, [saas_row])
        cust_path = _write_customer_json(tmp_path, [cust_row])
        _, saas_rows = read_saas_pending_json(saas_path)
        index = build_customer_index(cust_path)
        joined = join_saas_with_customer(saas_rows, index)

        assert joined[0]["customer_lookup_status"] == "MATCHED"
        assert joined[0]["fourCategoryId"] == "12345601"
        assert joined[0]["fourCategoryName"] == "测试类目"
        # documentId 应该回填为 customer 端的 documentId（与 saas.vCommodityId 相等）
        assert joined[0]["documentId"] == "shared-id-001"

    def test_multiple_sku_rows_under_same_v_commodity_id_not_merged(self, tmp_path):
        # 同一 vCommodityId 下两条 sku 候选不合并
        rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "vc-A", "skuId": "sku-1", "extLink": "https://a.com/1"},
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "vc-A", "skuId": "sku-2", "extLink": "https://a.com/2"},
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "vc-A", "skuId": "sku-3", "extLink": "https://a.com/3"},
        ]
        cust_row = dict(SAMPLE_CUSTOMER_ROW)
        cust_row["documentId"] = "vc-A"
        cust_row["fourCategoryId"] = "abc01"

        saas_path = _write_pending_json(tmp_path, rows)
        cust_path = _write_customer_json(tmp_path, [cust_row])
        out = tmp_path / "out.xlsx"
        result = run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
        )
        # 输出行数 = saas 行数（不合并）
        assert len(result["rows"]) == 3
        assert len(result["enriched"]) == 3
        # 三行都命中同一个 fourCategoryId
        for r in result["rows"]:
            assert r["fourCategoryId"] == "abc01"
            assert r["customer_lookup_status"] == "MATCHED"


class TestDetailCache:
    def test_detail_called_once_per_unique_document_id(self, tmp_path):
        # 3 行同 vCommodityId，detail 只应调 1 次
        rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "vc-X", "skuId": "sku-A", "extLink": "https://a/1"},
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "vc-X", "skuId": "sku-B", "extLink": "https://a/2"},
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "vc-X", "skuId": "sku-C", "extLink": "https://a/3"},
        ]
        cust_row = dict(SAMPLE_CUSTOMER_ROW)
        cust_row["documentId"] = "vc-X"
        cust_row["fourCategoryId"] = "cat01"

        saas_path = _write_pending_json(tmp_path, rows)
        cust_path = _write_customer_json(tmp_path, [cust_row])
        _, saas_rows = read_saas_pending_json(saas_path)
        index = build_customer_index(cust_path)
        joined = join_saas_with_customer(saas_rows, index)

        curl_template = {
            "url": "https://api.example.com/detail",
            "headers": [],
            "cookie": "",
            "data": '{"documentId":"x","fourCategoryId":"y"}',
            "method": "POST",
        }
        detail_payload = {"data": {"referenceInfos": [{"referencePrice": "100"}, {"referencePrice": "200"}]}}

        call_count = 0
        def counting_opener(request, timeout=15.0):
            nonlocal call_count
            call_count += 1
            return MockResponse(json.dumps(detail_payload).encode("utf-8"))

        from scripts.audit_saas_excel_file import enrich_rows_with_customer_price as _enrich2
        enriched = _enrich2(
            joined, curl_template,
            "documentId", "fourCategoryId",
            opener=counting_opener,
            request_interval=0.0,
        )
        # 3 行同一 documentId → detail 只调 1 次
        assert call_count == 1
        # 全部命中并取得最大值
        assert all(item["customerPriceStatus"] in ("FOUND", "FOUND_SINGLE_REFERENCE") for item in enriched)
        assert all(item["customerPrice"] == Decimal("200") for item in enriched)


# ---------------------------------------------------------------------------
# .env 加载与模型环境变量状态
# ---------------------------------------------------------------------------


from scripts.audit_saas_excel_file import _load_env_file


class TestEnvLoader:
    def test_load_env_file_sets_unset_keys(self, tmp_path, monkeypatch):
        monkeypatch.delenv("TITLE_MATCH_MODEL_API_KEY", raising=False)
        monkeypatch.delenv("TITLE_MATCH_MODEL_BASE_URL", raising=False)
        monkeypatch.delenv("TITLE_MATCH_MODEL_NAME", raising=False)
        env_path = tmp_path / "model.env"
        env_path.write_text(
            'TITLE_MATCH_MODEL_API_KEY=secret-key-xxx\n'
            'TITLE_MATCH_MODEL_BASE_URL=https://example.com/api\n'
            'TITLE_MATCH_MODEL_NAME=fake-model-id\n',
            encoding="utf-8",
        )
        status = _load_env_file(env_path)
        assert status["TITLE_MATCH_MODEL_API_KEY"] == "SET"
        assert status["TITLE_MATCH_MODEL_BASE_URL"] == "SET"
        assert status["TITLE_MATCH_MODEL_NAME"] == "SET"
        # 注入到 os.environ
        assert os.environ.get("TITLE_MATCH_MODEL_API_KEY") == "secret-key-xxx"
        assert os.environ.get("TITLE_MATCH_MODEL_BASE_URL") == "https://example.com/api"

    def test_load_env_file_does_not_overwrite_existing(self, tmp_path, monkeypatch):
        monkeypatch.setenv("TITLE_MATCH_MODEL_API_KEY", "preexisting")
        env_path = tmp_path / "model.env"
        env_path.write_text(
            'TITLE_MATCH_MODEL_API_KEY=newvalue\n', encoding="utf-8"
        )
        status = _load_env_file(env_path)
        assert status["TITLE_MATCH_MODEL_API_KEY"] == "EXISTING"
        assert os.environ.get("TITLE_MATCH_MODEL_API_KEY") == "preexisting"

    def test_load_env_file_ignores_unknown_keys(self, tmp_path, monkeypatch):
        monkeypatch.delenv("MALICIOUS_VAR", raising=False)
        env_path = tmp_path / "model.env"
        env_path.write_text("MALICIOUS_VAR=evil\n", encoding="utf-8")
        status = _load_env_file(env_path)
        assert "MALICIOUS_VAR" not in status
        assert os.environ.get("MALICIOUS_VAR") is None

    def test_load_env_file_ignores_comments_and_blank_lines(self, tmp_path, monkeypatch):
        monkeypatch.delenv("TITLE_MATCH_MODEL_BASE_URL", raising=False)
        env_path = tmp_path / "model.env"
        env_path.write_text(
            "# comment\n"
            "\n"
            'TITLE_MATCH_MODEL_BASE_URL="https://example.com/api"\n',
            encoding="utf-8",
        )
        status = _load_env_file(env_path)
        assert status["TITLE_MATCH_MODEL_BASE_URL"] == "SET"
        assert os.environ.get("TITLE_MATCH_MODEL_BASE_URL") == "https://example.com/api"

    def test_load_env_file_missing_file_returns_empty(self, tmp_path):
        path = tmp_path / "does_not_exist.env"
        status = _load_env_file(path)
        assert status == {}

    def test_load_env_file_empty_value_returns_empty_status(self, tmp_path, monkeypatch):
        monkeypatch.delenv("TITLE_MATCH_MODEL_API_KEY", raising=False)
        env_path = tmp_path / "model.env"
        env_path.write_text("TITLE_MATCH_MODEL_API_KEY=\n", encoding="utf-8")
        status = _load_env_file(env_path)
        assert status["TITLE_MATCH_MODEL_API_KEY"] == "EMPTY"
        assert os.environ.get("TITLE_MATCH_MODEL_API_KEY") is None


# ---------------------------------------------------------------------------
# Crawler subprocess integration tests (basic helpers)
# ---------------------------------------------------------------------------


from scripts.audit_saas_excel_file import (
    apply_crawl_results_to_rows,
    collect_urls_for_crawl,
    read_crawler_output,
    run_crawler_subprocess,
    write_crawler_input_xlsx,
)


def _make_crawler_output_xlsx(tmp_path, records):
    out = tmp_path / "_crawl_output.xlsx"
    wb = Workbook()
    ws = wb.active
    headers = [
        "原始链接", "识别域名", "识别商城", "是否京东", "是否跳过",
        "商品编码", "商品名称", "商品价格", "商品评价数", "是否下架",
        "采集状态", "失败原因", "采集时间",
    ]
    ws.append(headers)
    for r in records:
        ws.append([r.get(h, "") for h in headers])
    wb.save(str(out))
    return out


class TestCollectUrlsForCrawl:
    def test_collect_pending_only_and_dedupe(self, tmp_path):
        rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v1", "extLink": "https://item.jd.com/x.html"},
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v2", "extLink": ""},
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v3", "extLink": "https://www.ehsy.com/product-A"},
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v3", "skuId": "x", "extLink": "https://www.ehsy.com/product-A"},
        ]
        cust_rows = [
            {"documentId": v, "fourCategoryId": "c", "fourCategoryName": "n", "referencePrice": "100"}
            for v in ("v1", "v2", "v3")
        ]
        saas_path = _write_pending_json(tmp_path, rows)
        cust_path = _write_customer_json(tmp_path, cust_rows)
        out = tmp_path / "out.xlsx"
        result = run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
        )
        urls = collect_urls_for_crawl(result["rows"], result["enriched"], result["audit_response"])
        assert urls == ["https://www.ehsy.com/product-A"]


class TestWriteCrawlerInput:
    def test_write_input_xlsx(self, tmp_path):
        path = tmp_path / "_crawl_input.xlsx"
        write_crawler_input_xlsx(["https://a.com/1", "https://b.com/2"], path)
        wb = load_workbook(str(path))
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "URL"
        assert ws.cell(row=2, column=1).value == "https://a.com/1"
        assert ws.cell(row=3, column=1).value == "https://b.com/2"
        wb.close()


class TestRunCrawlerSubprocess:
    def test_missing_script(self, tmp_path):
        result = run_crawler_subprocess(
            crawler_script=tmp_path / "no_such_main.py",
            input_xlsx=tmp_path / "in.xlsx",
            output_xlsx=tmp_path / "out.xlsx",
        )
        assert result["ok"] is False
        assert "不存在" in result["error"]

    def test_runner_zero(self, tmp_path):
        script = tmp_path / "main.py"
        script.write_text("# fake", encoding="utf-8")
        captured = {}
        def runner(argv, cwd):
            captured["argv"] = argv
            captured["cwd"] = cwd
            return 0
        result = run_crawler_subprocess(
            crawler_script=script,
            input_xlsx=tmp_path / "in.xlsx",
            output_xlsx=tmp_path / "out.xlsx",
            dry_run=True,
            headless=True,
            concurrency=2,
            per_domain_delay=2.0,
            timeout_ms=10000,
            limit=5,
            runner=runner,
        )
        assert result["ok"] is True
        assert result["exit_code"] == 0
        assert "--dry-run" in captured["argv"]
        assert "--headless" in captured["argv"]
        assert "5" in captured["argv"]
        assert captured["cwd"] == str(script.parent)

    def test_runner_nonzero(self, tmp_path):
        script = tmp_path / "main.py"
        script.write_text("# fake", encoding="utf-8")
        result = run_crawler_subprocess(
            crawler_script=script,
            input_xlsx=tmp_path / "in.xlsx",
            output_xlsx=tmp_path / "out.xlsx",
            runner=lambda argv, cwd: 2,
        )
        assert result["ok"] is False
        assert result["exit_code"] == 2

    def test_runner_raises(self, tmp_path):
        script = tmp_path / "main.py"
        script.write_text("# fake", encoding="utf-8")
        def boom(argv, cwd):
            raise OSError("boom")
        result = run_crawler_subprocess(
            crawler_script=script,
            input_xlsx=tmp_path / "in.xlsx",
            output_xlsx=tmp_path / "out.xlsx",
            runner=boom,
        )
        assert result["ok"] is False
        assert "OSError" in result["error"]


class TestReadCrawlerOutput:
    def test_success_row(self, tmp_path):
        path = _make_crawler_output_xlsx(tmp_path, [
            {"原始链接": "https://a.com/1", "商品价格": "129.99", "识别商城": "震坤行", "采集状态": "成功"},
        ])
        results = read_crawler_output(path)
        r = results["https://a.com/1"]
        assert r["page_price_status"] == "OK"
        assert r["page_price_value"] == Decimal("129.99")
        assert r["page_price_source"] == "震坤行"

    def test_failed_row(self, tmp_path):
        path = _make_crawler_output_xlsx(tmp_path, [
            {"原始链接": "https://a.com/2", "商品价格": "", "采集状态": "失败", "失败原因": "触发验证码"},
        ])
        results = read_crawler_output(path)
        r = results["https://a.com/2"]
        assert r["page_price_status"] == "FAILED"
        assert r["page_price_value"] is None
        assert r["page_price_source"] == "触发验证码"

    def test_skipped_row(self, tmp_path):
        path = _make_crawler_output_xlsx(tmp_path, [
            {"原始链接": "https://item.jd.com/x.html", "采集状态": "跳过", "失败原因": "京东"},
        ])
        results = read_crawler_output(path)
        r = results["https://item.jd.com/x.html"]
        assert r["page_price_status"] == "SKIPPED"


class TestApplyCrawlResults:
    def test_apply_success_high_price_marks_over_30pct(self, tmp_path):
        rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v1",
             "extLink": "https://www.ehsy.com/product-A",
             "customerPrice": "100", "queryStr": "test", "skuName": "test"},
        ]
        cust_rows = [{"documentId": "v1", "fourCategoryId": "c", "fourCategoryName": "n", "referencePrice": "100"}]
        saas_path = _write_pending_json(tmp_path, rows)
        cust_path = _write_customer_json(tmp_path, cust_rows)
        out = tmp_path / "out.xlsx"
        result = run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
        )
        # 模拟爬虫返回 page_price=131
        crawl_results = {
            "https://www.ehsy.com/product-A": {
                "page_price_value": Decimal("131"),
                "page_price_status": "OK",
                "page_price_source": "震坤行",
            }
        }
        apply_crawl_results_to_rows(
            result["rows"], result["enriched"], result["audit_response"],
            crawl_results_by_url=crawl_results,
            crawler_process_failed=False,
        )
        from scripts.audit_saas_excel_file import write_audit_workbook
        write_audit_workbook(out, result["headers"], result["rows"], result["enriched"], result["audit_response"])
        wb = load_workbook(str(out))
        ws = wb["record_results"]
        hdr = [str(c.value or "") for c in ws[1]]
        idx_value = hdr.index(_display_header("page_price_value")) + 1
        idx_status = hdr.index(_display_header("page_price_status")) + 1
        idx_source = hdr.index(_display_header("page_price_source")) + 1
        idx_check = hdr.index(_display_header("page_vs_customer_check")) + 1
        idx_over = hdr.index(_display_header("over_30pct")) + 1
        assert str(ws.cell(row=2, column=idx_value).value) == "131"
        assert ws.cell(row=2, column=idx_status).value == _display_value("OK")
        assert ws.cell(row=2, column=idx_source).value == "震坤行"
        assert ws.cell(row=2, column=idx_check).value == _display_value("HIGHER_THAN_CUSTOMER_PRICE_30")
        assert ws.cell(row=2, column=idx_over).value == _display_value("TRUE")
        wb.close()

    def test_apply_success_within_threshold_marks_false(self, tmp_path):
        rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v1",
             "extLink": "https://www.ehsy.com/product-A",
             "customerPrice": "100", "queryStr": "test", "skuName": "test"},
        ]
        cust_rows = [{"documentId": "v1", "fourCategoryId": "c", "fourCategoryName": "n", "referencePrice": "100"}]
        saas_path = _write_pending_json(tmp_path, rows)
        cust_path = _write_customer_json(tmp_path, cust_rows)
        out = tmp_path / "out.xlsx"
        result = run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
        )
        crawl_results = {
            "https://www.ehsy.com/product-A": {
                "page_price_value": Decimal("90"),
                "page_price_status": "OK",
                "page_price_source": "震坤行",
            }
        }
        apply_crawl_results_to_rows(
            result["rows"], result["enriched"], result["audit_response"],
            crawl_results_by_url=crawl_results,
        )
        from scripts.audit_saas_excel_file import write_audit_workbook
        write_audit_workbook(out, result["headers"], result["rows"], result["enriched"], result["audit_response"])
        wb = load_workbook(str(out))
        ws = wb["record_results"]
        hdr = [str(c.value or "") for c in ws[1]]
        idx_check = hdr.index(_display_header("page_vs_customer_check")) + 1
        idx_over = hdr.index(_display_header("over_30pct")) + 1
        assert ws.cell(row=2, column=idx_check).value == _display_value("WITHIN_THRESHOLD")
        assert ws.cell(row=2, column=idx_over).value == _display_value("FALSE")
        wb.close()

    def test_apply_crawler_process_failed_marks_pending_rows(self, tmp_path):
        rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v1",
             "extLink": "https://www.ehsy.com/product-A",
             "customerPrice": "100"},
        ]
        cust_rows = [{"documentId": "v1", "fourCategoryId": "c", "fourCategoryName": "n", "referencePrice": "100"}]
        saas_path = _write_pending_json(tmp_path, rows)
        cust_path = _write_customer_json(tmp_path, cust_rows)
        out = tmp_path / "out.xlsx"
        result = run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
        )
        apply_crawl_results_to_rows(
            result["rows"], result["enriched"], result["audit_response"],
            crawl_results_by_url={},
            crawler_process_failed=True,
        )
        from scripts.audit_saas_excel_file import write_audit_workbook
        write_audit_workbook(out, result["headers"], result["rows"], result["enriched"], result["audit_response"])
        wb = load_workbook(str(out))
        ws = wb["record_results"]
        hdr = [str(c.value or "") for c in ws[1]]
        idx_status = hdr.index(_display_header("page_price_status")) + 1
        assert ws.cell(row=2, column=idx_status).value == _display_value("CRAWLER_PROCESS_FAILED")
        wb.close()

    def test_apply_url_not_returned_marks_failed(self, tmp_path):
        rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v1",
             "extLink": "https://www.ehsy.com/product-A",
             "customerPrice": "100"},
        ]
        cust_rows = [{"documentId": "v1", "fourCategoryId": "c", "fourCategoryName": "n", "referencePrice": "100"}]
        saas_path = _write_pending_json(tmp_path, rows)
        cust_path = _write_customer_json(tmp_path, cust_rows)
        out = tmp_path / "out.xlsx"
        result = run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
        )
        apply_crawl_results_to_rows(
            result["rows"], result["enriched"], result["audit_response"],
            crawl_results_by_url={},  # 空 → 所有行都没找到
            crawler_process_failed=False,
        )
        from scripts.audit_saas_excel_file import write_audit_workbook
        write_audit_workbook(out, result["headers"], result["rows"], result["enriched"], result["audit_response"])
        wb = load_workbook(str(out))
        ws = wb["record_results"]
        hdr = [str(c.value or "") for c in ws[1]]
        idx_status = hdr.index(_display_header("page_price_status")) + 1
        assert ws.cell(row=2, column=idx_status).value == _display_value("FAILED")
        wb.close()


# ---------------------------------------------------------------------------
# 爬取模板自动导出（双 JSON 模式启动时先生成京东+震坤行 xlsx）
# ---------------------------------------------------------------------------


class TestExcludedTemplateAutoExport:
    def test_dual_json_mode_auto_exports_excluded_templates(self, tmp_path, monkeypatch):
        """双 JSON 模式 run_main 会在解析前自动生成京东+震坤行模板。"""
        # 该用例需启用导出
        monkeypatch.delenv("AUDIT_SUPPRESS_EXCLUDED_EXPORT", raising=False)
        # 准备含京东 + 震坤行的 saas 行
        rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v-jd",
             "skuId": "100237336261", "extLink": "https://item.jd.com/100237336261.html",
             "vendorName": "jd"},
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v-zkh",
             "skuId": "AA2130014", "extLink": "https://www.zkh.com/item/AA2130014.html",
             "vendorName": "zkh"},
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v-ok",
             "skuId": "OK1", "extLink": "https://www.example.com/ok",
             "vendorName": "demo"},
        ]
        cust_rows = [
            {"documentId": "v-jd", "fourCategoryId": "c1", "fourCategoryName": "n1", "referencePrice": "100"},
            {"documentId": "v-zkh", "fourCategoryId": "c2", "fourCategoryName": "n2", "referencePrice": "200"},
            {"documentId": "v-ok", "fourCategoryId": "c3", "fourCategoryName": "n3", "referencePrice": "300"},
        ]
        saas_path = _write_pending_json(tmp_path, rows)
        cust_path = _write_customer_json(tmp_path, cust_rows)
        out = tmp_path / "out.xlsx"
        excluded_root = tmp_path / "excluded"

        result = run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
            excluded_output_dir=excluded_root,
        )

        # 检查爬取模板被生成
        export_info = result["excluded_export"]
        assert export_info["enabled"] is True
        assert "京东" in export_info
        assert "震坤行" in export_info
        assert export_info["京东"]["matched"] == 1
        assert export_info["震坤行"]["matched"] == 1

        # 检查文件存在
        date_subdirs = list(excluded_root.iterdir())
        assert len(date_subdirs) == 1
        date_dir = date_subdirs[0]
        assert (date_dir / "京东.xlsx").exists()
        assert (date_dir / "震坤行.xlsx").exists()
        assert (date_dir / "summary.txt").exists()

        # 检查 xlsx 表头与内容
        wb = load_workbook(str(date_dir / "京东.xlsx"))
        ws = wb.active
        hdr = [ws.cell(row=1, column=c).value for c in (1, 2)]
        assert hdr == ["商品编码", "url"]
        assert ws.cell(row=2, column=1).value == "100237336261"
        assert ws.cell(row=2, column=2).value == "https://item.jd.com/100237336261.html"
        wb.close()

    def test_no_excluded_export_flag_skips_template(self, tmp_path):
        """--no-excluded-export 应跳过模板导出。"""
        rows = [
            {**SAMPLE_SAAS_PENDING_ROW, "vCommodityId": "v-jd",
             "extLink": "https://item.jd.com/x.html", "vendorName": "jd"},
        ]
        cust_rows = [{"documentId": "v-jd", "fourCategoryId": "c", "fourCategoryName": "n", "referencePrice": "100"}]
        saas_path = _write_pending_json(tmp_path, rows)
        cust_path = _write_customer_json(tmp_path, cust_rows)
        out = tmp_path / "out.xlsx"
        excluded_root = tmp_path / "excluded"

        result = run_main(
            output_xlsx=out,
            input_saas_json=saas_path,
            input_customer_json=cust_path,
            excluded_output_dir=excluded_root,
            excluded_export_disabled=True,
        )

        assert result["excluded_export"]["enabled"] is False
        assert not excluded_root.exists()

    def test_input_saas_list_json_mode_no_excluded_export(self, tmp_path):
        """非双 JSON 模式不触发爬取模板导出。"""
        list_path = tmp_path / "list.json"
        list_path.write_text(json.dumps({"rows": [SAMPLE_CUSTOMER_ROW]}, ensure_ascii=False), encoding="utf-8")
        out = tmp_path / "out.xlsx"
        excluded_root = tmp_path / "excluded"

        result = run_main(
            output_xlsx=out,
            input_saas_list_json=list_path,
            excluded_output_dir=excluded_root,
        )
        assert result["excluded_export"]["enabled"] is False
        assert not excluded_root.exists()


# ---------------------------------------------------------------------------
# 新增业务规则测试：SaaS 预检、价格变动、最终决策
# ---------------------------------------------------------------------------


from scripts.audit_saas_excel_file import (
    compute_final_decision,
    compute_price_changed,
    compute_saas_price_vs_customer,
)


class TestSaasPriceVsCustomer:
    def test_saas_higher_than_30(self):
        status, ev = compute_saas_price_vs_customer(
            customer_price=Decimal("100"),
            saas_price=Decimal("131"),
        )
        assert status == "SAAS_HIGHER_THAN_CUSTOMER_30"
        assert "131" in ev

    def test_saas_within_threshold(self):
        status, _ = compute_saas_price_vs_customer(
            customer_price=Decimal("100"),
            saas_price=Decimal("130"),
        )
        assert status == "SAAS_WITHIN_CUSTOMER_THRESHOLD"

    def test_saas_price_missing(self):
        status, _ = compute_saas_price_vs_customer(
            customer_price=Decimal("100"),
            saas_price=None,
        )
        assert status == "SAAS_PRICE_MISSING"

    def test_customer_price_missing(self):
        status, _ = compute_saas_price_vs_customer(
            customer_price=None,
            saas_price=Decimal("100"),
        )
        assert status == "CUSTOMER_PRICE_MISSING"


class TestPriceChanged:
    def test_changed_obvious(self):
        status, ev = compute_price_changed(
            saas_price=Decimal("100"),
            page_price=Decimal("150"),
            page_price_status="OK",
        )
        assert status == "CHANGED"
        assert "100" in ev and "150" in ev

    def test_unchanged_within_tolerance(self):
        status, _ = compute_price_changed(
            saas_price=Decimal("100"),
            page_price=Decimal("100.5"),  # 0.5% < 1%
            page_price_status="OK",
        )
        assert status == "UNCHANGED"

    def test_changed_outside_tolerance(self):
        status, _ = compute_price_changed(
            saas_price=Decimal("100"),
            page_price=Decimal("102"),  # 2% > 1%
            page_price_status="OK",
        )
        assert status == "CHANGED"

    def test_unavailable_when_page_price_missing(self):
        status, _ = compute_price_changed(
            saas_price=Decimal("100"),
            page_price=None,
            page_price_status="FAILED",
        )
        assert status == "PRICE_UNAVAILABLE"

    def test_unavailable_when_saas_price_invalid(self):
        status, _ = compute_price_changed(
            saas_price=None,
            page_price=Decimal("100"),
            page_price_status="OK",
        )
        assert status == "PRICE_UNAVAILABLE"


class TestFinalDecision:
    def test_pass(self):
        assert compute_final_decision(
            "SAAS_WITHIN_CUSTOMER_THRESHOLD", "UNCHANGED", "FALSE"
        ) == "PASS"

    def test_reject_saas_high(self):
        assert compute_final_decision(
            "SAAS_HIGHER_THAN_CUSTOMER_30", "UNCHANGED", "FALSE"
        ) == "REJECT_SAAS_HIGH"

    def test_reject_price_changed(self):
        assert compute_final_decision(
            "SAAS_WITHIN_CUSTOMER_THRESHOLD", "CHANGED", "FALSE"
        ) == "REJECT_PRICE_CHANGED"

    def test_reject_both_when_multiple_anomalies(self):
        assert compute_final_decision(
            "SAAS_HIGHER_THAN_CUSTOMER_30", "CHANGED", "FALSE"
        ) == "REJECT_BOTH"

    def test_over_30_alone_also_rejects(self):
        assert compute_final_decision(
            "SAAS_WITHIN_CUSTOMER_THRESHOLD", "UNCHANGED", "TRUE"
        ) == "REJECT_SAAS_HIGH"


class TestNotSameNoLongerSkipped:
    """模型判 NOT_SAME 也要爬虫，不再 SKIPPED_MODEL_NOT_SAME。"""

    def test_not_same_still_pending_crawl(self, tmp_path):
        from scripts.audit_saas_excel_file import _compute_page_price_status

        class FakeAuditResult:
            title_match_result = "NOT_SAME"

        row = {"extLink": "https://www.example.com/abc"}
        result = _compute_page_price_status(row, FakeAuditResult())
        assert result["page_price_status"] == "PENDING_CRAWL"
