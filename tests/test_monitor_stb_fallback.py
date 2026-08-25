from types import SimpleNamespace
import subprocess
import sys

import openpyxl
import pytest

import scripts.monitor_stb_fallback as monitor_stb_fallback
from scripts.monitor_stb_fallback import (
    apply_detail_price_results,
    apply_detail_reference_results,
    apply_stbchina_results,
    extract_detail_reference_records,
    find_header,
    is_stbchina_url,
    load_zhaoshang_db_env,
)


def test_extract_detail_reference_records_pairs_each_url_with_its_price_and_detail_title():
    records = extract_detail_reference_records(
        "doc-1",
        {
            "commodityName": "detail 商品标题",
            "referenceInfos": [
                {"referenceLink": "https://example.com/a?id=1", "referencePrice": "12.50"},
                {"referenceLink": "https://example.com/b?id=2", "referencePrice": "0"},
            ],
        },
    )

    assert records[("doc-1", "https://example.com/a?id=1")] == {
        "price": 12.5,
        "title": "detail 商品标题",
    }
    assert records[("doc-1", "https://example.com/b?id=2")]["price"] == 0.0


def test_extract_detail_reference_records_keeps_url_price_when_detail_has_no_title():
    records = extract_detail_reference_records(
        "doc-1",
        {"referenceInfos": [{"referenceLink": "https://example.com/a?id=1", "referencePrice": "12.50"}]},
    )

    assert records[("doc-1", "https://example.com/a?id=1")] == {"price": 12.5, "title": ""}


def test_apply_detail_price_results_fills_price_without_setting_internet_title():
    row = SimpleNamespace(
        单据id="doc-1",
        URL="https://example.com/a?id=1",
        价格="",
        互联网标题="",
        爬取状态="待爬取",
        爬取失败原因="无价格",
    )

    changed = apply_detail_price_results(
        [row],
        {("doc-1", row.URL): {"price": 12.5, "title": "中核商品名称"}},
    )

    assert changed == 1
    assert row.价格 == 12.5
    assert row.互联网标题 == ""
    assert row.爬取状态 == "中核detail参考价回填"
    assert row.爬取失败原因 == ""


def test_apply_detail_reference_results_only_fills_rows_missing_both_fields():
    empty = SimpleNamespace(
        单据id="doc-1",
        URL="https://example.com/a?id=1",
        价格="",
        互联网标题="",
        爬取状态="待爬取",
        爬取失败原因="无价格",
    )
    existing_price = SimpleNamespace(
        单据id="doc-1",
        URL="https://example.com/b?id=2",
        价格=9.9,
        互联网标题="",
        爬取状态="成功",
        爬取失败原因="",
    )
    existing_title = SimpleNamespace(
        单据id="doc-1",
        URL="https://example.com/c?id=3",
        价格="",
        互联网标题="已有标题",
        爬取状态="待爬取",
        爬取失败原因="",
    )
    records = {
        ("doc-1", "https://example.com/a?id=1"): {"price": 12.5, "title": "detail 标题 A"},
        ("doc-1", "https://example.com/b?id=2"): {"price": 13.5, "title": "detail 标题 B"},
        ("doc-1", "https://example.com/c?id=3"): {"price": 14.5, "title": "detail 标题 C"},
    }

    changed = apply_detail_reference_results([empty, existing_price, existing_title], records)

    assert changed == 1
    assert empty.价格 == 12.5
    assert empty.互联网标题 == "detail 标题 A"
    assert empty.爬取状态 == "中核detail参考回填"
    assert empty.爬取失败原因 == ""
    assert existing_price.价格 == 9.9
    assert existing_price.互联网标题 == ""
    assert existing_title.价格 == ""
    assert existing_title.互联网标题 == "已有标题"


def test_detail_reference_results_require_exact_document_and_url_match():
    row = SimpleNamespace(
        单据id="doc-1",
        URL="https://example.com/a?id=1",
        价格="",
        互联网标题="",
        爬取状态="待爬取",
        爬取失败原因="",
    )
    records = {
        ("doc-2", row.URL): {"price": 12.5, "title": "wrong document"},
        ("doc-1", "https://example.com/a?id=2"): {"price": 13.5, "title": "wrong URL"},
    }

    assert apply_detail_reference_results([row], records) == 0
    assert row.价格 == ""
    assert row.互联网标题 == ""


def test_load_zhaoshang_db_env_requires_all_variables(tmp_path):
    env_file = tmp_path / "db.env"
    env_file.write_text("ZHAOSHANG_DB_HOST=db.example\n", encoding="utf-8")

    with pytest.raises(RuntimeError, match="ZHAOSHANG_DB_PASSWORD"):
        load_zhaoshang_db_env(env_file)


def test_load_zhaoshang_db_env_removes_outer_quotes(tmp_path):
    env_file = tmp_path / "db.env"
    env_file.write_text("\n".join([
        "ZHAOSHANG_DB_HOST='db.example'",
        "ZHAOSHANG_DB_PORT='3306'",
        "ZHAOSHANG_DB_USER='user'",
        "ZHAOSHANG_DB_PASSWORD='password'",
    ]), encoding="utf-8")

    assert load_zhaoshang_db_env(env_file)["ZHAOSHANG_DB_PORT"] == "3306"


def test_is_stbchina_url_matches_only_stbchina_host():
    assert is_stbchina_url("https://www.stbchina.cn/product?itemId=1")
    assert not is_stbchina_url("https://example.com/?site=stbchina.cn")


def test_apply_stbchina_results_updates_only_rows_with_a_price():
    row = SimpleNamespace(
        URL="https://www.stbchina.cn/product?itemId=1",
        价格="",
        互联网标题="",
        SKUID="",
        销量_评论数="",
    )

    changed = apply_stbchina_results(
        [row],
        {row.URL: {"price": 12.5, "title": "商城标题", "uid": "1", "comment": "2"}},
    )

    assert changed == 1
    assert row.价格 == 12.5
    assert row.互联网标题 == "商城标题"
    assert row.SKUID == "1"


def test_apply_stbchina_results_can_leave_internet_title_blank():
    row = SimpleNamespace(
        URL="https://www.stbchina.cn/product?itemId=1",
        价格="",
        互联网标题="",
        SKUID="",
        销量_评论数="",
    )

    changed = apply_stbchina_results(
        [row],
        {row.URL: {"price": 12.5, "title": "史泰博标题", "uid": "1", "comment": "2"}},
        fill_title=False,
    )

    assert changed == 1
    assert row.价格 == 12.5
    assert row.互联网标题 == ""


def test_fixed_entrypoint_exposes_opt_in_stb_and_db_env_options():
    result = subprocess.run(
        [sys.executable, "scripts/process_monitor_no_push_zhaoshang.py", "--help"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )

    assert result.returncode == 0
    assert "--zhaoshang-db-env" in result.stdout
    assert "--zs-wait-concurrency" in result.stdout
    assert "--auto-stbchina-fallback" in result.stdout
    assert "--detail-reference-fallback" in result.stdout
    assert "--leave-internet-title-blank" in result.stdout


def test_stb_fallback_reuses_complete_existing_result(tmp_path, monkeypatch):
    url = "https://www.stbchina.cn/product?itemId=1"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["商品链接", "商品价格", "商品标题", "采集状态"])
    sheet.append([url, 12.5, "史泰博标题", "成功"])
    workbook.save(tmp_path / "stbchina_missing_result.xlsx")

    def fail_if_called(*_args, **_kwargs):
        raise AssertionError("完整的既有史泰博结果不应重新运行脚本")

    monkeypatch.setattr(monitor_stb_fallback.subprocess, "run", fail_if_called)

    records = monitor_stb_fallback.run_stbchina_fallback([url], tmp_path)

    assert records[url]["price"] == 12.5


def test_find_header_accepts_current_monitor_workbook_aliases():
    headers = ["文档ID", "中核商品名称", "物料编码"]

    assert find_header(headers, "商品名称", "中核商品名称") == "中核商品名称"
    assert find_header(headers, "准标准商品编码", "物料编码") == "物料编码"


def test_completed_unconfirmed_priority_is_not_a_hard_failure():
    classify = getattr(monitor_stb_fallback, "classify_priority_output", lambda _text: "missing")
    state = classify("[优先级调整] COMPLETED_UNCONFIRMED: 任务已完成、未确认提优。")

    assert state == "completed_unconfirmed"


def test_already_highest_priority_is_an_idempotent_success():
    state = monitor_stb_fallback.classify_priority_output(
        "[优先级调整] ALREADY_HIGHEST: 已按当前任务主键范围验证。"
    )

    assert state == "already_highest_priority"


def test_active_task_without_priority_evidence_is_a_hard_failure():
    state = monitor_stb_fallback.classify_priority_output(
        "[优先级调整] ACTIVE_UNCONFIRMED: 活动任务未取得提优证据。"
    )

    assert state == "failed"


def test_priority_warning_is_a_hard_failure():
    classify = getattr(monitor_stb_fallback, "classify_priority_output", lambda _text: "missing")
    state = classify("[优先级调整] WARNING: 未取得 affected_rows>0")

    assert state == "failed"


def test_priority_output_without_positive_evidence_is_not_treated_as_updated():
    assert monitor_stb_fallback.classify_priority_output(
        "已有任务，跳过重复启用和优先级调整"
    ) == "missing"


def test_priority_output_accepts_only_positive_updated_rows_evidence():
    assert monitor_stb_fallback.classify_priority_output(
        "【优先级调整】updated rows=28"
    ) == "updated"


def test_locked_output_is_reused_when_existing_workbook_row_count_matches(tmp_path):
    path = tmp_path / "locked.xlsx"
    existing = openpyxl.Workbook()
    existing.active.append(["header"])
    existing.active.append(["row1"])
    existing.active.append(["row2"])
    existing.save(path)

    class LockedWorkbook:
        def save(self, _path):
            raise PermissionError("locked")

    save_workbook = getattr(monitor_stb_fallback, "save_or_reuse_locked_workbook", lambda *_args, **_kwargs: "missing")
    assert save_workbook(LockedWorkbook(), path, expected_rows=3) == "reused"


def test_locked_output_is_not_reused_when_existing_workbook_row_count_differs(tmp_path):
    path = tmp_path / "locked.xlsx"
    existing = openpyxl.Workbook()
    existing.active.append(["header"])
    existing.save(path)

    class LockedWorkbook:
        def save(self, _path):
            raise PermissionError("locked")

    save_workbook = getattr(monitor_stb_fallback, "save_or_reuse_locked_workbook", lambda *_args, **_kwargs: "missing")
    with pytest.raises(PermissionError):
        save_workbook(LockedWorkbook(), path, expected_rows=3)
