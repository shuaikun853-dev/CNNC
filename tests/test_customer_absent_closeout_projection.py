"""Final-projection tests: customer-absent rows must not just skip finalize —
the frozen artifact's audit_response and the published Excel must carry the
NO_ACTION decision too (reviewer gap: before this projection the workbook
still showed MANUAL_CHECK / 客户价缺失 for CUSTOMER_ABSENT_SKIPPED rows).

Fixture pattern is reused from tests/test_incremental_price_audit_force_finalize.py
(_seed_force / _rework_detail / _proven_absent_result).
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.schemas.workflow import WorkflowAction
from app.services.daily_audit_store import DailyAuditStore, ProcessLock
from app.workflows.incremental_price_audit_finalize import (
    build_frozen_audit_artifact,
    materialize_frozen_daily_run,
)
from scripts.audit_saas_excel_file import RESULT_APPENDED_COLUMNS


DATE = "2026-08-24"


def _row(v="v-1", sku="sku-1", url="https://shop.example/item", **extra):
    return {
        "vCommodityId": v,
        "documentId": v,
        "skuId": sku,
        "extLink": url,
        "queryStr": "客户标题",
        "sassVCommodityName": "客户标题",
        "skuName": "互联网标题",
        "vendorName": "商城",
        "shopName": "店铺",
        "price": "10",
        "fourCategoryId": "cat-1",
        "fourCategoryName": "分类",
        "list_reference_price": "10",
        "customer_lookup_status": "MATCHED",
        "__crawl_required": bool(url),
        **extra,
    }


def _seed_force(tmp_path: Path, *, rows=None, freeze=True):
    """Frozen run with completed Detail/Model and crawl work left PENDING."""
    rows = list(rows or [_row(), _row("v-2", "sku-2", "https://shop.example/other")])
    runtime_root = tmp_path / "runtime"
    run_dir = runtime_root / DATE
    run_dir.mkdir(parents=True, exist_ok=True)
    store = DailyAuditStore(run_dir / "state.sqlite3")
    store.ensure_run(DATE, cutoff_at=f"{DATE}T10:00:00+00:00")
    token = store.acquire_lease(DATE, "seed", 60)
    # Mirror the real source export: audit-appended columns are never part of
    # the frozen header matrix the final writer preserves positionally.
    source_headers = [
        name for name in rows[0]
        if not name.startswith("__") and name not in RESULT_APPENDED_COLUMNS
    ]
    snapshot = store.ingest_snapshot(
        token,
        rows,
        f"{DATE}T10:01:00+00:00",
        headers=source_headers,
        evidence={"observed_count": len(rows)},
    )["snapshot"]
    if freeze:
        store.freeze_snapshot(token, snapshot["id"])
    for item in {item["id"]: item for item in store.ordered_snapshot_items(snapshot["id"])}.values():
        detail = store.ensure_stage_work(token, item["id"], "DETAIL", f"detail-{item['id']}")
        store.claim_stage_work(token, detail["id"], max_attempts=3, max_elapsed_seconds=300)
        store.complete_stage_work(token, detail["id"], {
            "customerPrice": "10",
            "customerPriceStatus": "FOUND",
            "customerPriceSource": "DETAIL_REFERENCE_INFOS",
            "customerPriceWarning": "",
            "customerPriceEvidence": "customer=10",
        })
        model = store.ensure_stage_work(token, item["id"], "MODEL", f"model-{item['id']}")
        store.claim_stage_work(token, model["id"], max_attempts=1, max_elapsed_seconds=300)
        store.complete_stage_work(token, model["id"], {
            "match_result": "SAME", "confidence": 0.95, "model_match": "exact", "reason": "一致",
        })
    crawls = []
    for row in rows:
        if row.get("__crawl_required"):
            crawls.append(store.ensure_crawl_work(
                token, "mall", row["skuId"], row["extLink"], "legacy-v1",
                f"crawl-fp-{row['skuId']}",
                source_id=row["vCommodityId"],
            ))
    store.release_lease(token)
    return store, snapshot, crawls


def _rework_detail(store, item_id, *, result=None, error=None):
    """Replace one DETAIL row's durable evidence with a new outcome."""
    token = store.acquire_lease(DATE, "rework-detail", 60)
    detail = next(
        row for row in store.list_stage_work(token, stage="DETAIL")
        if row["audit_item_id"] == item_id
    )
    detail = store.ensure_stage_work(
        token, item_id, "DETAIL", str(detail["fingerprint"]) + "-rework"
    )
    claimed = store.claim_stage_work(
        token, detail["id"], max_attempts=1, max_elapsed_seconds=300
    )
    assert claimed is not None
    if error is not None:
        store.fail_stage_work(
            token, detail["id"], elapsed_seconds=1, error=error, retryable=True,
            max_attempts=1, max_elapsed_seconds=1,
        )
    else:
        store.complete_stage_work(token, detail["id"], result)
    store.release_lease(token)


def _proven_absent_result(status, *, evidence=""):
    return {
        "customerPrice": None,
        "customerPriceStatus": status,
        "customerPriceSource": "",
        "customerPriceWarning": "authoritative customer export has no such document",
        "customerPriceEvidence": evidence,
    }


def _absent_rows(count):
    return [
        _row(f"v-absent-{i}", f"sku-absent-{i}", "")
        for i in range(count)
    ]


def _rework_items_to_absent(store, snapshot, status, *, evidence=""):
    for item in store.ordered_snapshot_items(snapshot["id"]):
        if str(item["row"]["vCommodityId"]).startswith("v-absent-"):
            _rework_detail(
                store, item["id"], result=_proven_absent_result(status, evidence=evidence),
            )


# ---- T1: single absent row projects to NO_ACTION in the frozen artifact ----


def test_single_customer_absent_row_projects_no_action_in_artifact(tmp_path):
    """T1: one frozen row whose authoritative customer evidence is
    CUSTOMER_NOT_FOUND keeps finalize ready, and the artifact's audit_response
    records the final NO_ACTION projection — never MANUAL_CHECK / 客户价缺失."""
    store, snapshot, _crawls = _seed_force(
        tmp_path, rows=_absent_rows(1),
    )
    _rework_items_to_absent(
        store, snapshot, "CUSTOMER_NOT_FOUND",
        evidence="authoritative customer snapshot round complete; documentId absent",
    )
    token = store.acquire_lease(DATE, "t1", 60)

    built = build_frozen_audit_artifact(
        store, token, output_identity=f"daily-final-{DATE.replace('-', '')}",
    )

    assert built.readiness.ready is True
    assert built.readiness.customer_absent_skipped == 1
    assert built.artifact is not None
    result = built.artifact.audit_response.record_results[0]
    assert result.price_vs_customer_check == "CUSTOMER_ABSENT_SKIPPED"
    assert result.primary_action == WorkflowAction.NO_ACTION
    assert result.suggested_actions == [WorkflowAction.NO_ACTION]
    assert result.needs_human_approval is False
    assert result.execution_allowed is False
    assert result.executed is False
    assert result.reason == "客户侧权威快照当前无此单据，本轮无需处理"
    # The row's real evidence is appended, not invented.
    assert "customerPriceStatus=CUSTOMER_NOT_FOUND" in result.evidence
    assert any("authoritative customer export has no such document" in item for item in result.evidence)
    assert any(
        "authoritative customer snapshot round complete; documentId absent" in item
        for item in result.evidence
    )
    # The pre-projection verdict must not leak anywhere on the row.
    assert "CUSTOMER_PRICE_MISSING" != result.price_vs_customer_check
    assert "缺少有效客户价" not in result.reason
    store.release_lease(token)


# ---- T2: legacy LOOKUP_FAILED evidence projects identically, state untouched --


def test_legacy_lookup_failed_row_projects_no_action_without_rewriting_state(tmp_path):
    """T2: durable legacy LOOKUP_FAILED evidence gets the same final NO_ACTION
    projection, and the projection never rewrites the persisted state evidence."""
    store, snapshot, _crawls = _seed_force(tmp_path, rows=_absent_rows(1))
    _rework_items_to_absent(store, snapshot, "LOOKUP_FAILED")
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    token = store.acquire_lease(DATE, "t2", 60)

    built = build_frozen_audit_artifact(
        store, token, output_identity=f"daily-final-{DATE.replace('-', '')}",
    )

    assert built.artifact is not None
    result = built.artifact.audit_response.record_results[0]
    assert result.price_vs_customer_check == "CUSTOMER_ABSENT_SKIPPED"
    assert result.primary_action == WorkflowAction.NO_ACTION
    assert result.needs_human_approval is False
    assert "customerPriceStatus=LOOKUP_FAILED" in result.evidence

    # The durable state keeps its real historical evidence.
    reworked = next(
        row for row in store.list_stage_work(token, stage="DETAIL")
        if row["audit_item_id"] == item["id"]
        and str(row["fingerprint"]).endswith("-rework")
    )
    persisted = json.loads(reworked["result_json"])
    assert persisted["customerPriceStatus"] == "LOOKUP_FAILED"
    store.release_lease(token)


# ---- T3: mixed 100 ready + 3 absent day: counts corrected, others untouched --


def test_mixed_day_summary_counts_exclude_customer_absent_rows(tmp_path):
    """T3: 100 normal rows + 3 customer-absent rows.  The projected response
    keeps 103 record_results, marks exactly the 3 absent rows NO_ACTION with
    needs_human_approval False, corrects manual_check_count and
    customer_price_missing_count, and adds customer_absent_skipped_count
    (global + per affected vCommodityId group)."""
    rows = [_row(f"v-{i}", f"sku-{i}", "") for i in range(100)] + _absent_rows(3)
    store, snapshot, _crawls = _seed_force(tmp_path, rows=rows)
    _rework_items_to_absent(store, snapshot, "CUSTOMER_NOT_FOUND")
    token = store.acquire_lease(DATE, "t3", 60)

    built = build_frozen_audit_artifact(
        store, token, output_identity=f"daily-final-{DATE.replace('-', '')}",
    )

    assert built.artifact is not None
    response = built.artifact.audit_response
    results = response.record_results
    assert len(results) == 103
    absent = [r for r in results if r.vCommodityId.startswith("v-absent-")]
    assert len(absent) == 3
    for result in absent:
        assert result.price_vs_customer_check == "CUSTOMER_ABSENT_SKIPPED"
        assert result.primary_action == WorkflowAction.NO_ACTION
        assert result.needs_human_approval is False
    normal = [r for r in results if not r.vCommodityId.startswith("v-absent-")]
    assert len(normal) == 100
    assert all(r.needs_human_approval is True for r in normal)

    summary = response.summary
    assert summary["customer_absent_skipped_count"] == 3
    assert summary["customer_price_missing_count"] == 0
    # Recount independently: no absent row may be counted as manual check.
    manual_results = [r for r in results if r.primary_action == WorkflowAction.MANUAL_CHECK]
    assert all(not r.vCommodityId.startswith("v-absent-") for r in manual_results)
    assert summary["manual_check_count"] == len(manual_results) == 100

    for i in range(3):
        v = f"v-absent-{i}"
        assert summary[f"group.{v}.customer_absent_skipped_count"] == 1
        assert summary[f"group.{v}.manual_check_count"] == 0
    # Unaffected groups keep their original keys intact.
    assert summary["total_records"] == 103
    store.release_lease(token)


# ---- T4: proven-absent (NOT_PROVIDED / DETAIL_REQUEST_SKIPPED) never projects --


@pytest.mark.parametrize("status", ["NOT_PROVIDED", "DETAIL_REQUEST_SKIPPED"])
def test_proven_absent_rows_are_not_projected(tmp_path, status):
    """T4: identity-missing / price-not-provided is not document absence —
    the projection must stay away from these rows (MANUAL_CHECK, needs human,
    blocking readiness), proving customerPrice=None is not globally whitelisted."""
    store, snapshot, _crawls = _seed_force(tmp_path, rows=_absent_rows(1))
    item = store.ordered_snapshot_items(snapshot["id"])[0]
    _rework_detail(store, item["id"], result=_proven_absent_result(status))
    token = store.acquire_lease(DATE, "t4", 60)

    built = build_frozen_audit_artifact(
        store, token, output_identity=f"daily-final-{DATE.replace('-', '')}",
        tolerate_pending=True,
    )

    assert built.readiness.ready is False
    assert built.readiness.customer_price_missing == 1
    assert built.artifact is not None
    result = built.artifact.audit_response.record_results[0]
    assert result.price_vs_customer_check == "CUSTOMER_PRICE_MISSING"
    assert result.primary_action == WorkflowAction.MANUAL_CHECK
    assert result.needs_human_approval is True
    # The skip reason must never be used for a proven-absent row.
    assert result.reason != "客户侧权威快照当前无此单据，本轮无需处理"
    summary = built.artifact.audit_response.summary
    assert summary["customer_absent_skipped_count"] == 0
    assert summary["customer_price_missing_count"] == 1
    store.release_lease(token)


# ---- Guard: the projection fails closed on misaligned inputs ----


def test_projection_guard_raises_on_length_mismatch():
    """Fail-closed: record_results / price_states / enriched must align 1:1 or
    the projection refuses to guess which row is absent."""
    from app.services.customer_price_semantics import CustomerPriceBusinessState
    from app.workflows.incremental_price_audit_finalize import _project_customer_absent_results

    class _Response:
        record_results = [object(), object()]

    price_states = [CustomerPriceBusinessState("READY", "FOUND", "SUCCESS", False, "NONE")]
    with pytest.raises(ValueError):
        _project_customer_absent_results(_Response(), price_states, [{}])


# ---- T10: the published Excel carries the projected final decision ----


def test_published_workbook_shows_customer_absent_no_action(tmp_path):
    """T10: materialize the formal workbook for a mixed 1-absent + 1-normal day.
    The absent row shows 建议操作=无操作, 需人工审批=FALSE, the skip verdict in
    价格对比客户价, and its real customer evidence columns; the normal row is
    unchanged; the sheet has exactly one row per frozen row."""
    rows = _absent_rows(1) + [_row("v-ok", "sku-ok", "")]
    store, snapshot, _crawls = _seed_force(tmp_path, rows=rows)
    _rework_items_to_absent(
        store, snapshot, "CUSTOMER_NOT_FOUND",
        evidence="authoritative customer snapshot round complete; documentId absent",
    )
    output = tmp_path / "formal.xlsx"

    publication = materialize_frozen_daily_run(
        store, ProcessLock(tmp_path / "finalize.lock"), DATE, output,
    )

    assert publication.materialization is not None
    workbook = load_workbook(filename=str(publication.materialization.output_path), data_only=True)
    sheet = workbook["record_results"]
    headers = {
        str(cell.value): index
        for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1)), start=1)
    }
    data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    # One row per frozen row, nothing added or dropped.
    assert len(data_rows) == 2

    by_v = {row[headers["SaaS商品ID"] - 1]: row for row in data_rows}
    absent = by_v["v-absent-0"]
    normal = by_v["v-ok"]

    assert absent[headers["建议操作"] - 1] == "无操作"
    assert absent[headers["需人工审批"] - 1] == "FALSE"
    assert absent[headers["价格对比客户价"] - 1] == "客户侧无此单据，本轮无需处理"
    assert absent[headers["客户价状态"] - 1] == "客户侧查无此单据"
    assert absent[headers["客户价警告"] - 1]
    assert absent[headers["客户价证据"] - 1]

    # The normal row keeps its real (still human-gated) verdict.
    assert normal[headers["需人工审批"] - 1] == "TRUE"
    assert normal[headers["建议操作"] - 1] == "人工核查"
    workbook.close()
