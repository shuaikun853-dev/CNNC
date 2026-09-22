import ast
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "process_monitor_no_push_zhaoshang.py"


def test_batch_state_path_is_absolute_for_external_zhaoshang_cwd():
    tree = ast.parse(SCRIPT.read_text(encoding="utf-8"))
    assignment = next(
        node
        for node in tree.body
        if isinstance(node, ast.Assign)
        and any(isinstance(target, ast.Name) and target.id == "OUT_DIR" for target in node.targets)
    )

    assert isinstance(assignment.value, ast.Call)
    assert isinstance(assignment.value.func, ast.Attribute)
    assert assignment.value.func.attr == "resolve"


def test_cli_exposes_detail_price_preference_before_zhaoshang():
    result = subprocess.run(
        [sys.executable, str(SCRIPT), "--help"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )

    assert result.returncode == 0
    assert "--prefer-detail-price" in result.stdout
    assert "--shared-cache-path" in result.stdout


def test_step3_consumes_prewarm_cache_before_building_zhaoshang_queue():
    source = SCRIPT.read_text(encoding="utf-8")
    assert source.index("apply_cached_customer_results(") < source.index('print("[5] 招商局-only...")')


def test_step3_normalizes_all_comment_counts_before_import_filtering():
    source = SCRIPT.read_text(encoding="utf-8")

    normalize = source.index("normalize_comment_count(ir.销量_评论数)")
    required_filter = source.index("required_comment_missing = [")
    write_template = source.index("write_relaxed_template(out_template")

    assert normalize < required_filter < write_template


def load_read_zs_result():
    tree = ast.parse(SCRIPT.read_text(encoding="utf-8"))
    function = next(
        node
        for node in tree.body
        if isinstance(node, ast.FunctionDef) and node.name == "read_zs_result"
    )
    namespace = {
        "openpyxl": openpyxl,
        "config": SimpleNamespace(
            ZS_COL_CODE="\u5546\u54c1\u7f16\u7801",
            ZS_COL_TITLE="\u5546\u54c1\u6807\u9898/\u540d\u79f0",
            ZS_COL_PRICE="\u5546\u54c1\u4ef7\u683c",
            ZS_COL_URL="\u5546\u54c1\u8be6\u60c5\u5730\u5740",
            ZS_COL_COMMENT="\u8bc4\u8bba",
        ),
        "is_bad_title": lambda _title: False,
    }
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(SCRIPT), "exec"), namespace)
    return namespace["read_zs_result"]


def load_is_reusable_zs_result():
    tree = ast.parse(SCRIPT.read_text(encoding="utf-8"))
    function = next(
        node
        for node in tree.body
        if isinstance(node, ast.FunctionDef) and node.name == "is_reusable_zs_result"
    )
    namespace = {
        "openpyxl": openpyxl,
        "config": SimpleNamespace(
            ZS_COL_CODE="商品编码",
            ZS_COL_PRICE="商品价格",
            ZS_COL_URL="商品详情地址",
        ),
    }
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(SCRIPT), "exec"), namespace)
    return namespace["is_reusable_zs_result"]


class ReadZhaoshangResultTests(unittest.TestCase):
    def test_prefers_original_input_url_when_export_url_is_rewritten(self):
        read_zs_result = load_read_zs_result()
        original_url = (
            "https://www.qstbg.com/gd/detail.do?"
            "goods_id=21482219&formid=searchform"
        )
        exported_url = (
            "https://www.qstbg.com/gd/initParameter.do?"
            "goods_id=21482219&formid=form1#title_detail"
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "result.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(
                [
                    "\u5546\u54c1\u7f16\u7801",
                    "\u5546\u54c1\u6807\u9898/\u540d\u79f0",
                    "\u5546\u54c1\u4ef7\u683c",
                    "\u5546\u54c1\u8be6\u60c5\u5730\u5740",
                    "\u8bc4\u8bba",
                ]
            )
            sheet.append(["21482219", "title", 12.5, exported_url, 3])
            workbook.save(path)

            records = read_zs_result(path, {"21482219": original_url})

        self.assertIn(original_url, records)
        self.assertNotIn(exported_url, records)
        self.assertEqual(12.5, records[original_url]["price"])


    def test_prefers_page_sku_and_keeps_comment_and_channel(self):
        read_zs_result = load_read_zs_result()
        url = "https://item.jd.com/100311804305.html"
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "result.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append([
                "商品编码", "商品最小品类单元id", "商品标题/名称",
                "商品价格", "商品详情地址", "评论", "商城标识",
            ])
            sheet.append([
                "zhpq-placeholder", "100311804305", "亿联 MeetingEye 500",
                63420.5, url, 321, "jd",
            ])
            workbook.save(path)

            record = read_zs_result(path, {"zhpq-placeholder": url})[url]

        self.assertEqual("100311804305", record["code"])
        self.assertEqual("321", record["comment"])
        self.assertEqual("jd", record["channel"])

    def test_keeps_numeric_zero_comment_as_real_comment_evidence(self):
        read_zs_result = load_read_zs_result()
        url = "https://item.jd.com/100311804305.html"
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "result.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append([
                "商品编码", "商品最小品类单元id", "商品标题/名称",
                "商品价格", "商品详情地址", "评论", "商城标识",
            ])
            sheet.append([
                "zhpq-placeholder", "100311804305", "亿联 MeetingEye 500",
                63420.5, url, 0, "jd",
            ])
            workbook.save(path)

            record = read_zs_result(path, {"zhpq-placeholder": url})[url]

        self.assertEqual("0", record["comment"])


    def test_result_validation_uses_input_rows_when_product_codes_collide(self):
        is_reusable = load_is_reusable_zs_result()
        with tempfile.TemporaryDirectory() as tmp_dir:
            input_path = Path(tmp_dir) / "input.xlsx"
            result_path = Path(tmp_dir) / "result.xlsx"

            input_book = openpyxl.Workbook()
            input_sheet = input_book.active
            input_sheet.append(["商品编码", "url"])
            input_sheet.append(["SAME", "https://example.com/1"])
            input_sheet.append(["SAME", "https://example.com/2"])
            input_book.save(input_path)

            result_book = openpyxl.Workbook()
            result_sheet = result_book.active
            result_sheet.append(["商品编码", "商品价格", "商品详情地址"])
            result_sheet.append(["SAME", 1, "https://example.com/1"])
            result_sheet.append(["SAME", 2, "https://example.com/2"])
            result_book.save(result_path)

            job = SimpleNamespace(
                input_path=input_path,
                result_path=result_path,
                code_to_url={"SAME": "https://example.com/2"},
            )

            self.assertTrue(is_reusable(job))


if __name__ == "__main__":
    unittest.main()
