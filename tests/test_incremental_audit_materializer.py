import importlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.services.incremental_audit_materializer import (
    ArtifactConflictError,
    FrozenAuditArtifact,
    LegacyWorkbookMaterializer,
    MaterializationValidationError,
)


HEADERS = ("vCommodityId", "skuId", "extLink", "queryStr")


def _rows():
    return (
        {"vCommodityId": "v-1", "skuId": "sku-1", "extLink": " https://example/a ", "queryStr": "one"},
        {"vCommodityId": "v-1", "skuId": "sku-1", "extLink": " https://example/a ", "queryStr": "one"},
        {"vCommodityId": "v-2", "skuId": "sku-2", "extLink": "https://example/b", "queryStr": "two"},
    )


def _artifact(**changes):
    rows = changes.pop("rows", _rows())
    base = {
        "business_date": "2026-08-11",
        "tenant_id": "188605",
        "commodity_pool_id": "222",
        "output_identity": "daily-final-20260811",
        "frozen_snapshot_id": "snapshot-1",
        "frozen_snapshot_hash": "snapshot-hash-1",
        "frozen_snapshot_evidence": {"page": 3, "count": len(rows)},
        "headers": HEADERS,
        "frozen_rows": rows,
        "enriched_rows": tuple({"customerPrice": "10"} for _ in rows),
        "audit_response": {"record_results": tuple({"primary_action": "REVIEW"} for _ in rows)},
    }
    base.update(changes)
    return FrozenAuditArtifact(**base)


def _writer(path, headers, rows, enriched, audit_response):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "summary"
    summary.append(["total"])
    summary.append([len(rows)])
    sheet = workbook.create_sheet("record_results")
    sheet.append(list(headers))
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    workbook.save(path)
    workbook.close()


def _wrong_count_writer(path, headers, rows, enriched, audit_response):
    _writer(path, headers, rows[:-1], enriched[:-1], audit_response)


def _reordered_writer(path, headers, rows, enriched, audit_response):
    _writer(path, headers, tuple(reversed(rows)), enriched, audit_response)


def _source_corrupting_writer(path, headers, rows, enriched, audit_response):
    changed = [dict(row) for row in rows]
    changed[0]["queryStr"] = "wrong title"
    _writer(path, headers, changed, enriched, audit_response)


def test_successful_write_validates_and_publishes_sidecar(tmp_path):
    output = tmp_path / "final.xlsx"
    result = LegacyWorkbookMaterializer(writer=_writer).materialize(_artifact(), output)

    assert result.output_path == output
    assert result.manifest_path.exists()
    assert result.row_count == 3
    assert result.reused is False
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["input_manifest_sha256"] == result.input_manifest_sha256
    assert manifest["row_count"] == 3
    assert "rows" not in manifest


def test_same_input_reuses_byte_identical_output_without_second_writer_call(tmp_path):
    calls = []

    def writer(*args):
        calls.append(args[0])
        _writer(*args)

    materializer = LegacyWorkbookMaterializer(writer=writer)
    output = tmp_path / "final.xlsx"
    first = materializer.materialize(_artifact(), output)
    original = output.read_bytes()
    second = materializer.materialize(_artifact(), output)

    assert first.reused is False
    assert second.reused is True
    assert output.read_bytes() == original
    assert len(calls) == 1


def test_changed_input_conflicts_without_overwriting_published_output(tmp_path):
    output = tmp_path / "final.xlsx"
    materializer = LegacyWorkbookMaterializer(writer=_writer)
    materializer.materialize(_artifact(), output)
    original = output.read_bytes()

    with pytest.raises(ArtifactConflictError):
        materializer.materialize(_artifact(frozen_snapshot_hash="changed"), output)

    assert output.read_bytes() == original


@pytest.mark.parametrize("writer", [_wrong_count_writer, _reordered_writer, _source_corrupting_writer])
def test_invalid_writer_output_never_publishes(tmp_path, writer):
    output = tmp_path / "final.xlsx"

    with pytest.raises(MaterializationValidationError):
        LegacyWorkbookMaterializer(writer=writer).materialize(_artifact(), output)

    assert not output.exists()
    assert not list(tmp_path.glob(".*.tmp.xlsx"))


def test_headers_are_part_of_immutable_input_manifest(tmp_path):
    output = tmp_path / "final.xlsx"
    materializer = LegacyWorkbookMaterializer(writer=_writer)
    materializer.materialize(_artifact(), output)

    with pytest.raises(ArtifactConflictError):
        materializer.materialize(_artifact(headers=("vCommodityId", "skuId", "extLink")), output)


def test_reuse_rejects_changed_record_or_summary_projection(tmp_path):
    output = tmp_path / "final.xlsx"
    materializer = LegacyWorkbookMaterializer(writer=_writer)
    materializer.materialize(_artifact(), output)
    workbook = load_workbook(output)
    try:
        workbook["record_results"].cell(row=2, column=4).value = "tampered"
        workbook.save(output)
    finally:
        workbook.close()

    with pytest.raises(ArtifactConflictError):
        materializer.materialize(_artifact(), output)


def test_writer_failure_does_not_publish_or_leave_own_temp_file(tmp_path):
    output = tmp_path / "final.xlsx"

    def failing_writer(path, *_args):
        path.write_bytes(b"partial")
        raise RuntimeError("writer broke")

    with pytest.raises(RuntimeError, match="writer broke"):
        LegacyWorkbookMaterializer(writer=failing_writer).materialize(_artifact(), output)

    assert not output.exists()
    assert not list(tmp_path.glob(".*.tmp.xlsx"))


def test_crash_after_workbook_replace_recovers_sidecar_without_regenerating(tmp_path):
    output = tmp_path / "final.xlsx"
    artifact = _artifact()
    calls = []

    def writer(*args):
        calls.append(args[0])
        _writer(*args)

    def crash():
        raise RuntimeError("simulated crash")

    with pytest.raises(RuntimeError, match="simulated crash"):
        LegacyWorkbookMaterializer(writer=writer, after_workbook_replace=crash).materialize(artifact, output)

    published = output.read_bytes()
    recovered = LegacyWorkbookMaterializer(writer=writer).materialize(artifact, output)

    assert recovered.reused is True
    assert recovered.recovered_sidecar is True
    assert output.read_bytes() == published
    assert len(calls) == 1


def test_duplicate_source_rows_and_order_are_preserved_in_record_results(tmp_path):
    output = tmp_path / "final.xlsx"
    LegacyWorkbookMaterializer(writer=_writer).materialize(_artifact(), output)

    from openpyxl import load_workbook

    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        sheet = workbook["record_results"]
        values = list(sheet.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    assert [row[:3] for row in values] == [
        ("v-1", "sku-1", " https://example/a "),
        ("v-1", "sku-1", " https://example/a "),
        ("v-2", "sku-2", "https://example/b"),
    ]


@pytest.mark.parametrize(
    "changes",
    [
        {"tenant_id": "wrong"},
        {"commodity_pool_id": "wrong"},
        {"business_date": "not-a-date"},
        {"output_identity": ""},
        {"enriched_rows": ()},
        {"audit_response": SimpleNamespace(record_results=())},
    ],
)
def test_scope_date_identity_and_length_validation(changes, tmp_path):
    with pytest.raises(MaterializationValidationError):
        LegacyWorkbookMaterializer(writer=_writer).materialize(_artifact(**changes), tmp_path / "final.xlsx")


def test_module_import_has_no_production_side_effects(monkeypatch):
    import subprocess

    monkeypatch.setattr(subprocess, "run", lambda *_args, **_kwargs: pytest.fail("no subprocess"))
    module = importlib.import_module("app.services.incremental_audit_materializer")
    importlib.reload(module)
    assert hasattr(module, "LegacyWorkbookMaterializer")
