from __future__ import annotations

import json
from email.message import Message
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from scripts.reject_saas_audit_excel import (
    build_reject_payload,
    candidate_manifest_sha256,
    execute_reject_batches,
    fetch_pending_pairs,
    load_reject_candidates,
    reject_url_from_curl,
    run_main,
    select_pending_candidates,
    validate_candidates_are_pending,
    validate_execution_confirmation,
)


HEADERS = [
    "SaaS商品ID",
    "文档ID",
    "客户标题",
    "SaaS商品名称",
    "SKU编号",
    "互联网标题",
    "供应商",
    "店铺名称",
    "SaaS价格",
    "页面实际价格",
    "互联网链接",
    "评论数",
    "创建时间",
    "更新时间",
    "模型审核结果",
    "是否匹配",
]


def test_script_help_runs_from_repository_root() -> None:
    """Package-only imports would make the documented script entry point unusable."""
    result = subprocess.run(
        [sys.executable, "scripts/reject_saas_audit_excel.py", "--help"],
        cwd=Path(__file__).resolve().parents[1],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert "--input-xlsx" in result.stdout
    assert "--fill-colors" in result.stdout
    assert "--require-customer-price" in result.stdout
    assert "--require-customer-match" in result.stdout


def _write_audit_workbook(path: Path) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    summary.append(["总记录数"])
    summary.append([3])

    ws = wb.create_sheet("record_results")
    ws.append(HEADERS)
    rows = [
        ["101", None, None, None, "sku-red", "red title", None, None, None, None, None, None, None, None, "1", "1"],
        ["102", None, None, None, "sku-yellow", "yellow title", None, None, None, None, None, None, None, None, "1", "1"],
        ["103", None, None, None, "sku-clear", "clear title", None, None, None, None, None, None, None, None, "1", "1"],
    ]
    for row in rows:
        ws.append(row)
    ws["F2"].fill = PatternFill("solid", fgColor="FFFFC7CE")
    ws["F3"].fill = PatternFill("solid", fgColor="FFFFFF00")
    wb.save(path)


def test_load_reject_candidates_selects_only_red_and_yellow_f_cells(tmp_path: Path) -> None:
    """A missing fill check would incorrectly reject the uncoloured row."""
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)

    candidates = load_reject_candidates(workbook)

    assert [item.as_payload_item() for item in candidates] == [
        {"vCommodityId": "101", "skuId": "sku-red", "modelAuditResult": "1"},
        {"vCommodityId": "102", "skuId": "sku-yellow", "modelAuditResult": "1"},
    ]
    assert [item.color for item in candidates] == ["red", "yellow"]
    assert [item.excel_row for item in candidates] == [2, 3]


def test_load_reject_candidates_can_select_only_yellow_with_nonempty_customer_price(
    tmp_path: Path,
) -> None:
    """A broad colour filter or truthiness check could reject red rows or exclude price 0."""
    workbook = tmp_path / "audit.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    ws = wb.create_sheet("record_results")
    ws.append(
        HEADERS
        + ["四级类目ID", "四级类目名称", "客户数据匹配状态", "客户列表参考价", "客户价"]
    )
    rows = [
        ["201", None, None, None, "red-priced", "red", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, 50],
        ["202", None, None, None, "yellow-priced", "yellow", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, 88],
        ["203", None, None, None, "yellow-blank", "yellow", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, "   "],
        ["204", None, None, None, "yellow-zero", "yellow", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, 0],
        ["205", None, None, None, "clear-priced", "clear", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, 99],
    ]
    for row in rows:
        ws.append(row)
    ws["F2"].fill = PatternFill("solid", fgColor="FFFFC7CE")
    for row in (3, 4, 5):
        ws.cell(row, 6).fill = PatternFill("solid", fgColor="FFFFFF00")
    wb.save(workbook)

    candidates = load_reject_candidates(
        workbook,
        allowed_colors={"yellow"},
        require_customer_price=True,
    )

    assert [item.sku_id for item in candidates] == ["yellow-priced", "yellow-zero"]
    assert [item.color for item in candidates] == ["yellow", "yellow"]


def test_customer_price_filter_requires_customer_price_header(tmp_path: Path) -> None:
    """Falling back to a column position could silently read the wrong workbook field."""
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)

    with pytest.raises(ValueError, match="客户价"):
        load_reject_candidates(
            workbook,
            allowed_colors={"yellow"},
            require_customer_price=True,
        )


def test_load_reject_candidates_requires_matched_customer_rows(tmp_path: Path) -> None:
    """The S-column customer match filter must exclude unmatched rows."""
    workbook = tmp_path / "audit.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    ws = wb.create_sheet("record_results")
    ws.append(
        HEADERS
        + ["四级类目ID", "四级类目名称", "客户数据匹配状态", "客户列表参考价", "客户价"]
    )
    rows = [
        ["301", None, None, None, "red-matched", "red", None, None, None, None, None, None, None, None, "1", "1", None, None, "已匹配", None, 50],
        ["302", None, None, None, "yellow-unmatched", "yellow", None, None, None, None, None, None, None, None, "1", "1", None, None, "未匹配", None, 80],
        ["303", None, None, None, "yellow-matched-zero", "yellow", None, None, None, None, None, None, None, None, "1", "1", None, None, "已匹配", None, 0],
    ]
    for row in rows:
        ws.append(row)
    ws["F2"].fill = PatternFill("solid", fgColor="FFFFC7CE")
    ws["F3"].fill = PatternFill("solid", fgColor="FFFFFF00")
    ws["F4"].fill = PatternFill("solid", fgColor="FFFFFF00")
    wb.save(workbook)

    candidates = load_reject_candidates(
        workbook,
        require_customer_price=True,
        require_customer_match=True,
    )

    assert [item.sku_id for item in candidates] == ["red-matched", "yellow-matched-zero"]


def test_load_reject_candidates_coalesces_duplicate_business_keys(tmp_path: Path) -> None:
    """Duplicate Excel rows must remain auditable but only submit one business key."""
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)
    wb = __import__("openpyxl").load_workbook(workbook)
    ws = wb["record_results"]
    ws["A3"] = "101"
    ws["E3"] = "sku-red"
    wb.save(workbook)

    candidates = load_reject_candidates(workbook)

    assert len(candidates) == 1
    assert candidates[0].as_payload_item() == {
        "vCommodityId": "101",
        "skuId": "sku-red",
        "modelAuditResult": "1",
    }
    assert candidates[0].as_manifest_item()["sourceExcelRows"] == [2, 3]


def test_build_reject_payload_locks_cnnc_scope() -> None:
    """Changing the scope constants would send a valid reject list to another pool."""
    from scripts.reject_saas_audit_excel import RejectCandidate

    payload = build_reject_payload(
        [RejectCandidate("101", "sku-red", "1", "red", 2)]
    )

    assert payload == {
        "tenantId": "188605",
        "commodityPoolId": "222",
        "commodityList": [
            {"vCommodityId": "101", "skuId": "sku-red", "modelAuditResult": "1"}
        ],
    }


def test_reject_url_requires_saas_same_audit_curl() -> None:
    """Weak URL validation could reuse credentials against an unrelated endpoint."""
    parsed = {
        "url": "http://saasportal.iio.com/improxy/saas/iio2/sameAudit/queryAll",
        "cookie": "session=secret",
        "method": "POST",
    }

    assert reject_url_from_curl(parsed) == (
        "http://saasportal.iio.com/improxy/saas/iio2/sameAudit/reject"
    )
    with pytest.raises(ValueError, match="host"):
        reject_url_from_curl({**parsed, "url": "https://example.com/sameAudit/queryAll"})
    with pytest.raises(ValueError, match="Cookie"):
        reject_url_from_curl({**parsed, "cookie": ""})


def test_execution_confirmation_binds_count_and_manifest_hash(tmp_path: Path) -> None:
    """A stale confirmation must not authorize a changed candidate set."""
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)
    candidates = load_reject_candidates(workbook)
    digest = candidate_manifest_sha256(candidates)

    validate_execution_confirmation(candidates, expected_count=2, expected_sha256=digest)
    with pytest.raises(ValueError, match="count"):
        validate_execution_confirmation(candidates, expected_count=3, expected_sha256=digest)
    with pytest.raises(ValueError, match="SHA-256"):
        validate_execution_confirmation(candidates, expected_count=2, expected_sha256="0" * 64)


def test_manifest_hash_is_order_sensitive_and_uses_payload_fields() -> None:
    """The execution hash must bind the exact ordered records sent to SaaS."""
    from scripts.reject_saas_audit_excel import RejectCandidate

    first = RejectCandidate("101", "sku-a", "1", "red", 2)
    second = RejectCandidate("102", "sku-b", "1", "yellow", 3)
    digest = candidate_manifest_sha256([first, second])

    literal = json.dumps(
        [first.as_manifest_item(), second.as_manifest_item()],
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    import hashlib

    assert digest == hashlib.sha256(literal).hexdigest()
    assert digest != candidate_manifest_sha256([second, first])


class _FakeResponse:
    def __init__(self, payload: dict, status: int = 200) -> None:
        self._payload = json.dumps(payload).encode("utf-8")
        self.status = status
        self.headers = Message()
        self.headers["Content-Type"] = "application/json; charset=utf-8"

    def read(self) -> bytes:
        return self._payload

    def __enter__(self) -> "_FakeResponse":
        return self

    def __exit__(self, *_args: object) -> None:
        return None


class _FakeOpener:
    def __init__(self, responses: list[dict]) -> None:
        self.responses = list(responses)
        self.requests = []

    def open(self, request, timeout: int):
        self.requests.append((request, timeout))
        return _FakeResponse(self.responses.pop(0))


def _parsed_query_curl() -> dict:
    return {
        "url": "http://saasportal.iio.com/improxy/saas/iio2/sameAudit/queryAll",
        "cookie": "session=secret",
        "headers": [("Content-Type", "application/json")],
        "data": json.dumps(
            {
                "commodityId": "",
                "commodityName": "",
                "syncState": "1",
                "isReferUrl": "0",
                "page": 1,
                "limit": 10,
                "tenantId": "188605",
                "commodityPoolId": "222",
            }
        ),
        "method": "POST",
    }


class _PassThroughAuthManager:
    def __init__(self) -> None:
        self.calls = []

    def prepare_curl(self, system, parsed_curl, *, refresh=True):
        self.calls.append((system, refresh))
        prepared = dict(parsed_curl)
        prepared["cookie"] = "x-access-token=state-access; refreshToken=state-refresh"
        return prepared


def test_preview_refreshes_saas_auth_without_sending_a_write(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)
    curl_file = tmp_path / "query.curl"
    curl_file.write_text(
        "curl 'http://saasportal.iio.com/improxy/saas/iio2/sameAudit/queryAll' "
        "-H 'Content-Type: application/json' -b 'x-access-token=old; refreshToken=old-r' "
        "--data-raw '{\"tenantId\":\"188605\",\"commodityPoolId\":\"222\"}'",
        encoding="utf-8",
    )
    auth_manager = _PassThroughAuthManager()

    assert run_main(
        ["--input-xlsx", str(workbook), "--curl-file", str(curl_file)],
        auth_manager=auth_manager,
    ) == 0
    assert auth_manager.calls == [("saas", True)]


def test_fetch_pending_pairs_pages_to_reported_count() -> None:
    """Stopping after the first page could reject stale workbook rows."""
    opener = _FakeOpener(
        [
            {
                "rspCode": "0000",
                "rspDesc": "SUCCESS",
                "count": 3,
                "rows": [
                    {"vCommodityId": "101", "skuId": "sku-a"},
                    {"vCommodityId": "102", "skuId": "sku-b"},
                ],
            },
            {
                "rspCode": "0000",
                "rspDesc": "SUCCESS",
                "count": 3,
                "rows": [{"vCommodityId": "103", "skuId": "sku-c"}],
            },
        ]
    )

    result = fetch_pending_pairs(_parsed_query_curl(), opener=opener, page_size=2)

    assert result == {
        ("101", "sku-a"),
        ("102", "sku-b"),
        ("103", "sku-c"),
    }
    bodies = [json.loads(request.data) for request, _ in opener.requests]
    assert [(body["page"], body["limit"]) for body in bodies] == [(1, 2), (2, 2)]
    assert all(body["tenantId"] == "188605" for body in bodies)
    assert all(body["commodityPoolId"] == "222" for body in bodies)


def test_fetch_pending_pairs_can_include_all_reference_states() -> None:
    """Keeping the captured isReferUrl filter could hide a record that is still pending."""
    opener = _FakeOpener(
        [
            {
                "rspCode": "0000",
                "rspDesc": "SUCCESS",
                "count": 1,
                "rows": [{"vCommodityId": "101", "skuId": "sku-a"}],
            }
        ]
    )

    fetch_pending_pairs(
        _parsed_query_curl(),
        opener=opener,
        page_size=200,
        include_all_reference_states=True,
    )

    body = json.loads(opener.requests[0][0].data)
    assert body["syncState"] == "1"
    assert body["isReferUrl"] == ""


def test_validate_candidates_are_pending_reports_missing_keys() -> None:
    """Submitting workbook rows no longer pending would make the batch stale."""
    from scripts.reject_saas_audit_excel import RejectCandidate

    candidates = [
        RejectCandidate("101", "sku-a", "1", "red", 2),
        RejectCandidate("102", "sku-b", "1", "yellow", 3),
    ]

    with pytest.raises(ValueError, match="1 candidate"):
        validate_candidates_are_pending(candidates, {("101", "sku-a")})


def test_select_pending_candidates_excludes_exactly_reauthorized_missing_count() -> None:
    """Continuation authorization must not silently exclude more stale rows than approved."""
    from scripts.reject_saas_audit_excel import RejectCandidate

    candidates = [
        RejectCandidate("101", "sku-a", "1", "red", 2),
        RejectCandidate("102", "sku-b", "1", "yellow", 3),
    ]

    selected, missing = select_pending_candidates(
        candidates,
        {("102", "sku-b")},
        expected_missing_count=1,
    )

    assert selected == [candidates[1]]
    assert missing == [candidates[0]]
    with pytest.raises(ValueError, match="expected 0"):
        select_pending_candidates(
            candidates,
            {("102", "sku-b")},
            expected_missing_count=0,
        )


def test_execute_reject_batches_uses_fixed_payload_and_stops_on_failed_response() -> None:
    """Continuing after a failed write response would hide a partial reject batch."""
    from scripts.reject_saas_audit_excel import RejectCandidate

    candidates = [
        RejectCandidate("101", "sku-a", "1", "red", 2),
        RejectCandidate("102", "sku-b", "1", "yellow", 3),
        RejectCandidate("103", "sku-c", "1", "red", 4),
    ]
    opener = _FakeOpener(
        [
            {"rspCode": "0000", "rspDesc": "SUCCESS", "success": True},
            {"rspCode": "3333", "rspDesc": "login required", "success": False},
        ]
    )
    successes = []

    with pytest.raises(RuntimeError, match="rspCode=3333"):
        execute_reject_batches(
            _parsed_query_curl(),
            candidates,
            batch_size=2,
            opener=opener,
            on_batch_success=successes.append,
        )

    assert successes == [
        {
            "batchIndex": 1,
            "batchCount": 2,
            "firstExcelRow": 2,
            "lastExcelRow": 3,
            "rspCode": "0000",
            "rspDesc": "SUCCESS",
        }
    ]
    first_body = json.loads(opener.requests[0][0].data)
    assert first_body == {
        "tenantId": "188605",
        "commodityPoolId": "222",
        "commodityList": [
            {"vCommodityId": "101", "skuId": "sku-a", "modelAuditResult": "1"},
            {"vCommodityId": "102", "skuId": "sku-b", "modelAuditResult": "1"},
        ],
    }
    assert opener.requests[0][0].full_url.endswith("/sameAudit/reject")


def test_run_main_execute_writes_manifest_events_and_postflight_summary(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An executed batch must leave enough evidence to prevent an unsafe repeat."""
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)
    candidates = load_reject_candidates(workbook)
    digest = candidate_manifest_sha256(candidates)
    curl_file = tmp_path / "query.curl"
    curl_file.write_text(
        "curl 'http://saasportal.iio.com/improxy/saas/iio2/sameAudit/queryAll' "
        "-H 'Content-Type: application/json' -b 'session=secret' "
        "--data-raw '{\"tenantId\":\"188605\",\"commodityPoolId\":\"222\"}'",
        encoding="utf-8",
    )
    output_dir = tmp_path / "run"
    pending_before = {
        (candidate.v_commodity_id, candidate.sku_id) for candidate in candidates
    }
    pending_results = iter([pending_before, set()])
    monkeypatch.setattr(
        "scripts.reject_saas_audit_excel.fetch_pending_pairs",
        lambda *_args, **_kwargs: next(pending_results),
    )

    def _execute(_curl, items, **kwargs):
        event = {
            "batchIndex": 1,
            "batchCount": len(items),
            "firstExcelRow": 2,
            "lastExcelRow": 3,
            "rspCode": "0000",
            "rspDesc": "SUCCESS",
        }
        kwargs["on_batch_success"](event)
        return [event]

    monkeypatch.setattr(
        "scripts.reject_saas_audit_excel.execute_reject_batches", _execute
    )

    exit_code = run_main(
        [
            "--input-xlsx",
            str(workbook),
            "--curl-file",
            str(curl_file),
            "--output-dir",
            str(output_dir),
            "--execute",
            "--expected-count",
            "2",
            "--expected-sha256",
            digest,
        ],
        auth_manager=_PassThroughAuthManager(),
    )

    assert exit_code == 0
    manifest = json.loads((output_dir / "reject_manifest.json").read_text("utf-8"))
    summary = json.loads((output_dir / "reject_summary.json").read_text("utf-8"))
    events = [
        json.loads(line)
        for line in (output_dir / "reject_events.jsonl").read_text("utf-8").splitlines()
    ]
    assert manifest["candidateCount"] == 2
    assert manifest["candidateManifestSha256"] == digest
    assert "cookie" not in json.dumps(manifest).lower()
    assert events == [
        {
            "batchIndex": 1,
            "batchCount": 2,
            "firstExcelRow": 2,
            "lastExcelRow": 3,
            "rspCode": "0000",
            "rspDesc": "SUCCESS",
        }
    ]
    assert summary["status"] == "COMPLETED"
    assert summary["rejectedCount"] == 2
    assert summary["remainingPendingCount"] == 0


def test_run_main_can_continue_with_exactly_one_already_processed_candidate(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A continuation may exclude only the exact stale count the user reauthorized."""
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)
    all_candidates = load_reject_candidates(workbook)
    selected = [all_candidates[1]]
    digest = candidate_manifest_sha256(selected)
    curl_file = tmp_path / "query.curl"
    curl_file.write_text(
        "curl 'http://saasportal.iio.com/improxy/saas/iio2/sameAudit/queryAll' "
        "-H 'Content-Type: application/json' -b 'session=secret' "
        "--data-raw '{\"tenantId\":\"188605\",\"commodityPoolId\":\"222\"}'",
        encoding="utf-8",
    )
    output_dir = tmp_path / "run"
    pending_results = iter([{("102", "sku-yellow")}, set()])
    monkeypatch.setattr(
        "scripts.reject_saas_audit_excel.fetch_pending_pairs",
        lambda *_args, **_kwargs: next(pending_results),
    )

    def _execute(_curl, items, **kwargs):
        event = {
            "batchIndex": 1,
            "batchCount": 1,
            "firstExcelRow": 3,
            "lastExcelRow": 3,
            "rspCode": "0000",
            "rspDesc": "SUCCESS",
        }
        assert list(items) == selected
        kwargs["on_batch_success"](event)
        return [event]

    monkeypatch.setattr(
        "scripts.reject_saas_audit_excel.execute_reject_batches", _execute
    )

    assert run_main(
        [
            "--input-xlsx",
            str(workbook),
            "--curl-file",
            str(curl_file),
            "--output-dir",
            str(output_dir),
            "--execute",
            "--expected-count",
            "1",
            "--expected-sha256",
            digest,
            "--expected-already-processed-count",
            "1",
        ],
        auth_manager=_PassThroughAuthManager(),
    ) == 0

    manifest = json.loads((output_dir / "reject_manifest.json").read_text("utf-8"))
    assert manifest["sourceCandidateCount"] == 2
    assert manifest["candidateCount"] == 1
    assert manifest["alreadyProcessedCount"] == 1
    assert manifest["alreadyProcessedCandidates"] == [
        all_candidates[0].as_manifest_item()
    ]
