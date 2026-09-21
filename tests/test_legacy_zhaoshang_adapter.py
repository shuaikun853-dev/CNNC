import json
import os
import subprocess
import sys
import time
from dataclasses import replace
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from pymysql.err import InterfaceError, OperationalError

from app.services.legacy_zhaoshang_adapter import (
    CrawlRequest,
    CnncLegacyZhaoshangAdapter,
    HarvestStatus,
    LegacyPreSendFailure,
    PriorityStatus,
    ReconcileStatus,
    SubmitStatus,
    TemplateConflictError,
)
from app.services import legacy_zhaoshang_adapter as legacy_module
from scripts.zhaoshang_mall_mapping import ZhaoshangMallMapping


MAPPING = ZhaoshangMallMapping(
    mall_name="京东",
    source_id="1",
    mall_key="jd",
    domains=("jd.com",),
    vendor_keywords=("jd", "京东"),
)
DEFAULT_TASK_NAME = "cnnc_current_opt_京东_20260811180000000000"


def _adapter(**kwargs):
    kwargs.setdefault("db_env_path", None)
    return CnncLegacyZhaoshangAdapter(mappings=[MAPPING], **kwargs)


def _request(sku_id="sku-1", url="https://item.jd.com/a", backend_version="legacy-v1"):
    return CrawlRequest(
        mall_key="jd",
        mall_source_id="1",
        mall_name="京东",
        sku_id=sku_id,
        url=url.strip(),
        backend_version=backend_version,
    )


def _state_file(tmp_path):
    path = tmp_path / "BATCH_STATE.md"
    path.write_text("# state\n", encoding="utf-8")
    return path


def _prepare(tmp_path, requests=(_request(),), **kwargs):
    return _adapter(**kwargs).prepare_job(
        business_date="2026-08-11",
        requests=requests,
        template_path=tmp_path / "template.xlsx",
        result_path=tmp_path / "result.xlsx",
        task_name=DEFAULT_TASK_NAME,
        batch_state_path=_state_file(tmp_path),
    )


def _append_event(path, action, **details):
    event = {"mall": "jd", "task_name": "cnnc_current_opt_京东_20260811180000000000", "action": action, "task_id": "id-default", "task_number": "no-default", **details}
    with path.open("a", encoding="utf-8") as stream:
        stream.write(f"- `{json.dumps(event, ensure_ascii=False)}`\n")


def _append_submission(path, *, task_id="id-default", task_number="no-default", expected_subtask_count=1):
    _append_event(
        path,
        "CREATED",
        task_id=task_id,
        task_number=task_number,
        expected_subtask_count=expected_subtask_count,
    )
    _append_event(path, "ENABLED", task_id=task_id, task_number=task_number)


def _write_result(path: Path, rows, *, task_name=DEFAULT_TASK_NAME, mall_key="jd"):
    _write_result_with_headers(
        path,
        ["商品编码", "原始链接", "商品价格"],
        rows,
        task_name=task_name,
        mall_key=mall_key,
    )


def _write_result_with_headers(
    path: Path,
    headers,
    rows,
    *,
    task_name=DEFAULT_TASK_NAME,
    mall_key="jd",
    include_provenance=True,
):
    headers = list(headers)
    rows = [list(row) for row in rows]
    if include_provenance and "批次号" not in headers:
        headers.extend(["执行编号", "批次号", "商城标识"])
        rows = [
            [*row, f"no-default-execution-{index}", task_name, mall_key]
            for index, row in enumerate(rows)
        ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _write_provenance_result(path: Path, rows):
    _write_result_with_headers(
        path,
        ["商品编码", "商品详情地址", "商品价格", "执行编号", "批次号", "商城标识"],
        rows,
    )


def test_mapping_prefers_url_domain_and_uses_vendor_only_for_blank_url():
    adapter = _adapter()

    from_url = adapter.request_from_saas_row({"skuId": "s", "extLink": "https://item.jd.com/a", "vendorName": "other"}, backend_version="v1")
    nonblank_unknown = adapter.request_from_saas_row({"skuId": "s", "extLink": "https://other.example/a", "vendorName": "京东"}, backend_version="v1")
    from_vendor = adapter.request_from_saas_row({"skuId": "s", "extLink": "", "vendorName": "京东"}, backend_version="v1")

    assert from_url is not None and from_url.mall_key == "jd"
    assert nonblank_unknown is None
    assert from_vendor is not None and from_vendor.url == ""


def test_mapping_parses_each_row_url_once_and_caches_repeated_host(monkeypatch):
    calls = []
    original = legacy_module.urlparse

    def counted(value):
        calls.append(value)
        return original(value)

    monkeypatch.setattr(legacy_module, "urlparse", counted)
    adapter = _adapter()

    assert adapter.request_from_saas_row(
        {"skuId": "s1", "extLink": "https://item.jd.com/a"}, backend_version="v1"
    ) is not None
    assert adapter.request_from_saas_row(
        {"skuId": "s2", "extLink": "https://item.jd.com/b"}, backend_version="v1"
    ) is not None

    assert calls == ["https://item.jd.com/a", "https://item.jd.com/b"]
    assert list(adapter._host_mapping_cache) == ["item.jd.com"]


def test_request_identity_keeps_same_url_with_two_skus_and_hashes_are_order_independent(tmp_path):
    first = _request("sku-1")
    second = _request("sku-2")
    adapter = _adapter()
    job_one = adapter.prepare_job(
        "2026-08-11", (first, second), tmp_path / "one.xlsx", tmp_path / "one-result.xlsx", "task-one", _state_file(tmp_path)
    )
    job_two = adapter.prepare_job(
        "2026-08-11", (second, first), tmp_path / "two.xlsx", tmp_path / "two-result.xlsx", "task-two", _state_file(tmp_path)
    )

    assert len(job_one.requests) == 2
    assert job_one.logical_job_key == job_two.logical_job_key
    assert job_one.payload_sha256 == job_two.payload_sha256


def test_immutable_template_reuses_matching_payload_and_rejects_conflict(tmp_path):
    adapter = _adapter()
    state = _state_file(tmp_path)
    first = adapter.prepare_job("2026-08-11", (_request(),), tmp_path / "template.xlsx", tmp_path / "result.xlsx", "task", state)
    second = adapter.prepare_job("2026-08-11", (_request(),), tmp_path / "template.xlsx", tmp_path / "result.xlsx", "task", state)

    assert first.payload_sha256 == second.payload_sha256
    with pytest.raises(TemplateConflictError):
        adapter.prepare_job("2026-08-11", (_request("sku-2"),), tmp_path / "template.xlsx", tmp_path / "result.xlsx", "task", state)


def test_template_publication_retries_transient_windows_permission_error(tmp_path, monkeypatch):
    adapter = _adapter()
    state = _state_file(tmp_path)
    target = tmp_path / "template.xlsx"
    real_replace = os.replace
    replace_calls = 0

    def transient_replace(source, destination):
        nonlocal replace_calls
        if Path(destination) == target:
            replace_calls += 1
            if replace_calls == 1:
                raise PermissionError("transient scanner lock")
        return real_replace(source, destination)

    monkeypatch.setattr(legacy_module.os, "replace", transient_replace)

    adapter.prepare_job("2026-08-11", (_request(),), target, tmp_path / "result.xlsx", "task", state)

    assert replace_calls == 2
    assert target.is_file()
    assert not list(tmp_path.glob(".template.xlsx.*.tmp"))


def test_template_requires_embedded_payload_evidence_and_never_overwrites_unreadable_file(tmp_path):
    adapter = _adapter()
    state = _state_file(tmp_path)
    path = tmp_path / "template.xlsx"
    adapter.prepare_job("2026-08-11", (_request(),), path, tmp_path / "result.xlsx", "task", state)
    workbook = load_workbook(path)
    try:
        workbook.properties.description = "legacy_zhaoshang_payload_sha256=wrong"
        workbook.save(path)
    finally:
        workbook.close()
    with pytest.raises(TemplateConflictError):
        adapter.prepare_job("2026-08-11", (_request(),), path, tmp_path / "result.xlsx", "task", state)

    unreadable = tmp_path / "unreadable.xlsx"
    unreadable.write_bytes(b"")
    with pytest.raises(TemplateConflictError):
        adapter.prepare_job("2026-08-11", (_request(),), unreadable, tmp_path / "result.xlsx", "other-task", state)
    assert unreadable.read_bytes() == b""


def test_submit_uses_exact_create_only_cli_once_and_preserves_priority_evidence(tmp_path):
    calls = []

    def runner(args, timeout_seconds, env):
        calls.append((args, timeout_seconds, env))
        state = Path(args[args.index("--batch-state-file") + 1])
        _append_submission(state, task_id="id-1", task_number="no-1")
        _append_event(state, "PRIORITY_UPDATED", task_id="id-1", task_number="no-1", affected_rows=1)
        return subprocess.CompletedProcess(args, 0, stdout="ok", stderr="")

    adapter = _adapter(runner=runner)
    job = _prepare(tmp_path)
    outcome = adapter.submit(job, timeout_seconds=30)

    assert outcome.status is SubmitStatus.SUBMITTED
    assert outcome.task_id == "id-1"
    assert len(calls) == 1
    args = calls[0][0]
    assert "--create-only" in args
    assert "--wait-export-only" not in args
    assert args[args.index("--task-name") + 1] == job.task_name
    assert args[args.index("--task-weight") + 1] == "0"


def test_submit_confirms_created_and_enabled_without_priority_evidence(tmp_path):
    calls = []

    def runner(args, timeout_seconds, env):
        calls.append((args, timeout_seconds, env))
        state = Path(args[args.index("--batch-state-file") + 1])
        _append_event(state, "CREATED", task_id="id-1", task_number="no-1", expected_subtask_count=1)
        _append_event(state, "ENABLED", task_id="id-1", task_number="no-1")
        return subprocess.CompletedProcess(args, 1, stdout="", stderr="priority failed")

    outcome = _adapter(runner=runner).submit(_prepare(tmp_path), timeout_seconds=30)

    assert outcome.status is SubmitStatus.SUBMITTED
    assert (outcome.task_id, outcome.task_number) == ("id-1", "no-1")
    assert outcome.priority_evidence is None
    assert len(calls) == 1


@pytest.mark.parametrize(
    "action, details, expected",
    [
        ("PRIORITY_UPDATED", {"affected_rows": 2}, "updated"),
        ("PRIORITY_ALREADY_HIGHEST", {"already_highest_count": 2}, "already_highest_priority"),
        ("PRIORITY_UNCONFIRMED", {}, "completed_unconfirmed"),
    ],
)
def test_submit_accepts_only_supported_priority_evidence(tmp_path, action, details, expected):
    def runner(args, _timeout, _env):
        state = Path(args[args.index("--batch-state-file") + 1])
        _append_submission(state)
        _append_event(state, action, **details)
        return subprocess.CompletedProcess(args, 0, stdout="", stderr="")

    outcome = _adapter(runner=runner).submit(_prepare(tmp_path), timeout_seconds=30)

    assert outcome.status is SubmitStatus.SUBMITTED
    assert outcome.priority_evidence == expected


def test_submit_binds_priority_evidence_to_identity_from_the_same_event(tmp_path):
    def runner(args, _timeout, _env):
        state = Path(args[args.index("--batch-state-file") + 1])
        _append_submission(state, task_id="priority-id", task_number="priority-number")
        _append_event(state, "PRIORITY_UPDATED", task_id="priority-id", task_number="priority-number", affected_rows=1)
        _append_event(state, "EXPORTED", task_id="later-id", task_number="later-number")
        return subprocess.CompletedProcess(args, 0, stdout="", stderr="")

    outcome = _adapter(runner=runner).submit(_prepare(tmp_path), timeout_seconds=30)

    assert outcome.status is SubmitStatus.SUBMITTED
    assert (outcome.task_id, outcome.task_number) == ("priority-id", "priority-number")


def test_submit_fails_closed_on_conflicting_accepted_priority_identities(tmp_path):
    def runner(args, _timeout, _env):
        state = Path(args[args.index("--batch-state-file") + 1])
        _append_submission(state)
        _append_event(state, "PRIORITY_UPDATED", task_id="first-id", task_number="first-number", affected_rows=1)
        _append_event(
            state,
            "PRIORITY_ALREADY_HIGHEST",
            task_id="second-id",
            task_number="second-number",
            already_highest_count=1,
        )
        return subprocess.CompletedProcess(args, 0, stdout="", stderr="")

    outcome = _adapter(runner=runner).submit(_prepare(tmp_path), timeout_seconds=30)

    assert outcome.status is SubmitStatus.SUBMITTED
    assert outcome.priority_evidence is None
    assert "conflict" in outcome.diagnostic


def test_production_runner_shares_deadline_and_strips_model_secrets(tmp_path, monkeypatch):
    tool = tmp_path / "auto_crawl_task.py"
    tool.write_text("# test tool\n", encoding="utf-8")
    monkeypatch.setenv("TITLE_MATCH_MODEL_API_KEY", "must-not-leak")
    monkeypatch.setenv("TITLE_MATCH_MODEL_NAME", "must-not-leak")
    captured = {}

    class Auth:
        def export_environment(self, *_args, **_kwargs):
            return {
                "BLADE_AUTH": "blade",
                "COOKIE": "cookie",
                "TITLE_MATCH_MODEL_API_KEY": "must-still-not-leak",
            }

    ticks = iter((100.0, 104.0))
    adapter = _adapter(
        auth_manager=Auth(),
        external_tool_path=tool,
        monotonic=lambda: next(ticks),
    )

    process = _FakeSpawnedProcess()

    def popen(*args, **kwargs):
        captured.update(kwargs)
        return process

    monkeypatch.setattr(legacy_module.subprocess, "Popen", popen)
    adapter._production_runner([sys.executable, str(legacy_module.EXTERNAL_TOOL_PATH)], 10)

    assert process.communicate_timeouts == [pytest.approx(6)]
    assert captured["env"]["BLADE_AUTH"] == "blade"
    assert not any(name.startswith("TITLE_MATCH_MODEL_") for name in captured["env"])
    assert os.environ["TITLE_MATCH_MODEL_API_KEY"] == "must-not-leak"


def test_production_runner_loads_required_priority_database_environment(tmp_path, monkeypatch):
    tool = tmp_path / "auto_crawl_task.py"
    tool.write_text("# test tool\n", encoding="utf-8")
    db_env = tmp_path / "ZHAOSHANG_DB.env"
    db_env.write_text(
        "ZHAOSHANG_DB_HOST=host\n"
        "ZHAOSHANG_DB_PORT=3306\n"
        "ZHAOSHANG_DB_USER=user\n"
        "ZHAOSHANG_DB_PASSWORD=password\n",
        encoding="utf-8",
    )
    captured = {}

    class Auth:
        def export_environment(self, *_args, **_kwargs):
            return {"BLADE_AUTH": "blade"}

    def popen(*args, **kwargs):
        captured.update(kwargs)
        return _FakeSpawnedProcess()

    monkeypatch.setattr(legacy_module.subprocess, "Popen", popen)
    adapter = _adapter(auth_manager=Auth(), external_tool_path=tool, db_env_path=db_env)
    adapter._production_runner([sys.executable, str(legacy_module.EXTERNAL_TOOL_PATH)], 10)

    assert {name for name in captured["env"] if name.startswith("ZHAOSHANG_DB_")} == {
        "ZHAOSHANG_DB_HOST",
        "ZHAOSHANG_DB_PORT",
        "ZHAOSHANG_DB_USER",
        "ZHAOSHANG_DB_PASSWORD",
    }
    assert not any(name.startswith("ZHAOSHANG_DB_") for name in os.environ)


def test_production_runner_never_starts_subprocess_after_auth_uses_deadline(tmp_path, monkeypatch):
    tool = tmp_path / "auto_crawl_task.py"
    tool.write_text("# test tool\n", encoding="utf-8")

    class Auth:
        def export_environment(self, *_args, **_kwargs):
            return {"BLADE_AUTH": "blade", "COOKIE": "cookie"}

    ticks = iter((100.0, 111.0))
    adapter = _adapter(
        auth_manager=Auth(),
        external_tool_path=tool,
        monotonic=lambda: next(ticks),
    )
    monkeypatch.setattr(
        legacy_module.subprocess,
        "Popen",
        lambda *_args, **_kwargs: pytest.fail("subprocess must not start after deadline"),
    )

    with pytest.raises(LegacyPreSendFailure) as excinfo:
        adapter._production_runner([sys.executable, str(legacy_module.EXTERNAL_TOOL_PATH)], 10)

    assert "TimeoutError" in str(excinfo.value)


def test_default_reconciler_shares_auth_deadline_with_all_read_requests(tmp_path, monkeypatch):
    tool = tmp_path / "auto_crawl_task.py"
    tool.write_text("# test tool\n", encoding="utf-8")
    request_timeouts = []

    class Auth:
        def export_environment(self, *_args, **_kwargs):
            return {"BLADE_AUTH": "blade", "COOKIE": "cookie"}

    class Session:
        def get(self, *_args, **kwargs):
            request_timeouts.append(kwargs["timeout"])
            return object()

    class Loader:
        def exec_module(self, module):
            module.build_session = lambda: Session()
            module.priority_mall_code_for_key = lambda mall_key: mall_key

            def check_existing_task(session, *_args):
                session.get("https://read-only.invalid", timeout=60)
                return {"taskId": "id", "taskNumber": "number"}

            module.check_existing_task = check_existing_task

    class Spec:
        loader = Loader()

    class Module:
        pass

    monkeypatch.setattr(legacy_module.importlib.util, "spec_from_file_location", lambda *_args: Spec())
    monkeypatch.setattr(legacy_module.importlib.util, "module_from_spec", lambda _spec: Module())
    ticks = iter((100.0, 104.0, 105.0))
    adapter = _adapter(
        auth_manager=Auth(),
        external_tool_path=tool,
        monotonic=lambda: next(ticks),
    )

    remote = adapter._default_reconciler("task", "jd", 10)

    assert remote == {"taskId": "id", "taskNumber": "number"}
    assert request_timeouts == [pytest.approx(5)]


def test_default_reconciler_never_builds_session_after_auth_exhausts_deadline(tmp_path, monkeypatch):
    tool = tmp_path / "auto_crawl_task.py"
    tool.write_text("# test tool\n", encoding="utf-8")

    class Auth:
        def export_environment(self, *_args, **_kwargs):
            return {"BLADE_AUTH": "blade"}

    class Loader:
        def exec_module(self, module):
            module.build_session = lambda: pytest.fail("session must not be built after deadline")
            module.priority_mall_code_for_key = lambda mall_key: mall_key
            module.check_existing_task = lambda *_args: None

    class Spec:
        loader = Loader()

    class Module:
        pass

    monkeypatch.setattr(legacy_module.importlib.util, "spec_from_file_location", lambda *_args: Spec())
    monkeypatch.setattr(legacy_module.importlib.util, "module_from_spec", lambda _spec: Module())
    ticks = iter((100.0, 111.0))
    adapter = _adapter(
        auth_manager=Auth(),
        external_tool_path=tool,
        monotonic=lambda: next(ticks),
    )

    with pytest.raises(TimeoutError):
        adapter._default_reconciler("task", "jd", 10)


def test_submit_timeout_is_unknown_and_never_retried(tmp_path):
    calls = []

    def runner(args, _timeout, _env):
        calls.append(args)
        raise subprocess.TimeoutExpired(args, 30)

    outcome = _adapter(runner=runner).submit(_prepare(tmp_path), timeout_seconds=30)

    assert outcome.status is SubmitStatus.SUBMIT_UNKNOWN
    assert len(calls) == 1


def test_reconcile_unknown_uses_exact_identity_without_write_calls(tmp_path):
    calls = []

    def reconciler(task_name, mall_key, timeout_seconds):
        calls.append((task_name, mall_key, timeout_seconds))
        return {"taskId": "id-1", "taskNumber": "no-1", "mallLogo": "jd"}

    job = _prepare(tmp_path)
    _append_submission(job.batch_state_path, task_id="id-1", task_number="no-1")
    _append_event(
        job.batch_state_path,
        "PRIORITY_ALREADY_HIGHEST",
        task_id="id-1",
        task_number="no-1",
        already_highest_count=1,
    )
    outcome = _adapter(reconciler=reconciler).reconcile_unknown(job, timeout_seconds=12)

    assert outcome.status is ReconcileStatus.FOUND
    assert outcome.task_id == "id-1"
    assert calls == [("cnnc_current_opt_京东_20260811180000000000", "jd", 12)]


def test_reconcile_unknown_confirms_created_enabled_without_priority_or_post(tmp_path):
    reconcile_calls = []
    post_calls = []

    def reconciler(task_name, mall_key, timeout_seconds):
        reconcile_calls.append((task_name, mall_key, timeout_seconds))
        return {"taskId": "id-1", "taskNumber": "no-1", "mallLogo": "jd"}

    job = _prepare(tmp_path)
    _append_event(job.batch_state_path, "CREATED", task_id="id-1", task_number="no-1", expected_subtask_count=1)
    _append_event(job.batch_state_path, "ENABLED", task_id="id-1", task_number="no-1")
    adapter = _adapter(
        reconciler=reconciler,
        runner=lambda *_args: post_calls.append(True),
    )

    first = adapter.reconcile_unknown(job, timeout_seconds=12)
    second = adapter.reconcile_unknown(job, timeout_seconds=12)

    assert first.status is ReconcileStatus.FOUND
    assert second.status is ReconcileStatus.FOUND
    assert (first.task_id, first.task_number) == ("id-1", "no-1")
    assert post_calls == []
    assert len(reconcile_calls) == 2


def test_reconcile_unknown_adopts_one_exact_remote_after_durable_intent(tmp_path):
    job = _prepare(tmp_path)
    _append_event(
        job.batch_state_path,
        "ADD_INTENT",
        task_id="",
        task_number="",
        expected_subtask_count=1,
        timestamp="2026-08-14T09:01:02+08:00",
    )
    recorded = []
    post_calls = []

    def record_event(path, **event):
        recorded.append(dict(event))
        _append_event(
            Path(path),
            event.pop("action"),
            task_id=event.pop("task_id", ""),
            task_number=event.pop("task_number", ""),
            **event,
        )

    adapter = _adapter(
        reconciler=lambda *_args: {
            "taskId": "id-1",
            "taskNumber": "no-1",
            "taskName": job.task_name,
            "mallLogo": job.mall_key,
            "expectedSubtaskCount": 1,
            "state": "1",
        },
        event_recorder=record_event,
        runner=lambda *_args: post_calls.append(True),
    )

    first = adapter.reconcile_unknown(job, timeout_seconds=12)
    second = adapter.reconcile_unknown(job, timeout_seconds=12)

    assert first.status is ReconcileStatus.FOUND
    assert second.status is ReconcileStatus.FOUND
    assert (first.task_id, first.task_number) == ("id-1", "no-1")
    assert [event["action"] for event in recorded] == ["CREATED", "ENABLED"]
    assert recorded[0]["expected_subtask_count"] == 1
    assert recorded[0]["submission_marker_time"] == "2026-08-14T09:01:02+08:00"
    assert post_calls == []
    events = adapter._read_batch_events(job)
    assert adapter._submission_evidence(job, events) == ("id-1", "no-1")
    assert adapter._submission_marker_time(events, "id-1", "no-1") == datetime(
        2026, 8, 14, 9, 1, 2
    )


def test_submit_connection_reset_then_reconcile_retry_never_posts_twice(tmp_path):
    job = _prepare(tmp_path)
    _append_event(
        job.batch_state_path,
        "ADD_INTENT",
        task_id="",
        task_number="",
        timestamp="2026-08-14T09:01:02+08:00",
    )
    add_calls = []

    def runner(*_args):
        add_calls.append("/add")
        raise ConnectionResetError("response connection reset")

    def record_event(path, **event):
        _append_event(
            Path(path),
            event.pop("action"),
            task_id=event.pop("task_id", ""),
            task_number=event.pop("task_number", ""),
            **event,
        )

    adapter = _adapter(
        runner=runner,
        reconciler=lambda *_args: {
            "taskId": "id-1",
            "taskNumber": "no-1",
            "taskName": job.task_name,
            "mallLogo": job.mall_key,
            "expectedSubtaskCount": 1,
            "state": "1",
        },
        event_recorder=record_event,
    )

    submitted = adapter.submit(job, timeout_seconds=12)
    first = adapter.reconcile_unknown(job, timeout_seconds=12)
    second = adapter.reconcile_unknown(job, timeout_seconds=12)

    assert submitted.status is SubmitStatus.SUBMIT_UNKNOWN
    assert first.status is ReconcileStatus.FOUND
    assert second.status is ReconcileStatus.FOUND
    assert add_calls == ["/add"]


def test_exact_takeover_finishes_partial_created_evidence_after_process_crash(tmp_path):
    job = _prepare(tmp_path)
    _append_event(
        job.batch_state_path,
        "ADD_INTENT",
        task_id="",
        task_number="",
        timestamp="2026-08-14T09:01:02+08:00",
    )
    _append_event(
        job.batch_state_path,
        "CREATED",
        task_id="id-1",
        task_number="no-1",
        expected_subtask_count=1,
        timestamp="2026-08-14T09:01:05+08:00",
    )

    def record_event(path, **event):
        _append_event(
            Path(path),
            event.pop("action"),
            task_id=event.pop("task_id", ""),
            task_number=event.pop("task_number", ""),
            **event,
        )

    adapter = _adapter(
        reconciler=lambda *_args: {
            "taskId": "id-1",
            "taskNumber": "no-1",
            "taskName": job.task_name,
            "mallLogo": job.mall_key,
            "expectedSubtaskCount": 1,
            "state": "1",
        },
        event_recorder=record_event,
    )

    outcome = adapter.reconcile_unknown(job, timeout_seconds=12)
    events = adapter._read_batch_events(job)

    assert outcome.status is ReconcileStatus.FOUND
    assert adapter._submission_marker_time(events, "id-1", "no-1") == datetime(
        2026, 8, 14, 9, 1, 2
    )


@pytest.mark.parametrize(
    "remote",
    (
        {
            "taskId": "id-1",
            "taskNumber": "no-1",
            "taskName": "wrong-task",
            "mallLogo": "jd",
            "expectedSubtaskCount": 1,
        },
        {
            "taskId": "id-1",
            "taskNumber": "no-1",
            "taskName": DEFAULT_TASK_NAME,
            "mallLogo": "wrong-mall",
            "expectedSubtaskCount": 1,
        },
        {
            "taskId": "id-1",
            "taskNumber": "no-1",
            "taskName": DEFAULT_TASK_NAME,
            "mallLogo": "jd",
            "expectedSubtaskCount": 2,
        },
    ),
)
def test_reconcile_unknown_rejects_exact_remote_identity_or_count_mismatch(
    tmp_path, remote
):
    job = _prepare(tmp_path)
    _append_event(
        job.batch_state_path,
        "ADD_INTENT",
        task_id="",
        task_number="",
        expected_subtask_count=1,
        timestamp="2026-08-14T09:01:02+08:00",
    )
    recorded = []

    outcome = _adapter(
        reconciler=lambda *_args: remote,
        event_recorder=lambda *_args, **event: recorded.append(event),
    ).reconcile_unknown(job, timeout_seconds=12)

    assert outcome.status is ReconcileStatus.MANUAL
    assert recorded == []


def test_reconcile_unknown_rejects_multiple_exact_remote_matches(tmp_path):
    job = _prepare(tmp_path)
    _append_event(
        job.batch_state_path,
        "ADD_INTENT",
        task_id="",
        task_number="",
        expected_subtask_count=1,
        timestamp="2026-08-14T09:01:02+08:00",
    )
    remote = {
        "taskId": "id-1",
        "taskNumber": "no-1",
        "taskName": job.task_name,
        "mallLogo": job.mall_key,
        "expectedSubtaskCount": 1,
        "state": "1",
    }

    outcome = _adapter(reconciler=lambda *_args: [remote, dict(remote)]).reconcile_unknown(
        job, timeout_seconds=12
    )

    assert outcome.status is ReconcileStatus.MANUAL


def test_reconcile_unknown_does_not_adopt_when_evidence_write_fails(tmp_path):
    job = _prepare(tmp_path)
    _append_event(
        job.batch_state_path,
        "ADD_INTENT",
        task_id="",
        task_number="",
        expected_subtask_count=1,
        timestamp="2026-08-14T09:01:02+08:00",
    )

    def fail_record(*_args, **_kwargs):
        raise OSError("simulated durable event failure")

    outcome = _adapter(
        reconciler=lambda *_args: {
            "taskId": "id-1",
            "taskNumber": "no-1",
            "taskName": job.task_name,
            "mallLogo": job.mall_key,
            "expectedSubtaskCount": 1,
            "state": "1",
        },
        event_recorder=fail_record,
    ).reconcile_unknown(job, timeout_seconds=12)

    assert outcome.status is ReconcileStatus.NOT_FOUND
    assert "durable" in outcome.diagnostic


def test_priority_update_is_verified_evidence_and_idempotent(tmp_path):
    calls = []
    job = _prepare(tmp_path)
    _append_event(job.batch_state_path, "CREATED", task_id="id-1", task_number="no-1", expected_subtask_count=1)
    _append_event(job.batch_state_path, "ENABLED", task_id="id-1", task_number="no-1")

    def prioritize(prepared, task_id, task_number, timeout_seconds):
        calls.append((prepared.task_name, task_id, task_number, timeout_seconds))
        _append_event(
            prepared.batch_state_path,
            "PRIORITY_UPDATED",
            task_id=task_id,
            task_number=task_number,
            affected_rows=1,
        )

    adapter = _adapter(priority_runner=prioritize)

    first = adapter.prioritize(job, task_id="id-1", task_number="no-1", timeout_seconds=12)
    second = adapter.prioritize(job, task_id="id-1", task_number="no-1", timeout_seconds=12)

    assert first.status is PriorityStatus.UPDATED
    assert second.status is PriorityStatus.UPDATED
    assert first.weight_type == -1
    assert calls == [(job.task_name, "id-1", "no-1", 12)]


def test_priority_already_highest_is_idempotent_without_update(tmp_path):
    calls = []
    job = _prepare(tmp_path)
    _append_event(job.batch_state_path, "CREATED", task_id="id-1", task_number="no-1", expected_subtask_count=1)
    _append_event(job.batch_state_path, "ENABLED", task_id="id-1", task_number="no-1")

    def prioritize(prepared, task_id, task_number, _timeout_seconds):
        calls.append((task_id, task_number))
        _append_event(
            prepared.batch_state_path,
            "PRIORITY_ALREADY_HIGHEST",
            task_id=task_id,
            task_number=task_number,
            already_highest_count=1,
        )

    adapter = _adapter(priority_runner=prioritize)
    first = adapter.prioritize(job, task_id="id-1", task_number="no-1", timeout_seconds=12)
    second = adapter.prioritize(job, task_id="id-1", task_number="no-1", timeout_seconds=12)

    assert first.status is PriorityStatus.ALREADY_HIGHEST
    assert second.status is PriorityStatus.ALREADY_HIGHEST
    assert first.weight_type == -1
    assert calls == [("id-1", "no-1")]


def test_priority_unconfirmed_never_reports_highest_or_retries_automatically(tmp_path):
    calls = []
    job = _prepare(tmp_path)
    _append_event(job.batch_state_path, "CREATED", task_id="id-1", task_number="no-1", expected_subtask_count=1)
    _append_event(job.batch_state_path, "ENABLED", task_id="id-1", task_number="no-1")

    def prioritize(prepared, task_id, task_number, _timeout_seconds):
        calls.append((task_id, task_number))
        _append_event(
            prepared.batch_state_path,
            "PRIORITY_UNCONFIRMED",
            task_id=task_id,
            task_number=task_number,
            reason="verification_unavailable",
        )

    adapter = _adapter(priority_runner=prioritize)
    first = adapter.prioritize(job, task_id="id-1", task_number="no-1", timeout_seconds=12)
    second = adapter.prioritize(job, task_id="id-1", task_number="no-1", timeout_seconds=12)

    assert first.status is PriorityStatus.UNCONFIRMED
    assert second.status is PriorityStatus.UNCONFIRMED
    assert first.weight_type is None
    assert calls == [("id-1", "no-1")]


@pytest.mark.parametrize(
    "error",
    [
        TimeoutError("priority timeout"),
        ConnectionError("priority connection lost"),
        InterfaceError("priority interface unavailable"),
    ],
)
def test_priority_transient_infrastructure_errors_are_explicitly_retryable(
    tmp_path, error
):
    job = _prepare(tmp_path)
    _append_event(
        job.batch_state_path,
        "CREATED",
        task_id="id-1",
        task_number="no-1",
        expected_subtask_count=1,
    )
    _append_event(
        job.batch_state_path,
        "ENABLED",
        task_id="id-1",
        task_number="no-1",
    )

    def prioritize(*_args):
        raise error

    outcome = _adapter(priority_runner=prioritize).prioritize(
        job,
        task_id="id-1",
        task_number="no-1",
        timeout_seconds=12,
    )

    assert outcome.status is PriorityStatus.UNCONFIRMED
    assert outcome.retryable is True


@pytest.mark.parametrize(
    "error_code",
    [1040, 1053, 1203, 1205, 1206, 1213, 2002, 2003, 2004, 2006, 2013, 2055],
)
def test_priority_transient_operational_error_codes_are_retryable(
    tmp_path, error_code
):
    job = _prepare(tmp_path)
    _append_submission(job.batch_state_path, task_id="id-1", task_number="no-1")

    def prioritize(*_args):
        raise OperationalError(error_code, "message is not used for classification")

    outcome = _adapter(priority_runner=prioritize).prioritize(
        job,
        task_id="id-1",
        task_number="no-1",
        timeout_seconds=12,
    )

    assert outcome.status is PriorityStatus.UNCONFIRMED
    assert outcome.retryable is True
    assert outcome.diagnostic == f"OperationalError[{error_code}]"
    assert legacy_module.is_retryable_priority_error_diagnostic(outcome.diagnostic)


@pytest.mark.parametrize("error_code", [1044, 1045, 1049, 1142, 1143, 1227, 1251])
def test_priority_deterministic_operational_error_codes_are_not_retryable(
    tmp_path, error_code
):
    job = _prepare(tmp_path)
    _append_submission(job.batch_state_path, task_id="id-1", task_number="no-1")

    def prioritize(*_args):
        raise OperationalError(error_code, "message is not used for classification")

    outcome = _adapter(priority_runner=prioritize).prioritize(
        job,
        task_id="id-1",
        task_number="no-1",
        timeout_seconds=12,
    )

    assert outcome.status is PriorityStatus.UNCONFIRMED
    assert outcome.retryable is False
    assert outcome.diagnostic == "OperationalError"
    assert not legacy_module.is_retryable_priority_error_diagnostic(
        outcome.diagnostic
    )


def test_priority_identity_conflict_is_not_retryable(tmp_path):
    job = _prepare(tmp_path)
    _append_event(
        job.batch_state_path,
        "CREATED",
        task_id="id-1",
        task_number="no-1",
        expected_subtask_count=1,
    )
    _append_event(
        job.batch_state_path,
        "ENABLED",
        task_id="id-1",
        task_number="no-1",
    )

    outcome = _adapter(
        priority_runner=lambda *_args: pytest.fail(
            "identity conflict must fail before priority infrastructure"
        )
    ).prioritize(
        job,
        task_id="different-id",
        task_number="no-1",
        timeout_seconds=12,
    )

    assert outcome.status is PriorityStatus.UNCONFIRMED
    assert outcome.retryable is False


def test_priority_error_class_name_alone_does_not_make_failure_retryable(tmp_path):
    class InterfaceError(Exception):
        pass

    job = _prepare(tmp_path)
    _append_submission(job.batch_state_path, task_id="id-1", task_number="no-1")

    def prioritize(*_args):
        raise InterfaceError("not the priority backend DB-API type")

    outcome = _adapter(priority_runner=prioritize).prioritize(
        job,
        task_id="id-1",
        task_number="no-1",
        timeout_seconds=12,
    )

    assert outcome.status is PriorityStatus.UNCONFIRMED
    assert outcome.retryable is False


def test_durable_retryable_priority_failure_is_retried_without_explicit_override(
    tmp_path,
):
    calls = []
    job = _prepare(tmp_path)
    _append_submission(job.batch_state_path, task_id="id-1", task_number="no-1")
    _append_event(
        job.batch_state_path,
        "PRIORITY_UNCONFIRMED",
        task_id="id-1",
        task_number="no-1",
        reason="InterfaceError",
        retryable=True,
    )

    def prioritize(prepared, task_id, task_number, _timeout_seconds):
        calls.append((task_id, task_number))
        _append_event(
            prepared.batch_state_path,
            "PRIORITY_UPDATED",
            task_id=task_id,
            task_number=task_number,
            affected_rows=1,
        )

    outcome = _adapter(priority_runner=prioritize).prioritize(
        job,
        task_id="id-1",
        task_number="no-1",
        timeout_seconds=12,
    )

    assert outcome.status is PriorityStatus.UPDATED
    assert calls == [("id-1", "no-1")]


def test_production_priority_runner_uses_exact_identity_and_persists_verified_event(tmp_path, monkeypatch):
    db_env = tmp_path / "ZHAOSHANG_DB.env"
    db_env.write_text(
        "ZHAOSHANG_DB_HOST=host\n"
        "ZHAOSHANG_DB_PORT=3306\n"
        "ZHAOSHANG_DB_USER=user\n"
        "ZHAOSHANG_DB_PASSWORD=password\n",
        encoding="utf-8",
    )
    for name in ("ZHAOSHANG_DB_HOST", "ZHAOSHANG_DB_PORT", "ZHAOSHANG_DB_USER", "ZHAOSHANG_DB_PASSWORD"):
        monkeypatch.delenv(name, raising=False)
    observed = {}

    class Connection:
        _read_timeout = 30

        def close(self):
            observed["read_timeout_seconds"] = self._read_timeout
            observed["closed"] = True

    class AutoModule:
        @staticmethod
        def priority_mall_code_for_key(mall_key):
            observed["mall_key"] = mall_key
            return "jd"

        @staticmethod
        def record_batch_state_event(path, **event):
            _append_event(
                Path(path),
                event.pop("action"),
                task_id=event.pop("task_id"),
                task_number=event.pop("task_number"),
                **event,
            )

    class PriorityModule:
        class PriorityTarget:
            def __init__(self, **identity):
                observed["identity"] = identity
                self.identity = identity

        @staticmethod
        def connect_from_env():
            observed["db_keys"] = {name for name in os.environ if name.startswith("ZHAOSHANG_DB_")}
            return Connection()

        @staticmethod
        def boost_zhaoshang_task_priority(_connection, target, **kwargs):
            observed["target"] = target
            observed["kwargs"] = kwargs
            return {"priority_state": "updated", "affected_rows": 1}

    job = _prepare(tmp_path)
    job = replace(job, task_name="cnnc_current_opt_京东_68577368911017")
    _append_event(
        job.batch_state_path,
        "CREATED",
        task_name=job.task_name,
        task_id="id-1",
        task_number="no-1",
        expected_subtask_count=1,
        timestamp="2026-08-12T16:59:05",
    )
    _append_event(
        job.batch_state_path,
        "ENABLED",
        task_name=job.task_name,
        task_id="id-1",
        task_number="no-1",
    )
    adapter = _adapter(external_tool_path=tmp_path / "auto_crawl_task.py", db_env_path=db_env)
    adapter._external_tool_path.write_text("# test tool\n", encoding="utf-8")
    monkeypatch.setattr(
        adapter,
        "_load_module",
        lambda path, _name: AutoModule if path.name == "auto_crawl_task.py" else PriorityModule,
    )

    outcome = adapter.prioritize(job, task_id="id-1", task_number="no-1", timeout_seconds=180)

    assert outcome.status is PriorityStatus.UPDATED
    identity = observed["identity"]
    assert {
        key: identity[key]
        for key in ("task_name", "task_number", "mall_code", "marker_time", "window_start")
    } == {
        "task_name": job.task_name,
        "task_number": "no-1",
        "mall_code": "jd",
        "marker_time": datetime(2026, 8, 12, 16, 59, 5),
        "window_start": datetime(2026, 8, 12, 16, 57, 5),
    }
    # The exact identity is unchanged; only the create-time upper bound
    # stretches to now so a late-materialized subtask is still found on a
    # re-check instead of being permanently unreachable after marker+15min.
    assert identity["window_end"] >= datetime(2026, 8, 12, 17, 14, 5)
    assert identity["window_end"] >= datetime.now()
    assert observed["kwargs"]["expected_subtask_count"] == 1
    assert observed["read_timeout_seconds"] == 180
    assert observed["closed"] is True
    assert not any(name.startswith("ZHAOSHANG_DB_") for name in os.environ)


def test_production_priority_failure_persists_unconfirmed_and_is_not_retried(tmp_path, monkeypatch):
    db_env = tmp_path / "ZHAOSHANG_DB.env"
    db_env.write_text(
        "ZHAOSHANG_DB_HOST=host\n"
        "ZHAOSHANG_DB_PORT=3306\n"
        "ZHAOSHANG_DB_USER=user\n"
        "ZHAOSHANG_DB_PASSWORD=password\n",
        encoding="utf-8",
    )
    calls = []
    observed = {}

    class Connection:
        _read_timeout = 30

        def close(self):
            observed["read_timeout_seconds"] = self._read_timeout

    class AutoModule:
        @staticmethod
        def priority_mall_code_for_key(_mall_key):
            return "jd"

        @staticmethod
        def record_batch_state_event(path, **event):
            _append_event(
                Path(path),
                event.pop("action"),
                task_id=event.pop("task_id"),
                task_number=event.pop("task_number"),
                **event,
            )

    class PriorityModule:
        class PriorityTarget:
            def __init__(self, **identity):
                self.identity = identity

        @staticmethod
        def connect_from_env():
            return Connection()

        @staticmethod
        def boost_zhaoshang_task_priority(*_args, **_kwargs):
            calls.append("priority")
            raise RuntimeError("verification failed")

    job = _prepare(tmp_path)
    _append_event(
        job.batch_state_path,
        "CREATED",
        task_id="id-1",
        task_number="no-1",
        expected_subtask_count=1,
        timestamp="2026-08-12T16:59:05",
    )
    _append_event(job.batch_state_path, "ENABLED", task_id="id-1", task_number="no-1")
    adapter = _adapter(external_tool_path=tmp_path / "auto_crawl_task.py", db_env_path=db_env)
    adapter._external_tool_path.write_text("# test tool\n", encoding="utf-8")
    monkeypatch.setattr(
        adapter,
        "_load_module",
        lambda path, _name: AutoModule if path.name == "auto_crawl_task.py" else PriorityModule,
    )

    first = adapter.prioritize(job, task_id="id-1", task_number="no-1", timeout_seconds=12)
    second = adapter.prioritize(job, task_id="id-1", task_number="no-1", timeout_seconds=12)

    assert first.status is PriorityStatus.UNCONFIRMED
    assert second.status is PriorityStatus.UNCONFIRMED
    assert calls == ["priority"]
    assert observed["read_timeout_seconds"] == 120
    events = adapter._read_batch_events(job)
    evidence = [event for event in events if event.get("action") == "PRIORITY_UNCONFIRMED"]
    assert len(evidence) == 1
    assert evidence[0]["reason"] == "RuntimeError"
    assert evidence[0]["retryable"] is False


def test_production_priority_interface_error_persists_retryable_and_retries(
    tmp_path, monkeypatch
):
    db_env = tmp_path / "ZHAOSHANG_DB.env"
    db_env.write_text(
        "ZHAOSHANG_DB_HOST=host\n"
        "ZHAOSHANG_DB_PORT=3306\n"
        "ZHAOSHANG_DB_USER=user\n"
        "ZHAOSHANG_DB_PASSWORD=password\n",
        encoding="utf-8",
    )
    calls = []

    class Connection:
        _read_timeout = 30

        def close(self):
            pass

    class AutoModule:
        @staticmethod
        def priority_mall_code_for_key(_mall_key):
            return "jd"

        @staticmethod
        def record_batch_state_event(path, **event):
            _append_event(
                Path(path),
                event.pop("action"),
                task_id=event.pop("task_id"),
                task_number=event.pop("task_number"),
                **event,
            )

    class PriorityModule:
        class PriorityTarget:
            def __init__(self, **identity):
                self.identity = identity

        @staticmethod
        def connect_from_env():
            return Connection()

        @staticmethod
        def boost_zhaoshang_task_priority(*_args, **_kwargs):
            calls.append("priority")
            if len(calls) == 1:
                raise InterfaceError("database connection interrupted")
            return {"priority_state": "updated", "affected_rows": 1}

    job = _prepare(tmp_path)
    _append_event(
        job.batch_state_path,
        "CREATED",
        task_id="id-1",
        task_number="no-1",
        expected_subtask_count=1,
        timestamp="2026-08-12T16:59:05",
    )
    _append_event(job.batch_state_path, "ENABLED", task_id="id-1", task_number="no-1")
    adapter = _adapter(
        external_tool_path=tmp_path / "auto_crawl_task.py",
        db_env_path=db_env,
    )
    adapter._external_tool_path.write_text("# test tool\n", encoding="utf-8")
    monkeypatch.setattr(
        adapter,
        "_load_module",
        lambda path, _name: AutoModule
        if path.name == "auto_crawl_task.py"
        else PriorityModule,
    )

    first = adapter.prioritize(
        job, task_id="id-1", task_number="no-1", timeout_seconds=12
    )
    second = adapter.prioritize(
        job, task_id="id-1", task_number="no-1", timeout_seconds=12
    )

    assert first.status is PriorityStatus.UNCONFIRMED
    assert first.retryable is True
    assert second.status is PriorityStatus.UPDATED
    assert calls == ["priority", "priority"]
    retryable_evidence = [
        event
        for event in adapter._read_batch_events(job)
        if event.get("action") == "PRIORITY_UNCONFIRMED"
    ]
    assert retryable_evidence[0]["reason"] == "InterfaceError"
    assert retryable_evidence[0]["retryable"] is True


def test_explicit_priority_retry_can_supersede_durable_unconfirmed_evidence(tmp_path):
    calls = []
    job = _prepare(tmp_path)
    _append_submission(job.batch_state_path, task_id="id-1", task_number="no-1")
    _append_event(
        job.batch_state_path,
        "PRIORITY_UNCONFIRMED",
        task_id="id-1",
        task_number="no-1",
        reason="invalid_task_name_time_marker",
    )

    def prioritize(prepared, task_id, task_number, _timeout_seconds):
        calls.append((task_id, task_number))
        _append_event(
            prepared.batch_state_path,
            "PRIORITY_UPDATED",
            task_id=task_id,
            task_number=task_number,
            affected_rows=1,
        )

    adapter = _adapter(priority_runner=prioritize)
    blocked = adapter.prioritize(job, task_id="id-1", task_number="no-1", timeout_seconds=12)
    recovered = adapter.prioritize(
        job,
        task_id="id-1",
        task_number="no-1",
        timeout_seconds=12,
        retry_unconfirmed=True,
    )

    assert blocked.status is PriorityStatus.UNCONFIRMED
    assert recovered.status is PriorityStatus.UPDATED
    assert recovered.weight_type == -1
    assert calls == [("id-1", "no-1")]


def test_reconcile_unknown_rejects_priority_evidence_for_another_remote_identity(tmp_path):
    job = _prepare(tmp_path)
    _append_submission(job.batch_state_path)
    _append_event(job.batch_state_path, "PRIORITY_ALREADY_HIGHEST", already_highest_count=1)
    adapter = _adapter(
        reconciler=lambda *_args: {"taskId": "different-id", "taskNumber": "different-number"}
    )

    outcome = adapter.reconcile_unknown(job, timeout_seconds=12)

    assert outcome.status is ReconcileStatus.MANUAL


def test_transient_unknown_reconciliation_stays_recoverable(tmp_path):
    job = _prepare(tmp_path)
    adapter = _adapter(reconciler=lambda *_args: (_ for _ in ()).throw(TimeoutError("temporary")))

    outcome = adapter.reconcile_unknown(job, timeout_seconds=12)

    assert outcome.status is ReconcileStatus.NOT_FOUND


def test_harvest_timeout_is_running_and_uses_wait_export_only(tmp_path):
    calls = []

    def runner(args, _timeout, _env):
        calls.append(args)
        raise subprocess.TimeoutExpired(args, 15)

    outcome = _adapter(runner=runner).harvest(_prepare(tmp_path), timeout_seconds=15)

    assert outcome.status is HarvestStatus.RUNNING
    assert len(calls) == 1
    assert "--wait-export-only" in calls[0]
    assert "--create-only" not in calls[0]


def test_harvest_reuses_valid_export_without_runner(tmp_path):
    job = _prepare(tmp_path)
    _write_result(job.result_path, [["sku-1", "https://item.jd.com/a", "12.5"]])
    calls = []

    outcome = _adapter(runner=lambda *_args: calls.append(1)).harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.reused_existing is True
    assert outcome.result_map["https://item.jd.com/a"]["page_price_value"] == 12.5
    assert calls == []


def test_harvest_nonzero_reuses_fresh_export_only_after_formal_validation(tmp_path):
    job = _prepare(tmp_path)

    def runner(args, _timeout, _env):
        _write_result(
            job.result_path,
            [["sku-1", "https://item.jd.com/a", "12.5"]],
        )
        return subprocess.CompletedProcess(args, returncode=1)

    outcome = _adapter(runner=runner).harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.reused_existing is True
    assert outcome.result_map["https://item.jd.com/a"]["page_price_value"] == 12.5


def test_harvest_nonzero_without_result_keeps_running_for_retry(tmp_path):
    # A child exit without any workbook is a transient execution failure
    # against a confirmed remote identity; a bounded wait-export retry may
    # still finish it, so it must not strand the job in manual review.
    job = _prepare(tmp_path)
    outcome = _adapter(
        runner=lambda args, _timeout, _env: subprocess.CompletedProcess(
            args, returncode=1
        )
    ).harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.RUNNING


def test_harvest_nonzero_with_invalid_result_keeps_manual_failure(tmp_path):
    job = _prepare(tmp_path)

    def runner(args, _timeout, _env):
        _write_result(
            job.result_path,
            [["wrong-sku", "https://item.jd.com/a", "12.5"]],
        )
        return subprocess.CompletedProcess(args, returncode=1)

    outcome = _adapter(runner=runner).harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.MANUAL


def test_harvest_zero_return_keeps_normal_export_path(tmp_path):
    job = _prepare(tmp_path)

    def runner(args, _timeout, _env):
        _write_result(
            job.result_path,
            [["sku-1", "https://item.jd.com/a", "12.5"]],
        )
        return subprocess.CompletedProcess(args, returncode=0)

    outcome = _adapter(runner=runner).harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.result_map["https://item.jd.com/a"]["page_price_value"] == 12.5


def test_harvest_timeout_reuses_fresh_export_only_after_formal_validation(tmp_path):
    job = _prepare(tmp_path)

    def runner(args, _timeout, _env):
        _write_result(
            job.result_path,
            [["sku-1", "https://item.jd.com/a", "12.5"]],
        )
        raise subprocess.TimeoutExpired(args, 15)

    outcome = _adapter(runner=runner).harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.reused_existing is True
    assert outcome.result_map["https://item.jd.com/a"]["page_price_value"] == 12.5


def test_harvest_preserves_two_skus_sharing_one_url_as_distinct_results(tmp_path):
    job = _prepare(tmp_path, requests=(_request("sku-1"), _request("sku-2")))
    _write_result(
        job.result_path,
        [["sku-1", "https://item.jd.com/a", "12.5"], ["sku-2", "https://item.jd.com/a", "19.5"]],
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.result_by_identity is not None
    assert outcome.result_by_identity[("sku-1", "https://item.jd.com/a")]["page_price_value"] == 12.5
    assert outcome.result_by_identity[("sku-2", "https://item.jd.com/a")]["page_price_value"] == 19.5


def test_harvest_allows_unique_sku_blank_request_url_when_export_supplies_url(tmp_path):
    job = _prepare(tmp_path, requests=(_request("sku-blank", ""),))
    _write_result(job.result_path, [["sku-blank", "https://item.jd.com/discovered", "12.5"]])

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.result_by_identity is not None
    assert outcome.result_by_identity[("sku-blank", "")]["page_price_value"] == 12.5


def test_harvest_rejects_mismatched_existing_export(tmp_path):
    job = _prepare(tmp_path, requests=(_request("sku-1"), _request("sku-2")))
    _write_result(job.result_path, [["sku-1", "https://item.jd.com/a", "12.5"]])
    calls = []

    outcome = _adapter(runner=lambda *_args: calls.append(1)).harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.FAILED
    assert calls == []


def test_harvest_filters_foreign_batch_before_exact_identity_validation(tmp_path):
    requests = tuple(
        _request(f"sku-{index}", f"https://item.jd.com/{index}")
        for index in range(35)
    )
    job = _prepare(tmp_path, requests=requests)
    current_rows = [
        [
            request.sku_id,
            request.url,
            str(index + 1),
            f"execution-{index}",
            job.task_name,
            job.mall_key,
        ]
        for index, request in enumerate(requests)
    ]
    foreign = [
        requests[0].sku_id,
        requests[0].url,
        "0.01",
        "RW-lookalike-foreign",
        "reference-crawl-other-batch",
        job.mall_key,
    ]
    _write_provenance_result(job.result_path, [*current_rows, foreign])

    outcome = _adapter().harvest_existing(job)

    assert outcome.status is HarvestStatus.EXPORTED
    assert len(outcome.result_by_identity or {}) == 35
    assert outcome.validation_summary == {
        "expected": 35,
        "matched": 35,
        "missing": 0,
        "duplicate": 0,
        "unexpected": 0,
        "foreign_filtered": 1,
    }


def test_harvest_rejects_duplicate_identity_inside_current_batch(tmp_path):
    job = _prepare(tmp_path)
    rows = [
        ["sku-1", "https://item.jd.com/a", "12.5", "execution-1", job.task_name, job.mall_key],
        ["sku-1", "https://item.jd.com/a", "13.5", "execution-2", job.task_name, job.mall_key],
    ]
    _write_provenance_result(job.result_path, rows)

    outcome = _adapter().harvest_existing(job)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_marks_no_result_when_provenance_filter_leaves_one_request(tmp_path):
    requests = tuple(
        _request(f"sku-{index}", f"https://item.jd.com/{index}")
        for index in range(35)
    )
    job = _prepare(tmp_path, requests=requests)
    _append_submission(job.batch_state_path, expected_subtask_count=35)
    _append_event(job.batch_state_path, "EXPORTED", output_path=str(job.result_path.resolve()))
    rows = [
            [request.sku_id, request.url, "12.5", f"no-default-execution-{index}", job.task_name, job.mall_key]
        for index, request in enumerate(requests[:-1])
    ]
    rows.append(
        [requests[-1].sku_id, requests[-1].url, "0.01", "foreign", "other-batch", job.mall_key]
    )
    _write_provenance_result(job.result_path, rows)

    outcome = _adapter().harvest_existing(job)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.no_result_identities == frozenset(
        {(requests[-1].sku_id, requests[-1].url)}
    )


def test_harvest_rejects_when_current_batch_has_one_unexpected_extra_row(tmp_path):
    requests = tuple(
        _request(f"sku-{index}", f"https://item.jd.com/{index}")
        for index in range(35)
    )
    job = _prepare(tmp_path, requests=requests)
    rows = [
        [request.sku_id, request.url, "12.5", f"execution-{index}", job.task_name, job.mall_key]
        for index, request in enumerate(requests)
    ]
    rows.append(
        ["unexpected", "https://item.jd.com/unexpected", "1", "execution-extra", job.task_name, job.mall_key]
    )
    _write_provenance_result(job.result_path, rows)

    outcome = _adapter().harvest_existing(job)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_rejects_export_without_required_provenance_columns(tmp_path):
    job = _prepare(tmp_path)
    _write_result_with_headers(
        job.result_path,
        ["商品编码", "原始链接", "商品价格"],
        [["sku-1", "https://item.jd.com/a", "12.5"]],
        include_provenance=False,
    )

    outcome = _adapter().harvest_existing(job)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_rejects_current_batch_row_without_execution_identifier(tmp_path):
    job = _prepare(tmp_path)
    _write_provenance_result(
        job.result_path,
        [["sku-1", "https://item.jd.com/a", "12.5", "", job.task_name, job.mall_key]],
    )

    outcome = _adapter().harvest_existing(job)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_rejects_foreign_batch_row_without_execution_identifier(tmp_path):
    job = _prepare(tmp_path)
    _write_provenance_result(
        job.result_path,
        [
            ["sku-1", "https://item.jd.com/a", "12.5", "execution-1", job.task_name, job.mall_key],
            ["foreign", "https://item.jd.com/foreign", "99", "", "other-batch", job.mall_key],
        ],
    )

    outcome = _adapter().harvest_existing(job)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_rejects_current_batch_row_from_another_mall(tmp_path):
    job = _prepare(tmp_path)
    _write_provenance_result(
        job.result_path,
        [["sku-1", "https://item.jd.com/a", "12.5", "execution-1", job.task_name, "other-mall"]],
    )

    outcome = _adapter().harvest_existing(job)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_rejects_duplicate_execution_inside_current_batch(tmp_path):
    requests = (
        _request("sku-1", "https://item.jd.com/a"),
        _request("sku-2", "https://item.jd.com/b"),
    )
    job = _prepare(tmp_path, requests=requests)
    _write_provenance_result(
        job.result_path,
        [
            ["sku-1", "https://item.jd.com/a", "12.5", "same-execution", job.task_name, job.mall_key],
            ["sku-2", "https://item.jd.com/b", "13.5", "same-execution", job.task_name, job.mall_key],
        ],
    )

    outcome = _adapter().harvest_existing(job)

    assert outcome.status is HarvestStatus.FAILED


def test_partial_export_preserves_page_product_fields_for_customer_prewarm(tmp_path):
    url = "https://item.jd.com/100311804305.html"
    placeholder = "zhpq-8c63a56daa7ac188"
    job = _prepare(tmp_path, requests=(_request(placeholder, url),))
    _append_submission(job.batch_state_path, expected_subtask_count=1)
    _write_result_with_headers(
        job.result_path,
        ["商品编码", "商品最小品类单元id", "商品标题/名称", "商品价格", "评论", "商品详情地址"],
        [[placeholder, "100311804305", "亿联 MeetingEye 500", 63420.5, 321, url]],
    )

    outcome = _adapter().classify_existing_export_partial(job)

    result = outcome.exact_matches[(placeholder, url)]
    assert result["page_sku_id"] == "100311804305"
    assert result["page_title"] == "亿联 MeetingEye 500"
    assert result["comment_count"] == 321
    assert result["comment_count_raw"] == "321"
    assert result["channel"] == "jd"


def test_harvest_selects_the_only_sku_column_matching_durable_identities(tmp_path):
    job = _prepare(tmp_path, requests=(_request("11858822"),))
    _write_result_with_headers(
        job.result_path,
        ["商品编码", "商品最小品类单元id", "原始链接", "商品价格"],
        [["2541608", "11858822", "https://item.jd.com/a", "104"]],
    )
    calls = []

    outcome = _adapter(runner=lambda *_args: calls.append(1)).harvest(
        job, timeout_seconds=15
    )

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.reused_existing is True
    assert outcome.result_by_identity is not None
    assert set(outcome.result_by_identity) == {
        ("11858822", "https://item.jd.com/a")
    }
    assert calls == []


def test_harvest_rejects_when_no_sku_candidate_column_matches(tmp_path):
    job = _prepare(tmp_path, requests=(_request("sku-1"),))
    _write_result_with_headers(
        job.result_path,
        ["商品编码", "商品最小品类单元id", "原始链接", "商品价格"],
        [["wrong-a", "wrong-b", "https://item.jd.com/a", "12.5"]],
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_accepts_rowwise_identical_matching_sku_candidate_columns(tmp_path):
    job = _prepare(
        tmp_path,
        requests=(
            _request("sku-1", "https://item.jd.com/a"),
            _request("sku-2", "https://item.jd.com/b"),
            _request("sku-3", "https://item.jd.com/c"),
        ),
    )
    _write_result_with_headers(
        job.result_path,
        ["商品编码", "商品最小品类单元id", "原始链接", "商品价格"],
        [
            ["sku-1", "sku-1", "https://item.jd.com/a", "12.5"],
            ["sku-2", "sku-2", "https://item.jd.com/b", "13.5"],
            ["sku-3", "sku-3", "https://item.jd.com/c", "14.5"],
        ],
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.result_by_identity is not None
    assert set(outcome.result_by_identity) == {
        ("sku-1", "https://item.jd.com/a"),
        ("sku-2", "https://item.jd.com/b"),
        ("sku-3", "https://item.jd.com/c"),
    }


def test_harvest_accepts_distinct_exported_sku_aliases_for_one_result(tmp_path):
    shared_url = "https://item.jd.com/a"
    job = _prepare(
        tmp_path,
        requests=(_request("sku-a", shared_url), _request("sku-b", shared_url)),
    )
    _append_submission(job.batch_state_path, expected_subtask_count=2)
    _append_event(
        job.batch_state_path,
        "EXPORTED",
        output_path=str(job.result_path.resolve()),
    )
    _write_result_with_headers(
        job.result_path,
        ["商品编码", "商品最小品类单元id", "原始链接", "商品价格"],
        [["sku-a", "sku-b", shared_url, "12.5"]],
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert set(outcome.result_by_identity or {}) == {
        ("sku-a", shared_url),
        ("sku-b", shared_url),
    }


def test_harvest_does_not_collapse_sku_candidate_vector_with_null(tmp_path):
    job = _prepare(tmp_path, requests=(_request("sku-a"),))
    _write_result_with_headers(
        job.result_path,
        ["商品编码", "商品最小品类单元id", "原始链接", "商品价格"],
        [["sku-a", None, "https://item.jd.com/a", "12.5"]],
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_does_not_collapse_reordered_sku_candidate_vector(tmp_path):
    job = _prepare(
        tmp_path,
        requests=(_request("sku-a", ""), _request("sku-b", "")),
    )
    _write_result_with_headers(
        job.result_path,
        ["商品编码", "商品最小品类单元id", "原始链接", "商品价格"],
        [
            ["sku-a", "sku-b", "", "12.5"],
            ["sku-b", "sku-a", "", "13.5"],
        ],
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_accepts_all_exported_sku_alias_columns(tmp_path):
    shared_url = "https://item.jd.com/a"
    job = _prepare(
        tmp_path,
        requests=(_request("sku-a", shared_url), _request("sku-b", shared_url)),
    )
    _append_submission(job.batch_state_path, expected_subtask_count=2)
    _append_event(
        job.batch_state_path,
        "EXPORTED",
        output_path=str(job.result_path.resolve()),
    )
    _write_result_with_headers(
        job.result_path,
        ["商品编码", "skuId", "商品最小品类单元id", "原始链接", "商品价格"],
        [["sku-a", "sku-a", "sku-b", shared_url, "12.5"]],
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert set(outcome.result_by_identity or {}) == {
        ("sku-a", shared_url),
        ("sku-b", shared_url),
    }


def test_harvest_keeps_exact_url_validation_after_sku_candidate_collapse(tmp_path):
    job = _prepare(
        tmp_path,
        requests=(_request("sku-a", "https://item.jd.com/expected"),),
    )
    _write_result_with_headers(
        job.result_path,
        ["商品编码", "商品最小品类单元id", "原始链接", "商品价格"],
        [["sku-a", "sku-a", "https://item.jd.com/other", "12.5"]],
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_rejects_sku_candidate_that_conflicts_with_another_exact_url(tmp_path):
    job = _prepare(
        tmp_path,
        requests=(
            _request("sku-1", "https://item.jd.com/a"),
            _request("sku-2", "https://item.jd.com/b"),
        ),
    )
    _write_result(
        job.result_path,
        [
            ["sku-1", "https://item.jd.com/b", "12.5"],
            ["sku-2", "https://item.jd.com/a", "13.5"],
        ],
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_rejects_matching_sku_with_unrelated_export_url(tmp_path):
    job = _prepare(tmp_path, requests=(_request("sku-1", "https://item.jd.com/a"),))
    _write_result(
        job.result_path,
        [["sku-1", "https://unrelated.example/not-the-request", "12.5"]],
    )
    _append_event(
        job.batch_state_path,
        "CREATED",
        mall="qstbg",
        expected_subtask_count=1,
    )
    _append_event(job.batch_state_path, "ENABLED", mall="qstbg")
    _append_event(
        job.batch_state_path,
        "EXPORTED",
        mall="qstbg",
        output_path=str(job.result_path),
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.FAILED


def test_harvest_classifies_partial_export_with_exact_export_evidence(tmp_path):
    job = _prepare(tmp_path, requests=(_request("sku-1"), _request("sku-2")))
    _write_result(job.result_path, [["sku-1", "https://item.jd.com/a", "12.5"]])
    _append_submission(job.batch_state_path, expected_subtask_count=2)
    _append_event(
        job.batch_state_path,
        "EXPORTED",
        output_path=str(job.result_path),
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.no_result_identities == frozenset(
        {("sku-2", "https://item.jd.com/a")}
    )


def test_harvest_maps_canonicalized_export_url_by_unique_sku(tmp_path):
    durable_url = "https://www.lxwlmall.com/detail/inner?skuNo=sku-1"
    exported_url = "https://www.lxwlmall.com/detail?skuNo=sku-1"
    job = _prepare(tmp_path, requests=(_request("sku-1", durable_url),))
    _write_result(job.result_path, [["sku-1", exported_url, "12.5"]])
    _append_submission(job.batch_state_path)
    _append_event(
        job.batch_state_path,
        "EXPORTED",
        output_path=str(job.result_path),
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.result_by_identity is not None
    assert set(outcome.result_by_identity) == {("sku-1", durable_url)}


def test_harvest_maps_exact_qstbg_detail_redirect_by_goods_id(tmp_path):
    durable_url = "https://www.qstbg.com/gd/detail.do?goods_id=20904038"
    exported_url = (
        "https://www.qstbg.com/gd/initParameter.do?"
        "goods_id=20904038&formid=form1#title_detail"
    )
    request = replace(
        _request("20904038", durable_url),
        mall_key="qstbg",
        mall_source_id="42",
        mall_name="企事通商城",
    )
    job = _prepare(tmp_path, requests=(request,))
    _write_result(
        job.result_path,
        [["20904038", exported_url, "1596.58"]],
        task_name=job.task_name,
        mall_key=job.mall_key,
    )
    _append_event(
        job.batch_state_path,
        "CREATED",
        mall="qstbg",
        expected_subtask_count=1,
    )
    _append_event(job.batch_state_path, "ENABLED", mall="qstbg")
    _append_event(
        job.batch_state_path,
        "EXPORTED",
        mall="qstbg",
        output_path=str(job.result_path),
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.result_by_identity is not None
    assert set(outcome.result_by_identity) == {("20904038", durable_url)}


@pytest.mark.parametrize(
    "exported_url",
    [
        "https://www.qstbg.com/gd/initParameter.do?goods_id=other&formid=form1#title_detail",
        "https://www.qstbg.com/gd/initParameter.do?goods_id=20904038&formid=other#title_detail",
        "https://www.qstbg.com/gd/initParameter.do?goods_id=20904038&formid=form1&extra=1#title_detail",
        "https://qstbg.com/gd/initParameter.do?goods_id=20904038&formid=form1#title_detail",
    ],
)
def test_harvest_trusts_other_qstbg_exported_urls_after_exact_provenance(
    tmp_path, exported_url
):
    durable_url = "https://www.qstbg.com/gd/detail.do?goods_id=20904038"
    request = replace(
        _request("20904038", durable_url),
        mall_key="qstbg",
        mall_source_id="42",
        mall_name="企事通商城",
    )
    job = _prepare(tmp_path, requests=(request,))
    _write_result(
        job.result_path,
        [["20904038", exported_url, "1596.58"]],
        task_name=job.task_name,
        mall_key=job.mall_key,
    )
    _append_event(
        job.batch_state_path,
        "CREATED",
        mall="qstbg",
        expected_subtask_count=1,
    )
    _append_event(job.batch_state_path, "ENABLED", mall="qstbg")
    _append_event(
        job.batch_state_path,
        "EXPORTED",
        mall="qstbg",
        output_path=str(job.result_path),
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    result = (outcome.result_by_identity or {})[("20904038", durable_url)]
    assert result["exported_url"] == exported_url


def test_harvest_maps_exact_jd_product_host_alias(tmp_path):
    durable_url = "https://i-item.jd.com/10132026242240.html"
    exported_url = "https://item.jd.com/10132026242240.html"
    job = _prepare(tmp_path, requests=(_request("10132026242240", durable_url),))
    _write_result(job.result_path, [["10132026242240", exported_url, "0"]])
    _append_submission(job.batch_state_path)
    _append_event(
        job.batch_state_path,
        "EXPORTED",
        output_path=str(job.result_path),
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    assert outcome.result_by_identity is not None
    assert set(outcome.result_by_identity) == {("10132026242240", durable_url)}
    assert outcome.result_by_identity[("10132026242240", durable_url)][
        "page_price_status"
    ] == "PRICE_IS_ZERO"


@pytest.mark.parametrize(
    "exported_url",
    [
        "https://item.jd.com/10132026242241.html",
        "https://www.jd.com/10132026242240.html",
        "https://item.jd.com/10132026242240.html?other=1",
        "https://item.jd.com/product/10132026242240.html",
        "http://item.jd.com/10132026242240.html",
    ],
)
def test_harvest_trusts_other_jd_exported_urls_after_exact_provenance(
    tmp_path, exported_url
):
    durable_url = "https://i-item.jd.com/10132026242240.html"
    job = _prepare(tmp_path, requests=(_request("10132026242240", durable_url),))
    _write_result(job.result_path, [["10132026242240", exported_url, "0"]])
    _append_submission(job.batch_state_path)
    _append_event(
        job.batch_state_path,
        "EXPORTED",
        output_path=str(job.result_path),
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.EXPORTED
    result = (outcome.result_by_identity or {})[("10132026242240", durable_url)]
    assert result["exported_url"] == exported_url


def test_harvest_rejects_same_count_export_with_wrong_crawl_identity(tmp_path):
    job = _prepare(tmp_path, requests=(_request("sku-1"), _request("sku-2", "https://item.jd.com/b")))
    _write_result(
        job.result_path,
        [["sku-1", "https://item.jd.com/a", "12.5"], ["sku-2", "https://item.jd.com/wrong", "13"]],
    )

    outcome = _adapter().harvest(job, timeout_seconds=15)

    assert outcome.status is HarvestStatus.FAILED


def test_read_only_session_clamps_legacy_helper_timeout():
    from app.services.legacy_zhaoshang_adapter import _BoundedReadOnlySession

    calls = []

    class Session:
        def post(self, *_args, **kwargs):
            calls.append(kwargs["timeout"])
            return object()

    _BoundedReadOnlySession(Session(), 3).post("https://read-only.invalid", timeout=60)
    assert len(calls) == 1
    assert 0 < calls[0] <= 3


def test_read_only_session_uses_one_aggregate_deadline_across_requests():
    from app.services.legacy_zhaoshang_adapter import _BoundedReadOnlySession

    calls = []

    class Session:
        def post(self, *_args, **kwargs):
            calls.append(kwargs["timeout"])
            return object()

    ticks = iter((100.0, 101.0, 104.0, 111.0))
    session = _BoundedReadOnlySession(Session(), 10, monotonic=lambda: next(ticks))

    session.post("https://read-only.invalid/1", timeout=60)
    session.post("https://read-only.invalid/2", timeout=60)
    with pytest.raises(TimeoutError):
        session.post("https://read-only.invalid/3", timeout=60)

    assert calls == [pytest.approx(9), pytest.approx(6)]


class _PreSendAuth:
    def __init__(self, exc: BaseException | None = None):
        self.exc = exc

    def export_environment(self, *_args, **_kwargs):
        if self.exc is not None:
            raise self.exc
        return {"BLADE_AUTH": "blade"}


class _FakeSpawnedProcess:
    def __init__(self, *, first_communicate_raises: BaseException | None = None):
        self.first_communicate_raises = first_communicate_raises
        self.communicate_calls = 0
        self.communicate_timeouts = []
        self.killed = False
        self.returncode = 0

    def __enter__(self):
        return self

    def __exit__(self, *_exc):
        return False

    def kill(self):
        self.killed = True

    def communicate(self, input=None, timeout=None):
        self.communicate_calls += 1
        self.communicate_timeouts.append(timeout)
        if self.first_communicate_raises is not None and self.communicate_calls == 1:
            raise self.first_communicate_raises
        return ("", "")

    def poll(self):
        return self.returncode


def _patch_spawn(monkeypatch, *, process=None, raises=None):
    calls = []

    def fake_popen(*args, **kwargs):
        calls.append((args, kwargs))
        if raises is not None:
            raise raises
        return process

    monkeypatch.setattr(legacy_module.subprocess, "Popen", fake_popen)
    return calls


def _submit_with_production_runner(tmp_path, monkeypatch, *, auth=None, tool=None, monotonic=time.monotonic, batch_state=None):
    adapter = _adapter(
        auth_manager=auth if auth is not None else _PreSendAuth(),
        external_tool_path=tool if tool is not None else _tool_file(tmp_path),
        monotonic=monotonic,
    )
    job = adapter.prepare_job(
        business_date="2026-08-11",
        requests=(_request(),),
        template_path=tmp_path / "template.xlsx",
        result_path=tmp_path / "result.xlsx",
        task_name=DEFAULT_TASK_NAME,
        batch_state_path=batch_state if batch_state is not None else _state_file(tmp_path),
    )
    return adapter.submit(job, timeout_seconds=30)


def _tool_file(tmp_path):
    tool = tmp_path / "auto_crawl_task.py"
    tool.write_text("# test tool\n", encoding="utf-8")
    return tool


def test_submit_auth_refresh_failure_is_pre_send_not_unknown(tmp_path, monkeypatch):
    from scripts.auth_manager import RefreshTokenExpiredError

    popen_calls = _patch_spawn(monkeypatch)

    outcome = _submit_with_production_runner(
        tmp_path, monkeypatch, auth=_PreSendAuth(RefreshTokenExpiredError("refresh token expired"))
    )

    assert popen_calls == []
    assert outcome.status is SubmitStatus.NOT_ATTEMPTED
    assert "pre-send" in outcome.diagnostic
    assert "RefreshTokenExpiredError" in outcome.diagnostic


def test_submit_spawn_file_not_found_is_pre_send_not_unknown(tmp_path, monkeypatch):
    popen_calls = _patch_spawn(monkeypatch, raises=FileNotFoundError("no such executable"))

    outcome = _submit_with_production_runner(tmp_path, monkeypatch)

    assert len(popen_calls) == 1
    assert outcome.status is SubmitStatus.NOT_ATTEMPTED
    assert "could not be started" in outcome.diagnostic


def test_submit_spawn_oserror_is_pre_send_not_unknown(tmp_path, monkeypatch):
    popen_calls = _patch_spawn(monkeypatch, raises=PermissionError("spawn denied"))

    outcome = _submit_with_production_runner(tmp_path, monkeypatch)

    assert len(popen_calls) == 1
    assert outcome.status is SubmitStatus.NOT_ATTEMPTED
    assert "could not be started" in outcome.diagnostic


def test_submit_missing_external_tool_is_pre_send_not_unknown(tmp_path, monkeypatch):
    popen_calls = _patch_spawn(monkeypatch)
    missing_tool = tmp_path / "missing" / "auto_crawl_task.py"

    outcome = _submit_with_production_runner(tmp_path, monkeypatch, tool=missing_tool)

    assert popen_calls == []
    assert outcome.status is SubmitStatus.NOT_ATTEMPTED
    assert "external tool is unavailable" in outcome.diagnostic


def test_submit_deadline_exhausted_before_spawn_is_pre_send_not_unknown(tmp_path, monkeypatch):
    popen_calls = _patch_spawn(monkeypatch)
    ticks = iter((100.0, 200.0))

    outcome = _submit_with_production_runner(
        tmp_path, monkeypatch, monotonic=lambda: next(ticks)
    )

    assert popen_calls == []
    assert outcome.status is SubmitStatus.NOT_ATTEMPTED
    assert outcome.diagnostic


def test_submit_timeout_after_successful_spawn_stays_unknown(tmp_path, monkeypatch):
    process = _FakeSpawnedProcess(first_communicate_raises=subprocess.TimeoutExpired("cmd", 1))
    popen_calls = _patch_spawn(monkeypatch, process=process)

    outcome = _submit_with_production_runner(tmp_path, monkeypatch)

    assert len(popen_calls) == 1
    assert process.killed is True
    assert outcome.status is SubmitStatus.SUBMIT_UNKNOWN
    assert outcome.diagnostic == "legacy create-only completion is unknown"


def test_submit_injected_runner_oserror_stays_unknown(tmp_path):
    def runner(args, _timeout, _env):
        raise FileNotFoundError("runner-level missing file")

    outcome = _adapter(runner=runner).submit(_prepare(tmp_path), timeout_seconds=30)

    assert outcome.status is SubmitStatus.SUBMIT_UNKNOWN


def test_submit_child_died_after_possible_post_stays_unknown(tmp_path, monkeypatch):
    process = _FakeSpawnedProcess()
    process.returncode = 1
    _patch_spawn(monkeypatch, process=process)

    outcome = _submit_with_production_runner(tmp_path, monkeypatch)

    assert outcome.status is SubmitStatus.SUBMIT_UNKNOWN
    assert "nonzero exit" in outcome.diagnostic


def test_submit_child_completed_without_evidence_stays_unknown(tmp_path, monkeypatch):
    process = _FakeSpawnedProcess()
    process.returncode = 0
    _patch_spawn(monkeypatch, process=process)

    outcome = _submit_with_production_runner(tmp_path, monkeypatch)

    assert outcome.status is SubmitStatus.SUBMIT_UNKNOWN
    assert "lacks CREATED + ENABLED evidence" in outcome.diagnostic


def test_submit_existing_created_evidence_wins_over_pre_send_failure(tmp_path, monkeypatch):
    from scripts.auth_manager import RefreshTokenExpiredError

    popen_calls = _patch_spawn(monkeypatch)
    batch_state = _state_file(tmp_path)
    _append_submission(batch_state, task_id="id-9", task_number="no-9")

    outcome = _submit_with_production_runner(
        tmp_path,
        monkeypatch,
        auth=_PreSendAuth(RefreshTokenExpiredError("refresh token expired")),
        batch_state=batch_state,
    )

    assert popen_calls == []
    assert outcome.status is SubmitStatus.SUBMITTED
    assert outcome.task_id == "id-9"
    assert outcome.task_number == "no-9"
