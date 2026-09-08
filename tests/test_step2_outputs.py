"""STEP2 一次显式调用产出三表 + workday manifest 的端到端回归（全 fake，不触真实 API）。"""
import hashlib
import importlib.util
import json
import re
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parents[1] / "scripts"
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPT_DIR))
SPEC = importlib.util.spec_from_file_location(
    "export_current_monitor_status_step2",
    SCRIPT_DIR / "export_current_monitor_status.py",
)
assert SPEC and SPEC.loader
monitor_export = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(monitor_export)

from app.services.workday_context import WorkdayContextError


def cnnc_row(document_id, name="商品"):
    return {
        "documentId": document_id,
        "commodityName": name,
        "categoryName": "办公用品",
        "fourCategoryId": 4,
        "fourCategoryName": "四級",
        "referencePrice": 10,
        "referenceLink": "https://example.com/item",
        "materialCode": "M-1",
        "materialDesc": "desc",
        "createTime": "2026-09-08 10:00:00",
        "updateTime": "2026-09-08 10:00:00",
        "sassCheckStatus": 0,
        "quasiStandardCommodityCode": "",
    }


def saas_row(document_id, *, is_sync=1, is_standard=1, price_status=1, new_price="19.90"):
    return {
        "commodityId": str(document_id),
        "isSync": is_sync,
        "isSyncStandard": is_standard,
        "syncPriceStatus": price_status,
        "newPrice": new_price,
    }


class FakeHarness:
    """一次 main() 端到端所需的全部替身；cnnc_docs 顺序即 authoritative 顺序。"""

    def __init__(self, monkeypatch, tmp_path, cnnc_docs, saas_monitors, extra_cnnc_after_first_pass=None):
        self.tmp_path = tmp_path
        self.calls = []
        self.pass_count = 0
        self.extra = extra_cnnc_after_first_pass

        def fake_load_curl(path):
            name = Path(path).name
            if name == "latest_cnnc_query.curl":
                body = {"sassCheckStatuss": ["0"], "page": 1, "limit": 200}
            else:
                body = {"tenantId": "188605", "commodityPoolId": "222"}
            return {
                "url": f"https://fake.local/{name}",
                "headers": [("Content-Type", "application/json")],
                "cookie": "x-access-token=fake",
                "method": "POST",
                "data": json.dumps(body),
            }

        class FakeAuthManager:
            def prepare_curl(self, system, curl, refresh=False):
                return curl

        def fake_request(config, payload, **kwargs):
            self.calls.append((config["url"], dict(payload)))
            if "commodityId" in payload:
                ids = payload["commodityId"].split()
                rows = [row for row in saas_monitors if str(row["commodityId"]) in ids]
                return {"rspCode": "0000", "count": len(rows), "rows": rows}
            docs = list(cnnc_docs)
            if self.pass_count >= 1 and self.extra:
                docs = docs + self.extra  # 第二轮完整分页才会看到：证明单轮内不复用第二次数据
            self.pass_count += 1
            return {"rspCode": "0000", "count": len(docs), "rows": docs}

        monkeypatch.setattr(monitor_export, "ROOT", tmp_path)
        monkeypatch.setattr(monitor_export, "load_curl_command", fake_load_curl)
        monkeypatch.setattr(monitor_export, "AuthManager", FakeAuthManager)
        monkeypatch.setattr(monitor_export, "_request", fake_request)
        monkeypatch.setattr(monitor_export, "apply_internal_no_proxy_env", lambda: None)

    def run(self, argv):
        monitor_export.main(argv)
        return self.calls


BASIC_DOCS = [cnnc_row("A"), cnnc_row("B"), cnnc_row("C")]
BASIC_SAAS = [
    saas_row("A", is_sync=1, is_standard=1, price_status=1, new_price="10.00"),
    saas_row("B", is_sync=1, is_standard=1, price_status=5, new_price="5.00"),
    # C 无 SaaS monitor → SaaS监控缺失
]


def _step2_dir(tmp_path, business_date="2026-09-08"):
    return tmp_path / "workdays" / business_date / "03_step2"


def _rows(path, sheet):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(r) for r in book[sheet].iter_rows(values_only=True)]
    finally:
        book.close()


def _cells(row):
    """openpyxl 把空字符串写盘后读回为 None；期望行按同一规则归一后再逐字段比较。"""
    return [None if value == "" else value for value in row]


def _expected_row(document_id, monitor_values):
    return [cnnc_row(document_id)[key] for key in monitor_export.CNNC_KEYS] + monitor_values


def _read_manifest(tmp_path, business_date="2026-09-08"):
    return json.loads(
        (tmp_path / "workdays" / business_date / "manifest.json").read_text(encoding="utf-8")
    )


def _file_bytes(step2):
    return {name: (step2 / name).read_bytes() for name in monitor_export.STEP2_OUTPUT_NAMES.values()}


def test_case1_three_fixed_outputs_and_manifest(monkeypatch, tmp_path):
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run(["--business-date", "2026-09-08"])
    step2 = _step2_dir(tmp_path)
    assert (step2 / "STEP2_监控结果.xlsx").is_file()
    assert (step2 / "STEP2_可直接推送.xlsx").is_file()
    assert (step2 / "STEP2_不可直接推送.xlsx").is_file()
    for stage in ("01_step1", "02_reference_prewarm", "03_step2", "04_step3"):
        assert (tmp_path / "workdays" / "2026-09-08" / stage).is_dir()
    manifest = _read_manifest(tmp_path)
    step2_node = manifest["stages"]["step2"]
    assert step2_node["status"] == "COMPLETED"
    assert step2_node["counts"] == {"monitor_total": 3, "pushable": 1, "non_pushable": 2}
    assert step2_node["generation"] == 1
    assert re.match(r"^\d{8}T\d{6}-\d{6}$", step2_node["run_id"])
    assert step2_node["outputs"]["monitor"] == "03_step2/STEP2_监控结果.xlsx"
    for key, name in monitor_export.STEP2_OUTPUT_NAMES.items():
        info = step2_node["files"][key]
        assert info["path"] == f"03_step2/{name}"
        assert info["sha256"] == hashlib.sha256((step2 / name).read_bytes()).hexdigest()
        assert info["size"] == (step2 / name).stat().st_size
    assert manifest["stages"]["step3"] == {}


def test_case2_pushable_sheet_is_field_by_field_projection(monkeypatch, tmp_path):
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run(["--business-date", "2026-09-08"])
    row_a = _expected_row(
        "A",
        ["采集成功", "同款匹配成功", "价格计算成功", "10.00", "是", "可推送", "", "已匹配"],
    )
    rows = _rows(_step2_dir(tmp_path) / "STEP2_可直接推送.xlsx", "可直接推送")
    assert rows == [monitor_export.MONITOR_HEADERS, _cells(row_a)]


def test_case3_non_pushable_sheet_is_field_by_field_projection(monkeypatch, tmp_path):
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run(["--business-date", "2026-09-08"])
    row_b = _expected_row(
        "B",
        ["采集成功", "同款匹配成功", "价格计算失败", "5.00", "否", "价格检测校验失败", "价格检测校验失败", "已匹配"],
    )
    row_c = _expected_row(
        "C",
        ["", "", "", "", "否", "SaaS监控缺失", "中核存在但SaaS监控未返回", "缺失"],
    )
    rows = _rows(_step2_dir(tmp_path) / "STEP2_不可直接推送.xlsx", "不可直接推送")
    assert rows == [monitor_export.MONITOR_HEADERS, _cells(row_b), _cells(row_c)]


def test_case4_saas_missing_document_lands_in_non_pushable(monkeypatch, tmp_path):
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run(["--business-date", "2026-09-08"])
    rows = _rows(_step2_dir(tmp_path) / "STEP2_不可直接推送.xlsx", "不可直接推送")
    assert "C" in {row[0] for row in rows[1:]}


def test_case5_pushable_and_non_pushable_ids_do_not_intersect(monkeypatch, tmp_path):
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run(["--business-date", "2026-09-08"])
    step2 = _step2_dir(tmp_path)
    pushable_ids = {row[0] for row in _rows(step2 / "STEP2_可直接推送.xlsx", "可直接推送")[1:]}
    non_pushable_ids = {row[0] for row in _rows(step2 / "STEP2_不可直接推送.xlsx", "不可直接推送")[1:]}
    assert pushable_ids & non_pushable_ids == set()


def test_case6_projections_union_covers_monitor(monkeypatch, tmp_path):
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run(["--business-date", "2026-09-08"])
    step2 = _step2_dir(tmp_path)
    pushable_ids = {row[0] for row in _rows(step2 / "STEP2_可直接推送.xlsx", "可直接推送")[1:]}
    non_pushable_ids = {row[0] for row in _rows(step2 / "STEP2_不可直接推送.xlsx", "不可直接推送")[1:]}
    monitor_ids = {row[0] for row in _rows(step2 / "STEP2_监控结果.xlsx", "全量监控结果")[1:]}
    assert pushable_ids | non_pushable_ids == monitor_ids


def test_case7_projection_preserves_authoritative_order(monkeypatch, tmp_path):
    docs = [cnnc_row("A"), cnnc_row("B"), cnnc_row("D"), cnnc_row("C")]
    saas = [
        saas_row("A", is_sync=1, is_standard=1, price_status=1, new_price="1.00"),
        saas_row("B", is_sync=1, is_standard=1, price_status=5, new_price="5.00"),
        saas_row("D", is_sync=1, is_standard=1, price_status=1, new_price="4.00"),
    ]
    harness = FakeHarness(monkeypatch, tmp_path, docs, saas)
    harness.run(["--business-date", "2026-09-08"])
    step2 = _step2_dir(tmp_path)
    pushable_ids = [row[0] for row in _rows(step2 / "STEP2_可直接推送.xlsx", "可直接推送")[1:]]
    non_pushable_ids = [row[0] for row in _rows(step2 / "STEP2_不可直接推送.xlsx", "不可直接推送")[1:]]
    assert pushable_ids == ["A", "D"]
    assert non_pushable_ids == ["B", "C"]


def test_case8_single_snapshot_one_cnnc_pass_no_late_data(monkeypatch, tmp_path):
    harness = FakeHarness(
        monkeypatch,
        tmp_path,
        BASIC_DOCS,
        BASIC_SAAS,
        extra_cnnc_after_first_pass=[cnnc_row("X")],
    )
    harness.run(["--business-date", "2026-09-08"])
    monitor_ids = [row[0] for row in _rows(_step2_dir(tmp_path) / "STEP2_监控结果.xlsx", "全量监控结果")[1:]]
    assert monitor_ids == ["A", "B", "C"]
    assert len(harness.calls) == 2  # 1 次中核完整分页 + 3 个 id 一个 SaaS 批次
    cnnc_calls = [payload for url, payload in harness.calls if "latest_cnnc_query" in url]
    assert len(cnnc_calls) == 1


def test_explicit_business_date_overrides_default_helper(monkeypatch, tmp_path):
    monkeypatch.setattr(monitor_export, "business_date_in_shanghai", lambda: "2026-09-09")
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run(["--business-date", "2026-09-08"])
    assert (_step2_dir(tmp_path, "2026-09-08") / "STEP2_监控结果.xlsx").is_file()
    assert not (tmp_path / "workdays" / "2026-09-09").exists()
    assert _read_manifest(tmp_path, "2026-09-08")["business_date"] == "2026-09-08"


def test_default_business_date_comes_from_shanghai_helper(monkeypatch, tmp_path):
    monkeypatch.setattr(monitor_export, "business_date_in_shanghai", lambda: "2026-09-07")
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run([])
    assert (_step2_dir(tmp_path, "2026-09-07") / "STEP2_监控结果.xlsx").is_file()
    assert _read_manifest(tmp_path, "2026-09-07")["business_date"] == "2026-09-07"


def test_same_day_rerun_bumps_generation_and_atomically_replaces_outputs(monkeypatch, tmp_path):
    first = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    first.run(["--business-date", "2026-09-08"])
    step2 = _step2_dir(tmp_path)
    first_node = _read_manifest(tmp_path)["stages"]["step2"]
    first_run_id = first_node["run_id"]
    first_bytes = _file_bytes(step2)

    docs2 = BASIC_DOCS + [cnnc_row("D")]
    saas2 = BASIC_SAAS + [saas_row("D", is_sync=1, is_standard=1, price_status=1, new_price="7.00")]
    second = FakeHarness(monkeypatch, tmp_path, docs2, saas2)
    second.run(["--business-date", "2026-09-08"])

    node = _read_manifest(tmp_path)["stages"]["step2"]
    assert node["generation"] == 2
    assert node["run_id"] != first_run_id
    monitor_ids = [row[0] for row in _rows(step2 / "STEP2_监控结果.xlsx", "全量监控结果")[1:]]
    assert monitor_ids == ["A", "B", "C", "D"]  # 三文件整体为第二代内容
    assert _file_bytes(step2) != first_bytes
    for key, name in monitor_export.STEP2_OUTPUT_NAMES.items():
        assert node["files"][key]["sha256"] == hashlib.sha256((step2 / name).read_bytes()).hexdigest()
    staging = step2 / ".staging"
    assert not staging.exists() or not any(staging.iterdir())


def test_publication_failure_marks_failed_and_keeps_previous_generation(monkeypatch, tmp_path):
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run(["--business-date", "2026-09-08"])
    step2 = _step2_dir(tmp_path)
    before = _file_bytes(step2)

    def boom(records, output, sheet_name):
        raise RuntimeError("projection exploded")

    monkeypatch.setattr(monitor_export, "_write_projection_workbook", boom)
    with pytest.raises(RuntimeError, match="projection exploded"):
        monitor_export.main(["--business-date", "2026-09-08"])

    node = _read_manifest(tmp_path)["stages"]["step2"]
    assert node["status"] == "FAILED"
    assert node["error_type"] == "RuntimeError"
    assert "projection exploded" in node["error_summary"]
    assert _file_bytes(step2) == before  # 失败代不产出"看似完整"的新一代
    staging = step2 / ".staging"
    assert not staging.exists() or not any(staging.iterdir())


def test_query_failure_marks_failed_without_official_outputs(monkeypatch, tmp_path):
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)

    def exploding_request(config, payload, **kwargs):
        raise RuntimeError("network down")

    monkeypatch.setattr(monitor_export, "_request", exploding_request)
    with pytest.raises(RuntimeError, match="network down"):
        monitor_export.main(["--business-date", "2026-09-08"])

    node = _read_manifest(tmp_path)["stages"]["step2"]
    assert node["status"] == "FAILED"
    assert node["error_type"] == "RuntimeError"
    assert not list(_step2_dir(tmp_path).glob("*.xlsx"))


def test_malformed_manifest_fails_closed_without_touching_anything(monkeypatch, tmp_path):
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    manifest_path = tmp_path / "workdays" / "2026-09-08" / "manifest.json"
    manifest_path.parent.mkdir(parents=True)
    manifest_path.write_text("{ not json", encoding="utf-8")
    with pytest.raises(WorkdayContextError):
        monitor_export.main(["--business-date", "2026-09-08"])
    assert manifest_path.read_text(encoding="utf-8") == "{ not json"
    assert not list(_step2_dir(tmp_path).glob("*.xlsx"))


def test_step2_never_spawns_subprocess_or_step3(monkeypatch, tmp_path):
    def forbidden(*args, **kwargs):
        raise AssertionError("STEP2 must not spawn subprocesses")

    monkeypatch.setattr(subprocess, "run", forbidden)
    monkeypatch.setattr(subprocess, "Popen", forbidden)
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run(["--business-date", "2026-09-08"])
    for path in (tmp_path / "workdays").rglob("*"):
        assert "zhImport" not in path.name
    manifest = _read_manifest(tmp_path)
    assert manifest["stages"]["step2"]["status"] == "COMPLETED"
    assert manifest["stages"]["step3"] == {}


def test_monitor_workbook_keeps_five_sheets_and_reconciles_with_projections(monkeypatch, tmp_path):
    harness = FakeHarness(monkeypatch, tmp_path, BASIC_DOCS, BASIC_SAAS)
    harness.run(["--business-date", "2026-09-08"])
    step2 = _step2_dir(tmp_path)
    monitor_path = step2 / "STEP2_监控结果.xlsx"
    book = load_workbook(monitor_path, read_only=True, data_only=True)
    try:
        assert book.sheetnames == [
            "全量监控结果",
            "可推送数据",
            "需处理数据",
            "SaaS监控缺失",
            "状态统计",
        ]
        in_monitor_pushable = [list(r) for r in book["可推送数据"].iter_rows(values_only=True)]
        needed_ids = {row[0] for row in book["需处理数据"].iter_rows(values_only=True)} - {"文档ID"}
        missing_ids = {row[0] for row in book["SaaS监控缺失"].iter_rows(values_only=True)} - {"文档ID"}
    finally:
        book.close()
    pushable_file_rows = _rows(step2 / "STEP2_可直接推送.xlsx", "可直接推送")
    assert in_monitor_pushable == pushable_file_rows
    non_pushable_ids = {row[0] for row in _rows(step2 / "STEP2_不可直接推送.xlsx", "不可直接推送")[1:]}
    assert not (needed_ids & missing_ids)
    assert needed_ids | missing_ids == non_pushable_ids
