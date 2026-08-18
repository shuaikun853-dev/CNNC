from __future__ import annotations

import hashlib
import importlib
import json
from email.message import Message
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill


def _module():
    try:
        return importlib.import_module("scripts.approve_saas_audit_excel")
    except ModuleNotFoundError:
        pytest.fail("scripts.approve_saas_audit_excel has not been implemented")


HEADERS = [
    "SaaS商品ID",
    "文档ID",
    "客户标题",
    "SaaS商品名称",
    "SKU编号",
    "互联网标题",
    "供应商",
    "店铺名称",
    "SaaS价格",
    "页面实际价格",
    "互联网链接",
    "评论数",
    "创建时间",
    "更新时间",
    "模型审核结果",
    "是否匹配",
]


def _write_audit_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "summary"
    worksheet = workbook.create_sheet("record_results")
    worksheet.append(HEADERS)
    worksheet.append(
        ["101", None, None, None, "sku-clear", "clear", None, None, None,
         None, None, None, None, None, "1", "1"]
    )
    worksheet.append(
        ["102", None, None, None, "sku-red", "red", None, None, None,
         None, None, None, None, None, "1", "1"]
    )
    worksheet.append(
        ["103", None, None, None, "sku-yellow", "yellow", None, None, None,
         None, None, None, None, None, "1", "1"]
    )
    worksheet.append(
        ["104", None, None, None, "sku-other", "other", None, None, None,
         None, None, None, None, None, "1", "1"]
    )
    worksheet.append(
        ["101", None, None, None, "sku-clear", "clear duplicate", None, None,
         None, None, None, None, None, None, "1", "1"]
    )
    worksheet["F3"].fill = PatternFill("solid", fgColor="FFFFC7CE")
    worksheet["F4"].fill = PatternFill("solid", fgColor="FFFFFF00")
    worksheet["F5"].fill = PatternFill("solid", fgColor="FF00FF00")
    workbook.save(path)


class _FakeResponse:
    def __init__(self, payload: dict) -> None:
        self._payload = json.dumps(payload).encode("utf-8")
        self.status = 200
        self.headers = Message()
        self.headers["Content-Type"] = "application/json; charset=utf-8"

    def read(self) -> bytes:
        return self._payload

    def __enter__(self):
        return self

    def __exit__(self, *_args: object) -> None:
        return None


class _FakeOpener:
    def __init__(self, responses: list[dict]) -> None:
        self.responses = list(responses)
        self.requests = []

    def open(self, request, timeout: int):
        self.requests.append((request, timeout))
        return _FakeResponse(self.responses.pop(0))


def _parsed_query_curl() -> dict:
    return {
        "url": "http://saasportal.iio.com/improxy/saas/iio2/sameAudit/queryAll",
        "cookie": "session=secret",
        "headers": [("Content-Type", "application/json")],
        "data": json.dumps(
            {
                "syncState": "1",
                "isReferUrl": "0",
                "tenantId": "188605",
                "commodityPoolId": "222",
            }
        ),
        "method": "POST",
    }


class _PassThroughAuthManager:
    def __init__(self) -> None:
        self.calls = []

    def prepare_curl(self, system, parsed_curl, *, refresh=True):
        self.calls.append((system, refresh))
        prepared = dict(parsed_curl)
        prepared["cookie"] = "x-access-token=state-access; refreshToken=state-refresh"
        return prepared


def test_preview_refreshes_saas_auth_before_read_only_preflight(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    module = _module()
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)
    curl_file = tmp_path / "query.curl"
    curl_file.write_text(
        "curl 'http://saasportal.iio.com/improxy/saas/iio2/sameAudit/queryAll' "
        "-H 'Content-Type: application/json' -b 'x-access-token=old; refreshToken=old-r' "
        "--data-raw '{\"tenantId\":\"188605\",\"commodityPoolId\":\"222\"}'",
        encoding="utf-8",
    )
    auth_manager = _PassThroughAuthManager()
    monkeypatch.setattr(
        module,
        "fetch_pending_pairs",
        lambda parsed, **_kwargs: {
            ("101", "sku-clear"),
            ("104", "sku-other"),
        }
        if "state-access" in parsed["cookie"]
        else pytest.fail("preflight did not receive state-backed authentication"),
    )

    assert module.run_main(
        ["--input-xlsx", str(workbook), "--curl-file", str(curl_file)],
        auth_manager=auth_manager,
    ) == 0
    assert auth_manager.calls == [("saas", True)]


def test_load_approve_candidates_excludes_only_red_and_yellow_f_cells(
    tmp_path: Path,
) -> None:
    module = _module()
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)

    candidates = module.load_approve_candidates(workbook)

    assert [candidate.as_payload_item() for candidate in candidates] == [
        {"vCommodityId": "101", "skuId": "sku-clear", "modelAuditResult": "1"},
        {"vCommodityId": "104", "skuId": "sku-other", "modelAuditResult": "1"},
    ]
    assert [candidate.excel_row for candidate in candidates] == [2, 5]
    assert candidates[0].as_manifest_item()["sourceCells"] == [
        {"excelRow": 2, "fillColor": "00000000"},
        {"excelRow": 6, "fillColor": "00000000"},
    ]


def test_customer_price_filter_keeps_non_red_non_yellow_nonempty_values(
    tmp_path: Path,
) -> None:
    module = _module()
    workbook = tmp_path / "audit-with-customer-price.xlsx"
    wb = Workbook()
    wb.active.title = "summary"
    worksheet = wb.create_sheet("record_results")
    worksheet.append(
        HEADERS
        + ["四级类目ID", "四级类目名称", "客户数据匹配状态", "客户列表参考价", "客户价"]
    )
    rows = [
        ["201", None, None, None, "clear-priced", "clear", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, 50],
        ["202", None, None, None, "clear-blank", "clear", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, "   "],
        ["203", None, None, None, "clear-none", "clear", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, None],
        ["204", None, None, None, "clear-zero", "clear", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, 0],
        ["205", None, None, None, "red-priced", "red", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, 60],
        ["206", None, None, None, "yellow-priced", "yellow", None, None, None, None, None, None, None, None, "1", "1", None, None, None, None, 70],
    ]
    for row in rows:
        worksheet.append(row)
    worksheet["F6"].fill = PatternFill("solid", fgColor="FFFFC7CE")
    worksheet["F7"].fill = PatternFill("solid", fgColor="FFFFFF00")
    wb.save(workbook)

    candidates = module.load_approve_candidates(
        workbook, require_customer_price=True
    )

    assert [candidate.sku_id for candidate in candidates] == [
        "clear-priced",
        "clear-zero",
    ]


def test_customer_price_filter_requires_customer_price_header(tmp_path: Path) -> None:
    module = _module()
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)

    with pytest.raises(ValueError, match="客户价"):
        module.load_approve_candidates(workbook, require_customer_price=True)


def test_help_exposes_customer_price_filter() -> None:
    module = _module()

    assert "--require-customer-price" in module.build_parser().format_help()


def test_approve_url_is_locked_to_same_audit_approve() -> None:
    module = _module()
    parsed = _parsed_query_curl()

    assert module.approve_url_from_curl(parsed).endswith("/sameAudit/approve")
    with pytest.raises(ValueError, match="host"):
        module.approve_url_from_curl(
            {**parsed, "url": "https://example.com/sameAudit/queryAll"}
        )
    with pytest.raises(ValueError, match="path"):
        module.approve_url_from_curl(
            {**parsed, "url": "http://saasportal.iio.com/improxy/saas/iio2/sameAudit/reject"}
        )


def test_fixed_approve_curl_can_drive_unfiltered_pending_query() -> None:
    module = _module()
    parsed = {
        **_parsed_query_curl(),
        "url": "http://saasportal.iio.com/improxy/saas/iio2/sameAudit/approve",
        "data": json.dumps(
            {
                "tenantId": "188605",
                "commodityPoolId": "222",
                "commodityList": [
                    {
                        "vCommodityId": "stale",
                        "skuId": "stale",
                        "modelAuditResult": "1",
                    }
                ],
            }
        ),
    }
    opener = _FakeOpener(
        [{"rspCode": "0000", "rspDesc": "SUCCESS", "count": 0, "rows": []}]
    )

    assert module.fetch_pending_pairs(
        parsed,
        opener=opener,
        include_all_reference_states=True,
    ) == set()

    request = opener.requests[0][0]
    body = json.loads(request.data)
    assert request.full_url.endswith("/sameAudit/queryAll")
    assert body == {
        "commodityId": "",
        "commodityName": "",
        "syncState": "1",
        "isReferUrl": "",
        "page": 1,
        "limit": 200,
        "tenantId": "188605",
        "commodityPoolId": "222",
    }


def test_candidate_manifest_hash_binds_current_pending_intersection(
    tmp_path: Path,
) -> None:
    module = _module()
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)
    source = module.load_approve_candidates(workbook)

    pending, already_processed = module.select_pending_candidates(
        source, {("104", "sku-other")}
    )
    digest = module.candidate_manifest_sha256(pending)
    literal = json.dumps(
        [pending[0].as_manifest_item()], ensure_ascii=False, separators=(",", ":")
    ).encode("utf-8")

    assert pending == [source[1]]
    assert already_processed == [source[0]]
    assert digest == hashlib.sha256(literal).hexdigest()


def test_execute_approve_batches_stops_after_non_success_response() -> None:
    module = _module()
    candidates = [
        module.ApproveCandidate("101", "sku-a", "1", "00000000", 2),
        module.ApproveCandidate("102", "sku-b", "1", "00000000", 3),
        module.ApproveCandidate("103", "sku-c", "1", "00000000", 4),
    ]
    opener = _FakeOpener(
        [
            {"rspCode": "0000", "rspDesc": "SUCCESS"},
            {"rspCode": "3333", "rspDesc": "login required"},
        ]
    )
    successes = []

    with pytest.raises(RuntimeError, match="rspCode=3333"):
        module.execute_approve_batches(
            _parsed_query_curl(),
            candidates,
            batch_size=2,
            opener=opener,
            on_batch_success=successes.append,
        )

    assert [event["batchCount"] for event in successes] == [2]
    assert opener.requests[0][0].full_url.endswith("/sameAudit/approve")
    assert json.loads(opener.requests[0][0].data) == {
        "tenantId": "188605",
        "commodityPoolId": "222",
        "commodityList": [
            {"vCommodityId": "101", "skuId": "sku-a", "modelAuditResult": "1"},
            {"vCommodityId": "102", "skuId": "sku-b", "modelAuditResult": "1"},
        ],
    }


def test_execute_writes_approve_evidence_and_requires_zero_postflight(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    module = _module()
    workbook = tmp_path / "audit.xlsx"
    _write_audit_workbook(workbook)
    source = module.load_approve_candidates(workbook)
    digest = module.candidate_manifest_sha256(source)
    curl_file = tmp_path / "query.curl"
    curl_file.write_text(
        "curl 'http://saasportal.iio.com/improxy/saas/iio2/sameAudit/queryAll' "
        "-H 'Content-Type: application/json' -b 'session=secret' "
        "--data-raw '{\"tenantId\":\"188605\",\"commodityPoolId\":\"222\"}'",
        encoding="utf-8",
    )
    output_dir = tmp_path / "run"
    pending_results = iter(
        [
            {(candidate.v_commodity_id, candidate.sku_id) for candidate in source},
            set(),
        ]
    )
    monkeypatch.setattr(
        module, "fetch_pending_pairs", lambda *_args, **_kwargs: next(pending_results)
    )

    def _execute(_curl, items, **kwargs):
        event = {
            "batchIndex": 1,
            "batchCount": len(items),
            "firstExcelRow": 2,
            "lastExcelRow": 5,
            "rspCode": "0000",
            "rspDesc": "SUCCESS",
        }
        kwargs["on_batch_success"](event)
        return [event]

    monkeypatch.setattr(module, "execute_approve_batches", _execute)

    assert module.run_main(
        [
            "--input-xlsx", str(workbook),
            "--curl-file", str(curl_file),
            "--output-dir", str(output_dir),
            "--execute",
            "--expected-count", str(len(source)),
            "--expected-sha256", digest,
        ],
        auth_manager=_PassThroughAuthManager(),
    ) == 0

    manifest = json.loads((output_dir / "approve_manifest.json").read_text("utf-8"))
    summary = json.loads((output_dir / "approve_summary.json").read_text("utf-8"))
    assert manifest["candidateCount"] == 2
    assert "cookie" not in json.dumps(manifest).lower()
    assert summary["status"] == "COMPLETED"
    assert summary["approvedCount"] == 2
    assert summary["remainingPendingCount"] == 0
