import importlib.util
import json
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from scripts.export_excluded_malls import export_excluded_templates
from scripts.zhaoshang_mall_mapping import ZhaoshangMallMapping


EXTERNAL_TOOL = Path(
    r"E:\.A2026\招商局自动建任务\external_tools\crawl_task_tool\auto_crawl_task.py"
)


def _load_external_tool():
    tool_dir = str(EXTERNAL_TOOL.parent)
    if tool_dir not in sys.path:
        sys.path.insert(0, tool_dir)
    spec = importlib.util.spec_from_file_location("zhaoshang_auto_crawl_task", EXTERNAL_TOOL)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def _write_template(path, code, url):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品编码", "url"])
    sheet.append([code, url])
    workbook.save(path)


def _template_codes(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
        return [str(row[0]) for row in rows[1:]]
    finally:
        workbook.close()


def test_single_mall_template_allows_product_code_without_url(tmp_path):
    module = _load_external_tool()
    template = tmp_path / "haoxiny.xlsx"
    _write_template(template, "HX-001", "")

    assert module.read_subtasks_from_excel(template, "haoxiny") == [
        {"productCode": "HX-001", "url": "", "subtaskName": "HX-001"}
    ]


def test_single_mall_template_rejects_url_without_product_code(tmp_path):
    module = _load_external_tool()
    template = tmp_path / "haoxiny.xlsx"
    domain = module.MALL_CONFIG["haoxiny"]["domains"][0]
    _write_template(template, "", f"https://{domain}/product/1")

    with pytest.raises(ValueError, match="商品编码"):
        module.read_subtasks_from_excel(template, "haoxiny")


def test_lxwl_templates_prefer_url_domain_over_shared_vendor_keyword(tmp_path, monkeypatch):
    # 商城映射表（商城sourceId映射表/商城sourceId映射表.xlsx）是 gitignored
    # 的本地生产输入；注入结构等价的合成映射。两行共用厂商名“领先未来商城”
    # （命中 lxwl 的厂商关键字），若按厂商关键字匹配会全部落入“领先未来商城”，
    # 断言要求按 URL 域名分别路由到两个商城，在任何机器上都成立。
    monkeypatch.setattr(
        "scripts.export_excluded_malls.EXCLUDED_MALLS",
        [
            ZhaoshangMallMapping(
                mall_name="领先未来商城",
                source_id="34",
                mall_key="lxwl",
                domains=("www.66123123.com", "66123123.com"),
                vendor_keywords=("lxwl", "领先未来商城"),
            ),
            ZhaoshangMallMapping(
                mall_name="领先未来新商城",
                source_id="83",
                mall_key="lxwlmall",
                domains=("www.lxwlmall.com", "lxwlmall.com"),
                vendor_keywords=("lxwlmall", "领先未来新商城"),
            ),
        ],
    )
    source = tmp_path / "saas.json"
    source.write_text(
        json.dumps(
            {
                "rows": [
                    {
                        "skuId": "OLD-1",
                        "extLink": "https://www.66123123.com/product/1",
                        "vendorName": "领先未来商城",
                    },
                    {
                        "skuId": "NEW-1",
                        "extLink": "https://www.lxwlmall.com/detail/1",
                        "vendorName": "领先未来商城",
                    },
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    summary = export_excluded_templates(source, tmp_path / "out", "batch")

    assert _template_codes(summary["领先未来商城"]["path"]) == ["OLD-1"]
    assert _template_codes(summary["领先未来新商城"]["path"]) == ["NEW-1"]
