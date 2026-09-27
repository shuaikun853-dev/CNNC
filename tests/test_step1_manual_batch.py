from pathlib import Path

from openpyxl import Workbook

from scripts import audit_saas_excel_file as step1


def _template(path: Path, rows: list[tuple[str, str]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品编码", "url"])
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


def test_step1_supported_request_uses_batch_and_never_loads_legacy_tool(
    tmp_path, monkeypatch, offline_manual_mall_mappings
):
    url = "https://item.jd.com/100123.html"
    template = _template(tmp_path / "jd.xlsx", [("100123", url)])
    seen = []

    def batch_runner(items, **kwargs):
        seen.append((items, kwargs))
        return {
            items[0].source_id: {
                "state": "SUCCESS",
                "result": {
                    "page_price_status": "PRICE_IS_ZERO",
                    "page_price_value": 0,
                    "page_title": "已下架",
                    "comment_count": 0,
                    "page_sku_id": "100123",
                    "returned_url": url,
                },
            }
        }

    monkeypatch.setattr(
        step1,
        "_load_zhaoshang_run_crawl_task",
        lambda _path: (_ for _ in ()).throw(
            AssertionError("normal supported STEP1 request must not load legacy create")
        ),
    )

    result = step1.run_zhaoshang_tasks(
        {
            "京东": {
                "deduped": 1,
                "path": template,
                "mall_key": "jd",
                "source_id": "3",
            }
        },
        batch_runner=batch_runner,
        batch_state_path=tmp_path / "BATCH_STATE.md",
        max_workers=1,
    )

    assert len(seen) == 1
    assert seen[0][1]["purpose"] == "legacy-step1-manual"
    assert result["京东"]["ok"] is True
    assert result["京东"]["backend_version"] == "batch-url-v1"
    assert result["京东"]["normalized_results"][url]["page_price_status"] == "PRICE_IS_ZERO"
    assert result["京东"]["normalized_results"][url]["page_price_value"] == 0


def test_step1_unsupported_url_uses_explicit_legacy_fallback(tmp_path, monkeypatch):
    url = "https://www.mymro.cn/not-proven/98S6010"
    template = _template(tmp_path / "mymro.xlsx", [("98S6010", url)])
    legacy_calls = []

    def legacy_tool(**kwargs):
        legacy_calls.append(kwargs)
        return {
            "result_path": str(tmp_path / "legacy-result.xlsx"),
            "main_task_id": 1,
            "task_number": "RW-1",
            "priority_state": "already_highest_priority",
            "affected_rows": 0,
            "already_highest_count": 1,
        }

    monkeypatch.setattr(
        step1, "_load_zhaoshang_run_crawl_task", lambda _path: legacy_tool
    )

    result = step1.run_zhaoshang_tasks(
        {
            "固安捷": {
                "deduped": 1,
                "path": template,
                "mall_key": "mymro",
                "source_id": "20",
            }
        },
        batch_runner=lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("unsupported URL must not be sent to Batch")
        ),
        max_workers=1,
    )

    assert len(legacy_calls) == 1
    assert result["固安捷"]["backend_version"] == "legacy-v1"


def test_step1_url_without_verified_mall_mapping_uses_legacy_fallback(
    tmp_path, monkeypatch
):
    url = "https://unverified.example/product/100123"
    template = _template(tmp_path / "jd-unknown-host.xlsx", [("100123", url)])
    legacy_calls = []
    monkeypatch.setattr(
        step1,
        "_load_zhaoshang_run_crawl_task",
        lambda _path: lambda **kwargs: legacy_calls.append(kwargs)
        or {
            "result_path": str(tmp_path / "legacy-result.xlsx"),
            "main_task_id": 1,
            "task_number": "RW-1",
            "priority_state": "already_highest_priority",
            "already_highest_count": 1,
        },
    )

    result = step1.run_zhaoshang_tasks(
        {
            "京东": {
                "deduped": 1,
                "path": template,
                "mall_key": "jd",
                "source_id": "3",
            }
        },
        batch_runner=lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("unverified host must not be sent to Batch")
        ),
        max_workers=1,
    )

    assert len(legacy_calls) == 1
    assert result["京东"]["backend_version"] == "legacy-v1"


def test_step1_batch_identity_conflict_is_not_reported_as_success(tmp_path, offline_manual_mall_mappings):
    url = "https://item.jd.com/100123.html"
    template = _template(tmp_path / "jd.xlsx", [("100123", url)])

    result = step1.run_zhaoshang_tasks(
        {
            "京东": {
                "deduped": 1,
                "path": template,
                "mall_key": "jd",
                "source_id": "3",
            }
        },
        batch_runner=lambda items, **_kwargs: {
            items[0].source_id: {
                "state": "MANUAL_REQUIRED",
                "error": "IDENTITY_CONFLICT: duplicate executeNumber result",
            }
        },
        max_workers=1,
    )

    assert result["京东"]["ok"] is False
    assert "IDENTITY_CONFLICT" in result["京东"]["error"]


def test_step1_batch_runner_error_is_returned_in_compatible_failure_shape(tmp_path, offline_manual_mall_mappings):
    url = "https://item.jd.com/100123.html"
    template = _template(tmp_path / "jd.xlsx", [("100123", url)])

    result = step1.run_zhaoshang_tasks(
        {
            "京东": {
                "deduped": 1,
                "path": template,
                "mall_key": "jd",
                "source_id": "3",
            }
        },
        batch_runner=lambda *_args, **_kwargs: (_ for _ in ()).throw(
            RuntimeError("batch unavailable")
        ),
        max_workers=1,
    )

    assert result["京东"]["ok"] is False
    assert result["京东"]["backend_version"] == "batch-url-v1"
    assert "batch unavailable" in result["京东"]["error"]


def test_step1_terminal_batch_failure_preserves_confirmed_priority_evidence(tmp_path, offline_manual_mall_mappings):
    url = "https://item.jd.com/100123.html"
    template = _template(tmp_path / "jd.xlsx", [("100123", url)])

    result = step1.run_zhaoshang_tasks(
        {
            "京东": {
                "deduped": 1,
                "path": template,
                "mall_key": "jd",
                "source_id": "3",
            }
        },
        batch_runner=lambda items, **_kwargs: {
            items[0].source_id: {
                "state": "TERMINAL_FAILED",
                "error": "crawler failed",
                "priority_confirmed": True,
            }
        },
        max_workers=1,
    )

    assert result["京东"]["ok"] is False
    assert result["京东"]["priority_state"] == "batch_lifecycle_confirmed"


def test_step1_mixed_mall_legacy_template_contains_only_unsupported_rows(
    tmp_path, monkeypatch, offline_manual_mall_mappings
):
    supported = "https://www.mymro.cn/u-98S6010.html"
    unsupported = "https://www.mymro.cn/not-proven/98S6011"
    template = _template(
        tmp_path / "mymro.xlsx",
        [("98S6010", supported), ("98S6011", unsupported)],
    )
    legacy_urls = []

    def legacy_tool(**kwargs):
        from openpyxl import load_workbook

        workbook = load_workbook(kwargs["excel_path"], data_only=True, read_only=True)
        try:
            rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
            legacy_urls.extend(row[1] for row in rows)
        finally:
            workbook.close()
        return {
            "result_path": str(tmp_path / "legacy-result.xlsx"),
            "main_task_id": 1,
            "task_number": "RW-1",
            "priority_state": "already_highest_priority",
            "already_highest_count": 1,
        }

    monkeypatch.setattr(
        step1, "_load_zhaoshang_run_crawl_task", lambda _path: legacy_tool
    )

    result = step1.run_zhaoshang_tasks(
        {
            "固安捷": {
                "deduped": 2,
                "path": template,
                "mall_key": "mymro",
                "source_id": "20",
            }
        },
        batch_runner=lambda items, **_kwargs: {
            items[0].source_id: {
                "state": "SUCCESS",
                "result": {
                    "page_price_status": "OK",
                    "page_price_value": 12.5,
                    "page_title": "title",
                    "page_sku_id": "changed",
                    "returned_url": supported,
                },
            }
        },
        batch_state_path=tmp_path / "BATCH_STATE.md",
        max_workers=1,
    )

    assert legacy_urls == [unsupported]


def test_step1_mixed_mall_does_not_hide_batch_failure_behind_legacy_success(
    tmp_path, monkeypatch, offline_manual_mall_mappings
):
    supported = "https://www.mymro.cn/u-98S6010.html"
    unsupported = "https://www.mymro.cn/not-proven/98S6011"
    template = _template(
        tmp_path / "mymro.xlsx",
        [("98S6010", supported), ("98S6011", unsupported)],
    )
    monkeypatch.setattr(
        step1,
        "_load_zhaoshang_run_crawl_task",
        lambda _path: lambda **_kwargs: {
            "result_path": str(tmp_path / "legacy-result.xlsx"),
            "main_task_id": 1,
            "task_number": "RW-1",
            "priority_state": "already_highest_priority",
            "already_highest_count": 1,
        },
    )

    result = step1.run_zhaoshang_tasks(
        {
            "固安捷": {
                "deduped": 2,
                "path": template,
                "mall_key": "mymro",
                "source_id": "20",
            }
        },
        batch_runner=lambda items, **_kwargs: {
            items[0].source_id: {
                "state": "MANUAL_REQUIRED",
                "error": "IDENTITY_CONFLICT",
            }
        },
        batch_state_path=tmp_path / "BATCH_STATE.md",
        max_workers=1,
    )

    assert result["固安捷"]["ok"] is False
    assert "IDENTITY_CONFLICT" in result["固安捷"]["error"]


def test_step1_jslink_is_canonicalized_to_nbdeli_before_batch_routing(
    tmp_path, monkeypatch, offline_manual_mall_mappings
):
    url = "https://www.jslink.com/pcweb/detail?sku=122262592"
    template = _template(tmp_path / "jslink.xlsx", [("122262592", url)])
    seen = []
    monkeypatch.setattr(
        step1,
        "_load_zhaoshang_run_crawl_task",
        lambda _path: (_ for _ in ()).throw(
            AssertionError("canonical jslink work must not use legacy create")
        ),
    )

    result = step1.run_zhaoshang_tasks(
        {
            "得力集实采": {
                "deduped": 1,
                "path": template,
                "mall_key": "jslink",
                "source_id": "legacy-jslink",
            }
        },
        batch_runner=lambda items, **_kwargs: seen.extend(items)
        or {
            items[0].source_id: {
                "state": "SUCCESS",
                "result": {
                    "page_price_status": "OK",
                    "page_price_value": 10,
                    "page_title": "title",
                    "page_sku_id": "122262592",
                    "returned_url": items[0].request.url,
                },
            }
        },
        batch_state_path=tmp_path / "BATCH_STATE.md",
        max_workers=1,
    )

    assert seen[0].request.mall_key == "nbdeli"
    assert "b2b.nbdeli.com" in seen[0].request.url
    assert url in result["得力集实采"]["normalized_results"]
    rows = [{"extLink": url, "skuId": "122262592"}]
    assert step1.apply_zhaoshang_results_to_rows(rows, [], {}, result) == 1


def test_step1_apply_consumes_normalized_batch_result_without_fake_workbook():
    url = "https://item.jd.com/100123.html"
    rows = [{"extLink": url, "skuId": "100123"}]

    count = step1.apply_zhaoshang_results_to_rows(
        rows,
        [],
        {},
        {
            "京东": {
                "ok": True,
                "normalized_results": {
                    url: {
                        "product_code": "100123",
                        "page_price_status": "PRICE_IS_ZERO",
                        "page_price_value": 0,
                        "page_price_source": "zhaoshang_batch",
                    }
                },
            }
        },
    )

    assert count == 1
    assert rows[0]["__page_price_value"] == 0
    assert rows[0]["__page_price_status"] == "PRICE_IS_ZERO"


def test_step1_batch_normalized_results_satisfy_the_upper_result_gate():
    assert step1._zhaoshang_result_has_payload(
        {
            "ok": True,
            "backend_version": "batch-url-v1",
            "normalized_results": {
                "https://item.jd.com/100123.html": {
                    "page_price_status": "OK",
                    "page_price_value": 10,
                }
            },
        }
    )


def test_step1_apply_merges_batch_and_legacy_fallback_results(monkeypatch):
    batch_url = "https://www.mymro.cn/u-98S6010.html"
    legacy_url = "https://www.mymro.cn/not-proven/98S6011"
    rows = [
        {"extLink": batch_url, "skuId": "98S6010"},
        {"extLink": legacy_url, "skuId": "98S6011"},
    ]
    monkeypatch.setattr(
        step1,
        "read_zhaoshang_result_xlsx",
        lambda *_args, **_kwargs: {
            legacy_url: {
                "product_code": "98S6011",
                "page_price_status": "OK",
                "page_price_value": 20,
                "page_price_source": "legacy",
            }
        },
    )

    count = step1.apply_zhaoshang_results_to_rows(
        rows,
        [],
        {},
        {
            "固安捷": {
                "ok": True,
                "result_path": "legacy.xlsx",
                "normalized_results": {
                    batch_url: {
                        "product_code": "98S6010",
                        "page_price_status": "OK",
                        "page_price_value": 10,
                        "page_price_source": "zhaoshang_batch",
                    }
                },
            }
        },
    )

    assert count == 2
    assert [row["__page_price_value"] for row in rows] == [10, 20]


def test_step1_batch_results_keep_same_url_different_sku_identity_exact():
    url = "https://item.jd.com/shared.html"
    rows = [
        {"extLink": url, "skuId": "SKU-A"},
        {"extLink": url, "skuId": "SKU-B"},
    ]
    summary = {
        "京东": {
            "ok": True,
            "normalized_results_by_identity": {
                url: {
                    "SKU-A": {
                        "product_code": "SKU-A",
                        "page_price_status": "OK",
                        "page_price_value": 10,
                        "page_price_source": "zhaoshang_batch",
                    },
                    "SKU-B": {
                        "product_code": "SKU-B",
                        "page_price_status": "OK",
                        "page_price_value": 20,
                        "page_price_source": "zhaoshang_batch",
                    },
                }
            },
        }
    }

    count = step1.apply_zhaoshang_results_to_rows(rows, [], {}, summary)

    assert count == 2
    assert rows[0]["__page_price_value"] == 10
    assert rows[1]["__page_price_value"] == 20
