from __future__ import annotations

from datetime import UTC, datetime, timedelta
import json
import sqlite3
import subprocess
from decimal import Decimal
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from app.services.daily_audit_store import DailyAuditStore
from app.services import incremental_local_fallback as fallback_module
from app.workflows import incremental_price_audit as audit_workflow
from app.services.incremental_local_fallback import (
    AFTER_1800,
    DAYTIME_60M,
    PEAK_3M,
    SUPPORTED_MALLS,
    eligibility_reason,
    run_incremental_local_fallback,
    read_crawler_output,
    execute_pinned_crawler,
    verify_crawler_revision,
)
from app.workflows.incremental_price_audit import (
    TickBudget, TickConfig, TickDependencies, _backlog,
    _final_remote_check_for_local_fallback, _harvest,
)


def at(hour, minute=0, second=0):
    return datetime(2026, 9, 17, hour, minute, second, tzinfo=UTC)


@pytest.mark.parametrize(
    "now,started,last_check,expected",
    [
        (at(8, 59, 59), at(8), None, None),       # 16:59:59 CST, 59m59s
        (at(9), at(8), None, None),               # 17:00 CST is peak, no prior peak check
        (at(8, 59, 59), at(7, 59, 59), None, DAYTIME_60M),
        (at(9), at(7), None, None),
        (at(9, 2, 59), at(7), at(9), None),
        (at(9, 3), at(7), at(9), PEAK_3M),
        (at(9, 57), at(7), at(9, 54), PEAK_3M),
        (at(9, 59, 59), at(7), at(9, 56, 59), PEAK_3M),
        (at(10), at(9, 59), None, AFTER_1800),
    ],
)
def test_time_policy_boundaries(now, started, last_check, expected):
    assert eligibility_reason(now, started, last_check) == expected


def setup_work(tmp_path, *, mall="nbdeli", count=1, state="SUBMITTED", started=at(7)):
    store = DailyAuditStore(tmp_path / "state.sqlite3")
    store.ensure_run("2026-09-17")
    token = store.acquire_lease("2026-09-17", "test", ttl_seconds=600)
    job = store.prepare_remote_intent(token, "job", "task", "hash", mall_key=mall)
    rows = []
    for index in range(count):
        work = store.ensure_crawl_work(
            token, mall, f"sku-{index}", f"https://{mall}.test/item/{index}",
            "legacy-v1", f"fp-{index}", source_id=f"doc-{index}",
        )
        store.link_crawl_work_to_remote_job(token, work["id"], job["id"])
        rows.append(work)
    with sqlite3.connect(store.path) as conn:
        conn.execute("UPDATE remote_jobs SET state=?, created_at=? WHERE id=?", (state, started.isoformat(), job["id"]))
        conn.execute("UPDATE crawl_work SET state=? WHERE remote_job_id=?", (state, job["id"]))
    return store, token, rows, job


def success_result(price="12.30"):
    return {
        "page_price_value": Decimal(price),
        "page_price_status": "OK" if Decimal(price) else "PRICE_IS_ZERO",
        "page_price_source": "本地兜底:nbdeli",
    }


def checked_jobs(candidates):
    return {int(candidate["remote_job_id"]) for candidate in candidates}


def test_schema_and_atomic_claim_allow_only_one_owner(tmp_path):
    store, token, rows, _job = setup_work(tmp_path)
    observed = store.observe_local_fallback_remote_checks(
        token, at(8, 30).isoformat(), SUPPORTED_MALLS, checked_remote_job_ids=set()
    )
    assert observed[0]["last_remote_check_at"] is None
    first = store.claim_local_fallback(token, rows[0]["id"], DAYTIME_60M, "tick-a", at(8, 30).isoformat())
    second = store.claim_local_fallback(token, rows[0]["id"], DAYTIME_60M, "tick-b", at(8, 30).isoformat())
    assert first is not None and second is None


def test_peak_first_check_then_second_check_runs_once(tmp_path):
    store, token, rows, _job = setup_work(tmp_path)
    calls = []
    first = run_incremental_local_fallback(
        store, token, now=at(9), tick_id="17:00",
        final_remote_check=lambda candidates: calls.append("remote") or checked_jobs(candidates),
        execute=lambda claims: calls.append("local") or {},
    )
    second = run_incremental_local_fallback(
        store, token, now=at(9, 3), tick_id="17:03",
        final_remote_check=lambda candidates: calls.append("remote") or checked_jobs(candidates),
        execute=lambda claims: {claims[0]["id"]: {"result": success_result()}},
    )
    assert first["pending"] == 1 and "local" not in calls
    assert second["success"] == 1 and calls.count("remote") == 2
    assert store.get_crawl_work(rows[0]["id"])["state"] == "SUCCESS"


def test_peak_timestamp_requires_completed_check_for_exact_job(tmp_path):
    store, token, _rows, job = setup_work(tmp_path)
    run_incremental_local_fallback(
        store, token, now=at(9), tick_id="failed-check",
        final_remote_check=lambda _candidates: set(), execute=lambda _claims: {},
    )
    ledger = store.list_local_fallback_work(token)[0]
    assert ledger["last_remote_check_at"] is None

    run_incremental_local_fallback(
        store, token, now=at(9, 3), tick_id="actual-check",
        final_remote_check=lambda _candidates: {int(job["id"])}, execute=lambda _claims: {},
    )
    ledger = store.list_local_fallback_work(token)[0]
    assert ledger["last_remote_check_at"] == at(9, 3).isoformat()


def test_failed_final_remote_check_neither_stamps_nor_falls_back(tmp_path):
    store, token, rows, _job = setup_work(tmp_path)
    local_calls = []

    def failed_check(_candidates):
        raise TimeoutError("remote query timed out")

    metrics = run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="failed-final-check",
        final_remote_check=failed_check,
        execute=lambda claims: local_calls.append(claims) or {},
    )
    assert metrics["pending"] == 1
    assert local_calls == []
    assert store.list_local_fallback_work(token)[0]["last_remote_check_at"] is None
    assert store.get_crawl_work(rows[0]["id"])["state"] == "SUBMITTED"


def test_final_remote_result_wins_before_local_claim(tmp_path):
    store, token, rows, _job = setup_work(tmp_path)
    local_calls = []
    def remote(_candidates):
        result = success_result("9.99")
        result["page_price_source"] = "招商局"
        store.complete_crawl_work(token, rows[0]["id"], result)
        return checked_jobs(_candidates)
    metrics = run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="day", final_remote_check=remote,
        execute=lambda claims: local_calls.append(claims) or {},
    )
    assert metrics["success"] == 0 and metrics["remote_success"] == 1 and local_calls == []
    assert json.loads(store.get_crawl_work(rows[0]["id"])["result_json"])["page_price_source"] == "招商局"


def test_partial_job_only_unresolved_urls_are_executed(tmp_path):
    store, token, rows, _job = setup_work(tmp_path, count=20)
    for row in rows[:15]:
        store.complete_crawl_work(token, row["id"], success_result())
    received = []
    def execute(claims):
        received.extend(claims)
        return {claim["id"]: {"result": success_result("4.56")} for claim in claims}
    metrics = run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="partial", final_remote_check=checked_jobs,
        execute=execute,
    )
    assert metrics["success"] == 5
    assert len(received) == 5
    assert len({claim["exact_url"] for claim in received}) == 5


def test_duplicate_exact_url_is_claimed_once_and_fans_out(tmp_path):
    store, token, rows, job = setup_work(tmp_path)
    duplicate = store.ensure_crawl_work(
        token, "nbdeli", "sku-0", rows[0]["exact_url"], "legacy-v1", "other-fp", source_id="doc-other"
    )
    store.link_crawl_work_to_remote_job(token, duplicate["id"], job["id"])
    with sqlite3.connect(store.path) as conn:
        conn.execute("UPDATE crawl_work SET state='SUBMITTED' WHERE id=?", (duplicate["id"],))
    received = []
    metrics = run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="dedupe", final_remote_check=checked_jobs,
        execute=lambda claims: received.extend(claims) or {claims[0]["id"]: {"result": success_result()}},
    )
    assert len(received) == 1
    assert store.get_crawl_work(rows[0]["id"])["state"] == "SUCCESS"
    assert store.get_crawl_work(duplicate["id"])["state"] == "SUCCESS"
    assert metrics["pending"] == 0 and metrics["success"] == 1 and metrics["total"] == 1
    assert metrics["links"] == [{
        "mall_key": "nbdeli", "exact_url": rows[0]["exact_url"], "state": "SUCCESS",
        "source": "本地兜底:nbdeli", "remote_job_id": job["id"],
    }]


@pytest.mark.parametrize("terminal_state", ["NO_RESULT", "TERMINAL_FAILED"])
def test_local_success_supersedes_terminal_duplicate_exact_url(tmp_path, terminal_state):
    store, token, rows, job = setup_work(tmp_path)
    duplicate = store.ensure_crawl_work(
        token, "nbdeli", "sku-0", rows[0]["exact_url"], "legacy-v1",
        "other-fp", source_id="doc-other",
    )
    store.link_crawl_work_to_remote_job(token, duplicate["id"], job["id"])
    with sqlite3.connect(store.path) as conn:
        conn.execute("UPDATE crawl_work SET state='SUBMITTED' WHERE id=?", (duplicate["id"],))
    store.observe_local_fallback_remote_checks(
        token, at(8, 30).isoformat(), SUPPORTED_MALLS,
        checked_remote_job_ids={int(job["id"])},
    )
    claim = store.claim_local_fallback(
        token, rows[0]["id"], DAYTIME_60M, "local", at(8, 30).isoformat()
    )
    assert claim is not None

    if terminal_state == "TERMINAL_FAILED":
        store.fail_crawl_work(token, duplicate["id"], "remote terminal", retryable=False)
    else:
        with sqlite3.connect(store.path) as conn:
            conn.execute(
                """UPDATE remote_jobs
                      SET state='COMPLETED', remote_identity='remote-job',
                          post_attempted=1, result_path='remote.xlsx'
                    WHERE id=?""",
                (job["id"],),
            )
        store.complete_partial_exact_no_result(
            token, duplicate["id"], remote_job_id=job["id"],
            remote_identity="remote-job", result_path="remote.xlsx",
            reason="remote export has no matching row",
        )

    store.complete_local_fallback(
        token, claim["id"], success_result("4.56"),
        crawler_revision="crawler-revision", completed_at=at(8, 31).isoformat(),
    )

    for crawl_id in (rows[0]["id"], duplicate["id"]):
        crawl = store.get_crawl_work(crawl_id)
        assert crawl["state"] == "SUCCESS"
        assert json.loads(crawl["result_json"])["page_price_value"] == "4.56"


def test_late_remote_cannot_overwrite_local_success(tmp_path):
    store, token, rows, _job = setup_work(tmp_path)
    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="local", final_remote_check=checked_jobs,
        execute=lambda claims: {claims[0]["id"]: {"result": success_result("4.56")}},
    )
    with pytest.raises(Exception, match="immutable"):
        store.complete_crawl_work(token, rows[0]["id"], success_result("99.99"))


def test_opening_harvest_ignores_late_remote_after_local_success(tmp_path):
    store, token, rows, job = setup_work(tmp_path)
    with sqlite3.connect(store.path) as conn:
        conn.execute(
            "UPDATE remote_jobs SET remote_identity='remote-job', post_attempted=1 WHERE id=?",
            (job["id"],),
        )
    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="local", final_remote_check=checked_jobs,
        execute=lambda claims: {claims[0]["id"]: {"result": success_result("4.56")}},
    )
    remote = {"page_price_value": "99.99", "page_price_status": "OK", "page_price_source": "招商局"}
    deps = TickDependencies(
        discover=lambda _timeout: None,
        harvest=lambda _jobs, _timeout: {
            int(job["id"]): {
                "state": "COMPLETED", "remote_identity": "remote-job",
                "crawl_results": {int(rows[0]["id"]): remote},
            }
        },
        monotonic=lambda: 0,
    )
    counters = {"crawl_harvested": 0}
    checked = _harvest(
        store, token, deps, TickBudget(0, 60, 5, lambda: 0), counters,
        TickConfig(business_date="2026-09-17"),
    )
    assert checked == {int(job["id"])}
    assert counters["late_remote_ignored"] == 1
    assert json.loads(store.get_crawl_work(rows[0]["id"])["result_json"])["page_price_value"] == "4.56"


def test_remote_check_proof_requires_per_job_evidence(tmp_path):
    store, token, rows, job = setup_work(tmp_path)
    deps = TickDependencies(
        discover=lambda _timeout: None,
        harvest=lambda _jobs, _timeout: {},
        monotonic=lambda: 0,
    )
    checked: set[int] = set()
    swept = _harvest(
        store, token, deps, TickBudget(0, 60, 5, lambda: 0),
        {"crawl_harvested": 0}, TickConfig(business_date="2026-09-17"),
        checked_job_ids=checked,
    )
    assert swept == {int(job["id"])} and checked == set()
    assert _final_remote_check_for_local_fallback(
        store, token, deps, TickBudget(0, 60, 5, lambda: 0),
        [store.get_crawl_work(rows[0]["id"])],
    ) == set()


@pytest.mark.parametrize(
    "evidence",
    [
        {"manual_error": "remote query failed"},
        {"state": "PRIORITY_PENDING", "remote_identity": "remote-job"},
        {
            "state": "PRIORITY_PENDING",
            "remote_identity": "remote-job",
            "retryable_error": "priority rows are not ready",
        },
    ],
)
def test_final_remote_check_rejects_non_query_evidence(tmp_path, evidence):
    store, token, rows, _job = setup_work(tmp_path)
    deps = TickDependencies(
        discover=lambda _timeout: None,
        harvest=lambda jobs, _timeout: {int(jobs[0]["id"]): evidence},
        monotonic=lambda: 0,
    )

    checked = _final_remote_check_for_local_fallback(
        store, token, deps, TickBudget(0, 60, 5, lambda: 0),
        [store.get_crawl_work(rows[0]["id"])],
    )

    assert checked == set()


def test_failed_final_remote_check_cannot_persist_embedded_price(tmp_path):
    store, token, rows, job = setup_work(tmp_path)
    deps = TickDependencies(
        discover=lambda _timeout: None,
        harvest=lambda _jobs, _timeout: {
            int(job["id"]): {
                "state": "COMPLETED", "manual_error": "export incomplete",
                "crawl_results": {
                    int(rows[0]["id"]): {
                        "page_price_value": "19.20", "page_price_status": "OK",
                        "page_price_source": "招商局",
                    }
                },
            }
        },
    )
    checked = _final_remote_check_for_local_fallback(
        store, token, deps, TickBudget(0, 60, 5, lambda: 0),
        [store.get_crawl_work(rows[0]["id"])],
    )
    assert checked == set()
    assert store.get_crawl_work(rows[0]["id"])["state"] == "SUBMITTED"


@pytest.mark.parametrize("state", ["RUNNING", "EXPORTED", "COMPLETED"])
def test_final_remote_check_accepts_successful_remote_query_states(tmp_path, state):
    store, token, rows, job = setup_work(tmp_path)
    deps = TickDependencies(
        discover=lambda _timeout: None,
        harvest=lambda _jobs, _timeout: {
            int(job["id"]): {"state": state, "remote_identity": "remote-job"}
        },
        monotonic=lambda: 0,
    )

    checked = _final_remote_check_for_local_fallback(
        store, token, deps, TickBudget(0, 60, 5, lambda: 0),
        [store.get_crawl_work(rows[0]["id"])],
    )

    assert checked == {int(job["id"])}


@pytest.mark.parametrize("terminal_state", ["NO_RESULT", "TERMINAL_FAILED"])
def test_final_remote_price_rescues_terminal_before_local_claim(tmp_path, terminal_state):
    store, token, rows, job = setup_work(tmp_path)
    seed_remote_terminal(store, token, rows[0], job, terminal_state)
    remote_price = {"page_price_value": "19.20", "page_price_status": "OK", "page_price_source": "招商局"}
    deps = TickDependencies(
        discover=lambda _timeout: None,
        harvest=lambda _jobs, _timeout: {
            int(job["id"]): {
                "state": "COMPLETED", "remote_identity": "remote-job",
                "crawl_results": {int(rows[0]["id"]): remote_price},
            }
        },
    )

    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="remote-rescue",
        final_remote_check=lambda candidates: _final_remote_check_for_local_fallback(
            store, token, deps, TickBudget(0, 60, 5, lambda: 0), candidates
        ),
        execute=lambda _claims: pytest.fail("valid final remote price must win"),
    )

    crawl = store.get_crawl_work(rows[0]["id"])
    assert crawl["state"] == "SUCCESS"
    assert json.loads(crawl["result_json"]) == remote_price
    assert store.list_local_fallback_work(token)[0]["attempt_count"] == 0


def test_opening_harvest_manual_error_is_not_a_completed_remote_check(tmp_path):
    store, token, _rows, job = setup_work(tmp_path)
    deps = TickDependencies(
        discover=lambda _timeout: None,
        harvest=lambda _jobs, _timeout: {
            int(job["id"]): {"manual_error": "remote query failed"}
        },
        monotonic=lambda: 0,
    )
    checked: set[int] = set()

    _harvest(
        store, token, deps, TickBudget(0, 60, 5, lambda: 0),
        {"crawl_harvested": 0}, TickConfig(business_date="2026-09-17"),
        checked_job_ids=checked,
    )

    assert checked == set()


def seed_remote_terminal(store, token, row, job, state):
    if state == "TERMINAL_FAILED":
        store.fail_crawl_work(token, row["id"], "remote terminal", retryable=False)
    else:
        with sqlite3.connect(store.path) as conn:
            conn.execute(
                """UPDATE remote_jobs SET state='COMPLETED', remote_identity='remote-job',
                          post_attempted=1, result_path='remote.xlsx' WHERE id=?""",
                (job["id"],),
            )
        store.complete_partial_exact_no_result(
            token, row["id"], remote_job_id=job["id"], remote_identity="remote-job",
            result_path="remote.xlsx", reason="remote export has no matching row",
        )


def _assert_remote_terminal_can_be_claimed_for_local_fallback(tmp_path, terminal_state):
    store, token, rows, job = setup_work(tmp_path)
    seed_remote_terminal(store, token, rows[0], job, terminal_state)
    observed = store.observe_local_fallback_remote_checks(
        token, at(8, 30).isoformat(), SUPPORTED_MALLS,
        checked_remote_job_ids={int(job["id"])},
    )
    assert [row["id"] for row in observed] == [rows[0]["id"]]
    claim = store.claim_local_fallback(
        token, rows[0]["id"], DAYTIME_60M, "late-claim", at(8, 31).isoformat()
    )
    assert claim is not None and claim["attempt_count"] == 1


def test_no_result_can_be_claimed_for_local_fallback(tmp_path):
    _assert_remote_terminal_can_be_claimed_for_local_fallback(tmp_path, "NO_RESULT")


def test_terminal_failed_can_be_claimed_for_local_fallback(tmp_path):
    _assert_remote_terminal_can_be_claimed_for_local_fallback(tmp_path, "TERMINAL_FAILED")


def _assert_remote_terminal_local_success_supersedes_exact_price(tmp_path, terminal_state):
    store, token, rows, job = setup_work(tmp_path)
    seed_remote_terminal(store, token, rows[0], job, terminal_state)
    metrics = run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="local",
        final_remote_check=checked_jobs,
        execute=lambda claims: {claims[0]["id"]: {"result": success_result("18.80")}},
    )
    crawl = store.get_crawl_work(rows[0]["id"])
    assert metrics["success"] == 1
    assert crawl["state"] == "SUCCESS"
    assert json.loads(crawl["result_json"])["page_price_value"] == "18.80"
    assert json.loads(crawl["result_json"])["page_price_source"].startswith("本地兜底:")


def test_no_result_local_success_supersedes_with_local_price(tmp_path):
    _assert_remote_terminal_local_success_supersedes_exact_price(tmp_path, "NO_RESULT")


def test_terminal_failed_local_success_supersedes_with_local_price(tmp_path):
    _assert_remote_terminal_local_success_supersedes_exact_price(tmp_path, "TERMINAL_FAILED")


def _assert_remote_terminal_local_failure_preserves_remote_evidence(tmp_path, terminal_state):
    store, token, rows, job = setup_work(tmp_path)
    seed_remote_terminal(store, token, rows[0], job, terminal_state)
    before = store.get_crawl_work(rows[0]["id"])
    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="local",
        final_remote_check=checked_jobs,
        execute=lambda claims: {claims[0]["id"]: {"error": "RISK_STOP", "retryable": False}},
    )
    after = store.get_crawl_work(rows[0]["id"])
    assert (after["state"], after["result_json"], after["error_text"]) == (
        before["state"], before["result_json"], before["error_text"]
    )
    assert store.list_local_fallback_work(token)[0]["state"] == "TERMINAL_FAILED"


def test_no_result_local_failure_preserves_no_result(tmp_path):
    _assert_remote_terminal_local_failure_preserves_remote_evidence(tmp_path, "NO_RESULT")


def test_terminal_failed_local_failure_preserves_terminal_evidence(tmp_path):
    _assert_remote_terminal_local_failure_preserves_remote_evidence(tmp_path, "TERMINAL_FAILED")


def _assert_remote_terminal_local_crash_preserves_remote_evidence(tmp_path, terminal_state):
    store, token, rows, job = setup_work(tmp_path)
    seed_remote_terminal(store, token, rows[0], job, terminal_state)
    store.observe_local_fallback_remote_checks(
        token, at(8, 30).isoformat(), SUPPORTED_MALLS,
        checked_remote_job_ids={int(job["id"])},
    )
    claim = store.claim_local_fallback(
        token, rows[0]["id"], DAYTIME_60M, "crashed", at(8, 30).isoformat()
    )
    assert claim is not None
    before = store.get_crawl_work(rows[0]["id"])
    assert store.recover_interrupted_local_fallback_claims(
        token, "next", at(8, 30, 2).isoformat(), max_attempts=1
    ) == 1
    after = store.get_crawl_work(rows[0]["id"])
    assert (after["state"], after["result_json"], after["error_text"]) == (
        before["state"], before["result_json"], before["error_text"]
    )


def test_no_result_local_crash_preserves_no_result(tmp_path):
    _assert_remote_terminal_local_crash_preserves_remote_evidence(tmp_path, "NO_RESULT")


def test_terminal_failed_local_crash_preserves_terminal_evidence(tmp_path):
    _assert_remote_terminal_local_crash_preserves_remote_evidence(tmp_path, "TERMINAL_FAILED")


@pytest.mark.parametrize("terminal_state", ["NO_RESULT", "TERMINAL_FAILED"])
def test_remote_terminal_local_retry_caps_at_two(tmp_path, terminal_state):
    store, token, rows, job = setup_work(tmp_path)
    seed_remote_terminal(store, token, rows[0], job, terminal_state)
    attempts = []
    for index, when in enumerate((at(8, 30), at(8, 31), at(8, 32))):
        run_incremental_local_fallback(
            store, token, now=when, tick_id=f"tick-{index}",
            final_remote_check=checked_jobs,
            execute=lambda claims: attempts.append(claims[0]["id"]) or {
                claims[0]["id"]: {"error": "NETWORK_TIMEOUT", "retryable": True}
            },
        )
    assert len(attempts) == 2
    assert store.list_local_fallback_work(token)[0]["attempt_count"] == 2
    assert store.list_local_fallback_work(token)[0]["state"] == "TERMINAL_FAILED"


@pytest.mark.parametrize("terminal_state", ["NO_RESULT", "TERMINAL_FAILED"])
@pytest.mark.parametrize("settle", ["fail", "recover"])
def test_local_failure_paths_preserve_concurrent_terminal_remote_evidence(
    tmp_path, terminal_state, settle
):
    store, token, rows, job = setup_work(tmp_path)
    store.observe_local_fallback_remote_checks(
        token, at(8, 30).isoformat(), SUPPORTED_MALLS,
        checked_remote_job_ids={int(job["id"])},
    )
    claim = store.claim_local_fallback(
        token, rows[0]["id"], DAYTIME_60M, "claim", at(8, 30).isoformat()
    )
    assert claim is not None

    if terminal_state == "TERMINAL_FAILED":
        store.fail_crawl_work(token, rows[0]["id"], "remote terminal", retryable=False)
    else:
        with sqlite3.connect(store.path) as conn:
            conn.execute(
                """UPDATE remote_jobs
                      SET state='COMPLETED', remote_identity='remote-job',
                          post_attempted=1, result_path='remote.xlsx'
                    WHERE id=?""",
                (job["id"],),
            )
        store.complete_partial_exact_no_result(
            token, rows[0]["id"], remote_job_id=job["id"],
            remote_identity="remote-job", result_path="remote.xlsx",
            reason="remote export has no matching row",
        )
    before = store.get_crawl_work(rows[0]["id"])

    if settle == "fail":
        store.fail_local_fallback(
            token, claim["id"], "local terminal", retryable=False
        )
    else:
        assert store.recover_interrupted_local_fallback_claims(
            token, "next-tick", at(8, 30, 2).isoformat(), max_attempts=1
        ) == 1

    after = store.get_crawl_work(rows[0]["id"])
    assert after["state"] == terminal_state
    assert after["result_json"] == before["result_json"]
    assert after["error_text"] == before["error_text"]


def test_stale_submit_unknown_does_not_block_after_local_success(tmp_path):
    store, token, rows, job = setup_work(tmp_path, state="SUBMIT_UNKNOWN")
    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="local", final_remote_check=checked_jobs,
        execute=lambda claims: {claims[0]["id"]: {"result": success_result("4.56")}},
    )
    assert store.get_remote_job(job["id"])["state"] == "SUBMIT_UNKNOWN"
    backlog = _backlog(store, token, [])
    assert backlog["remote_unknown"] == 0
    assert backlog["crawl_pending"] == 0


def test_mixed_remote_and_local_success_does_not_keep_job_pending(tmp_path):
    store, token, rows, job = setup_work(tmp_path, count=2, state="SUBMIT_UNKNOWN")
    store.complete_crawl_work(token, rows[0]["id"], {
        "page_price_value": "1.23", "page_price_status": "OK",
        "page_price_source": "招商局",
    })
    store.complete_crawl_work(token, rows[1]["id"], success_result("4.56"))

    assert store.get_remote_job(job["id"])["state"] == "SUBMIT_UNKNOWN"
    assert _backlog(store, token, [])["remote_unknown"] == 0


@pytest.mark.parametrize("bad_price", ["not-a-price", "NaN", "Infinity"])
def test_malformed_sibling_success_keeps_stale_remote_job_pending(tmp_path, bad_price):
    store, token, rows, job = setup_work(tmp_path, count=2, state="SUBMIT_UNKNOWN")
    store.complete_crawl_work(token, rows[0]["id"], {
        "page_price_value": bad_price, "page_price_status": "OK",
        "page_price_source": "招商局",
    })
    store.complete_crawl_work(token, rows[1]["id"], success_result("4.56"))

    assert store.get_remote_job(job["id"])["state"] == "SUBMIT_UNKNOWN"
    assert _backlog(store, token, [])["remote_unknown"] == 1


def test_retryable_failure_waits_for_future_tick_and_caps_at_two(tmp_path):
    store, token, rows, _job = setup_work(tmp_path)
    calls = []
    def fail(claims):
        calls.append(claims[0]["claim_tick_id"])
        return {claims[0]["id"]: {"error": "NETWORK_ERROR", "retryable": True}}
    one = run_incremental_local_fallback(store, token, now=at(8, 30), tick_id="a", final_remote_check=checked_jobs, execute=fail)
    same = run_incremental_local_fallback(store, token, now=at(8, 31), tick_id="a", final_remote_check=checked_jobs, execute=fail)
    two = run_incremental_local_fallback(store, token, now=at(8, 32), tick_id="b", final_remote_check=checked_jobs, execute=fail)
    three = run_incremental_local_fallback(store, token, now=at(8, 33), tick_id="c", final_remote_check=checked_jobs, execute=fail)
    assert [one["failed"], same["failed"], two["failed"], three["failed"]] == [1, 1, 1, 1]
    assert calls == ["a", "b"]
    ledger = store.list_local_fallback_work(token)[0]
    assert ledger["attempt_count"] == 2 and ledger["state"] == "TERMINAL_FAILED"


def test_unsupported_mall_and_completed_are_never_candidates(tmp_path):
    store, token, rows, _job = setup_work(tmp_path, mall="other")
    assert run_incremental_local_fallback(
        store, token, now=at(10), tick_id="x", final_remote_check=checked_jobs,
        execute=lambda c: pytest.fail("must not execute"),
    )["total"] == 0


def test_historical_replay_candidate_counts_are_unique(tmp_path):
    store, token, rows, _job = setup_work(tmp_path, count=171, started=at(6))
    received = []
    run_incremental_local_fallback(
        store, token, now=at(8), tick_id="2016-09-15", final_remote_check=checked_jobs,
        execute=lambda claims: received.extend(claims) or {
            claim["id"]: {"error": "RISK_STOP", "retryable": False} for claim in claims
        },
    )
    assert len(received) == 171
    assert len({claim["exact_url"] for claim in received}) == 171


def test_20260916_replay_manual_qstbg_does_not_wait_forever(tmp_path):
    store, token, _rows, _job = setup_work(
        tmp_path, mall="qstbg", count=50, state="MANUAL_REQUIRED", started=at(7)
    )
    received = []
    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="2026-09-16", final_remote_check=checked_jobs,
        execute=lambda claims: received.extend(claims) or {
            claim["id"]: {"error": "RISK_STOP", "retryable": False} for claim in claims
        },
    )
    assert len(received) == 50
    assert {claim["eligible_reason"] for claim in received} == {DAYTIME_60M}


def write_output(path, rows):
    book = Workbook()
    sheet = book.active
    sheet.append(["商城商品编码", "商品标题", "商品价格", "商品链接", "采集状态", "失败原因"])
    for row in rows:
        sheet.append(row)
    book.save(path)


def fake_crawler_run(monkeypatch, *, site_outputs, exit_code):
    monkeypatch.setattr(fallback_module, "verify_crawler_revision", lambda _root: fallback_module.EXPECTED_CRAWLER_REVISION)

    def run(command, **kwargs):
        manifest = load_workbook(command[command.index("--manifest") + 1], data_only=True)
        try:
            for site, _source, target in list(manifest.active.values)[1:]:
                rows = site_outputs.get(site)
                if rows == "MALFORMED":
                    fallback_module.Path(target).write_bytes(b"not an xlsx")
                elif rows is not None:
                    write_output(fallback_module.Path(target), rows)
        finally:
            manifest.close()
        return subprocess.CompletedProcess(command, exit_code, "", "")

    monkeypatch.setattr(fallback_module.subprocess, "run", run)


def multi_site_claims():
    return [
        {"id": 1, "mall_key": "nbdeli", "exact_url": "https://b2b.nbdeli.com/item/1"},
        {"id": 2, "mall_key": "nbdeli", "exact_url": "https://b2b.nbdeli.com/item/2"},
        {"id": 3, "mall_key": "guangbo", "exact_url": "https://btob.guangbo.net/item/3"},
        {"id": 4, "mall_key": "comix", "exact_url": "https://www.qx.com/item/4"},
    ]


def test_zdet_is_excluded_from_incremental_fallback_but_uses_pinned_crawler_revision():
    url = "https://b.zdet.net/mall/item?id=1400061078916"

    assert "zdet" not in SUPPORTED_MALLS
    assert "zdet" not in fallback_module.SITE_HOSTS
    assert not fallback_module._url_matches_site(url, "zdet")
    assert fallback_module.EXPECTED_CRAWLER_REVISION == (
        "08abb74089cfa9007dc6f6ddc850f9fd41990a24"
    )
    assert fallback_module.DEFAULT_CRAWLER_ROOT == fallback_module.Path(
        r"E:\.A中核运维\cnnc-mall-crawler-prod"
    )


def test_execute_pinned_crawler_exit3_preserves_completed_other_site_success(tmp_path, monkeypatch):
    claims = multi_site_claims()
    fake_crawler_run(monkeypatch, exit_code=3, site_outputs={
        "nbdeli": [
            ["1", "first", "18.80", claims[0]["exact_url"], "SUCCESS", ""],
            ["2", "second", "", claims[1]["exact_url"], "SKIPPED_RISK", "RISK_STOP"],
        ],
        "guangbo": [["3", "third", "19.20", claims[2]["exact_url"], "SUCCESS", ""]],
    })
    result = execute_pinned_crawler(claims, run_dir=tmp_path / "run", tick_id="risk", crawler_root=tmp_path / "crawler")
    assert result[1]["result"]["page_price_value"] == Decimal("18.80")
    assert result[3]["result"]["page_price_value"] == Decimal("19.20")


def test_execute_pinned_crawler_exit3_marks_only_risk_or_unstarted_urls_as_risk(tmp_path, monkeypatch):
    claims = multi_site_claims()
    fake_crawler_run(monkeypatch, exit_code=3, site_outputs={
        "nbdeli": [
            ["1", "first", 1, claims[0]["exact_url"], "SUCCESS", ""],
            ["2", "second", "", claims[1]["exact_url"], "SKIPPED_RISK", "RISK_STOP"],
        ],
        "guangbo": [["3", "third", 2, claims[2]["exact_url"], "SUCCESS", ""]],
    })
    result = execute_pinned_crawler(claims, run_dir=tmp_path / "run", tick_id="risk", crawler_root=tmp_path / "crawler")
    assert result[2]["error"].startswith("RISK_STOP")
    assert result[4]["error"].startswith("RISK_STOP")
    assert "result" in result[1] and "result" in result[3]


def test_execute_pinned_crawler_exit2_preserves_valid_other_site_outputs(tmp_path, monkeypatch):
    claims = multi_site_claims()
    fake_crawler_run(monkeypatch, exit_code=2, site_outputs={
        "nbdeli": [["1", "first", 1, claims[0]["exact_url"], "SUCCESS", ""],
                   ["2", "second", 2, claims[1]["exact_url"], "SUCCESS", ""]],
        "comix": [["4", "fourth", 4, claims[3]["exact_url"], "SUCCESS", ""]],
    })
    result = execute_pinned_crawler(claims, run_dir=tmp_path / "run", tick_id="error", crawler_root=tmp_path / "crawler")
    assert {key for key, outcome in result.items() if "result" in outcome} == {1, 2, 4}
    assert result[3]["error"].startswith("PROCESS_FAILED")


def _assert_execute_pinned_crawler_bad_one_site_does_not_poison_others(tmp_path, monkeypatch, bad_site):
    claims = multi_site_claims()
    bad_rows = {
        "MISSING": None,
        "MALFORMED": "MALFORMED",
        "FOREIGN": [["3", "third", 3, "https://btob.guangbo.net/foreign", "SUCCESS", ""]],
        "DUPLICATE": [["3", "third", 3, claims[2]["exact_url"], "SUCCESS", ""],
                      ["3", "third", 3, claims[2]["exact_url"], "SUCCESS", ""]],
    }[bad_site]
    fake_crawler_run(monkeypatch, exit_code=2, site_outputs={
        "nbdeli": [["1", "first", 1, claims[0]["exact_url"], "SUCCESS", ""],
                   ["2", "second", 2, claims[1]["exact_url"], "SUCCESS", ""]],
        "guangbo": bad_rows,
        "comix": [["4", "fourth", 4, claims[3]["exact_url"], "SUCCESS", ""]],
    })
    result = execute_pinned_crawler(claims, run_dir=tmp_path / "run", tick_id="bad", crawler_root=tmp_path / "crawler")
    assert {key for key, outcome in result.items() if "result" in outcome} == {1, 2, 4}
    assert result[3]["error"].startswith("PROCESS_FAILED" if bad_site == "MISSING" else "MANIFEST_MISMATCH")


def test_execute_pinned_crawler_missing_one_site_output_does_not_poison_other_sites(tmp_path, monkeypatch):
    _assert_execute_pinned_crawler_bad_one_site_does_not_poison_others(tmp_path, monkeypatch, "MISSING")


def test_execute_pinned_crawler_malformed_one_site_output_does_not_poison_other_sites(tmp_path, monkeypatch):
    _assert_execute_pinned_crawler_bad_one_site_does_not_poison_others(tmp_path, monkeypatch, "MALFORMED")


def test_execute_pinned_crawler_foreign_url_is_site_scoped_failure(tmp_path, monkeypatch):
    _assert_execute_pinned_crawler_bad_one_site_does_not_poison_others(tmp_path, monkeypatch, "FOREIGN")


def test_execute_pinned_crawler_duplicate_url_is_site_scoped_failure(tmp_path, monkeypatch):
    _assert_execute_pinned_crawler_bad_one_site_does_not_poison_others(tmp_path, monkeypatch, "DUPLICATE")


def test_execute_pinned_crawler_wrong_host_preflight_is_site_scoped(tmp_path, monkeypatch):
    claims = [multi_site_claims()[0], {
        "id": 3, "mall_key": "guangbo", "exact_url": "https://wrong.invalid/item/3",
    }]
    fake_crawler_run(monkeypatch, exit_code=0, site_outputs={
        "nbdeli": [["1", "first", "18.80", claims[0]["exact_url"], "SUCCESS", ""]],
    })
    result = execute_pinned_crawler(claims, run_dir=tmp_path / "run", tick_id="scope", crawler_root=tmp_path / "crawler")
    assert result[1]["result"]["page_price_value"] == Decimal("18.80")
    assert result[3]["error"] == "INVALID_SCOPE"


def test_execute_pinned_crawler_duplicate_claim_preflight_is_site_scoped(tmp_path, monkeypatch):
    first = multi_site_claims()[0]
    claims = [first, {**first, "id": 2}, multi_site_claims()[2]]
    fake_crawler_run(monkeypatch, exit_code=0, site_outputs={
        "guangbo": [["3", "third", "19.20", claims[2]["exact_url"], "SUCCESS", ""]],
    })
    result = execute_pinned_crawler(claims, run_dir=tmp_path / "run", tick_id="duplicate", crawler_root=tmp_path / "crawler")
    assert result[1]["error"].startswith("MANIFEST_MISMATCH")
    assert result[2]["error"].startswith("MANIFEST_MISMATCH")
    assert result[3]["result"]["page_price_value"] == Decimal("19.20")


def test_execute_pinned_crawler_rechecks_live_budget_at_child_launch(tmp_path, monkeypatch):
    claim = multi_site_claims()[0]
    available = [180.0]
    monkeypatch.setattr(
        fallback_module, "verify_crawler_revision",
        lambda _root: available.__setitem__(0, 75.0) or fallback_module.EXPECTED_CRAWLER_REVISION,
    )
    seen_timeout = []

    def run(command, **kwargs):
        seen_timeout.append(kwargs["timeout"])
        manifest = load_workbook(command[command.index("--manifest") + 1], data_only=True)
        try:
            _site, _source, target = list(manifest.active.values)[1]
            write_output(fallback_module.Path(target), [
                ["1", "first", "18.80", claim["exact_url"], "SUCCESS", ""],
            ])
        finally:
            manifest.close()
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(fallback_module.subprocess, "run", run)
    result = execute_pinned_crawler(
        [claim], run_dir=tmp_path / "run", tick_id="live-budget", crawler_root=tmp_path / "crawler",
        timeout_seconds=lambda: available[0],
    )
    assert seen_timeout == [75.0]
    assert result[1]["result"]["page_price_value"] == Decimal("18.80")


def test_budget_exhausted_during_claim_loop_rolls_back_all_attempts(tmp_path):
    store, token, _rows, _job = setup_work(tmp_path, count=2)
    available = iter((180.0, 180.0, 0.0))
    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="budget-loop",
        final_remote_check=checked_jobs,
        execute=lambda _claims: pytest.fail("budget exhaustion must defer before child"),
        execution_budget_seconds=lambda: next(available),
    )
    assert [row["attempt_count"] for row in store.list_local_fallback_work(token)] == [0, 0]


def test_budget_exhausted_during_crawler_preparation_does_not_charge_attempt(tmp_path, monkeypatch):
    store, token, _rows, _job = setup_work(tmp_path)
    with sqlite3.connect(store.path) as connection:
        connection.execute(
            "UPDATE crawl_work SET exact_url='https://b2b.nbdeli.com/item/0' WHERE id=?",
            (_rows[0]["id"],),
        )
    available = [180.0]
    monkeypatch.setattr(
        fallback_module, "verify_crawler_revision",
        lambda _root: available.__setitem__(0, 0.0) or fallback_module.EXPECTED_CRAWLER_REVISION,
    )
    monkeypatch.setattr(
        fallback_module.subprocess, "run",
        lambda *_args, **_kwargs: pytest.fail("no child may start after budget exhaustion"),
    )
    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="budget-prepare",
        final_remote_check=checked_jobs,
        execute=lambda claims: execute_pinned_crawler(
            claims, run_dir=tmp_path / "run", tick_id="budget-prepare",
            crawler_root=tmp_path / "crawler", timeout_seconds=lambda: available[0],
        ),
        execution_budget_seconds=lambda: available[0],
    )
    ledger = store.list_local_fallback_work(token)[0]
    assert ledger["attempt_count"] == 0 and ledger["state"] == "PENDING"


def test_prelaunch_budget_defer_preserves_retry_count_for_next_tick(tmp_path):
    store, token, rows, _job = setup_work(tmp_path)
    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="first",
        final_remote_check=checked_jobs,
        execute=lambda claims: {claims[0]["id"]: {"error": "TEMPORARY_TIMEOUT", "retryable": True}},
    )
    run_incremental_local_fallback(
        store, token, now=at(8, 31), tick_id="deferred",
        final_remote_check=checked_jobs,
        execute=lambda claims: {claims[0]["id"]: {"deferred": True}},
    )
    ledger = store.list_local_fallback_work(token)[0]
    assert ledger["attempt_count"] == 1 and ledger["state"] == "RETRYABLE_FAILED"
    run_incremental_local_fallback(
        store, token, now=at(8, 32), tick_id="next",
        final_remote_check=checked_jobs,
        execute=lambda claims: {claims[0]["id"]: {"result": success_result("18.80")}},
    )
    assert store.list_local_fallback_work(token)[0]["attempt_count"] == 2
    assert store.get_crawl_work(rows[0]["id"])["state"] == "SUCCESS"


def _assert_verify_crawler_revision_rejects_wrong_sha_or_dirty_worktree(tmp_path, dirty):
    root = tmp_path / "crawler"
    (root / "mall_crawler").mkdir(parents=True)
    (root / "mall_crawler" / "cli.py").write_text("# test fixture\n", encoding="utf-8")

    def runner(command, **_kwargs):
        output = (fallback_module.EXPECTED_CRAWLER_REVISION if dirty else "0" * 40) if command[1] == "rev-parse" else (" M cli.py" if dirty else "")
        return subprocess.CompletedProcess(command, 0, output, "")

    with pytest.raises(ValueError, match="工作区不干净" if dirty else "revision 不匹配"):
        verify_crawler_revision(root, runner=runner)


def test_verify_crawler_revision_rejects_wrong_sha(tmp_path):
    _assert_verify_crawler_revision_rejects_wrong_sha_or_dirty_worktree(tmp_path, False)


def test_verify_crawler_revision_rejects_dirty_worktree(tmp_path):
    _assert_verify_crawler_revision_rejects_wrong_sha_or_dirty_worktree(tmp_path, True)


def test_local_fallback_does_not_claim_when_budget_exhausted_after_final_remote_check(tmp_path):
    store, token, rows, _job = setup_work(tmp_path)
    available = [90.0]
    executed = []

    def check(candidates):
        available[0] = 30.0
        return checked_jobs(candidates)

    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="deferred",
        final_remote_check=check,
        execute=lambda claims: executed.append(claims) or {},
        execution_budget_seconds=lambda: available[0],
    )
    ledger = store.list_local_fallback_work(token)[0]
    assert ledger["state"] == "PENDING"
    assert ledger["attempt_count"] == 0
    assert executed == []
    assert store.get_crawl_work(rows[0]["id"])["state"] == "SUBMITTED"


def test_budget_defer_does_not_consume_fallback_attempt(tmp_path):
    store, token, _rows, _job = setup_work(tmp_path)
    run_incremental_local_fallback(
        store, token, now=at(8, 30), tick_id="deferred",
        final_remote_check=checked_jobs,
        execute=lambda _claims: pytest.fail("budget defer must not execute"),
        execution_budget_seconds=lambda: 0.0,
    )
    assert store.list_local_fallback_work(token)[0]["attempt_count"] == 0


def test_next_scheduler_tick_can_claim_after_budget_defer(tmp_path):
    store, token, rows, _job = setup_work(tmp_path)
    for tick_id, available in (("deferred", 30.0), ("next", 180.0)):
        run_incremental_local_fallback(
            store, token, now=at(8, 30) if tick_id == "deferred" else at(8, 31),
            tick_id=tick_id, final_remote_check=checked_jobs,
            execute=lambda claims: {claims[0]["id"]: {"result": success_result("18.80")}},
            execution_budget_seconds=lambda available=available: available,
        )
    assert store.get_crawl_work(rows[0]["id"])["state"] == "SUCCESS"
    assert store.list_local_fallback_work(token)[0]["attempt_count"] == 1


def test_local_fallback_child_timeout_is_bounded_by_live_tick_budget():
    time_now = [70.0]
    budget = TickBudget(0.0, 240.0, 120.0, lambda: time_now[0])
    observed = []
    deps = TickDependencies(
        discover=lambda _timeout: None,
        local_fallback=lambda _claims, _tick_id, timeout: observed.append(timeout) or {},
    )
    audit_workflow._execute_local_fallback_with_budget(deps, budget, (), "tick")
    assert len(observed) == 1 and observed[0]() == 50.0
    time_now[0] = 80.0
    assert observed[0]() == 40.0


@pytest.mark.parametrize("price", ["NaN", "Infinity", "-1"])
def test_crawler_output_rejects_non_finite_and_negative_prices(tmp_path, price):
    path = tmp_path / "out.xlsx"
    url = "https://b2b.nbdeli.com/pcweb/detail?sku=1"
    write_output(path, [["1", "item", price, url, "SUCCESS", ""]])
    with pytest.raises(ValueError, match="价格"):
        read_crawler_output(path, "nbdeli", [url])


def test_crawler_output_maps_zero_and_rejects_duplicate_foreign_or_missing(tmp_path):
    url = "https://b2b.nbdeli.com/pcweb/detail?sku=1"
    path = tmp_path / "zero.xlsx"
    write_output(path, [["1", "item", 0, url, "SUCCESS", ""]])
    assert read_crawler_output(path, "nbdeli", [url])[url]["result"]["page_price_status"] == "PRICE_IS_ZERO"

    duplicate = tmp_path / "duplicate.xlsx"
    write_output(duplicate, [["1", "item", 1, url, "SUCCESS", ""], ["1", "item", 1, url, "SUCCESS", ""]])
    with pytest.raises(ValueError, match="重复"):
        read_crawler_output(duplicate, "nbdeli", [url])

    foreign = tmp_path / "foreign.xlsx"
    write_output(foreign, [["1", "item", 1, "https://evil.example/1", "SUCCESS", ""]])
    with pytest.raises(ValueError, match="站点|锁定"):
        read_crawler_output(foreign, "nbdeli", [url])

    missing = tmp_path / "missing.xlsx"
    write_output(missing, [])
    with pytest.raises(ValueError, match="锁定"):
        read_crawler_output(missing, "nbdeli", [url])


def test_crawler_output_rejects_malformed_xlsx(tmp_path):
    path = tmp_path / "malformed.xlsx"
    path.write_bytes(b"not a zip workbook")
    with pytest.raises(ValueError, match="不是有效 XLSX"):
        read_crawler_output(path, "nbdeli", ["https://b2b.nbdeli.com/pcweb/detail?sku=1"])


def test_crawler_output_rejects_structurally_invalid_xlsx(tmp_path):
    path = tmp_path / "missing-content-types.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
    with pytest.raises(ValueError, match="不是有效 XLSX"):
        read_crawler_output(path, "nbdeli", ["https://b2b.nbdeli.com/pcweb/detail?sku=1"])
