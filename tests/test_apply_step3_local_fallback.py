from pathlib import Path
import importlib

import openpyxl
import pytest


def _load_apply_local_fallback():
    try:
        module = importlib.import_module("scripts.apply_step3_local_fallback")
    except ModuleNotFoundError:
        pytest.fail("apply_step3_local_fallback implementation is missing")
    return module.apply_local_fallback


def _write_import(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet2"
    sheet.append(["单据id", "互联网标题", "URL", "SKUID", "价格", "评论数"])
    sheet.append(["700000000000000001", "", "https://shop.test/success", "", "", ""])
    sheet.append(["700000000000000002", "", "https://shop.test/failure", "", "", ""])
    workbook.save(path)


def _write_crawl_result(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "采集结果"
    sheet.append([
        "原始链接", "商品编码", "商品名称", "商品价格", "商品评价数", "采集状态", "失败原因",
    ])
    sheet.append([
        "https://shop.test/success", "SKU-1", "有效标题", 12.5, 8, "成功", "",
    ])
    sheet.append([
        "https://shop.test/failure", "SKU-2", "失败标题", "", "", "失败", "验证码",
    ])
    workbook.save(path)


def test_local_fallback_fills_only_successful_positive_price_and_keeps_failures_blank(tmp_path):
    apply_local_fallback = _load_apply_local_fallback()
    input_path = tmp_path / "input.xlsx"
    result_path = tmp_path / "crawler.xlsx"
    output_path = tmp_path / "output.xlsx"
    log_path = tmp_path / "log.xlsx"
    remaining_path = tmp_path / "remaining.xlsx"
    _write_import(input_path)
    _write_crawl_result(result_path)

    summary = apply_local_fallback(
        input_path=input_path,
        crawler_result_path=result_path,
        output_path=output_path,
        log_path=log_path,
        remaining_path=remaining_path,
    )

    workbook = openpyxl.load_workbook(output_path, data_only=True)
    sheet = workbook["Sheet2"]
    assert sheet["A2"].value == "700000000000000001"
    assert sheet["A2"].number_format == "@"
    assert sheet["A3"].number_format == "@"
    assert sheet["D2"].number_format == "@"
    assert sheet["D3"].number_format == "@"
    assert sheet["B2"].value == "有效标题"
    assert sheet["D2"].value == "SKU-1"
    assert sheet["E2"].value == 12.5
    assert sheet["F2"].value == 8
    assert sheet["B3"].value in (None, "")
    assert sheet["E3"].value in (None, "")
    workbook.close()

    remaining = openpyxl.load_workbook(remaining_path, read_only=True, data_only=True)
    assert remaining.active.max_row == 2
    assert remaining.active.cell(2, 1).value == "700000000000000002"
    remaining.close()

    assert summary == {
        "input_rows": 2,
        "crawler_success_records": 1,
        "filled_rows": 1,
        "remaining_rows": 1,
    }


def test_local_fallback_keeps_only_numeric_comment_counts(tmp_path):
    apply_local_fallback = _load_apply_local_fallback()
    input_path = tmp_path / "input.xlsx"
    result_path = tmp_path / "crawler.xlsx"
    output_path = tmp_path / "output.xlsx"
    log_path = tmp_path / "log.xlsx"
    remaining_path = tmp_path / "remaining.xlsx"

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["单据id", "互联网标题", "URL", "SKUID", "价格", "评论数"])
    sheet.append(["1", "", "https://shop.test/existing", "", 1, "1,234+"])
    sheet.append(["2", "", "https://shop.test/crawled", "", "", "旧的非数字值"])
    workbook.save(input_path)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "原始链接", "商品编码", "商品名称", "商品价格", "商品评价数", "采集状态", "失败原因",
    ])
    sheet.append([
        "https://shop.test/crawled", "SKU-2", "有效标题", 2, "未识别", "成功", "",
    ])
    workbook.save(result_path)

    apply_local_fallback(
        input_path=input_path,
        crawler_result_path=result_path,
        output_path=output_path,
        log_path=log_path,
        remaining_path=remaining_path,
    )

    output = openpyxl.load_workbook(output_path, data_only=True)
    assert output.active["F2"].value == 1234
    assert isinstance(output.active["F2"].value, int)
    assert output.active["F3"].value is None
    output.close()
