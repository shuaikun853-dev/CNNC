"""Offline contracts for the finite incremental price-audit tick."""

from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import UTC, datetime, timedelta
import json
from urllib import error

import pytest
from pymysql.err import OperationalError

from app.services.daily_audit_store import DailyAuditStore, ProcessLock
from app.services.legacy_zhaoshang_adapter import (
    retryable_priority_error_diagnostic,
)
from app.services.saas_audit_core import build_records
from app.services.title_match_model_adapter import ModelCallRetryExhaustedError
from app.workflows.work_01_saas_approval_json_audit import audit_saas_approval_records
from openpyxl import load_workbook
import scripts.audit_saas_excel_file as legacy_excel
from app.workflows.incremental_price_audit import (
    CrawlPlan,
    CrawlMember,
    DiscoveredState,
    TickBudget,
    TickConfig,
    TickDependencies,
    _backlog,
    _safe_error,
    _harvest,
    run_existing_remote_drain,
    run_incremental_tick,
)


@dataclass
class Clock:
    value: float = 0

    def monotonic(self) -> float:
        return self.value

    def advance(self, seconds: float) -> None:
        self.value += seconds


def _row(v="doc-1", sku="sku-1", **extra):
    return {
        "vCommodityId": v,
        "documentId": v,
        "skuId": sku,
        "extLink": " https://shop.example/item ",
        "fourCategoryId": "cat",
        "list_reference_price": "10",
        "queryStr": "title",
        "skuName": "title",
        "mall_key": "mall-a",
        **extra,
    }


def _config(clock, *, date="2026-08-11", budget=60, reserve=10, stability=0.0):
    # Older contracts freeze on one complete observation; stability-window
    # behaviour has its own dedicated closeout tests.
    return TickConfig(
        business_date=date,
        runtime_budget_seconds=budget,
        heavy_work_reserve_seconds=reserve,
        lease_ttl_seconds=90,
        cutoff_at=datetime(2026, 8, 11, 10, 0, tzinfo=UTC),
        source_stability_seconds=stability,
    )


def _deps(events, rows, clock, *, post_raises=False):
    def harvest(jobs, timeout):
        events.append("harvest")
        return {}

    def discover(timeout):
        events.append("discover")
        return DiscoveredState(
            observed_at=datetime(2026, 8, 11, 10, 1, tzinfo=UTC),
            headers=["vCommodityId"],
            rows=rows,
            source_fingerprints=[f"source-{row_value['vCommodityId']}-{row_value['skuId']}" for row_value in rows],
            customer_fingerprints=[f"customer-{row_value['fourCategoryId']}-{row_value['list_reference_price']}" for row_value in rows],
            evidence={"saas_count": len(rows), "customer_count": len(rows)},
        )

    def detail(row, timeout):
        events.append(f"detail:{row['vCommodityId']}")
        # Round-3 contract: a healthy detail callback yields READY evidence
        # (positive price AND a FOUND-family status).  A bare price with no
        # status is contradictory evidence (INVALID) and blocks the tick.
        return dict(FOUND_PAYLOAD)

    def model(row, timeout):
        events.append(f"model:{row['vCommodityId']}")
        return {"match_result": "SAME"}

    def plan(items, timeout):
        members = [
            CrawlMember(
                mall_key=item["row"]["mall_key"], source_id=item["row"]["vCommodityId"], sku_id=item["row"]["skuId"],
                exact_url=item["row"]["extLink"], backend_version="legacy-v1", row=item["row"],
            )
            for item in items
        ]
        return [CrawlPlan(
            mall_key="mall-a", logical_key="mall-a:batch", task_name="tick-mall-a", payload_hash="payload-a",
            payload_json={"mall": "mall-a"}, members=members,
        )]

    def submit(job, crawl_items, timeout):
        events.append("submit")
        if post_raises:
            raise TimeoutError("post timed out")
        return {"remote_identity": "task-1", "state": "SUBMITTED"}

    def finalize(plan, members, timeout):
        payload = {"members": [member.sku_id for member in members]}
        suffix = "-".join(payload["members"])
        return CrawlPlan(
            mall_key=plan.mall_key,
            logical_key=f"{plan.mall_key}:{suffix}",
            task_name=f"tick-{suffix}",
            payload_hash=f"payload-{suffix}",
            payload_json=payload,
            members=tuple(members),
        )

    return TickDependencies(
        discover=discover, harvest=harvest, detail=detail, model=model, plan_crawl=plan,
        submit_crawl=submit, finalize_crawl_plan=finalize, monotonic=clock.monotonic, wall_clock=lambda: datetime(2026, 8, 11, 10, 2, tzinfo=UTC),
    )


def test_tick_order_idempotency_and_duplicate_membership(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    events: list[str] = []
    clock = Clock()
    deps = _deps(events, [_row(), _row()], clock)
    config = _config(clock)

    first = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    second = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)

    assert first.outcome == "WAITING_REMOTE"
    # Crawl submission now precedes MODEL: remote work is already pending while
    # the model retry backlog keeps the stage rows claimable.
    assert list(first.phase_trace[:5]) == ["recover", "harvest", "discover", "submit", "detail"]
    assert events.count("harvest") == 3
    assert events.count("detail:doc-1") == 1
    assert events.count("model:doc-1") == 1
    assert events.count("submit") == 1
    snapshot = store.latest_snapshot("2026-08-11")
    assert snapshot is not None
    assert len(store.ordered_snapshot_items(snapshot["id"])) == 2
    assert second.counters["detail_requests"] == 0


def test_overlap_returns_before_dependencies_and_next_date_is_independent(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    events: list[str] = []
    clock = Clock()
    lock = ProcessLock(tmp_path / "tick.lock").acquire()
    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), _deps(events, [_row()], clock))
    lock.release()
    assert summary.outcome == "ALREADY_RUNNING"
    assert events == []
    run_incremental_tick(store, ProcessLock(tmp_path / "other.lock"), _config(clock, date="2026-08-12"), _deps(events, [_row()], clock))
    assert store.get_run("2026-08-11") is None
    assert store.get_run("2026-08-12") is not None


def test_crash_claim_is_recovered_and_budget_retains_backlog(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    events: list[str] = []
    clock = Clock()
    deps = _deps(events, [_row()], clock)
    original_discover = deps.discover

    def late_discover(timeout):
        value = original_discover(timeout)
        clock.advance(55)
        return value

    deps.discover = late_discover
    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock, budget=60, reserve=10), deps)
    assert summary.outcome == "BUDGET_EXHAUSTED"
    assert events == ["harvest", "discover"]
    token = store.acquire_lease("2026-08-11", "test", 30)
    item = store.ordered_snapshot_items(store.latest_snapshot("2026-08-11")["id"])[0]
    work = store.ensure_stage_work(token, item["id"], "DETAIL", "manual-crash")
    store.claim_stage_work(token, work["id"], max_attempts=1, max_elapsed_seconds=10)
    store.release_lease(token)
    clock.value = 0
    resumed = run_incremental_tick(store, ProcessLock(tmp_path / "tick2.lock"), _config(clock), deps)
    assert resumed.counters["recovered_stage_claims"] == 1


def test_submit_timeout_is_unknown_and_never_blindly_resent(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    events: list[str] = []
    clock = Clock()
    config = _config(clock)
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, _deps(events, [_row()], clock, post_raises=True))
    assert events.count("submit") == 1
    rerun = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, _deps(events, [_row()], clock, post_raises=True))
    assert events.count("submit") == 1
    inspect = store.acquire_lease("2026-08-11", "inspect", 30)
    assert store.list_remote_jobs(inspect, states={"SUBMIT_UNKNOWN"})
    store.release_lease(inspect)


def test_priority_pending_is_harvested_before_crawl_can_run_and_never_reposts(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    events: list[str] = []
    clock = Clock()
    config = _config(clock)
    first_deps = _deps(events, [_row()], clock)
    first_deps.submit_crawl = lambda _job, _crawls, _timeout: (
        events.append("submit")
        or {"remote_identity": "mall-a|task-id|task-number", "state": "PRIORITY_PENDING"}
    )

    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, first_deps)
    inspect = store.acquire_lease("2026-08-11", "inspect", 30)
    job = store.list_remote_jobs(inspect)[0]
    crawl = store.crawl_work_for_job(job["id"])[0]
    assert job["state"] == "PRIORITY_PENDING"
    assert crawl["state"] == "PENDING"
    store.release_lease(inspect)

    second_deps = _deps(events, [_row()], clock)

    def harvest(jobs, _timeout):
        assert [job["state"] for job in jobs] == ["PRIORITY_PENDING"]
        return {
            jobs[0]["id"]: {
                "state": "SUBMITTED",
                "remote_identity": "mall-a|task-id|task-number",
                "result": {"priority_status": "PRIORITY_UPDATED", "weight_type": -1},
            }
        }

    second_deps.harvest = harvest
    second_deps.submit_crawl = lambda *_args: pytest.fail("priority recovery must never POST")
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, second_deps)

    verify = store.acquire_lease("2026-08-11", "verify", 30)
    assert store.get_remote_job(job["id"])["state"] == "SUBMITTED"
    assert store.get_crawl_work(crawl["id"])["state"] == "SUBMITTED"
    store.release_lease(verify)
    assert events.count("submit") == 1


def test_stage_limits_and_finalize_freeze_behavior(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    deps = _deps(events, [_row()], clock)
    before = TickDependencies(**{**deps.__dict__, "wall_clock": lambda: datetime(2026, 8, 11, 9, 59, tzinfo=UTC)})
    config = _config(clock)
    assert run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, before, finalize=True).outcome == "FINALIZE_NOT_READY"
    final = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps, finalize=True)
    frozen = final.frozen_snapshot_id
    assert frozen is not None
    changed = _deps(events, [_row("doc-2", "sku-2")], clock)
    rerun = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, changed, finalize=True)
    assert rerun.frozen_snapshot_id == frozen
    assert events.count("discover") == 1


def test_customer_change_reprocesses_changed_item_and_elapsed_cap_is_terminal(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    config = TickConfig(
        business_date="2026-08-11", runtime_budget_seconds=60, heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30, detail_max_attempts=1, detail_max_elapsed_seconds=1,
        model_max_attempts=1, model_max_elapsed_seconds=1,
    )
    deps = _deps(events, [_row()], clock)

    def slow_detail(row, timeout):
        clock.advance(2)
        return {"late": True}

    deps.detail = slow_detail
    first = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    assert store.list_stage_work(token, stage="DETAIL")[0]["state"] == "TERMINAL_FAILED"
    store.release_lease(token)
    clock.value = 0
    changed = _deps(events, [_row(list_reference_price="11")], clock)
    changed.detail = lambda row, timeout: {"customerPrice": "11"}
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, changed)
    token = store.acquire_lease("2026-08-11", "inspect-2", 30)
    assert len(store.list_stage_work(token, stage="DETAIL")) == 2
    store.release_lease(token)
    assert first.counters["detail_requests"] == 0


def test_detail_terminal_requires_manual_review_like_model_terminal(tmp_path):
    """Flipped with the customer-price gates: an exhausted DETAIL retry budget
    means the row's price-review baseline is unresolved, so the tick ends in
    MANUAL_REVIEW_REQUIRED exactly like a terminal MODEL row."""
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    detail_deps = _deps([], [_row()], clock)
    detail_deps.plan_crawl = None
    detail_deps.finalize_crawl_plan = None
    detail_deps.submit_crawl = None
    detail_deps.detail = lambda _row, _timeout: (_ for _ in ()).throw(
        ValueError("permanent detail failure")
    )
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30,
        detail_max_attempts=1,
        model_max_attempts=1,
    )

    detail_summary = run_incremental_tick(
        store, ProcessLock(tmp_path / "tick.lock"), config, detail_deps
    )

    assert detail_summary.outcome == "MANUAL_REVIEW_REQUIRED"
    assert detail_summary.backlog["detail_terminal"] == 1
    # A plain exception carries no business prefix, so the row classifies as
    # customer-price FAILED rather than MISSING.
    assert detail_summary.backlog["customer_price_failed"] == 1
    assert detail_summary.backlog["model_terminal"] == 0

    model_store = DailyAuditStore(tmp_path / "model-state.sqlite")
    model_deps = _deps([], [_row()], clock)
    model_deps.plan_crawl = None
    model_deps.finalize_crawl_plan = None
    model_deps.submit_crawl = None
    model_deps.model = lambda _row, _timeout: (_ for _ in ()).throw(
        ValueError("permanent model failure")
    )

    model_summary = run_incremental_tick(
        model_store, ProcessLock(tmp_path / "model-tick.lock"), config, model_deps
    )

    assert model_summary.outcome == "MANUAL_REVIEW_REQUIRED"
    assert model_summary.backlog["model_terminal"] == 1


def test_unexpected_discovery_error_releases_lock_and_lease(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    deps = _deps(events, [_row()], clock)
    deps.discover = lambda timeout: (_ for _ in ()).throw(RuntimeError("discovery exploded"))
    lock = ProcessLock(tmp_path / "tick.lock")
    failed = run_incremental_tick(store, lock, _config(clock), deps)
    assert failed.outcome == "FAILED"
    assert lock.held is False
    healthy = run_incremental_tick(store, lock, _config(clock), _deps(events, [_row()], clock))
    assert healthy.outcome == "WAITING_REMOTE"


def test_hard_post_crash_is_recovered_unknown_without_resend_and_zero_budget_calls_nothing(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    deps = _deps(events, [_row()], clock)
    deps.submit_crawl = lambda job, rows, timeout: (_ for _ in ()).throw(SystemExit("hard stop"))
    with pytest.raises(SystemExit):
        run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), deps)
    resumed = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), _deps(events, [_row()], clock))
    assert resumed.counters["recovered_submit_unknown"] == 1
    assert events.count("submit") == 0
    zero_events: list[str] = []
    zero = TickConfig(business_date="2026-08-12", runtime_budget_seconds=0, heavy_work_reserve_seconds=0, lease_ttl_seconds=1)
    summary = run_incremental_tick(store, ProcessLock(tmp_path / "zero.lock"), zero, _deps(zero_events, [_row()], clock))
    assert summary.outcome == "BUDGET_EXHAUSTED"
    assert zero_events == []


def test_finalize_freezes_current_saas_complete_snapshot_despite_customer_count(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11", cutoff_at="2026-08-11T10:00:00+00:00")
    token = store.acquire_lease("2026-08-11", "seed", 30)
    old = store.ingest_snapshot(
        token,
        [_row(f"old-{index}", f"old-s-{index}") for index in range(340)],
        "2026-08-11T10:01:00+00:00",
        headers=["vCommodityId"],
        evidence={"saas_count": 340, "customer_count": 82},
    )["snapshot"]
    store.release_lease(token)

    current_rows = [_row(f"doc-{index}", f"sku-{index}") for index in range(349)]
    events: list[str] = []
    deps = _deps(events, current_rows, clock)
    deps.plan_crawl = None
    deps.finalize_crawl_plan = None
    deps.submit_crawl = None
    deps.discover = lambda _timeout: DiscoveredState(
        observed_at=datetime(2026, 8, 11, 10, 2, tzinfo=UTC),
        headers=["vCommodityId"],
        rows=current_rows,
        source_fingerprints=[f"source-{index}" for index in range(349)],
        customer_fingerprints=[f"customer-{index}" for index in range(349)],
        evidence={"saas_count": 349, "customer_count": 82},
    )

    summary = run_incremental_tick(
        store, ProcessLock(tmp_path / "tick.lock"), _config(clock), deps, finalize=True
    )
    latest = store.list_snapshots("2026-08-11")[-1]

    assert summary.frozen_snapshot_id == latest["id"]
    assert summary.frozen_snapshot_id != old["id"]
    assert latest["row_count"] == 349
    assert summary.counters["saas_count"] == 349
    assert summary.counters["customer_count"] == 82
    assert "detail:doc-0" in events
    assert any(event.startswith("model:doc-") for event in events)


def test_finalize_does_not_freeze_when_current_row_count_differs_from_saas_count(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    rows = [_row(f"doc-{index}", f"sku-{index}") for index in range(348)]
    deps = _deps([], rows, clock)
    deps.plan_crawl = None
    deps.finalize_crawl_plan = None
    deps.submit_crawl = None
    deps.discover = lambda _timeout: DiscoveredState(
        observed_at=datetime(2026, 8, 11, 10, 2, tzinfo=UTC),
        headers=["vCommodityId"],
        rows=rows,
        source_fingerprints=[f"source-{index}" for index in range(348)],
        customer_fingerprints=[f"customer-{index}" for index in range(348)],
        evidence={"saas_count": 349, "customer_count": 82},
    )

    summary = run_incremental_tick(
        store, ProcessLock(tmp_path / "tick.lock"), _config(clock), deps, finalize=True
    )

    assert summary.outcome == "WAITING_SOURCE_COUNT"
    assert summary.frozen_snapshot_id is None


def test_unchanged_snapshot_is_observed_again_and_can_freeze_across_cutoff(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    config = _config(clock, stability=600)
    row = _row()
    first = _deps([], [row], clock)
    first.discover = lambda _timeout: DiscoveredState(
        observed_at=datetime(2026, 8, 11, 9, 51, tzinfo=UTC),
        headers=["vCommodityId"],
        rows=[row],
        source_fingerprints=["same-source"],
        customer_fingerprints=["same-customer"],
    )
    second = _deps([], [row], clock)
    second.discover = lambda _timeout: DiscoveredState(
        observed_at=datetime(2026, 8, 11, 10, 1, tzinfo=UTC),
        headers=["vCommodityId"],
        rows=[row],
        source_fingerprints=["same-source"],
        customer_fingerprints=["same-customer"],
        evidence={"saas_count": 1, "customer_count": 1},
    )

    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, first)
    finalized = run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        config,
        second,
        finalize=True,
    )

    snapshots = store.list_snapshots("2026-08-11")
    assert len(snapshots) == 2
    assert snapshots[0]["ordered_hash"] == snapshots[1]["ordered_hash"]
    assert snapshots[0]["observed_at"] == "2026-08-11T09:51:00+00:00"
    assert snapshots[1]["observed_at"] == "2026-08-11T10:01:00+00:00"
    assert finalized.frozen_snapshot_id == snapshots[1]["id"]


def test_persisted_external_error_is_type_only_and_does_not_store_secrets():
    message = _safe_error(RuntimeError("https://service.invalid/path?token=do-not-store"))

    assert message == "RuntimeError"
    assert "do-not-store" not in message


def test_hard_stage_crash_remains_claimed_then_next_tick_recovers(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    crashing = _deps(events, [_row()], clock)
    crashing.detail = lambda row, timeout: (_ for _ in ()).throw(SystemExit("hard stop"))
    with pytest.raises(SystemExit):
        run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), crashing)
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    assert store.list_stage_work(token, stage="DETAIL")[0]["state"] == "CLAIMED"
    store.release_lease(token)
    resumed = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), _deps(events, [_row()], clock))
    assert resumed.counters["recovered_stage_claims"] == 1


def test_unknown_submit_harvest_reconciles_exact_identity_and_persists_crawl_success(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    config = _config(clock)
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, _deps(events, [_row()], clock, post_raises=True))
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    job = store.list_remote_jobs(token, states={"SUBMIT_UNKNOWN"})[0]
    crawl = store.crawl_work_for_job(job["id"])[0]
    store.release_lease(token)
    deps = _deps(events, [_row()], clock)
    deps.harvest = lambda jobs, timeout: {
        job["id"]: {"state": "COMPLETED", "remote_identity": "exact-task", "crawl_results": {crawl["id"]: {"price": "10"}}}
    }
    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    token = store.acquire_lease("2026-08-11", "verify", 30)
    assert store.get_remote_job(job["id"])["remote_identity"] == "exact-task"
    assert store.get_crawl_work(crawl["id"])["state"] == "SUCCESS"
    store.release_lease(token)
    assert summary.outcome in {"COMPLETED", "PENDING_RETRY"}


def test_harvest_persists_request_level_no_result_without_terminal_failure(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    config = _config(clock)
    seed_deps = _deps(events, [_row()], clock)
    seed_finalize = seed_deps.finalize_crawl_plan
    seed_deps.finalize_crawl_plan = lambda plan, members, timeout: replace(
        seed_finalize(plan, members, timeout),
        input_path=str(tmp_path / "input.xlsx"),
        result_path=str(tmp_path / "result.xlsx"),
    )
    run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        config,
        seed_deps,
    )
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    job = store.list_remote_jobs(token)[0]
    crawl = store.crawl_work_for_job(job["id"])[0]
    store.release_lease(token)
    deps = _deps(events, [_row()], clock)
    deps.harvest = lambda jobs, timeout: {
        job["id"]: {
            "state": "COMPLETED",
            "remote_identity": job["remote_identity"],
            "result": {},
            "crawl_results": {
                crawl["id"]: {
                    "no_result": True,
                    "reason": "provenance-valid remote request has no result row",
                }
            },
        }
    }

    run_incremental_tick(
        store, ProcessLock(tmp_path / "tick.lock"), config, deps
    )

    no_result = store.get_crawl_work(crawl["id"])
    assert no_result["state"] == "NO_RESULT"
    assert json.loads(no_result["result_json"])["page_price_value"] is None


def test_harvest_exception_is_bounded_and_becomes_manual_not_tick_failure(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    config = TickConfig(business_date="2026-08-11", runtime_budget_seconds=60, heavy_work_reserve_seconds=5, lease_ttl_seconds=30, harvest_max_attempts=2, harvest_max_elapsed_seconds=1)
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, _deps(events, [_row()], clock))
    failing = _deps(events, [_row()], clock)
    failing.harvest = lambda jobs, timeout: (_ for _ in ()).throw(TimeoutError("harvest timeout"))
    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, failing)
    token = store.acquire_lease("2026-08-11", "verify", 30)
    assert store.list_remote_jobs(token)[0]["state"] == "MANUAL_REQUIRED"
    store.release_lease(token)
    assert summary.outcome == "MANUAL_REVIEW_REQUIRED"


def _seed_priority_pending_job(store, business_date, identity):
    store.ensure_run(business_date)
    token = store.acquire_lease(business_date, "seed-priority-pending", 60)
    job = store.prepare_remote_intent(
        token,
        "priority-pending-logical",
        "priority-pending-task",
        "priority-pending-payload",
    )
    store.claim_remote_submission(token, int(job["id"]))
    job = store.reconcile_remote_job(
        token,
        int(job["id"]),
        remote_identity=identity,
        state="PRIORITY_PENDING",
        result={"state": "PRIORITY_PENDING", "remote_identity": identity},
    )
    store.release_lease(token)
    return job


def test_retryable_priority_failure_stays_pending_then_succeeds_on_same_job(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    identity = "jd|1242978|RW-20260818110643"
    job = _seed_priority_pending_job(store, "2026-08-11", identity)
    attempts = []
    deps = _deps([], [], clock)

    def harvest(jobs, _timeout):
        current = jobs[0]
        attempts.append((int(current["id"]), str(current["remote_identity"])))
        if len(attempts) == 1:
            return {
                int(job["id"]): {
                    "state": "PRIORITY_PENDING",
                    "remote_identity": identity,
                    "retryable_error": "InterfaceError",
                    "result": {
                        "state": "PRIORITY_PENDING",
                        "priority_status": "PRIORITY_UNCONFIRMED",
                    },
                }
            }
        return {
            int(job["id"]): {
                "state": "SUBMITTED",
                "remote_identity": identity,
                "result": {"priority_status": "PRIORITY_UPDATED"},
            }
        }

    deps.harvest = harvest
    deps.submit_crawl = lambda *_args: pytest.fail(
        "retrying existing priority must never call /add"
    )
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30,
        harvest_max_attempts=3,
        harvest_max_elapsed_seconds=60,
    )

    run_incremental_tick(store, ProcessLock(tmp_path / "tick-1.lock"), config, deps)
    first = store.get_remote_job(int(job["id"]))
    assert first["state"] == "PRIORITY_PENDING"
    assert first["remote_identity"] == identity
    assert first["harvest_attempts"] == 1
    assert first["error_text"] == "InterfaceError"

    run_incremental_tick(store, ProcessLock(tmp_path / "tick-2.lock"), config, deps)
    second = store.get_remote_job(int(job["id"]))
    assert second["state"] == "SUBMITTED"
    assert second["remote_identity"] == identity
    assert second["harvest_attempts"] == 2
    assert second["error_text"] is None
    assert attempts == [(int(job["id"]), identity), (int(job["id"]), identity)]


def test_terminal_priority_pending_export_completes_without_priority_confirmation(
    tmp_path,
):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    identity = "jd|1242978|RW-20260818110643"
    job = _seed_priority_pending_job(store, "2026-08-11", identity)
    deps = _deps([], [], clock)
    deps.harvest = lambda _jobs, _timeout: {
        int(job["id"]): {
            "state": "COMPLETED",
            "remote_identity": identity,
            "priority_terminal_bypass": True,
            "remote_terminal_progress": {
                "total": 1,
                "pending": 0,
                "running": 0,
                "failed": 0,
                "waiting_db": 0,
                "finished": 1,
            },
            "result": {"https://same.invalid/p": {"page_price_value": 11}},
            "crawl_results": {},
        }
    }
    deps.submit_crawl = lambda *_args: pytest.fail(
        "terminal priority bypass must never call /add"
    )
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30,
        harvest_max_attempts=3,
        harvest_max_elapsed_seconds=60,
    )

    run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        config,
        deps,
    )

    completed = store.get_remote_job(int(job["id"]))
    assert completed["state"] == "COMPLETED"
    assert completed["remote_identity"] == identity
    assert completed["error_text"] is None


@pytest.mark.parametrize(
    (
        "max_attempts",
        "max_elapsed_seconds",
        "attempt_elapsed",
        "pending_attempts",
    ),
    [
        (3, 60, 0.1, 2),
        (10, 1, 0.6, 1),
    ],
)
def test_retryable_priority_failure_becomes_manual_only_at_harvest_limit(
    tmp_path,
    max_attempts,
    max_elapsed_seconds,
    attempt_elapsed,
    pending_attempts,
):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    identity = "jd|1242978|RW-20260818110643"
    job = _seed_priority_pending_job(store, "2026-08-11", identity)
    deps = _deps([], [], clock)
    deps.harvest = lambda _jobs, _timeout: {
        int(job["id"]): {
            "state": "PRIORITY_PENDING",
            "remote_identity": identity,
            "retryable_error": "InterfaceError",
            "elapsed_seconds": attempt_elapsed,
            "result": {"state": "PRIORITY_PENDING"},
        }
    }
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=0,
        lease_ttl_seconds=60,
        harvest_max_attempts=max_attempts,
        harvest_max_elapsed_seconds=max_elapsed_seconds,
    )
    token = store.acquire_lease("2026-08-11", "priority-limit", 60)
    try:
        for _ in range(pending_attempts):
            _harvest(
                store,
                token,
                deps,
                TickBudget(clock.monotonic(), 60, 0, clock.monotonic),
                {"crawl_harvested": 0},
                config,
            )
            pending = store.get_remote_job(int(job["id"]))
            assert pending["state"] == "PRIORITY_PENDING"
            assert pending["remote_identity"] == identity

        _harvest(
            store,
            token,
            deps,
            TickBudget(clock.monotonic(), 60, 0, clock.monotonic),
            {"crawl_harvested": 0},
            config,
        )
        limited = store.get_remote_job(int(job["id"]))
    finally:
        store.release_lease(token)

    assert limited["state"] == "MANUAL_REQUIRED"
    assert limited["remote_identity"] == identity
    assert limited["error_text"] == "InterfaceError"


def test_default_priority_retry_bound_matches_remote_harvest_bound(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    identity = "jd|1242978|RW-20260818110643"
    job = _seed_priority_pending_job(store, "2026-08-11", identity)
    deps = _deps([], [], clock)
    deps.harvest = lambda _jobs, _timeout: {
        int(job["id"]): {
            "state": "PRIORITY_PENDING",
            "remote_identity": identity,
            "retryable_error": "OperationalError[2006]",
            "elapsed_seconds": 0.01,
            "result": {"state": "PRIORITY_PENDING"},
        }
    }
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=0,
        lease_ttl_seconds=60,
    )
    token = store.acquire_lease("2026-08-11", "priority-default-limit", 60)
    try:
        for expected_attempts in (1, 2, 3, 4, 5):
            _harvest(
                store,
                token,
                deps,
                TickBudget(clock.monotonic(), 60, 0, clock.monotonic),
                {"crawl_harvested": 0},
                config,
            )
            pending = store.get_remote_job(int(job["id"]))
            assert pending["state"] == "PRIORITY_PENDING"
            assert pending["harvest_attempts"] == expected_attempts
    finally:
        store.release_lease(token)

    assert config.priority_retry_max_attempts == config.harvest_max_attempts == 120
    assert (
        config.priority_retry_max_elapsed_seconds
        == config.harvest_max_elapsed_seconds
    )


def test_priority_success_on_last_permitted_attempt_wins_over_retry_bound(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    identity = "jd|1242978|RW-20260818110643"
    job = _seed_priority_pending_job(store, "2026-08-11", identity)
    calls = 0
    deps = _deps([], [], clock)

    def harvest(_jobs, _timeout):
        nonlocal calls
        calls += 1
        if calls < 3:
            return {
                int(job["id"]): {
                    "state": "PRIORITY_PENDING",
                    "remote_identity": identity,
                    "retryable_error": "InterfaceError",
                    "elapsed_seconds": 0.01,
                    "result": {"state": "PRIORITY_PENDING"},
                }
            }
        return {
            int(job["id"]): {
                "state": "SUBMITTED",
                "remote_identity": identity,
                "elapsed_seconds": 0.01,
                "result": {"priority_status": "PRIORITY_UPDATED"},
            }
        }

    deps.harvest = harvest
    deps.submit_crawl = lambda *_args: pytest.fail(
        "retrying existing priority must never call /add"
    )
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=0,
        lease_ttl_seconds=60,
    )
    token = store.acquire_lease("2026-08-11", "priority-success-at-limit", 60)
    try:
        for _ in range(3):
            _harvest(
                store,
                token,
                deps,
                TickBudget(clock.monotonic(), 60, 0, clock.monotonic),
                {"crawl_harvested": 0},
                config,
            )
        completed = store.get_remote_job(int(job["id"]))
    finally:
        store.release_lease(token)

    assert completed["harvest_attempts"] == 3
    assert completed["state"] == "SUBMITTED"
    assert completed["remote_identity"] == identity


def test_deterministic_operational_error_code_is_manual_on_first_priority_attempt(
    tmp_path,
):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    identity = "jd|1242978|RW-20260818110643"
    job = _seed_priority_pending_job(store, "2026-08-11", identity)
    error = OperationalError(1045, "message must not affect classification")
    deps = _deps([], [], clock)

    def harvest(_jobs, _timeout):
        diagnostic = retryable_priority_error_diagnostic(error)
        if diagnostic is not None:
            return {
                int(job["id"]): {
                    "state": "PRIORITY_PENDING",
                    "remote_identity": identity,
                    "retryable_error": diagnostic,
                }
            }
        return {int(job["id"]): {"manual_error": type(error).__name__}}

    deps.harvest = harvest
    deps.submit_crawl = lambda *_args: pytest.fail(
        "deterministic priority failure must never call /add"
    )
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=0,
        lease_ttl_seconds=60,
    )
    token = store.acquire_lease("2026-08-11", "priority-deterministic", 60)
    try:
        _harvest(
            store,
            token,
            deps,
            TickBudget(clock.monotonic(), 60, 0, clock.monotonic),
            {"crawl_harvested": 0},
            config,
        )
        failed = store.get_remote_job(int(job["id"]))
    finally:
        store.release_lease(token)

    assert failed["harvest_attempts"] == 1
    assert failed["state"] == "MANUAL_REQUIRED"
    assert failed["remote_identity"] == identity
    assert failed["error_text"] == "OperationalError"


def test_second_tick_submits_only_delta_and_preserves_completed_old_crawl(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    submitted: list[list[str]] = []
    first_rows = [_row("old", "old-sku"), _row("old", "old-sku")]
    second_rows = [*first_rows, _row("new", "new-sku")]

    def dependencies(rows, harvest=None):
        deps = _deps([], rows, clock)
        def plan(items, timeout):
            members = [CrawlMember("mall-a", item["row"]["vCommodityId"], item["row"]["skuId"], item["row"]["extLink"], "legacy-v1", item["row"]) for item in items]
            suffix = str(len(members))
            return [CrawlPlan("mall-a", f"mall-a:{suffix}", f"task-{suffix}", f"payload-{suffix}", members=members)]
        deps.plan_crawl = plan
        deps.submit_crawl = lambda job, crawls, timeout: (submitted.append([crawl["sku_id"] for crawl in crawls]) or {"remote_identity": f"remote-{job['id']}", "state": "SUBMITTED"})
        if harvest is not None:
            deps.harvest = harvest
        return deps

    config = _config(clock)
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, dependencies(first_rows))
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    old_job = store.list_remote_jobs(token)[0]
    old_crawl = store.crawl_work_for_job(old_job["id"])[0]
    store.release_lease(token)
    harvest = lambda jobs, timeout: {old_job["id"]: {"state": "COMPLETED", "remote_identity": old_job["remote_identity"], "crawl_results": {old_crawl["id"]: {"price": "10"}}}}
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, dependencies(second_rows, harvest))
    token = store.acquire_lease("2026-08-11", "verify", 30)
    assert submitted == [["old-sku"], ["new-sku"]]
    assert store.get_crawl_work(old_crawl["id"])["remote_job_id"] == old_job["id"]
    assert store.get_crawl_work(old_crawl["id"])["state"] == "SUCCESS"
    jobs = store.list_remote_jobs(token)
    assert len(jobs) == 2
    new_job = next(job for job in jobs if job["id"] != old_job["id"])
    assert new_job["payload_hash"] == "payload-new-sku"
    assert new_job["task_name"] == "tick-new-sku"
    assert new_job["payload_json"] == '{"members":["new-sku"]}'
    store.release_lease(token)


def test_frozen_ordered_duplicates_materialize_equivalent_legacy_record_results(tmp_path):
    rows = [_row("doc", "sku"), _row("doc", "sku"), _row("other", "other-sku")]
    enriched = [{"customerPrice": "10", "customerPriceStatus": "EXISTS", "customerPriceSource": "fixture", "customerPriceWarning": "", "customerPriceEvidence": ""} for _ in rows]
    records = build_records(rows, enriched)
    response = audit_saas_approval_records(records, use_model=False, use_link_audit=False, crawl_results=[], use_crawl_result=False)
    direct = tmp_path / "direct.xlsx"
    legacy_excel.write_audit_workbook(direct, list(rows[0]), rows, enriched, response)
    store = DailyAuditStore(tmp_path / "state.sqlite")
    token = store.acquire_lease("2026-08-11", "materialize", 30)
    snapshot = store.ingest_snapshot(token, rows)["snapshot"]
    store.freeze_snapshot(token, snapshot["id"])
    frozen_rows = [item["row"] for item in store.ordered_snapshot_items(snapshot["id"])]
    frozen_records = build_records(frozen_rows, enriched)
    frozen_response = audit_saas_approval_records(frozen_records, use_model=False, use_link_audit=False, crawl_results=[], use_crawl_result=False)
    materialized = tmp_path / "materialized.xlsx"
    legacy_excel.write_audit_workbook(materialized, list(rows[0]), frozen_rows, enriched, frozen_response)
    def values(path):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb["record_results"]
            return [tuple(row) for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    assert values(materialized) == values(direct)


def test_harvest_at_reserve_boundary_checkpoints_without_discovery_or_new_work(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    deps = _deps(events, [_row()], clock)
    def exhausting_harvest(jobs, timeout):
        events.append("harvest")
        clock.advance(50)
        return {}
    deps.harvest = exhausting_harvest
    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock, budget=60, reserve=10), deps, finalize=True)
    assert summary.outcome == "BUDGET_EXHAUSTED"
    assert summary.phase_trace == ("recover", "harvest", "checkpoint")
    assert events == ["harvest"]
    assert store.get_run("2026-08-11") is not None


def test_post_discovery_reserve_stops_stage_preprocessing_but_counts_uninitialized_backlog(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    rows = [_row(f"doc-{index}", f"sku-{index}") for index in range(3)]
    deps = _deps([], rows, clock)
    original_discover = deps.discover

    def exhausting_discover(timeout):
        discovered = original_discover(timeout)
        clock.advance(55)
        return discovered

    deps.discover = exhausting_discover

    summary = run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        _config(clock, budget=60, reserve=10),
        deps,
    )

    token = store.acquire_lease("2026-08-11", "inspect", 30)
    stages = store.list_stage_work(token)
    store.release_lease(token)
    assert stages == []
    assert summary.backlog["stage_pending"] == 6
    assert summary.outcome == "BUDGET_EXHAUSTED"


def test_model_preprocessing_stops_traversing_when_local_work_reaches_reserve(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    rows = [_row(f"doc-{index}", f"sku-{index}", queryStr=f"title-{index}") for index in range(4)]
    deps = _deps([], rows, clock)
    deps.detail = None
    requests: list[list[str]] = []
    deps.model_batch = lambda batch, timeout: (
        requests.append([str(row["skuId"]) for row in batch])
        or {index: {"match_result": "SAME"} for index, _row_value in enumerate(batch)}
    )
    original_ensure_batch = store.ensure_stage_work_batch

    def advancing_ensure_batch(token, entries):
        works = original_ensure_batch(token, entries)
        clock.advance(6 * len(entries))
        return works

    store.ensure_stage_work_batch = advancing_ensure_batch

    summary = run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=20,
            heavy_work_reserve_seconds=5,
            lease_ttl_seconds=30,
        ),
        deps,
    )

    token = store.acquire_lease("2026-08-11", "inspect", 30)
    models = store.list_stage_work(token, stage="MODEL")
    store.release_lease(token)
    # The batched grouping transaction advances the clock past the reserve, so
    # no MODEL batch is dispatched at all; the durable rows stay claimable for
    # the next tick.
    assert len(models) == 4
    assert requests == []
    assert summary.backlog["stage_pending"] == 8
    # MODEL traversal hit the reserve, but the crawl submissions made before
    # MODEL leave durable remote work pending, so the tick reports the remote
    # wait rather than a pure budget stop.
    assert summary.outcome == "WAITING_REMOTE"


@pytest.mark.parametrize(
    ("runtime", "external", "stage_limit", "expected"),
    [(30, 7, 20, 7), (30, 60, 11, 11), (5, 60, 20, 5)],
)
def test_stage_callback_timeout_is_minimum_of_external_stage_and_tick_limits(
    tmp_path, runtime, external, stage_limit, expected,
):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    observed: list[float] = []
    deps = _deps([], [_row()], clock)
    deps.detail = lambda row, timeout: (observed.append(timeout) or {"customerPrice": "10"})
    deps.model = lambda row, timeout: (observed.append(timeout) or {"match_result": "SAME"})
    config = TickConfig(
        business_date="2026-08-11", runtime_budget_seconds=runtime, heavy_work_reserve_seconds=0,
        lease_ttl_seconds=30, external_timeout_seconds=external,
        detail_max_elapsed_seconds=stage_limit, model_max_elapsed_seconds=stage_limit,
    )

    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)

    assert observed[:2] == [expected, expected]


def test_snapshot_discovery_has_180_second_budget_without_expanding_other_callbacks(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    observed: dict[str, list[float]] = {
        "discover": [], "harvest": [], "detail": [], "model": [],
        "plan": [], "finalize_plan": [], "submit": [],
    }
    row = _row()

    def discover(timeout):
        observed["discover"].append(timeout)
        return DiscoveredState(
            observed_at=datetime(2026, 8, 11, 10, 1, tzinfo=UTC),
            headers=["vCommodityId"],
            rows=[row],
            source_fingerprints=["source"],
            customer_fingerprints=["customer"],
            evidence={"saas_count": 1, "customer_count": 1},
        )

    deps = _deps([], [row], clock)
    deps.discover = discover
    deps.harvest = lambda jobs, timeout: observed["harvest"].append(timeout) or {}
    deps.detail = lambda _row, timeout: observed["detail"].append(timeout) or {"customerPrice": "10"}
    deps.model = lambda _row, timeout: observed["model"].append(timeout) or {"match_result": "SAME"}
    original_plan = deps.plan_crawl
    original_finalize = deps.finalize_crawl_plan
    assert original_plan is not None and original_finalize is not None
    deps.plan_crawl = lambda items, timeout: observed["plan"].append(timeout) or original_plan(items, timeout)
    deps.finalize_crawl_plan = lambda plan, members, timeout: (
        observed["finalize_plan"].append(timeout)
        or original_finalize(plan, members, timeout)
    )
    deps.submit_crawl = lambda _job, _crawls, timeout: (
        observed["submit"].append(timeout)
        or {"remote_identity": "task-1", "state": "SUBMITTED"}
    )

    run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=720,
            heavy_work_reserve_seconds=120,
            lease_ttl_seconds=900,
        ),
        deps,
    )

    assert observed["discover"] == [180]
    assert observed["detail"] == [60]
    assert observed["model"] == [60]
    assert observed["plan"] == [60]
    assert observed["finalize_plan"] == [60]
    assert observed["submit"] == [60]
    assert observed["harvest"] == [60, 60]


def test_plan_finalization_crossing_reserve_does_not_prepare_or_post_remote_intent(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    prepared: list[tuple[str, list[str]]] = []
    posted: list[int] = []
    deps = _deps([], [_row()], clock)

    def finalize(plan, members, timeout):
        prepared.append((plan.mall_key, [member.sku_id for member in members]))
        clock.advance(5)
        return CrawlPlan(
            plan.mall_key, "prepared-only", "prepared-only", "prepared-only",
            payload_json={"members": [member.sku_id for member in members]}, members=tuple(members),
        )

    deps.finalize_crawl_plan = finalize
    deps.submit_crawl = lambda job, crawls, timeout: (posted.append(int(job["id"])) or {"remote_identity": "never", "state": "SUBMITTED"})
    config = TickConfig(
        business_date="2026-08-11", runtime_budget_seconds=10, heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30,
    )

    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)

    token = store.acquire_lease("2026-08-11", "inspect", 30)
    crawls = store.list_crawl_work(token)
    assert prepared == [("mall-a", ["sku-1"])]
    assert posted == []
    assert store.list_remote_jobs(token) == []
    assert crawls[0]["remote_job_id"] is None
    assert crawls[0]["state"] == "PENDING"
    store.release_lease(token)
    assert summary.outcome == "BUDGET_EXHAUSTED"


def test_linking_remote_intent_at_reserve_checkpoints_without_dispatching_post(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    posted: list[int] = []
    deps = _deps([], [_row()], clock)
    deps.submit_crawl = lambda job, crawls, timeout: (
        posted.append(int(job["id"]))
        or {"remote_identity": "must-not-post", "state": "SUBMITTED"}
    )
    original_link = store.link_crawl_work_batch_to_remote_job

    def exhausting_link(token, crawl_work_ids, job_id):
        linked = original_link(token, crawl_work_ids, job_id)
        clock.advance(5)
        return linked

    store.link_crawl_work_batch_to_remote_job = exhausting_link

    summary = run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=10,
            heavy_work_reserve_seconds=5,
            lease_ttl_seconds=30,
        ),
        deps,
    )

    token = store.acquire_lease("2026-08-11", "inspect", 30)
    jobs = store.list_remote_jobs(token)
    crawls = store.list_crawl_work(token)
    store.release_lease(token)
    assert posted == []
    assert len(jobs) == 1 and jobs[0]["state"] == "SUBMIT_INTENT"
    assert len(crawls) == 1 and crawls[0]["remote_job_id"] == jobs[0]["id"]
    assert summary.outcome == "BUDGET_EXHAUSTED"


def test_model_batch_deduplicates_stage_fingerprints_and_uses_bounded_network_calls(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    rows = [_row(f"doc-{index}", f"sku-{index}", queryStr=f"title-{index}") for index in range(25)]
    rows.append(dict(rows[0]))
    requests: list[list[str]] = []
    deps = _deps([], rows, clock)

    def model_batch(batch, timeout):
        requests.append([str(row["skuId"]) for row in batch])
        return {
            index: {"match_result": "SAME", "confidence": 0.9, "model_match": "exact"}
            for index, _row_value in enumerate(batch)
        }

    deps.model_batch = model_batch
    config = TickConfig(
        business_date="2026-08-11", runtime_budget_seconds=60, heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30, model_batch_size=10,
    )
    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)

    token = store.acquire_lease("2026-08-11", "inspect", 30)
    models = store.list_stage_work(token, stage="MODEL")
    store.release_lease(token)
    assert [len(batch) for batch in requests] == [10, 10, 5]
    assert sum(len(batch) for batch in requests) == 25
    # Exact duplicate snapshot members preserve output multiplicity but share
    # one durable audit item and therefore one model result.
    assert len(models) == 25
    assert all(work["state"] == "SUCCESS" for work in models)
    assert summary.counters["model_attempts"] == 3


def test_model_hard_crash_uses_one_durable_replay_then_succeeds(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    calls = 0
    crashing = _deps([], [_row()], clock)

    def hard_crash(batch, timeout):
        nonlocal calls
        calls += 1
        raise SystemExit("hard model crash")

    crashing.model_batch = hard_crash
    with pytest.raises(SystemExit):
        run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), crashing)

    token = store.acquire_lease("2026-08-11", "inspect-crash", 30)
    interrupted = store.list_stage_work(token, stage="MODEL")[0]
    store.release_lease(token)
    assert interrupted["state"] == "CLAIMED"
    assert interrupted["attempts"] == 1

    resumed = _deps([], [_row()], clock)

    def succeed(batch, timeout):
        nonlocal calls
        calls += 1
        return {0: {"match_result": "SAME"}}

    resumed.model_batch = succeed
    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), resumed)

    token = store.acquire_lease("2026-08-11", "inspect-resume", 30)
    completed = store.list_stage_work(token, stage="MODEL")[0]
    store.release_lease(token)
    assert calls == 2
    assert summary.counters["recovered_stage_claims"] == 1
    assert completed["state"] == "SUCCESS"
    assert completed["attempts"] == 2


@pytest.mark.parametrize("batched", [False, True])
def test_caught_model_failure_is_terminal_after_one_durable_claim(tmp_path, batched):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    calls = 0
    failing = _deps([], [_row()], clock)

    def fail(*_args):
        nonlocal calls
        calls += 1
        raise RuntimeError("bounded adapter failure")

    if batched:
        failing.model_batch = fail
    else:
        failing.model = fail
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30,
        model_max_attempts=2,
    )

    first = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, failing)
    second = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, failing)

    token = store.acquire_lease("2026-08-11", "inspect", 30)
    model = store.list_stage_work(token, stage="MODEL")[0]
    store.release_lease(token)
    assert calls == 1
    assert model["state"] == "TERMINAL_FAILED"
    assert model["attempts"] == 1
    assert first.outcome == "MANUAL_REVIEW_REQUIRED"
    assert second.outcome == "MANUAL_REVIEW_REQUIRED"


@pytest.mark.parametrize(
    "failure",
    [
        TimeoutError("model timeout"),
        ModelCallRetryExhaustedError(
            "batch model call", 1, 60.0, TimeoutError("model timeout")
        ),
        error.HTTPError("https://model.invalid", 429, "rate limited", {}, None),
        error.HTTPError("https://model.invalid", 500, "unavailable", {}, None),
        ConnectionResetError("connection reset"),
    ],
    ids=("timeout", "retry-exhausted-timeout", "http-429", "http-500", "connection-reset"),
)
def test_transient_model_failure_is_retryable_by_the_next_tick(tmp_path, failure):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    calls = 0
    deps = _deps([], [_row()], clock)

    def model_batch(_batch, _timeout):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise failure
        return {0: {"match_result": "SAME"}}

    deps.model_batch = model_batch
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30,
        model_max_attempts=2,
    )

    first = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    token = store.acquire_lease("2026-08-11", "inspect-first", 30)
    after_first = store.list_stage_work(token, stage="MODEL")[0]
    store.release_lease(token)
    second = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    token = store.acquire_lease("2026-08-11", "inspect-second", 30)
    after_second = store.list_stage_work(token, stage="MODEL")[0]
    store.release_lease(token)

    assert first.outcome == "WAITING_REMOTE"
    assert after_first["state"] == "RETRYABLE_FAILED"
    assert after_first["retryable"] == 1
    assert after_first["attempts"] == 1
    assert second.outcome == "WAITING_REMOTE"
    assert after_second["state"] == "SUCCESS"
    assert after_second["attempts"] == 2
    assert calls == 2


@pytest.mark.parametrize(
    "failure",
    [
        error.HTTPError("https://model.invalid", 400, "bad request", {}, None),
        ValueError("permanent business validation"),
        json.JSONDecodeError("invalid model contract", "not-json", 0),
    ],
    ids=("http-400", "business-validation", "parser-contract"),
)
def test_permanent_model_failure_remains_terminal(tmp_path, failure):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    calls = 0
    deps = _deps([], [_row()], clock)

    def model_batch(_batch, _timeout):
        nonlocal calls
        calls += 1
        raise failure

    deps.model_batch = model_batch
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30,
        model_max_attempts=2,
    )

    first = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    second = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    model = store.list_stage_work(token, stage="MODEL")[0]
    store.release_lease(token)

    assert first.outcome == "MANUAL_REVIEW_REQUIRED"
    assert second.outcome == "MANUAL_REVIEW_REQUIRED"
    assert model["state"] == "TERMINAL_FAILED"
    assert model["attempts"] == 1
    assert calls == 1


def test_transient_model_retry_stops_at_the_durable_attempt_ceiling(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    calls = 0
    deps = _deps([], [_row()], clock)

    def model_batch(_batch, _timeout):
        nonlocal calls
        calls += 1
        raise TimeoutError("still unavailable")

    deps.model_batch = model_batch
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30,
        model_max_attempts=2,
    )

    first = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    second = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    third = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    model = store.list_stage_work(token, stage="MODEL")[0]
    store.release_lease(token)

    assert first.outcome == "WAITING_REMOTE"
    assert second.outcome == "MANUAL_REVIEW_REQUIRED"
    assert third.outcome == "MANUAL_REVIEW_REQUIRED"
    assert model["state"] == "TERMINAL_FAILED"
    assert model["attempts"] == 2
    assert calls == 2


def test_model_retry_success_plans_once_without_reclaiming_existing_success(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    model_batches: list[list[str]] = []
    submit_calls = 0
    deps = _deps(
        [],
        [
            _row("doc-retry", "sku-retry", queryStr="retry title"),
            _row("doc-ok", "sku-ok", queryStr="already successful title"),
        ],
        clock,
    )

    retry_calls = 0

    def model_batch(batch, _timeout):
        nonlocal retry_calls
        identities = [str(row["vCommodityId"]) for row in batch]
        model_batches.append(identities)
        if identities == ["doc-retry"]:
            retry_calls += 1
            if retry_calls == 1:
                raise TimeoutError("transient")
        return {0: {"match_result": "SAME"}}

    deps.model_batch = model_batch
    def submit(job, crawls, timeout):
        nonlocal submit_calls
        submit_calls += 1
        return {"remote_identity": f"task-{job['id']}", "state": "SUBMITTED"}

    deps.submit_crawl = submit
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30,
        model_max_attempts=2,
        model_batch_size=1,
    )

    first = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    second = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    third = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    models = store.list_stage_work(token, stage="MODEL")
    crawls = store.list_crawl_work(token)
    jobs = store.list_remote_jobs(token)
    store.release_lease(token)

    assert first.outcome == "WAITING_REMOTE"
    assert second.outcome == "WAITING_REMOTE"
    assert third.outcome == "WAITING_REMOTE"
    assert model_batches == [["doc-retry"], ["doc-ok"], ["doc-retry"]]
    assert all(work["state"] == "SUCCESS" for work in models)
    assert len(crawls) == 2
    # Crawl planning happens before MODEL, so both rows were planned and
    # submitted together in the first tick; the later MODEL retry success must
    # not re-plan or re-submit anything.
    assert len(jobs) == 1
    assert submit_calls == 1
    assert all(crawl["remote_job_id"] == jobs[0]["id"] for crawl in crawls)


def test_retryable_model_backlog_survives_a_later_snapshot_and_plans_once(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    retry_row = _row("doc-retry", "sku-retry", queryStr="retry title")
    current_row = _row("doc-current", "sku-current", queryStr="current title")
    snapshots = [[retry_row], [current_row], [current_row]]
    model_batches: list[list[str]] = []
    submit_calls = 0
    deps = _deps([], snapshots[0], clock)

    def discover(_timeout):
        rows = snapshots.pop(0)
        return DiscoveredState(
            observed_at=datetime(2026, 8, 11, 10, len(model_batches) + 1, tzinfo=UTC),
            headers=["vCommodityId"],
            rows=rows,
            source_fingerprints=[f"source-{row['vCommodityId']}" for row in rows],
            customer_fingerprints=[f"customer-{row['vCommodityId']}" for row in rows],
        )

    def model_batch(batch, _timeout):
        identities = [str(row["vCommodityId"]) for row in batch]
        model_batches.append(identities)
        if model_batches == [["doc-retry"]]:
            raise TimeoutError("transient")
        return {
            index: {"match_result": "SAME"}
            for index, _row_value in enumerate(batch)
        }

    def submit(job, _crawls, _timeout):
        nonlocal submit_calls
        submit_calls += 1
        return {"remote_identity": f"task-{job['id']}", "state": "SUBMITTED"}

    deps.discover = discover
    deps.model_batch = model_batch
    deps.submit_crawl = submit
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30,
        model_max_attempts=2,
    )

    first = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    second = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)
    third = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)

    token = store.acquire_lease("2026-08-11", "inspect-cross-snapshot", 30)
    models = store.list_stage_work(token, stage="MODEL")
    crawls = store.list_crawl_work(token)
    jobs = store.list_remote_jobs(token)
    store.release_lease(token)

    assert first.outcome == "WAITING_REMOTE"
    assert second.outcome == "WAITING_REMOTE"
    assert third.outcome == "WAITING_REMOTE"
    assert model_batches == [["doc-retry"], ["doc-current", "doc-retry"]]
    assert sorted((work["attempts"], work["state"]) for work in models) == [
        (1, "SUCCESS"),
        (2, "SUCCESS"),
    ]
    assert {(crawl["source_id"], crawl["sku_id"]) for crawl in crawls} == {
        ("doc-current", "sku-current"),
        ("doc-retry", "sku-retry"),
    }
    # Each snapshot planned its own crawl work exactly once (tick 1: the retry
    # row, tick 2: the newly discovered current row); the MODEL retry backlog
    # never re-triggers a plan or a second remote submit for either row.
    assert len(jobs) == 2
    assert submit_calls == 2


def test_linked_submit_intent_crash_resumes_exactly_once_without_replanning(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    deps = _deps(events, [_row()], clock)
    original_claim = store.claim_remote_submission
    store.claim_remote_submission = lambda *args, **kwargs: (_ for _ in ()).throw(SystemExit("after durable links"))
    with pytest.raises(SystemExit):
        run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), deps)
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    interrupted = store.list_remote_jobs(token, states={"SUBMIT_INTENT"})[0]
    bound = store.crawl_work_for_job(interrupted["id"])
    assert len(bound) == 1 and bound[0]["remote_job_id"] == interrupted["id"]
    store.release_lease(token)
    store.claim_remote_submission = original_claim

    resumed = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), deps)
    token = store.acquire_lease("2026-08-11", "verify", 30)
    jobs = store.list_remote_jobs(token)
    store.release_lease(token)
    assert events.count("submit") == 1
    assert len(jobs) == 1
    assert jobs[0]["id"] == interrupted["id"]
    assert jobs[0]["payload_hash"] == interrupted["payload_hash"]
    assert resumed.outcome == "WAITING_REMOTE"


def test_last_allowed_harvest_persists_terminal_result_instead_of_manual_review(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    config = TickConfig(
        business_date="2026-08-11", runtime_budget_seconds=60, heavy_work_reserve_seconds=5,
        lease_ttl_seconds=30, harvest_max_attempts=2, harvest_max_elapsed_seconds=60,
    )
    first = _deps([], [_row()], clock)
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, first)
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    job = store.list_remote_jobs(token)[0]
    crawl = store.crawl_work_for_job(job["id"])[0]
    store.release_lease(token)
    completed = _deps([], [_row()], clock)
    completed.harvest = lambda jobs, timeout: {
        job["id"]: {"state": "COMPLETED", "remote_identity": job["remote_identity"], "crawl_results": {crawl["id"]: {"price": "10"}}}
    }

    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, completed)

    token = store.acquire_lease("2026-08-11", "verify", 30)
    assert store.get_remote_job(job["id"])["state"] == "COMPLETED"
    assert store.get_crawl_work(crawl["id"])["state"] == "SUCCESS"
    store.release_lease(token)


def test_harvest_claims_only_bounded_batch_and_uses_one_callback_deadline(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    for index in range(100):
        job = store.prepare_remote_intent(token, f"job-{index}", f"job-{index}", f"payload-{index}")
        store.claim_remote_submission(token, job["id"])
        store.reconcile_remote_job(token, job["id"], remote_identity=f"remote-{index}", state="SUBMITTED")
    store.release_lease(token)
    observed: list[tuple[int, float]] = []
    deps = _deps([], [], clock)
    def harvest(jobs, timeout):
        observed.append((len(jobs), timeout))
        clock.advance(timeout)
        return {}
    deps.harvest = harvest
    config = TickConfig(
        business_date="2026-08-11", runtime_budget_seconds=30, heavy_work_reserve_seconds=1,
        lease_ttl_seconds=60, external_timeout_seconds=10, harvest_batch_size=6,
    )

    # A sweep keeps issuing bounded batches (never more than 6 jobs and one
    # shared deadline per remote request) until the tick budget is spent; the
    # old behavior stopped the whole round after the first batch of 6.
    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)

    token = store.acquire_lease("2026-08-11", "verify", 60)
    attempts = [job["harvest_attempts"] for job in store.list_remote_jobs(token)]
    store.release_lease(token)
    assert observed == [(6, 10), (6, 10), (6, 10)]
    assert sum(attempts) == 18
    assert summary.outcome == "BUDGET_EXHAUSTED"
    assert summary.counters["harvest_batches"] == 3
    assert summary.counters["harvest_swept"] == 18


def test_harvest_fairly_selects_less_attempted_jobs_without_exceeding_batch(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    jobs = []
    for index in range(12):
        job = store.prepare_remote_intent(
            token, f"job-{index}", f"job-{index}", f"payload-{index}"
        )
        store.claim_remote_submission(token, int(job["id"]))
        job = store.reconcile_remote_job(
            token,
            int(job["id"]),
            remote_identity=f"remote-{index}",
            state="SUBMITTED",
        )
        jobs.append(job)
    for job in jobs[:6]:
        claimed = store.claim_harvest(
            token, int(job["id"]), max_attempts=120, max_elapsed_seconds=21600
        )
        assert claimed is not None
        store.reconcile_remote_job(
            token,
            int(job["id"]),
            remote_identity=str(job["remote_identity"]),
            state="RUNNING",
        )

    observed: list[list[int]] = []
    deps = TickDependencies(
        discover=lambda _timeout: pytest.fail("isolated harvest must not discover"),
        harvest=lambda selected, _timeout: observed.append(
            [int(job["id"]) for job in selected]
        ) or {},
        monotonic=clock.monotonic,
    )
    # One sweep still prefers the least-attempted jobs first, but now checks
    # every eligible job once (two bounded batches) instead of only six.
    _harvest(
        store,
        token,
        deps,
        TickBudget(clock.monotonic(), 60, 5, clock.monotonic),
        {},
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=60,
            heavy_work_reserve_seconds=5,
            harvest_batch_size=6,
        ),
    )

    assert observed == [
        [int(job["id"]) for job in jobs[6:]],
        [int(job["id"]) for job in jobs[:6]],
    ]
    assert all(len(batch) <= 6 for batch in observed)
    store.release_lease(token)
    token = store.acquire_lease("2026-08-11", "verify", 60)
    attempts = {int(job["id"]): int(job["harvest_attempts"]) for job in store.list_remote_jobs(token)}
    store.release_lease(token)
    # The lead six entered the sweep with one prior attempt and were checked
    # again (attempts == 2); the untouched six were reached for the first time.
    lead_ids = {int(job["id"]) for job in jobs[:6]}
    assert all(
        attempts[job_id] == (2 if job_id in lead_ids else 1) for job_id in attempts
    )


def _seed_submitted_jobs(store, business_date, count, *, prefix="job"):
    store.ensure_run(business_date)
    token = store.acquire_lease(business_date, "seed", 60)
    jobs = []
    for index in range(count):
        job = store.prepare_remote_intent(
            token, f"{prefix}-{index}", f"{prefix}-{index}", f"payload-{index}"
        )
        store.claim_remote_submission(token, int(job["id"]))
        jobs.append(
            store.reconcile_remote_job(
                token,
                int(job["id"]),
                remote_identity=f"remote-{prefix}-{index}",
                state="SUBMITTED",
            )
        )
    store.release_lease(token)
    return jobs


def _sweep_deps(clock, harvest):
    return TickDependencies(
        discover=lambda _timeout: pytest.fail("isolated sweep must not discover"),
        harvest=harvest,
        monotonic=clock.monotonic,
    )


def test_harvest_sweep_checks_every_eligible_job_exactly_once_in_batches(tmp_path):
    # Case 1: one sweep over 34 eligible jobs issues 6+6+6+6+6+4 checks and
    # every remote job is checked exactly once.
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    jobs = _seed_submitted_jobs(store, "2026-08-11", 34)
    observed: list[list[int]] = []

    def harvest(selected, _timeout):
        observed.append([int(job["id"]) for job in selected])
        return {}

    token = store.acquire_lease("2026-08-11", "sweep", 60)
    _harvest(
        store,
        token,
        _sweep_deps(clock, harvest),
        TickBudget(clock.monotonic(), 600, 10, clock.monotonic),
        {"crawl_harvested": 0},
        TickConfig(business_date="2026-08-11", runtime_budget_seconds=600),
    )
    store.release_lease(token)

    assert [len(batch) for batch in observed] == [6, 6, 6, 6, 6, 4]
    flattened = [job_id for batch in observed for job_id in batch]
    assert sorted(flattened) == sorted(int(job["id"]) for job in jobs)
    assert len(flattened) == len(set(flattened)) == 34


def test_harvest_sweep_continues_after_unchanged_running_lead_batch(tmp_path):
    # Case 2: when the first batch of six is still RUNNING with no further
    # progress, the remaining 28 eligible jobs must still be checked in the
    # same sweep instead of waiting for the next scheduler slot.
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    jobs = _seed_submitted_jobs(store, "2026-08-11", 34)
    observed: list[list[int]] = []
    calls = 0

    def harvest(selected, _timeout):
        nonlocal calls
        calls += 1
        observed.append([int(job["id"]) for job in selected])
        if calls == 1:
            return {
                int(job["id"]): {
                    "state": "RUNNING",
                    "remote_identity": str(job["remote_identity"]),
                }
                for job in selected
            }
        return {}

    token = store.acquire_lease("2026-08-11", "sweep", 60)
    _harvest(
        store,
        token,
        _sweep_deps(clock, harvest),
        TickBudget(clock.monotonic(), 600, 10, clock.monotonic),
        {"crawl_harvested": 0},
        TickConfig(business_date="2026-08-11", runtime_budget_seconds=600),
    )
    store.release_lease(token)

    assert [len(batch) for batch in observed] == [6, 6, 6, 6, 6, 4]
    lead_ids = {int(job["id"]) for job in jobs[:6]}
    token = store.acquire_lease("2026-08-11", "verify", 60)
    states = {int(job["id"]): str(job["state"]) for job in store.list_remote_jobs(token)}
    store.release_lease(token)
    assert all(states[job_id] == "RUNNING" for job_id in lead_ids)
    assert sum(state == "RUNNING" for state in states.values()) == 6


def test_harvest_sweep_completes_jobs_through_existing_export_ingest(tmp_path):
    # Case 3: jobs that complete mid-sweep flow through the existing
    # reconcile + complete_crawl_work path with durable per-crawl evidence.
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    jobs = _seed_submitted_jobs(store, "2026-08-11", 12)
    token = store.acquire_lease("2026-08-11", "seed-crawls", 60)
    crawl_ids = {}
    for index, job in enumerate(jobs):
        crawl = store.ensure_crawl_work(
            token,
            f"mall-{index % 3}",
            f"sku-{index}",
            f"https://shop.example/item/{index}",
            "legacy-v1",
            f"fingerprint-{index}",
            source_id=f"v-{index}",
        )
        store.link_crawl_work_to_remote_job(token, int(crawl["id"]), int(job["id"]))
        crawl_ids[int(job["id"])] = int(crawl["id"])
    store.release_lease(token)
    completing = {int(job["id"]) for job in jobs[:6]}

    def harvest(selected, _timeout):
        results = {}
        for job in selected:
            job_id = int(job["id"])
            if job_id not in completing:
                continue
            results[job_id] = {
                "state": "COMPLETED",
                "remote_identity": str(job["remote_identity"]),
                "export_path": str(tmp_path / f"result-{job_id}.xlsx"),
                "crawl_results": {
                    crawl_ids[job_id]: {
                        "page_price_status": "OK",
                        "page_price_value": "12.30",
                        "page_price_source": "招商局",
                    }
                },
            }
        return results

    token = store.acquire_lease("2026-08-11", "sweep", 60)
    _harvest(
        store,
        token,
        _sweep_deps(clock, harvest),
        TickBudget(clock.monotonic(), 600, 10, clock.monotonic),
        {"crawl_harvested": 0},
        TickConfig(business_date="2026-08-11", runtime_budget_seconds=600),
    )
    store.release_lease(token)

    for job in jobs:
        job_state = store.get_remote_job(int(job["id"]))["state"]
        crawl_state = store.get_crawl_work(crawl_ids[int(job["id"])])["state"]
        if int(job["id"]) in completing:
            assert job_state == "COMPLETED"
            assert crawl_state == "SUCCESS"
            result = json.loads(
                str(store.get_crawl_work(crawl_ids[int(job["id"])])["result_json"])
            )
            assert result["page_price_value"] == "12.30"
        else:
            assert job_state == "SUBMITTED"
            assert crawl_state == "PENDING"


def test_harvest_sweep_never_polls_one_job_twice_in_the_same_round(tmp_path):
    # Case 4: a PRIORITY_PENDING job prioritized into SUBMITTED by an early
    # batch stays harvestable, but the same sweep must not poll it again.
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    jobs = []
    for index in range(7):
        job = store.prepare_remote_intent(
            token, f"mixed-{index}", f"mixed-{index}", f"payload-{index}"
        )
        store.claim_remote_submission(token, int(job["id"]))
        state = "PRIORITY_PENDING" if index == 0 else "SUBMITTED"
        jobs.append(
            store.reconcile_remote_job(
                token,
                int(job["id"]),
                remote_identity=f"remote-mixed-{index}",
                state=state,
            )
        )
    store.release_lease(token)
    observed: list[list[int]] = []

    def harvest(selected, _timeout):
        observed.append([int(job["id"]) for job in selected])
        return {
            int(job["id"]): {
                "state": "SUBMITTED",
                "remote_identity": str(job["remote_identity"]),
            }
            for job in selected
            if str(job["state"]) == "PRIORITY_PENDING"
        }

    token = store.acquire_lease("2026-08-11", "sweep", 60)
    _harvest(
        store,
        token,
        _sweep_deps(clock, harvest),
        TickBudget(clock.monotonic(), 600, 10, clock.monotonic),
        {"crawl_harvested": 0},
        TickConfig(business_date="2026-08-11", runtime_budget_seconds=600),
    )
    store.release_lease(token)

    prioritized_id = int(jobs[0]["id"])
    flattened = [job_id for batch in observed for job_id in batch]
    assert [len(batch) for batch in observed] == [6, 1]
    assert flattened.count(prioritized_id) == 1
    assert sorted(flattened) == sorted(int(job["id"]) for job in jobs)
    assert store.get_remote_job(prioritized_id)["state"] == "SUBMITTED"


def test_next_sweep_rechecks_jobs_still_running(tmp_path):
    # Case 5: a later scheduler round (a new sweep with a fresh budget) may
    # poll a still-RUNNING remote job again.
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    jobs = _seed_submitted_jobs(store, "2026-08-11", 3)
    polls: list[list[int]] = []

    def harvest(selected, _timeout):
        polls.append([int(job["id"]) for job in selected])
        return {
            int(job["id"]): {
                "state": "RUNNING",
                "remote_identity": str(job["remote_identity"]),
            }
            for job in selected
        }

    deps = _sweep_deps(clock, harvest)
    token = store.acquire_lease("2026-08-11", "sweep", 60)
    for _ in range(2):
        _harvest(
            store,
            token,
            deps,
            TickBudget(clock.monotonic(), 600, 10, clock.monotonic),
            {},
            TickConfig(business_date="2026-08-11", runtime_budget_seconds=600),
        )
    store.release_lease(token)

    expected_ids = [int(job["id"]) for job in jobs]
    assert polls == [expected_ids, expected_ids]
    token = store.acquire_lease("2026-08-11", "verify", 60)
    assert all(job["harvest_attempts"] == 2 for job in store.list_remote_jobs(token))
    store.release_lease(token)


def test_harvest_sweep_stops_at_budget_checkpoint_without_breaking_state(tmp_path):
    # Case 6: exhausting the runtime budget mid-sweep exits at a safe
    # checkpoint: swept jobs keep their recorded attempts, unswept jobs stay
    # untouched, and the tick reports BUDGET_EXHAUSTED without discovery.
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    jobs = _seed_submitted_jobs(store, "2026-08-11", 34)
    events: list[str] = []
    deps = _deps(events, [], clock)

    def harvest(_jobs, timeout):
        events.append("harvest")
        clock.advance(30)
        return {}

    deps.harvest = harvest
    summary = run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=100,
            heavy_work_reserve_seconds=10,
            lease_ttl_seconds=90,
        ),
        deps,
    )

    assert summary.outcome == "BUDGET_EXHAUSTED"
    assert list(summary.phase_trace) == ["recover", "harvest", "checkpoint"]
    assert events.count("harvest") == 3
    token = store.acquire_lease("2026-08-11", "verify", 60)
    listed = store.list_remote_jobs(token)
    store.release_lease(token)
    assert sum(int(job["harvest_attempts"]) for job in listed) == 18
    assert all(str(job["state"]) == "SUBMITTED" for job in listed)


def test_tick_uses_same_harvest_after_submit_when_budget_remains(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    deps = _deps([], [_row()], clock)
    observed: list[list[int]] = []

    def harvest(jobs, _timeout):
        observed.append([int(job["id"]) for job in jobs])
        return {
            int(job["id"]): {
                "state": "RUNNING",
                "remote_identity": str(job["remote_identity"]),
            }
            for job in jobs
        }

    deps.harvest = harvest

    run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        _config(clock),
        deps,
    )

    assert observed[0] == []
    assert len(observed[1]) == 1


def test_existing_remote_drain_crosses_internal_batches_in_one_command(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    for index in range(20):
        job = store.prepare_remote_intent(
            token, f"job-{index}", f"job-{index}", f"payload-{index}"
        )
        store.claim_remote_submission(token, job["id"])
        store.reconcile_remote_job(
            token,
            job["id"],
            remote_identity=f"remote-{index}",
            state="SUBMITTED",
        )
    store.release_lease(token)
    observed: list[int] = []
    deps = _deps([], [], clock)

    def harvest(jobs, _timeout):
        observed.append(len(jobs))
        return {
            job["id"]: {
                "state": "COMPLETED",
                "remote_identity": job["remote_identity"],
                "crawl_results": {},
            }
            for job in jobs
        }

    deps.harvest = harvest
    config = TickConfig(
        business_date="2026-08-11",
        runtime_budget_seconds=60,
        heavy_work_reserve_seconds=5,
        lease_ttl_seconds=60,
        harvest_batch_size=6,
    )

    summary = run_existing_remote_drain(
        store,
        ProcessLock(tmp_path / "drain.lock"),
        config,
        deps,
        max_iterations=10,
    )

    token = store.acquire_lease("2026-08-11", "verify", 60)
    states = [job["state"] for job in store.list_remote_jobs(token)]
    store.release_lease(token)
    assert observed == [6, 6, 6, 2]
    assert states == ["COMPLETED"] * 20
    assert summary.outcome == "DRAINED"


def test_existing_remote_drain_is_existing_only_and_skips_completed_jobs(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    completed = store.prepare_remote_intent(token, "done", "done", "done-payload")
    store.claim_remote_submission(token, completed["id"])
    completed = store.reconcile_remote_job(
        token, completed["id"], remote_identity="remote-done", state="COMPLETED"
    )
    pending = store.prepare_remote_intent(token, "pending", "pending", "pending-payload")
    store.claim_remote_submission(token, pending["id"])
    pending = store.reconcile_remote_job(
        token, pending["id"], remote_identity="remote-pending", state="SUBMITTED"
    )
    store.release_lease(token)
    forbidden = lambda *_args, **_kwargs: pytest.fail(
        "drain-only must not discover, stage, plan, or submit"
    )
    deps = TickDependencies(
        discover=forbidden,
        detail=forbidden,
        model=forbidden,
        model_batch=forbidden,
        plan_crawl=forbidden,
        finalize_crawl_plan=forbidden,
        submit_crawl=forbidden,
        monotonic=clock.monotonic,
        wall_clock=lambda: datetime(2026, 8, 11, 20, tzinfo=UTC),
    )
    observed: list[list[int]] = []

    def harvest(jobs, _timeout):
        observed.append([int(job["id"]) for job in jobs])
        return {
            pending["id"]: {
                "state": "COMPLETED",
                "remote_identity": pending["remote_identity"],
                "crawl_results": {},
            }
        }

    deps.harvest = harvest
    summary = run_existing_remote_drain(
        store,
        ProcessLock(tmp_path / "drain.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=60,
            heavy_work_reserve_seconds=5,
            lease_ttl_seconds=60,
        ),
        deps,
    )

    token = store.acquire_lease("2026-08-11", "verify", 60)
    jobs = store.list_remote_jobs(token)
    snapshots = store.list_snapshots("2026-08-11")
    stages = store.list_stage_work(token)
    store.release_lease(token)
    assert summary.outcome == "DRAINED"
    assert observed == [[pending["id"]]]
    assert len(jobs) == 2
    assert jobs[0]["state"] == "COMPLETED"
    assert jobs[0]["harvest_attempts"] == 0
    assert snapshots == []
    assert stages == []


def test_existing_remote_drain_after_formal_final_does_not_touch_completed_job(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    snapshot = store.ingest_snapshot(
        token,
        [_row()],
        "2026-08-11T10:00:00+00:00",
        source_fingerprints=["source"],
        customer_fingerprints=["customer"],
        headers=["vCommodityId"],
    )["snapshot"]
    job = store.prepare_remote_intent(token, "done", "done", "payload")
    store.claim_remote_submission(token, job["id"])
    job = store.reconcile_remote_job(
        token, job["id"], remote_identity="remote-done", state="COMPLETED"
    )
    store.freeze_snapshot(token, int(snapshot["id"]))
    store.prepare_finalization(token, str(tmp_path / "final.xlsx"))
    store.mark_finalized(token, str(tmp_path / "final.xlsx"), "manifest")
    store.release_lease(token)
    deps = _deps([], [], Clock())
    deps.harvest = lambda *_args: pytest.fail(
        "an already-finalized run must not re-harvest completed work"
    )

    summary = run_existing_remote_drain(
        store,
        ProcessLock(tmp_path / "drain.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=60,
            heavy_work_reserve_seconds=5,
            lease_ttl_seconds=60,
        ),
        deps,
    )

    assert summary.outcome == "ALREADY_FINALIZED"
    assert store.get_remote_job(job["id"])["harvest_attempts"] == 0


def test_manual_drain_does_not_freeze_an_unverified_post_cutoff_snapshot(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    store.ensure_run("2026-08-11", cutoff_at="2026-08-11T10:00:00+00:00")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    snapshot = store.ingest_snapshot(
        token,
        [_row()],
        "2026-08-11T10:01:00+00:00",
        headers=["vCommodityId"],
        evidence={"saas_count": 1, "customer_count": 2},
    )["snapshot"]
    store.release_lease(token)
    deps = _deps([], [], Clock())
    deps.harvest = lambda *_args: pytest.fail(
        "drain without a source-count-verified freeze must not harvest"
    )

    summary = run_existing_remote_drain(
        store,
        ProcessLock(tmp_path / "drain.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=60,
            heavy_work_reserve_seconds=5,
            cutoff_at=datetime(2026, 8, 11, 10, 0, tzinfo=UTC),
        ),
        deps,
    )

    assert summary.outcome == "DRAIN_SNAPSHOT_NOT_READY"
    assert summary.frozen_snapshot_id is None
    assert store.get_run("2026-08-11")["frozen_snapshot_id"] is None
    assert store.get_snapshot(int(snapshot["id"])) is not None


def test_tick_after_formal_final_is_noop_before_any_external_operation(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    snapshot = store.ingest_snapshot(
        token,
        [_row()],
        "2026-08-11T10:00:00+00:00",
        source_fingerprints=["source"],
        customer_fingerprints=["customer"],
        headers=["vCommodityId"],
        evidence={"saas_count": 1, "customer_count": 1},
    )["snapshot"]
    store.freeze_snapshot(token, int(snapshot["id"]))
    store.prepare_finalization(token, str(tmp_path / "final.xlsx"))
    store.mark_finalized(token, str(tmp_path / "final.xlsx"), "manifest")
    store.release_lease(token)
    forbidden = lambda *_args, **_kwargs: pytest.fail(
        "an already-finalized tick must not call external operations"
    )
    deps = TickDependencies(
        discover=forbidden,
        harvest=forbidden,
        detail=forbidden,
        model=forbidden,
        plan_crawl=forbidden,
        submit_crawl=forbidden,
        monotonic=Clock().monotonic,
        wall_clock=lambda: datetime(2026, 8, 11, 20, tzinfo=UTC),
    )

    summary = run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=60,
            heavy_work_reserve_seconds=5,
        ),
        deps,
        finalize=True,
    )

    assert summary.outcome == "ALREADY_FINALIZED"
    assert summary.phase_trace == ("noop:already_finalized",)


def test_existing_remote_drain_advances_priority_then_harvests_same_command(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    job = store.prepare_remote_intent(token, "priority", "priority", "payload")
    store.claim_remote_submission(token, job["id"])
    job = store.reconcile_remote_job(
        token, job["id"], remote_identity="remote-priority", state="PRIORITY_PENDING"
    )
    store.release_lease(token)
    observed: list[str] = []
    deps = _deps([], [], clock)

    def harvest(jobs, _timeout):
        state = str(jobs[0]["state"])
        observed.append(state)
        if state == "PRIORITY_PENDING":
            return {
                job["id"]: {
                    "state": "SUBMITTED",
                    "remote_identity": job["remote_identity"],
                    "result": {"priority_status": "PRIORITY_UPDATED", "weight_type": -1},
                }
            }
        return {
            job["id"]: {
                "state": "COMPLETED",
                "remote_identity": job["remote_identity"],
                "crawl_results": {},
            }
        }

    deps.harvest = harvest
    summary = run_existing_remote_drain(
        store,
        ProcessLock(tmp_path / "drain.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=60,
            heavy_work_reserve_seconds=5,
            lease_ttl_seconds=60,
        ),
        deps,
    )

    assert observed == ["PRIORITY_PENDING", "SUBMITTED"]
    assert summary.outcome == "DRAINED"
    assert summary.counters["drain_iterations"] == 2


def test_existing_remote_drain_stops_on_no_progress(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    job = store.prepare_remote_intent(token, "pending", "pending", "payload")
    store.claim_remote_submission(token, job["id"])
    store.reconcile_remote_job(token, job["id"], remote_identity="remote", state="SUBMITTED")
    store.release_lease(token)
    calls = 0
    deps = _deps([], [], clock)

    def harvest(_jobs, _timeout):
        nonlocal calls
        calls += 1
        return {}

    deps.harvest = harvest
    summary = run_existing_remote_drain(
        store,
        ProcessLock(tmp_path / "drain.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=60,
            heavy_work_reserve_seconds=5,
            lease_ttl_seconds=60,
        ),
        deps,
    )

    assert summary.outcome == "DRAIN_NO_PROGRESS"
    assert calls == 1
    assert summary.phase_trace[-2] == "stop:no_progress"


def test_existing_remote_drain_stops_at_deadline_after_progress(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    job = store.prepare_remote_intent(token, "pending", "pending", "payload")
    store.claim_remote_submission(token, job["id"])
    job = store.reconcile_remote_job(
        token, job["id"], remote_identity="remote", state="SUBMITTED"
    )
    store.release_lease(token)
    deps = _deps([], [], clock)

    def harvest(_jobs, _timeout):
        clock.advance(5.1)
        return {
            job["id"]: {
                "state": "RUNNING",
                "remote_identity": job["remote_identity"],
            }
        }

    deps.harvest = harvest
    summary = run_existing_remote_drain(
        store,
        ProcessLock(tmp_path / "drain.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=6,
            heavy_work_reserve_seconds=1,
            lease_ttl_seconds=60,
        ),
        deps,
    )

    assert summary.outcome == "BUDGET_EXHAUSTED"
    assert summary.phase_trace[-2] == "stop:deadline"


def test_existing_remote_drain_stops_at_max_iterations_after_progress(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    job = store.prepare_remote_intent(token, "pending", "pending", "payload")
    store.claim_remote_submission(token, job["id"])
    job = store.reconcile_remote_job(
        token, job["id"], remote_identity="remote", state="SUBMITTED"
    )
    store.release_lease(token)
    polls = 0
    deps = _deps([], [], clock)

    def harvest(_jobs, _timeout):
        nonlocal polls
        polls += 1
        return {
            job["id"]: {
                "state": "RUNNING",
                "remote_identity": job["remote_identity"],
                "result": {"poll": polls},
            }
        }

    deps.harvest = harvest
    summary = run_existing_remote_drain(
        store,
        ProcessLock(tmp_path / "drain.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=60,
            heavy_work_reserve_seconds=5,
            lease_ttl_seconds=60,
        ),
        deps,
        max_iterations=2,
    )

    assert summary.outcome == "DRAIN_MAX_ITERATIONS"
    assert polls == 2
    assert summary.phase_trace[-2] == "stop:max_iterations"


def test_existing_remote_drain_fails_closed_on_identity_mismatch(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    job = store.prepare_remote_intent(token, "pending", "pending", "payload")
    store.claim_remote_submission(token, job["id"])
    job = store.reconcile_remote_job(
        token, job["id"], remote_identity="expected", state="SUBMITTED"
    )
    store.release_lease(token)
    deps = _deps([], [], clock)
    deps.harvest = lambda _jobs, _timeout: {
        job["id"]: {
            "state": "COMPLETED",
            "remote_identity": "different",
            "crawl_results": {},
        }
    }

    summary = run_existing_remote_drain(
        store,
        ProcessLock(tmp_path / "drain.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=60,
            heavy_work_reserve_seconds=5,
            lease_ttl_seconds=60,
        ),
        deps,
    )

    assert summary.outcome == "FAILED"
    assert summary.error == "RemoteIntentConflict"
    assert store.get_remote_job(job["id"])["state"] == "SUBMITTED"


def test_existing_remote_drain_keeps_validator_or_provenance_failure_manual(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    job = store.prepare_remote_intent(token, "pending", "pending", "payload")
    store.claim_remote_submission(token, job["id"])
    job = store.reconcile_remote_job(
        token, job["id"], remote_identity="remote", state="SUBMITTED"
    )
    store.release_lease(token)
    deps = _deps([], [], clock)
    deps.harvest = lambda _jobs, _timeout: {
        job["id"]: {"manual_error": "validator or provenance failed"}
    }

    summary = run_existing_remote_drain(
        store,
        ProcessLock(tmp_path / "drain.lock"),
        TickConfig(
            business_date="2026-08-11",
            runtime_budget_seconds=60,
            heavy_work_reserve_seconds=5,
            lease_ttl_seconds=60,
        ),
        deps,
    )

    assert summary.outcome == "MANUAL_REVIEW_REQUIRED"
    assert store.get_remote_job(job["id"])["state"] == "MANUAL_REQUIRED"


def test_harvest_persists_per_job_elapsed_evidence_not_whole_batch_elapsed(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    store.ensure_run("2026-08-11")
    token = store.acquire_lease("2026-08-11", "seed", 60)
    jobs = []
    for index in range(2):
        job = store.prepare_remote_intent(token, f"job-{index}", f"job-{index}", f"payload-{index}")
        store.claim_remote_submission(token, job["id"])
        jobs.append(store.reconcile_remote_job(token, job["id"], remote_identity=f"remote-{index}", state="SUBMITTED"))
    store.release_lease(token)
    deps = _deps([], [], clock)
    def harvest(claimed, timeout):
        clock.advance(1)
        first = {"state": "RUNNING", "remote_identity": jobs[0]["remote_identity"], "elapsed_seconds": 1}
        clock.advance(10)
        second = {"state": "RUNNING", "remote_identity": jobs[1]["remote_identity"], "elapsed_seconds": 10}
        return {jobs[0]["id"]: first, jobs[1]["id"]: second}
    deps.harvest = harvest
    config = TickConfig(business_date="2026-08-11", runtime_budget_seconds=60, heavy_work_reserve_seconds=1, lease_ttl_seconds=60)

    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)

    token = store.acquire_lease("2026-08-11", "verify", 60)
    elapsed = [store.get_remote_job(job["id"])["harvest_elapsed_seconds"] for job in jobs]
    store.release_lease(token)
    assert elapsed == [1, 10]


def test_new_source_with_same_backend_sku_url_reuses_completed_crawl_without_post(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    first_rows = [_row("source-a", "same-sku")]
    second_rows = [*first_rows, _row("source-b", "same-sku")]
    config = _config(clock)
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, _deps(events, first_rows, clock))
    token = store.acquire_lease("2026-08-11", "inspect", 30)
    job = store.list_remote_jobs(token)[0]
    crawl = store.crawl_work_for_job(job["id"])[0]
    store.release_lease(token)
    second = _deps(events, second_rows, clock)
    second.harvest = lambda jobs, timeout: {
        job["id"]: {"state": "COMPLETED", "remote_identity": job["remote_identity"], "crawl_results": {crawl["id"]: {"price": "10"}}}
    }
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, second)

    token = store.acquire_lease("2026-08-11", "verify", 30)
    crawls = store.list_crawl_work(token)
    jobs = store.list_remote_jobs(token)
    store.release_lease(token)
    assert events.count("submit") == 1
    assert len(jobs) == 1
    assert {(crawl["source_id"], crawl["sku_id"]) for crawl in crawls} == {("source-a", "same-sku"), ("source-b", "same-sku")}
    assert all(crawl["state"] == "SUCCESS" for crawl in crawls)


def test_new_source_follows_same_pending_backend_request_without_second_post(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    events: list[str] = []
    first_rows = [_row("source-a", "same-sku")]
    second_rows = [
        *first_rows,
        _row("source-b", "same-sku"),
        _row("source-c", "new-sku", extLink="https://shop.example/new-item"),
    ]
    config = _config(clock)

    run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        config,
        _deps(events, first_rows, clock),
    )
    run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        config,
        _deps(events, second_rows, clock),
    )

    token = store.acquire_lease("2026-08-11", "verify", 30)
    crawls = store.list_crawl_work(token)
    jobs = store.list_remote_jobs(token)
    store.release_lease(token)
    assert events.count("submit") == 2
    assert len(jobs) == 2
    assert len(crawls) == 3
    first_job = min(jobs, key=lambda job: int(job["id"]))
    second_job = max(jobs, key=lambda job: int(job["id"]))
    by_source = {crawl["source_id"]: crawl for crawl in crawls}
    assert by_source["source-a"]["remote_job_id"] == first_job["id"]
    assert by_source["source-b"]["remote_job_id"] == first_job["id"]
    assert by_source["source-c"]["remote_job_id"] == second_job["id"]


def test_submit_pre_send_failure_releases_claim_and_next_tick_retries_safely(tmp_path):
    from app.workflows.incremental_price_audit import SubmitNotAttemptedError

    store = DailyAuditStore(tmp_path / "state.sqlite")
    events: list[str] = []
    clock = Clock()
    config = _config(clock)

    failing = _deps(events, [_row()], clock)

    def pre_send_submit(_job, _crawls, _timeout):
        events.append("submit")
        raise SubmitNotAttemptedError("legacy pre-send: auth refresh failed")

    failing.submit_crawl = pre_send_submit
    first = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, failing)

    assert first.counters["submit_not_attempted"] == 1
    assert first.counters["submit_unknown"] == 0
    inspect = store.acquire_lease("2026-08-11", "inspect", 30)
    jobs = store.list_remote_jobs(inspect)
    assert [job["state"] for job in jobs] == ["SUBMIT_INTENT"]
    assert jobs[0]["post_attempted"] is False
    assert "auth refresh failed" in jobs[0]["error_text"]
    store.release_lease(inspect)

    working = _deps(events, [_row()], clock)
    run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, working)

    assert events.count("submit") == 2
    verify = store.acquire_lease("2026-08-11", "verify", 30)
    final = store.list_remote_jobs(verify)[0]
    assert final["state"] == "SUBMITTED"
    assert final["remote_identity"] == "task-1"
    store.release_lease(verify)


def test_repeated_pre_send_failures_keep_releasing_claim_without_unknown(tmp_path):
    from app.workflows.incremental_price_audit import SubmitNotAttemptedError

    store = DailyAuditStore(tmp_path / "state.sqlite")
    events: list[str] = []
    clock = Clock()
    config = _config(clock)
    deps = _deps(events, [_row()], clock)

    def pre_send_submit(_job, _crawls, _timeout):
        events.append("submit")
        raise SubmitNotAttemptedError("legacy pre-send: tool missing")

    deps.submit_crawl = pre_send_submit
    for _ in range(3):
        run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), config, deps)

    assert events.count("submit") == 3
    inspect = store.acquire_lease("2026-08-11", "inspect", 30)
    jobs = store.list_remote_jobs(inspect)
    assert [job["state"] for job in jobs] == ["SUBMIT_INTENT"]
    assert jobs[0]["post_attempted"] is False
    assert jobs[0]["error_text"]
    store.release_lease(inspect)


# ---- requeue / backlog snapshot scoping (round 2, matrix C) ----------------

INVALID_SUCCESS_PAYLOAD = {
    "customerPrice": None,
    "customerPriceStatus": "NOT_FOUND",
    "customerPriceSource": "",
    "customerPriceWarning": "详情接口未返回有效参考价",
    "customerPriceEvidence": "",
}

FOUND_PAYLOAD = {
    "customerPrice": "10",
    "customerPriceStatus": "FOUND",
    "customerPriceSource": "DETAIL_REFERENCE_INFOS",
    "customerPriceWarning": "",
    "customerPriceEvidence": "参考价1=10",
}


def _quiet_row(v, sku):
    return _row(v, sku, extLink="")


def _ingest(store, token, rows, observed_at, evidence=None):
    kwargs = {
        "source_fingerprints": [f"source-{row['vCommodityId']}" for row in rows],
        "customer_fingerprints": [f"customer-{row['fourCategoryId']}" for row in rows],
        "headers": ["vCommodityId"],
    }
    if evidence is not None:
        kwargs["evidence"] = evidence
    return store.ingest_snapshot(token, rows, observed_at, **kwargs)["snapshot"]


def _quiet_deps(rows, *, forbidden=False):
    """Crawl-free deps: no plan/submit, so outcomes reflect stage work only."""
    def discover(timeout):
        if forbidden:
            raise AssertionError("frozen tick must not rediscover")
        return DiscoveredState(
            observed_at=datetime(2026, 8, 11, 10, 1, tzinfo=UTC),
            headers=["vCommodityId"],
            rows=rows,
            source_fingerprints=[f"source-{row['vCommodityId']}" for row in rows],
            customer_fingerprints=[f"customer-{row['fourCategoryId']}" for row in rows],
            evidence={"saas_count": len(rows), "customer_count": len(rows)},
        )

    if forbidden:
        def guard(*_args, **_kwargs):
            raise AssertionError("this tick must not run external stage work")
        # The opening harvest sweep legitimately invokes the callback once with
        # no eligible jobs; everything else must stay untouched.
        return TickDependencies(
            discover=guard, harvest=lambda jobs, timeout: {}, detail=guard, model=guard,
            monotonic=Clock().monotonic,
            wall_clock=lambda: datetime(2026, 8, 11, 10, 2, tzinfo=UTC),
        )

    def detail(row, timeout):
        return dict(FOUND_PAYLOAD)

    def model(row, timeout):
        return {"match_result": "SAME", "confidence": 1, "model_match": "exact"}

    return TickDependencies(
        discover=discover, harvest=lambda jobs, timeout: {}, detail=detail, model=model,
        monotonic=Clock().monotonic,
        wall_clock=lambda: datetime(2026, 8, 11, 10, 2, tzinfo=UTC),
    )


def _complete_stage(store, token, item_id, stage, fingerprint, payload):
    work = store.ensure_stage_work(token, item_id, stage, fingerprint)
    claimed = store.claim_stage_work(token, work["id"], max_attempts=3, max_elapsed_seconds=300)
    assert claimed is not None
    store.complete_stage_work(token, work["id"], payload)


def _item_id_by_v(store, snapshot_id, v):
    for item in store.ordered_snapshot_items(snapshot_id):
        if item["row"]["vCommodityId"] == v:
            return int(item["id"])
    raise AssertionError(f"no snapshot item for {v}")


def _write_invalid_success(store, item_id):
    """Checkpoint the exact legacy invalid-SUCCESS the old worker produced."""
    token = store.acquire_lease("2026-08-11", "seed-invalid-success", 60)
    work = store.ensure_stage_work(token, item_id, "DETAIL", f"manual-{item_id}")
    claimed = store.claim_stage_work(token, work["id"], max_attempts=3, max_elapsed_seconds=300)
    assert claimed is not None
    store.complete_stage_work(token, work["id"], dict(INVALID_SUCCESS_PAYLOAD))
    store.release_lease(token)
    return int(work["id"])


def _detail_works(store):
    token = store.acquire_lease("2026-08-11", "inspect-detail", 30)
    works = store.list_stage_work(token, stage="DETAIL")
    store.release_lease(token)
    return works


def test_requeue_skips_rows_outside_the_current_snapshot_membership(tmp_path):
    """Snapshot A held row A (invalid-SUCCESS), snapshot B only row B: the tick
    requeues nothing outside its authoritative snapshot, row A's legacy state
    stays untouched, and the backlog is not polluted by the dropped row."""
    store = DailyAuditStore(tmp_path / "state.sqlite")
    store.ensure_run("2026-08-11")
    seed = store.acquire_lease("2026-08-11", "seed-snapshots", 60)
    snapshot_a = _ingest(store, seed, [_quiet_row("doc-a", "sku-a")], "2026-08-11T09:00:00+00:00")
    snapshot_b = _ingest(store, seed, [_quiet_row("doc-b", "sku-b")], "2026-08-11T09:30:00+00:00")
    store.release_lease(seed)
    bad_work_id = _write_invalid_success(store, _item_id_by_v(store, snapshot_a["id"], "doc-a"))

    clock = Clock()
    summary = run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        _config(clock),
        _quiet_deps([_quiet_row("doc-b", "sku-b")]),
    )

    assert summary.counters["detail_business_requeued"] == 0
    works = {int(work["id"]): work for work in _detail_works(store)}
    assert int(bad_work_id) in works
    dropped = works[bad_work_id]
    # Not requeued: still the legacy SUCCESS with its result payload intact.
    assert dropped["state"] == "SUCCESS"
    assert json.loads(dropped["result_json"])["customerPriceStatus"] == "NOT_FOUND"
    assert dropped["error_text"] is None
    # Row B processed normally in the same tick.
    good = next(work for work in works.values() if int(work["id"]) != bad_work_id)
    assert good["state"] == "SUCCESS"
    assert json.loads(good["result_json"])["customerPriceStatus"] == "FOUND"
    # No stale-row pollution: the outcome is not stuck on retry-backlog.
    assert summary.outcome != "PENDING_RETRY"
    assert summary.backlog["stage_pending"] == 0
    assert summary.backlog["customer_price_missing"] == 0
    assert summary.backlog["customer_price_failed"] == 0


def test_requeue_still_recovers_invalid_success_inside_current_snapshot(tmp_path):
    """A single-snapshot run (the common case): invalid-SUCCESS rows inside the
    authoritative snapshot are requeued exactly once and retried to SUCCESS."""
    store = DailyAuditStore(tmp_path / "state.sqlite")
    store.ensure_run("2026-08-11")
    seed = store.acquire_lease("2026-08-11", "seed-snapshot", 60)
    snapshot = _ingest(store, seed, [_quiet_row("doc-bad", "sku-bad")], "2026-08-11T09:00:00+00:00")
    store.release_lease(seed)
    _write_invalid_success(store, _item_id_by_v(store, snapshot["id"], "doc-bad"))

    events: list[str] = []
    clock = Clock()
    summary = run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        _config(clock),
        _quiet_deps([_quiet_row("doc-bad", "sku-bad")]),
    )

    assert summary.counters["detail_business_requeued"] == 1
    works = _detail_works(store)
    assert len(works) == 1
    assert works[0]["state"] == "SUCCESS"
    assert json.loads(works[0]["result_json"])["customerPriceStatus"] == "FOUND"
    assert summary.outcome == "COMPLETED"


def test_requeue_on_frozen_run_uses_frozen_membership(tmp_path):
    """A frozen run requeues strictly within its frozen snapshot: an
    invalid-SUCCESS row outside that membership is never resurrected."""
    store = DailyAuditStore(tmp_path / "state.sqlite")
    store.ensure_run("2026-08-11")
    seed = store.acquire_lease("2026-08-11", "seed-frozen", 60)
    frozen_snapshot = _ingest(store, seed, [_quiet_row("doc-keep", "sku-keep")], "2026-08-11T09:00:00+00:00")
    keep_item = _item_id_by_v(store, frozen_snapshot["id"], "doc-keep")
    # A later observation the run never froze (both ingested pre-freeze).
    other_snapshot = _ingest(store, seed, [_quiet_row("doc-late", "sku-late")], "2026-08-11T09:30:00+00:00")
    store.freeze_snapshot(seed, int(frozen_snapshot["id"]))
    # Frozen membership work is already complete; the tick must find nothing to do.
    _complete_stage(store, seed, keep_item, "DETAIL", f"detail-{keep_item}", dict(FOUND_PAYLOAD))
    _complete_stage(store, seed, keep_item, "MODEL", f"model-{keep_item}", {"match_result": "SAME"})
    store.release_lease(seed)
    late_work_id = _write_invalid_success(store, _item_id_by_v(store, other_snapshot["id"], "doc-late"))

    summary = run_incremental_tick(
        store, ProcessLock(tmp_path / "tick.lock"), _config(Clock()),
        _quiet_deps([], forbidden=True),
    )

    assert summary.counters["detail_business_requeued"] == 0
    works = {int(work["id"]): work for work in _detail_works(store)}
    late = works[late_work_id]
    assert late["state"] == "SUCCESS"
    assert json.loads(late["result_json"])["customerPriceStatus"] == "NOT_FOUND"


# ---- FINALIZED runs are a full durable no-op (round 2, matrix F) -----------


def test_finalized_run_tick_neither_requeues_nor_touches_durable_state(tmp_path):
    """FINALIZED history stays immutable: a plain tick exits before any detail
    work, so even a legacy invalid-SUCCESS row is left exactly as it was, and
    final_output_path never moves.  (Corrects the round-1 claim that rerunning
    the tick would repair historical days.)"""
    store = DailyAuditStore(tmp_path / "state.sqlite")
    store.ensure_run("2026-08-11")
    seed = store.acquire_lease("2026-08-11", "seed-final", 60)
    snapshot = _ingest(
        store, seed, [_quiet_row("doc-1", "sku-1")], "2026-08-11T09:00:00+00:00",
        evidence={"saas_count": 1, "customer_count": 1},
    )
    item_id = _item_id_by_v(store, snapshot["id"], "doc-1")
    store.freeze_snapshot(seed, int(snapshot["id"]))
    _complete_stage(store, seed, item_id, "DETAIL", f"detail-{item_id}", dict(FOUND_PAYLOAD))
    _complete_stage(store, seed, item_id, "MODEL", f"model-{item_id}", {"match_result": "SAME"})
    final_output = tmp_path / "final.xlsx"
    store.prepare_finalization(seed, str(final_output))
    store.mark_finalized(seed, str(final_output), "manifest")
    store.release_lease(seed)

    # Poison left behind from the legacy era: an invalid-SUCCESS DETAIL row.
    poisoned_id = _write_invalid_success(store, item_id)

    def before(work):
        return (work["id"], work["state"], work["attempts"], work["result_json"], work["error_text"], work["elapsed_seconds"], work["retryable"])

    stage_before = sorted(before(work) for work in _detail_works(store))
    run_before = store.get_run("2026-08-11")

    deps = TickDependencies(
        discover=lambda *_a: pytest.fail("already-finalized tick must not discover"),
        harvest=lambda *_a: pytest.fail("already-finalized tick must not harvest"),
        detail=lambda *_a: pytest.fail("already-finalized tick must not run detail"),
        model=lambda *_a: pytest.fail("already-finalized tick must not run model"),
        plan_crawl=lambda *_a: pytest.fail("already-finalized tick must not plan"),
        submit_crawl=lambda *_a: pytest.fail("already-finalized tick must not submit"),
        monotonic=Clock().monotonic, wall_clock=lambda: datetime(2026, 8, 11, 20, tzinfo=UTC),
    )
    summary = run_incremental_tick(
        store, ProcessLock(tmp_path / "tick.lock"),
        TickConfig(business_date="2026-08-11", runtime_budget_seconds=60, heavy_work_reserve_seconds=5),
        deps,
    )

    assert summary.outcome == "ALREADY_FINALIZED"
    assert summary.phase_trace == ("noop:already_finalized",)
    assert summary.counters["detail_business_requeued"] == 0
    stage_after = sorted(before(work) for work in _detail_works(store))
    assert stage_after == stage_before
    poisoned = next(work for work in _detail_works(store) if int(work["id"]) == poisoned_id)
    assert poisoned["state"] == "SUCCESS"
    assert json.loads(poisoned["result_json"])["customerPriceStatus"] == "NOT_FOUND"
    run_after = store.get_run("2026-08-11")
    assert run_after["final_output_path"] == run_before["final_output_path"] == str(final_output.resolve())
    assert run_after["state"] == "FINALIZED"


# ---- Round 3 final-state consistency (tick/backlog/outcome/metric) --------


def _seed_terminal_detail(store, item_id, *, error="permanent detail failure"):
    """Checkpoint one DETAIL row as TERMINAL_FAILED (retry budget exhausted)."""
    token = store.acquire_lease("2026-08-11", "seed-terminal-detail", 60)
    try:
        work = store.ensure_stage_work(token, item_id, "DETAIL", f"manual-{item_id}")
        claimed = store.claim_stage_work(token, work["id"], max_attempts=1, max_elapsed_seconds=300)
        assert claimed is not None
        return store.fail_stage_work(
            token, work["id"], elapsed_seconds=0.1, error=error, retryable=False,
            max_attempts=1, max_elapsed_seconds=300,
        )
    finally:
        store.release_lease(token)


def test_empty_authoritative_snapshot_clears_dropped_detail_blockers(tmp_path):
    """Snapshot A held row A (DETAIL terminal), snapshot B is the current
    authoritative snapshot with zero rows: the DETAIL backlog — technical and
    customer-price alike — must be scoped to nothing (final-state language:
    nothing blocks the final sheet), while a caller without an authoritative
    snapshot (drain/harvest) still sees the whole-run view."""
    store = DailyAuditStore(tmp_path / "state.sqlite")
    store.ensure_run("2026-08-11")
    seed = store.acquire_lease("2026-08-11", "seed-empty-snapshot", 60)
    snapshot_a = _ingest(store, seed, [_quiet_row("doc-a", "sku-a")], "2026-08-11T09:00:00+00:00")
    item_a = _item_id_by_v(store, snapshot_a["id"], "doc-a")
    store.release_lease(seed)
    _seed_terminal_detail(store, item_a)

    clock = Clock()
    summary = run_incremental_tick(
        store,
        ProcessLock(tmp_path / "tick.lock"),
        _config(clock),
        _quiet_deps([]),
    )

    # Snapshot B (0 rows) is the authoritative scope: the dropped row A must
    # not pollute any DETAIL backlog bucket.
    assert summary.backlog["detail_terminal"] == 0
    assert summary.backlog["customer_price_missing"] == 0
    assert summary.backlog["customer_price_failed"] == 0
    assert summary.backlog["customer_price_invalid"] == 0
    assert summary.backlog["customer_absent_skipped"] == 0
    assert summary.backlog["customer_price_success_blockers"] == 0
    assert summary.backlog["stage_pending"] == 0
    assert summary.outcome == "COMPLETED"

    # Whole-run fallback: a caller without an authoritative snapshot
    # (drain/harvest paths call _backlog without items) still counts row A.
    token = store.acquire_lease("2026-08-11", "inspect-backlog", 30)
    try:
        whole_run = _backlog(store, token)
    finally:
        store.release_lease(token)
    assert whole_run["detail_terminal"] == 1
    assert whole_run["customer_price_failed"] == 1


ABSENT_SKIP_PAYLOAD = {
    "customerPrice": None,
    "customerPriceStatus": "LOOKUP_FAILED",
    "customerPriceSource": "",
    "customerPriceWarning": "customer export did not contain this documentId",
    "customerPriceEvidence": "authoritative customer snapshot round complete; documentId absent",
}

CONTRADICTORY_PRICE_PAYLOAD = {
    "customerPrice": "100",
    "customerPriceStatus": "NOT_FOUND",
    "customerPriceSource": "",
    "customerPriceWarning": "详情接口未返回有效参考价",
    "customerPriceEvidence": "",
}


@pytest.mark.parametrize("status", ["CUSTOMER_NOT_FOUND", "LOOKUP_FAILED"])
def test_success_customer_absent_skip_is_nonblocking_tick_completes(tmp_path, status):
    """2026-09-11 business authorization (rewritten from the R2-B manual-review
    contract): the round's authoritative customer query succeeded and the
    documentId is absent — a non-blocking skip.  The row is SUCCESS evidence
    (not terminal, not retrying), the backlog shows customer_absent_skipped,
    customer_price_success_blockers stays 0, and a day with only such rows
    completes instead of demanding manual review.  LOOKUP_FAILED is the legacy
    spelling of the same deterministic join outcome."""
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    deps = _quiet_deps([_quiet_row("doc-1", "sku-1")])
    deps.detail = lambda _row, _timeout: {**ABSENT_SKIP_PAYLOAD, "customerPriceStatus": status}

    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), deps)

    assert summary.backlog["customer_absent_skipped"] == 1
    assert summary.backlog["customer_price_success_blockers"] == 0
    assert summary.backlog["detail_terminal"] == 0
    assert summary.outcome != "MANUAL_REVIEW_REQUIRED"
    assert summary.outcome == "COMPLETED"


def test_success_contradictory_price_detail_blocks_tick_as_manual_review(tmp_path):
    """SUCCESS + positive price + NOT_FOUND status: contradictory evidence is
    INVALID and blocks — the tick must not report COMPLETED."""
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    deps = _quiet_deps([_quiet_row("doc-1", "sku-1")])
    deps.detail = lambda _row, _timeout: dict(CONTRADICTORY_PRICE_PAYLOAD)

    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), deps)

    assert summary.backlog["customer_price_invalid"] == 1
    assert summary.backlog["customer_price_success_blockers"] == 1
    assert summary.outcome == "MANUAL_REVIEW_REQUIRED"
    assert summary.outcome != "COMPLETED"


def test_retryable_detail_failure_stays_pending_retry_not_manual(tmp_path):
    """RETRYING must not be escalated to manual review: it is ordinary retry
    backlog and surfaces as PENDING_RETRY."""
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    deps = _quiet_deps([_quiet_row("doc-1", "sku-1")])
    deps.detail = lambda _row, _timeout: (_ for _ in ()).throw(ValueError("transient detail failure"))

    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), deps)

    assert summary.outcome == "PENDING_RETRY"
    assert summary.backlog["stage_pending"] == 1
    assert summary.backlog["customer_price_missing"] == 0
    assert summary.backlog["customer_price_failed"] == 0
    assert summary.backlog["customer_price_invalid"] == 0
    assert summary.backlog["customer_absent_skipped"] == 0


def test_found_price_completes_without_customer_price_blockers(tmp_path):
    store = DailyAuditStore(tmp_path / "state.sqlite")
    clock = Clock()
    deps = _quiet_deps([_quiet_row("doc-1", "sku-1")])

    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), deps)

    assert summary.outcome == "COMPLETED"
    assert summary.backlog["customer_price_missing"] == 0
    assert summary.backlog["customer_price_failed"] == 0
    assert summary.backlog["customer_price_invalid"] == 0
    assert summary.backlog["customer_absent_skipped"] == 0
    assert summary.backlog["customer_price_success_blockers"] == 0


def test_final_pending_metric_matches_status_computed_pending(tmp_path):
    """The persisted final_pending metric and the status projection must speak
    the same language: one SUCCESS-stage blocking customer-price row (here the
    contradictory INVALID evidence) is exactly one unit of final pending in
    both views, while a customer-absent skip row counts as zero in both."""
    from app.services.incremental_audit_runtime import runtime_paths
    from app.services.incremental_audit_status import collect_runtime_status

    root = tmp_path / "runtime"
    store = DailyAuditStore(runtime_paths("2026-08-11", root).state_path)
    clock = Clock()
    deps = _quiet_deps([_quiet_row("doc-1", "sku-1"), _quiet_row("doc-2", "sku-2")])
    deps.detail = lambda row, _timeout: (
        dict(CONTRADICTORY_PRICE_PAYLOAD)
        if row.get("documentId") == "doc-1" else dict(ABSENT_SKIP_PAYLOAD)
    )

    summary = run_incremental_tick(store, ProcessLock(tmp_path / "tick.lock"), _config(clock), deps)

    status = collect_runtime_status(business_date="2026-08-11", runtime_root=root)
    assert summary.outcome == "MANUAL_REVIEW_REQUIRED"
    assert summary.backlog["customer_price_invalid"] == 1
    assert summary.backlog["customer_absent_skipped"] == 1
    assert status["metrics"]["final_pending"] == 1
    assert status["final"]["computed_pending"] == 1
    assert status["final"]["computed_pending"] == status["metrics"]["final_pending"]
