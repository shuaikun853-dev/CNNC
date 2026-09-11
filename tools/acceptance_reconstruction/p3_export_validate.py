# -*- coding: utf-8 -*-
"""p3_export_validate.py — F3 导出+验证层主脚本（验收回溯重建）。

规格: reports/SPEC_ACCEPTANCE_RECONSTRUCTION.md §4/§6/§7（冻结版契约）。

用法:
    python.exe -X utf8 p3_export_validate.py --config <config.local.json> [--part export|validate|recalc|all]

输入（只读, F1/F2 产物）: temp\\step{1..5}.parquet, temp\\final_main.parquet,
    lineage\\lineage.parquet, inventory\\p1_stats.json, reports\\stage_stats.json,
    inventory\\source_inventory_round1.csv, FINAL 三件套（sha 复核）。
输出:
    output_customer\\01..05 客户 xlsx（SPEC §4 中文列, 冻结/筛选/文本防科学计数/'0.00'）
    inventory\\source_inventory.xlsx（round1 清单 + 17 文件分类清单）
    lineage\\lineage.xlsx（全量 2.49M 行, 自动拆 sheet）
    reports\\stage_statistics.xlsx（漏斗/剔除/冲突/评论/断言）
    reports\\price_recalc_check.xlsx（价格指数复算, 仅记录不改交付物）
    reports\\validation_report.xlsx + reports\\reconstruction_report.md（结论占位给主 Agent）

硬门禁任一 FAIL → exit 2; 磁盘/单文件超限 → DISK_SAFETY_STOP exit 3。
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import shutil
import statistics
import sys
import time
from datetime import datetime
from decimal import Decimal, InvalidOperation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from p1_common import (  # noqa: E402
    connect_duckdb,
    dir_size,
    file_meta,
    fmt_bytes,
    load_config,
    sha256_file,
)
from p3_format import (  # noqa: E402
    MISSING_COMMENT_STATES,
    SHEET_MAX_DATA_ROWS,
    build_sensitive_pattern,
    customer_fields,
    customer_headers,
    customer_kinds,
    dec_of,
    find_sensitive_hit,
    index_val_equal,
    recalc_avg_price,
    recalc_power_index,
    write_xlsx,
)

# ---------------------------------------------------------------------------
# 冻结常量
# ---------------------------------------------------------------------------

EXPECTED_STEP5_CODES = 120047
WORKDIR_NEW_LIMIT_BYTES = 5 * 1024 ** 3      # workdir 新增写盘上限 5GB
XLSX_MAX_BYTES = 400 * 1024 ** 3 // 1024     # 单 xlsx 上限 400MB
FREE_SPACE_MIN = 10 * 1024 ** 3              # 导出前 E 盘剩余空间下限 10GB

STEP_PARQUET = {"01": "step1.parquet", "02": "step2.parquet", "03": "step3.parquet",
                "04": "step4.parquet", "05": "step5.parquet"}
CUSTOMER_FILES = {
    "01": ("01_同款校验结果数据.xlsx", "同款校验结果数据"),
    "02": ("02_去极值后结果数据.xlsx", "去极值后结果数据"),
    "03": ("03_一物多码去重后结果数据.xlsx", "一物多码去重后结果数据"),
    "04": ("04_人工审核通过结果数据.xlsx", "人工审核通过结果数据"),
    "05": ("05_价格指数计算结果数据.xlsx", "价格指数计算结果数据"),
}
LINEAGE_COLS = ["stage", "standard_code", "shop_code", "sku_id", "url", "source_file",
                "source_sheet", "source_row", "record_origin", "verdict_source",
                "snapshot_source", "next_stage_action", "internal_reason",
                "comment_source", "stage_drop_reason", "conflict_type", "sources"]

#: 17 文件分类清单（p1_stats.inputs 16 外部源 + round1 清单 CSV 本体 = 17 行）
CATEGORY_MAP = {
    "final_authority": ("FINAL_AUTHORITY", ""),
    "final_authority_equiv": ("FINAL_AUTHORITY", "等价重存副本"),
    "final_authority_orig": ("FINAL_AUTHORITY", "0825 原始字节"),
    "p25_pool": ("SAME_PRODUCT_CANDIDATE", "平台匹配候选池"),
    "s3_flags": ("SAME_PRODUCT_CANDIDATE|OUTLIER_RELATED|DEDUP_RELATED",
                 "补采批, 兼带 去极值/一物多码/京东 标志"),
    "s3_raw": ("SAME_PRODUCT_CANDIDATE", "补采批未打标志版(对照)"),
    "drop80": ("OUTLIER_RELATED", "偏离超80 剔除清单"),
    "drop_new": ("OUTLIER_RELATED", "0821 并入剔除(真实原因)"),
    "dify_recheck": ("SAME_PRODUCT_CANDIDATE|MANUAL_AUDIT", "0908 验收表同款再校验"),
    "recall_csv": ("COMMENT_SOURCE|PRICE_INDEX_HISTORY", "0909 召回明细+评论数"),
    "comment_index_parquet": ("COMMENT_SOURCE", "评论索引"),
    "delivery_coverage_parquet": ("COMMENT_SOURCE", "交付行评论覆盖"),
    "power_match_parquet": ("COMMENT_SOURCE", "幂次子集评论匹配"),
    "delivery_0825_parquet": ("PRICE_INDEX_HISTORY", "FINAL 已归一化历史交付(对照)"),
    "zs_result": ("COMMENT_SOURCE", "0910 苏宁评论(最新)"),
    "algo_src": ("ALGO_SOURCE", "算法包(31测试通过); 7给定标签外补充标签, 无可归入项"),
}
CATEGORY_INVENTORY_ROW = "SOURCE_INVENTORY"

#: 复算口径与偏差说明（0909 重算实证锚点, 报告引用）
RECALC_NOTE = (
    "口径来源: 0909 重算实证（幂次 CEIL 2位 / 平均价 HALF_UP 2位, 写死）。"
    "已知偏差原因: 一级数据源分级取数不可复现 / 评论时间点不同（评论数快照差异）。"
    "历史锚点: 0909 mingxi_sku5 acceptance median_abs_diff = 1.005。"
    "mismatch 仅记录, 不修改任何交付物。"
)
COMMENT_STATES_ORDER = ["POSITIVE", "ZERO", "EMPTY_SRC", "NOT_FOUND", "NO_KEY"]


def _log(msg: str) -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)


def _q(path: str) -> str:
    return path.replace("'", "''")


def _wd(cfg: dict, *parts: str) -> str:
    return os.path.join(cfg["workdir"], *parts)


def _require(path: str, what: str) -> None:
    if not os.path.isfile(path):
        print(f"MISSING_INPUT: {what} 不存在: {path}（先运行前置阶段）", flush=True)
        sys.exit(2)


def _fetch_rows(con, sql: str, batch: int = 50_000):
    """流式取数生成器（元组）。"""
    cur = con.execute(sql)
    while True:
        rows = cur.fetchmany(batch)
        if not rows:
            break
        yield from rows


def _check_free_space(cfg: dict, min_free: int = FREE_SPACE_MIN) -> int:
    drive = os.path.splitdrive(cfg["workdir"])[0] + os.sep
    free = shutil.disk_usage(drive).free
    if free < min_free:
        print(f"DISK_SAFETY_STOP: {drive} 剩余 {free} 字节 < 下限 {min_free}", flush=True)
        sys.exit(3)
    return free


def _check_file_size(path: str) -> int:
    size = os.path.getsize(path)
    if size > XLSX_MAX_BYTES:
        print(f"DISK_SAFETY_STOP: {path} = {size} 字节 > 单文件上限 {XLSX_MAX_BYTES}", flush=True)
        sys.exit(3)
    return size


def _setup_scratch_tmp(cfg: dict) -> str:
    """openpyxl write_only 把 sheet XML 流式写入系统临时目录（C 盘可能不足,
    2026-09-11 实测 C 盘满导致 lxml IO_WRITE）→ 重定向到 workdir 内 scratch, 结束后清理。"""
    import tempfile

    scratch = _wd(cfg, ".tmp_xlsx")
    os.makedirs(scratch, exist_ok=True)
    os.environ["TMP"] = scratch
    os.environ["TEMP"] = scratch
    tempfile.tempdir = scratch
    return scratch


def _cleanup_scratch(scratch: str) -> None:
    if os.path.isdir(scratch):
        shutil.rmtree(scratch, ignore_errors=True)


# ---------------------------------------------------------------------------
# export
# ---------------------------------------------------------------------------

def part_export(cfg: dict, con) -> list[dict]:
    """5 个客户 xlsx + 3 个内部产物。返回 manifest（含 sheets 行数）。"""
    _check_free_space(cfg)
    manifest: list[dict] = []

    # ---- 5 个客户文件 ----
    outdir = _wd(cfg, "output_customer")
    os.makedirs(outdir, exist_ok=True)
    for step in ("01", "02", "03", "04", "05"):
        fname, sheet_base = CUSTOMER_FILES[step]
        path = os.path.join(outdir, fname)
        parquet = os.path.join(cfg["temp_dir"], STEP_PARQUET[step])
        _require(parquet, f"客户{step}源数据")
        fields = customer_fields(step)
        sql = f"SELECT {', '.join(fields)} FROM '{_q(parquet)}'"
        counter = {"n": 0}

        def _prog(name: str, n: int, _f=fname, _c=counter) -> None:
            _c["n"] = n
            if n % 500_000 == 0:
                _log(f"  {_f} {name}: 已写 {n} 行")

        _log(f"导出客户{step}: {fname} ← {os.path.basename(parquet)}")
        sheets = [{"name": sheet_base, "headers": customer_headers(step),
                   "kinds": customer_kinds(step), "rows": _fetch_rows(con, sql)}]
        res = write_xlsx(path, sheets, on_progress=_prog)
        size = _check_file_size(path)
        manifest.append({"file": path, "kind": "customer", "step": step,
                         "sheets": res, "size": size})
        _log(f"  完成: {fname} 行数={sum(s['rows'] for s in res)} 大小={fmt_bytes(size)}")
        _check_free_space(cfg)

    # ---- inventory/source_inventory.xlsx ----
    p1_stats_path = _wd(cfg, "inventory", "p1_stats.json")
    _require(p1_stats_path, "p1_stats.json")
    with open(p1_stats_path, "r", encoding="utf-8") as f:
        p1_stats = json.load(f)
    inv_path = _wd(cfg, "inventory", "source_inventory.xlsx")
    _log("导出内部产物: inventory/source_inventory.xlsx")
    res = write_xlsx(inv_path, _inventory_sheets(cfg, p1_stats))
    size = _check_file_size(inv_path)
    manifest.append({"file": inv_path, "kind": "internal", "sheets": res, "size": size})
    _log(f"  完成: sheets={res} 大小={fmt_bytes(size)}")

    # ---- lineage/lineage.xlsx ----
    lineage_pq = _wd(cfg, "lineage", "lineage.parquet")
    _require(lineage_pq, "lineage.parquet")
    lin_path = _wd(cfg, "lineage", "lineage.xlsx")
    _log("导出内部产物: lineage/lineage.xlsx（全量, 自动拆 sheet）")
    sql = f"SELECT {', '.join(LINEAGE_COLS)} FROM '{_q(lineage_pq)}'"
    counter = {"n": 0}

    def _prog_lin(name: str, n: int, _c=counter) -> None:
        _c["n"] = n
        if n % 500_000 == 0:
            _log(f"  lineage.xlsx {name}: 已写 {n} 行")

    res = write_xlsx(lin_path, [{"name": "lineage", "headers": LINEAGE_COLS,
                                 "kinds": ["text"] * len(LINEAGE_COLS),
                                 "rows": _fetch_rows(con, sql, 100_000)}],
                     on_progress=_prog_lin)
    size = _check_file_size(lin_path)
    manifest.append({"file": lin_path, "kind": "internal", "sheets": res, "size": size})
    _log(f"  完成: sheets={res} 大小={fmt_bytes(size)}")
    _check_free_space(cfg)

    # ---- reports/stage_statistics.xlsx ----
    stats_path = _wd(cfg, "reports", "stage_stats.json")
    _require(stats_path, "stage_stats.json")
    with open(stats_path, "r", encoding="utf-8") as f:
        stage_stats = json.load(f)
    st_path = _wd(cfg, "reports", "stage_statistics.xlsx")
    _log("导出内部产物: reports/stage_statistics.xlsx")
    res = write_xlsx(st_path, _stage_stats_sheets(stage_stats))
    size = _check_file_size(st_path)
    manifest.append({"file": st_path, "kind": "internal", "sheets": res, "size": size})
    _log(f"  完成: sheets={[s['name'] for s in res]} 大小={fmt_bytes(size)}")

    return manifest


def _inventory_sheets(cfg: dict, p1_stats: dict) -> list[dict]:
    """sheet1=round1 清单 CSV; sheet 分类清单=17 文件（16 外部源+CSV 本体）。"""
    csv_path = cfg.get("inventory_round1_csv") or _wd(cfg, "inventory", "source_inventory_round1.csv")
    _require(csv_path, "inventory_round1_csv")

    def _csv_rows():
        with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            next(reader, None)  # 跳过 CSV 自身表头行（sheet 已有表头）
            for row in reader:
                yield tuple(row)

    cls_rows = []
    for key, meta in p1_stats.get("inputs", {}).items():
        cat, note = CATEGORY_MAP.get(key, ("UNCLASSIFIED", ""))
        cls_rows.append((key, cat, meta.get("path", ""), meta.get("sha256", ""),
                         str(meta.get("size", "")), meta.get("mtime", ""), note))
    m = file_meta(csv_path)
    cls_rows.append(("inventory_round1_csv", CATEGORY_INVENTORY_ROW, csv_path,
                     sha256_file(csv_path), str(m["size"]), m["mtime"],
                     "F1 产物清单 CSV 本体; p1_stats.inputs 实为 16 外部源, 加本行=17"))
    return [
        {"name": "round1_inventory", "headers": ["path", "size_bytes", "mtime"],
         "kinds": ["text", "text", "text"], "rows": _csv_rows()},
        {"name": "分类清单",
         "headers": ["文件键", "分类标签", "路径", "sha256", "size_bytes", "mtime", "备注"],
         "kinds": ["text", "text", "text", "text", "text", "text", "text"],
         "rows": cls_rows},
    ]


def _stage_stats_sheets(st: dict) -> list[dict]:
    """stage_stats.json → 多 sheet 表格（漏斗/剔除明细/冲突/评论四态/断言）。"""
    funnel_rows = []
    detail_rows = []
    for step in ("step1", "step2", "step3", "step4", "step5"):
        d = st.get("funnel", {}).get(step, {})
        funnel_rows.append((step, d.get("rows", ""), d.get("codes", ""), d.get("malls", ""),
                            d.get("unique_skus", ""), d.get("rows_in_prev", ""),
                            d.get("filtered", ""), d.get("rows_out", ""),
                            d.get("yes", ""), d.get("no", "")))
        for reason, n in (d.get("filtered_breakdown") or {}).items():
            detail_rows.append((step, reason, n))

    conflict_rows = []
    for k, v in (st.get("conflicts") or {}).items():
        if isinstance(v, dict):
            for k2, v2 in v.items():
                conflict_rows.append((f"{k}.{k2}", v2, ""))
        else:
            conflict_rows.append((k, v, ""))

    cm = st.get("comments", {})
    s4, s5 = cm.get("step4", {}), cm.get("step5", {})
    comment_rows = [(state, s4.get(state, ""), s5.get(state), "")
                    for state in COMMENT_STATES_ORDER]
    comment_rows.append(("comment_filled", s4.get("comment_filled", ""),
                         s5.get("comment_filled", ""), ""))
    comment_rows.append(("coverage(填充率)", None, None,
                         f"step4={s4.get('coverage')}; step5={s5.get('coverage')}"))

    assert_rows = []
    for a in st.get("asserts", []):
        act = a.get("actual")
        exp = a.get("expected")
        act_s = json.dumps(act, ensure_ascii=False) if isinstance(act, list) else str(act)
        exp_s = json.dumps(exp, ensure_ascii=False) if isinstance(exp, list) else str(exp)
        assert_rows.append((a.get("stage", ""), a.get("name", ""), act_s, exp_s,
                            "PASS" if a.get("pass") else "FAIL"))

    return [
        {"name": "漏斗",
         "headers": ["阶段", "总行数", "标品数", "商城数", "唯一SKU数", "上行数", "剔除数",
                     "输出数", "是", "否"],
         "kinds": ["text"] + ["int"] * 9, "rows": funnel_rows},
        {"name": "剔除明细", "headers": ["阶段", "剔除原因", "数量"],
         "kinds": ["text", "text", "int"], "rows": detail_rows},
        {"name": "冲突", "headers": ["冲突/事件", "数量", "备注"],
         "kinds": ["text", "int", "text"], "rows": conflict_rows},
        {"name": "评论四态", "headers": ["状态", "step4", "step5", "备注"],
         "kinds": ["text", "int", "int", "text"], "rows": comment_rows},
        {"name": "断言", "headers": ["阶段", "断言", "实际", "期望", "结果"],
         "kinds": ["int", "text", "text", "text", "text"], "rows": assert_rows},
    ]


# ---------------------------------------------------------------------------
# recalc（价格指数复算检查, 只记录不改交付物）
# ---------------------------------------------------------------------------

def _classify_method(methods: set[str]) -> str:
    joined = " | ".join(sorted(m for m in methods if m))
    if any("幂次" in m for m in methods):
        return "幂次"
    if any("平均价" in m for m in methods):
        return "平均价"
    if any("不计算" in m for m in methods):
        return "不计算"
    return "未知"


def part_recalc(cfg: dict, con) -> dict:
    """按 code 分组复算 step5 价格指数（幂次 p=1 CEIL2 / 平均价 HALF_UP2 / 不计算验空）。"""
    s5 = os.path.join(cfg["temp_dir"], "step5.parquet")
    _require(s5, "step5.parquet")
    _log("价格指数复算: 读取 step5（ORDER BY code 流式分组）")
    sql = (f"SELECT code, mall, sku, price, comment_value, comment_state, index_val, method "
           f"FROM '{_q(s5)}' ORDER BY code")

    classes = {"幂次": _new_cls(), "平均价": _new_cls(), "不计算": _new_cls(), "未知": _new_cls()}
    mixed_method_codes = 0
    cur_group: list | None = None
    prev_code = None

    def _flush(group: list) -> None:
        nonlocal mixed_method_codes
        if not group:
            return
        code = group[0][0]
        methods = {r[7] for r in group}
        if len({m for m in methods if m}) > 1:
            mixed_method_codes += 1
        cls_name = _classify_method(methods)
        c = classes[cls_name]
        c["codes"] += 1
        c["rows"] += len(group)
        rows3 = [(r[3], r[4], r[5]) for r in group]
        if cls_name == "不计算":
            for r in group:
                iv = (r[6] or "").strip() if r[6] is not None else ""
                if iv != "":
                    c["mismatch"] += 1
                    if len(c["samples"]) < 20:
                        c["samples"].append((code, r[1], r[2], r[3], r[4], r[5],
                                             "应为空", r[6], None))
            return
        computed = (recalc_power_index(rows3) if cls_name == "幂次"
                    else recalc_avg_price(rows3))
        skipped = len(rows3) - sum(1 for r in rows3 if r[0] is not None)
        c["skipped_price"] += skipped
        if computed is None:
            c["no_pool_codes"] += 1
            c["mismatch"] += len(group)
            return
        for r in group:
            iv = (r[6] or "").strip() if r[6] is not None else ""
            ok = index_match(computed, iv)
            if ok:
                c["match"] += 1
            else:
                c["mismatch"] += 1
                diff = _abs_diff(computed, iv)
                if diff is not None:
                    c["diffs"].append(float(diff))
                if len(c["samples"]) < 20:
                    c["samples"].append((code, r[1], r[2], r[3], r[4], r[5],
                                         str(computed), r[6] or "",
                                         str(diff) if diff is not None else "不可比"))

    for row in _fetch_rows(con, sql, 100_000):
        code = row[0]
        if cur_group is None or code != prev_code:
            _flush(cur_group or [])
            cur_group = [row]
            prev_code = code
        else:
            cur_group.append(row)
    _flush(cur_group or [])

    for c in classes.values():
        c["max_diff"] = max(c["diffs"]) if c["diffs"] else None
        c["median_diff"] = statistics.median(c["diffs"]) if c["diffs"] else None

    summary = {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "classes": {k: {kk: vv for kk, vv in v.items() if kk != "diffs"}
                    for k, v in classes.items()},
        "mixed_method_codes": mixed_method_codes,
        "note": RECALC_NOTE,
    }
    _write_recalc_xlsx(cfg, summary)
    for k, v in summary["classes"].items():
        _log(f"  复算[{k}]: rows={v['rows']} codes={v['codes']} match={v['match']} "
             f"mismatch={v['mismatch']} max_diff={v['max_diff']} median_diff={v['median_diff']}")
    _log(f"  复算: method 混合 code 数={mixed_method_codes}")
    return summary


def _new_cls() -> dict:
    return {"rows": 0, "codes": 0, "match": 0, "mismatch": 0, "skipped_price": 0,
            "no_pool_codes": 0, "diffs": [], "samples": [], "max_diff": None,
            "median_diff": None}


def index_match(computed: Decimal, index_val: str) -> bool:
    return index_val_equal(computed, index_val)


def _abs_diff(computed: Decimal, index_val: str) -> Decimal | None:
    if not index_val:
        return None
    try:
        return abs(dec_of(index_val) - computed)
    except (InvalidOperation, ValueError):
        return None


def _write_recalc_xlsx(cfg: dict, summary: dict) -> None:
    path = _wd(cfg, "reports", "price_recalc_check.xlsx")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    sum_rows = []
    for cls_name, c in summary["classes"].items():
        sum_rows.append((cls_name, c["rows"], c["codes"], c["match"], c["mismatch"],
                         c["skipped_price"], c["no_pool_codes"],
                         str(c["max_diff"]) if c["max_diff"] is not None else "",
                         str(c["median_diff"]) if c["median_diff"] is not None else ""))
    sum_rows.append(("method混合code数", summary["mixed_method_codes"], "", "", "", "", "",
                     "", ""))
    note_rows = [("口径与偏差说明", summary["note"])]

    def _cls_sheet(name: str) -> dict:
        c = summary["classes"][name]
        rows = [(s[0], s[1], s[2], s[3], s[4], s[5], s[6], s[7], str(s[8] or ""))
                for s in c.get("samples", [])]
        return {"name": name,
                "headers": ["code", "mall", "sku", "price", "comment_value", "comment_state",
                            "复算值", "原值", "abs_diff"],
                "kinds": ["text_at", "text", "text_at", "price", "int", "text",
                          "text", "text", "text"],
                "rows": rows}

    write_xlsx(path, [
        {"name": "汇总",
         "headers": ["类别", "行数", "code数", "match", "mismatch", "价格无效跳过行",
                     "无池code", "max_abs_diff", "median_abs_diff"],
         "kinds": ["text", "int", "int", "int", "int", "int", "int", "text", "text"],
         "rows": sum_rows},
        {"name": "口径说明", "headers": ["项目", "说明"],
         "kinds": ["text", "text"], "rows": note_rows},
        _cls_sheet("幂次"),
        _cls_sheet("平均价"),
        _cls_sheet("不计算"),
    ])


# ---------------------------------------------------------------------------
# validate（全部门禁）
# ---------------------------------------------------------------------------

def gate_final_sha(cfg: dict, p1_stats: dict) -> dict:
    items = []
    ok = True
    for key in ("final_authority", "final_authority_equiv", "final_authority_orig"):
        expect = (p1_stats.get("inputs", {}).get(key, {}) or {}).get("sha256", "")
        path = cfg[key]
        actual = sha256_file(path)
        same = actual == expect
        ok = ok and same
        items.append({"key": key, "path": path, "expected": expect, "actual": actual,
                      "pass": same})
    return {"idx": "G1", "name": "FINAL 三件套 SHA256 与 p1_stats.inputs 完全一致",
            "pass": ok, "actual": "一致" if ok else "存在差异", "expected": "3/3 一致",
            "details": items,
            "note": "；".join(f"{i['key']}={'OK' if i['pass'] else 'MISMATCH'}" for i in items)}


def gate_step5_distinct_code(con, temp_dir: str) -> dict:
    s5 = os.path.join(temp_dir, "step5.parquet")
    actual = con.execute(f"SELECT COUNT(DISTINCT code) FROM '{_q(s5)}'").fetchone()[0]
    return {"idx": "G2", "name": "Step5 distinct 标准商品编码 = 120047",
            "pass": actual == EXPECTED_STEP5_CODES, "actual": actual,
            "expected": EXPECTED_STEP5_CODES, "note": ""}


def gate_step5_vs_final(con, temp_dir: str) -> dict:
    s5 = os.path.join(temp_dir, "step5.parquet")
    fm = os.path.join(temp_dir, "final_main.parquet")
    final_only = con.execute(
        f"SELECT COUNT(*) FROM (SELECT code, mall, sku FROM '{_q(fm)}' "
        f"EXCEPT SELECT code, mall, sku FROM '{_q(s5)}')").fetchone()[0]
    step5_only = con.execute(
        f"SELECT COUNT(*) FROM (SELECT code, mall, sku FROM '{_q(s5)}' "
        f"EXCEPT SELECT code, mall, sku FROM '{_q(fm)}')").fetchone()[0]
    value_diff = con.execute(
        f"SELECT COUNT(*) FROM '{_q(s5)}' s5 JOIN '{_q(fm)}' fm "
        f"ON s5.code=fm.code AND s5.mall=fm.mall AND s5.sku=fm.sku "
        f"WHERE s5.price IS DISTINCT FROM fm.price "
        f"OR s5.index_val IS DISTINCT FROM fm.index_val "
        f"OR s5.method IS DISTINCT FROM fm.method").fetchone()[0]
    ok = final_only == 0 and step5_only == 0 and value_diff == 0
    return {"idx": "G3", "name": "Step5 vs FINAL 键集合双向=0 且 价格/指数/方式逐行 diff=0",
            "pass": ok,
            "actual": f"FINAL_ONLY={final_only}, STEP5_ONLY={step5_only}, VALUE_DIFF={value_diff}",
            "expected": "0/0/0", "note": "键=(code,mall,sku)"}


def gate_subset_chain(con, temp_dir: str) -> dict:
    t = {n: os.path.join(temp_dir, f"{n}.parquet") for n in ("step1", "step2", "step3", "step4", "step5")}
    layers = [
        ("S2⊆S1_yes",
         f"SELECT COUNT(*) FROM (SELECT code,mall,sku FROM '{_q(t['step2'])}' "
         f"EXCEPT SELECT code,mall,sku FROM '{_q(t['step1'])}' WHERE same_verdict='是')"),
        ("S3⊆S2",
         f"SELECT COUNT(*) FROM (SELECT code,mall,sku FROM '{_q(t['step3'])}' "
         f"EXCEPT SELECT code,mall,sku FROM '{_q(t['step2'])}')"),
        ("S4⊆S3",
         f"SELECT COUNT(*) FROM (SELECT code,mall,sku FROM '{_q(t['step4'])}' "
         f"EXCEPT SELECT code,mall,sku FROM '{_q(t['step3'])}')"),
        ("S5⊆S4",
         f"SELECT COUNT(*) FROM (SELECT code,mall,sku FROM '{_q(t['step5'])}' "
         f"EXCEPT SELECT code,mall,sku FROM '{_q(t['step4'])}')"),
    ]
    details = []
    ok = True
    for name, sql in layers:
        n = con.execute(sql).fetchone()[0]
        ok = ok and n == 0
        details.append(f"{name}={n}")
    return {"idx": "G4", "name": "键集合包含链 S2⊆S1_yes⊆…/S3⊆S2/S4⊆S3/S5⊆S4",
            "pass": ok, "actual": "；".join(details), "expected": "各层差=0", "note": ""}


def gate_step3_unique_sku(con, temp_dir: str) -> dict:
    t3 = os.path.join(temp_dir, "step3.parquet")
    mx = con.execute(
        f"SELECT MAX(cnt) FROM (SELECT code, mall, COUNT(DISTINCT sku) AS cnt "
        f"FROM '{_q(t3)}' GROUP BY code, mall)").fetchone()[0]
    return {"idx": "G5", "name": "Step3 (code,mall) 组内 distinct sku 最大值 ≤ 1",
            "pass": mx is not None and mx <= 1, "actual": mx, "expected": "≤1", "note": ""}


def scan_customer_files(paths: list[str], pattern) -> dict:
    """全量扫描 5 个客户 xlsx: sheet 名 + 全部单元格文本; 同时统计行数。"""
    from openpyxl import load_workbook

    out: dict[str, dict] = {}
    for path in paths:
        entry = {"sheets": [], "rows_total": 0, "hits": [], "hit_total": 0}
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                nrows = 0
                w = find_sensitive_hit(pattern, ws.title)
                if w:
                    entry["hits"].append({"sheet": ws.title, "loc": "sheet名", "word": w})
                    entry["hit_total"] += 1
                for ri, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    nrows += 1
                    for ci, v in enumerate(row, start=1):
                        if v is None:
                            continue
                        w = find_sensitive_hit(pattern, v)
                        if w:
                            entry["hit_total"] += 1
                            if len(entry["hits"]) < 50:
                                entry["hits"].append({"sheet": ws.title,
                                                      "loc": f"R{ri}C{ci}", "word": w})
                entry["sheets"].append({"name": ws.title, "rows": nrows})
                entry["rows_total"] += max(0, nrows - 1)  # 去掉每 sheet 表头
        finally:
            wb.close()
        out[path] = entry
    return out


def gate_leak_scan(scan: dict) -> dict:
    total = sum(e["hit_total"] for e in scan.values())
    files_bad = [os.path.basename(p) for p, e in scan.items() if e["hit_total"] > 0]
    return {"idx": "G6", "name": "客户 xlsx 敏感词扫描命中=0（sheet名+表头+全数据）",
            "pass": total == 0, "actual": f"命中总数={total}; 命中文件={files_bad or '无'}",
            "expected": 0, "details": {os.path.basename(p): e["hits"] for p, e in scan.items()},
            "note": "词表大小写不敏感; “去极值”不在词表（流程名允许）"}


def gate_comment_states(con, temp_dir: str, stage_stats: dict) -> dict:
    rows = []
    ok4 = ok5 = True
    for step in ("step4", "step5"):
        path = os.path.join(temp_dir, f"{step}.parquet")
        dist = dict(con.execute(
            f"SELECT comment_state, COUNT(*) FROM '{_q(path)}' GROUP BY comment_state"
        ).fetchall())
        ref = stage_stats.get("comments", {}).get(step, {})
        same = all(dist.get(k, 0) == ref.get(k, 0) for k in COMMENT_STATES_ORDER)
        if step == "step4":
            ok4 = same
        else:
            ok5 = same
        rows.append({"step": step, "counts": {k: dist.get(k, 0) for k in COMMENT_STATES_ORDER},
                     "not_found_total": dist.get("NOT_FOUND", 0) + dist.get("NO_KEY", 0),
                     "coverage_ref": ref.get("coverage"), "matches_stage_stats": same})
    return {"idx": "G7", "name": "评论四态统计（step4/step5, 与 stage_stats 对照）",
            "pass": ok4 and ok5, "actual": json.dumps(rows, ensure_ascii=False),
            "expected": "与 stage_stats.json 评论四态一致", "details": rows,
            "note": "NO_KEY 单列; not_found_total=NOT_FOUND+NO_KEY"}


def gate_row_counts(cfg: dict, scan: dict, con, temp_dir: str) -> tuple[dict, list]:
    details = []
    ok = True
    for step, (fname, _base) in CUSTOMER_FILES.items():
        path = os.path.join(_wd(cfg, "output_customer"), fname)
        parquet = os.path.join(temp_dir, STEP_PARQUET[step])
        src_rows = con.execute(f"SELECT COUNT(*) FROM '{_q(parquet)}'").fetchone()[0]
        xlsx_rows = scan.get(path, {}).get("rows_total", -1)
        same = xlsx_rows == src_rows
        ok = ok and same
        details.append({"file": fname, "src_parquet": parquet, "src_rows": src_rows,
                        "xlsx_rows": xlsx_rows, "pass": same})
    return {"idx": "G8", "name": "客户 xlsx 行数与源 parquet 一致（全量核对, 非抽查）",
            "pass": ok, "actual": "；".join(f"{d['file']}={d['xlsx_rows']}(源{d['src_rows']})"
                                           for d in details),
            "expected": "全部一致", "details": details, "note": ""}, details


def read_05_sample(path: str, n: int = 5) -> list[list]:
    """05 文件前 n 个数据行（人工可读性抽查用, 只取首个 sheet）。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            rows.append(list(row))
            if i >= n:  # 表头 + n 数据行
                break
        return rows
    finally:
        wb.close()


def gate_disk(cfg: dict) -> dict:
    drive = os.path.splitdrive(cfg["workdir"])[0] + os.sep
    du = shutil.disk_usage(drive)
    parts = {}
    oversized = []
    for sub in ("output_customer", "inventory", "lineage", "reports"):
        d = _wd(cfg, sub)
        parts[sub] = dir_size(d) if os.path.isdir(d) else 0
    total_new = sum(parts.values())
    for root, _dirs, files in os.walk(cfg["workdir"]):
        for fn in files:
            if fn.lower().endswith(".xlsx"):
                p = os.path.join(root, fn)
                if os.path.getsize(p) > XLSX_MAX_BYTES:
                    oversized.append(p)
    cap_ok = total_new <= WORKDIR_NEW_LIMIT_BYTES and not oversized
    return {"idx": "G9", "name": "磁盘与写盘量（E 盘剩余 / workdir 新增 < 5GB / 单 xlsx ≤400MB）",
            "pass": cap_ok,
            "actual": (f"E盘剩余={fmt_bytes(du.free)}; 新增写盘={fmt_bytes(total_new)} "
                       f"({json.dumps(parts, ensure_ascii=False)}); 超限xlsx={oversized or '无'}"),
            "expected": f"新增<{fmt_bytes(WORKDIR_NEW_LIMIT_BYTES)}; 无超限xlsx",
            "details": {"free_bytes": du.free, "parts": parts, "total_new_bytes": total_new,
                        "oversized": oversized},
            "note": f"基线口径: 直接报告当前值（temp 基线 105.9MB 不严格）; 上限 {fmt_bytes(WORKDIR_NEW_LIMIT_BYTES)}"}


def output_file_manifest(cfg: dict) -> list[dict]:
    files = []
    for sub in ("output_customer", "inventory", "lineage", "reports"):
        d = _wd(cfg, sub)
        if not os.path.isdir(d):
            continue
        for fn in sorted(os.listdir(d)):
            p = os.path.join(d, fn)
            if os.path.isfile(p):
                st = os.stat(p)
                files.append({"path": p, "size": st.st_size,
                              "sha256": sha256_file(p) if st.st_size else ""})
    return files


def part_validate(cfg: dict, con, recalc_summary: dict | None) -> tuple[list[dict], dict]:
    """全部 9 项门禁; 写 validation_report.xlsx 与 reconstruction_report.md。"""
    temp_dir = cfg["temp_dir"]
    p1_stats_path = _wd(cfg, "inventory", "p1_stats.json")
    stats_path = _wd(cfg, "reports", "stage_stats.json")
    _require(p1_stats_path, "p1_stats.json")
    _require(stats_path, "stage_stats.json")
    with open(p1_stats_path, "r", encoding="utf-8") as f:
        p1_stats = json.load(f)
    with open(stats_path, "r", encoding="utf-8") as f:
        stage_stats = json.load(f)
    if recalc_summary is None:
        recalc_summary = read_recalc_summary(_wd(cfg, "reports", "price_recalc_check.xlsx"))

    gates = []
    _log("门禁 G1: FINAL 三件套 SHA256")
    gates.append(gate_final_sha(cfg, p1_stats))
    _log("门禁 G2/G3/G4/G5: step5 键值与链包含")
    gates.append(gate_step5_distinct_code(con, temp_dir))
    gates.append(gate_step5_vs_final(con, temp_dir))
    gates.append(gate_subset_chain(con, temp_dir))
    gates.append(gate_step3_unique_sku(con, temp_dir))

    _log("门禁 G6/G8: 客户 xlsx 全量扫描（敏感词+行数）")
    pattern = build_sensitive_pattern()
    paths = [os.path.join(_wd(cfg, "output_customer"), CUSTOMER_FILES[s][0])
             for s in ("01", "02", "03", "04", "05")]
    missing = [p for p in paths if not os.path.isfile(p)]
    if missing:
        print(f"MISSING_INPUT: 客户文件未导出: {missing}（先 --part export）", flush=True)
        sys.exit(2)
    scan = scan_customer_files(paths, pattern)
    gates.append(gate_leak_scan(scan))
    _log("门禁 G7: 评论四态")
    gates.append(gate_comment_states(con, temp_dir, stage_stats))
    g8, row_details = gate_row_counts(cfg, scan, con, temp_dir)
    gates.append(g8)
    _log("门禁 G9: 磁盘与写盘量")
    gates.append(gate_disk(cfg))

    sample_rows = read_05_sample(paths[4], 5)
    manifest = output_file_manifest(cfg)

    _write_validation_xlsx(cfg, gates, row_details, sample_rows, scan, manifest)
    _write_reconstruction_md(cfg, gates, stage_stats, manifest, recalc_summary)

    _log("=== VALIDATION GATES ===")
    for g in gates:
        _log(f"{g['idx']} {'PASS' if g['pass'] else 'FAIL'}: {g['name']} | 实际={g['actual']}")
    hard_fail = [g for g in gates if not g["pass"]]
    if hard_fail:
        print(f"GATE_FAIL: {len(hard_fail)} 项门禁未通过 → exit 2", flush=True)
        return gates, {"scan": scan, "manifest": manifest}
    _log("RESULT: ALL 9 GATES PASS")
    return gates, {"scan": scan, "manifest": manifest}


def _write_validation_xlsx(cfg: dict, gates: list[dict], row_details: list,
                           sample_rows: list, scan: dict, manifest: list) -> None:
    path = _wd(cfg, "reports", "validation_report.xlsx")
    gate_rows = [(g["idx"], g["name"], "PASS" if g["pass"] else "FAIL",
                  str(g["actual"]), str(g["expected"]), g.get("note", "")) for g in gates]
    leak_rows = []
    for p, e in scan.items():
        for h in e["hits"]:
            leak_rows.append((os.path.basename(p), h["sheet"], h["loc"], h["word"]))
    sample = sample_rows or []
    header = sample[0] if sample else customer_headers("05")
    body = sample[1:] if sample else []
    man_rows = [(m["path"], str(m["size"]), m["sha256"]) for m in manifest]
    write_xlsx(path, [
        {"name": "门禁结果", "headers": ["门禁", "名称", "结果", "实际", "期望", "说明"],
         "kinds": ["text", "text", "text", "text", "text", "text"], "rows": gate_rows},
        {"name": "评论四态",
         "headers": ["阶段", "POSITIVE", "ZERO", "EMPTY_SRC", "NOT_FOUND", "NO_KEY",
                     "与stage_stats一致"],
         "kinds": ["text"] + ["int"] * 5 + ["text"],
         "rows": [(d["step"], d["counts"]["POSITIVE"], d["counts"]["ZERO"],
                   d["counts"]["EMPTY_SRC"], d["counts"]["NOT_FOUND"], d["counts"]["NO_KEY"],
                   "是" if d["matches_stage_stats"] else "否")
                  for d in gates[6]["details"]] if len(gates) > 6 else []},
        {"name": "行数核对", "headers": ["文件", "源parquet行数", "客户xlsx行数", "一致"],
         "kinds": ["text", "int", "int", "text"],
         "rows": [(d["file"], d["src_rows"], d["xlsx_rows"],
                   "是" if d["pass"] else "否") for d in row_details]},
        {"name": "抽样5行", "headers": [str(h) for h in header],
         "kinds": customer_kinds("05"), "rows": body},
        {"name": "泄漏扫描", "headers": ["文件", "sheet", "位置", "命中词"],
         "kinds": ["text", "text", "text", "text"], "rows": leak_rows},
        {"name": "产物清单", "headers": ["文件", "size_bytes", "sha256"],
         "kinds": ["text", "text", "text"], "rows": man_rows},
    ])


def _write_reconstruction_md(cfg: dict, gates: list[dict], stage_stats: dict,
                             manifest: list, recalc_summary: dict | None) -> None:
    path = _wd(cfg, "reports", "reconstruction_report.md")
    lines = []
    add = lines.append
    add("# 验收回溯重建报告（F1–F3）")
    add("")
    add(f"- 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    add(f"- 脚本: {os.path.abspath(__file__)}")
    add(f"- workdir: {cfg['workdir']}")
    add("- 数据真值优先级: FINAL_AUTHORITY > 真实历史过程文件 > 已确认规则重判")
    add("")
    add("## A. 概述")
    add("")
    add("本报告为验收回溯重建（F1 抽取 / F2 重建 / F3 导出+验证）的汇总骨架。"
        "客户交付 5 文件见 D; 全部门禁见 E; 复算见 G。规则来源: docs/acceptance 4 文档 + "
        "cnnc_price_algorithms_review_v1; 不对方差公式全量重跑（SPEC §8.3/§23 禁止）。")
    add("")

    g1 = next(g for g in gates if g["idx"] == "G1")
    add("## B. 输入完整性（FINAL 三件套 SHA256, 处理前后一致）")
    add("")
    add("| 文件 | 结果 |")
    add("|---|---|")
    for it in g1["details"]:
        add(f"| {os.path.basename(it['path'])} | {'一致' if it['pass'] else '不一致'} |")
    add("")
    add("17 文件分类清单见 `inventory/source_inventory.xlsx`「分类清单」"
        "（p1_stats.inputs 16 个外部源 + round1 清单 CSV 本体）。")
    add("")

    add("## C. 重建漏斗（F2 stage_stats.json）")
    add("")
    add("| 阶段 | 总行数 | 标品数 | 商城数 | 唯一SKU数 | 上行数 | 剔除数 | 输出数 |")
    add("|---|---|---|---|---|---|---|---|")
    for step in ("step1", "step2", "step3", "step4", "step5"):
        d = stage_stats.get("funnel", {}).get(step, {})
        add(f"| {step} | {d.get('rows')} | {d.get('codes')} | {d.get('malls')} | "
            f"{d.get('unique_skus')} | {d.get('rows_in_prev')} | {d.get('filtered')} | "
            f"{d.get('rows_out')} |")
    add("")
    add("剔除明细: 见 `reports/stage_statistics.xlsx`「剔除明细」sheet。")
    add("")

    add("## D. 客户导出（output_customer）")
    add("")
    add("| 文件 | 大小 | sha256 |")
    add("|---|---|---|")
    for m in manifest:
        if "output_customer" in m["path"]:
            add(f"| {os.path.basename(m['path'])} | {m['size']} | {m['sha256']} |")
    add("")

    add("## E. 门禁结果（G1–G9）")
    add("")
    add("| 门禁 | 名称 | 结果 | 实际 |")
    add("|---|---|---|---|")
    for g in gates:
        add(f"| {g['idx']} | {g['name']} | {'PASS' if g['pass'] else 'FAIL'} | {g['actual']} |")
    add("")

    add("## F. 评论四态（step4 = step5）")
    add("")
    add("| 状态 | 行数 |")
    add("|---|---|")
    cm = stage_stats.get("comments", {}).get("step5", {})
    for k in COMMENT_STATES_ORDER:
        add(f"| {k} | {cm.get(k)} |")
    add(f"| coverage(填充率) | {cm.get('coverage')} |")
    add("")
    add("评论合并优先级与四态定义: SPEC §5（ZS → RECCSV → CVPOW → CVDEL → CVIDX; 禁止无来源填数）。")
    add("")

    add("## G. 价格指数复算（内部检查, 不改交付物）")
    add("")
    if recalc_summary:
        add("| 类别 | 行数 | code数 | match | mismatch | 价格无效跳过行 | max_abs_diff | median_abs_diff |")
        add("|---|---|---|---|---|---|---|---|")
        for cls_name, c in recalc_summary["classes"].items():
            add(f"| {cls_name} | {c['rows']} | {c['codes']} | {c['match']} | {c['mismatch']} | "
                f"{c['skipped_price']} | {c['max_diff']} | {c['median_diff']} |")
        add("")
        add(f"说明: {recalc_summary['note']}")
    else:
        add("状态: 本次运行未包含 recalc（运行 --part recalc 或 all 后更新本节）。")
    add("")

    add("## H. 冲突与已知偏差记录")
    add("")
    add("| 冲突/事件 | 数量 |")
    add("|---|---|")
    for k, v in (stage_stats.get("conflicts") or {}).items():
        if isinstance(v, dict):
            for k2, v2 in v.items():
                add(f"| {k}.{k2} | {v2} |")
        else:
            add(f"| {k} | {v} |")
    add("")
    add("已知偏差与说明（源自 F2 stage_stats.json notes）:")
    for k, v in (stage_stats.get("notes") or {}).items():
        add(f"- {k}: {v}")
    fn = stage_stats.get("funnel", {}).get("step3", {})
    if fn.get("note_min_num"):
        add(f"- min_num 语义: {fn['note_min_num']}")
    if fn.get("s3_comment_unnormalized_note"):
        add(f"- S3 评论: {fn['s3_comment_unnormalized_note']}")
    add("- 价格复算已知偏差: 一级数据源分级取数不可复现 / 评论时间点不同; "
        "历史锚点 0909 mingxi_sku5 acceptance median_abs_diff = 1.005。")
    add("")
    add("## 结论")
    add("")
    add("【结论占位 — 由主 Agent 填写】")
    add("")

    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def read_recalc_summary(path: str) -> dict | None:
    """从 price_recalc_check.xlsx 读回汇总（供独立 validate 使用）。"""
    if not os.path.isfile(path):
        return None
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "汇总" not in wb.sheetnames:
            return None
        ws = wb["汇总"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    classes: dict[str, dict] = {}
    for r in rows[1:]:
        if not r or not r[0] or r[0] == "method混合code数":
            continue
        r = tuple(r) + (None,) * (9 - len(r))  # 尾部空单元格读回会截短
        classes[str(r[0])] = {"rows": r[1], "codes": r[2], "match": r[3], "mismatch": r[4],
                              "skipped_price": r[5], "no_pool_codes": r[6],
                              "max_diff": r[7], "median_diff": r[8]}
    return {"classes": classes, "note": RECALC_NOTE, "from_xlsx": True}


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

WORKDIR_CACHE = ""


def main() -> int:
    global WORKDIR_CACHE
    ap = argparse.ArgumentParser(description="F3 导出+验证层（验收回溯重建）")
    ap.add_argument("--config", required=True)
    ap.add_argument("--part", default="all", choices=["export", "validate", "recalc", "all"])
    args = ap.parse_args()

    cfg = load_config(args.config)
    WORKDIR_CACHE = cfg["workdir"]
    scratch = _setup_scratch_tmp(cfg)
    t0 = time.time()
    _log(f"F3 启动: part={args.part} workdir={cfg['workdir']}")

    con = connect_duckdb(cfg)
    recalc_summary = None
    try:
        if args.part in ("export", "all"):
            part_export(cfg, con)
        if args.part in ("recalc", "all"):
            recalc_summary = part_recalc(cfg, con)
        if args.part in ("validate", "all"):
            gates, _extra = part_validate(cfg, con, recalc_summary)
            if any(not g["pass"] for g in gates):
                return 2
    finally:
        con.close()
        _cleanup_scratch(scratch)
        _log(f"F3 结束: 耗时 {time.time() - t0:.1f}s")
    return 0


if __name__ == "__main__":
    sys.exit(main())
