# -*- coding: utf-8 -*-
"""p1_common.py — F1 抽取层公共工具（验收回溯重建）。

对应规格: reports/SPEC_ACCEPTANCE_RECONSTRUCTION.md §2/§3/§7（冻结版契约）。

职责:
- 配置加载 (config.local.json)
- 键/价格/单元格归一化 (norm_key / norm_price / norm_cell): float 整数去 .0、strip、None→''
- xlsx 流式迭代 (openpyxl read_only=True, 输入文件只读)
- duckdb 连接 (sys.path.insert(libs_dir) 后 import duckdb, 内存库、只读查询)
- sha256 / 文件元信息
- parquet 写出 (pyarrow)
- temp 目录磁盘护栏 (超 3GB → DISK_SAFETY_STOP)

硬红线: 所有输入只读; 只写 workdir 下的 temp\\ 与 inventory\\。
"""
from __future__ import annotations

import hashlib
import json
import os
import sys
from datetime import date, datetime

# ---------------------------------------------------------------------------
# 配置
# ---------------------------------------------------------------------------

REQUIRED_CFG_KEYS = [
    "workdir", "temp_dir", "final_authority", "final_authority_equiv",
    "final_authority_orig", "p25_pool", "s3_flags", "s3_raw", "drop80",
    "drop_new", "dify_recheck", "recall_csv", "comment_index_parquet",
    "delivery_coverage_parquet", "power_match_parquet",
    "delivery_0825_parquet", "zs_result", "algo_src", "libs_dir",
]

_LIBS_DIR = [None]


def load_config(path: str) -> dict:
    """加载 config.local.json 并校验必需键存在。"""
    with open(path, "r", encoding="utf-8") as f:
        cfg = json.load(f)
    missing = [k for k in REQUIRED_CFG_KEYS if k not in cfg]
    if missing:
        raise ValueError(f"config 缺少必需键: {missing}")
    return cfg


def ensure_libs(libs_dir: str) -> None:
    """把 duckdb/pyarrow 所在目录插入 sys.path（幂等）。"""
    if libs_dir and libs_dir not in sys.path:
        sys.path.insert(0, libs_dir)
    _LIBS_DIR[0] = libs_dir or _LIBS_DIR[0]


# ---------------------------------------------------------------------------
# 归一化（SPEC §2）
# ---------------------------------------------------------------------------

def norm_cell(v) -> str:
    """单元格 → 归一化字符串。None→''；float 整数去 .0（21038747.0→'21038747'）；其余 str+strip。"""
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return str(v).strip()
    if isinstance(v, bool):
        return "是" if v else "否"
    if isinstance(v, (datetime, date)):
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d %H:%M:%S").strip()
        return v.strftime("%Y-%m-%d").strip()
    return str(v).strip()


def norm_key(v) -> str:
    """ckey 分量归一化 = norm_cell。"""
    return norm_cell(v)


def norm_price(v):
    """价格解析为 float 或 None。空/非法 → None；'1,234' 去千分位。"""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


# ---------------------------------------------------------------------------
# 文件元信息
# ---------------------------------------------------------------------------

def sha256_file(path: str, chunk: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            b = f.read(chunk)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def file_meta(path: str) -> dict:
    st = os.stat(path)
    return {
        "path": path,
        "size": st.st_size,
        "mtime": datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
    }


# ---------------------------------------------------------------------------
# xlsx 流式迭代（输入只读: load_workbook(read_only=True)）
# ---------------------------------------------------------------------------

def iter_xlsx_rows(path: str, sheet_name: str | None = None):
    """逐行 yield 原始值元组（values_only）。sheet_name=None 取第一个 sheet。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            yield row
    finally:
        wb.close()


def xlsx_header_rows(path: str, sheet_name: str | None = None):
    """返回 (header: list[str], data_rows: 迭代器 of (excel_row_no:int, tuple))。

    header = 第一个非全空行; data_rows 从其后一行开始。excel_row_no 从 1 计
    （含表头行自身），即表头=1、首个数据行=2。
    """
    gen = iter_xlsx_rows(path, sheet_name)
    row_no = 0
    header = None
    for row in gen:
        row_no += 1
        if row and any(v is not None and str(v).strip() != "" for v in row):
            header = [norm_cell(v) for v in row]
            break
    if header is None:
        raise ValueError(f"{path} sheet={sheet_name}: 未找到表头行")

    def _data():
        nonlocal row_no
        for row in gen:
            row_no += 1
            yield row_no, row

    return header, _data()


def column_index(header: list[str], wanted: dict) -> dict:
    """按归一化表头文本精确匹配，返回 field→列下标。缺失即报错（列名必须全对上）。"""
    norm_hdr = [norm_cell(h) for h in header]
    idx = {}
    missing = []
    for field, want in wanted.items():
        w = norm_cell(want)
        found = None
        for i, h in enumerate(norm_hdr):
            if h == w:
                found = i
                break
        if found is None:
            missing.append(f"{field}←{want!r}")
        else:
            idx[field] = found
    if missing:
        raise ValueError(f"表头缺少列: {missing}; 实际表头={header}")
    return idx


def cell(row: tuple, i: int):
    """取第 i 列，行短于列数时返回 None。"""
    return row[i] if i < len(row) else None


# ---------------------------------------------------------------------------
# duckdb
# ---------------------------------------------------------------------------

def connect_duckdb(cfg: dict):
    """内存 duckdb 连接。仅用于只读查询与向 temp\\ 写 parquet（SPEC §7）。"""
    ensure_libs(cfg["libs_dir"])
    import duckdb

    return duckdb.connect()


# ---------------------------------------------------------------------------
# parquet 写出
# ---------------------------------------------------------------------------

_PA_TYPES = {
    "string": "string",
    "double": "float64",
    "float": "float32",
    "int64": "int64",
    "int32": "int32",
}


def write_parquet(columns: dict, types: dict, out_path: str) -> str:
    """按列写出 parquet。types: field→'string'|'double'|'int64'|'int32'。"""
    try:
        import pyarrow as pa
        import pyarrow.parquet as pq
    except ImportError:
        if _LIBS_DIR[0] and _LIBS_DIR[0] not in sys.path:
            sys.path.insert(0, _LIBS_DIR[0])
        import pyarrow as pa
        import pyarrow.parquet as pq

    arrays = {}
    for name, values in columns.items():
        t = types[name]
        arrays[name] = pa.array(values, type=pa.type_for_alias(_PA_TYPES[t]))
    table = pa.table(arrays)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    pq.write_table(table, out_path)
    return out_path


# ---------------------------------------------------------------------------
# 磁盘护栏（SPEC §7: temp 中间产物 >3GB → DISK_SAFETY_STOP）
# ---------------------------------------------------------------------------

DISK_LIMIT_BYTES = 3 * 1024 ** 3


def dir_size(path: str) -> int:
    total = 0
    for root, _dirs, files in os.walk(path):
        for fn in files:
            try:
                total += os.path.getsize(os.path.join(root, fn))
            except OSError:
                pass
    return total


def guard_disk_limit(temp_dir: str, limit_bytes: int = DISK_LIMIT_BYTES) -> int:
    size = dir_size(temp_dir)
    if size > limit_bytes:
        print(f"DISK_SAFETY_STOP: temp 目录已达 {size} 字节 > 上限 {limit_bytes} 字节", flush=True)
        sys.exit(3)
    return size


def fmt_bytes(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.2f}{unit}" if unit != "B" else f"{n}B"
        n /= 1024.0
    return f"{n}GB"
