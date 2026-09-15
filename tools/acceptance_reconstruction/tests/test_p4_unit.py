# -*- coding: utf-8 -*-
"""p4 单元测试: 插列位置 / 评论值解析与来源链 / map 构建 / 05 常规插列与回退重建。

全部使用合成小数据（mini parquet / mini xlsx / 内存工作簿）, 不触碰真实客户数据文件。
"""
from __future__ import annotations

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import p4_add_comments as P4

# ---------------------------------------------------------------------------
# 列规格与插列位置
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("step", ["01", "02", "03"])
def test_headers_kinds_insert_position(step):
    hdr = P4.headers_with_comment(step)
    kinds = P4.kinds_with_comment(step)
    orig_hdr = ["标准商品编码", "标准商品名称", "互联网商品编码（SKUID）", "互联网商城编码",
                "互联网商品名称", "互联网链接", "互联网商品价格"]
    assert hdr[:7] == orig_hdr and hdr[7] == "评论数"
    assert kinds[:7] == ["text_at", "text", "text_at", "text", "text", "text_at", "price"]
    assert kinds[7] == "int"
    if step == "01":
        assert hdr[8] == "同款判断" and kinds[8] == "text" and len(hdr) == 9
    else:
        assert len(hdr) == 8


def test_headers_bad_step_raises():
    # 第7列不是互联网商品价格时必须报错（防列序漂移）
    import p4_add_comments as m
    orig = m.customer_headers
    m.customer_headers = lambda s: ["a", "b", "c", "d", "e", "f", "g"]
    try:
        with pytest.raises(ValueError):
            m.headers_with_comment("01")
    finally:
        m.customer_headers = orig


# ---------------------------------------------------------------------------
# 评论值解析
# ---------------------------------------------------------------------------


def test_zs_comment_to_int():
    assert P4.zs_comment_to_int("0") == 0
    assert P4.zs_comment_to_int("12") == 12
    assert P4.zs_comment_to_int(7) == 7
    assert P4.zs_comment_to_int(4.0) == 4
    assert P4.zs_comment_to_int(None) is None
    assert P4.zs_comment_to_int("") is None
    assert P4.zs_comment_to_int("500+") is None      # 非数字展示串 → 跳过
    assert P4.zs_comment_to_int("1万+") is None
    assert P4.zs_comment_to_int("abc") is None


def _maps():
    final_map = {
        ("c1", "s1", "jd"): ("POSITIVE", 11.0),
        ("c2", "s2", "jd"): ("ZERO", 0.0),
        ("c3", "s3", "jd"): ("EMPTY_SRC", None),
        ("c4", "s4", "jd"): ("NOT_FOUND", None),
        ("c5", "s5", "jd"): ("NO_KEY", None),
        # k6 同时存在于 FINAL(缺失态) 与 CVIDX → is_final 行不得回退到候选链
        ("c6", "s6", "jd"): ("EMPTY_SRC", None),
    }
    cvidx_map = {("c6", "s6", "jd"): 7}
    zs_map = {("jd", "c9"): 0, ("suning", "c10"): 5}
    return final_map, cvidx_map, zs_map


def test_resolve_comment_final_chain():
    fm, cm, zm = _maps()
    r = lambda c, s: P4.resolve_comment(True, c, s, "jd", fm, cm, zm)  # noqa: E731
    assert r("c1", "s1") == (11, "FINAL")
    assert r("c2", "s2") == (0, "FINAL")
    assert r("c3", "s3") == (None, "MISS")
    assert r("c4", "s4") == (None, "MISS")
    assert r("c5", "s5") == (None, "MISS")


def test_resolve_comment_final_never_falls_back():
    fm, cm, zm = _maps()
    # is_final 行即便键在候选链有值, 也必须按 FINAL 缺失态写空
    assert P4.resolve_comment(True, "c6", "s6", "jd", fm, cm, zm) == (None, "MISS")


def test_resolve_comment_candidate_chain():
    fm, cm, zm = _maps()
    r = lambda c, s, m: P4.resolve_comment(False, c, s, m, fm, cm, zm)  # noqa: E731
    assert r("c6", "s6", "jd") == (7, "CVIDX")
    assert r("c9", "s9", "jd") == (0, "ZS")           # CVIDX miss → ZS '0'→0
    assert r("c10", "s10", "suning") == (5, "ZS")
    assert r("c99", "s99", "jd") == (None, "MISS")    # 两层都 miss → 空


def test_assemble_row():
    row = ("c", "n", "s", "m", "t", "u", 1.5, "是", True)
    out = P4.assemble_row(row, 8, 5)
    assert out == ("c", "n", "s", "m", "t", "u", 1.5, 5, "是")
    assert out[7] == 5
    row2 = ("c", "n", "s", "m", "t", "u", 1.5, None)  # 02/03 无同款判断
    assert P4.assemble_row(row2, 7, None) == ("c", "n", "s", "m", "t", "u", 1.5, None)


# ---------------------------------------------------------------------------
# map 构建（mini parquet / mini xlsx）
# ---------------------------------------------------------------------------


def test_build_final_comment_map(tmp_path):
    p = tmp_path / "comments_final.parquet"
    pq.write_table(pa.table({
        "code": pa.array(["c1", "c2", "c3"], type=pa.string()),
        "sku": pa.array(["s1", "s2", "s3"], type=pa.string()),
        "mall": pa.array(["jd", "jd", "jd"], type=pa.string()),
        "state": pa.array(["POSITIVE", "ZERO", "EMPTY_SRC"], type=pa.string()),
        "comment_value": pa.array([11.0, 0.0, None], type=pa.float64()),
    }), str(p))
    import duckdb
    con = duckdb.connect()
    try:
        m = P4.build_final_comment_map(con, str(p))
        assert m[("c1", "s1", "jd")] == ("POSITIVE", 11.0)
        assert m[("c2", "s2", "jd")] == ("ZERO", 0.0)
        assert m[("c3", "s3", "jd")] == ("EMPTY_SRC", None)
    finally:
        con.close()


def test_build_cvidx_map(tmp_path):
    p = tmp_path / "comment_index.parquet"
    pq.write_table(pa.table({
        "标准商品编码": pa.array(["c1", "c2", "c3"], type=pa.string()),
        "SKUID": pa.array(["s1", "s2", "s3"], type=pa.string()),
        "商城编码": pa.array(["jd", "jd", "jd"], type=pa.string()),
        "评论数": pa.array([11, None, 5], type=pa.int64()),
    }), str(p))
    import duckdb
    con = duckdb.connect()
    try:
        m = P4.build_cvidx_map(con, [("c1", "s1", "jd"), ("c2", "s2", "jd"),
                                     ("c3", "s3", "jd"), ("cX", "sX", "jd")], str(p))
        assert m == {("c1", "s1", "jd"): 11, ("c3", "s3", "jd"): 5}  # NULL 评论=未命中
        assert P4.build_cvidx_map(con, [], str(p)) == {}
    finally:
        con.close()


def test_build_zs_map(tmp_path):
    p = tmp_path / "zs.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["商品编码", "商城标识", "评论"])
    ws.append(["c1", "suning", "0"])
    ws.append(["c2", "jd", None])
    ws.append(["c3", "jd", "500+"])     # 非数字 → 不入 map
    ws.append(["c4", "jd", 7])
    ws.append(["c5", "jd", "12"])
    wb.save(p)
    m = P4.build_zs_map(str(p))
    assert m == {("suning", "c1"): 0, ("jd", "c4"): 7, ("jd", "c5"): 12}


# ---------------------------------------------------------------------------
# 05 常规插列 / 列宽平移 / 回退重建
# ---------------------------------------------------------------------------

HDR_05 = ["标准商品编码", "标准商品名称", "互联网商品编码（SKUID）", "互联网商城编码",
          "互联网商品名称", "互联网链接", "互联网商品价格", "本次处理后的价格指数",
          "价格指数计算方式"]

FINAL_MAP_05 = {
    ("100000011587", "111", "jd"): ("POSITIVE", 5.0),
    ("100000011588", "222", "jd"): ("EMPTY_SRC", None),
    ("100000011589", "333", "zjmi"): ("ZERO", 0.0),
}


def _mini_05_workbook(wb: Workbook) -> None:
    """构造 05 缩样: 主表 9列×3数据行 + 总结 sheet（合并单元格+公式）。"""
    ws = wb.create_sheet("最终处理结果")
    ws.append(HDR_05)
    ws.append(["100000011587", "名称A", "111", "jd", "标题A", "https://x/1", 95.19, 96.09,
               "平均价计算指导价"])
    ws.append(["100000011588", "名称B", "222", "jd", "标题B", "https://x/2", 10.0, None,
               "不计算"])
    ws.append(["100000011589", "名称C", "333", "zjmi", "标题C", "https://x/3", 7.5, 8.0,
               "幂次归一化"])
    g1 = ws.cell(1, 7)
    g1.font = Font(name="微软雅黑", size=11, bold=True)
    g1.fill = PatternFill(fill_type="solid", fgColor="FF1F4E78")
    g1.alignment = Alignment(horizontal="center")
    ws["G2"].number_format = "#,##0.00"
    ws["H1"].font = Font(name="微软雅黑", size=11, bold=True)
    ws.auto_filter.ref = "A1:I4"
    ws.column_dimensions["G"].width = 24.05
    ws.column_dimensions["H"].width = 19.475
    ws.column_dimensions["I"].width = 24.0

    ws2 = wb.create_sheet("四项问题处理总结")
    ws2["A1"] = "四项问题处理总结"
    ws2["A2"] = "汇总行"
    ws2.merge_cells("A2:F2")
    ws2["A3"] = "项1"
    ws2["B3"] = 3
    ws2["A4"] = "项2"
    ws2["B4"] = 4
    ws2["A5"] = "=SUM(B3:B4)"


def test_shift_column_dimensions():
    wb = Workbook()
    ws = wb.active
    ws.column_dimensions["C"].width = 30.0
    ws.column_dimensions["E"].width = 40.0
    P4.shift_column_dimensions(ws, 3, 12.0)
    assert ws.column_dimensions["C"].width == 12.0       # 新列
    assert ws.column_dimensions["D"].width == 30.0       # C → D
    assert ws.column_dimensions["F"].width == 40.0       # E → F
    assert "E" not in ws.column_dimensions or ws.column_dimensions["E"].width is None


def test_apply_comment_column_normal():
    wb = Workbook()
    _mini_05_workbook(wb)
    ws = wb["最终处理结果"]
    stats = P4.apply_comment_column_normal(ws, FINAL_MAP_05)
    assert stats == {"filled": 2, "empty": 1, "rows": 3}
    assert ws.max_column == 10
    h = ws.cell(1, 8)
    assert h.value == "评论数"
    g1, old_h1 = ws.cell(1, 7), ws.cell(1, 9)
    assert h.font.bold and h.font.name == g1.font.name
    assert h.fill.fgColor.rgb == g1.fill.fgColor.rgb
    assert h.alignment.horizontal == g1.alignment.horizontal
    assert old_h1.value == "本次处理后的价格指数"
    assert ws.cell(2, 8).value == 5
    assert ws.cell(3, 8).value is None
    assert ws.cell(4, 8).value == 0
    assert ws.auto_filter.ref == "A1:J4"
    assert ws.column_dimensions["H"].width == 12.0
    assert ws.column_dimensions["I"].width == 19.475     # 原 H 列宽右移
    assert ws.column_dimensions["J"].width == 24.0       # 原 I 列宽右移
    # 总结 sheet 不受影响（含公式/合并）
    ws2 = wb["四项问题处理总结"]
    assert ws2["A5"].value == "=SUM(B3:B4)"
    assert "A2:F2" in [str(r) for r in ws2.merged_cells.ranges]


def test_process_05_normal_roundtrip(tmp_path):
    src = tmp_path / "05.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    _mini_05_workbook(wb)
    wb.save(src)
    wb.close()
    out = tmp_path / "05_new.xlsx"
    stats = P4.process_05_normal(str(src), str(out), FINAL_MAP_05, rss_limit=None)
    assert stats == {"filled": 2, "empty": 1, "rows": 3}
    wb2 = load_workbook(out)
    try:
        ws = wb2["最终处理结果"]
        assert ws.max_column == 10 and ws.max_row == 4
        assert ws.cell(1, 8).value == "评论数"
        assert ws.cell(2, 8).value == 5 and ws.cell(4, 8).value == 0
        assert ws.cell(2, 7).value == 95.19
        assert ws.cell(2, 7).number_format == "#,##0.00"
        assert wb2["四项问题处理总结"]["A5"].value == "=SUM(B3:B4)"
        assert "A2:F2" in [str(r) for r in wb2["四项问题处理总结"].merged_cells.ranges]
    finally:
        wb2.close()


def test_rebuild_05_write_only(tmp_path):
    src = tmp_path / "05_backup.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    _mini_05_workbook(wb)
    wb.save(src)
    wb.close()
    out = tmp_path / "05_rebuilt.xlsx"
    stats = P4.rebuild_05_write_only(str(src), str(out), FINAL_MAP_05)
    assert stats == {"filled": 2, "empty": 1, "rows": 3}

    wb2 = load_workbook(out)
    try:
        assert wb2.sheetnames == ["最终处理结果", "四项问题处理总结"]
        ws = wb2["最终处理结果"]
        assert ws.max_column == 10
        assert [ws.cell(1, i).value for i in range(1, 11)] == \
            HDR_05[:7] + ["评论数"] + HDR_05[7:]
        assert ws.cell(2, 8).value == 5
        assert ws.cell(3, 8).value is None
        assert ws.cell(4, 8).value == 0
        assert ws.cell(2, 7).value == 95.19
        assert ws.cell(2, 7).number_format == "0.00"      # p3 价格口径
        assert ws.cell(2, 1).number_format == "@"          # 编码文本格式
        assert ws.auto_filter.ref == "A1:J4"
        ws2 = wb2["四项问题处理总结"]
        assert ws2["A5"].value == "=SUM(B3:B4)"
        assert ws2["B3"].value == 3 and ws2["B4"].value == 4
    finally:
        wb2.close()


# ---------------------------------------------------------------------------
# 验证辅助
# ---------------------------------------------------------------------------


def test_cells_eq():
    assert P4._cells_eq(1.5, 1.5)
    assert P4._cells_eq(None, None)
    assert P4._cells_eq(float("nan"), float("nan"))
    assert not P4._cells_eq(1.5, 1.6)
    assert not P4._cells_eq("a", "b")
    assert not P4._cells_eq(None, "")


def test_val_eq_empty_string_equivalence():
    # openpyxl 往返把 '' 单元格写成空单元格 → None 与 '' 等价
    assert P4._val_eq(None, "")
    assert P4._val_eq("", None)
    assert P4._val_eq("x", "x")
    assert not P4._val_eq(None, "x")
    assert not P4._val_eq(" ", None)     # 空格不等价于空
    assert not P4._val_eq(0, None)
    assert P4._val_eq(float("nan"), float("nan"))
