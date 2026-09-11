# -*- coding: utf-8 -*-
"""p3 单元测试: 格式化写出 / sheet 拆分 / 敏感词扫描 / 复算口径 / 门禁 FAIL 路径。

全部使用合成小数据（mini parquet / mini xlsx）, 不触碰真实数据文件。
"""
from __future__ import annotations

import json
import os
import sys

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

import p3_format as F
import p3_export_validate as P


# ---------------------------------------------------------------------------
# 拆分计划
# ---------------------------------------------------------------------------

def test_split_plan_single():
    assert F.split_sheet_names("数据", 545163) == [("数据", 545163)]
    assert F.split_sheet_names("数据", 0) == [("数据", 0)]


def test_split_plan_boundary():
    # 恰好等于上限 → 单 sheet
    assert F.split_sheet_names("d", F.SHEET_MAX_DATA_ROWS) == [("d", F.SHEET_MAX_DATA_ROWS)]


def test_split_plan_overflow():
    plan = F.split_sheet_names("d", F.SHEET_MAX_DATA_ROWS + 1)
    assert plan[0] == ("d_1", F.SHEET_MAX_DATA_ROWS)
    assert plan[1] == ("d_2", 1)
    plan2 = F.split_sheet_names("d", F.SHEET_MAX_DATA_ROWS * 2 + 100)
    assert [n for n, _ in plan2] == ["d_1", "d_2", "d_3"]
    assert sum(n for _, n in plan2) == F.SHEET_MAX_DATA_ROWS * 2 + 100


# ---------------------------------------------------------------------------
# convert_cell（写入前值/格式转换）
# ---------------------------------------------------------------------------

def test_convert_text_at_strips_and_none():
    assert F.convert_cell("text_at", " 21038747 ") == ("21038747", "@")
    assert F.convert_cell("text_at", "") == (None, None)
    assert F.convert_cell("text_at", None) == (None, None)


def test_convert_price_format():
    assert F.convert_cell("price", 24.0) == (24.0, "0.00")
    assert F.convert_cell("price", None) == (None, None)


def test_convert_int_from_float():
    assert F.convert_cell("int", 1928.0) == (1928, None)
    assert F.convert_cell("int", None) == (None, None)
    assert F.convert_cell("int", "12") == (12, None)


def test_convert_index_empty_is_none_numeric_is_float():
    assert F.convert_cell("index", "") == (None, None)          # 空串→写空
    assert F.convert_cell("index", "22.98") == (22.98, "0.00")
    assert F.convert_cell("index", "abc") == ("abc", "@")       # 非数值兜底为文本


# ---------------------------------------------------------------------------
# xlsx 写出: 冻结/筛选/格式/拆 sheet
# ---------------------------------------------------------------------------

def _rows(n: int, width: int = 3):
    for i in range(n):
        if width == 3:
            yield (f"10{i:06d}", 10.0 + i, "22.98" if i % 2 == 0 else "")
        else:
            yield tuple(f"v{i}-{j}" for j in range(width))


def test_write_xlsx_styles_and_values(tmp_path):
    path = str(tmp_path / "t.xlsx")
    res = F.write_xlsx(path, [{
        "name": "T", "headers": ["编码", "价格", "指数"],
        "kinds": ["text_at", "price", "index"], "rows": _rows(4),
    }])
    assert res == [{"name": "T", "rows": 4}]
    wb = load_workbook(path)
    ws = wb["T"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:C5"
    assert ws["A2"].value == "10000000" and ws["A2"].number_format == "@"
    assert ws["B2"].value == 10.0 and ws["B2"].number_format == "0.00"
    assert ws["C2"].value == 22.98 and ws["C2"].number_format == "0.00"
    assert ws["C3"].value is None  # 空指数写空
    wb.close()


def test_write_xlsx_split_sheets(tmp_path):
    path = str(tmp_path / "t.xlsx")
    res = F.write_xlsx(path, [{
        "name": "T", "headers": ["a", "b", "c"],
        "kinds": ["text", "text", "text"], "rows": _rows(5, width=3),
    }], max_rows=2)
    assert [r["name"] for r in res] == ["T_1", "T_2", "T_3"]
    assert [r["rows"] for r in res] == [2, 2, 1]
    wb = load_workbook(path)
    assert wb.sheetnames == ["T_1", "T_2", "T_3"]
    for name in wb.sheetnames:
        ws = wb[name]
        assert [c.value for c in ws[1]] == ["a", "b", "c"]  # 表头一致
        assert ws.freeze_panes == "A2"
    assert wb["T_1"].auto_filter.ref == "A1:C3"
    assert wb["T_3"].auto_filter.ref == "A1:C2"
    wb.close()


def test_write_xlsx_customer_headers_frozen_spec():
    assert F.customer_headers("01") == [
        "标准商品编码", "标准商品名称", "互联网商品编码（SKUID）", "互联网商城编码",
        "互联网商品名称", "互联网链接", "互联网商品价格", "同款判断(是/否)"]
    assert F.customer_headers("02") == F.customer_headers("01")[:7]
    assert F.customer_headers("03") == F.customer_headers("02")
    assert F.customer_headers("04") == F.customer_headers("02") + ["评论数", "人工审核结果"]
    assert F.customer_headers("05") == (F.customer_headers("02")
                                        + ["评论数", "本次处理后的价格指数", "价格指数计算方式"])


# ---------------------------------------------------------------------------
# 敏感词扫描
# ---------------------------------------------------------------------------

def test_scanner_hits_case_insensitive():
    pat = F.build_sensitive_pattern()
    hit = F.find_sensitive_hit(pat, "内部基准价明细")
    assert hit in ("内部基准", "基准价")  # 长词优先, 任一命中即拦截
    assert F.find_sensitive_hit(pat, "See BASELINE value") == "BASELINE"  # 命中原文片段
    assert F.find_sensitive_hit(pat, "xverdict_sourcey") == "verdict_source"
    assert F.find_sensitive_hit(pat, "RECONSTRUCTED_FINAL") == "RECONSTRUCTED"
    assert F.find_sensitive_hit(pat, "ReferencePrice") == "ReferencePrice"  # 大小写不敏感命中
    assert F.find_sensitive_hit(pat, None) is None


def test_scanner_clean_and_whitelist():
    pat = F.build_sensitive_pattern()
    assert F.find_sensitive_hit(pat, "去极值后结果数据") is None   # 流程名允许
    assert F.find_sensitive_hit(pat, "价格指数计算结果数据") is None
    assert F.find_sensitive_hit(pat, "同款判断(是/否)") is None
    assert F.find_sensitive_hit(pat, "https://item.jd.com/1.html") is None


def test_scan_customer_files_hit_and_clean(tmp_path):
    pat = F.build_sensitive_pattern()
    bad = str(tmp_path / "bad.xlsx")
    F.write_xlsx(bad, [{"name": "S", "headers": ["a", "b"], "kinds": ["text", "text"],
                        "rows": iter([("ok", "value with 内部参考值 inside")])}])
    good = str(tmp_path / "good.xlsx")
    F.write_xlsx(good, [{"name": "数据", "headers": ["a", "b"], "kinds": ["text", "text"],
                         "rows": iter([("10", "正常文本")])}])
    scan = P.scan_customer_files([bad, good], pat)
    assert scan[bad]["hit_total"] == 1
    assert scan[bad]["rows_total"] == 1
    assert scan[good]["hit_total"] == 0
    # 敏感词出现在 sheet 名
    bad2 = str(tmp_path / "bad2.xlsx")
    F.write_xlsx(bad2, [{"name": "source_sheet", "headers": ["a"], "kinds": ["text"],
                         "rows": iter([("x",)])}])
    scan2 = P.scan_customer_files([bad2], pat)
    assert scan2[bad2]["hit_total"] == 1


def test_gate_leak_scan_fail_path():
    g = P.gate_leak_scan({"a.xlsx": {"hit_total": 2, "hits": [{"word": "baseline"}]},
                          "b.xlsx": {"hit_total": 0, "hits": []}})
    assert g["pass"] is False
    g2 = P.gate_leak_scan({"a.xlsx": {"hit_total": 0, "hits": []}})
    assert g2["pass"] is True


# ---------------------------------------------------------------------------
# 复算口径（Decimal CEIL2 / HALF_UP2 / 缺失评论按 0）
# ---------------------------------------------------------------------------

def test_comment_weight_missing_states():
    for st in ("EMPTY_SRC", "NOT_FOUND", "NO_KEY"):
        assert F.comment_weight(None, st) == F.Decimal(1)
    assert F.comment_weight(1928.0, "POSITIVE") == F.Decimal(1929)
    assert F.comment_weight(0.0, "ZERO") == F.Decimal(1)
    assert F.comment_weight(None, "POSITIVE") == F.Decimal(1)  # 无值按 0, 不造数


def test_recalc_power_ceil2():
    # (10+20+25)/3 = 18.333... → CEIL2 = 18.34
    rows = [(10.0, None, "NOT_FOUND"), (20.0, 0.0, "ZERO"), (25.0, None, "EMPTY_SRC")]
    assert F.recalc_power_index(rows) == F.Decimal("18.34")


def test_recalc_weighted_and_ceil():
    # (10×1 + 20×2 + 30×1) / (1+2+1) = 80/4 = 20.00
    rows = [(10.0, 0.0, "ZERO"), (20.0, 1.0, "POSITIVE"), (30.0, None, "NOT_FOUND")]
    assert F.recalc_power_index(rows) == F.Decimal("20.00")


def test_recalc_avg_halfup2():
    # (10+20+25)/3 = 18.333... → HALF_UP2 = 18.33
    rows = [(10.0, None, "NOT_FOUND"), (20.0, 0.0, "ZERO"), (25.0, None, "EMPTY_SRC")]
    assert F.recalc_avg_price(rows) == F.Decimal("18.33")


def test_recalc_invalid_price_skipped_and_empty_pool():
    assert F.recalc_power_index([(None, 1.0, "POSITIVE")]) is None
    assert F.recalc_avg_price([]) is None
    ok, bad = F.partition_valid_price([(10.0, 0, "ZERO"), (None, 0, "ZERO"),
                                       ("abc", 0, "ZERO")])
    assert len(ok) == 1 and bad == 2


def test_index_val_equal():
    assert F.index_val_equal(F.Decimal("22.98"), "22.98")
    assert F.index_val_equal(F.Decimal("22.98"), " 22.980 ")  # 数值相等口径
    assert not F.index_val_equal(F.Decimal("22.98"), "")
    assert not F.index_val_equal(F.Decimal("22.98"), None)
    assert not F.index_val_equal(F.Decimal("22.98"), "abc")


# ---------------------------------------------------------------------------
# 门禁函数（合成 mini parquet; 覆盖 FAIL 路径）
# ---------------------------------------------------------------------------

KEY_COLS = ["code", "mall", "sku"]


def _pq(tmp_path, name: str, cols: dict, rows: list[tuple]):
    path = str(tmp_path / name)
    pq.write_table(pa.table({c: pa.array([r[i] for r in rows], type=pa.string())
                             for i, c in enumerate(cols)}), path)
    return path


def _make_step_chain(tmp_path, *, break_subset: bool = False, dup_sku: bool = False,
                     price_diff: bool = False, codes=("1001", "1002", "1003")):
    """合成 5 步 + final_main。默认: S2⊆S1yes⊆… 成立, step5 与 final_main 一致。"""
    s1 = [(codes[0], "jd", "1", "是"), (codes[1], "jd", "2", "是"), (codes[2], "jd", "3", "否")]
    s2 = s1[:2]
    s3 = s2 + ([(codes[1], "tb", "9")] if break_subset else [])
    s4 = list(s2)
    s5 = [(codes[0], "jd", "1", "10.0", "22.98", "幂次归一化（评论数加权平均价）计算指导价"),
          (codes[1], "jd", "2", "20.0", "18.33", "平均价计算指导价")]
    if dup_sku:
        s3 = s3 + [(codes[0], "jd", "77")]  # 同 (code,mall) 第二个 sku → G5 FAIL
    fm = list(s5)
    if price_diff:
        fm[0] = (codes[0], "jd", "1", "99.0", "22.98", "幂次归一化（评论数加权平均价）计算指导价")

    _pq(tmp_path, "step1.parquet", KEY_COLS + ["same_verdict"], s1)
    _pq(tmp_path, "step2.parquet", KEY_COLS, s2)
    _pq(tmp_path, "step3.parquet", KEY_COLS, s3)
    _pq(tmp_path, "step4.parquet", KEY_COLS, s4)
    _pq(tmp_path, "step5.parquet", KEY_COLS + ["price", "index_val", "method"], s5)
    _pq(tmp_path, "final_main.parquet", KEY_COLS + ["price", "index_val", "method"], fm)


def test_gate_subset_chain_pass_and_fail(tmp_path):
    import duckdb
    con = duckdb.connect()
    _make_step_chain(tmp_path, break_subset=False)
    g = P.gate_subset_chain(con, str(tmp_path))
    assert g["pass"] is True
    _make_step_chain(tmp_path, break_subset=True)
    g2 = P.gate_subset_chain(con, str(tmp_path))
    assert g2["pass"] is False
    assert "S3⊆S2=1" in g2["actual"]
    con.close()


def test_gate_step5_vs_final_pass_and_fail(tmp_path):
    import duckdb
    con = duckdb.connect()
    _make_step_chain(tmp_path, price_diff=False)
    g = P.gate_step5_vs_final(con, str(tmp_path))
    assert g["pass"] is True
    _make_step_chain(tmp_path, price_diff=True)
    g2 = P.gate_step5_vs_final(con, str(tmp_path))
    assert g2["pass"] is False
    assert "VALUE_DIFF=1" in g2["actual"]
    con.close()


def test_gate_step3_unique_sku_fail(tmp_path):
    import duckdb
    con = duckdb.connect()
    _make_step_chain(tmp_path, dup_sku=True)
    g = P.gate_step3_unique_sku(con, str(tmp_path))
    assert g["pass"] is False and g["actual"] == 2
    _make_step_chain(tmp_path, dup_sku=False)
    g2 = P.gate_step3_unique_sku(con, str(tmp_path))
    assert g2["pass"] is True and g2["actual"] == 1
    con.close()


def test_gate_step5_distinct_code(tmp_path):
    import duckdb
    con = duckdb.connect()
    _make_step_chain(tmp_path)
    g = P.gate_step5_distinct_code(con, str(tmp_path))
    assert g["pass"] is False          # 合成数据只有 2 个 code ≠ 120047
    assert g["actual"] == 2
    con.close()


def test_gate_final_sha_pass_and_fail(tmp_path):
    fa = tmp_path / "fa.xlsx"
    fa.write_bytes(b"AAA")
    other = tmp_path / "other.xlsx"
    other.write_bytes(b"BBB")
    from p1_common import sha256_file
    p1_stats = {"inputs": {
        "final_authority": {"sha256": sha256_file(str(fa))},
        "final_authority_equiv": {"sha256": sha256_file(str(fa))},
        "final_authority_orig": {"sha256": sha256_file(str(fa))},
    }}
    cfg = {"final_authority": str(fa), "final_authority_equiv": str(fa),
           "final_authority_orig": str(fa)}
    g = P.gate_final_sha(cfg, p1_stats)
    assert g["pass"] is True
    cfg_bad = dict(cfg, final_authority_orig=str(other))
    g2 = P.gate_final_sha(cfg_bad, p1_stats)
    assert g2["pass"] is False


def test_gate_row_counts_mismatch(tmp_path):
    import duckdb
    con = duckdb.connect()
    _make_step_chain(tmp_path)
    cfg = {"workdir": str(tmp_path)}
    outdir = os.path.join(str(tmp_path), "output_customer")
    os.makedirs(outdir, exist_ok=True)
    # step1=3 行(含否), step2..step5 各 2 行
    expect = {"01": 3, "02": 2, "03": 2, "04": 2, "05": 2}
    scan = {}
    for step, (fname, _base) in P.CUSTOMER_FILES.items():
        path = os.path.join(outdir, fname)
        scan[path] = {"rows_total": expect[step], "hit_total": 0, "hits": [], "sheets": []}
    g, _det = P.gate_row_counts(cfg, scan, con, str(tmp_path))
    assert g["pass"] is True
    first = next(iter(scan))
    scan[first]["rows_total"] = 1
    g2, _det2 = P.gate_row_counts(cfg, scan, con, str(tmp_path))
    assert g2["pass"] is False
    con.close()


def test_read_recalc_summary_roundtrip(tmp_path):
    summary = {"generated_at": "t", "classes": {
        "幂次": {"rows": 287163, "codes": 100, "match": 99, "mismatch": 1,
                "skipped_price": 0, "no_pool_codes": 0, "max_diff": "0.01",
                "median_diff": "0.01"},
        "平均价": {"rows": 1, "codes": 1, "match": 1, "mismatch": 0,
                 "skipped_price": 0, "no_pool_codes": 0, "max_diff": None,
                 "median_diff": None},
        "不计算": {"rows": 1, "codes": 1, "match": 0, "mismatch": 0,
                  "skipped_price": 0, "no_pool_codes": 0, "max_diff": None,
                  "median_diff": None},
        "未知": {"rows": 0, "codes": 0, "match": 0, "mismatch": 0,
                "skipped_price": 0, "no_pool_codes": 0, "max_diff": None,
                "median_diff": None},
    }, "mixed_method_codes": 0, "note": "n"}
    P._write_recalc_xlsx({"workdir": str(tmp_path)}, summary)
    got = P.read_recalc_summary(str(tmp_path / "reports" / "price_recalc_check.xlsx"))
    assert got is not None
    assert got["classes"]["幂次"]["match"] == 99
    assert got["classes"]["幂次"]["mismatch"] == 1


def test_stage_stats_sheets_structure():
    st = {"funnel": {"step1": {"rows": 10, "codes": 5, "malls": 2, "unique_skus": 9,
                               "rows_in_prev": None, "filtered": None, "rows_out": 10,
                               "yes": 8, "no": 2,
                               "filtered_breakdown": {"R-a_dropped": 1}}},
          "conflicts": {"WINNER_RULE_CONFLICT": 3, "final_row_rule_conflicts_by_rule":
                        {"R-a": 1, "R-b": 0}},
          "comments": {"step4": {"POSITIVE": 1, "ZERO": 2, "EMPTY_SRC": 3, "NOT_FOUND": 4,
                                 "NO_KEY": 0, "comment_filled": 3, "coverage": 0.3},
                       "step5": {}},
          "asserts": [{"name": "x", "actual": [0, 0], "expected": [0, 0],
                       "pass": True, "stage": 5}]}
    sheets = P._stage_stats_sheets(st)
    names = [s["name"] for s in sheets]
    assert names == ["漏斗", "剔除明细", "冲突", "评论四态", "断言"]
    fun = sheets[0]
    rows = list(fun["rows"])
    assert rows[0][0] == "step1" and rows[0][1] == 10
    assert list(sheets[1]["rows"]) == [("step1", "R-a_dropped", 1)]
    assert ("WINNER_RULE_CONFLICT", 3, "") in list(sheets[2]["rows"])
    assert ("final_row_rule_conflicts_by_rule.R-a", 1, "") in list(sheets[2]["rows"])
    assert sheets[4]["rows"][0][4] == "PASS"


def test_inventory_sheets_category_labels(tmp_path):
    cfg = {"workdir": str(tmp_path),
           "inventory_round1_csv": str(tmp_path / "inv.csv")}
    csvp = tmp_path / "inv.csv"
    csvp.write_text("path,size_bytes,mtime\nx,1,2026-09-11\n", encoding="utf-8")
    p1_stats = {"inputs": {
        "final_authority": {"path": "fa", "sha256": "s1", "size": 1, "mtime": "m1"},
        "algo_src": {"path": "algo", "sha256": "s2", "size": 2, "mtime": "m2"},
    }}
    sheets = P._inventory_sheets(cfg, p1_stats)
    assert sheets[0]["name"] == "round1_inventory"
    cls = sheets[1]
    rows = list(cls["rows"])
    assert len(rows) == 3  # 2 inputs + CSV 本体
    by_key = {r[0]: r for r in rows}
    assert by_key["final_authority"][1] == "FINAL_AUTHORITY"
    assert "ALGO_SOURCE" in by_key["algo_src"][1]
    assert by_key["inventory_round1_csv"][1] == "SOURCE_INVENTORY"
