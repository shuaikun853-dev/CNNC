"""Offline evidence and request-count benchmark for incremental audit history.

This tool never calls CNNC, SaaS, a model, Zhaoshang, or a database. It reads
already archived SaaS JSON snapshots and, optionally, a legacy final workbook.
The numbers are request-count evidence, not runtime predictions.
"""

from __future__ import annotations

import argparse
from dataclasses import asdict, dataclass
from datetime import date, datetime, time as day_time
import json
import math
from pathlib import Path
import re
import sys
from typing import Any, Iterable, Mapping, Sequence
from zoneinfo import ZoneInfo


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.services.saas_audit_core import (  # noqa: E402
    model_fingerprint,
    normalize_cell_str,
    read_saas_pending_json,
)
from app.services.title_match_model_adapter import (  # noqa: E402
    TITLE_MATCH_PROMPT_VERSION,
    TITLE_MATCH_SCHEMA_VERSION,
)


SHANGHAI = ZoneInfo("Asia/Shanghai")
DEFAULT_MODEL = "title-match"
DEFAULT_PROMPT_VERSION = TITLE_MATCH_PROMPT_VERSION
DEFAULT_SCHEMA_VERSION = TITLE_MATCH_SCHEMA_VERSION


@dataclass(frozen=True)
class HistoryObservation:
    business_date: str
    observed_at: str
    source_path: Path
    headers: tuple[str, ...]
    rows: tuple[dict[str, Any], ...]


@dataclass(frozen=True)
class ObservationMetric:
    observed_at: str
    source_path: str
    rows: int
    new_detail_identity_proxies: int
    new_model_fingerprints: int
    new_crawl_identity_proxies: int
    model_calls: int


def _validated_business_date(value: str) -> str:
    return date.fromisoformat(value).isoformat()


def _snapshot_path(directory: Path) -> Path | None:
    for name in ("saas_json.json", "saas_current.json"):
        candidate = directory / name
        if candidate.is_file():
            return candidate
    return None


def _recorded_observed_at(directory: Path, fallback: datetime) -> datetime:
    summary_path = directory / "fetch_summary.json"
    if not summary_path.is_file():
        return fallback
    try:
        payload = json.loads(summary_path.read_text(encoding="utf-8"))
        value = datetime.fromisoformat(str(payload.get("completed_at", "")))
    except (OSError, TypeError, ValueError, json.JSONDecodeError):
        return fallback
    return value if value.tzinfo is not None and value.utcoffset() is not None else fallback


def discover_history_observations(
    runtime_root: Path | str,
    business_date: str,
) -> list[HistoryObservation]:
    """Load timestamped, top-level legacy observations for one local day."""
    business_date = _validated_business_date(business_date)
    compact = business_date.replace("-", "")
    pattern = re.compile(rf"^price_audit_{compact}_(\d{{6}})(?:_|$)")
    observations: list[HistoryObservation] = []
    for directory in sorted(Path(runtime_root).glob(f"price_audit_{compact}_*")):
        match = pattern.match(directory.name)
        if match is None:
            continue
        hhmmss = match.group(1)
        fallback_observed = datetime.strptime(
            f"{business_date} {hhmmss}", "%Y-%m-%d %H%M%S"
        ).replace(tzinfo=SHANGHAI)
        snapshot_directories = [
            *sorted(directory.glob("superseded_*")),
            directory,
        ]
        for snapshot_directory in snapshot_directories:
            source_path = _snapshot_path(snapshot_directory)
            if source_path is None:
                continue
            observed = _recorded_observed_at(snapshot_directory, fallback_observed)
            headers, rows = read_saas_pending_json(source_path)
            observations.append(
                HistoryObservation(
                    business_date=business_date,
                    observed_at=observed.isoformat(timespec="seconds"),
                    source_path=source_path.resolve(),
                    headers=tuple(headers),
                    rows=tuple(rows),
                )
            )
    return sorted(observations, key=lambda item: (item.observed_at, str(item.source_path)))


def row_identity(row: Mapping[str, Any]) -> tuple[str, str, str]:
    return (
        normalize_cell_str(row.get("vCommodityId")),
        normalize_cell_str(row.get("skuId")),
        normalize_cell_str(row.get("extLink")),
    )


def _detail_identity_proxy(row: Mapping[str, Any]) -> str:
    # Historical SaaS-only snapshots do not retain the corresponding Detail
    # response or every customer fingerprint input. documentId is therefore a
    # lower-bound proxy, never reported as an exact Detail request count.
    return normalize_cell_str(row.get("vCommodityId"))


def _model_key(row: Mapping[str, Any]) -> str:
    return model_fingerprint(
        dict(row), DEFAULT_MODEL, DEFAULT_PROMPT_VERSION, DEFAULT_SCHEMA_VERSION
    )


def _crawl_identity_proxy(row: Mapping[str, Any]) -> tuple[str, str, str] | None:
    identity = row_identity(row)
    return identity if identity[2] else None


def analyze_observations(
    observations: Sequence[HistoryObservation],
    *,
    model_batch_size: int = 10,
    capacity_per_tick: int | None = None,
) -> dict[str, Any]:
    """Compare repeated full refresh with durable cross-tick reuse.

    ``capacity_per_tick`` is an optional work-unit scenario. It deliberately
    does not pretend that one work unit has a known wall-clock duration.
    """
    if not observations:
        raise ValueError("at least one historical observation is required")
    if model_batch_size <= 0:
        raise ValueError("model_batch_size must be positive")
    if capacity_per_tick is not None and capacity_per_tick <= 0:
        raise ValueError("capacity_per_tick must be positive")

    seen_detail: set[str] = set()
    seen_model: set[str] = set()
    seen_crawl: set[tuple[str, str, str]] = set()
    observation_metrics: list[ObservationMetric] = []
    replayed_detail_proxies = 0
    replayed_model_fingerprints = 0
    replayed_crawl_proxies = 0
    scenario_backlog = 0

    for observation in observations:
        current_detail = {
            value
            for row in observation.rows
            if (value := _detail_identity_proxy(row))
        }
        current_model = {_model_key(row) for row in observation.rows}
        current_crawl = {
            value
            for row in observation.rows
            if (value := _crawl_identity_proxy(row)) is not None
        }
        new_detail = current_detail - seen_detail
        new_model = current_model - seen_model
        new_crawl = current_crawl - seen_crawl
        model_calls = math.ceil(len(new_model) / model_batch_size) if new_model else 0
        observation_metrics.append(
            ObservationMetric(
                observed_at=observation.observed_at,
                source_path=str(observation.source_path),
                rows=len(observation.rows),
                new_detail_identity_proxies=len(new_detail),
                new_model_fingerprints=len(new_model),
                new_crawl_identity_proxies=len(new_crawl),
                model_calls=model_calls,
            )
        )
        replayed_detail_proxies += len(current_detail)
        replayed_model_fingerprints += len(current_model)
        replayed_crawl_proxies += len(current_crawl)
        seen_detail.update(current_detail)
        seen_model.update(current_model)
        seen_crawl.update(current_crawl)
        if capacity_per_tick is not None:
            scenario_backlog = max(
                0,
                scenario_backlog + len(new_detail) + len(new_model) + len(new_crawl)
                - capacity_per_tick,
            )

    final = observations[-1]
    final_model = {_model_key(row) for row in final.rows}
    final_identities = [row_identity(row) for row in final.rows]
    return {
        "business_date": final.business_date,
        "observation_count": len(observations),
        "observation_window": {
            "first": observations[0].observed_at,
            "last": final.observed_at,
            "inside_09_18_local": all(
                day_time(9, 0)
                <= datetime.fromisoformat(item.observed_at).astimezone(SHANGHAI).time()
                <= day_time(18, 0)
                for item in observations
            ),
        },
        "final_snapshot_rows": len(final.rows),
        "final_snapshot_ordered_identity_count": len(final_identities),
        "final_snapshot_ordered_identity_unique": len(set(final_identities)),
        "legacy_repeated_refresh": {
            "detail_identity_proxy_requests": replayed_detail_proxies,
            "model_fingerprint_inputs": replayed_model_fingerprints,
            "crawl_identity_proxy_requests": replayed_crawl_proxies,
        },
        "durable_incremental": {
            "detail_identity_proxy_requests": len(seen_detail),
            "model_fingerprint_inputs": len(seen_model),
            "model_calls": sum(item.model_calls for item in observation_metrics),
            "crawl_identity_proxy_requests": len(seen_crawl),
            "peak_new_stage_inputs": max(
                item.new_detail_identity_proxies
                + item.new_model_fingerprints
                + item.new_crawl_identity_proxies
                for item in observation_metrics
            ),
        },
        "legacy_final_only": {
            "rows": len(final.rows),
            "model_fingerprint_inputs": len(final_model),
            "model_calls": math.ceil(len(final_model) / model_batch_size)
            if final_model else 0,
        },
        "scenario": {
            "capacity_per_tick_work_units": capacity_per_tick,
            "final_pending_work_units": scenario_backlog
            if capacity_per_tick is not None else None,
            "warning": (
                "Capacity is an explicit work-unit scenario, not a measured runtime or SLA."
            ),
        },
        "observations": [asdict(item) for item in observation_metrics],
        "evidence_boundaries": [
            "Snapshot rows and their order are historical evidence.",
            "Model fingerprints are exact for the configured model/prompt/schema labels.",
            "Detail counts are documentId lower-bound proxies because archived Detail/customer inputs are incomplete.",
            "Crawl counts omit mall-policy and remote export behavior because those are external boundaries.",
            "Model/page outputs and elapsed time are not reconstructed or estimated.",
        ],
    }


def _legacy_workbook_identities(path: Path) -> list[tuple[str, str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["record_results"]
        iterator = sheet.iter_rows(values_only=True)
        headers = [normalize_cell_str(value) for value in next(iterator)]
        aliases = {
            "vCommodityId": {
                "vCommodityId", "vCommodityId(物料ID)", "SaaS商品ID", "文档ID"
            },
            "skuId": {"skuId", "skuId(SKU ID)", "SKU编号"},
            "extLink": {"extLink", "extLink(商品链接)", "互联网链接"},
        }
        indexes = []
        for field in ("vCommodityId", "skuId", "extLink"):
            indexes.append(next(index for index, value in enumerate(headers) if value in aliases[field]))
        return [
            tuple(normalize_cell_str(row[index]) for index in indexes)
            for row in iterator
        ]
    finally:
        workbook.close()


def compare_final_snapshot_to_legacy(
    observation: HistoryObservation,
    legacy_workbook: Path | str,
) -> dict[str, Any]:
    snapshot = [row_identity(row) for row in observation.rows]
    workbook = _legacy_workbook_identities(Path(legacy_workbook))
    return {
        "snapshot_rows": len(snapshot),
        "workbook_rows": len(workbook),
        "ordered_identity_equal": snapshot == workbook,
        "snapshot_duplicate_rows": len(snapshot) - len(set(snapshot)),
        "workbook_duplicate_rows": len(workbook) - len(set(workbook)),
    }


def _jsonable(value: Any) -> Any:
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, Mapping):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, Iterable) and not isinstance(value, (str, bytes)):
        return [_jsonable(item) for item in value]
    return value


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--runtime-root", type=Path, default=PROJECT_ROOT / "runtime")
    parser.add_argument("--business-date", required=True)
    parser.add_argument("--model-batch-size", type=int, default=10)
    parser.add_argument("--capacity-per-tick", type=int)
    parser.add_argument("--legacy-final", type=Path)
    args = parser.parse_args(argv)

    observations = discover_history_observations(args.runtime_root, args.business_date)
    report = analyze_observations(
        observations,
        model_batch_size=args.model_batch_size,
        capacity_per_tick=args.capacity_per_tick,
    )
    if args.legacy_final:
        report["legacy_final_equivalence"] = compare_final_snapshot_to_legacy(
            observations[-1], args.legacy_final
        )
    print(json.dumps(_jsonable(report), ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
