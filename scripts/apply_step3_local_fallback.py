from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from scripts.comment_count import normalize_comment_count
except ModuleNotFoundError:
    from comment_count import normalize_comment_count


BAD_TITLE_TERMS = ("云采APP", "超低价格", "登录", "验证码", "采购平台首页")


def _headers(sheet):
    return {
        str(cell.value or "").strip(): cell.column
        for cell in sheet[1]
        if str(cell.value or "").strip()
    }


def _positive_price(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(str(value).strip())
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) and number > 0 else None


def _usable_title(value):
    text = str(value or "").strip()
    return text if text and not any(term in text for term in BAD_TITLE_TERMS) else ""


def read_crawler_results(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = _headers(sheet)
        required = (
            "原始链接",
            "商品编码",
            "商品名称",
            "商品价格",
            "商品评价数",
            "采集状态",
            "失败原因",
        )
        missing = [name for name in required if name not in headers]
        if missing:
            raise ValueError("多商城脚本结果缺少列: {}".format(", ".join(missing)))
        records = {}
        for row in range(2, sheet.max_row + 1):
            url = str(sheet.cell(row, headers["原始链接"]).value or "").strip()
            if not url:
                continue
            status = str(sheet.cell(row, headers["采集状态"]).value or "").strip()
            reason = str(sheet.cell(row, headers["失败原因"]).value or "").strip()
            price = _positive_price(sheet.cell(row, headers["商品价格"]).value)
            records[url] = {
                "status": status,
                "reason": reason,
                "price": price,
                "title": _usable_title(sheet.cell(row, headers["商品名称"]).value),
                "sku": str(sheet.cell(row, headers["商品编码"]).value or "").strip(),
                "comment": normalize_comment_count(
                    sheet.cell(row, headers["商品评价数"]).value
                ),
            }
        return records
    finally:
        workbook.close()


def _write_rows(path, sheet_name, headers, rows, text_columns=()):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for row in rows:
        sheet.append(row)
    for column in text_columns:
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row, column).number_format = "@"
    sheet.freeze_panes = "A2"
    for column in range(1, len(headers) + 1):
        values = [
            str(sheet.cell(row, column).value or "")
            for row in range(1, min(sheet.max_row, 100) + 1)
        ]
        sheet.column_dimensions[get_column_letter(column)].width = min(
            max(max(map(len, values), default=8) + 2, 12), 64
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def apply_local_fallback(
    *, input_path, crawler_result_path, output_path, log_path, remaining_path
):
    input_path = Path(input_path)
    crawler_result_path = Path(crawler_result_path)
    output_path = Path(output_path)
    log_path = Path(log_path)
    remaining_path = Path(remaining_path)
    if output_path.resolve() == input_path.resolve():
        raise ValueError("本地脚本兜底必须另存新文件，不能覆盖输入模板")

    records = read_crawler_results(crawler_result_path)
    successful = {
        url: record
        for url, record in records.items()
        if record["status"] == "成功" and record["price"] is not None
    }
    workbook = openpyxl.load_workbook(input_path)
    sheet = workbook.active
    headers = _headers(sheet)
    required = ("单据id", "互联网标题", "URL", "SKUID", "价格", "评论数")
    missing = [name for name in required if name not in headers]
    if missing:
        workbook.close()
        raise ValueError("导入模板缺少列: {}".format(", ".join(missing)))

    filled_rows = 0
    log_rows = []
    remaining_rows = []
    input_rows = sheet.max_row - 1
    for row in range(2, sheet.max_row + 1):
        document_id_cell = sheet.cell(row, headers["单据id"])
        document_id = str(document_id_cell.value or "").strip()
        document_id_cell.value = document_id
        document_id_cell.number_format = "@"
        sku_cell = sheet.cell(row, headers["SKUID"])
        sku_cell.value = str(sku_cell.value or "").strip()
        sku_cell.number_format = "@"
        comment_cell = sheet.cell(row, headers["评论数"])
        comment_cell.value = normalize_comment_count(comment_cell.value)
        url = str(sheet.cell(row, headers["URL"]).value or "").strip()
        current_price = _positive_price(sheet.cell(row, headers["价格"]).value)
        record = records.get(url)
        if (
            current_price is None
            and record
            and record["status"] == "成功"
            and record["price"] is not None
        ):
            sheet.cell(row, headers["价格"]).value = record["price"]
            if record["title"]:
                sheet.cell(row, headers["互联网标题"]).value = record["title"]
            if record["sku"]:
                sku_cell.value = record["sku"]
            comment_cell.value = record["comment"]
            filled_rows += 1
            current_price = record["price"]
            log_rows.append(
                [
                    document_id,
                    url,
                    "已回填",
                    record["price"],
                    record["title"],
                    "多商城本地脚本",
                    "",
                ]
            )
        if current_price is None:
            reason = "脚本未返回结果"
            if record:
                reason = record["reason"] or (
                    "脚本未返回有效正价格"
                    if record["status"] == "成功"
                    else record["status"]
                )
            remaining_rows.append([document_id, url, reason])
            log_rows.append(
                [document_id, url, "留空", "", "", "多商城本地脚本", reason]
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    _write_rows(
        log_path,
        "本地脚本兜底日志",
        ["单据id", "URL", "结果", "价格", "互联网标题", "来源", "原因"],
        log_rows,
        text_columns=(1,),
    )
    _write_rows(
        remaining_path,
        "脚本后仍无价格",
        ["单据id", "URL", "原因"],
        remaining_rows,
        text_columns=(1,),
    )
    return {
        "input_rows": input_rows,
        "crawler_success_records": len(successful),
        "filled_rows": filled_rows,
        "remaining_rows": len(remaining_rows),
    }


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="将第三步多商城本地脚本结果精确回填到新导入模板"
    )
    parser.add_argument("--input-xlsx", required=True)
    parser.add_argument("--crawler-result-xlsx", required=True)
    parser.add_argument("--output-xlsx", required=True)
    parser.add_argument("--log-xlsx", required=True)
    parser.add_argument("--remaining-xlsx", required=True)
    args = parser.parse_args(argv)
    summary = apply_local_fallback(
        input_path=args.input_xlsx,
        crawler_result_path=args.crawler_result_xlsx,
        output_path=args.output_xlsx,
        log_path=args.log_xlsx,
        remaining_path=args.remaining_xlsx,
    )
    print(json.dumps(summary, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
