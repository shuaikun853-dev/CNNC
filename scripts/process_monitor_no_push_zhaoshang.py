import json
import os
import pickle
import re
import shlex
import ssl
import subprocess
import sys
import time
import urllib.request
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse, urlunparse

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from scripts.auth_manager import AuthManager
    from scripts.comment_count import normalize_comment_count
    from scripts.curl_utils import apply_curl_headers, internal_direct_opener, internal_subprocess_env, parse_curl_command
    from scripts.monitor_stb_fallback import apply_detail_price_results, apply_detail_reference_results, apply_stbchina_results, classify_priority_output, extract_detail_reference_records, find_header, is_stbchina_url, load_zhaoshang_db_env, run_stbchina_fallback, save_or_reuse_locked_workbook
    from scripts.zhaoshang_task_phases import CreatePhaseError, MallJob, build_zhaoshang_task_name, run_create_phase, run_wait_export_phase
except ModuleNotFoundError:
    from auth_manager import AuthManager
    from comment_count import normalize_comment_count
    from curl_utils import apply_curl_headers, internal_direct_opener, internal_subprocess_env, parse_curl_command
    from monitor_stb_fallback import apply_detail_price_results, apply_detail_reference_results, apply_stbchina_results, classify_priority_output, extract_detail_reference_records, find_header, is_stbchina_url, load_zhaoshang_db_env, run_stbchina_fallback, save_or_reuse_locked_workbook
    from zhaoshang_task_phases import CreatePhaseError, MallJob, build_zhaoshang_task_name, run_create_phase, run_wait_export_phase

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, r"E:\.A中核运维")

from app.services.shared_crawl_cache import SharedCrawlCache
from app.services.step3_prewarm_cache import apply_cached_customer_results, missing_required_comment
from cnnc_same_product_import import config
from cnnc_same_product_import.link_parser import extract_jd_sku
from cnnc_same_product_import.mall_mapper import MallMapper
from cnnc_same_product_import.template_writer import ImportRow, write_log, write_template
import cnnc_same_product_import.template_writer as tw


import argparse


def parse_args():
    parser = argparse.ArgumentParser(
        description="处理 SaaS 监控表中不能直接推送的数据：中核 detail 取链接、招商局优先爬价、生成三文件和待确认脚本兜底清单。"
    )
    parser.add_argument("--monitor-xlsx", required=True, help="runtime/saas状态查询/<批次>/中核待处理_SaaS监控状态_<批次>.xlsx")
    parser.add_argument("--auth-curl-file", help="兼容模式：包含中核 query/detail 和招商局 exportTask cURL 的文本文件")
    parser.add_argument("--cnnc-query-curl-file", help="中核 query cURL 文件；与 --detail-curl-file、--zhaoshang-export-curl-file 一起使用")
    parser.add_argument("--detail-curl-file", help="中核 detail cURL 文件；与 --cnnc-query-curl-file、--zhaoshang-export-curl-file 一起使用")
    parser.add_argument("--zhaoshang-export-curl-file", help="招商局 exportTask cURL 文件；与 --cnnc-query-curl-file、--detail-curl-file 一起使用")
    parser.add_argument("--timestamp", default=None, help="输出批次名；默认当前时间戳")
    parser.add_argument("--output-base", default=r"E:\.A中核运维\cnnc-ops-agent\runtime\模板以及推送失败数据")
    parser.add_argument("--cnnc-query-limit", type=int, default=100)
    parser.add_argument("--zs-wait-timeout", type=int, default=600)
    parser.add_argument("--zs-wait-concurrency", type=int, default=6, help="招商局任务全部顺序创建后，并发等待和导出的最大任务数")
    parser.add_argument("--detail-retries", type=int, default=5, help="中核 detail 网络异常重试次数")
    parser.add_argument("--zhaoshang-db-env", help="招商局优先级数据库环境变量文件，仅在当前进程加载")
    parser.add_argument("--auto-stbchina-fallback", action="store_true", help="招商局后自动运行史泰博专用脚本回填仍缺价的 stbchina.cn 链接")
    parser.add_argument("--prefer-detail-price", action="store_true", help="中核 detail 有有效 URL 参考价时直接回填价格，并跳过该链接的招商局任务")
    parser.add_argument(
        "--shared-cache-path",
        default=str(Path(__file__).resolve().parents[1] / "runtime" / "incremental_price_audit" / "shared_crawl_cache.sqlite3"),
        help="客户链接预爬共享缓存；命中完整结果时第三步直接复用，京东/苏宁额外要求评论数",
    )
    parser.add_argument("--leave-internet-title-blank", action="store_true", help="只回填价格，所有互联网标题保留为空")
    parser.add_argument("--detail-reference-fallback", action="store_true", help="经用户明确授权后，以同单据同URL的中核detail参考价和商品标题补全仍同时缺少价格、标题的行，并另存模板")
    parser.add_argument("--skip-zhaoshang", action="store_true", help="不创建招商局任务；仅生成当前可处理的三类文件和缺价清单")
    parser.add_argument("--zs-create-only", action="store_true", help="只批量创建、启用招商局任务并尝试提优，不等待或导出")
    args = parser.parse_args()
    if args.zs_wait_concurrency < 1:
        parser.error("--zs-wait-concurrency 必须大于等于 1")
    return args


ARGS = parse_args()
if ARGS.zhaoshang_db_env:
    os.environ.update(load_zhaoshang_db_env(ARGS.zhaoshang_db_env))
MONITOR = ARGS.monitor_xlsx
ATTACHMENT = ARGS.auth_curl_file
TIMESTAMP = ARGS.timestamp or datetime.now().strftime("%Y%m%d%H%M%S")
OUT_BASE = Path(ARGS.output_base)
OUT_DIR = (OUT_BASE / TIMESTAMP).resolve()
OUT_DIR.mkdir(parents=True, exist_ok=True)
BATCH_STATE_PATH = OUT_DIR / "BATCH_STATE.md"
if not BATCH_STATE_PATH.is_file():
    raise RuntimeError("第三步生产运行前必须先创建批次状态文件: {}".format(BATCH_STATE_PATH))

PYTHON = sys.executable
ZS_SCRIPT = r"E:\.A2026\招商局自动建任务\external_tools\crawl_task_tool\auto_crawl_task.py"
ZS_CWD = r"E:\.A2026\招商局自动建任务\external_tools\crawl_task_tool"
ZS_WAIT_TIMEOUT = ARGS.zs_wait_timeout
URL_RE = re.compile(r"https?://[^\s,，;；\n\r\t]+")

MALL_CN_NAME = {
    "jd": "京东", "zkh": "震坤行商城", "zjmi": "浙江物产商城", "tima": "浙江宏伟商城",
    "xfs": "鑫方盛商城", "ycw": "阳采集团", "nbdeli": "得力商城", "mymro": "固安捷",
    "lxwl": "领先未来商城", "colipu": "科力普商城", "comix": "齐心商城",
    "suning": "苏宁商城",
    "qstbg": "企事通商城", "ofs": "欧菲斯办公用品", "epec": "易派客",
    "ehsy": "西域", "haoxiny": "云中鹤商城", "guangbo": "广博商城",
    "xhgj": "咸亨国际商城", "lxwlmall": "领先未来新商城", "jslink": "得力集实采",
}
BAD_TITLE_TERMS = ["云采APP", "超低价格", "登录", "验证码", "采购平台首页"]


def collect_curl_blocks(text):
    try:
        from scripts.curl_utils import collect_curl_blocks as _collect
    except ModuleNotFoundError:
        from curl_utils import collect_curl_blocks as _collect
    return _collect(text)


def parse_curl(block):
    parsed = parse_curl_command(block)
    return {
        "url": parsed["url"],
        "headers": dict(parsed["headers_dict"]),
        "body": parsed.get("body") or parsed.get("data"),
    }


def load_curls(path=None, *, query_path=None, detail_path=None, zhaoshang_export_path=None):
    if query_path or detail_path or zhaoshang_export_path:
        if not (query_path and detail_path and zhaoshang_export_path):
            raise RuntimeError("分别传入 cURL 时，必须同时提供 --cnnc-query-curl-file、--detail-curl-file 和 --zhaoshang-export-curl-file")
        query = parse_curl(Path(query_path).read_text(encoding="utf-8"))
        detail = parse_curl(Path(detail_path).read_text(encoding="utf-8"))
        z = parse_curl(Path(zhaoshang_export_path).read_text(encoding="utf-8"))
        cookie = z["headers"].get("Cookie", "")
        blade = z["headers"].get("Blade-Auth", "")
        if not blade or not cookie:
            raise RuntimeError("招商局 exportTask cURL 中缺少 Blade-Auth 或 Cookie")
        return query, detail, blade, cookie
    if not path:
        raise RuntimeError("必须提供 --auth-curl-file，或同时提供三份独立 cURL 文件")
    text = Path(path).read_text(encoding="utf-8")
    query = detail = None
    blade = cookie = ""
    for block in collect_curl_blocks(text):
        if "quasiStandardCommodityDocumentModel/query" in block:
            query = parse_curl(block)
        elif "quasiStandardCommodityDocumentModel/detail" in block:
            detail = parse_curl(block)
        elif "crawlTaskManage/exportTask" in block:
            z = parse_curl(block)
            cookie = z["headers"].get("Cookie", "")
            blade = z["headers"].get("Blade-Auth", "")
    if not query or not detail or not blade or not cookie:
        raise RuntimeError("附件中缺少中核 query/detail 或招商局认证 cURL")
    return query, detail, blade, cookie


query_curl, detail_curl, blade_auth, zs_cookie = load_curls(
    ATTACHMENT,
    query_path=ARGS.cnnc_query_curl_file,
    detail_path=ARGS.detail_curl_file,
    zhaoshang_export_path=ARGS.zhaoshang_export_curl_file,
)
ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE
DIRECT_OPENER = urllib.request.build_opener(
    urllib.request.ProxyHandler({}),
    urllib.request.HTTPSHandler(context=ctx),
)


def cnnc_post(kind, body):
    cfg = query_curl if kind == "query" else detail_curl
    req = urllib.request.Request(
        cfg["url"],
        data=json.dumps(body, ensure_ascii=False).encode("utf-8"),
        method="POST",
    )
    apply_curl_headers(req, cfg["headers"])
    with DIRECT_OPENER.open(req, timeout=120) as resp:
        return json.loads(resp.read().decode("utf-8", errors="replace"))


def extract_sku_from_url(url):
    if not url:
        return ""
    parsed = urlparse(url)
    qs = parse_qs(parsed.query)
    for key in ["sku", "skuId", "skuNo", "itemId", "id", "goods_id", "productId"]:
        if key in qs and qs[key] and str(qs[key][0]).strip():
            return str(qs[key][0]).strip()
    if "jd.com" in parsed.netloc:
        jd = extract_jd_sku(url)
        if jd:
            return jd
    for pat in [
        r"/itemcode-?([A-Za-z0-9_-]+)\.html$",
        r"/(\d{4,})\.html$",
        r"/product-(\d{4,})\.html$",
        r"/item-(\d{4,})\.html$",
        r"/goods-(\d{4,})\.html$",
        r"/([A-Za-z0-9][A-Za-z0-9_-]{3,})/?$",
    ]:
        m = re.search(pat, parsed.path or "", re.I)
        if m:
            return m.group(1)
    return ""


def is_bad_title(title):
    s = str(title or "").strip()
    return (not s) or s.upper() in ("APP", "NA", "NONE", "NULL") or any(t in s for t in BAD_TITLE_TERMS)


def read_zs_result(path, code_to_url):
    records = {}
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    def col(name):
        return headers.index(name) + 1 if name in headers else -1

    ci = {k: col(v) for k, v in {
        "code": config.ZS_COL_CODE, "title": config.ZS_COL_TITLE,
        "price": config.ZS_COL_PRICE, "url": config.ZS_COL_URL,
        "comment": config.ZS_COL_COMMENT, "channel": "商城标识",
    }.items()}
    ci["page_sku"] = col("商品最小品类单元id")
    if ci["page_sku"] < 0:
        ci["page_sku"] = col("商品最小品类单元Id")
    for r in range(2, ws.max_row + 1):
        def gv(k, *, preserve_zero=False):
            if ci[k] <= 0:
                return ""
            value = ws.cell(r, ci[k]).value
            if value is None:
                return ""
            if not preserve_zero and not value:
                return ""
            return str(value).strip()
        raw_code = gv("code")
        code = raw_code[5:] if raw_code.startswith("code-") else raw_code
        url = code_to_url.get(raw_code) or code_to_url.get(code) or gv("url")
        if not url:
            continue
        try:
            price = float(gv("price")) if gv("price") else ""
        except Exception:
            price = ""
        title = gv("title")
        page_sku = gv("page_sku")
        records[url] = {
            "code": page_sku or code,
            "title": "" if is_bad_title(title) else title,
            "price": price,
            "comment": gv("comment", preserve_zero=True),
            "channel": gv("channel"),
        }
    return records


def is_reusable_zs_result(job):
    if not job.result_path.is_file():
        return False
    workbook = None
    input_workbook = None
    try:
        input_workbook = openpyxl.load_workbook(
            str(job.input_path), read_only=True, data_only=True
        )
        expected_rows = sum(
            1
            for row in input_workbook.active.iter_rows(min_row=2, values_only=True)
            if any(value not in (None, "") for value in row)
        )
        workbook = openpyxl.load_workbook(
            str(job.result_path), read_only=True, data_only=True
        )
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        required = {config.ZS_COL_CODE, config.ZS_COL_PRICE, config.ZS_COL_URL}
        if not required.issubset(set(headers)):
            return False
        code_index = headers.index(config.ZS_COL_CODE)
        url_index = headers.index(config.ZS_COL_URL)
        data_rows = sum(
            1
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[code_index] not in (None, "") or row[url_index] not in (None, "")
        )
        return data_rows == expected_rows
    except Exception as exc:
        print("[招商局] 已有结果无效，将按任务状态恢复导出: {} ({})".format(
            job.result_path, exc
        ))
        return False
    finally:
        if workbook is not None:
            workbook.close()
        if input_workbook is not None:
            input_workbook.close()


print("[1] 读监控表否的行...")
wb_m = openpyxl.load_workbook(MONITOR, data_only=True)
ws_m = wb_m.active
h_m = [ws_m.cell(1, c).value for c in range(1, ws_m.max_column + 1)]
doc_header = find_header(h_m, "文档ID")
name_header = find_header(h_m, "商品名称", "中核商品名称")
code_header = find_header(h_m, "准标准商品编码", "物料编码")
no_rows = []
for r in range(2, ws_m.max_row + 1):
    if str(ws_m.cell(r, h_m.index("是否可推送") + 1).value or "") == "否":
        no_rows.append({
            "doc": str(ws_m.cell(r, h_m.index(doc_header) + 1).value or "").strip(),
            "name": str(ws_m.cell(r, h_m.index(name_header) + 1).value or ""),
            "code": str(ws_m.cell(r, h_m.index(code_header) + 1).value or ""),
            "four_cat": str(ws_m.cell(r, h_m.index("四级类目ID") + 1).value or "").strip() if "四级类目ID" in h_m else "",
        })
print("    否的行: {}".format(len(no_rows)))

print("[2] 调中核 query...")
base_body = json.loads(query_curl["body"])
all_cnnc, page = [], 1
limit = ARGS.cnnc_query_limit
query_failed_reason = ""
query_needed_docs = {row["doc"] for row in no_rows if not row.get("four_cat")}
try:
    while query_needed_docs:
        body = dict(base_body)
        body["page"] = page
        body["limit"] = limit
        data = cnnc_post("query", body)
        rsp_code = str(data.get("rspCode") or "")
        if rsp_code and rsp_code != "0000":
            query_failed_reason = "rspCode={} rspDesc={}".format(rsp_code, data.get("rspDesc"))
            break
        rows = data.get("rows") or data.get("data") or []
        if isinstance(rows, dict):
            rows = rows.get("rows") or rows.get("records") or []
        all_cnnc.extend(rows)
        query_needed_docs.difference_update(
            str(row.get("documentId", "")).strip() for row in rows
        )
        total = int(data.get("total") or data.get("count") or 0)
        print("    page={} rows={} total={} loaded={} remaining_docs={}".format(
            page, len(rows), total, len(all_cnnc), len(query_needed_docs)
        ))
        if not rows or len(rows) < limit or (total and len(all_cnnc) >= total):
            break
        page += 1
except Exception as exc:
    query_failed_reason = str(exc)
if not query_needed_docs:
    print("    监控表四级类目ID齐全，跳过全量中核 query。")
elif query_failed_reason:
    print("    中核 query 不可用，改用监控表四级类目ID兜底: {}".format(query_failed_reason))
doc_to_row = {str(r.get("documentId", "")).strip(): r for r in all_cnnc}

print("[3] 调 detail 查链接...")
doc_to_links, doc_to_info, detail_reference_records = {}, {}, {}
ok = fail = 0
for i, nd in enumerate(no_rows):
    doc_id = nd["doc"]
    four_cat = str((doc_to_row.get(doc_id) or {}).get("fourCategoryId", "")).strip() or nd.get("four_cat", "")
    if not four_cat:
        fail += 1
        doc_to_links[doc_id] = []
        doc_to_info[doc_id] = {"commodityName": nd["name"]}
        continue
    try:
        last_error = None
        for attempt in range(ARGS.detail_retries):
            try:
                data = cnnc_post("detail", {"documentId": doc_id, "fourCategoryId": four_cat})
                if isinstance(data, dict):
                    break
                last_error = RuntimeError("detail 未返回 JSON 对象")
            except Exception as exc:
                last_error = exc
            if attempt + 1 < ARGS.detail_retries:
                time.sleep(min(2 * (attempt + 1), 10))
        else:
            raise last_error or RuntimeError("detail 重试后仍无有效响应")
        d = data.get("data", {})
        if not isinstance(d, dict):
            d = {}
        links = []
        for item in d.get("referenceInfos", []) if isinstance(d.get("referenceInfos", []), list) else []:
            links += [u for u in URL_RE.findall(str(item.get("referenceLink", "") or "")) if u.startswith("http")]
        links += [u for u in URL_RE.findall(str(d.get("referenceLink", "") or "")) if u.startswith("http")]
        seen, unique = set(), []
        for u in links:
            if u not in seen:
                seen.add(u)
                unique.append(u)
        doc_to_links[doc_id] = unique
        doc_to_info[doc_id] = {"commodityName": str(d.get("commodityName", "") or nd["name"])}
        detail_reference_records.update(extract_detail_reference_records(doc_id, d))
        ok += 1
    except Exception:
        fail += 1
        doc_to_links[doc_id] = []
        doc_to_info[doc_id] = {"commodityName": nd["name"]}
    if (i + 1) % 20 == 0:
        print("    进度: {}/{} ok={} fail={}".format(i + 1, len(no_rows), ok, fail))

multi = sum(1 for v in doc_to_links.values() if len(v) >= 2)
single = sum(1 for v in doc_to_links.values() if len(v) == 1)
zero = sum(1 for v in doc_to_links.values() if len(v) == 0)
print("    >=2链接:{} 1链接:{} 0链接:{}".format(multi, single, zero))

print("[4] 构建导入记录...")
mapper = MallMapper()
all_import_rows, single_rows = [], []
for nd in no_rows:
    doc_id = nd["doc"]
    links = doc_to_links.get(doc_id, [])
    info = doc_to_info.get(doc_id, {})
    if len(links) >= 2:
        for url in links:
            url = url.replace("www.jslink.com", "b2b.nbdeli.com")
            domain, shop_code, shop_name = mapper.identify_by_url(url)
            if not shop_code:
                continue
            all_import_rows.append(ImportRow(
                商品名称=info.get("commodityName", ""), 标品名称=info.get("commodityName", ""),
                单据id=doc_id, URL=url, SKUID=extract_sku_from_url(url),
                供应商编码=shop_code, 供应商名称=shop_name,
                潜在商品编码=nd["code"], 域名=domain,
                商城识别结果="{}/{}".format(shop_code, shop_name), 爬取状态="待爬取",
            ))
    else:
        single_rows.append({"doc": doc_id, "code": nd["code"], "name": nd["name"], "links": links,
                            "reason": "只有{}个链接".format(len(links)) if links else "未查到链接"})
print("    导入:{}条 推送失败:{}条".format(len(all_import_rows), len(single_rows)))

prewarm_cache_summary = {
    "scanned": 0, "cache_miss": 0, "applied": 0, "incomplete": 0,
    "missing_required_comment": 0, "channel_mismatch": 0,
}
shared_cache_path = Path(ARGS.shared_cache_path).resolve()
if shared_cache_path.is_file():
    try:
        prewarm_cache_summary = apply_cached_customer_results(
            all_import_rows,
            shared_cache=SharedCrawlCache.open_read_only(shared_cache_path),
            mall_resolver=config.get_zhaoshang_mall_key,
        )
        print(
            "    客户预爬缓存复用:{}条；不完整:{}；京东/苏宁缺评论:{}；未命中:{}".format(
                prewarm_cache_summary["applied"],
                prewarm_cache_summary["incomplete"],
                prewarm_cache_summary["missing_required_comment"],
                prewarm_cache_summary["cache_miss"],
            )
        )
        if ARGS.leave_internet_title_blank:
            for cached_row in all_import_rows:
                cached_row.互联网标题 = ""
    except Exception as exc:
        prewarm_cache_summary["error"] = str(exc)
        print("    客户预爬缓存不可用，保持原招商局流程: {}".format(exc))
else:
    prewarm_cache_summary["error"] = "shared cache missing"
    print("    客户预爬缓存不存在，保持原招商局流程: {}".format(shared_cache_path))

detail_direct_backfilled = 0
if ARGS.prefer_detail_price:
    detail_direct_backfilled = apply_detail_price_results(
        all_import_rows, detail_reference_records
    )
    print("中核detail直接回填价格:{}条；互联网标题保留为空；其余链接再进入招商局".format(
        detail_direct_backfilled
    ))

print("[5] 招商局-only...")
all_urls = list(dict.fromkeys(
    r.URL for r in all_import_rows if r.URL and r.价格 in ("", None)
))
zs_urls = [u for u in all_urls if config.get_zhaoshang_mall_key(u)]
non_zs_urls = [u for u in all_urls if not config.get_zhaoshang_mall_key(u)]
print("    总URL:{} 招商局:{} 非招商局:{}".format(len(all_urls), len(zs_urls), len(non_zs_urls)))
if ARGS.skip_zhaoshang:
    print("    已按参数跳过招商局建任务和爬取，仅生成可处理文件。")
    zs_urls = []
else:
    auth_template = {
        "cookie": zs_cookie,
        "headers": [("Cookie", zs_cookie), ("Blade-Auth", blade_auth)],
        "headers_dict": {"Cookie": zs_cookie, "Blade-Auth": blade_auth},
    }
    prepared_auth = AuthManager().prepare_curl(
        "zhaoshang", auth_template, refresh=True
    )
    zs_cookie = prepared_auth["cookie"]
    blade_auth = prepared_auth["headers_dict"]["Blade-Auth"]
groups = {}
for u in zs_urls:
    groups.setdefault(config.get_zhaoshang_mall_key(u), []).append(u)
records, priority_failures, priority_unconfirmed, priority_idempotent, task_failures = {}, [], [], [], []
zs_tmp = OUT_DIR / "_zs_tmp_priority"
zs_tmp.mkdir(parents=True, exist_ok=True)
jobs = []
for mall, urls in groups.items():
    cn_name = MALL_CN_NAME.get(mall, mall)
    input_path = zs_tmp / "zs_input_{}.xlsx".format(mall)
    result_path = zs_tmp / "zs_result_{}.xlsx".format(mall)
    print("[招商局] {} {}条，任务名将在正式创建前生成".format(mall, len(urls)))
    wb_in = openpyxl.Workbook()
    ws_in = wb_in.active
    ws_in.title = "products"
    ws_in.cell(1, 1, "商品编码")
    ws_in.cell(1, 2, "url")
    code_to_url = {}
    for i, u in enumerate(urls, 2):
        code = extract_sku_from_url(u) or "ROW{}".format(i - 1)
        code_to_url[code] = u
        code_to_url["code-" + code] = u
        ws_in.cell(i, 1, code)
        ws_in.cell(i, 2, u)
    wb_in.save(input_path)
    jobs.append(MallJob(
        mall=mall,
        task_name="",
        input_path=input_path,
        result_path=result_path,
        code_to_url=code_to_url,
    ))


zs_env = internal_subprocess_env(os.environ)
zs_env["PYTHONIOENCODING"] = "utf-8"
if not ARGS.skip_zhaoshang:
    zs_env["BLADE_AUTH"] = blade_auth
    zs_env["COOKIE"] = zs_cookie


def run_zhaoshang_job(job, mode):
    cmd = [
        PYTHON, ZS_SCRIPT,
        "--task-name", job.task_name,
        "--mall", job.mall,
        "--excel", str(job.input_path),
        "--output", str(job.result_path),
        "--wait-timeout", str(ZS_WAIT_TIMEOUT),
        "--progress-interval", "30",
        "--batch-state-file", str(BATCH_STATE_PATH),
        "--" + mode,
    ]
    proc = subprocess.run(
        cmd,
        cwd=ZS_CWD,
        env=zs_env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=ZS_WAIT_TIMEOUT + 240,
    )
    if mode == "create-only" and proc.returncode == 0:
        out_text = (proc.stdout or "") + "\n" + (proc.stderr or "")
        if classify_priority_output(out_text) in ("failed", "missing"):
            return subprocess.CompletedProcess(
                args=proc.args,
                returncode=3,
                stdout=proc.stdout,
                stderr=(proc.stderr or "") + "\n优先级调整未确认成功",
            )
    return proc


def prepare_zhaoshang_create_job(job):
    cn_name = MALL_CN_NAME.get(job.mall, job.mall)
    return replace(job, task_name=build_zhaoshang_task_name(cn_name))


def report_zhaoshang_outcome(outcome):
    print("[招商局:{}] mall={} exit={}".format(
        outcome.mode, outcome.job.mall, outcome.returncode
    ))
    if outcome.stdout:
        print(outcome.stdout.rstrip())
    if outcome.stderr:
        print(outcome.stderr.rstrip())
    if outcome.exception:
        print(outcome.exception)
    if outcome.mode != "create-only":
        return
    out_text = outcome.stdout + "\n" + outcome.stderr
    priority_state = classify_priority_output(out_text)
    if priority_state in ("failed", "missing"):
        priority_failures.append({
            "mall": outcome.job.mall,
            "task_name": outcome.job.task_name,
            "reason": "优先级调整未取得 affected_rows>0 证据",
        })
    elif priority_state == "completed_unconfirmed":
        priority_unconfirmed.append({
            "mall": outcome.job.mall,
            "task_name": outcome.job.task_name,
            "reason": "任务已完成、未确认提优",
        })
    elif priority_state == "already_highest_priority":
        priority_idempotent.append({
            "mall": outcome.job.mall,
            "task_name": outcome.job.task_name,
            "reason": "已按当前任务主键范围确认最高优先级",
        })


try:
    create_phase = run_create_phase(
        jobs,
        run_zhaoshang_job,
        on_outcome=report_zhaoshang_outcome,
        is_reusable=is_reusable_zs_result,
        prepare_job=prepare_zhaoshang_create_job,
    )
except CreatePhaseError as exc:
    raise RuntimeError(
        "招商局创建阶段失败，已停止后续商城 /add；商城={}，原因={}".format(
            exc.outcome.job.mall,
            exc.outcome.exception or exc.outcome.stderr or exc.outcome.returncode,
        )
    ) from exc

for job in create_phase.reused:
    if job.result_path.exists():
        print("[招商局] 复用已导出结果: {}".format(job.result_path))
        records.update(read_zs_result(job.result_path, job.code_to_url))

if ARGS.zs_create_only:
    print("[招商局] 仅建任务模式完成；待等待/导出商城数={}".format(len(create_phase.pending)))
else:
    wait_phase = run_wait_export_phase(
        create_phase.pending,
        run_zhaoshang_job,
        max_workers=ARGS.zs_wait_concurrency,
        on_outcome=report_zhaoshang_outcome,
        is_reusable=is_reusable_zs_result,
    )
    wait_failures = {outcome.job.mall: outcome for outcome in wait_phase.failures}
    for job in create_phase.pending:
        if is_reusable_zs_result(job):
            records.update(read_zs_result(job.result_path, job.code_to_url))
            continue
        outcome = wait_failures.get(job.mall)
        reason = "未生成招商局结果文件"
        if outcome:
            reason = outcome.exception or outcome.stderr or "等待/导出退出码={}".format(outcome.returncode)
        task_failures.append({"mall": job.mall, "reason": reason})

print("[6] 回填...")
price_count = title_count = comment_count = 0
missing_rows = []
for ir in all_import_rows:
    parsed = urlparse(ir.URL)
    norm = urlunparse((parsed.scheme, parsed.netloc, parsed.path, "", "", ""))
    rec = records.get(ir.URL) or records.get(norm)
    if rec:
        if rec.get("code") and not str(rec["code"]).startswith("ROW"):
            ir.SKUID = str(rec["code"])[5:] if str(rec["code"]).startswith("code-") else rec["code"]
        if not ir.SKUID or str(ir.SKUID).startswith("-"):
            ir.SKUID = extract_sku_from_url(ir.URL)
        if rec.get("price") not in ("", None):
            ir.价格 = rec["price"]
            price_count += 1
            ir.互联网标题 = "" if ARGS.leave_internet_title_blank else (rec.get("title") or "")
            if ir.互联网标题:
                title_count += 1
        else:
            ir.价格 = ""
            ir.互联网标题 = ""
        comment = normalize_comment_count(rec.get("comment"))
        if comment is not None:
            ir.销量_评论数 = comment
            comment_count += 1
        else:
            ir.销量_评论数 = ""
    else:
        if ir.价格 in ("", None):
            if not ir.SKUID or str(ir.SKUID).startswith("-"):
                ir.SKUID = extract_sku_from_url(ir.URL)
            ir.价格 = ""
            ir.互联网标题 = ""
            ir.销量_评论数 = ""
    if ir.价格 in ("", None):
        missing_rows.append(ir)

for ir in all_import_rows:
    comment = normalize_comment_count(ir.销量_评论数)
    ir.销量_评论数 = "" if comment is None else comment

price_count = sum(1 for row in all_import_rows if row.价格 not in ("", None))
title_count = sum(1 for row in all_import_rows if str(row.互联网标题 or "").strip())
comment_count = sum(1 for row in all_import_rows if str(row.销量_评论数 or "").strip())
missing_rows = [row for row in all_import_rows if row.价格 in ("", None)]
required_comment_missing = [
    row for row in all_import_rows
    if config.get_zhaoshang_mall_key(row.URL) in {"jd", "suning"}
    and row.价格 not in ("", None)
    and not str(row.销量_评论数 or "").strip()
]
for row in required_comment_missing:
    row.爬取状态 = "评论数缺失"
    row.爬取失败原因 = "京东/苏宁要求评论数，但招商局结果未返回可用评论数"
print(
    "    有价格:{} 有标题:{} 有评论:{} 无价格:{} 京东/苏宁缺评论:{}".format(
        price_count, title_count, comment_count, len(missing_rows), len(required_comment_missing)
    )
)

print("[7] 写文件...")
old_validate = tw.validate_row
def relaxed_validate(row):
    ok, reason = old_validate(row)
    if not ok and "爬取未成功" in reason:
        return True, ""
    return ok, reason


def write_relaxed_template(output_path, rows):
    importable_rows = [
        row for row in rows
        if not missing_required_comment(
            row, mall_resolver=config.get_zhaoshang_mall_key
        )
    ]
    tw.validate_row = relaxed_validate
    try:
        try:
            return write_template(config.DEFAULT_TEMPLATE, output_path, importable_rows)
        except PermissionError:
            existing = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
            try:
                expected_rows = len({(row.单据id, row.URL) for row in importable_rows}) + 1
                if existing.active.max_row != expected_rows:
                    raise
            finally:
                existing.close()
            print("    模板文件被占用，已核对行数并复用现有文件: {}".format(output_path))
            return output_path, {"written": expected_rows - 1, "reused_locked": True}
    finally:
        tw.validate_row = old_validate


out_template = OUT_DIR / "导入zhimport模版_{}.xlsx".format(TIMESTAMP)
out_log = OUT_DIR / "处理日志_{}.xlsx".format(TIMESTAMP)
_, wstats = write_relaxed_template(out_template, all_import_rows)
write_log(out_log, all_import_rows)

fail_path = OUT_DIR / "直接推送失败的数据_{}.xlsx".format(TIMESTAMP)
wb_f = openpyxl.Workbook()
ws_f = wb_f.active
ws_f.title = "推送失败"
fh = ["文档ID(saasId)", "准标准商品编码", "商品名称", "链接数", "链接", "失败原因"]
for c, h in enumerate(fh, 1):
    cell = ws_f.cell(1, c, h)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
for i, row in enumerate(single_rows, 2):
    ws_f.cell(i, 1, row["doc"])
    ws_f.cell(i, 2, row["code"])
    ws_f.cell(i, 3, row["name"])
    ws_f.cell(i, 4, len(row["links"]))
    ws_f.cell(i, 5, " | ".join(row["links"]) if row["links"] else "")
    ws_f.cell(i, 6, row["reason"])
for c, w in enumerate([25, 20, 40, 8, 60, 20], 1):
    ws_f.column_dimensions[get_column_letter(c)].width = w
ws_f.freeze_panes = "A2"
fail_save_state = save_or_reuse_locked_workbook(wb_f, fail_path, len(single_rows) + 1)
if fail_save_state == "reused":
    print("    直接推送失败文件被占用，已核对行数并复用现有文件。")

review_path = OUT_DIR / "待确认脚本兜底清单_{}.xlsx".format(TIMESTAMP)
wb_r = openpyxl.Workbook()
ws_r = wb_r.active
ws_r.title = "待确认脚本兜底"
ws_r.append(["单据id", "商品名称", "SKUID", "供应商编码", "供应商名称", "URL", "原因"])
review_rows = list(missing_rows)
review_keys = {(row.单据id, row.URL) for row in review_rows}
for row in required_comment_missing:
    if (row.单据id, row.URL) not in review_keys:
        review_rows.append(row)
        review_keys.add((row.单据id, row.URL))
for ir in review_rows:
    mall = config.get_zhaoshang_mall_key(ir.URL)
    if missing_required_comment(ir, mall_resolver=config.get_zhaoshang_mall_key):
        reason = "京东/苏宁缺少评论数，不能作为第三步完整结果"
    else:
        reason = "招商局未返回有效价格" if mall else "非招商局URL，需确认是否跑脚本"
    ws_r.append([ir.单据id, ir.商品名称, ir.SKUID, ir.供应商编码, ir.供应商名称, ir.URL, reason])
wb_r.save(review_path)

stb_backfilled = 0
stb_error = ""
stb_template_path = ""
stb_remaining_path = ""
if ARGS.auto_stbchina_fallback:
    stb_urls = [row.URL for row in missing_rows if is_stbchina_url(row.URL)]
    if stb_urls:
        print("[8] 史泰博自动兜底: {}条...".format(len(stb_urls)))
        try:
            stb_records = run_stbchina_fallback(stb_urls, OUT_DIR)
            stb_backfilled = apply_stbchina_results(
                all_import_rows, stb_records, fill_title=not ARGS.leave_internet_title_blank
            )
        except RuntimeError as exc:
            stb_error = str(exc)
            print("    史泰博兜底已停止: {}".format(stb_error))

        missing_rows = [row for row in all_import_rows if row.价格 in ("", None)]
        stb_template = OUT_DIR / "导入zhimport模版_{}_招商局_STB补价后.xlsx".format(TIMESTAMP)
        _, _ = write_relaxed_template(stb_template, all_import_rows)
        stb_template_path = str(stb_template)

        stb_remaining = OUT_DIR / "STB补价后仍无价格清单_{}.xlsx".format(TIMESTAMP)
        wb_stb = openpyxl.Workbook()
        ws_stb = wb_stb.active
        ws_stb.title = "仍无价格"
        ws_stb.append(["单据id", "商品名称", "SKUID", "供应商编码", "供应商名称", "URL", "原因"])
        for row in missing_rows:
            reason = stb_error if is_stbchina_url(row.URL) and stb_error else (
                "史泰博未返回有效价格" if is_stbchina_url(row.URL) else "非史泰博链接，仍需用户确认"
            )
            ws_stb.append([row.单据id, row.商品名称, row.SKUID, row.供应商编码, row.供应商名称, row.URL, reason])
        wb_stb.save(stb_remaining)
        stb_remaining_path = str(stb_remaining)
        print("    史泰博回填:{} 剩余无价格:{}".format(stb_backfilled, len(missing_rows)))

detail_reference_backfilled = 0
detail_reference_template_path = ""
detail_reference_log_path = ""
detail_reference_remaining_path = ""
if ARGS.detail_reference_fallback:
    print("[9] 中核detail参考数据补全仍缺价格和标题的行...")
    detail_reference_backfilled = apply_detail_reference_results(all_import_rows, detail_reference_records)
    missing_rows = [
        row for row in all_import_rows
        if row.价格 in ("", None) and row.互联网标题 in ("", None)
    ]
    detail_template = OUT_DIR / "导入zhimport模版_{}_招商局_STB_DETAIL参考补全后.xlsx".format(TIMESTAMP)
    _, _ = write_relaxed_template(detail_template, all_import_rows)
    detail_reference_template_path = str(detail_template)

    detail_log = OUT_DIR / "处理日志_{}_招商局_STB_DETAIL参考补全后.xlsx".format(TIMESTAMP)
    write_log(detail_log, all_import_rows)
    detail_reference_log_path = str(detail_log)

    detail_remaining = OUT_DIR / "DETAIL参考补全后仍无价格标题清单_{}.xlsx".format(TIMESTAMP)
    wb_detail = openpyxl.Workbook()
    ws_detail = wb_detail.active
    ws_detail.title = "仍无价格和标题"
    ws_detail.append(["单据id", "商品名称", "SKUID", "供应商编码", "供应商名称", "URL", "原因"])
    for row in missing_rows:
        ws_detail.append([
            row.单据id, row.商品名称, row.SKUID, row.供应商编码, row.供应商名称, row.URL,
            "同单据同URL的detail参考信息缺少有效价格或商品标题",
        ])
    wb_detail.save(detail_remaining)
    detail_reference_remaining_path = str(detail_remaining)
    print("    detail参考回填:{} 仍同时缺价格和标题:{}".format(
        detail_reference_backfilled, len(missing_rows)
    ))

summary = {
    "out_dir": str(OUT_DIR), "template": str(out_template), "log": str(out_log),
    "fail": str(fail_path), "review": str(review_path), "import_rows": len(all_import_rows),
    "single_rows": len(single_rows), "zs_urls": len(zs_urls), "non_zs_urls": len(non_zs_urls),
    "detail_direct_backfilled": detail_direct_backfilled,
    "prewarm_cache": prewarm_cache_summary,
    "price_count": price_count, "missing_count": len(missing_rows),
    "required_comment_missing_count": len(required_comment_missing),
    "priority_failures": priority_failures, "priority_unconfirmed": priority_unconfirmed,
    "priority_idempotent": priority_idempotent,
    "task_failures": task_failures,
    "stb_backfilled": stb_backfilled, "stb_error": stb_error,
    "stb_template": stb_template_path, "stb_remaining": stb_remaining_path,
    "detail_reference_backfilled": detail_reference_backfilled,
    "detail_reference_template": detail_reference_template_path,
    "detail_reference_log": detail_reference_log_path,
    "detail_reference_remaining": detail_reference_remaining_path,
}
(OUT_DIR / "step2_zhaoshang_only_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
if priority_failures or task_failures:
    raise SystemExit(2)
