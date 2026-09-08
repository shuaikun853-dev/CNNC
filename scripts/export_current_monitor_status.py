from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import sys
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.request import Request

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from curl_utils import (
    apply_curl_headers,
    apply_internal_no_proxy_env,
    internal_direct_opener,
    load_curl_command,
)
try:
    from scripts.auth_manager import AuthManager
except ModuleNotFoundError:  # Direct invocation from scripts/.
    from auth_manager import AuthManager


ROOT = Path(__file__).resolve().parents[1]

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.services.incremental_audit_runtime import business_date_in_shanghai
from app.services.workday_context import (
    ensure_workday_layout,
    now_in_shanghai,
    read_manifest,
    update_stage,
    workday_root,
)

CNNC_KEYS = [
    "documentId",
    "commodityName",
    "categoryName",
    "fourCategoryId",
    "fourCategoryName",
    "referencePrice",
    "referenceLink",
    "materialCode",
    "materialDesc",
    "createTime",
    "updateTime",
    "sassCheckStatus",
    "quasiStandardCommodityCode",
]
MONITOR_HEADERS = [
    "文档ID",
    "中核商品名称",
    "中核类目名称",
    "四级类目ID",
    "四级类目名称",
    "客户参考价",
    "参考链接",
    "物料编码",
    "物料描述",
    "中核创建时间",
    "中核更新时间",
    "SaaS校验状态",
    "准标准商品编码",
    "采集状态",
    "同款匹配状态",
    "价格计算状态",
    "newPrice",
    "是否可推送",
    "最终分类",
    "问题说明",
    "SaaS监控状态",
]
PUSHABLE_COL = 17
STEP2_DIR_NAME = "03_step2"
STEP2_OUTPUT_NAMES = {
    "monitor": "STEP2_监控结果.xlsx",
    "pushable": "STEP2_可直接推送.xlsx",
    "non_pushable": "STEP2_不可直接推送.xlsx",
}


def _is_valid_new_price(value: object) -> bool:
    if value is None or isinstance(value, bool):
        return False
    text = str(value).strip()
    if not text or text.lower() in {"null", "none"}:
        return False
    try:
        price = Decimal(text)
        return price.is_finite() and price > 0
    except InvalidOperation:
        return False


def _request(
    config: dict,
    payload: dict,
    *,
    auth_manager: AuthManager | None = None,
    auth_system: str | None = None,
) -> dict:
    def request_factory() -> Request:
        request = Request(
            config["url"],
            data=json.dumps(payload, separators=(",", ":")).encode(),
            method=config["method"],
        )
        apply_curl_headers(request, config["headers"], config["cookie"])
        return request

    if auth_manager is not None and auth_system is not None:
        data = auth_manager.request_json(
            auth_system,
            request_factory,
            read_only=True,
            timeout=60,
        )
    else:
        with internal_direct_opener().open(request_factory(), timeout=60) as response:
            data = json.loads(response.read().decode(response.headers.get_content_charset() or "utf-8"))
    if str(data.get("rspCode")) != "0000":
        raise RuntimeError(f"rspCode={data.get('rspCode')}, rspDesc={data.get('rspDesc')}")
    return data


def _rows(data: dict) -> list[dict]:
    value = data.get("data", data.get("rows", []))
    if isinstance(value, dict):
        value = value.get("rows", value.get("records", []))
    return value if isinstance(value, list) else []


def _status(row: dict) -> tuple[str, str, str, str, str, str]:
    sync, standard, price = (str(row.get(key) or "") for key in ("isSync", "isSyncStandard", "syncPriceStatus"))
    collect = {"1": "采集成功", "0": "采集失败", "2": "采集中"}.get(sync, "未知")
    same = {"1": "同款匹配成功", "0": "同款匹配中", "2": "同款匹配异常/未成功"}.get(standard, "未知")
    price_text = {"1": "价格计算成功", "2": "价格计算中/未完成", "5": "价格计算失败", "": "无价格"}.get(price, "未知")
    if sync != "1":
        problem = "采集未成功"
    elif standard == "0":
        problem = "同款匹配中"
    elif standard != "1":
        problem = "同款匹配异常/未成功"
    elif price == "5":
        problem = "价格检测校验失败"
    elif not _is_valid_new_price(row.get("newPrice")):
        problem = "推送价格为空或无效"
    else:
        return collect, same, price_text, "是", "可推送", ""
    return collect, same, price_text, "否", problem, problem


def _build_monitor_values(row: dict) -> list[object]:
    collect, same, price, pushable, category, problem = _status(row)
    return [
        collect,
        same,
        price,
        row.get("newPrice"),
        pushable,
        category,
        problem,
        "已匹配",
    ]


def _format_sheet(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col in range(1, sheet.max_column + 1):
        width = max(
            len(str(sheet.cell(row, col).value or ""))
            for row in range(1, sheet.max_row + 1)
        )
        sheet.column_dimensions[get_column_letter(col)].width = min(
            max(10, width + 2),
            50,
        )


def _write_workbook(
    records: list[list[object]],
    output: Path,
) -> dict[str, object]:
    pushable = [row for row in records if row[17] == "是"]
    missing = [row for row in records if row[20] == "缺失"]
    needed = [
        row
        for row in records
        if row[17] == "否" and row[20] != "缺失"
    ]
    categories = Counter(row[18] for row in records)
    stats = [
        ["指标", "分类", "数量"],
        ["中核待处理总数", "", len(records)],
        ["SaaS匹配数", "", len(records) - len(missing)],
        ["SaaS监控缺失数", "", len(missing)],
        ["可推送数", "", len(pushable)],
        ["需处理数", "", len(needed)],
        *[
            ["最终分类", key, value]
            for key, value in categories.most_common()
        ],
    ]

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = (
        ("全量监控结果", [MONITOR_HEADERS, *records]),
        ("可推送数据", [MONITOR_HEADERS, *pushable]),
        ("需处理数据", [MONITOR_HEADERS, *needed]),
        ("SaaS监控缺失", [MONITOR_HEADERS, *missing]),
        ("状态统计", stats),
    )
    for name, data in sheets:
        sheet = workbook.create_sheet(name)
        for values in data:
            sheet.append(values)
        _format_sheet(sheet)
    workbook.save(output)
    return {
        "cnnc_total": len(records),
        "saas_matched": len(records) - len(missing),
        "saas_missing": len(missing),
        "pushable": len(pushable),
        "needs_processing": len(needed),
        "categories": dict(categories),
    }


def _collect_records(
    cnnc: dict,
    cnnc_body: dict,
    saas: dict,
    saas_body: dict,
    auth_manager: AuthManager,
) -> list[list[object]]:
    cnnc_rows: list[dict] = []
    page, total = 1, None
    while True:
        data = _request(cnnc, {**cnnc_body, "page": page, "limit": 200})
        expected, page_rows = int(data["count"]), _rows(data)
        total = expected if total is None else total
        if expected != total or not page_rows and len(cnnc_rows) < total:
            raise RuntimeError("中核分页 count 异常")
        cnnc_rows.extend(page_rows)
        if len(cnnc_rows) >= total:
            break
        page += 1
    if len(cnnc_rows) != total:
        raise RuntimeError("中核实际行数与 count 不一致")
    ids = [str(row.get("documentId") or "").strip() for row in cnnc_rows]
    if not all(ids) or len(ids) != len(set(ids)):
        raise RuntimeError("documentId 为空或重复")
    saas_rows: list[dict] = []
    for start in range(0, len(ids), 200):
        data = _request(
            saas,
            {**saas_body, "page": 1, "limit": 200, "commodityId": " ".join(ids[start:start + 200])},
            auth_manager=auth_manager,
            auth_system="saas",
        )
        batch = _rows(data)
        if len(batch) != int(data["count"]):
            raise RuntimeError("SaaS 分批返回 count 异常")
        saas_rows.extend(batch)
    index = {str(row.get("commodityId") or "").strip(): row for row in saas_rows}
    records = []
    for row in cnnc_rows:
        monitor = index.get(str(row["documentId"]).strip())
        monitor_values = (
            _build_monitor_values(monitor)
            if monitor
            else [
                "",
                "",
                "",
                "",
                "否",
                "SaaS监控缺失",
                "中核存在但SaaS监控未返回",
                "缺失",
            ]
        )
        records.append(
            [row.get(key, "") for key in CNNC_KEYS] + monitor_values
        )
    return records


def _write_projection_workbook(records, output: Path, sheet_name: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(MONITOR_HEADERS)
    for row in records:
        sheet.append(row)
    _format_sheet(sheet)
    workbook.save(output)


def _read_sheet_rows(path: Path, sheet_name: str) -> list[list[object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _as_xlsx_value(value: object) -> object:
    # openpyxl 把空字符串写盘后读回为 None；校验按同一规则归一，其余类型逐相等。
    return None if value == "" else value


def _validate_split_projection(records, monitor_path: Path, pushable_path: Path, non_pushable_path: Path) -> None:
    normalized = [[_as_xlsx_value(value) for value in row] for row in records]
    expected_monitor = [MONITOR_HEADERS, *normalized]
    expected_pushable = [MONITOR_HEADERS, *[row for row in normalized if row[PUSHABLE_COL] == "是"]]
    expected_non_pushable = [MONITOR_HEADERS, *[row for row in normalized if row[PUSHABLE_COL] == "否"]]
    monitor_rows = _read_sheet_rows(monitor_path, "全量监控结果")
    if monitor_rows != expected_monitor:
        raise RuntimeError(f"monitor workbook rows do not match authoritative records: {monitor_path}")
    pushable_rows = _read_sheet_rows(pushable_path, "可直接推送")
    if pushable_rows != expected_pushable:
        raise RuntimeError(f"pushable projection mismatch: {pushable_path}")
    non_pushable_rows = _read_sheet_rows(non_pushable_path, "不可直接推送")
    if non_pushable_rows != expected_non_pushable:
        raise RuntimeError(f"non-pushable projection mismatch: {non_pushable_path}")
    pushable_ids = {row[0] for row in pushable_rows[1:]}
    non_pushable_ids = {row[0] for row in non_pushable_rows[1:]}
    monitor_ids = {row[0] for row in monitor_rows[1:]}
    if pushable_ids & non_pushable_ids:
        raise RuntimeError("pushable and non-pushable documentIds intersect")
    if pushable_ids | non_pushable_ids != monitor_ids:
        raise RuntimeError("pushable + non-pushable do not cover monitor documentIds")


def _begin_step2_run(business_date: str, base=None) -> dict:
    ensure_workday_layout(business_date, base)
    manifest = read_manifest(business_date, base)
    previous = manifest["stages"].get("step2") or {}
    now = now_in_shanghai()
    run_info = {
        "run_id": now.strftime("%Y%m%dT%H%M%S") + f"-{now.microsecond:06d}",
        "generation": int(previous.get("generation") or 0) + 1,
        "started_at": now.isoformat(timespec="seconds"),
    }
    update_stage(
        business_date,
        "step2",
        {"status": "RUNNING", **run_info},
        base,
    )
    return run_info


def _fail_step2_run(business_date: str, error: BaseException, base=None) -> None:
    try:
        update_stage(
            business_date,
            "step2",
            {
                "status": "FAILED",
                "error_type": type(error).__name__,
                "error_summary": str(error)[:500],
            },
            base,
        )
    except Exception:
        pass  # 失败记录失败不得掩盖原始业务错误


def _file_record(relative: str, path: Path) -> dict:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return {"path": relative, "sha256": digest.hexdigest(), "size": path.stat().st_size}


def _complete_step2_publication(records, business_date: str, run_info: dict, base=None) -> dict:
    step2_dir = workday_root(business_date, base) / STEP2_DIR_NAME
    staging_dir = step2_dir / ".staging" / run_info["run_id"]
    staging_dir.mkdir(parents=True, exist_ok=False)
    pushable_rows = [row for row in records if row[PUSHABLE_COL] == "是"]
    non_pushable_rows = [row for row in records if row[PUSHABLE_COL] == "否"]
    staged = {key: staging_dir / name for key, name in STEP2_OUTPUT_NAMES.items()}
    try:
        summary = _write_workbook(records, staged["monitor"])
        _write_projection_workbook(pushable_rows, staged["pushable"], "可直接推送")
        _write_projection_workbook(non_pushable_rows, staged["non_pushable"], "不可直接推送")
        _validate_split_projection(records, staged["monitor"], staged["pushable"], staged["non_pushable"])
        for name in STEP2_OUTPUT_NAMES.values():
            os.replace(staging_dir / name, step2_dir / name)
    except BaseException:
        shutil.rmtree(staging_dir, ignore_errors=True)
        raise
    files = {
        key: _file_record(f"{STEP2_DIR_NAME}/{name}", step2_dir / name)
        for key, name in STEP2_OUTPUT_NAMES.items()
    }
    counts = {
        "monitor_total": len(records),
        "pushable": len(pushable_rows),
        "non_pushable": len(non_pushable_rows),
    }
    update_stage(
        business_date,
        "step2",
        {
            "status": "COMPLETED",
            "completed_at": now_in_shanghai().isoformat(timespec="seconds"),
            "outputs": {key: f"{STEP2_DIR_NAME}/{name}" for key, name in STEP2_OUTPUT_NAMES.items()},
            "files": files,
            "counts": counts,
        },
        base,
    )
    shutil.rmtree(staging_dir, ignore_errors=True)
    return {
        "business_date": business_date,
        "run_id": run_info["run_id"],
        "generation": run_info["generation"],
        "outputs": {key: str(step2_dir / name) for key, name in STEP2_OUTPUT_NAMES.items()},
        "counts": counts,
        "summary": summary,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="查询截至当前的中核 SaaS 监控状态并生成五 Sheet Excel。"
    )
    parser.add_argument(
        "--business-date",
        default=None,
        help="业务日期 YYYY-MM-DD；缺省使用上海时区当前业务日",
    )
    return parser


def main(argv: list[str] | None = None) -> None:
    args = build_parser().parse_args(argv)
    business_date = args.business_date or business_date_in_shanghai()
    run_info = _begin_step2_run(business_date, ROOT)
    try:
        apply_internal_no_proxy_env()
        cnnc = load_curl_command(ROOT / "data/latest_cnnc_query.curl")
        saas = load_curl_command(ROOT / "data/saas_zhuangtai.curl")
        cnnc_body, saas_body = json.loads(cnnc["data"]), json.loads(saas["data"])
        if cnnc_body.get("sassCheckStatuss") != ["0"] or (str(saas_body.get("tenantId")), str(saas_body.get("commodityPoolId"))) != ("188605", "222"):
            raise RuntimeError("租户、商品池或中核待处理筛选不符合第二步范围")
        auth_manager = AuthManager()
        saas = auth_manager.prepare_curl("saas", saas, refresh=True)
        records = _collect_records(cnnc, cnnc_body, saas, saas_body, auth_manager)
        publication = _complete_step2_publication(records, business_date, run_info, ROOT)
    except BaseException as error:
        _fail_step2_run(business_date, error, ROOT)
        raise
    print(json.dumps(publication, ensure_ascii=False))


if __name__ == "__main__":
    main()
