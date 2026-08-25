from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Sequence
import urllib.request
from urllib.parse import urlparse, urlunparse

from openpyxl import load_workbook

try:
    from scripts.auth_manager import AuthManager
    from scripts.curl_utils import (
        apply_curl_headers,
        internal_direct_opener,
        load_curl_command,
    )
except ModuleNotFoundError:  # Direct invocation: python scripts/reject_saas_audit_excel.py
    from auth_manager import AuthManager
    from curl_utils import apply_curl_headers, internal_direct_opener, load_curl_command


TENANT_ID = "188605"
COMMODITY_POOL_ID = "222"
REJECT_FILL_COLORS = {
    "FFFFC7CE": "red",
    "FFFFFF00": "yellow",
}
REQUIRED_HEADERS = {
    "SaaS商品ID",
    "SKU编号",
    "互联网标题",
    "模型审核结果",
}
FILL_COLOR_MODES = {
    "red-yellow": {"red", "yellow"},
    "yellow-only": {"yellow"},
    "red-only": {"red"},
}


@dataclass(frozen=True)
class RejectCandidate:
    v_commodity_id: str
    sku_id: str
    model_audit_result: str
    color: str
    excel_row: int
    source_excel_rows: tuple[int, ...] = ()

    def as_payload_item(self) -> dict[str, str]:
        return {
            "vCommodityId": self.v_commodity_id,
            "skuId": self.sku_id,
            "modelAuditResult": self.model_audit_result,
        }

    def as_manifest_item(self) -> dict[str, Any]:
        return {
            **self.as_payload_item(),
            "color": self.color,
            "excelRow": self.excel_row,
            "sourceExcelRows": list(self.source_excel_rows or (self.excel_row,)),
        }


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_is_nonempty(value: Any) -> bool:
    if value is None:
        return False
    return not isinstance(value, str) or bool(value.strip())


def load_reject_candidates(
    input_xlsx: Path,
    *,
    allowed_colors: set[str] | None = None,
    require_customer_price: bool = False,
    require_customer_match: bool = False,
) -> list[RejectCandidate]:
    selected_colors = set(allowed_colors or REJECT_FILL_COLORS.values())
    unsupported_colors = selected_colors - set(REJECT_FILL_COLORS.values())
    if not selected_colors or unsupported_colors:
        raise ValueError(f"Unsupported reject colors: {sorted(unsupported_colors)}")
    workbook = load_workbook(input_xlsx, data_only=True, read_only=False)
    if "record_results" not in workbook.sheetnames:
        raise ValueError("Workbook is missing record_results sheet")
    worksheet = workbook["record_results"]
    header_columns = {
        _cell_text(worksheet.cell(1, column).value): column
        for column in range(1, worksheet.max_column + 1)
    }
    missing_headers = REQUIRED_HEADERS - set(header_columns)
    if require_customer_price and "客户价" not in header_columns:
        missing_headers.add("客户价")
    if require_customer_match and "客户数据匹配状态" not in header_columns:
        missing_headers.add("客户数据匹配状态")
    if missing_headers:
        raise ValueError(f"Workbook is missing required headers: {sorted(missing_headers)}")

    title_column = header_columns["互联网标题"]
    candidates: list[RejectCandidate] = []
    candidate_indexes: dict[tuple[str, str], int] = {}
    for row in range(2, worksheet.max_row + 1):
        cell = worksheet.cell(row, title_column)
        color = REJECT_FILL_COLORS.get(cell.fill.fgColor.rgb)
        if color not in selected_colors:
            continue
        if require_customer_price and not _cell_is_nonempty(
            worksheet.cell(row, header_columns["客户价"]).value
        ):
            continue
        if require_customer_match and _cell_text(
            worksheet.cell(row, header_columns["客户数据匹配状态"]).value
        ) != "已匹配":
            continue
        v_commodity_id = _cell_text(
            worksheet.cell(row, header_columns["SaaS商品ID"]).value
        )
        sku_id = _cell_text(worksheet.cell(row, header_columns["SKU编号"]).value)
        model_audit_result = _cell_text(
            worksheet.cell(row, header_columns["模型审核结果"]).value
        )
        if not v_commodity_id or not sku_id:
            raise ValueError(f"Row {row} has a blank vCommodityId or skuId")
        if model_audit_result != "1":
            raise ValueError(
                f"Row {row} has unsupported modelAuditResult={model_audit_result!r}"
            )
        business_key = (v_commodity_id, sku_id)
        existing_index = candidate_indexes.get(business_key)
        if existing_index is not None:
            existing = candidates[existing_index]
            candidates[existing_index] = RejectCandidate(
                v_commodity_id=existing.v_commodity_id,
                sku_id=existing.sku_id,
                model_audit_result=existing.model_audit_result,
                color=existing.color,
                excel_row=existing.excel_row,
                source_excel_rows=existing.source_excel_rows + (row,),
            )
            continue
        candidate_indexes[business_key] = len(candidates)
        candidates.append(
            RejectCandidate(
                v_commodity_id=v_commodity_id,
                sku_id=sku_id,
                model_audit_result=model_audit_result,
                color=color,
                excel_row=row,
                source_excel_rows=(row,),
            )
        )
    if not candidates:
        raise ValueError("Workbook has no reject candidates for the selected filters")
    return candidates


def build_reject_payload(candidates: Sequence[RejectCandidate]) -> dict[str, Any]:
    if not candidates:
        raise ValueError("Reject payload cannot be empty")
    return {
        "tenantId": TENANT_ID,
        "commodityPoolId": COMMODITY_POOL_ID,
        "commodityList": [candidate.as_payload_item() for candidate in candidates],
    }


def candidate_manifest_sha256(candidates: Sequence[RejectCandidate]) -> str:
    encoded = json.dumps(
        [candidate.as_manifest_item() for candidate in candidates],
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def validate_execution_confirmation(
    candidates: Sequence[RejectCandidate],
    *,
    expected_count: int,
    expected_sha256: str,
) -> None:
    if len(candidates) != expected_count:
        raise ValueError(
            f"Candidate count changed: expected {expected_count}, actual {len(candidates)}"
        )
    actual_sha256 = candidate_manifest_sha256(candidates)
    if actual_sha256.lower() != expected_sha256.strip().lower():
        raise ValueError(
            f"Candidate manifest SHA-256 changed: expected {expected_sha256}, "
            f"actual {actual_sha256}"
        )


def reject_url_from_curl(parsed_curl: dict[str, Any]) -> str:
    parsed_url = urlparse(str(parsed_curl.get("url") or ""))
    if parsed_url.hostname != "saasportal.iio.com":
        raise ValueError("cURL host must be saasportal.iio.com")
    if parsed_url.scheme not in {"http", "https"}:
        raise ValueError("cURL scheme must be HTTP or HTTPS")
    if not str(parsed_curl.get("cookie") or "").strip():
        raise ValueError("cURL must contain a Cookie header")
    allowed_paths = {
        "/improxy/saas/iio2/sameAudit/queryAll",
        "/improxy/saas/iio2/sameAudit/reject",
    }
    if parsed_url.path not in allowed_paths:
        raise ValueError("cURL path must be a sameAudit queryAll or reject endpoint")
    return urlunparse(parsed_url._replace(path="/improxy/saas/iio2/sameAudit/reject"))


def query_url_from_curl(parsed_curl: dict[str, Any]) -> str:
    parsed_url = urlparse(str(parsed_curl.get("url") or ""))
    if parsed_url.hostname != "saasportal.iio.com":
        raise ValueError("cURL host must be saasportal.iio.com")
    if parsed_url.scheme not in {"http", "https"}:
        raise ValueError("cURL scheme must be HTTP or HTTPS")
    if not str(parsed_curl.get("cookie") or "").strip():
        raise ValueError("cURL must contain a Cookie header")
    if parsed_url.path not in {
        "/improxy/saas/iio2/sameAudit/queryAll",
        "/improxy/saas/iio2/sameAudit/reject",
        "/improxy/saas/iio2/sameAudit/approve",
    }:
        raise ValueError(
            "cURL path must be a sameAudit queryAll, reject, or approve endpoint"
        )
    return urlunparse(parsed_url._replace(path="/improxy/saas/iio2/sameAudit/queryAll"))


def _decode_json_response(response: Any) -> dict[str, Any]:
    charset = response.headers.get_content_charset() or "utf-8"
    decoded = json.loads(response.read().decode(charset))
    if not isinstance(decoded, dict):
        raise RuntimeError("SaaS response must be a JSON object")
    return decoded


def fetch_pending_pairs(
    parsed_curl: dict[str, Any],
    *,
    opener: Any | None = None,
    page_size: int = 200,
    timeout: int = 30,
    include_all_reference_states: bool = False,
    auth_manager: AuthManager | None = None,
) -> set[tuple[str, str]]:
    if page_size <= 0 or page_size > 200:
        raise ValueError("page_size must be between 1 and 200")
    try:
        base_payload = json.loads(str(parsed_curl.get("data") or "{}"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"cURL request body is not valid JSON: {exc}") from exc
    if not isinstance(base_payload, dict):
        raise ValueError("cURL request body must be a JSON object")
    source_path = urlparse(str(parsed_curl.get("url") or "")).path
    if source_path != "/improxy/saas/iio2/sameAudit/queryAll":
        base_payload = {}
    query_url = query_url_from_curl(parsed_curl)
    direct_opener = opener or internal_direct_opener()
    rows: list[dict[str, Any]] = []
    expected_count: int | None = None
    page = 1
    while expected_count is None or len(rows) < expected_count:
        payload = dict(base_payload)
        payload.update(
            {
                "tenantId": TENANT_ID,
                "commodityPoolId": COMMODITY_POOL_ID,
                "commodityId": "",
                "commodityName": "",
                "syncState": "1",
                "isReferUrl": (
                    "" if include_all_reference_states else base_payload.get("isReferUrl", "")
                ),
                "page": page,
                "limit": page_size,
            }
        )
        def request_factory() -> urllib.request.Request:
            request = urllib.request.Request(
                query_url,
                data=json.dumps(
                    payload, ensure_ascii=False, separators=(",", ":")
                ).encode("utf-8"),
                method="POST",
            )
            apply_curl_headers(
                request,
                parsed_curl.get("headers") or [],
                parsed_curl["cookie"],
            )
            return request

        if auth_manager is not None:
            result = auth_manager.request_json(
                "saas",
                request_factory,
                read_only=True,
                opener=direct_opener,
                timeout=timeout,
            )
        else:
            with direct_opener.open(request_factory(), timeout=timeout) as response:
                result = _decode_json_response(response)
        if str(result.get("rspCode")) != "0000":
            raise RuntimeError(
                f"Pending preflight failed: rspCode={result.get('rspCode')}, "
                f"rspDesc={result.get('rspDesc')}"
            )
        page_rows = result.get("rows")
        if not isinstance(page_rows, list):
            raise RuntimeError("Pending preflight response has no rows list")
        current_count = int(result.get("count"))
        if expected_count is None:
            expected_count = current_count
        elif current_count != expected_count:
            raise RuntimeError(
                f"Pending count changed during pagination: {expected_count} -> {current_count}"
            )
        rows.extend(page_rows)
        if not page_rows and len(rows) < expected_count:
            raise RuntimeError(
                f"Pending pagination stopped early: actual {len(rows)}, expected {expected_count}"
            )
        page += 1
    if expected_count is None or len(rows) != expected_count:
        raise RuntimeError(
            f"Pending pagination mismatch: actual {len(rows)}, expected {expected_count}"
        )
    pairs: set[tuple[str, str]] = set()
    for row_number, row in enumerate(rows, start=1):
        pair = (_cell_text(row.get("vCommodityId")), _cell_text(row.get("skuId")))
        if not all(pair):
            raise RuntimeError(f"Pending row {row_number} has a blank business key")
        pairs.add(pair)
    return pairs


def validate_candidates_are_pending(
    candidates: Sequence[RejectCandidate], pending_pairs: set[tuple[str, str]]
) -> None:
    missing = [
        candidate
        for candidate in candidates
        if (candidate.v_commodity_id, candidate.sku_id) not in pending_pairs
    ]
    if missing:
        raise ValueError(
            f"{len(missing)} candidate(s) are no longer pending; first Excel row "
            f"is {missing[0].excel_row}"
        )


def select_pending_candidates(
    candidates: Sequence[RejectCandidate],
    pending_pairs: set[tuple[str, str]],
    *,
    expected_missing_count: int,
) -> tuple[list[RejectCandidate], list[RejectCandidate]]:
    selected: list[RejectCandidate] = []
    missing: list[RejectCandidate] = []
    for candidate in candidates:
        if (candidate.v_commodity_id, candidate.sku_id) in pending_pairs:
            selected.append(candidate)
        else:
            missing.append(candidate)
    if len(missing) != expected_missing_count:
        raise ValueError(
            f"Already-processed candidate count changed: expected "
            f"{expected_missing_count}, actual {len(missing)}"
        )
    if not selected:
        raise ValueError("No pending candidates remain for rejection")
    return selected, missing


def execute_reject_batches(
    parsed_curl: dict[str, Any],
    candidates: Sequence[RejectCandidate],
    *,
    batch_size: int = 50,
    opener: Any | None = None,
    timeout: int = 30,
    on_batch_success: Callable[[dict[str, Any]], None] | None = None,
) -> list[dict[str, Any]]:
    if batch_size <= 0 or batch_size > 100:
        raise ValueError("batch_size must be between 1 and 100")
    reject_url = reject_url_from_curl(parsed_curl)
    direct_opener = opener or internal_direct_opener()
    events: list[dict[str, Any]] = []
    for offset in range(0, len(candidates), batch_size):
        batch = candidates[offset : offset + batch_size]
        request = urllib.request.Request(
            reject_url,
            data=json.dumps(
                build_reject_payload(batch),
                ensure_ascii=False,
                separators=(",", ":"),
            ).encode("utf-8"),
            method="POST",
        )
        apply_curl_headers(request, parsed_curl.get("headers") or [], parsed_curl["cookie"])
        with direct_opener.open(request, timeout=timeout) as response:
            result = _decode_json_response(response)
        rsp_code = str(result.get("rspCode"))
        if rsp_code != "0000":
            raise RuntimeError(
                f"Reject batch {len(events) + 1} failed: rspCode={rsp_code}, "
                f"rspDesc={result.get('rspDesc')}"
            )
        event = {
            "batchIndex": len(events) + 1,
            "batchCount": len(batch),
            "firstExcelRow": batch[0].excel_row,
            "lastExcelRow": batch[-1].excel_row,
            "rspCode": rsp_code,
            "rspDesc": _cell_text(result.get("rspDesc")),
        }
        events.append(event)
        if on_batch_success:
            on_batch_success(event)
    return events


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Preview a colour-scoped SaaS reject batch. Execution is separately gated."
    )
    parser.add_argument("--input-xlsx", type=Path, required=True)
    parser.add_argument("--curl-file", type=Path, required=True)
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--expected-count", type=int)
    parser.add_argument("--expected-sha256")
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--batch-size", type=int, default=50)
    parser.add_argument("--expected-already-processed-count", type=int, default=0)
    parser.add_argument(
        "--fill-colors",
        choices=tuple(FILL_COLOR_MODES),
        default="red-yellow",
    )
    parser.add_argument("--require-customer-price", action="store_true")
    parser.add_argument("--require-customer-match", action="store_true")
    return parser


def run_main(
    argv: Sequence[str] | None = None,
    *,
    auth_manager: AuthManager | None = None,
) -> int:
    args = build_parser().parse_args(argv)
    allowed_colors = FILL_COLOR_MODES[args.fill_colors]
    candidates = load_reject_candidates(
        args.input_xlsx,
        allowed_colors=allowed_colors,
        require_customer_price=args.require_customer_price,
        require_customer_match=args.require_customer_match,
    )
    parsed_curl = load_curl_command(args.curl_file)
    reject_url_from_curl(parsed_curl)
    manager = auth_manager or AuthManager()
    parsed_curl = manager.prepare_curl("saas", parsed_curl, refresh=True)
    reject_url_from_curl(parsed_curl)
    source_digest = candidate_manifest_sha256(candidates)
    print(f"candidate_count={len(candidates)}")
    print(f"candidate_manifest_sha256={source_digest}")
    if args.execute:
        if args.expected_count is None or not args.expected_sha256:
            raise ValueError("--execute requires --expected-count and --expected-sha256")
        if args.output_dir is None:
            raise ValueError("--execute requires --output-dir")
        args.output_dir.mkdir(parents=True, exist_ok=False)
        events_path = args.output_dir / "reject_events.jsonl"
        summary_path = args.output_dir / "reject_summary.json"
        events: list[dict[str, Any]] = []
        execution_candidates = list(candidates)
        already_processed: list[RejectCandidate] = []
        execution_digest = source_digest

        def record_success(event: dict[str, Any]) -> None:
            events.append(event)
            with events_path.open("a", encoding="utf-8") as output:
                output.write(json.dumps(event, ensure_ascii=False, separators=(",", ":")))
                output.write("\n")

        try:
            pending_before = fetch_pending_pairs(
                parsed_curl,
                include_all_reference_states=True,
                auth_manager=manager,
            )
            execution_candidates, already_processed = select_pending_candidates(
                candidates,
                pending_before,
                expected_missing_count=args.expected_already_processed_count,
            )
            execution_digest = candidate_manifest_sha256(execution_candidates)
            validate_execution_confirmation(
                execution_candidates,
                expected_count=args.expected_count,
                expected_sha256=args.expected_sha256,
            )
            manifest = {
                "tenantId": TENANT_ID,
                "commodityPoolId": COMMODITY_POOL_ID,
                "inputXlsx": str(args.input_xlsx.resolve()),
                "sourceCandidateCount": len(candidates),
                "candidateCount": len(execution_candidates),
                "alreadyProcessedCount": len(already_processed),
                "candidateManifestSha256": execution_digest,
                "selection": {
                    "fillColors": sorted(allowed_colors),
                    "requireCustomerPrice": args.require_customer_price,
                    "requireCustomerMatch": args.require_customer_match,
                },
                "colors": dict(
                    Counter(candidate.color for candidate in execution_candidates)
                ),
                "candidates": [
                    candidate.as_manifest_item() for candidate in execution_candidates
                ],
                "alreadyProcessedCandidates": [
                    candidate.as_manifest_item() for candidate in already_processed
                ],
            }
            (args.output_dir / "reject_manifest.json").write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            execute_reject_batches(
                parsed_curl,
                execution_candidates,
                batch_size=args.batch_size,
                on_batch_success=record_success,
            )
            pending_after = fetch_pending_pairs(
                parsed_curl,
                include_all_reference_states=True,
                auth_manager=manager,
            )
            remaining = [
                candidate
                for candidate in execution_candidates
                if (candidate.v_commodity_id, candidate.sku_id) in pending_after
            ]
            if remaining:
                raise RuntimeError(
                    f"Postflight found {len(remaining)} candidate(s) still pending"
                )
            summary = {
                "status": "COMPLETED",
                "sourceCandidateCount": len(candidates),
                "candidateCount": len(execution_candidates),
                "alreadyProcessedCount": len(already_processed),
                "rejectedCount": sum(event["batchCount"] for event in events),
                "batchCount": len(events),
                "pendingBeforeCount": len(pending_before),
                "pendingAfterCount": len(pending_after),
                "remainingPendingCount": 0,
                "candidateManifestSha256": execution_digest,
            }
        except Exception as exc:
            summary = {
                "status": "BLOCKED",
                "sourceCandidateCount": len(candidates),
                "candidateCount": len(execution_candidates),
                "alreadyProcessedCount": len(already_processed),
                "confirmedBatchCount": len(events),
                "confirmedRejectedCount": sum(
                    event["batchCount"] for event in events
                ),
                "candidateManifestSha256": execution_digest,
                "errorType": type(exc).__name__,
                "error": str(exc),
            }
            summary_path.write_text(
                json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            raise
        summary_path.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print(f"status={summary['status']}")
        print(f"rejected_count={summary['rejectedCount']}")
        print(f"output_dir={args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(run_main())
