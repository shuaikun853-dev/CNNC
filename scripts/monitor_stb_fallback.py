import os
import re
import subprocess
import sys
from pathlib import Path
from urllib.parse import urlparse

import openpyxl


REQUIRED_DB_VARIABLES = (
    "ZHAOSHANG_DB_HOST",
    "ZHAOSHANG_DB_PORT",
    "ZHAOSHANG_DB_USER",
    "ZHAOSHANG_DB_PASSWORD",
)
STB_ROOT = Path(r"E:\.A2026\多商城爬取\multi_mall_crawler")
STB_SCRIPT = STB_ROOT / "scripts" / "stbchina_fast_crawler.py"
URL_RE = re.compile(r"https?://[^\s,，;；\n\r\t]+")


def find_header(headers, *candidates):
    for candidate in candidates:
        if candidate in headers:
            return candidate
    raise RuntimeError("监控表缺少列: " + " / ".join(candidates))


def load_zhaoshang_db_env(path):
    path = Path(path)
    values = {}
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        key, separator, value = line.partition("=")
        if separator:
            cleaned = value.strip()
            if len(cleaned) >= 2 and cleaned[0] == cleaned[-1] and cleaned[0] in "\"'":
                cleaned = cleaned[1:-1]
            values[key.strip()] = cleaned
    missing = [key for key in REQUIRED_DB_VARIABLES if not values.get(key)]
    if missing:
        raise RuntimeError("招商局数据库配置缺少: " + ", ".join(missing))
    return {key: values[key] for key in REQUIRED_DB_VARIABLES}


def is_stbchina_url(url):
    host = (urlparse(str(url or "")).hostname or "").lower()
    return host == "stbchina.cn" or host.endswith(".stbchina.cn")


def classify_priority_output(text):
    if "ACTIVE_UNCONFIRMED" in text:
        return "failed"
    if "COMPLETED_UNCONFIRMED" in text:
        return "completed_unconfirmed"
    if "ALREADY_HIGHEST" in text:
        return "already_highest_priority"
    if (
        "【优先级调整】WARNING" in text
        or "[优先级调整] WARNING" in text
        or "调整失败" in text
        or "update 影响行数为 0" in text
    ):
        return "failed"
    updated = re.search(r"updated rows=(\d+)", text)
    if updated and int(updated.group(1)) > 0:
        return "updated"
    affected = re.search(r"affected_rows\D+(\d+)", text)
    if affected and int(affected.group(1)) > 0:
        return "updated"
    return "missing"


def save_or_reuse_locked_workbook(workbook, path, expected_rows):
    path = Path(path)
    try:
        workbook.save(path)
        return "saved"
    except PermissionError:
        if not path.exists():
            raise
        existing = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if existing.active.max_row != expected_rows:
                raise
        finally:
            existing.close()
        return "reused"


def read_stbchina_results(path):
    records = {}
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        columns = {header: index for index, header in enumerate(headers)}
        required = ("商品链接", "商品价格", "采集状态")
        if any(name not in columns for name in required):
            raise RuntimeError("史泰博结果缺少必需列")
        for row in sheet.iter_rows(min_row=2, values_only=True):
            url = str(row[columns["商品链接"]] or "").strip()
            if not url or str(row[columns["采集状态"]] or "").strip() != "成功":
                continue
            try:
                price = float(row[columns["商品价格"]])
            except (TypeError, ValueError):
                continue
            if price < 0:
                continue
            records[url] = {
                "price": price,
                "title": str(row[columns.get("商品标题", -1)] or "").strip() if "商品标题" in columns else "",
                "uid": str(row[columns.get("uid", -1)] or "").strip() if "uid" in columns else "",
                "comment": str(row[columns.get("评论数", -1)] or "").strip() if "评论数" in columns else "",
            }
    finally:
        workbook.close()
    return records


def extract_detail_reference_records(document_id, detail_data):
    document_id = str(document_id or "").strip()
    if not document_id or not isinstance(detail_data, dict):
        return {}
    title = ""
    for key in ("commodityName", "sassVCommodityName", "vCommodityName", "itemName"):
        title = str(detail_data.get(key) or "").strip()
        if title:
            break
    if not title:
        return {}

    records = {}
    reference_infos = detail_data.get("referenceInfos")
    if not isinstance(reference_infos, list):
        return records
    for item in reference_infos:
        if not isinstance(item, dict):
            continue
        try:
            price = float(item.get("referencePrice"))
        except (TypeError, ValueError):
            continue
        if price < 0:
            continue
        for url in URL_RE.findall(str(item.get("referenceLink") or "")):
            records[(document_id, url)] = {"price": price, "title": title}
    return records


def apply_detail_reference_results(rows, records):
    changed = 0
    for row in rows:
        if row.价格 not in ("", None) or row.互联网标题 not in ("", None):
            continue
        record = records.get((str(row.单据id).strip(), str(row.URL).strip()))
        if not record or record.get("price") in ("", None) or not record.get("title"):
            continue
        row.价格 = record["price"]
        row.互联网标题 = record["title"]
        row.爬取状态 = "中核detail参考回填"
        row.爬取失败原因 = ""
        changed += 1
    return changed


def run_stbchina_fallback(urls, output_dir):
    urls = list(dict.fromkeys(url for url in urls if is_stbchina_url(url)))
    if not urls:
        return {}
    output_dir = Path(output_dir)
    input_path = output_dir / "stbchina_missing_urls.txt"
    result_path = output_dir / "stbchina_missing_result.xlsx"
    if result_path.exists():
        existing = read_stbchina_results(result_path)
        if all(url in existing for url in urls):
            return {url: existing[url] for url in urls}
    if not STB_SCRIPT.exists():
        raise RuntimeError("史泰博专用脚本不存在")
    input_path.write_text("\n".join(urls) + "\n", encoding="utf-8")
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    proc = subprocess.run(
        [sys.executable, str(STB_SCRIPT), "--input", str(input_path), "--output", str(result_path)],
        cwd=STB_ROOT,
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=900,
    )
    if proc.returncode != 0:
        raise RuntimeError("史泰博专用脚本失败")
    if not result_path.exists():
        raise RuntimeError("史泰博专用脚本未生成结果文件")
    return read_stbchina_results(result_path)


def apply_stbchina_results(rows, records):
    changed = 0
    for row in rows:
        record = records.get(row.URL)
        if not record:
            continue
        row.价格 = record["price"]
        row.互联网标题 = record["title"]
        if record["uid"]:
            row.SKUID = record["uid"]
        row.销量_评论数 = record["comment"]
        changed += 1
    return changed
