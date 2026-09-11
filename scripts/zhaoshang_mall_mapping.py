from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MAPPING_PATH = ROOT / "商城sourceId映射表" / "商城sourceId映射表.xlsx"

EXTRA_DOMAIN_ALIASES = {
    "jd": ("item.jd.com",),
    "zjmi": ("zjmi.com",),
    "nbdeli": ("nbdeli.cn", "delib2b.com"),
    "epec": ("emall.epec.com", "lt95.com"),
    "comix": ("qx.com",),
    "ofs": ("officemate.cn",),
}

SUPPORTED_ZHAOSHANG_MALL_KEYS = {
    "jd",
    "colipu",
    "zkh",
    "xfs",
    "comix",
    "xhgj",
    "epec",
    "ycw",
    "tima",
    "zjmi",
    "mymro",
    "nbdeli",
    "lxwl",
    "ofs",
    "qstbg",
    "lxwlmall",
    "jslink",
    "ehsy",
    "haoxiny",
    "guangbo",
}


@dataclass(frozen=True)
class ZhaoshangMallMapping:
    mall_name: str
    source_id: str
    mall_key: str
    domains: tuple[str, ...]
    vendor_keywords: tuple[str, ...]


def _norm(value: object) -> str:
    return str(value or "").strip()


def _host_from_url(value: object) -> str:
    host = (urlparse(_norm(value)).hostname or "").lower().strip()
    return host.lstrip(".")


def _split_domains(value: object) -> list[str]:
    text = _norm(value)
    if not text:
        return []
    parts = []
    for raw in text.replace("，", ",").replace("；", ",").replace(";", ",").split(","):
        item = raw.strip().lower().lstrip(".")
        if item:
            parts.append(item)
    return parts


def _dedupe(items: list[str]) -> tuple[str, ...]:
    seen = set()
    result = []
    for item in items:
        candidates = [item]
        if item.startswith("www."):
            candidates.append(item[4:])
        for candidate in candidates:
            if candidate and candidate not in seen:
                seen.add(candidate)
                result.append(candidate)
    return tuple(result)


def load_zhaoshang_mall_mappings(path: Path | None = None) -> list[ZhaoshangMallMapping]:
    mapping_path = path or DEFAULT_MAPPING_PATH
    if not mapping_path.exists():
        return []

    wb = load_workbook(mapping_path, data_only=True, read_only=True)
    try:
        ws = wb.active
        if ws is None:
            return []

        header_row_index = None
        header_map: dict[str, int] = {}
        for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [_norm(v) for v in row]
            if "商城名称" in values and "sourceId" in values:
                header_row_index = row_index
                header_map = {name: idx for idx, name in enumerate(values) if name}
                break
        if header_row_index is None:
            return []

        def col(*names: str) -> int:
            for name in names:
                if name in header_map:
                    return header_map[name]
            return -1

        name_idx = col("商城名称")
        source_idx = col("sourceId")
        key_idx = col("商城标识 identification", "identification")
        url_idx = col("商城URL")
        domain_idx = col("匹配域名 matchDomain", "matchDomain")
        if min(name_idx, source_idx, key_idx) < 0:
            return []

        mappings: list[ZhaoshangMallMapping] = []
        for row in ws.iter_rows(min_row=header_row_index + 1, values_only=True):
            values = list(row)
            mall_name = _norm(values[name_idx] if name_idx < len(values) else "")
            source_id = _norm(values[source_idx] if source_idx < len(values) else "")
            mall_key = _norm(values[key_idx] if key_idx < len(values) else "")
            if not mall_name or not source_id or not mall_key:
                continue
            if mall_key not in SUPPORTED_ZHAOSHANG_MALL_KEYS:
                continue

            domains: list[str] = []
            if url_idx >= 0 and url_idx < len(values):
                domains.append(_host_from_url(values[url_idx]))
            if domain_idx >= 0 and domain_idx < len(values):
                domains.extend(_split_domains(values[domain_idx]))
            domains.extend(EXTRA_DOMAIN_ALIASES.get(mall_key, ()))

            mappings.append(
                ZhaoshangMallMapping(
                    mall_name=mall_name,
                    source_id=source_id,
                    mall_key=mall_key,
                    domains=_dedupe(domains),
                    vendor_keywords=_dedupe([mall_key, mall_name.lower()]),
                )
            )
        return mappings
    finally:
        wb.close()


def is_zhaoshang_domain(url: str, mappings: list[ZhaoshangMallMapping] | None = None) -> bool:
    host = (urlparse(url or "").hostname or "").lower().lstrip(".")
    if not host:
        return False
    for mapping in mappings if mappings is not None else load_zhaoshang_mall_mappings():
        for domain in mapping.domains:
            domain = domain.lower().lstrip(".")
            if host == domain or host.endswith("." + domain):
                return True
    return False
