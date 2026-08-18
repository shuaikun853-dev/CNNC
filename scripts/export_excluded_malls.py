"""把走招商局任务系统的商城导出到爬取模板，按当日日期分目录。

用法：
  python scripts/export_excluded_malls.py \
    --input-saas-json ./data/saas_json.json \
    --output-dir      ./爬取模板

输出目录结构：
  爬取模板/
    YYYY-MM-DD/
      京东.xlsx          # 表头：商品编码 + url
      震坤行.xlsx
      summary.txt        # 数量统计

安全约束：
  - 不调任何外部接口
  - 不读 cookies/token
  - 不打印任何敏感信息
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from app.services.saas_audit_core import (  # noqa: E402
    _parse_json_segments,
    is_jd_link,
    normalize_cell_str,
)
from scripts.zhaoshang_mall_mapping import ZhaoshangMallMapping, load_zhaoshang_mall_mappings  # noqa: E402


EXCLUDED_MALLS = load_zhaoshang_mall_mappings()


def _mall_predicate(row: dict, domains: tuple[str, ...], vendor_keywords: tuple[str, ...]) -> bool:
    url = row.get("extLink", "") or ""
    if (urlparse(url).hostname or "").strip():
        return _host_match(url, domains)
    vendor = normalize_cell_str(row.get("vendorName")).lower()
    return any(kw in vendor for kw in vendor_keywords)


def _host_match(url: str, suffixes: tuple[str, ...]) -> bool:
    host = (urlparse(url or "").hostname or "").lower()
    if not host:
        return False
    for suf in suffixes:
        suf = suf.lower().lstrip(".")
        if host == suf or host.endswith("." + suf):
            return True
    return False


def read_saas_rows(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8")
    segments = _parse_json_segments(text)
    rows = []
    for seg in segments:
        if isinstance(seg, dict) and isinstance(seg.get("rows"), list):
            rows.extend(seg["rows"])
        elif isinstance(seg, list):
            rows.extend(item for item in seg if isinstance(item, dict))
    return rows


def write_template_xlsx(rows: list[dict], output_path: Path) -> None:
    """按爬取模板（商品编码 + url）生成 xlsx。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "商品编码导入模板"

    # 表头
    headers = ["商品编码", "url"]
    ws.append(headers)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据
    for row in rows:
        code = normalize_cell_str(row.get("skuId"))
        url = normalize_cell_str(row.get("extLink"))
        ws.append([code, url])

    # 列宽
    ws.column_dimensions[get_column_letter(1)].width = 28
    ws.column_dimensions[get_column_letter(2)].width = 80
    ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))


def export_excluded_templates(
    input_saas_json: Path,
    output_root: Path,
    date_label: str | None = None,
) -> dict[str, dict]:
    """导出京东/震坤行模板到 <output_root>/<date_label>/<商城>.xlsx。

    供 audit_saas_excel_file.run_main 启动时调用。
    返回 {mall_name: {"matched": N, "deduped": N, "path": Path}}
    """
    if not input_saas_json.exists():
        raise FileNotFoundError(f"saas json 不存在：{input_saas_json}")
    label = date_label or date.today().strftime("%Y-%m-%d")
    out_dir = output_root / label
    out_dir.mkdir(parents=True, exist_ok=True)

    saas_rows = read_saas_rows(input_saas_json)
    summary: dict[str, dict] = {
        "_meta": {
            "date_label": label,
            "input": str(input_saas_json),
            "saas_rows": len(saas_rows),
            "output_dir": str(out_dir),
        }
    }

    summary_lines = [
        f"日期: {label}",
        f"输入文件: {input_saas_json}",
        f"saas 行数: {len(saas_rows)}",
        "",
    ]

    for mapping in EXCLUDED_MALLS:
        if not isinstance(mapping, ZhaoshangMallMapping):
            continue
        mall_name = mapping.mall_name
        domains = mapping.domains
        vendor_keywords = mapping.vendor_keywords
        mall_key = mapping.mall_key
        matched = [row for row in saas_rows if _mall_predicate(row, domains, vendor_keywords)]
        seen: set[tuple[str, str]] = set()
        deduped: list[dict] = []
        for row in matched:
            key = (normalize_cell_str(row.get("skuId")), normalize_cell_str(row.get("extLink")))
            if key in seen:
                continue
            seen.add(key)
            deduped.append(row)

        out_path = out_dir / f"{mall_name}.xlsx"
        write_template_xlsx(deduped, out_path)
        summary[mall_name] = {
            "matched": len(matched),
            "deduped": len(deduped),
            "path": out_path,
            "mall_key": mall_key,
            "source_id": mapping.source_id,
            "domains": list(domains),
        }
        summary_lines.append(f"{mall_name}: 命中 {len(matched)} 行 / 去重后 {len(deduped)} 行")

    summary_path = out_dir / "summary.txt"
    summary_path.write_text("\n".join(summary_lines), encoding="utf-8")
    summary["_meta"]["summary_txt"] = str(summary_path)
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="将京东/震坤行等不参与爬虫的商城导出到模板文件")
    parser.add_argument("--input-saas-json", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path,
                        help="输出根目录（自动追加 YYYY-MM-DD 子目录）")
    parser.add_argument("--date", default=None,
                        help="日期标签（默认今天，格式 YYYY-MM-DD）")
    args = parser.parse_args(argv)

    if not args.input_saas_json.exists():
        print(f"输入文件不存在：{args.input_saas_json}")
        return 1

    try:
        summary = export_excluded_templates(
            input_saas_json=args.input_saas_json,
            output_root=args.output_dir,
            date_label=args.date,
        )
    except Exception as exc:
        print(f"导出失败：{exc}")
        return 1

    meta = summary.get("_meta", {})
    print(f"读取 saas 行数: {meta.get('saas_rows')}")
    for name, info in summary.items():
        if name == "_meta":
            continue
        print(f"{name}: 命中 {info['matched']} 行，去重后 {info['deduped']} 行 → {info['path']}")
    print(f"\n汇总: {meta.get('summary_txt')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
