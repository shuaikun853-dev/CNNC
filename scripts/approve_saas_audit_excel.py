from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass, replace
import hashlib
import json
from pathlib import Path
from typing import Any, Callable, Sequence
import urllib.request
from urllib.parse import urlparse, urlunparse

from openpyxl import load_workbook

try:
    from scripts.auth_manager import AuthManager
    from scripts.curl_utils import (
        apply_curl_headers,
        internal_direct_opener,
        load_curl_command,
    )
    from scripts.reject_saas_audit_excel import fetch_pending_pairs
except ModuleNotFoundError:  # Direct invocation: python scripts/approve_saas_audit_excel.py
    from auth_manager import AuthManager
    from curl_utils import apply_curl_headers, internal_direct_opener, load_curl_command
    from reject_saas_audit_excel import fetch_pending_pairs


TENANT_ID = "188605"
COMMODITY_POOL_ID = "222"
EXCLUDED_FILL_COLORS = {"FFFFC7CE", "FFFFFF00"}
REQUIRED_HEADERS = {
    "SaaS商品ID",
    "SKU编号",
    "互联网标题",
    "模型审核结果",
}


@dataclass(frozen=True)
class ApproveCandidate:
    v_commodity_id: str
    sku_id: str
    model_audit_result: str
    fill_color: str
    excel_row: int
    source_cells: tuple[tuple[int, str], ...] = ()

    def as_payload_item(self) -> dict[str, str]:
        return {
            "vCommodityId": self.v_commodity_id,
            "skuId": self.sku_id,
            "modelAuditResult": self.model_audit_result,
        }

    def as_manifest_item(self) -> dict[str, Any]:
        source_cells = self.source_cells or ((self.excel_row, self.fill_color),)
        return {
            **self.as_payload_item(),
            "sourceCells": [
                {"excelRow": excel_row, "fillColor": fill_color}
                for excel_row, fill_color in source_cells
            ],
        }


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_is_nonempty(value: Any) -> bool:
    if value is None:
        return False
    return not isinstance(value, str) or bool(value.strip())


def load_approve_candidates(
    input_xlsx: Path,
    *,
    require_customer_price: bool = False,
) -> list[ApproveCandidate]:
    workbook = load_workbook(input_xlsx, data_only=True, read_only=False)
    if "record_results" not in workbook.sheetnames:
        raise ValueError("Workbook is missing record_results sheet")
    worksheet = workbook["record_results"]
    header_columns = {
        _cell_text(worksheet.cell(1, column).value): column
        for column in range(1, worksheet.max_column + 1)
    }
    missing_headers = REQUIRED_HEADERS - set(header_columns)
    if require_customer_price and "客户价" not in header_columns:
        missing_headers.add("客户价")
    if missing_headers:
        raise ValueError(f"Workbook is missing required headers: {sorted(missing_headers)}")

    title_column = header_columns["互联网标题"]
    candidates: list[ApproveCandidate] = []
    candidate_indexes: dict[tuple[str, str], int] = {}
    for row in range(2, worksheet.max_row + 1):
        fill_color = _cell_text(worksheet.cell(row, title_column).fill.fgColor.rgb)
        if fill_color in EXCLUDED_FILL_COLORS:
            continue
        if require_customer_price and not _cell_is_nonempty(
            worksheet.cell(row, header_columns["客户价"]).value
        ):
            continue
        v_commodity_id = _cell_text(
            worksheet.cell(row, header_columns["SaaS商品ID"]).value
        )
        sku_id = _cell_text(worksheet.cell(row, header_columns["SKU编号"]).value)
        model_audit_result = _cell_text(
            worksheet.cell(row, header_columns["模型审核结果"]).value
        )
        if not v_commodity_id or not sku_id:
            raise ValueError(f"Row {row} has a blank vCommodityId or skuId")
        if model_audit_result != "1":
            raise ValueError(
                f"Row {row} has unsupported modelAuditResult={model_audit_result!r}"
            )
        business_key = (v_commodity_id, sku_id)
        if business_key in candidate_indexes:
            index = candidate_indexes[business_key]
            existing = candidates[index]
            candidates[index] = replace(
                existing,
                source_cells=existing.source_cells + ((row, fill_color),),
            )
            continue
        candidate_indexes[business_key] = len(candidates)
        candidates.append(
            ApproveCandidate(
                v_commodity_id=v_commodity_id,
                sku_id=sku_id,
                model_audit_result=model_audit_result,
                fill_color=fill_color,
                excel_row=row,
                source_cells=((row, fill_color),),
            )
        )
    if not candidates:
        raise ValueError("Workbook has no approve candidates for the selected filters")
    return candidates


def select_pending_candidates(
    candidates: Sequence[ApproveCandidate],
    pending_pairs: set[tuple[str, str]],
) -> tuple[list[ApproveCandidate], list[ApproveCandidate]]:
    pending: list[ApproveCandidate] = []
    already_processed: list[ApproveCandidate] = []
    for candidate in candidates:
        key = (candidate.v_commodity_id, candidate.sku_id)
        if key in pending_pairs:
            pending.append(candidate)
        else:
            already_processed.append(candidate)
    if not pending:
        raise ValueError("No current pending candidates remain for approval")
    return pending, already_processed


def candidate_manifest_sha256(candidates: Sequence[ApproveCandidate]) -> str:
    encoded = json.dumps(
        [candidate.as_manifest_item() for candidate in candidates],
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def validate_execution_confirmation(
    candidates: Sequence[ApproveCandidate],
    *,
    expected_count: int,
    expected_sha256: str,
) -> None:
    if len(candidates) != expected_count:
        raise ValueError(
            f"Candidate count changed: expected {expected_count}, actual {len(candidates)}"
        )
    actual_sha256 = candidate_manifest_sha256(candidates)
    if actual_sha256.lower() != expected_sha256.strip().lower():
        raise ValueError(
            f"Candidate manifest SHA-256 changed: expected {expected_sha256}, "
            f"actual {actual_sha256}"
        )


def approve_url_from_curl(parsed_curl: dict[str, Any]) -> str:
    parsed_url = urlparse(str(parsed_curl.get("url") or ""))
    if parsed_url.hostname != "saasportal.iio.com":
        raise ValueError("cURL host must be saasportal.iio.com")
    if parsed_url.scheme not in {"http", "https"}:
        raise ValueError("cURL scheme must be HTTP or HTTPS")
    if not str(parsed_curl.get("cookie") or "").strip():
        raise ValueError("cURL must contain a Cookie header")
    if parsed_url.path not in {
        "/improxy/saas/iio2/sameAudit/queryAll",
        "/improxy/saas/iio2/sameAudit/approve",
    }:
        raise ValueError("cURL path must be a sameAudit queryAll or approve endpoint")
    return urlunparse(parsed_url._replace(path="/improxy/saas/iio2/sameAudit/approve"))


def build_approve_payload(candidates: Sequence[ApproveCandidate]) -> dict[str, Any]:
    if not candidates:
        raise ValueError("Approve payload cannot be empty")
    return {
        "tenantId": TENANT_ID,
        "commodityPoolId": COMMODITY_POOL_ID,
        "commodityList": [candidate.as_payload_item() for candidate in candidates],
    }


def _decode_json_response(response: Any) -> dict[str, Any]:
    charset = response.headers.get_content_charset() or "utf-8"
    decoded = json.loads(response.read().decode(charset))
    if not isinstance(decoded, dict):
        raise RuntimeError("SaaS response must be a JSON object")
    return decoded


def execute_approve_batches(
    parsed_curl: dict[str, Any],
    candidates: Sequence[ApproveCandidate],
    *,
    batch_size: int = 50,
    opener: Any | None = None,
    timeout: int = 30,
    on_batch_success: Callable[[dict[str, Any]], None] | None = None,
) -> list[dict[str, Any]]:
    if batch_size <= 0 or batch_size > 100:
        raise ValueError("batch_size must be between 1 and 100")
    approve_url = approve_url_from_curl(parsed_curl)
    direct_opener = opener or internal_direct_opener()
    events: list[dict[str, Any]] = []
    for offset in range(0, len(candidates), batch_size):
        batch = candidates[offset : offset + batch_size]
        request = urllib.request.Request(
            approve_url,
            data=json.dumps(
                build_approve_payload(batch),
                ensure_ascii=False,
                separators=(",", ":"),
            ).encode("utf-8"),
            method="POST",
        )
        apply_curl_headers(request, parsed_curl.get("headers") or [], parsed_curl["cookie"])
        with direct_opener.open(request, timeout=timeout) as response:
            result = _decode_json_response(response)
        rsp_code = str(result.get("rspCode"))
        if rsp_code != "0000":
            raise RuntimeError(
                f"Approve batch {len(events) + 1} failed: rspCode={rsp_code}, "
                f"rspDesc={result.get('rspDesc')}"
            )
        event = {
            "batchIndex": len(events) + 1,
            "batchCount": len(batch),
            "firstExcelRow": batch[0].excel_row,
            "lastExcelRow": batch[-1].excel_row,
            "rspCode": rsp_code,
            "rspDesc": _cell_text(result.get("rspDesc")),
        }
        events.append(event)
        if on_batch_success:
            on_batch_success(event)
    return events


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Preview a current-pending, non-red/non-yellow SaaS approve batch."
    )
    parser.add_argument("--input-xlsx", type=Path, required=True)
    parser.add_argument("--curl-file", type=Path, required=True)
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--expected-count", type=int)
    parser.add_argument("--expected-sha256")
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--batch-size", type=int, default=50)
    parser.add_argument("--require-customer-price", action="store_true")
    return parser


def run_main(
    argv: Sequence[str] | None = None,
    *,
    auth_manager: AuthManager | None = None,
) -> int:
    args = build_parser().parse_args(argv)
    source_candidates = load_approve_candidates(
        args.input_xlsx,
        require_customer_price=args.require_customer_price,
    )
    parsed_curl = load_curl_command(args.curl_file)
    approve_url_from_curl(parsed_curl)
    manager = auth_manager or AuthManager()
    parsed_curl = manager.prepare_curl("saas", parsed_curl, refresh=True)
    approve_url_from_curl(parsed_curl)
    pending_pairs = fetch_pending_pairs(
        parsed_curl,
        include_all_reference_states=True,
        auth_manager=manager,
    )
    candidates, already_processed = select_pending_candidates(
        source_candidates, pending_pairs
    )
    digest = candidate_manifest_sha256(candidates)
    print(f"source_candidate_count={len(source_candidates)}")
    print(f"candidate_count={len(candidates)}")
    print(f"already_processed_count={len(already_processed)}")
    print(f"candidate_manifest_sha256={digest}")
    print(f"require_customer_price={str(args.require_customer_price).lower()}")

    if not args.execute:
        return 0
    if args.expected_count is None or not args.expected_sha256:
        raise ValueError("--execute requires --expected-count and --expected-sha256")
    if args.output_dir is None:
        raise ValueError("--execute requires --output-dir")
    validate_execution_confirmation(
        candidates,
        expected_count=args.expected_count,
        expected_sha256=args.expected_sha256,
    )
    args.output_dir.mkdir(parents=True, exist_ok=False)
    events_path = args.output_dir / "approve_events.jsonl"
    summary_path = args.output_dir / "approve_summary.json"
    events: list[dict[str, Any]] = []

    def record_success(event: dict[str, Any]) -> None:
        events.append(event)
        with events_path.open("a", encoding="utf-8") as output:
            output.write(json.dumps(event, ensure_ascii=False, separators=(",", ":")))
            output.write("\n")

    try:
        manifest = {
            "tenantId": TENANT_ID,
            "commodityPoolId": COMMODITY_POOL_ID,
            "inputXlsx": str(args.input_xlsx.resolve()),
            "sourceCandidateCount": len(source_candidates),
            "candidateCount": len(candidates),
            "alreadyProcessedCount": len(already_processed),
            "requireCustomerPrice": args.require_customer_price,
            "candidateManifestSha256": digest,
            "fillColors": dict(Counter(candidate.fill_color for candidate in candidates)),
            "candidates": [candidate.as_manifest_item() for candidate in candidates],
            "alreadyProcessedCandidates": [
                candidate.as_manifest_item() for candidate in already_processed
            ],
        }
        (args.output_dir / "approve_manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        execute_approve_batches(
            parsed_curl,
            candidates,
            batch_size=args.batch_size,
            on_batch_success=record_success,
        )
        pending_after = fetch_pending_pairs(
            parsed_curl,
            include_all_reference_states=True,
            auth_manager=manager,
        )
        remaining = [
            candidate
            for candidate in candidates
            if (candidate.v_commodity_id, candidate.sku_id) in pending_after
        ]
        if remaining:
            raise RuntimeError(
                f"Postflight found {len(remaining)} candidate(s) still pending"
            )
        summary = {
            "status": "COMPLETED",
            "sourceCandidateCount": len(source_candidates),
            "candidateCount": len(candidates),
            "alreadyProcessedCount": len(already_processed),
            "requireCustomerPrice": args.require_customer_price,
            "approvedCount": sum(event["batchCount"] for event in events),
            "batchCount": len(events),
            "pendingBeforeCount": len(pending_pairs),
            "pendingAfterCount": len(pending_after),
            "remainingPendingCount": 0,
            "candidateManifestSha256": digest,
        }
    except Exception as exc:
        summary = {
            "status": "BLOCKED",
            "sourceCandidateCount": len(source_candidates),
            "candidateCount": len(candidates),
            "alreadyProcessedCount": len(already_processed),
            "requireCustomerPrice": args.require_customer_price,
            "confirmedBatchCount": len(events),
            "confirmedApprovedCount": sum(
                event["batchCount"] for event in events
            ),
            "candidateManifestSha256": digest,
            "errorType": type(exc).__name__,
            "error": str(exc),
        }
        summary_path.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        raise
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"status={summary['status']}")
    print(f"approved_count={summary['approvedCount']}")
    print(f"output_dir={args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(run_main())
